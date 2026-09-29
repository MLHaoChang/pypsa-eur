"""
Untagged rendering: a corporate Word template without tags gets its body
rebuilt from a *mapping plan* — WP10 of
docs/superpowers/plans/2026-09-28-llm-report-generation-increment-1.md
(Increment 2, "Pinned interfaces").

Three entry points:

* ``default_mapping(outline, doc)`` — code only. Each template heading at
  or after ``body_start_index`` is matched to a report section by title
  (numbering and case ignored, a small synonym table, ``difflib`` for near
  matches) and ``kept``; a heading that looks like sample text ("Lorem
  ipsum", "[Chapter title]") is ``dropped``; a heading the report has
  nothing for (an "Introduction") is kept empty and noted. Report sections
  no heading matched are ``inserted`` after the last kept top-level
  heading — except ``skipped`` sections, which go to
  ``unmapped_sections``. A ``not_established`` section is still inserted:
  the omission is the finding. Cover placeholders (client, date, project)
  get a value or a note.
* ``propose_mapping(provider, ...)`` — one generation call through
  ``generator.generate_json`` (the JSON-in-text contract every profile
  follows), then the answer is validated as a ``MappingPlan`` and
  **sanitised** against the outline and the report: unknown heading
  indices, unknown section ids, duplicate coverage, bad insert anchors and
  unknown placeholders are dropped with a note each, and every report
  section ends up covered exactly once or listed. A ``SectionFailure`` is
  returned as such; the caller falls back to ``default_mapping``.
* ``render_untagged(doc, template, plan, *, figure_bytes)`` — keeps every
  body element before ``body_start_index`` untouched (cover page, TOC),
  removes the rest but the final ``sectPr``, and writes the plan: kept /
  renamed headings in the template's own heading style for their level, a
  ``sec:<section_id>`` bookmark per section, the section's blocks through
  the writer's block renderer (bullets in the template's list style,
  tables in its table style), inserted sections at their anchors, the
  "Numbers to check" appendix. Placeholders are replaced in the cover
  and in the header/footer (run-merged). ``w:updateFields`` is set in
  ``settings.xml`` when the template has a TOC so Word refreshes it on
  open.

Positions in a plan are the outline's (``TemplateHeading.index``: the
index of the top-level ``w:body`` child). ``after_heading_index == -1``
means "at the start of the body".
"""
from __future__ import annotations

import difflib
import io
import re
import threading
from collections.abc import Mapping
from typing import Any, Literal

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pydantic import BaseModel, ConfigDict, Field

from models.report import ReportDocument, Section
from services.reports import docx_writer as W
from services.reports import prompts
from services.reports.default_template import ensure_report_styles
from services.reports.docx_reader import (
    HEADER_FOOTER_INDEX,
    TemplateHeading,
    TemplateOutline,
    merged_paragraph_text,
    read_template,
)
from services.reports.generator import SectionFailure, generate_json

MappingAction = Literal["keep", "rename", "drop"]

#: ``after_heading_index`` for an insert at the start of the body.
BODY_START_ANCHOR = -1

#: Words that mark a template heading as sample text, not a chapter.
_SAMPLE_WORDS = ("lorem", "ipsum", "placeholder", "sample", "example",
                 "beispiel", "exemple", "ejemplo", "esempio", "voorbeeld")
_PLACEHOLDER_RE = re.compile(r"\[[^\[\]\n]{1,60}\]|<[^<>\n]{1,60}>|(?<![A-Za-z0-9])[Xx]{3,}(?![A-Za-z0-9])")
_NUMBERING_RE = re.compile(
    r"^\s*(?:(?:\d+(?:\.\d+)*\.?)|(?:[A-Za-z]\.)|(?:[IVXLCivxlc]+\.))(?:\s+|$)")
_PUNCT_RE = re.compile(r"[^\w\s]", re.UNICODE)
_MATCH_THRESHOLD = 0.8

#: Normalised heading words → section id. The English and German words of
#: the corporate fixtures plus the obvious short forms; the default writer's
#: ``SECTION_TITLES`` and the report's own headings are matched too.
_SYNONYMS: dict[str, str] = {
    "summary": "executive_summary",
    "executive summary": "executive_summary",
    "management summary": "executive_summary",
    "zusammenfassung": "executive_summary",
    "résumé": "executive_summary",
    "resumen": "executive_summary",
    "residual failure modes": "fmea_top",
    "failure modes": "fmea_top",
    "fmea": "fmea_top",
    "verbleibende ausfallarten": "fmea_top",
    "ausfallarten": "fmea_top",
    "availability target": "target",
    "target": "target",
    "adequacy": "target",
    "verfügbarkeitsziel": "target",
    "certification": "certification",
    "monte carlo": "certification",
    "zertifizierung": "certification",
    "cost": "cost",
    "cost at target": "cost",
    "kosten": "cost",
    "frontier": "frontier",
    "cost vs availability frontier": "frontier",
    "sizing": "sizing",
    "capacity": "sizing",
    "dimensionierung": "sizing",
    "redundancy": "redundancy",
    "redundanz": "redundancy",
    "levers": "levers",
    "optimisation levers": "levers",
    "optimization levers": "levers",
    "dtc": "dtc",
    "critical load": "dtc",
    "tea": "tea",
    "techno economic summary": "tea",
    "lcoe": "tea",
    "gates": "gates",
    "dynamics gates": "gates",
    "dynamics": "gates",
    "multi energy": "multi_energy",
    "multi energy adequacy": "multi_energy",
}

_CLIENT_WORDS = ("client", "customer", "kunde", "kundenname", "cliente")
_DATE_WORDS = ("date", "datum", "fecha", "data")
_PROJECT_WORDS = ("project", "projekt", "title", "titel", "report", "bericht",
                  "study", "studie", "projet", "proyecto")


# ── the plan ─────────────────────────────────────────────────────────────


class _Model(BaseModel):
    model_config = ConfigDict(extra="ignore")


class MappingEntry(_Model):
    heading_index: int
    action: MappingAction = "keep"
    new_text: str | None = None
    section_ids: list[str] = Field(default_factory=list)


class MappingPlan(_Model):
    entries: list[MappingEntry] = Field(default_factory=list)
    #: ``{"after_heading_index": int, "section_id": str, "heading": str}``
    inserted: list[dict] = Field(default_factory=list)
    #: placeholder text → value ("[Client name]" → "ACME")
    placeholders: dict[str, str] = Field(default_factory=dict)
    unmapped_sections: list[str] = Field(default_factory=list)
    notes: list[str] = Field(default_factory=list)


# ── matching ─────────────────────────────────────────────────────────────


def normalise_heading(text: str) -> str:
    """"1.2 Residual Failure-Modes:" → "residual failure modes"."""
    text = _NUMBERING_RE.sub("", text or "", count=1)
    text = _PUNCT_RE.sub(" ", text.lower())
    return " ".join(text.split())


def _looks_like_sample(text: str) -> bool:
    if _PLACEHOLDER_RE.search(text or ""):
        return True
    words = set(normalise_heading(text).split())
    return any(w in words for w in _SAMPLE_WORDS)


def _candidates(section: Section) -> list[str]:
    """
    The names a heading may match: the section's heading, its id, the
    default writer's title and the synonym-table words for it.
    """
    names = [section.heading, section.section_id.replace("_", " "),
             W.SECTION_TITLES.get(section.section_id, "")]
    names += [k for k, v in _SYNONYMS.items() if v == section.section_id]
    return [normalise_heading(n) for n in names if n]


def _score(heading: str, section: Section) -> float:
    """1.0 exact, 0.8–0.95 containment (shorter names score higher), else difflib."""
    best = 0.0
    for name in _candidates(section):
        if not name:
            continue
        if name == heading:
            return 1.0
        short, long_ = sorted((heading, name), key=len)
        if len(short) >= 3 and f" {short} " in f" {long_} ":
            best = max(best, 0.8 + 0.15 * len(short) / len(long_))
        best = max(best, difflib.SequenceMatcher(None, heading, name).ratio())
    return best


def _body_headings(outline: TemplateOutline) -> list[TemplateHeading]:
    start = outline.body_start_index
    return sorted((h for h in outline.headings
                   if h.level >= 1 and (start is None or h.index >= start)),
                  key=lambda h: h.index)


def _top_level(headings: list[TemplateHeading]) -> int:
    return min((h.level for h in headings), default=1)


def _last_kept_anchor(plan: MappingPlan, headings: list[TemplateHeading]) -> int:
    """The last kept/renamed heading at the body's top level, else -1."""
    level = _top_level(headings)
    by_index = {h.index: h for h in headings}
    anchor = BODY_START_ANCHOR
    for entry in sorted(plan.entries, key=lambda e: e.heading_index):
        h = by_index.get(entry.heading_index)
        if h is not None and entry.action != "drop" and h.level == level:
            anchor = h.index
    return anchor


def _placeholder_values(outline: TemplateOutline, doc: ReportDocument,
                        notes: list[str]) -> dict[str, str]:
    values: dict[str, str] = {}
    seen: list[str] = []
    for p in outline.placeholders:
        if p.text in seen:
            continue
        seen.append(p.text)
        cover = (outline.body_start_index is None
                 or p.paragraph_index < outline.body_start_index
                 or p.paragraph_index == HEADER_FOOTER_INDEX)
        inner = normalise_heading(p.text.strip("[]<>"))
        words = set(inner.split())
        if words & set(_CLIENT_WORDS):
            values[p.text] = "Client"
            notes.append(f"placeholder {p.text!r} filled with 'Client': set the "
                         "client name in the plan")
        elif words & set(_DATE_WORDS):
            values[p.text] = doc.created_at[:10]
        elif words & set(_PROJECT_WORDS):
            values[p.text] = doc.title
        elif cover:
            notes.append(f"placeholder {p.text!r} left as it is")
        else:
            notes.append(f"placeholder {p.text!r} is in the template body: "
                         "removed by the rebuild")
    return values


def default_mapping(outline: TemplateOutline, doc: ReportDocument) -> MappingPlan:
    """The code-only plan: match by title similarity, else append."""
    headings = _body_headings(outline)
    notes: list[str] = []
    entries: list[MappingEntry] = []
    covered: set[str] = set()
    for h in headings:
        if _looks_like_sample(h.text):
            entries.append(MappingEntry(heading_index=h.index, action="drop"))
            notes.append(f"heading {h.index} {h.text!r} dropped: sample text")
            continue
        norm = normalise_heading(h.text)
        scored = [(_score(norm, s), i, s) for i, s in enumerate(doc.sections)
                  if s.section_id not in covered]
        scored.sort(key=lambda t: (-t[0], t[1]))
        if scored and scored[0][0] >= _MATCH_THRESHOLD:
            section = scored[0][2]
            covered.add(section.section_id)
            entries.append(MappingEntry(heading_index=h.index, action="keep",
                                        section_ids=[section.section_id]))
        else:
            entries.append(MappingEntry(heading_index=h.index, action="keep"))
            notes.append(f"heading {h.index} {h.text!r}: template section kept "
                         "without report content")
    plan = MappingPlan(entries=entries, notes=notes)
    _append_uncovered(plan, headings, doc, covered, note=False)
    plan.placeholders = _placeholder_values(outline, doc, plan.notes)
    return plan


def _append_uncovered(plan: MappingPlan, headings: list[TemplateHeading],
                      doc: ReportDocument, covered: set[str], *, note: bool) -> None:
    """Every section not yet covered: inserted after the anchor, or listed."""
    anchor = _last_kept_anchor(plan, headings)
    for section in doc.sections:
        sid = section.section_id
        if sid in covered or sid in plan.unmapped_sections:
            continue
        if section.status == "skipped":
            plan.unmapped_sections.append(sid)
            continue
        plan.inserted.append({"after_heading_index": anchor, "section_id": sid,
                              "heading": section.heading})
        covered.add(sid)
        if note:
            plan.notes.append(f"section {sid!r} was not in the plan: inserted "
                              f"after heading {anchor}")


# ── the model's proposal ─────────────────────────────────────────────────


def sanitise_plan(plan: MappingPlan, outline: TemplateOutline,
                  doc: ReportDocument) -> MappingPlan:
    """
    ``plan`` checked against the outline and the report: bad references are
    dropped with a note each, and every report section is covered exactly
    once or listed in ``unmapped_sections``.
    """
    headings = _body_headings(outline)
    by_index = {h.index: h for h in headings}
    known = {s.section_id: s for s in doc.sections}
    notes = list(plan.notes)
    entries: list[MappingEntry] = []
    covered: set[str] = set()
    seen_indices: set[int] = set()
    for entry in plan.entries:
        if entry.heading_index not in by_index:
            notes.append(f"entry for heading {entry.heading_index} dropped: not a "
                         "body heading of the template")
            continue
        if entry.heading_index in seen_indices:
            notes.append(f"duplicate entry for heading {entry.heading_index} dropped")
            continue
        seen_indices.add(entry.heading_index)
        ids: list[str] = []
        for sid in entry.section_ids:
            if sid not in known:
                notes.append(f"unknown section {sid!r} removed from heading "
                             f"{entry.heading_index}")
            elif sid in covered or sid in ids:
                notes.append(f"section {sid!r} already covered: duplicate on heading "
                             f"{entry.heading_index} removed")
            else:
                ids.append(sid)
        action = entry.action
        new_text = (entry.new_text or "").strip() or None
        if action == "rename" and new_text is None:
            notes.append(f"rename of heading {entry.heading_index} without text: kept")
            action = "keep"
        if action != "rename":
            new_text = None
        if action == "drop" and ids:
            notes.append(f"heading {entry.heading_index} is dropped: its sections "
                         f"{ids} are re-inserted")
            ids = []
        covered.update(ids)
        entries.append(MappingEntry(heading_index=entry.heading_index, action=action,
                                    new_text=new_text, section_ids=ids))
    for h in headings:
        if h.index not in seen_indices:
            entries.append(MappingEntry(heading_index=h.index, action="keep"))
            notes.append(f"heading {h.index} {h.text!r} was not in the plan: kept")
    entries.sort(key=lambda e: e.heading_index)
    kept = {e.heading_index for e in entries if e.action != "drop"}

    inserted: list[dict] = []
    for item in plan.inserted:          # dicts: MappingPlan validated them
        sid = str(item.get("section_id", ""))
        try:
            anchor = int(item.get("after_heading_index", BODY_START_ANCHOR))
        except (TypeError, ValueError):
            anchor = None
        if sid not in known:
            notes.append(f"insert of unknown section {sid!r} dropped")
            continue
        if sid in covered:
            notes.append(f"section {sid!r} already covered: insert dropped")
            continue
        if anchor is None or (anchor != BODY_START_ANCHOR and anchor not in kept):
            notes.append(f"insert of {sid!r} after heading "
                         f"{item.get('after_heading_index')!r} dropped: not a kept "
                         "heading")
            continue
        heading = str(item.get("heading") or "").strip() or known[sid].heading
        inserted.append({"after_heading_index": anchor, "section_id": sid,
                         "heading": heading})
        covered.add(sid)

    placeholder_texts = {p.text for p in outline.placeholders}
    placeholders: dict[str, str] = {}
    for key, value in plan.placeholders.items():
        if key in placeholder_texts:
            placeholders[key] = str(value)
        else:
            notes.append(f"placeholder {key!r} dropped: not in the template")

    unmapped: list[str] = []
    for sid in plan.unmapped_sections:
        if sid not in known:
            notes.append(f"unknown section {sid!r} removed from unmapped_sections")
        elif sid in covered:
            notes.append(f"section {sid!r} is mapped: removed from unmapped_sections")
        elif known[sid].status != "skipped":
            notes.append(f"section {sid!r} is not skipped: it cannot be unmapped, "
                         "inserted instead")
        elif sid not in unmapped:
            unmapped.append(sid)

    out = MappingPlan(entries=entries, inserted=inserted, placeholders=placeholders,
                      unmapped_sections=unmapped, notes=notes)
    _append_uncovered(out, headings, doc, covered, note=True)
    return out


def propose_mapping(provider: Any, *, base_request: Mapping[str, Any],
                    outline: TemplateOutline, doc: ReportDocument, language: str,
                    stop_event: threading.Event | None = None,
                    ) -> MappingPlan | SectionFailure:
    """
    The model's mapping plan, sanitised — or the ``SectionFailure`` saying
    why there is none (the caller falls back to ``default_mapping``).
    """
    result = generate_json(
        provider, base_request=base_request,
        user_message=prompts.mapping_user_message(outline, doc, language=language),
        model_cls=MappingPlan, stop_event=stop_event)
    if isinstance(result, SectionFailure):
        return result
    plan, repairs = result
    if repairs:
        plan.notes.append(f"plan accepted after {repairs} repair turn")
    return sanitise_plan(plan, outline, doc)


# ── rendering ────────────────────────────────────────────────────────────


_TAG_SECTPR = qn("w:sectPr")
_TAG_P = qn("w:p")
_TAG_T = qn("w:t")
_TAG_DEL = qn("w:del")
_TAG_TXBX = qn("w:txbxContent")

#: ``w:updateFields`` must precede these in ``CT_Settings`` (schema order).
_SETTINGS_AFTER_UPDATE_FIELDS = tuple(qn(f"w:{name}") for name in (
    "hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids",
    "mathPr", "attachedSchema", "themeFontLang", "clrSchemeMapping",
    "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade",
    "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary",
    "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator"))


def _body_children(document) -> list:
    return [c for c in document.element.body.iterchildren()
            if isinstance(c.tag, str) and c.tag != _TAG_SECTPR]


def _text_elements(p_el) -> list:
    """The ``w:t`` elements of a paragraph in order, skipping deleted text and text boxes."""
    out: list = []

    def walk(el) -> None:
        for child in el.iterchildren():
            tag = child.tag
            if not isinstance(tag, str) or tag in (_TAG_DEL, _TAG_TXBX):
                continue
            if tag == _TAG_T:
                out.append(child)
            else:
                walk(child)

    walk(p_el)
    return out


def _replace_all(text: str, mapping: dict[str, str]) -> str:
    for key, value in mapping.items():
        if key:
            text = text.replace(key, value)
    return text


def _replace_in_paragraph(p_el, mapping: dict[str, str]) -> bool:
    """
    Replace every placeholder in one paragraph. A placeholder inside one
    run is replaced there (formatting of every run kept); one split across
    runs is replaced run-merged: the result goes into the first ``w:t``,
    the others are emptied.
    """
    merged = merged_paragraph_text(p_el)
    present = [key for key in mapping if key and key in merged]
    if not present:
        return False
    texts = _text_elements(p_el)
    if not texts:
        return False
    in_one_run = {key for key in present
                  if any(t.text and key in t.text for t in texts)}
    if in_one_run == set(present):
        for t in texts:
            if t.text and any(key in t.text for key in present):
                t.text = _replace_all(t.text, mapping)
                t.set(qn("xml:space"), "preserve")
    else:
        texts[0].text = _replace_all(merged, mapping)
        texts[0].set(qn("xml:space"), "preserve")
        for t in texts[1:]:
            t.text = ""
    return True


def _replace_placeholders(document, children_before_body: list,
                          mapping: dict[str, str]) -> None:
    if not mapping:
        return
    roots = list(children_before_body)
    for section in document.sections:
        for part in (section.header, section.footer, section.first_page_header,
                     section.first_page_footer, section.even_page_header,
                     section.even_page_footer):
            if not part.is_linked_to_previous:
                roots.append(part._element)
    for root in roots:
        if root.tag == _TAG_P:
            _replace_in_paragraph(root, mapping)
        else:
            for p_el in root.iter(_TAG_P):
                _replace_in_paragraph(p_el, mapping)


def _set_update_fields(document) -> None:
    settings = document.settings.element
    flag = settings.find(qn("w:updateFields"))
    if flag is None:
        flag = OxmlElement("w:updateFields")
        successor = next((c for c in settings.iterchildren()
                          if c.tag in _SETTINGS_AFTER_UPDATE_FIELDS), None)
        if successor is not None:
            successor.addprevious(flag)
        else:
            settings.append(flag)
    flag.set(qn("w:val"), "true")


def _max_bookmark_id(document) -> int:
    ids = []
    for el in document.element.iter(qn("w:bookmarkStart")):
        try:
            ids.append(int(el.get(qn("w:id"), "-1")))
        except ValueError:
            continue
    return max(ids, default=-1)


class _Styles:
    """The template's style names per heading level, with fallbacks."""

    def __init__(self, document, outline: TemplateOutline) -> None:
        self._doc = document
        self.by_level: dict[int, str] = {}
        for h in outline.headings:
            if h.level >= 1 and h.style and h.level not in self.by_level:
                self.by_level[h.level] = h.style
        self.top = _top_level(_body_headings(outline))
        self.bullet = ("List Bullet" if self._has("List Bullet")
                       else "List Paragraph" if self._has("List Paragraph")
                       else "List Bullet")
        template_tables = [t.style for t in outline.tables if t.style]
        self.table = template_tables[0] if template_tables else "Table Grid"

    def _has(self, name: str) -> bool:
        try:
            self._doc.styles[name]
            return True
        except KeyError:
            return False

    def heading(self, level: int) -> str | None:
        level = max(1, level)
        for name in (self.by_level.get(level), f"Heading {level}"):
            if name and self._has(name):
                return name
        return None


class _Renderer:
    def __init__(self, document, report: ReportDocument, outline: TemplateOutline,
                 figure_bytes: dict[str, bytes]) -> None:
        self.doc = document
        self.report = report
        self.figure_bytes = figure_bytes
        self.styles = _Styles(document, outline)
        self.counters = W._Counters(bullet_style=self.styles.bullet,
                                    table_style=self.styles.table)
        self.bookmarks = W._Bookmarks()
        self.bookmarks._next = _max_bookmark_id(document) + 1
        self.sections = {s.section_id: s for s in report.sections}
        self.written: list[str] = []

    def heading(self, text: str, level: int, section_id: str | None = None):
        style = self.styles.heading(level)
        if style is not None:
            p = W._para(self.doc, text, style)
        else:
            p = self.doc.add_paragraph()
            p.add_run(text).bold = True
        if section_id:
            self.bookmarks.wrap(p, f"{W.BOOKMARK_PREFIX}{section_id}")
        return p

    def section_body(self, section: Section) -> None:
        W._render_section_body(self.doc, section, self.report, self.figure_bytes,
                               self.counters)
        self.written.append(section.section_id)

    def sections_under(self, ids: list[str], level: int, *, first_on_heading: bool) -> None:
        for i, sid in enumerate(ids):
            section = self.sections.get(sid)
            if section is None:
                continue
            if i > 0 or not first_on_heading:
                self.heading(section.heading, level + 1 if first_on_heading else level,
                             sid)
            self.section_body(section)

    def appendix(self) -> None:
        W._appendix(self.doc, self.report, self.bookmarks,
                    heading_style=self.styles.heading(self.styles.top))


def render_untagged(doc: ReportDocument, template: bytes, plan: MappingPlan, *,
                    figure_bytes: dict[str, bytes]) -> bytes:
    """
    ``doc`` written into the untagged ``template`` as ``plan`` says; the
    ``.docx`` bytes. See the module docstring for what is kept and what is
    rebuilt.
    """
    outline = read_template(template)
    document = Document(io.BytesIO(bytes(template)))
    ensure_report_styles(document)
    children = _body_children(document)
    start = outline.body_start_index if outline.body_start_index is not None \
        else len(children)
    kept_children = children[:start]
    for child in children[start:]:
        document.element.body.remove(child)
    _replace_placeholders(document, kept_children, dict(plan.placeholders))

    headings = _body_headings(outline)
    by_index = {h.index: h for h in headings}
    entries = sorted((e for e in plan.entries if e.heading_index in by_index),
                     key=lambda e: e.heading_index)
    # Inserted sections are emitted after their anchor heading's whole
    # subtree (the entries that follow it at a deeper level).
    inserted_at: dict[int, list[dict]] = {}
    for item in plan.inserted:
        anchor = item.get("after_heading_index", BODY_START_ANCHOR)
        try:
            anchor = int(anchor)
        except (TypeError, ValueError):
            anchor = BODY_START_ANCHOR
        if anchor != BODY_START_ANCHOR and anchor not in by_index:
            anchor = BODY_START_ANCHOR
        inserted_at.setdefault(anchor, []).append(item)

    renderer = _Renderer(document, doc, outline, figure_bytes or {})
    top = renderer.styles.top

    def emit_inserted(anchor: int, level: int) -> None:
        for item in inserted_at.pop(anchor, []):
            section = renderer.sections.get(str(item.get("section_id")))
            if section is None:
                continue
            text = str(item.get("heading") or "").strip() or section.heading
            renderer.heading(text, level, section.section_id)
            renderer.section_body(section)

    emit_inserted(BODY_START_ANCHOR, top)
    open_anchors: list[tuple[int, int]] = []   # (heading index, level), nested
    for entry in entries:
        h = by_index[entry.heading_index]
        while open_anchors and open_anchors[-1][1] >= h.level:
            index, level = open_anchors.pop()
            emit_inserted(index, level)
        if entry.action == "drop":
            continue
        text = entry.new_text if entry.action == "rename" and entry.new_text else h.text
        ids = [sid for sid in entry.section_ids if sid in renderer.sections]
        renderer.heading(text, h.level, ids[0] if ids else None)
        renderer.sections_under(ids, h.level, first_on_heading=True)
        open_anchors.append((h.index, h.level))
    while open_anchors:
        index, level = open_anchors.pop()
        emit_inserted(index, level)
    for anchor in list(inserted_at):        # anchors that were dropped
        emit_inserted(anchor, top)

    # Never lose a section: whatever the plan neither wrote nor listed goes
    # at the end under its own heading.
    for section in doc.sections:
        sid = section.section_id
        if sid in renderer.written or sid in plan.unmapped_sections:
            continue
        renderer.heading(section.heading, top, sid)
        renderer.section_body(section)

    renderer.appendix()
    if outline.has_toc:
        _set_update_fields(document)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
