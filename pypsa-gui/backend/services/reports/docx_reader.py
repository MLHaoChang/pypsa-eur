"""
Read the outline of a user's Word template — WP8 of
docs/superpowers/plans/2026-09-28-llm-report-generation-increment-1.md
(Increment 2, "Pinned interfaces").

``read_template(data)`` opens a ``.docx`` from bytes and returns a
``TemplateOutline``: what the template's body contains, in body order, so
that

* the tagged renderer (WP9) knows which ``{{ … }}`` / ``{% … %}`` tags exist
  and where — even when Word split a tag across several runs, which
  ``merged_paragraph_text`` joins back;
* the untagged renderer (WP10) knows the headings (level, style name,
  position), where the body starts (after the cover page and the TOC), the
  ``[Client name]``-style placeholders, the tables with their header rows,
  and the header/footer text it must keep;
* the routes (WP11) can tell the user which content the renderers will not
  touch (``unsupported``: text boxes, SmartArt, content controls, embedded
  objects, other field codes).

**Positions.** Every top-level child of ``w:body`` (a paragraph or a table)
is one index step, in document order, so ``TemplateHeading.index``,
``TemplateTag.paragraph_index``, ``TemplatePlaceholder.paragraph_index`` and
``TemplateTable.index`` all live on the same axis. A tag or placeholder in
a table cell carries the table's index; one in the header or footer carries
``-1`` (there is no body position for it).

**Mode.** ``"tagged"`` iff at least one ``var`` or ``for`` tag exists —
those are the tags that fill something. ``if``/``endif`` alone leave the
template untagged.

**Language.** ``detect_language`` is a small stop-word / heading-word vote
over the headings, placeholders and header/footer text (no library): the
winner needs at least two votes and a strict lead, else ``None``.

Errors: anything python-docx cannot open (not a zip, a zip without the
document part, a truncated or password-protected file) raises
``TemplateReadError`` (a ``ValueError``), never a bare crash.
"""
from __future__ import annotations

import io
import re
from typing import Any, Literal

from docx import Document
from docx.oxml.ns import qn
from pydantic import BaseModel, ConfigDict

TemplateMode = Literal["tagged", "untagged"]
TagKind = Literal["var", "for", "endfor", "if", "endif", "other"]


class TemplateReadError(ValueError):
    """The bytes are not a Word document this reader can open."""


class _Model(BaseModel):
    model_config = ConfigDict(extra="ignore")


class TemplateHeading(_Model):
    index: int
    level: int
    text: str
    style: str
    is_body_start: bool


class TemplateTag(_Model):
    kind: TagKind
    text: str
    paragraph_index: int


class TemplatePlaceholder(_Model):
    text: str
    paragraph_index: int


class TemplateTable(_Model):
    index: int
    n_rows: int
    n_cols: int
    header: list[str]
    style: str | None


class TemplateOutline(_Model):
    mode: TemplateMode
    language: str | None
    headings: list[TemplateHeading]
    tags: list[TemplateTag]
    placeholders: list[TemplatePlaceholder]
    tables: list[TemplateTable]
    header_text: str
    footer_text: str
    body_start_index: int | None
    n_paragraphs: int
    has_toc: bool
    unsupported: list[str]


#: Position recorded for a tag or placeholder found in the header/footer.
HEADER_FOOTER_INDEX = -1

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_DGM = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"

_TAG_P = f"{{{_W}}}p"
_TAG_TBL = f"{{{_W}}}tbl"
_TAG_SECTPR = f"{{{_W}}}sectPr"
_TAG_SDT = f"{{{_W}}}sdt"
_TAG_T = f"{{{_W}}}t"
_TAG_TAB = f"{{{_W}}}tab"
_TAG_BR = f"{{{_W}}}br"
_TAG_CR = f"{{{_W}}}cr"
_TAG_DEL = f"{{{_W}}}del"
_TAG_TXBX = f"{{{_W}}}txbxContent"
_TAG_OBJECT = f"{{{_W}}}object"
_TAG_FLDSIMPLE = f"{{{_W}}}fldSimple"
_TAG_INSTRTEXT = f"{{{_W}}}instrText"
_TAG_DRAWING = f"{{{_W}}}drawing"
_TAG_GRAPHICDATA = f"{{{_A}}}graphicData"
_TAG_MC_FALLBACK = f"{{{_MC}}}Fallback"

#: Subtrees ``merged_paragraph_text`` does not descend into: deleted text
#: (tracked changes), text-box contents (a floating box, not the paragraph)
#: and the ``mc:Fallback`` duplicate of an ``mc:AlternateContent`` choice.
_TEXT_SKIP = frozenset({_TAG_DEL, _TAG_TXBX, _TAG_MC_FALLBACK})

_TAG_RE = re.compile(r"\{\{.*?\}\}|\{%.*?%\}", re.DOTALL)
_BRACKET_RE = re.compile(r"\[([^\[\]\n]{1,60})\]")
_ANGLE_RE = re.compile(r"<([^<>\n]{1,60})>")
_XXX_RE = re.compile(r"(?<![A-Za-z0-9])[Xx]{3,}(?![A-Za-z0-9])")
_LETTER_RE = re.compile(r"[^\W\d_]")
_WORD_RE = re.compile(r"[^\W\d_]+")

#: Field codes the renderers leave in place (page numbers, dates, the
#: TOC's own PAGEREF/HYPERLINK entries). Anything else is a "field code".
_BENIGN_FIELDS = frozenset({
    "TOC", "PAGE", "NUMPAGES", "SECTIONPAGES", "DATE", "TIME", "CREATEDATE",
    "SAVEDATE", "PRINTDATE", "PAGEREF", "HYPERLINK",
})

UNSUPPORTED_TEXT_BOX = "text box"
UNSUPPORTED_SMARTART = "SmartArt"
UNSUPPORTED_CONTENT_CONTROL = "content control"
UNSUPPORTED_EMBEDDED_OBJECT = "embedded object"
UNSUPPORTED_FIELD_CODE = "field code"


# ── paragraph text ───────────────────────────────────────────────────────


def merged_paragraph_text(paragraph: Any) -> str:
    """
    The paragraph's text with every run joined — including runs nested in
    ``w:hyperlink`` / ``w:smartTag`` / ``w:sdtContent`` / ``w:fldSimple``
    wrappers — so a ``{{ tag }}`` Word split across runs reads whole.

    Accepts a python-docx ``Paragraph`` or a raw ``w:p`` element. A tab
    reads as a tab character, a line break as a newline; deleted (tracked)
    text and the contents of text boxes are left out.
    """
    element = getattr(paragraph, "_p", paragraph)
    parts: list[str] = []
    _collect_text(element, parts)
    return "".join(parts)


def _collect_text(element, parts: list[str]) -> None:
    for child in element.iterchildren():
        tag = child.tag
        if not isinstance(tag, str) or tag in _TEXT_SKIP:
            continue
        if tag == _TAG_T:
            parts.append(child.text or "")
        elif tag == _TAG_TAB:
            parts.append("\t")
        elif tag in (_TAG_BR, _TAG_CR):
            parts.append("\n")
        else:
            _collect_text(child, parts)


# ── tags and placeholders ────────────────────────────────────────────────


def _tag_kind(text: str) -> TagKind:
    if text.startswith("{{"):
        return "var"
    inner = text[2:-2].strip("-+ \t\n")
    word = inner.split(None, 1)[0] if inner else ""
    if word in ("for", "endfor", "if", "endif"):
        return word  # type: ignore[return-value]
    return "other"


def _find_tags(text: str, index: int) -> list[TemplateTag]:
    return [TemplateTag(kind=_tag_kind(m.group(0)), text=m.group(0),
                        paragraph_index=index)
            for m in _TAG_RE.finditer(text)]


def _looks_like_xml(inner: str) -> bool:
    """``<w:p>``, ``<br/>``, ``<a href=…>``, ``</x>`` — markup, not a placeholder."""
    if inner[:1] in ("/", "?", "!"):
        return True
    head = inner.split(None, 1)[0]
    if any(ch in head for ch in ":/="):
        return True
    return "=" in inner and '"' in inner


def _find_placeholders(text: str, index: int) -> list[TemplatePlaceholder]:
    plain = _TAG_RE.sub(" ", text)
    found: list[tuple[int, str]] = []
    for m in _BRACKET_RE.finditer(plain):
        if _LETTER_RE.search(m.group(1)):
            found.append((m.start(), m.group(0)))
    for m in _ANGLE_RE.finditer(plain):
        inner = m.group(1)
        if _LETTER_RE.search(inner) and not _looks_like_xml(inner):
            found.append((m.start(), m.group(0)))
    for m in _XXX_RE.finditer(plain):
        found.append((m.start(), m.group(0)))
    found.sort(key=lambda item: item[0])
    return [TemplatePlaceholder(text=t, paragraph_index=index) for _, t in found]


# ── headings ─────────────────────────────────────────────────────────────

_HEADING_STYLE_RE = re.compile(r"^heading\s*(\d+)$", re.IGNORECASE)


def _style_outline_level(style) -> int | None:
    """``w:outlineLvl`` on the style or one of its base styles."""
    seen = 0
    while style is not None and seen < 12:
        p_pr = style.element.find(qn("w:pPr"))
        if p_pr is not None:
            lvl = p_pr.find(qn("w:outlineLvl"))
            if lvl is not None:
                try:
                    return int(lvl.get(qn("w:val")))
                except (TypeError, ValueError):
                    return None
        style = style.base_style
        seen += 1
    return None


def _heading_level(paragraph) -> int | None:
    """Level 0 for Title, n for "Heading n" / outlineLvl n-1, else None."""
    style = paragraph.style
    name = style.name if style is not None else ""
    if name == "Title":
        return 0
    m = _HEADING_STYLE_RE.match(name or "")
    if m:
        return int(m.group(1))
    p_pr = paragraph._p.pPr
    if p_pr is not None:
        lvl = p_pr.find(qn("w:outlineLvl"))
        if lvl is not None:
            try:
                return int(lvl.get(qn("w:val"))) + 1
            except (TypeError, ValueError):
                return None
    style_lvl = _style_outline_level(style) if style is not None else None
    if style_lvl is not None and style_lvl < 9:
        return style_lvl + 1
    return None


# ── fields and unsupported content ───────────────────────────────────────


def _field_instructions(root) -> list[str]:
    instrs: list[str] = []
    for el in root.iter():
        tag = el.tag
        if tag == _TAG_FLDSIMPLE:
            instrs.append(el.get(qn("w:instr")) or "")
        elif tag == _TAG_INSTRTEXT:
            instrs.append(el.text or "")
    return instrs


def _field_word(instr: str) -> str:
    stripped = instr.strip()
    return stripped.split(None, 1)[0].upper() if stripped else ""


def _paragraph_has_toc_field(element) -> bool:
    return any("TOC" == _field_word(i) for i in _field_instructions(element))


def _unsupported(doc) -> list[str]:
    found: list[str] = []
    roots = [doc.element.body]
    for section in doc.sections:
        for part in (section.header, section.footer):
            if not part.is_linked_to_previous:
                roots.append(part._element)
    tags_seen: set[str] = set()
    smartart = False
    for root in roots:
        for el in root.iter():
            tag = el.tag
            if not isinstance(tag, str):
                continue
            tags_seen.add(tag)
            if tag.startswith(f"{{{_DGM}}}"):
                smartart = True
            elif tag == _TAG_GRAPHICDATA and "diagram" in (el.get("uri") or ""):
                smartart = True
    if _TAG_TXBX in tags_seen:
        found.append(UNSUPPORTED_TEXT_BOX)
    if smartart:
        found.append(UNSUPPORTED_SMARTART)
    if any(el.tag == _TAG_SDT for el in doc.element.body.iter()):
        found.append(UNSUPPORTED_CONTENT_CONTROL)
    if _TAG_OBJECT in tags_seen:
        found.append(UNSUPPORTED_EMBEDDED_OBJECT)
    other_fields = [w for root in roots for w in map(_field_word, _field_instructions(root))
                    if w and w not in _BENIGN_FIELDS]
    if other_fields:
        found.append(UNSUPPORTED_FIELD_CODE)
    return found


# ── the walk ─────────────────────────────────────────────────────────────


def _open(data: bytes):
    if not data:
        raise TemplateReadError("empty file: not a Word document")
    try:
        return Document(io.BytesIO(bytes(data)))
    except Exception as exc:  # PackageNotFoundError, BadZipFile, KeyError, XML errors…
        raise TemplateReadError(f"not a Word (.docx) document: {exc}") from exc


def _cell_text(cell) -> str:
    return "\n".join(merged_paragraph_text(p) for p in cell.paragraphs).strip()


def _part_text(part) -> str:
    if part.is_linked_to_previous:
        return ""
    lines = [merged_paragraph_text(p) for p in part.paragraphs]
    for table in part.tables:
        for row in table.rows:
            lines.append("\t".join(_cell_text(c) for c in row.cells))
    return "\n".join(line for line in lines if line.strip()).strip()


def read_template(data: bytes) -> TemplateOutline:
    """Open ``data`` as a ``.docx`` and describe its body, header and footer."""
    doc = _open(data)
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = doc.element.body
    headings: list[TemplateHeading] = []
    tags: list[TemplateTag] = []
    placeholders: list[TemplatePlaceholder] = []
    tables: list[TemplateTable] = []
    toc_index: int | None = None
    n_paragraphs = 0
    children = [c for c in body.iterchildren() if isinstance(c.tag, str)
                and c.tag != _TAG_SECTPR]

    for index, child in enumerate(children):
        if child.tag == _TAG_P:
            n_paragraphs += 1
            paragraph = Paragraph(child, doc._body)
            text = merged_paragraph_text(paragraph)
            level = _heading_level(paragraph)
            if level is not None and text.strip():
                headings.append(TemplateHeading(
                    index=index, level=level, text=text.strip(),
                    style=paragraph.style.name if paragraph.style is not None else "",
                    is_body_start=False))
            tags.extend(_find_tags(text, index))
            placeholders.extend(_find_placeholders(text, index))
            if toc_index is None and _paragraph_has_toc_field(child):
                toc_index = index
        elif child.tag == _TAG_TBL:
            table = Table(child, doc._body)
            rows = table.rows
            n_cols = max((len(r.cells) for r in rows), default=0)
            header = [_cell_text(c) for c in rows[0].cells] if rows else []
            style_name = table.style.name if table.style is not None else None
            tables.append(TemplateTable(index=index, n_rows=len(rows),
                                        n_cols=n_cols, header=header,
                                        style=style_name))
            for row in rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        text = merged_paragraph_text(p)
                        tags.extend(_find_tags(text, index))
                        placeholders.extend(_find_placeholders(text, index))
            if toc_index is None and _paragraph_has_toc_field(child):
                toc_index = index
        else:
            # w:sdt (a content control) or anything else at body level: one
            # index step; its paragraphs are still scanned for tags so a tag
            # inside a control is not silently lost.
            for p_el in child.iter(_TAG_P):
                text = merged_paragraph_text(p_el)
                tags.extend(_find_tags(text, index))
                placeholders.extend(_find_placeholders(text, index))
            if toc_index is None and _paragraph_has_toc_field(child):
                toc_index = index

    header_text = footer_text = ""
    if doc.sections:
        first = doc.sections[0]
        header_text = _part_text(first.header)
        footer_text = _part_text(first.footer)
        for part in (first.header, first.footer):
            if part.is_linked_to_previous:
                continue
            for p in part.paragraphs:
                text = merged_paragraph_text(p)
                tags.extend(_find_tags(text, HEADER_FOOTER_INDEX))
                placeholders.extend(_find_placeholders(text, HEADER_FOOTER_INDEX))
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            text = merged_paragraph_text(p)
                            tags.extend(_find_tags(text, HEADER_FOOTER_INDEX))
                            placeholders.extend(
                                _find_placeholders(text, HEADER_FOOTER_INDEX))

    has_toc = any(_field_word(i) == "TOC"
                  for i in _field_instructions(doc.element.body))
    body_start_index = _body_start(headings, toc_index if has_toc else None)
    for h in headings:
        h.is_body_start = h.index == body_start_index

    mode: TemplateMode = ("tagged" if any(t.kind in ("var", "for") for t in tags)
                          else "untagged")
    outline = TemplateOutline(
        mode=mode, language=None, headings=headings, tags=tags,
        placeholders=placeholders, tables=tables,
        header_text=header_text, footer_text=footer_text,
        body_start_index=body_start_index, n_paragraphs=n_paragraphs,
        has_toc=has_toc, unsupported=_unsupported(doc),
    )
    outline.language = detect_language(outline)
    return outline


def _body_start(headings: list[TemplateHeading], toc_index: int | None) -> int | None:
    level_one = [h for h in headings if h.level == 1]
    if toc_index is not None:
        after = [h for h in level_one if h.index > toc_index]
        if after:
            return after[0].index
    if level_one:
        return level_one[0].index
    if headings:
        return headings[0].index
    return None


# ── language ─────────────────────────────────────────────────────────────

_LANGUAGE_WORDS: dict[str, frozenset[str]] = {
    "en": frozenset("""
        the and of for to in with on by from at this that is are
        summary executive introduction results result conclusion conclusions
        recommendations appendix background scope methodology method overview
        table figure contents report chapter section prepared client date name
        design target availability failure modes analysis confidential page
        reference study assessment
    """.split()),
    "de": frozenset("""
        die der das und für von mit zu im am auf bei nach über des dem den ist
        zusammenfassung einleitung ergebnisse ergebnis schlussfolgerung
        schlussfolgerungen fazit empfehlungen anhang hintergrund umfang methodik
        überblick tabelle abbildung inhalt inhaltsverzeichnis bericht kapitel
        abschnitt erstellt kunde kundenname datum verfügbarkeit
        verfügbarkeitsziel ausfallarten verbleibende analyse vertraulich seite
        referenzdesign studie bewertung
    """.split()),
    "fr": frozenset("""
        le la les et de des du pour à au aux avec sur par dans ce cette est
        résumé synthèse introduction résultats résultat conclusion conclusions
        recommandations annexe contexte périmètre méthodologie aperçu tableau
        figure sommaire rapport chapitre section préparé client date nom
        conception objectif disponibilité défaillance modes analyse
        confidentiel page référence étude évaluation
    """.split()),
    "es": frozenset("""
        el la los las y de del para a en con por sobre este esta es un una
        resumen ejecutivo introducción resultados resultado conclusión
        conclusiones recomendaciones anexo antecedentes alcance metodología
        visión tabla figura índice contenido informe capítulo sección preparado
        cliente fecha nombre diseño objetivo disponibilidad fallo modos análisis
        confidencial página referencia estudio evaluación
    """.split()),
    "it": frozenset("""
        il lo la i gli le e di del della dei per a al con su da in questo
        questa è un una sommario sintesi introduzione risultati risultato
        conclusione conclusioni raccomandazioni appendice contesto ambito
        metodologia panoramica tabella figura indice rapporto relazione capitolo
        sezione preparato cliente data nome progettazione obiettivo
        disponibilità guasto modalità analisi riservato pagina riferimento
        studio valutazione
    """.split()),
    "nl": frozenset("""
        de het een en van voor met op bij naar over in door dit deze is
        samenvatting inleiding resultaten resultaat conclusie conclusies
        aanbevelingen bijlage achtergrond reikwijdte methodologie overzicht
        tabel figuur inhoud inhoudsopgave rapport hoofdstuk sectie opgesteld
        klant datum naam ontwerp doel beschikbaarheid storing faalwijzen
        analyse vertrouwelijk pagina referentie studie beoordeling
    """.split()),
}


def detect_language(outline: TemplateOutline) -> str | None:
    """
    ``"en"`` / ``"de"`` / ``"fr"`` / ``"es"`` / ``"it"`` / ``"nl"`` from a
    word vote over the headings, placeholders and header/footer text;
    ``None`` when no language has at least two votes and a strict lead.
    """
    texts = [h.text for h in outline.headings]
    texts += [p.text for p in outline.placeholders]
    texts += [outline.header_text, outline.footer_text]
    words = [w.lower() for text in texts for w in _WORD_RE.findall(text)]
    if not words:
        return None
    votes = {lang: sum(1 for w in words if w in vocab)
             for lang, vocab in _LANGUAGE_WORDS.items()}
    ranked = sorted(votes.items(), key=lambda item: item[1], reverse=True)
    best, best_votes = ranked[0]
    runner_up = ranked[1][1] if len(ranked) > 1 else 0
    if best_votes >= 2 and best_votes > runner_up:
        return best
    return None
