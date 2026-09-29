"""
The round trip — an edited Word copy of a report read back into
``ReportDocument`` blocks and merged as a new version.

WP12 of docs/superpowers/plans/2026-09-28-llm-report-generation-increment-1.md
("Increment 3 — round trip", pinned interfaces).

``read_edited_docx(data, base)`` is the inverse of ``docx_writer``: it walks
the edited file's body and recovers, per section, the block list the writer
would have needed to produce what it sees. ``merge_round_trip(base, result)``
turns that into the next version of the document. Neither touches the store.

**Reading, in order.**

1. Tracked changes are accepted: every ``w:ins`` / ``w:moveTo`` is unwrapped
   in place, every ``w:del`` / ``w:moveFrom`` removed (with its
   ``w:delText``), a ``w:del`` on a table-row mark removes the row; each
   counts one in ``accepted_tracked_changes``. Marks inside ``w:rPr`` (a
   formatting change, a paragraph-mark insertion) are not content and are
   left alone.
2. The body is split into sections. A paragraph carrying a ``w:bookmarkStart``
   named ``sec:<id>`` (what the writer puts on every heading) opens section
   ``<id>``; a heading paragraph with no such bookmark — the user's copy of
   Word, or a save-as, may have dropped them — opens the base section whose
   heading matches its text (case-insensitive, leading numbering stripped,
   ``difflib`` ratio >= 0.85) or, when nothing matches, an UNMATCHED chunk
   (``section_id=None``, listed in ``unmatched``). A deeper heading inside an
   open section that matches nothing is that section's content (a bold
   paragraph), not a split. Everything before the first section (title,
   generated-from line, the writer's preamble) is ignored; the "Numbers to
   check" appendix (``sec:numbers_to_check`` or its heading text) ends the
   walk.
3. Paragraphs and tables become blocks (see ``_SectionBuilder``): list
   paragraphs → one ``Bullets`` block per run of them; the writer's callout
   renderings → the base ``Callout`` when the text still matches, else a new
   one of the same kind; the writer's ``Note:`` line and the status sentence
   ("Not established: this section was …") → dropped, they are rendered
   from ``section.note`` / ``section.status``, not from blocks; a picture →
   the base section's ``FigureRef`` in order; a table → the base
   ``TableRef`` in order when every cell still reads the same, else a
   ``TableRef`` to ``<id>_edited_v<next>`` with the edited table in
   ``edited_tables``; anything else → a ``Paragraph`` with the inline
   markdown rebuilt from the runs (``**bold**``, ``*italic*``, ``` `code` ``
   from the writer's code font, ``[text](url)`` from a hyperlink or the
   writer's link runs), adjacent runs of equal formatting merged.
4. A section is ``changed`` when its recovered blocks differ from the base
   section's — whitespace-insensitive, and with the base's multi-chunk
   paragraphs compared chunk by chunk, since that is how the writer laid
   them out. An unchanged section hands back the base's own block objects.
5. Comments (python-docx 1.2 ``document.comments``) are attributed by the
   position of their ``w:commentRangeStart`` (or ``w:commentReference``):
   inside a section → that section's ``comments``; before the first section
   or in the appendix → ``comments_global``.

``RoundTripResult.edited_tables`` is an additive extension of the pinned
model: the ``Table`` objects the edited ``TableRef`` blocks point at, keyed
by their new id, so the merge can add them to ``doc.tables`` (the literal
``Table`` is not a block type). A picture the user added has no block to
land in; it is stated as a disclosure callout rather than dropped.

Errors: bytes python-docx cannot open raise ``docx_reader.TemplateReadError``
(a ``ValueError``), the same class the template routes already map to 400.
"""
from __future__ import annotations

import difflib
import io
import re
from datetime import UTC, datetime
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from pydantic import BaseModel, ConfigDict, Field as PField

from models.report import (
    Block,
    Bullets,
    Callout,
    Field,
    FigureRef,
    Paragraph,
    ReportDocument,
    Section,
    SectionAudit,
    Table,
    TableRef,
)
from services.reports.docx_reader import (
    TemplateReadError,
    _heading_level,
    merged_paragraph_text,
)
from services.reports.docx_writer import (
    APPENDIX_HEADING,
    BOOKMARK_PREFIX,
    _CODE_FONT,
    _LINK_COLOR,
)

APPENDIX_ID = "numbers_to_check"
HEADING_MATCH_RATIO = 0.85
REMOVED_NOTE = "removed from the edited copy; kept from version {version}"
ADDED_PICTURE_TEXT = ("A picture added in the edited copy has no place in the "
                      "report model and was not carried over.")

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _w(tag: str) -> str:
    return f"{{{_W}}}{tag}"


_TAG_P, _TAG_TBL, _TAG_SDT, _TAG_SDT_CONTENT = _w("p"), _w("tbl"), _w("sdt"), _w("sdtContent")
_TAG_R, _TAG_HYPERLINK, _TAG_RPR, _TAG_PPR = _w("r"), _w("hyperlink"), _w("rPr"), _w("pPr")
_TAG_TRPR, _TAG_TR = _w("trPr"), _w("tr")
_TAG_INS, _TAG_DEL, _TAG_MOVE_TO, _TAG_MOVE_FROM = (
    _w("ins"), _w("del"), _w("moveTo"), _w("moveFrom"))
_TAG_T, _TAG_TAB, _TAG_BR, _TAG_CR = _w("t"), _w("tab"), _w("br"), _w("cr")
_TAG_DRAWING, _TAG_PICT = _w("drawing"), _w("pict")
_TAG_BOOKMARK_START = _w("bookmarkStart")
_TAG_COMMENT_START, _TAG_COMMENT_REF = _w("commentRangeStart"), _w("commentReference")
_TAG_SECTPR = _w("sectPr")

#: A tracked-change mark whose parent is one of these is a property change
#: or a paragraph mark, not content: left in place.
_PROPERTY_PARENTS = frozenset({_TAG_RPR, _TAG_PPR})

_NUMBERING_RE = re.compile(
    r"^\s*(?:(?:\d+|[A-Za-z]|[ivxlcdm]+)(?:[.)]|(?:\.\d+)+\.?)?\s+)+", re.IGNORECASE)
_WS_RE = re.compile(r"\s+")
_LINK_URL_RUN_RE = re.compile(r"^ \((?P<url>\S+)\)$")
_MISSING_TABLE_RE = re.compile(r"^Table '(?P<id>[^']+)' is not available in this version")
_MISSING_FIGURE_RE = re.compile(r"^Figure '(?P<id>[^']+)' was not produced for this version")
GAP_PREFIX = "Gap:"
NOTE_PREFIX = "Note:"
NOT_ESTABLISHED_PREFIX = "Not established:"
_STATUS_SENTENCE_PREFIX = "Not established: this section was "


# ── pinned models ────────────────────────────────────────────────────────


class _Model(BaseModel):
    model_config = ConfigDict(extra="ignore")


class RoundTripSection(_Model):
    section_id: str | None
    heading: str
    blocks: list[Block] = PField(default_factory=list)
    changed: bool
    comments: list[str] = PField(default_factory=list)


class RoundTripResult(_Model):
    sections: list[RoundTripSection] = PField(default_factory=list)
    unmatched: list[str] = PField(default_factory=list)
    comments_global: list[str] = PField(default_factory=list)
    accepted_tracked_changes: int = 0
    # Additive (see the module docstring): the tables the edited `TableRef`
    # blocks point at, keyed by their new id.
    edited_tables: dict[str, Table] = PField(default_factory=dict)


# ── opening and tracked changes ──────────────────────────────────────────


def _open(data: bytes):
    if not data:
        raise TemplateReadError("empty file: not a Word document")
    try:
        return Document(io.BytesIO(bytes(data)))
    except Exception as exc:  # PackageNotFoundError, BadZipFile, KeyError, XML errors…
        raise TemplateReadError(f"not a Word (.docx) document: {exc}") from exc


def _unwrap(element) -> None:
    parent = element.getparent()
    index = parent.index(element)
    for child in list(element):
        parent.insert(index, child)
        index += 1
    parent.remove(element)


def accept_tracked_changes(body) -> int:
    """
    Accept every tracked change under ``body`` in place; returns how many
    were accepted.
    """
    count = 0
    for tag in (_TAG_INS, _TAG_MOVE_TO):
        for el in list(body.iter(tag)):
            parent = el.getparent()
            if parent is None or parent.tag in _PROPERTY_PARENTS:
                continue
            _unwrap(el)
            count += 1
    for tag in (_TAG_DEL, _TAG_MOVE_FROM):
        for el in list(body.iter(tag)):
            parent = el.getparent()
            if parent is None:
                continue
            if parent.tag == _TAG_TRPR:
                row = parent.getparent()
                if row is not None and row.tag == _TAG_TR and row.getparent() is not None:
                    row.getparent().remove(row)
                    count += 1
                continue
            if parent.tag in _PROPERTY_PARENTS:
                continue
            parent.remove(el)
            count += 1
    return count


# ── inline markdown from runs ────────────────────────────────────────────


def _run_text(run_el) -> str:
    parts: list[str] = []
    for child in run_el.iter():
        tag = child.tag
        if tag == _TAG_T:
            parts.append(child.text or "")
        elif tag in (_TAG_TAB, _TAG_BR, _TAG_CR):
            parts.append(" ")
    return "".join(parts)


def _run_format(run_el) -> tuple[bool, bool, bool, bool]:
    """``(bold, italic, code, writer_link)`` from the run's own properties."""
    r_pr = run_el.find(_TAG_RPR)
    if r_pr is None:
        return (False, False, False, False)

    def on(name: str) -> bool:
        el = r_pr.find(_w(name))
        if el is None:
            return False
        val = el.get(qn("w:val"))
        return val is None or val.lower() in ("1", "true", "on")

    fonts = r_pr.find(_w("rFonts"))
    code = fonts is not None and _CODE_FONT in {
        fonts.get(qn("w:ascii")), fonts.get(qn("w:hAnsi"))}
    color = r_pr.find(_w("color"))
    underline = r_pr.find(_w("u"))
    underlined = underline is not None and (underline.get(qn("w:val")) or "single") != "none"
    link = (underlined and color is not None
            and (color.get(qn("w:val")) or "").upper() == str(_LINK_COLOR))
    return (on("b"), on("i"), code, link)


def _color_of(run_el) -> str:
    r_pr = run_el.find(_TAG_RPR)
    color = r_pr.find(_w("color")) if r_pr is not None else None
    return (color.get(qn("w:val")) or "").upper() if color is not None else ""


def _hyperlink_url(hyperlink_el, part) -> str | None:
    rid = hyperlink_el.get(qn("r:id"))
    if rid:
        try:
            return part.rels[rid].target_ref
        except (KeyError, AttributeError):
            return None
    anchor = hyperlink_el.get(qn("w:anchor"))
    return f"#{anchor}" if anchor else None


def _spans(p_el, part) -> list[tuple[str, str, str | None]]:
    """``[(kind, text, url)]`` in run order, adjacent equal kinds merged."""
    raw: list[tuple[str, str, str | None]] = []

    def visit(el) -> None:
        for child in el.iterchildren():
            tag = child.tag
            if not isinstance(tag, str):
                continue
            if tag == _TAG_R:
                text = _run_text(child)
                if not text:
                    continue
                bold, italic, code, link = _run_format(child)
                kind = ("code" if code else "link" if link else "bold" if bold
                        else "italic" if italic else "text")
                raw.append((kind, text, None))
            elif tag == _TAG_HYPERLINK:
                text = "".join(_run_text(r) for r in child.iter(_TAG_R))
                if text:
                    raw.append(("link", text, _hyperlink_url(child, part)))
            elif tag in (_TAG_PPR, _TAG_RPR):
                continue
            else:  # smartTag, sdt/sdtContent, fldSimple, bookmarks…
                visit(child)

    visit(p_el)
    # The writer's own link rendering: an underlined coloured run, then
    # " (url)" in the same colour when the url differs from the text.
    spans: list[tuple[str, str, str | None]] = []
    i = 0
    while i < len(raw):
        kind, text, url = raw[i]
        if kind == "link" and url is None:
            url = text
            if i + 1 < len(raw):
                m = _LINK_URL_RUN_RE.match(raw[i + 1][1])
                if m:
                    url = m.group("url")
                    i += 1
        if spans and spans[-1][0] == kind and kind != "link":
            spans[-1] = (kind, spans[-1][1] + text, None)
        else:
            spans.append((kind, text, url))
        i += 1
    return spans


def _md_span(kind: str, text: str, url: str | None) -> str:
    if kind == "text":
        return text
    lead = text[:len(text) - len(text.lstrip())]
    trail = text[len(text.rstrip()):]
    core = text.strip()
    if not core:
        return text
    if kind == "bold":
        inner = f"**{core}**"
    elif kind == "italic":
        inner = f"*{core}*"
    elif kind == "code":
        inner = f"`{core}`"
    else:
        inner = f"[{core}]({url or core})"
    return f"{lead}{inner}{trail}"


def paragraph_markdown(p_el, part) -> str:
    """The paragraph's text with the writer's inline markdown restored."""
    return "".join(_md_span(*s) for s in _spans(p_el, part)).strip()


# ── helpers ──────────────────────────────────────────────────────────────


def _norm(text: str | None) -> str:
    return _WS_RE.sub(" ", (text or "")).strip()


def _heading_key(text: str) -> str:
    return _norm(_NUMBERING_RE.sub("", text)).lower()


def _match_heading(text: str, candidates: dict[str, str]) -> str | None:
    """The section id whose heading best matches ``text``, or None."""
    key = _heading_key(text)
    if not key:
        return None
    best_id, best = None, 0.0
    for section_id, heading in candidates.items():
        other = _heading_key(heading)
        if not other:
            continue
        if other == key:
            return section_id
        ratio = difflib.SequenceMatcher(None, key, other).ratio()
        if ratio > best:
            best_id, best = section_id, ratio
    return best_id if best >= HEADING_MATCH_RATIO else None


def _bookmark_id(p_el) -> str | None:
    for el in p_el.iter(_TAG_BOOKMARK_START):
        name = el.get(qn("w:name")) or ""
        if name.startswith(BOOKMARK_PREFIX):
            return name[len(BOOKMARK_PREFIX):]
    return None


def _style_name(p_el, doc) -> str:
    p_pr = p_el.find(_TAG_PPR)
    style_el = p_pr.find(_w("pStyle")) if p_pr is not None else None
    if style_el is None:
        return ""
    style_id = style_el.get(qn("w:val")) or ""
    try:
        return doc.styles.get_by_id(style_id, 1).name or ""  # WD_STYLE_TYPE.PARAGRAPH
    except (KeyError, AttributeError):
        return style_id


def _is_list_paragraph(p_el, style: str) -> bool:
    if style.startswith("List"):
        return True
    p_pr = p_el.find(_TAG_PPR)
    return p_pr is not None and p_pr.find(_w("numPr")) is not None


def _has_picture(p_el) -> bool:
    return any(el.tag in (_TAG_DRAWING, _TAG_PICT) for el in p_el.iter())


def _callout_text(block: Callout) -> str:
    """The text the writer renders for ``block``, without its style prefix."""
    text = block.text.strip()
    if block.kind == "not_established" and text and text[-1] not in ".!?":
        text += "."
    return text


def _expand(blocks: list) -> list:
    """The base blocks the way the writer lays them out, for comparison."""
    out: list = []
    for block in blocks:
        if isinstance(block, Paragraph):
            out.extend(Paragraph(md=c.replace("\n", " ").strip())
                       for c in block.md.split("\n\n") if c.strip())
        elif isinstance(block, Bullets):
            if block.items:
                out.append(block)
        else:
            out.append(block)
    return out


def _block_key(block) -> tuple:
    if isinstance(block, Paragraph):
        return ("paragraph", _norm(block.md))
    if isinstance(block, Bullets):
        return ("bullets", tuple(_norm(i) for i in block.items))
    if isinstance(block, TableRef):
        return ("table_ref", block.table_id, _norm(block.caption))
    if isinstance(block, FigureRef):
        return ("figure_ref", block.figure_id, _norm(block.caption))
    if isinstance(block, Callout):
        return ("callout", block.kind, _norm(block.text).rstrip("."))
    if isinstance(block, Field):
        return ("field", _norm(block.key), _norm(block.value))
    return (getattr(block, "type", "?"), repr(block))


def _same_blocks(a: list, b: list) -> bool:
    return [_block_key(x) for x in a] == [_block_key(x) for x in b]


def _table_cells(table_el, doc) -> list[list[str]]:
    from docx.table import Table as DocxTable

    table = DocxTable(table_el, doc._body)
    rows: list[list[str]] = []
    for row in table.rows:
        cells: list[str] = []
        seen: set[int] = set()
        for cell in row.cells:
            if id(cell._tc) in seen:  # a merged cell repeats its element
                continue
            seen.add(id(cell._tc))
            cells.append(_norm("\n".join(merged_paragraph_text(p) for p in cell.paragraphs)))
        rows.append(cells)
    return rows


# ── per-section block recovery ───────────────────────────────────────────


class _SectionBuilder:
    """Collects the blocks of one section as the walk hands it body children."""

    def __init__(self, section_id: str | None, heading: str, base: Section | None,
                 document: ReportDocument, next_version: int) -> None:
        self.section_id = section_id
        self.heading = heading
        self.base = base
        self.document = document
        self.next_version = next_version
        self.blocks: list = []
        self.comments: list[str] = []
        self.heading_level = 1
        base_blocks = base.blocks if base is not None else []
        self._callouts = [b for b in base_blocks if isinstance(b, Callout)]
        self._fields = [b for b in base_blocks if isinstance(b, Field)]
        self._figure_refs = [b for b in base_blocks if isinstance(b, FigureRef)]
        self._table_refs = [b for b in base_blocks if isinstance(b, TableRef)
                            and b.table_id in document.tables]
        self._missing_table_refs = [b for b in base_blocks if isinstance(b, TableRef)
                                    and b.table_id not in document.tables]
        self._new_tables = 0

    # -- paragraphs --------------------------------------------------------

    def add_paragraph(self, p_el, doc, *, style: str, level: int | None) -> None:
        part = doc.part
        if _has_picture(p_el):
            self._add_picture()
            return
        text = _norm(merged_paragraph_text(p_el))
        if not text:
            return
        if style == "Report Caption":
            return  # captions belong to the table / figure block before them
        if level is not None and level >= 1:
            self._add_block(Paragraph(md=f"**{text}**"))
            return
        if _is_list_paragraph(p_el, style):
            item = paragraph_markdown(p_el, part)
            if item.startswith("• "):
                item = item[2:].strip()
            if self.blocks and isinstance(self.blocks[-1], Bullets):
                self.blocks[-1] = Bullets(items=[*self.blocks[-1].items, item])
            else:
                self.blocks.append(Bullets(items=[item]))
            return
        if style == "Disclosure":
            self._add_disclosure(text)
            return
        if text.startswith(NOT_ESTABLISHED_PREFIX):
            body = text[len(NOT_ESTABLISHED_PREFIX):].strip()
            kept = self._take_callout("not_established", body)
            if kept is not None:
                self._add_block(kept)
            elif not text.startswith(_STATUS_SENTENCE_PREFIX):
                self._add_block(Callout(kind="not_established", text=body))
            return  # the status sentence is rendered from `status`, not a block
        field = self._as_field(p_el, part)
        if field is not None:
            self._add_block(field)
            return
        self._add_block(Paragraph(md=paragraph_markdown(p_el, part)))

    def _add_disclosure(self, text: str) -> None:
        m = _MISSING_TABLE_RE.match(text)
        if m:
            ref = next((r for r in self._missing_table_refs if r.table_id == m.group("id")),
                       None)
            if ref is not None:
                self._missing_table_refs.remove(ref)
            self._add_block(ref or TableRef(table_id=m.group("id")))
            return
        m = _MISSING_FIGURE_RE.match(text)
        if m:
            ref = next((r for r in self._figure_refs if r.figure_id == m.group("id")), None)
            if ref is not None:
                self._figure_refs.remove(ref)
            self._add_block(ref or FigureRef(figure_id=m.group("id")))
            return
        if text.startswith(NOTE_PREFIX):
            if self.base is not None and _norm(self.base.note) == _norm(
                    text[len(NOTE_PREFIX):]):
                return  # rendered from `section.note`
            kept = self._take_callout("disclosure", text)
            self._add_block(kept or Callout(kind="disclosure", text=text))
            return
        if text.startswith(GAP_PREFIX):
            body = text[len(GAP_PREFIX):].strip()
            kept = self._take_callout("gap", body)
            self._add_block(kept or Callout(kind="gap", text=body))
            return
        kept = self._take_callout("disclosure", text)
        self._add_block(kept or Callout(kind="disclosure", text=text))

    def _take_callout(self, kind: str, text: str) -> Callout | None:
        key = _norm(text).rstrip(".")
        for callout in self._callouts:
            if callout.kind == kind and _norm(_callout_text(callout)).rstrip(".") == key:
                self._callouts.remove(callout)
                return callout
        return None

    def _as_field(self, p_el, part) -> Field | None:
        spans = _spans(p_el, part)
        if len(spans) < 2 or spans[0][0] != "bold":
            return None
        label = spans[0][1].strip()
        if not label.endswith(":"):
            return None
        key = label[:-1].strip()
        base_field = next((f for f in self._fields if f.key.strip().lower() == key.lower()),
                          None)
        if base_field is None:
            return None
        self._fields.remove(base_field)
        value = "".join(_md_span(*s) for s in spans[1:]).strip()
        if _norm(value) == _norm(base_field.value):
            return base_field
        return Field(key=base_field.key, value=value)

    def _add_picture(self) -> None:
        if self._figure_refs:
            self._add_block(self._figure_refs.pop(0))
        else:
            self._add_block(Callout(kind="disclosure", text=ADDED_PICTURE_TEXT))

    # -- tables ------------------------------------------------------------

    def add_table(self, table_el, doc, edited_tables: dict[str, Table]) -> None:
        cells = _table_cells(table_el, doc)
        ref = self._table_refs.pop(0) if self._table_refs else None
        original = self.document.tables.get(ref.table_id) if ref is not None else None
        if original is not None:
            expected = [[_norm(c) for c in original.columns]] + [
                [_norm(c) for c in row] for row in original.rows]
            if cells == expected:
                self._add_block(ref)
                return
        if ref is not None:
            table_id = f"{ref.table_id}_edited_v{self.next_version}"
            caption = ref.caption or (original.caption if original is not None else None)
        else:
            self._new_tables += 1
            table_id = f"{self.section_id or 'unmatched'}_table{self._new_tables}_v{self.next_version}"
            caption = None
        columns = cells[0] if cells else []
        table = Table(table_id=table_id, columns=columns, rows=cells[1:],
                      caption=caption, source_path=None)
        edited_tables[table_id] = table
        self._add_block(TableRef(table_id=table_id, caption=caption))

    # -- result ------------------------------------------------------------

    def _add_block(self, block) -> None:
        self.blocks.append(block)

    def finish(self) -> RoundTripSection:
        changed = True
        blocks = self.blocks
        if self.base is not None and _same_blocks(blocks, _expand(self.base.blocks)):
            changed = False
            blocks = list(self.base.blocks)
        elif self.base is None:
            changed = bool(blocks)
        return RoundTripSection(section_id=self.section_id, heading=self.heading,
                                blocks=blocks, changed=changed, comments=self.comments)


# ── the walk ─────────────────────────────────────────────────────────────


def _body_children(body) -> list:
    """Top-level paragraphs and tables, content controls flattened."""
    out: list = []
    for child in body.iterchildren():
        tag = child.tag
        if not isinstance(tag, str) or tag == _TAG_SECTPR:
            continue
        if tag in (_TAG_P, _TAG_TBL):
            out.append(child)
        elif tag == _TAG_SDT:
            content = child.find(_TAG_SDT_CONTENT)
            if content is not None:
                out.extend(c for c in content.iterchildren() if c.tag in (_TAG_P, _TAG_TBL))
    return out


def _comment_positions(children: list) -> dict[str, int]:
    """Comment id → index of the body child holding its anchor."""
    positions: dict[str, int] = {}
    for index, child in enumerate(children):
        for el in child.iter(_TAG_COMMENT_START, _TAG_COMMENT_REF):
            cid = el.get(qn("w:id"))
            if cid is not None and cid not in positions:
                positions[cid] = index
    return positions


def read_edited_docx(data: bytes, base: ReportDocument) -> RoundTripResult:
    """Read the edited copy ``data`` back against ``base`` (see the module docstring)."""
    doc = _open(data)
    body = doc.element.body
    accepted = accept_tracked_changes(body)
    children = _body_children(body)
    comment_positions = _comment_positions(children)
    next_version = base.version + 1

    by_id = {s.section_id: s for s in base.sections}
    candidates = {s.section_id: s.heading for s in base.sections}
    opened: set[str] = set()
    builders: list[_SectionBuilder] = []
    owner: list[_SectionBuilder | None] = []  # per body child, for comments
    current: _SectionBuilder | None = None
    edited_tables: dict[str, Table] = {}

    def open_section(section_id: str | None, heading: str, level: int) -> None:
        nonlocal current
        current = _SectionBuilder(section_id, heading, by_id.get(section_id) if section_id
                                  else None, base, next_version)
        current.heading_level = level
        builders.append(current)
        if section_id is not None:
            opened.add(section_id)

    stop = False
    for child in children:
        if stop:
            owner.append(None)
            continue
        if child.tag == _TAG_TBL:
            owner.append(current)
            if current is not None:
                current.add_table(child, doc, edited_tables)
            continue
        from docx.text.paragraph import Paragraph as DocxParagraph

        paragraph = DocxParagraph(child, doc._body)
        style = _style_name(child, doc)
        level = _heading_level(paragraph)
        text = _norm(merged_paragraph_text(child))
        bookmark = _bookmark_id(child)
        if bookmark == APPENDIX_ID or (
                level is not None and level >= 1 and text.lower() == APPENDIX_HEADING.lower()):
            stop = True
            owner.append(None)
            continue
        if bookmark is not None and bookmark in by_id and bookmark not in opened:
            open_section(bookmark, text or by_id[bookmark].heading, level or 1)
            owner.append(current)
            continue
        if level is not None and level >= 1 and text:
            matched = _match_heading(text, {k: v for k, v in candidates.items()
                                            if k not in opened})
            if matched is not None:
                open_section(matched, text, level)
                owner.append(current)
                continue
            if current is None or level <= current.heading_level:
                open_section(None, text, level)
                owner.append(current)
                continue
            # A deeper heading inside an open section: its content.
        owner.append(current)
        if current is not None:
            current.add_paragraph(child, doc, style=style, level=level)

    comments_global: list[str] = []
    for comment in doc.comments:
        text = _norm(comment.text)
        if not text:
            continue
        index = comment_positions.get(str(comment.comment_id))
        target = owner[index] if index is not None and index < len(owner) else None
        if target is None:
            comments_global.append(text)
        else:
            target.comments.append(text)

    sections = [b.finish() for b in builders]
    unmatched = [s.heading or _first_text(s)[:80] for s in sections
                 if s.section_id is None and (s.heading or s.blocks)]
    return RoundTripResult(sections=sections, unmatched=unmatched,
                           comments_global=comments_global,
                           accepted_tracked_changes=accepted,
                           edited_tables=edited_tables)


def _first_text(section: RoundTripSection) -> str:
    for block in section.blocks:
        if isinstance(block, Paragraph):
            return block.md
        if isinstance(block, Bullets) and block.items:
            return block.items[0]
        if isinstance(block, Callout):
            return block.text
    return ""


# ── the merge ────────────────────────────────────────────────────────────


def merge_round_trip(base: ReportDocument, result: RoundTripResult) -> ReportDocument:
    """
    The next version of ``base`` with the edited copy folded in: changed
    sections replaced (``source="user_edit"``, audit reset), comments
    carried as ``pending_instruction``, absent sections kept with a note,
    edited tables added. ``base`` is not mutated.
    """
    by_id: dict[str, RoundTripSection] = {
        s.section_id: s for s in result.sections if s.section_id is not None}
    sections: list[Section] = []
    for section in base.sections:
        edited = by_id.get(section.section_id)
        if edited is None:
            note = REMOVED_NOTE.format(version=base.version)
            if section.note and note not in section.note:
                note = f"{section.note}; {note}"
            sections.append(section.model_copy(update={"note": note}))
            continue
        update: dict[str, Any] = {}
        if edited.changed:
            update.update(blocks=list(edited.blocks), source="user_edit",
                          audit=SectionAudit())
        if edited.comments:
            update["comments"] = list(edited.comments)
            update["pending_instruction"] = "; ".join(edited.comments)
        sections.append(section.model_copy(update=update) if update else section)
    tables = dict(base.tables)
    tables.update(result.edited_tables)
    return base.model_copy(update={
        "version": base.version + 1,
        "created_at": datetime.now(UTC).isoformat(),
        "sections": sections,
        "tables": tables,
    })
