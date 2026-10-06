"""
Tagged rendering: Jinja2 tags inside a user's Word template, filled from a
``ReportDocument``.

WP9 of docs/superpowers/plans/2026-09-28-llm-report-generation-increment-1.md.

The template is an ordinary ``.docx`` whose text carries ``{{ … }}`` and
``{% … %}`` tags. Word splits a tag across runs whenever the author paused
while typing or the spell-checker touched it, so every paragraph is read
*run-merged*: the runs' texts are joined, the joined text is rendered, and
the result is written back into the FIRST run (its formatting wins); the
other runs are dropped. ``docxtpl`` would do the same; it is not a
dependency of this backend, and the subset needed here is small:

* a paragraph that is ONLY one ``{{ tag }}`` renders as content — a
  ``{{ figures.x }}`` becomes an inline picture, a ``{{ fields.x.text }}``
  becomes one paragraph per blank-line-separated chunk (each in the
  original paragraph's style) with the markdown inline subset as runs
  (the writer's tokenizer);
* a paragraph with a tag *inside* other text renders as plain text (a
  figure there is its caption);
* a table row whose first cell starts with ``{% for v in expr %}`` is the
  template for one row per item; ``{% endfor %}`` closes it at the end of
  the row's last cell or in a following row (removed);
* ``{% if … %}…{% endif %}`` on a paragraph hides the paragraph when it
  renders to nothing.

Rendering never crashes on a bad template: an unknown name, a missing id
or a syntax error is a ``TaggedRenderError`` whose message names the
paragraph, the known top-level names and, for ``fields.``/``tables.``/
``figures.`` lookups, the known ids. ``settings.xml`` and the section
properties are left untouched; bookmarks are WP10's job.

Context (every key listed so a template author can rely on it):

``meta``        title, version, evidence_hash, profile_id, model, created_at,
                mode, language
``fields``      section_id → {heading, status, paragraphs, bullets, text,
                note, <Field key>: value …}
``headline``    the executive summary's Field blocks, key → value
``tables``      table_id → {columns, rows, caption}
``figures``     figure_id → picture sentinel (caption as text)
``disclosures`` every ``Callout(kind="disclosure")`` text, in order
``gaps``        every ``Callout(kind="gap")`` text, in order
``unverified``  section_id → the audit's unverified numbers
"""
from __future__ import annotations

import copy
import io
import re
from collections.abc import Iterable, Iterator
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.table import Table as DocxTable, _Cell, _Row
from docx.text.paragraph import Paragraph as DocxParagraph
from docx.text.run import Run
from jinja2 import Environment, StrictUndefined, TemplateError, Undefined
from jinja2.exceptions import TemplateSyntaxError

from models.report import (
    Bullets,
    Callout,
    Field,
    Paragraph,
    ReportDocument,
)
from services.reports.docx_writer import _CODE_FONT, _LINK_COLOR, tokenize_inline

EXECUTIVE_SUMMARY_ID = "executive_summary"
PICTURE_WIDTH = Inches(6.0)
CONTEXT_NAMES = ("meta", "fields", "headline", "tables", "figures",
                 "disclosures", "gaps", "unverified")

_TAG_MARKERS = ("{{", "{%")
_ONLY_TAG = re.compile(r"^\s*\{\{\s*(?P<expr>.+?)\s*\}\}\s*$", re.DOTALL)
_FOR_ROW = re.compile(r"^\s*\{%-?\s*for\s+(?P<var>\w+)\s+in\s+(?P<expr>.+?)\s*-?%\}",
                      re.DOTALL)
_ENDFOR = re.compile(r"\{%-?\s*endfor\s*-?%\}\s*$")
_ENDFOR_ONLY = re.compile(r"^\s*\{%-?\s*endfor\s*-?%\}\s*$")
_TAG_SPAN = re.compile(r"(\{\{.*?\}\}|\{%.*?%\})", re.DOTALL)
_SMART_QUOTES = str.maketrans({"“": '"', "”": '"', "„": '"', "‘": "'", "’": "'"})


class TaggedRenderError(ValueError):
    """A template tag the renderer cannot fill: unknown name, bad syntax."""


# ── the sentinel a figure tag evaluates to ────────────────────────────────


class _FigureTag:
    """``{{ figures.x }}``: a picture when alone in a paragraph, else text."""

    def __init__(self, figure_id: str, caption: str | None, png: bytes | None):
        self.figure_id = figure_id
        self.caption = caption
        self.png = png

    def __str__(self) -> str:
        return self.caption or self.figure_id

    def missing_text(self) -> str:
        return (f"Figure {self.figure_id!r} was not produced for this version "
                f"of the report.")


# ── context ──────────────────────────────────────────────────────────────


def build_context(doc: ReportDocument, figure_bytes: dict[str, bytes]) -> dict[str, Any]:
    """The Jinja2 context for ``doc`` — the keys the module docstring lists."""
    fields: dict[str, dict[str, Any]] = {}
    headline: dict[str, str] = {}
    disclosures: list[str] = []
    gaps: list[str] = []
    unverified: dict[str, list[str]] = {}
    for section in doc.sections:
        paragraphs = [b.md for b in section.blocks if isinstance(b, Paragraph)]
        bullets = [item for b in section.blocks if isinstance(b, Bullets)
                   for item in b.items]
        entry: dict[str, Any] = {
            "heading": section.heading,
            "status": section.status,
            "paragraphs": paragraphs,
            "bullets": bullets,
            "text": "\n\n".join(paragraphs),
            "note": section.note or "",
        }
        for block in section.blocks:
            if isinstance(block, Field):
                entry.setdefault(block.key, block.value)
                if section.section_id == EXECUTIVE_SUMMARY_ID:
                    headline.setdefault(block.key, block.value)
            elif isinstance(block, Callout):
                if block.kind == "disclosure":
                    disclosures.append(block.text)
                elif block.kind == "gap":
                    gaps.append(block.text)
        fields[section.section_id] = entry
        unverified[section.section_id] = list(section.audit.unverified)
    tables = {tid: {"columns": list(t.columns), "rows": [list(r) for r in t.rows],
                    "caption": t.caption or ""}
              for tid, t in doc.tables.items()}
    figures = {fid: _FigureTag(fid, f.caption, figure_bytes.get(fid))
               for fid, f in doc.figures.items()}
    for fid, png in figure_bytes.items():
        figures.setdefault(fid, _FigureTag(fid, None, png))
    meta = {
        "title": doc.title, "version": doc.version,
        "evidence_hash": doc.evidence_hash, "profile_id": doc.profile_id or "",
        "model": doc.model or "", "created_at": doc.created_at,
        "mode": doc.mode, "language": doc.language,
    }
    return {"meta": meta, "fields": fields, "headline": headline,
            "tables": tables, "figures": figures, "disclosures": disclosures,
            "gaps": gaps, "unverified": unverified}


# ── run-merged paragraph text ────────────────────────────────────────────


def _merged_text(paragraph: DocxParagraph) -> str:
    """The runs' texts joined — a tag Word split across runs reads whole."""
    return "".join(run.text for run in paragraph.runs)


def _has_tag(text: str) -> bool:
    return any(marker in text for marker in _TAG_MARKERS)


def _straighten_quotes(text: str) -> str:
    """Word's smart quotes inside a tag → straight quotes; prose untouched."""
    return _TAG_SPAN.sub(lambda m: m.group(0).translate(_SMART_QUOTES), text)


# ── errors ───────────────────────────────────────────────────────────────


def _known(context: dict[str, Any], text: str) -> str:
    parts = ["known names: " + ", ".join(CONTEXT_NAMES)]
    for name in ("fields", "tables", "figures"):
        if f"{name}." in text or f"{name}[" in text:
            ids = ", ".join(sorted(context[name])) or "(none)"
            parts.append(f"known {name}: {ids}")
    return "; ".join(parts)


def _error(exc: Exception, text: str, context: dict[str, Any]) -> TaggedRenderError:
    kind = "syntax error" if isinstance(exc, TemplateSyntaxError) else "cannot render"
    message = getattr(exc, "message", None) or str(exc)
    return TaggedRenderError(
        f"{kind}: {message}; in paragraph {text.strip()!r}; {_known(context, text)}")


# ── paragraph editing ────────────────────────────────────────────────────


def _clone_base_run(paragraph: DocxParagraph):
    """A ``w:r`` carrying the first run's formatting and no content."""
    runs = paragraph.runs
    if runs:
        r = copy.deepcopy(runs[0]._r)
        for child in list(r):
            if child.tag != qn("w:rPr"):
                r.remove(child)
        return r
    return paragraph.add_run()._r


def _drop_runs(paragraph: DocxParagraph) -> None:
    for run in paragraph.runs:
        run._r.getparent().remove(run._r)


def _insert_runs(paragraph: DocxParagraph, base_r, tokens) -> list[Run]:
    """Fresh runs cloned from ``base_r`` in place of the paragraph's runs."""
    originals = [run._r for run in paragraph.runs]
    anchor = originals[0] if originals else None
    out: list[Run] = []

    def place(r) -> Run:
        if anchor is not None:
            anchor.addprevious(r)
        else:
            paragraph._p.append(r)
        return Run(r, paragraph)

    for kind, text, url in tokens:
        run = place(copy.deepcopy(base_r))
        run.text = text
        if kind == "bold":
            run.bold = True
        elif kind == "italic":
            run.italic = True
        elif kind == "code":
            run.font.name = _CODE_FONT
        elif kind == "link":
            run.underline = True
            run.font.color.rgb = _LINK_COLOR
        out.append(run)
        if kind == "link" and url and url != text:
            extra = place(copy.deepcopy(base_r))
            extra.text = f" ({url})"
            extra.font.color.rgb = _LINK_COLOR
            out.append(extra)
    for r in originals:
        r.getparent().remove(r)
    return out


def _set_plain_text(paragraph: DocxParagraph, text: str) -> None:
    """``text`` into the first run (its formatting kept); the others go."""
    base_r = _clone_base_run(paragraph)
    _insert_runs(paragraph, base_r, [("text", text, None)])


def _set_markdown(paragraph: DocxParagraph, md: str) -> None:
    base_r = _clone_base_run(paragraph)
    _insert_runs(paragraph, base_r, tokenize_inline(md))


def _set_picture(paragraph: DocxParagraph, figure: _FigureTag) -> None:
    if not figure.png:
        _set_plain_text(paragraph, figure.missing_text())
        return
    base_r = _clone_base_run(paragraph)
    (run,) = _insert_runs(paragraph, base_r, [("text", "", None)])
    run.add_picture(io.BytesIO(figure.png), width=PICTURE_WIDTH)


def _clone_paragraph_after(paragraph: DocxParagraph, source_p) -> DocxParagraph:
    """A copy of ``source_p`` (style, first run) inserted after ``paragraph``."""
    p = copy.deepcopy(source_p)
    for tag in ("w:bookmarkStart", "w:bookmarkEnd"):
        for el in list(p.iter(qn(tag))):
            el.getparent().remove(el)
    paragraph._p.addnext(p)
    return DocxParagraph(p, paragraph._parent)


def _remove_paragraph(paragraph: DocxParagraph) -> None:
    p = paragraph._p
    parent = p.getparent()
    # Word needs at least one paragraph per table cell.
    if parent.tag == qn("w:tc") and len(parent.findall(qn("w:p"))) == 1:
        _drop_runs(paragraph)
        return
    parent.remove(p)


# ── rendering one paragraph ──────────────────────────────────────────────


class _Renderer:
    def __init__(self, context: dict[str, Any]) -> None:
        self.env = Environment(undefined=StrictUndefined, autoescape=False)
        self.context = context

    def render_text(self, text: str, context: dict[str, Any]) -> str:
        try:
            return self.env.from_string(text).render(context)
        except TemplateError as exc:
            raise _error(exc, text, context) from exc

    def evaluate(self, expr: str, text: str, context: dict[str, Any]) -> Any:
        try:
            value = self.env.compile_expression(expr, undefined_to_none=False)(**context)
            if isinstance(value, Undefined):
                str(value)  # StrictUndefined raises here with its own message
            return value
        except TemplateError as exc:
            raise _error(exc, text, context) from exc

    def paragraph(self, paragraph: DocxParagraph,
                  context: dict[str, Any] | None = None) -> None:
        context = self.context if context is None else context
        original = _merged_text(paragraph)
        if not _has_tag(original):
            return
        text = _straighten_quotes(original)
        only = _ONLY_TAG.match(text)
        if only and text.count("{{") == 1 and "{%" not in text:
            value = self.evaluate(only.group("expr"), text, context)
            if isinstance(value, _FigureTag):
                _set_picture(paragraph, value)
                return
            rendered = self.render_text(text, context)
            chunks = [c.replace("\n", " ").strip()
                      for c in rendered.split("\n\n") if c.strip()] or [""]
            source_p = copy.deepcopy(paragraph._p)
            _set_markdown(paragraph, chunks[0])
            current = paragraph
            for chunk in chunks[1:]:
                current = _clone_paragraph_after(current, source_p)
                _set_markdown(current, chunk)
            return
        rendered = self.render_text(text, context)
        if "{%" in text and not rendered.strip():
            _remove_paragraph(paragraph)
            return
        _set_plain_text(paragraph, rendered)

    # ── tables: row loops ────────────────────────────────────────────────

    def table(self, table: DocxTable, context: dict[str, Any] | None = None) -> None:
        context = self.context if context is None else context
        index = 0
        while True:
            rows = table.rows
            if index >= len(rows):
                return
            row = rows[index]
            cells = _row_cells(row)
            first = _cell_text(cells[0]) if cells else ""
            match = _FOR_ROW.match(_straighten_quotes(first))
            if not match:
                index += 1
                continue
            index += self._expand_loop(table, index, match, context)

    def _expand_loop(self, table: DocxTable, start: int, match: re.Match,
                     context: dict[str, Any]) -> int:
        """Replace the loop rows by one copy per item; the rows inserted."""
        rows = table.rows
        var, expr = match.group("var"), match.group("expr")
        first_cell = _row_cells(rows[start])[0]
        head_text = _straighten_quotes(_cell_text(first_cell))
        _strip_prefix(first_cell, len(match.group(0)))
        # Locate the endfor: this row's last cell, or a following row.
        template_rows: list[_Row] = []
        endfor_row: _Row | None = None
        for i in range(start, len(rows)):
            row = rows[i]
            cells = _row_cells(row)
            if i > start and _ENDFOR_ONLY.match(_row_text(cells)):
                endfor_row = row
                break
            template_rows.append(row)
            last = cells[-1] if cells else None
            if last is not None and _ENDFOR.search(_cell_text(last)):
                _strip_suffix(last)
                break
        else:
            raise TaggedRenderError(
                f"syntax error: '{{% for {var} in {expr} %}}' in a table row has no "
                f"matching '{{% endfor %}}'; in paragraph {head_text.strip()!r}; "
                f"{_known(context, head_text)}")
        items = self.evaluate(expr, head_text, context)
        try:
            items = list(items)
        except TypeError as exc:
            raise TaggedRenderError(
                f"cannot render: {expr!r} is not a list; in paragraph "
                f"{head_text.strip()!r}; {_known(context, head_text)}") from exc
        anchor = template_rows[0]._tr
        inserted = 0
        for n, item in enumerate(items):
            loop = {"index": n + 1, "index0": n, "first": n == 0,
                    "last": n == len(items) - 1, "length": len(items)}
            scope = {**context, var: item, "loop": loop}
            for template in template_rows:
                tr = copy.deepcopy(template._tr)
                anchor.addprevious(tr)
                inserted += 1
                for cell in _row_cells(_Row(tr, table)):
                    self.container(cell, scope)
        for row in template_rows:
            row._tr.getparent().remove(row._tr)
        if endfor_row is not None:
            endfor_row._tr.getparent().remove(endfor_row._tr)
        return inserted

    # ── containers: body, cells, headers, footers ────────────────────────

    def container(self, container, context: dict[str, Any] | None = None) -> None:
        """Every table (row loops first) and paragraph under ``container``."""
        context = self.context if context is None else context
        for table in list(container.tables):
            self.table(table, context)
            for cell in _table_cells(table):
                self.container(cell, context)
        for paragraph in list(container.paragraphs):
            self.paragraph(paragraph, context)


def _row_cells(row: _Row) -> list[_Cell]:
    """The row's distinct cells (a merged cell is listed once)."""
    out: list[_Cell] = []
    seen: set[int] = set()
    for cell in row.cells:
        if id(cell._tc) not in seen:
            seen.add(id(cell._tc))
            out.append(cell)
    return out


def _table_cells(table: DocxTable) -> Iterator[_Cell]:
    for row in table.rows:
        yield from _row_cells(row)


def _cell_text(cell: _Cell) -> str:
    return "\n".join(_merged_text(p) for p in cell.paragraphs)


def _row_text(cells: Iterable[_Cell]) -> str:
    return "\n".join(_cell_text(c) for c in cells)


def _strip_prefix(cell: _Cell, n: int) -> None:
    """Drop the first ``n`` merged characters of the cell's text."""
    for paragraph in cell.paragraphs:
        if n <= 0:
            return
        text = _merged_text(paragraph)
        take = min(n, len(text))
        if take:
            _set_plain_text(paragraph, text[take:])
        n -= take + 1  # the newline between paragraphs


def _strip_suffix(cell: _Cell) -> None:
    for paragraph in reversed(cell.paragraphs):
        text = _merged_text(paragraph)
        if text.strip():
            _set_plain_text(paragraph, _ENDFOR.sub("", text))
            return


def _header_footer_parts(document) -> Iterator[Any]:
    for section in document.sections:
        for name in ("header", "footer", "first_page_header", "first_page_footer",
                     "even_page_header", "even_page_footer"):
            part = getattr(section, name)
            if not part.is_linked_to_previous:
                yield part


# ── entry point ──────────────────────────────────────────────────────────


def render_tagged(doc: ReportDocument, template: bytes, *,
                  figure_bytes: dict[str, bytes]) -> bytes:
    """
    ``template`` (a ``.docx``) with its tags filled from ``doc``; the
    ``.docx`` bytes. A template without tags comes back with its text
    unchanged. Raises ``TaggedRenderError`` for a tag that cannot be filled.
    """
    document = Document(io.BytesIO(bytes(template)))
    renderer = _Renderer(build_context(doc, figure_bytes))
    renderer.container(document)
    for part in _header_footer_parts(document):
        renderer.container(part)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
