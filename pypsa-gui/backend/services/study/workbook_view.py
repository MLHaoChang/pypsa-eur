"""
The guided study's workbook sheets (U2 WP9; owner decision 9).

Plan: docs/superpowers/plans/2026-10-05-guided-study-u2-engine-rewire.md WP9
(§2 C8). The case workbook is the Investment Case engine's
(`engine_adapter.case_workbook`: `build_workbook` first); these are the GS
sheets appended after it, and the report workbook's per-option sheets:

* ``Cash flows (formulas)`` — a LIVE-FORMULA VIEW over the engine's project
  pre-tax cash series (`engine_adapter.cash_series`). The engine's own
  columns (CAPEX, replacements, the incremental operating cash, the terminal
  value, grants) are values; the net cash flow is the engine's identity as a
  formula (operating + terminal − CAPEX − replacements + grants), and the
  discount factor, discounted and cumulative cash flow are formulas on the
  rate cell. The NPV is pinned as ``=CF0 + NPV(rate, CF1:CFn)``: Excel's
  ``NPV`` discounts its first value one period, so the year-0 cash flow is
  added outside it (a range starting at CF0 would understate the NPV by the
  factor 1 + r). The guided decomposition (bills, savings, O&M) sits beside
  it with a reconciliation check; market revenue at duals is a column of its
  own, labelled excluded, and no formula reads it.
* ``KPIs`` — the guided case's KPIs; the NPV cell refers to the formula
  sheet's (live), the rest are STATIC (the model's values at the case rate).
* ``Assumptions`` — the ledger, one row per key, with its `engine_path`, the
  tested range, and who set each default (`provenance_label`: a guided-study
  rule, the tariff, the defaults pack — never "library", gate WP8a N9).
* ``Provenance`` — written by `engine_adapter.case_workbook` (it reads the
  engine's digests); the text helpers are here.

Text is written as text: a string that starts with ``=`` (a ledger source a
user typed, say) is stored as a string cell, never as a formula. The pro
forma's workbook (`proforma_xlsx`, the fallback until WP10) reuses these
helpers.
"""
from __future__ import annotations

from collections.abc import Mapping
from typing import Any

from models.study import AssumptionsLedger, InvestmentCase, LedgerRow

__all__ = ["CASH_FLOWS_TITLE", "LEDGER_COLUMNS", "LIVE", "STATIC", "assumptions",
           "basis_label", "engine_cash_flows", "formula", "kpis", "provenance_label", "text"]

RATE_CELL = "$B$1"
CASH_FLOWS_TITLE = "Cash flows (formulas)"
_HEADER_ROW = 4


def text(ws, row: int, col: int, value: Any):
    """
    A value that is never a formula (labels, ledger text, provenance). A
    control character openpyxl refuses (a study name, a ledger source) is
    replaced, never an untyped 500 (gate S7 [N9]).
    """
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("�", value)
    cell = ws.cell(row=row, column=col)
    cell.value = value
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"
    return cell


def formula(ws, row: int, col: int, expr: str):
    cell = ws.cell(row=row, column=col)
    cell.value = expr
    return cell


def basis_label(case: InvestmentCase) -> str:
    b = case.basis
    terms = b.terms.value if hasattr(b.terms, "value") else str(b.terms)
    tax = "pre-tax" if b.tax == "pre" else "post-tax"
    subsidy = "without subsidy" if b.subsidy == "excl" else "with subsidy"
    return f"{terms}, {tax}, {subsidy}"


# ── the formula view over the engine's cash series ────────────────────────

#: (label, source): ``("engine", key)`` is the engine's series, ``("view",
#: field)`` the guided view's year field, None a formula.
_ENGINE_CF = (
    ("Year", ("engine", "years")),
    ("CAPEX", ("engine", "capex")),
    ("Replacements", ("engine", "replacements")),
    ("Operating cash (engine)", ("engine", "operating")),
    ("Terminal value", ("engine", "terminal")),
    ("Grants", ("engine", "grants")),
    ("Net cash flow", None),
    ("Discount factor", None),
    ("Discounted cash flow", None),
    ("Cumulative discounted", None),
    ("Bill (baseline)", ("view", "bill_baseline")),
    ("Bill (option)", ("view", "bill_option")),
    ("Savings", ("view", "savings")),
    ("Fixed O&M", ("view", "opex_fixed")),
    ("Variable O&M", ("view", "opex_variable")),
    ("Check: savings - O&M - operating cash", None),
    ("Market revenue at duals (excluded)", ("view", "market_revenue_at_duals")),
)


def _header(ws, case: InvestmentCase, series: Mapping[str, Any] | None, title: str) -> None:
    ws.title = title
    rate = series["rate"] if series is not None else case.discount_rate
    text(ws, 1, 1, f"Discount rate ({basis_label(case).split(',')[0]})")
    ws.cell(1, 2).value = rate
    text(ws, 1, 3, "Currency year")
    ws.cell(1, 4).value = case.currency_year
    text(ws, 2, 1, "Basis")
    text(ws, 2, 2, basis_label(case))
    text(ws, 2, 3, "Status")
    text(ws, 2, 4, case.status)


def engine_cash_flows(ws, case: InvestmentCase, series: Mapping[str, Any] | None,
                      title: str = CASH_FLOWS_TITLE) -> int | None:
    """
    The formula view (module docstring); returns the NPV row, None when the
    case is not established (no cash flow, no NPV: the reasons are listed).
    `series` is `engine_adapter.cash_series(bundle)`.
    """
    _header(ws, case, series, title)
    text(ws, 3, 1, "Source")
    text(ws, 3, 2, "finance engine: project pre-tax cash (operating + terminal value "
                   "- CAPEX - replacements + grants), one row per year")
    for j, (label, _src) in enumerate(_ENGINE_CF, start=1):
        text(ws, _HEADER_ROW, j, label)
    first = _HEADER_ROW + 1
    if series is None or case.status != "ok" or not case.years:
        text(ws, first + 1, 1, "Not established")
        for k, note in enumerate(case.honesty_notes):
            text(ws, first + 1 + k, 2, note)
        return None
    L = {label: ws.cell(_HEADER_ROW, j).column_letter
         for j, (label, _s) in enumerate(_ENGINE_CF, start=1)}
    col = {label: j for j, (label, _s) in enumerate(_ENGINE_CF, start=1)}
    years = list(series["years"])
    view_years = {y.year: y for y in case.years}
    for i, t in enumerate(years):
        r = first + i
        for label, src in _ENGINE_CF:
            if src is None:
                continue
            kind, key = src
            if kind == "engine":
                v = series[key][i]
            else:
                y = view_years.get(t)
                v = None if y is None else getattr(y, key)
            if v is not None:  # a null figure stays an empty cell (listed below)
                ws.cell(r, col[label]).value = v
        # The engine's identity (finance.engine: ebitda = operating + terminal;
        # project pre-tax = ebitda - capex - replacement + grant), in its order.
        formula(ws, r, col["Net cash flow"],
                f"={L['Operating cash (engine)']}{r}+{L['Terminal value']}{r}"
                f"-{L['CAPEX']}{r}-{L['Replacements']}{r}+{L['Grants']}{r}")
        formula(ws, r, col["Discount factor"], f"=1/(1+{RATE_CELL})^{L['Year']}{r}")
        formula(ws, r, col["Discounted cash flow"],
                f"={L['Net cash flow']}{r}*{L['Discount factor']}{r}")
        cum = L["Cumulative discounted"]
        formula(ws, r, col["Cumulative discounted"],
                f"={L['Discounted cash flow']}{r}" if i == 0
                else f"={cum}{r - 1}+{L['Discounted cash flow']}{r}")
        formula(ws, r, col["Check: savings - O&M - operating cash"],
                f"={L['Savings']}{r}-{L['Fixed O&M']}{r}-{L['Variable O&M']}{r}"
                f"-{L['Operating cash (engine)']}{r}")
    last = first + len(years) - 1
    net = L["Net cash flow"]
    npv_row = last + 2
    text(ws, npv_row, 1, "NPV")
    # Excel's NPV discounts its FIRST value one period: CF0 stays outside.
    formula(ws, npv_row, 2, f"={net}{first}+NPV({RATE_CELL},{net}{first + 1}:{net}{last})")
    text(ws, npv_row + 1, 1, "Check: sum of discounted cash flows")
    formula(ws, npv_row + 1, 2, f"=SUM({L['Discounted cash flow']}{first}:"
                                f"{L['Discounted cash flow']}{last})")
    text(ws, npv_row + 2, 1, "Engine NPV at the case rate")
    ws.cell(npv_row + 2, 2).value = series["npv"]
    text(ws, npv_row + 2, 3, "static value from the finance engine: it does not move when "
                             "the rate cell or a cash flow is edited")
    text(ws, npv_row + 3, 1, "Market revenue at duals")
    text(ws, npv_row + 3, 2, "reported, not in the net cash flow (the bill prices that energy)")
    unavailable = sorted({f"{k}: {v}" for y in case.years for k, v in y.unavailable.items()})
    for k, note in enumerate(unavailable):
        text(ws, npv_row + 5 + k, 1, "Not available" if k == 0 else None)
        text(ws, npv_row + 5 + k, 2, note)
    return npv_row


# ── KPIs ──────────────────────────────────────────────────────────────────

# Gate S5 [N5] (carried to S7): the IRR, the paybacks and the LCOS are the
# model's values at the case's rate; they are written as STATIC cells and
# labelled so. Only the NPV cell is a live formula (it reads the cash-flow
# sheet, which reads the rate cell).
LIVE = "live formula: follows the rate cell and the cash flows"
STATIC = ("static value from the model at the case's rate: it does not move when "
          "the rate cell or a cash flow is edited")

# U2 WP8 (§5.4): `levelised_cost` (with its basis) and `terminal_value_eur`.
KPI_UNITS = {"npv": "EUR", "irr": "per unit", "payback_simple": "years",
             "payback_discounted": "years", "lcoe": "EUR/MWh",
             "levelised_cost": "EUR/MWh", "levelised_cost_basis": "",
             "lcoh": "EUR/MWh", "dscr_min": "ratio", "capex_total": "EUR",
             "terminal_value_eur": "EUR"}


def kpis(ws, case: InvestmentCase, npv_row: int | None, title: str = "KPIs",
         cash_flows_title: str = CASH_FLOWS_TITLE) -> None:
    ws.title = title
    for j, h in enumerate(("kpi", "value", "unit", "not available because", "cell"), start=1):
        text(ws, 1, j, h)
    rows: list[tuple[str, Any, str, str | None]] = [
        ("status", case.status, "", None),
        ("option_id", case.option_id, "", None),
        ("horizon_years", case.horizon_years, "years", None),
        ("discount_rate", case.discount_rate, "per unit (real)", None),
    ]
    k = case.kpis
    if k is not None:
        for key, unit in KPI_UNITS.items():
            rows.append((key, getattr(k, key), unit, k.unavailable.get(key)))
    r = 2
    for key, value, unit, why in rows:
        text(ws, r, 1, key)
        live = key == "npv" and npv_row is not None
        if live:
            formula(ws, r, 2, f"='{cash_flows_title}'!B{npv_row}")
        elif isinstance(value, str):
            text(ws, r, 2, value)
        elif value is not None:
            ws.cell(r, 2).value = value
        text(ws, r, 3, unit)
        text(ws, r, 4, why)
        if k is not None and key in KPI_UNITS:
            text(ws, r, 5, LIVE if live else STATIC)
        r += 1
    if case.value_streams:
        r += 1
        for j, h in enumerate(("value stream (annual saving by bill component)",
                               "value", "share", "engine"), start=1):
            text(ws, r, j, h)
        for s in case.value_streams:
            r += 1
            text(ws, r, 1, s.key or s.label)
            ws.cell(r, 2).value = s.annual_value
            ws.cell(r, 3).value = s.share
            text(ws, r, 4, s.engine)
    mr = case.market_revenue_at_duals
    if mr is not None:
        r += 2
        text(ws, r, 1, "market_revenue_at_duals (annual, excluded from the cash flow)")
        ws.cell(r, 2).value = mr.annual_value
        text(ws, r, 3, "EUR")
        text(ws, r, 4, mr.engine)


# ── Assumptions ───────────────────────────────────────────────────────────

LEDGER_COLUMNS = ("key", "label", "value", "unit", "currency_year", "basis",
                  "provenance", "status", "sensitivity_flag", "range_low",
                  "range_high", "source", "source_year", "source_url",
                  # U2 (plan §1): the engine field the row compiles to.
                  "engine_path",
                  # Gate S9 [S4]: the range the tornado tests a key driver on
                  # (`findings.bounds_for`), re-centred on an edited value
                  # outside the default band.
                  "tested_low", "tested_high", "tested_range_note")

#: Gate U2-S1 C4 / WP8a N9: who set a default, in words (`provenance_label`).
GUIDED_RULE = "guided study rule"
TARIFF = "tariff"
DEFAULTS_PACK = "defaults pack"
EARLIER_LIBRARY = "study library (earlier version)"


def provenance_label(row: LedgerRow, ledger_version: str) -> str:
    """
    Who set the row, in words: the user's own provenance (`user`,
    `imported`, `measured`) as stated; a default (`provenance=library`) by
    its author — a rule of the guided study (rows 21-34: its source names
    the rule, `library.GUIDED_RULE_SOURCE`), the tariff (the rows that
    compile into the commercial tariff or its export price), the generic
    defaults pack, or, on a ledger seeded before U2, the study's earlier
    library. Never the bare word "library" (gate WP8a N9).
    """
    from services.study import library as L

    if row.provenance != "library":
        return row.provenance
    if (row.source or "").startswith(L.GUIDED_RULE_SOURCE.split("{", 1)[0]):
        return GUIDED_RULE
    if (row.engine_path or "").startswith(("commercial.import_tariff",
                                           "commercial.export_price_ref")):
        return TARIFF
    if (ledger_version or "").startswith(L.DEFAULTS_VERSION_PREFIX):
        return DEFAULTS_PACK
    return EARLIER_LIBRARY


def assumptions(ws, ledger: AssumptionsLedger) -> None:
    from services.study.findings import bounds_for

    ws.title = "Assumptions"
    for j, h in enumerate(LEDGER_COLUMNS, start=1):
        text(ws, 1, j, h)
    for i, row in enumerate(ledger.rows, start=2):
        values = {
            "key": row.key, "label": row.label, "value": row.value, "unit": row.unit,
            "currency_year": row.currency_year,
            "basis": row.basis.value if hasattr(row.basis, "value") else row.basis,
            "provenance": provenance_label(row, ledger.ledger_version), "status": row.status,
            "sensitivity_flag": bool(row.sensitivity_flag),
            "range_low": row.range.low if row.range else None,
            "range_high": row.range.high if row.range else None,
            "source": row.source, "source_year": row.source_year,
            "source_url": row.source_url, "engine_path": row.engine_path,
        }
        tested = bounds_for(row) if row.sensitivity_flag else None
        values["tested_low"] = tested[0] if tested else None
        values["tested_high"] = tested[1] if tested else None
        values["tested_range_note"] = tested[2][0] if tested and tested[2] else None
        for j, col in enumerate(LEDGER_COLUMNS, start=1):
            text(ws, i, j, values[col])
