"""
The decision report as a Word document (plan S7; spec decision 19 amended:
MVP-1 ships DOCX and printable HTML, native PDF is MVP-2).

python-docx: the stale banner and the disclosures first (spec §7: before
the first number), then one heading per section with its prose (through
:func:`services.study.report.validate_prose`), its facts table, the same
PNG charts as the HTML (``services/study/report_charts.py``) and its data
tables; the Assumptions table marks the rows a user edited; a Provenance
appendix closes it. python-docx writes text runs, never markup, so user text
cannot inject anything.
"""
from __future__ import annotations

import io

from models.study import DecisionReport
from services.study import report as R
from services.study.render_html import num, pct, provenance

__all__ = ["render_docx"]


def _table(doc, header: list[str], rows: list[list[str]], style: str = "Light Grid Accent 1"):
    t = doc.add_table(rows=1, cols=len(header))
    try:
        t.style = style
    except (KeyError, ValueError):
        t.style = "Table Grid"
    for j, h in enumerate(header):
        t.rows[0].cells[j].text = h
    for row in rows:
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = "" if v is None else str(v)
    return t


def render_docx(report: DecisionReport, charts: dict[str, bytes] | None = None, *,
                stale: bool | None = None, stale_reasons: list[str] | None = None) -> bytes:
    from docx import Document
    from docx.shared import Inches

    # Gate S7 [N9]: a control character in user text would make python-docx
    # raise (an untyped 500); it is replaced, never passed through.
    report = R.xml_safe_report(report)
    prose = R.validate_prose(report)
    money, money_yr = R.money_unit(report), R.money_unit(report, per_year=True)
    if charts is None:
        from services.study import report_charts

        charts = report_charts.render_all(report)
    stale = report.stale if stale is None else stale
    stale_reasons = [R.xml_safe(r) for r in (
        report.stale_reasons if stale_reasons is None else stale_reasons)]

    def help_text(code: str) -> str:
        return report.honesty_help.get(code) or R.help_for(code)[0]

    doc = Document()
    doc.core_properties.title = report.study_name or "Decision study"
    doc.add_heading(report.study_name or "Decision study", level=0)
    doc.add_paragraph(report.question_title or report.question_id)

    if stale:
        p = doc.add_paragraph()
        p.add_run("This report is stale. ").bold = True
        p.add_run("What it was computed from has changed since; re-run the study and assemble "
                  "the report again.")
        for r in stale_reasons:
            doc.add_paragraph(f"{help_text(r)} ({r})", style="List Bullet")

    doc.add_heading("Read this first", level=1)
    for d in report.required_disclosures:
        # Gate S7 [N2]: a tariff's sentence is its author's text, labelled so.
        label = ("From the tariff (its author's wording, not checked by the study): "
                 if d.source == "tariff" else "")
        doc.add_paragraph(f"{label}{d.text} ({d.code})", style="List Bullet")
    if report.evidence_gaps:
        doc.add_heading("Evidence gaps", level=2)
        for g in report.evidence_gaps:
            doc.add_paragraph(g, style="List Bullet")
    if report.not_established:
        doc.add_heading("Not established", level=2)
        for g in report.not_established:
            doc.add_paragraph(g, style="List Bullet")

    for i, (sid, title) in enumerate(R.SECTIONS, start=1):
        s = report.sections[sid]
        doc.add_heading(f"{i}. {title} ({s.status})", level=1)
        if s.status != "ok" and s.note:
            doc.add_paragraph(f"{help_text(s.note)} ({s.note})")
        for text in prose[sid]:
            doc.add_paragraph(text)
        p = s.payload or {}
        if sid == "executive_summary":
            keys = p.get("headline_kpis") or []
            if keys:
                _table(doc, ["Key figure", "Value", "Basis and provenance"],
                       [[report.facts[k].label, R.format_fact(report.facts[k]),
                         provenance(report.facts[k])] for k in keys])
            if p.get("main_caveat"):
                doc.add_paragraph(f"Main caveat: {help_text(p['main_caveat'])} "
                                  f"({p['main_caveat']})")
            doc.add_paragraph(f"Maturity: {p.get('maturity_class') or 'not established'}")
            for r in p.get("reasons") or []:
                doc.add_paragraph(f"{help_text(r)} ({r})", style="List Bullet")
        elif s.facts:
            _table(doc, ["Figure", "Value", "Basis and provenance"],
                   [[f.label, R.format_fact(f), provenance(f)] for f in s.facts.values()])
        for fig in s.figures:
            if charts.get(fig):
                doc.add_picture(io.BytesIO(charts[fig]), width=Inches(6.0))
        if sid == "question" and p:
            doc.add_paragraph(f"Baseline: {p.get('baseline')}")
            _table(doc, ["Option", "Solved", "Battery MW", "PV MW", f"Battery NPV ({money})",
                         f"Option NPV ({money})", "Judged"],
                   [[f"{o['label']} ({o['option_id']})", o["solve_status"],
                     num(o["battery_mw"], 3), num(o["pv_mw"], 3), num(o["battery_npv"]),
                     num(o["option_npv"]), " ".join([o["attribution"] or "", *o["notes"]])]
                    for o in p.get("options") or []])
        if sid == "recommended_system" and p.get("explain"):
            _table(doc, ["Asset", "What limited its size", "Reading notes"],
                   [[e["name"], e["binding_constraint"] or "", " ".join(e["reading_notes"])]
                    for e in p["explain"]])
        if sid == "economics" and p:
            _table(doc, ["Year"] + [f"{h} ({money})" for h in (
                "CAPEX", "Replacements", "Fixed O&M", "Savings", "Salvage", "Net cash flow",
                "Cumulative discounted")],
                   [[y["year"], num(y["capex"]), num(y["replacements"]), num(y["opex_fixed"]),
                     num(y["savings"]), num(y.get("salvage")), num(y["net_cash_flow"]),
                     num(y["cumulative_discounted"])] for y in p.get("cash_flow") or []])
        if sid == "drivers" and p:
            doc.add_heading(p.get("value_streams_label") or "Value streams", level=2)
            _table(doc, ["Stream", f"Annual value ({money_yr})", "Share"],
                   [[v["label"], num(v["annual_value"]), pct(v["share"])]
                    for v in p.get("value_streams") or []])
            doc.add_heading("Tornado (battery NPV, sizes fixed)", level=2)
            _table(doc, ["Driver", "Low value", "High value", f"Battery NPV at low ({money})",
                         f"Battery NPV at high ({money})", f"Swing ({money})"],
                   [[r["label"], f"{num(r['low_value'], 4)} {r.get('unit') or ''}".strip(),
                     f"{num(r['high_value'], 4)} {r.get('unit') or ''}".strip(),
                     num(r["npv_low"]), num(r["npv_high"]), num(r["swing"])]
                    for r in p.get("tornado") or []])
        if sid == "robustness" and p:
            if p.get("pending"):
                doc.add_paragraph("Not reached: " + ", ".join(p["pending"]))
            if p.get("skipped"):
                _table(doc, ["Driver skipped", "Why"],
                       [[k, f"{help_text(c)} ({c})"] for k, c in p["skipped"].items()])
        if sid == "assumptions" and p:
            doc.add_heading("Assumptions table", level=2)
            _table(doc, ["Assumption", "Key", "Value", "Unit", "Range", "Source", "Edited"],
                   [[r["label"], r["key"], num(r["value"], 4), r["unit"],
                     f"{num(r['range_low'], 4)} to {num(r['range_high'], 4)}", r["source"],
                     "edited" if r["edited"] else ""] for r in p.get("rows") or []])
        if sid == "appendix" and p:
            doc.add_paragraph(p.get("model") or "")
            doc.add_heading("Provenance", level=2)
            rows = [[k, p.get(k)] for k in ("tariff_id", "tariff_name", "library_version",
                                            "fidelity", "ledger_hash", "base_network_hash")]
            forks = p.get("option_forks") or {}
            rows += [[f"network hash, {forks.get(fk, fk)}", h]
                     for fk, h in (p.get("option_network_hashes") or {}).items()]
            rows.append(["generated_at", report.generated_at.isoformat()])
            _table(doc, ["Field", "Value"], rows)
            doc.add_heading("Honesty codes", level=2)
            _table(doc, ["Code", "Meaning"], [[c, t] for c, t in report.honesty_help.items()])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
