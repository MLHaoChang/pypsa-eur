"""
The decision report as a workbook (plan S7).

Sheets, in order:

* ``Verdict`` — the verdict class, its rendered sentence, every report fact
  (value, unit, basis, currency year, engine, fidelity, why not established),
  the disclosures with their sentences, the stale flag and its reasons;
* ``Value streams`` and ``Tornado`` — the drivers section's tables;
* ``Assumptions`` — the ledger (``workbook_view.assumptions``: its
  ``engine_path``, the tested range, who set each default) plus an
  ``edited`` column;
* ``Provenance`` — hashes, engines, fidelity, library version;
* per option with an established case, ``Cash flows <option>`` and
  ``KPIs <option>`` — U2 WP9: the case workbook's own sheets
  (``workbook_view.engine_cash_flows`` / ``kpis``, as
  ``engine_adapter.case_workbook`` writes them) over the engine's cash series
  stored with the record (``engine_cash``, `engine_adapter.cash_series`): the
  live-formula view, the NPV pinned as ``=CF0 + NPV(rate, CF1:CFn)``, the
  IRR and paybacks labelled STATIC. The engine's own sheets
  (``build_workbook``) are in each option's case workbook, not repeated here.
  A case with no stored series (a record assembled before WP9, or a pro
  forma case) keeps the earlier sheet (``proforma_xlsx._cash_flows``) over
  the stored case, and says so; WP10 deletes that path.

Every text cell goes through ``workbook_view.text``: a string that starts
with ``=`` (a study name, a ledger source) is stored as a string, never a
formula.
"""
from __future__ import annotations

import io
from typing import Any

from collections.abc import Mapping

from models.study import AssumptionsLedger, DecisionReport, InvestmentCase
from services.study import proforma_xlsx as PX
from services.study import report as R
from services.study import workbook_view as WV

__all__ = ["write_report_xlsx"]


def _rows(ws, title: str, header: list[str], rows: list[list[Any]]) -> None:
    ws.title = title
    for j, h in enumerate(header, start=1):
        WV.text(ws, 1, j, h)
    for i, row in enumerate(rows, start=2):
        for j, v in enumerate(row, start=1):
            if isinstance(v, str) or v is None:
                WV.text(ws, i, j, v)
            else:
                ws.cell(i, j).value = v


def _verdict(ws, report: DecisionReport, prose: dict[str, list[str]], stale: bool,
             stale_reasons: list[str]) -> None:
    es = report.sections["executive_summary"].payload or {}
    rows: list[list[Any]] = [
        ["study", report.study_name, "", "", "", "", "", ""],
        ["verdict_status", es.get("verdict_status"), "", "", "", "", "", ""],
        ["verdict_class", es.get("verdict_class"), "", "", "", "", "", ""],
        ["verdict", " ".join(prose.get("executive_summary") or []), "", "", "", "", "", ""],
    ]
    # Gate S9 BC-S9-1: the drivers the verdict sentence refers to, by label.
    rows += [["verdict_driver", d["label"], d["key"], es.get("drivers_heading") or "",
              "", "", "", ""] for d in es.get("driver_labels") or []]
    rows.append(["stale", "yes" if stale else "no", "", "", "", "", "", ""])
    rows += [["stale_reason", r, R.help_for(r)[0], "", "", "", "", ""] for r in stale_reasons]
    rows.append(["", "", "", "", "", "", "", ""])
    rows.append(["fact", "value", "unit", "basis", "currency_year", "engine", "fidelity",
                 "not established because"])
    for key, f in report.facts.items():
        b = f.basis
        rows.append([key, f.value, f.unit,
                     None if b is None else f"{b.terms.value}, {b.tax}-tax, subsidy {b.subsidy}",
                     f.currency_year, f.engine,
                     None if f.fidelity is None else f.fidelity.value, f.unavailable])
    rows.append(["", "", "", "", "", "", "", ""])
    rows.append(["disclosure", "sentence", "source", "", "", "", "", ""])
    rows += [[d.code, d.text, d.source, "", "", "", "", ""] for d in report.required_disclosures]
    _rows(ws, "Verdict", ["field", "value", "", "", "", "", "", ""], rows)


#: The earlier sheet's note (a case with no stored engine series).
EARLIER_SHEET = ("the cash flows as stored with a report assembled before the engine's cash "
                 "series were recorded: the earlier sheet over the stored case; assemble the "
                 "report again for the live-formula view over the engine's series")


def _engine_series(case: InvestmentCase, series: Mapping | None) -> Mapping | None:
    """The stored series when it is this case's (same case id, the engine's)."""
    if not series or case.engine != "finance_engine":
        return None
    if series.get("case_id") != case.case_id or series.get("engine") != "finance_engine":
        return None
    return series


def write_report_xlsx(report: DecisionReport, cases: dict[str, InvestmentCase],
                      ledger: AssumptionsLedger, *, stale: bool | None = None,
                      stale_reasons: list[str] | None = None,
                      engine_cash: Mapping[str, Mapping] | None = None) -> bytes:
    import openpyxl

    prose = R.validate_prose(report)
    stale = report.stale if stale is None else stale
    stale_reasons = report.stale_reasons if stale_reasons is None else stale_reasons
    wb = openpyxl.Workbook()
    _verdict(wb.active, report, prose, stale, stale_reasons)
    drivers = report.sections["drivers"].payload or {}
    _rows(wb.create_sheet(), "Value streams",
          ["stream", "key", "annual_value", "share", "engine", "basis"],
          [[s.get("label"), s.get("key"), s.get("annual_value"), s.get("share"),
            s.get("engine"), drivers.get("value_streams_label")]
           for s in drivers.get("value_streams") or []])
    _rows(wb.create_sheet(), "Tornado",
          ["driver", "key", "unit", "low_value", "high_value", "battery_npv_low",
           "battery_npv_high", "swing", "evaluation"],
          [[r.get("label"), r.get("key"), r.get("unit"), r.get("low_value"),
            r.get("high_value"), r.get("npv_low"), r.get("npv_high"), r.get("swing"),
            r.get("evaluation")] for r in drivers.get("tornado") or []])
    ws = wb.create_sheet()
    WV.assumptions(ws, ledger)
    col = len(WV.LEDGER_COLUMNS) + 1
    WV.text(ws, 1, col, "edited")
    for i, row in enumerate(ledger.rows, start=2):
        edited = row.status == "customised" or row.provenance in ("user", "measured")
        WV.text(ws, i, col, "edited" if edited else "")
    app = report.sections["appendix"].payload or {}
    prov = [["study_id", report.study_id], ["question_id", report.question_id],
            ["generated_at", report.generated_at.isoformat()],
            ["currency_year", report.currency_year],
            ["fidelity", None if report.fidelity is None else report.fidelity.value]]
    prov += [[k, app.get(k)] for k in ("tariff_id", "library_version", "ledger_hash",
                                       "base_network_hash")]
    forks = app.get("option_forks") or {}
    prov += [[f"network_hash {forks.get(k, k)}", h]
             for k, h in (app.get("option_network_hashes") or {}).items()]
    prov += [["honesty_note", n] for n in report.honesty_notes]
    _rows(wb.create_sheet(), "Provenance", ["field", "value"], prov)
    for oid, case in sorted(cases.items()):
        if case.status != "ok":
            continue
        cf_title, kpi_title = f"Cash flows {oid}"[:31], f"KPIs {oid}"[:31]
        series = _engine_series(case, (engine_cash or {}).get(oid))
        cf = wb.create_sheet()
        if series is not None:
            npv_row = WV.engine_cash_flows(cf, case, series, title=cf_title)
        else:
            npv_row = PX._cash_flows(cf, case, title=cf_title)
            WV.text(cf, 3, 1, "Source")
            WV.text(cf, 3, 2, EARLIER_SHEET)
        WV.kpis(wb.create_sheet(), case, npv_row, title=kpi_title, cash_flows_title=cf_title)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
