"""
The pro forma as a workbook (guided investment study MVP-1, phase S5).

Plan: docs/superpowers/plans/2026-09-28-guided-investment-study-mvp1-v2.md (S5;
review v1 S9). openpyxl, as `services/asset_results/export.py` already uses.

Four sheets:

* ``Cash flows`` — one row per year (0 .. horizon). The money columns are the
  case's values; the net cash flow, discount factor, discounted and
  cumulative discounted cash flow are FORMULAS, so a reader can audit or
  change them. The NPV is pinned as ``=CF0 + NPV(rate, CF1:CFn)``: Excel's
  ``NPV`` discounts its first value one period, so the year-0 cash flow is
  added outside it (a range starting at CF0 would understate the NPV by the
  factor 1 + r). Market revenue at duals is a column of its own, labelled
  excluded, and no formula reads it.
* ``KPIs`` — the case's KPIs; the NPV cell refers to the cash-flow sheet's.
  A null KPI is an empty cell with its reason beside it, never a zero. The
  ``cell`` column labels each KPI live (the NPV) or STATIC (the IRR, the
  paybacks, the LCOS: the model's values at the case's rate, which do not
  move with the rate cell; gate S5 [N5], S7).
* ``Assumptions`` — the ledger, one row per key.
* ``Provenance`` — engines, basis, currency year, fidelity, library version,
  ledger and model hashes, the honesty notes and the upfront-cost gap.

Text is written as text: a string that starts with ``=`` (a ledger source a
user typed, say) is stored as a string cell, never as a formula.

U2 WP9: the guided workbooks are the engine's (`engine_adapter.case_workbook`,
`services/study/workbook_view.py`). This pro forma workbook stays only for the
pro forma fallback (a stored report record assembled before WP9, whose cases
carry no engine series) until WP10 deletes it; its text, KPI and ledger
sheets are the view module's.
"""
from __future__ import annotations

import io
from datetime import UTC, datetime
from typing import Any

from models.study import AssumptionsLedger, InvestmentCase
from services.study import workbook_view as WV

__all__ = ["write_proforma_xlsx"]

RATE_CELL = "$B$1"
_CF_COLUMNS = (
    ("Year", "year"),
    ("CAPEX", "capex"),
    ("Replacements", "replacements"),
    ("Fixed O&M", "opex_fixed"),
    ("Variable O&M", "opex_variable"),
    ("Bill (baseline)", "bill_baseline"),
    ("Bill (option)", "bill_option"),
    ("Savings", "savings"),
    ("Salvage", "salvage"),
    ("Net cash flow", None),
    ("Discount factor", None),
    ("Discounted cash flow", None),
    ("Cumulative discounted", None),
    ("Market revenue at duals (excluded)", "market_revenue_at_duals"),
)
_HEADER_ROW = 4
_COL = {label: j for j, (label, _f) in enumerate(_CF_COLUMNS, start=1)}


_text = WV.text
_formula = WV.formula
_basis_label = WV.basis_label


def _cash_flows(ws, case: InvestmentCase, title: str = "Cash flows") -> int | None:
    """The cash-flow sheet; returns the NPV row (None when there is no cash flow)."""
    ws.title = title
    _text(ws, 1, 1, "Discount rate (real)")
    ws.cell(1, 2).value = case.discount_rate
    _text(ws, 1, 3, "Currency year")
    ws.cell(1, 4).value = case.currency_year
    _text(ws, 2, 1, "Basis")
    _text(ws, 2, 2, _basis_label(case))
    _text(ws, 2, 3, "Status")
    _text(ws, 2, 4, case.status)
    for j, (label, _field) in enumerate(_CF_COLUMNS, start=1):
        _text(ws, _HEADER_ROW, j, label)
    if not case.years:
        return None
    letter = {label: ws.cell(_HEADER_ROW, j).column_letter
              for j, (label, _f) in enumerate(_CF_COLUMNS, start=1)}
    first = _HEADER_ROW + 1
    for i, y in enumerate(case.years):
        r = first + i
        for j, (label, field) in enumerate(_CF_COLUMNS, start=1):
            if field is None:
                continue
            v = getattr(y, field)
            if v is not None:  # a null figure stays an empty cell (listed below)
                ws.cell(r, j).value = v
        L = letter
        _formula(ws, r, _COL["Net cash flow"], (f"={L['Savings']}{r}+{L['Salvage']}{r}-{L['CAPEX']}{r}"
                             f"-{L['Replacements']}{r}-{L['Fixed O&M']}{r}"
                             f"-{L['Variable O&M']}{r}"))
        _formula(ws, r, _COL["Discount factor"], f"=1/(1+{RATE_CELL})^{L['Year']}{r}")
        _formula(ws, r, _COL["Discounted cash flow"], f"={L['Net cash flow']}{r}*{L['Discount factor']}{r}")
        cum = L["Cumulative discounted"]
        _formula(ws, r, _COL["Cumulative discounted"], (f"={L['Discounted cash flow']}{r}" if i == 0
                             else f"={cum}{r - 1}+{L['Discounted cash flow']}{r}"))
    last = first + len(case.years) - 1
    net = letter["Net cash flow"]
    npv_row = last + 2
    _text(ws, npv_row, 1, "NPV")
    # Excel's NPV discounts its FIRST value one period: CF0 stays outside.
    _formula(ws, npv_row, 2, f"={net}{first}+NPV({RATE_CELL},{net}{first + 1}:{net}{last})")
    _text(ws, npv_row + 1, 1, "Check: sum of discounted cash flows")
    _formula(ws, npv_row + 1, 2, f"=SUM({letter['Discounted cash flow']}{first}:"
                                 f"{letter['Discounted cash flow']}{last})")
    _text(ws, npv_row + 2, 1, "Market revenue at duals")
    _text(ws, npv_row + 2, 2, "reported, not in the net cash flow (the bill prices that energy)")
    unavailable = sorted({f"{k}: {v}" for y in case.years for k, v in y.unavailable.items()})
    for k, note in enumerate(unavailable):
        _text(ws, npv_row + 4 + k, 1, "Not available" if k == 0 else None)
        _text(ws, npv_row + 4 + k, 2, note)
    return npv_row


LIVE, STATIC = WV.LIVE, WV.STATIC
_KPI_UNITS = WV.KPI_UNITS
_LEDGER_COLUMNS = WV.LEDGER_COLUMNS
_assumptions = WV.assumptions


def _kpis(ws, case: InvestmentCase, npv_row: int | None, title: str = "KPIs",
          cash_flows_title: str = "Cash flows") -> None:
    WV.kpis(ws, case, npv_row, title=title, cash_flows_title=cash_flows_title)


def _provenance(ws, case: InvestmentCase, ledger: AssumptionsLedger) -> None:
    ws.title = "Provenance"
    p = case.provenance
    fidelity = case.fidelity.value if case.fidelity is not None else None
    rows: list[tuple[str, Any]] = [
        ("study_id", case.study_id), ("option_id", case.option_id),
        ("case_id", case.case_id), ("status", case.status),
        ("engine", case.engine),
        ("engines", ", ".join(p.engines) if p else None),
        ("basis", _basis_label(case)),
        ("perspective", case.perspective.value),
        ("currency_year", case.currency_year),
        ("fidelity", fidelity),
        ("library_version", (p.library_version if p else None) or ledger.ledger_version),
        ("ledger_hash", p.ledger_hash if p else None),
        ("model_hash", p.model_hash if p else None),
        ("project_ref", p.project_ref if p else None),
        ("tariff_id", p.tariff_id if p else None),
        ("discount_rate", case.discount_rate),
        ("horizon_years", case.horizon_years),
        ("salvage_basis", case.salvage_basis),
        ("generated_at", datetime.now(tz=UTC).isoformat(timespec="seconds")),
    ]
    for gap in case.upfront_gaps:
        rows += [
            (f"upfront_{gap.asset}_ledger_eur", gap.ledger_upfront_eur),
            (f"upfront_{gap.asset}_back_calculated_single_lifetime_eur",
             gap.back_calculated_upfront_eur),
            (f"upfront_{gap.asset}_gap_eur", gap.gap_eur),
        ]
    for note in case.honesty_notes:
        rows.append(("honesty_note", note))
    for i, (key, value) in enumerate(rows, start=1):
        _text(ws, i, 1, key)
        _text(ws, i, 2, value)


def write_proforma_xlsx(case: InvestmentCase, ledger: AssumptionsLedger) -> bytes:
    """The case and its ledger as an .xlsx workbook (bytes)."""
    import openpyxl

    wb = openpyxl.Workbook()
    npv_row = _cash_flows(wb.active, case)
    _kpis(wb.create_sheet(), case, npv_row)
    _assumptions(wb.create_sheet(), ledger)
    _provenance(wb.create_sheet(), case, ledger)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
