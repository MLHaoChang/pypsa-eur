"""
Series upload parsing: the timestamp and zone rules of a Library series
(Edge Investment Case P1 WP1.1b; moved out of `routers/library.py` in P2
WP2.4b-0, P1 condition 4, so the P2 importers share them).

Timestamps are ISO strings. Either ALL carry an offset (instants; `timezone`,
if given, is the zone to keep them in, else UTC) or NONE do (naive wall-clock,
stored naive; a `timezone` is then refused, the tariff engine's rule: a naive
axis plus a zone is ambiguous across DST). A zone pandas recognises that the
offset pattern does not still makes instants.

Pure service: no FastAPI. `SeriesInputError` (a `ValueError`) carries the
message the route returns as a 422.
"""
from __future__ import annotations

import io

import numpy as np
import pandas as pd

# A 15-minute year is 35,040 points; this admits ~28 of them. JSON bodies have
# no upload guard in this repo (`upload_guard` covers UploadFile only).
MAX_POINTS = 1_000_000

# A zone marker at the end of an ISO time part: Z, UTC, GMT, +01, +0100,
# +01:00 — case-insensitive (pandas reads `...t00:00z` as UTC too).
OFFSET_RE = r"(?i)(?:Z|UTC|GMT|[+-]\d{2}(?::?\d{2})?)$"


class SeriesInputError(ValueError):
    """An upload the series rules refuse (the route's 422)."""


def series_from(timestamps: list[str], values: list[float], timezone: str | None = None,
                *, max_points: int = MAX_POINTS) -> pd.Series:
    """The float series `values` on the parsed `timestamps` (see module docstring)."""
    if max(len(timestamps), len(values)) > max_points:
        raise SeriesInputError(f"at most {max_points:,} points per series")
    if len(timestamps) != len(values):
        raise SeriesInputError("timestamps and values must have the same length")
    raw = pd.Series(timestamps, dtype=str).str.strip()
    time_part = raw.str.extract(r"[Tt ](.*)$")[0].fillna("")
    aware = time_part.str.contains(OFFSET_RE, regex=True)
    if aware.any() and not aware.all():
        raise SeriesInputError("timestamps mix offset-carrying and naive values; "
                               "send every timestamp with an offset, or none")
    try:
        idx = pd.DatetimeIndex(pd.to_datetime(list(raw), utc=bool(aware.all())))
    except (ValueError, TypeError) as exc:
        raise SeriesInputError(f"unparseable timestamp: {exc}") from exc
    if idx.isna().any():
        # An empty cell or an uncached formula is no instant (WP2.4b-ii review F2).
        raise SeriesInputError(f"timestamp #{int(np.flatnonzero(idx.isna())[0]) + 1} is empty")
    if idx.tz is None and timezone is not None:
        raise SeriesInputError("timestamps carry no UTC offset, so `timezone` is ambiguous "
                               "across DST; send ISO timestamps with an offset")
    if idx.tz is not None:
        try:
            idx = idx.tz_convert(timezone or "UTC")
        except Exception as exc:  # noqa: BLE001 — bad zone name
            raise SeriesInputError(f"unknown timezone {timezone!r}") from exc
    return pd.Series(np.asarray(values, dtype=float), index=idx)


# ── file uploads (P2 WP2.4b-ii) ──────────────────────────────────────────────


_ISO_DATE = __import__("re").compile(r"^\d{4}-\d{2}-\d{2}")


def _cell(v) -> str:
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return "" if v is None else str(v)


def _rows(data: bytes, filename: str, cap: int) -> list[list[str]]:
    """The file's non-blank rows as strings, at most `cap` of them."""
    name = (filename or "").lower()
    out: list[list[str]] = []
    if name.endswith(".csv"):
        import csv

        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise SeriesInputError("the CSV is not UTF-8 text") from exc
        try:
            for r in csv.reader(io.StringIO(text)):
                if any(c.strip() for c in r):
                    out.append(r)
                    if len(out) > cap:
                        break
        except csv.Error as exc:            # an unterminated quote, a huge field (review F2)
            raise SeriesInputError(f"malformed CSV: {exc}") from exc
    elif name.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            for r in wb.worksheets[0].iter_rows(values_only=True):
                if any(c is not None and str(c).strip() for c in r):
                    out.append([_cell(c) for c in r])
                    if len(out) > cap:
                        break
        except Exception as exc:  # noqa: BLE001 — a corrupt workbook
            raise SeriesInputError(f"unreadable xlsx: {type(exc).__name__}") from exc
    else:
        raise SeriesInputError("upload a csv or xlsx file (timestamp,value)")
    return out


def parse_upload(data: bytes, filename: str, timezone: str | None = None, *,
                 max_points: int = MAX_POINTS, allow_empty: bool = False) -> pd.Series:
    """A `timestamp,value` CSV or the first sheet of an xlsx (header row
    `timestamp`, `value`; ISO 8601 timestamps), under `series_from`'s
    timestamp and zone rules, sorted by time. An empty value is refused with
    its row, or read as NaN with `allow_empty` (meter data: its month is then
    not established)."""
    rows = _rows(data, filename, max_points + 1)   # the header + the cap, then stop
    if not rows or [c.strip().lower() for c in rows[0][:2]] != ["timestamp", "value"] \
            or any(str(c).strip() for c in rows[0][2:]):
        raise SeriesInputError("the first row must be the header timestamp,value")
    body = rows[1:]
    if not body:
        raise SeriesInputError("the upload has no rows")
    if len(body) > max_points:
        raise SeriesInputError(f"at most {max_points:,} points per series")
    stamps, values = [], []
    for i, r in enumerate(body, start=2):
        if len(r) < 2 or any(str(c).strip() for c in r[2:]):
            raise SeriesInputError(f"row {i}: expected two columns, timestamp and value")
        raw = str(r[1]).strip()
        if not raw:
            if not allow_empty:
                raise SeriesInputError(f"row {i}: the value is empty")
            values.append(float("nan"))
        else:
            try:
                v = float(raw)
            except ValueError as exc:
                raise SeriesInputError(f"row {i}: value {raw!r} is not a number") from exc
            if not np.isfinite(v):
                raise SeriesInputError(f"row {i}: value {raw!r} is not finite")
            values.append(v)
        stamp = str(r[0]).strip()
        if stamp and not _ISO_DATE.match(stamp):
            # `01/02/2030` would be read month-first, a bare time as today
            # (review F14): a file's timestamps start with an ISO date.
            raise SeriesInputError(f"row {i}: timestamp {stamp!r} does not start with an ISO "
                                   "date (YYYY-MM-DD)")
        stamps.append(stamp)
    return series_from(stamps, values, timezone, max_points=max_points).sort_index()


# ── meter data → monthly history (P2 WP2.4b-ii) ──────────────────────────────


SETTLEMENTS = {"15min": 0.25, "30min": 0.5, "h": 1.0}
METER_UNITS = ("kW", "W", "kWh_per_interval")
LABELS = ("start", "end")
_GAP_FACTOR = 1.5   # a step longer than this × the month's step is a gap, not coverage
_STEP_TOL = 0.02    # jitter tolerated on the step and on interval boundaries

# The Library's meter-data help (the metered-year sentence is carried from
# WP2.1a-iii review round 1).
METER_HISTORY_HELP = (
    "Meter history seeds demand ratchets and prices non-convex energy tiers. Upload the "
    "site's import meter as timestamp,value rows and say what the values are: average kW "
    "over each row's interval, W, or kWh per interval (converted to kW); and whether the "
    "timestamp labels the interval's start (default) or its end. Give the site's timezone: "
    "offset timestamps are placed on it, and plain local times are read as that zone's "
    "wall clock. Peaks are the highest settlement-interval mean of each complete month on "
    "the site clock; import the history at the SAME settlement as the tariff's demand "
    "item. Energy is each month's import; export rows (negative) count as no import, row by "
    "row. A cyclic or months-mode ratchet reads history under the SAME rate-year key, so it "
    "resolves only when the modelled rate year is a metered year: for a future year use "
    "non-cyclic range mode with meter history, or cyclic_year (which reads the modelled "
    "months themselves).")


def _durations(idx: pd.DatetimeIndex, months: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    """(hours each row covers, each row's month step): the step is the median
    gap inside the row's month, so a file mixing 15- and 5-min months keeps
    both (review F11); a gap longer than 1.5 steps covers one step."""
    d = np.diff(idx.asi8) / 3.6e12
    nxt = np.append(d, np.nan)
    frame = pd.DataFrame({"m": months, "d": nxt})
    step_of = frame.groupby("m")["d"].median()
    overall = float(np.median(d)) if len(d) else 1.0
    step = frame["m"].map(step_of).fillna(overall).to_numpy(dtype=float)
    dur = np.where(np.isfinite(nxt) & (nxt <= _GAP_FACTOR * step), nxt, step)
    return dur, step


def meter_series(series: pd.Series, *, timezone: str | None, unit: str = "kW",
                 label: str = "start") -> tuple[pd.Series, list[str]]:
    """The upload as average kW over each row's interval, each row labelled
    by its interval START, on the site clock (review F1, F4, F5, F6):

    * offset timestamps need the site's `timezone` (never silently UTC);
    * plain local times with a `timezone` are that zone's wall clock — the
      repeated fall-back hour resolved by row order, a nonexistent time refused;
    * `unit`: kW as is, W / 1000, kWh per interval / the interval's hours;
    * `label="end"` moves every row back by its step."""
    if unit not in METER_UNITS:
        raise SeriesInputError(f"unit is one of {', '.join(METER_UNITS)}, got {unit!r}")
    if label not in LABELS:
        raise SeriesInputError(f"label is one of {', '.join(LABELS)}, got {label!r}")
    notes: list[str] = []
    s = series.sort_index()
    idx = pd.DatetimeIndex(s.index)
    if idx.tz is not None:
        if timezone is None:
            raise SeriesInputError("the timestamps carry UTC offsets: give the site's timezone "
                                   "so months and settlement windows are on the site clock")
        try:
            idx = idx.tz_convert(timezone)
        except Exception as exc:  # noqa: BLE001 — bad zone name
            raise SeriesInputError(f"unknown timezone {timezone!r}") from exc
    elif timezone is not None:
        try:
            # The repeated fall-back hour: its first occurrence is summer time
            # (row order), the second winter time; a third is a duplicate.
            idx = idx.tz_localize(timezone, ambiguous=~idx.duplicated(keep="first"),
                                  nonexistent="raise")
        except Exception as exc:  # noqa: BLE001 — a time that does not exist, or no order
            raise SeriesInputError(f"the local times cannot be placed on {timezone}: "
                                   f"{str(exc)[:200]}") from exc
        notes.append("local_times_read_on_the_site_clock")
    if idx.has_duplicates:
        raise SeriesInputError("meter timestamps repeat")
    if len(idx) < 2:
        raise SeriesInputError("meter data needs at least two rows")
    months = np.asarray(idx.year * 100 + idx.month)
    dur, step = _durations(idx, months)
    if label == "end":
        idx = idx - pd.to_timedelta(step, unit="h")
        months = np.asarray(idx.year * 100 + idx.month)
    v = s.to_numpy(dtype=float)
    if unit == "W":
        v = v / 1000.0
    elif unit == "kWh_per_interval":
        v = v / dur
    return pd.Series(v, index=idx), notes


def meter_history(series: pd.Series, *, settlement: str = "15min",
                  timezone: str | None = None) -> dict:
    """Monthly import peaks (kW, on `settlement` interval means) and energy
    (kWh) of a kW meter series labelled by interval start (`meter_series`),
    on the site clock (`timezone` converts a tz-aware series). Only COMPLETE
    months are history; a month the rows do not cover, one with a NaN, and a
    month missing between the first and the last are named in `notes`.
    Peaks need rows that fit the settlement intervals: a coarser meter, or
    rows crossing an interval boundary, give energy only (named)."""
    from services.commercial.tariff_engine import interval_key

    if settlement not in SETTLEMENTS:
        raise SeriesInputError(f"settlement is one of {', '.join(SETTLEMENTS)}, "
                               f"got {settlement!r}")
    s = series.sort_index()
    if s.index.has_duplicates:
        raise SeriesInputError("meter timestamps repeat")
    idx = pd.DatetimeIndex(s.index)
    if idx.tz is not None and timezone is not None:
        idx = idx.tz_convert(timezone)
    if len(idx) < 2:
        raise SeriesInputError("meter data needs at least two rows")
    notes: list[str] = []
    d = np.diff(idx.asi8) / 3.6e12
    grid = pd.Timedelta(hours=float(np.median(d))).round("s")
    if grid > pd.Timedelta(0):
        utc = idx.tz_convert("UTC") if idx.tz is not None else idx
        snapped = utc.round(grid)
        if idx.tz is not None:
            snapped = snapped.tz_convert(idx.tz)
        dev = np.abs(snapped.asi8 - idx.asi8) / 3.6e12
        if dev.any() and (dev <= _STEP_TOL * grid.total_seconds() / 3600.0).all() \
                and not snapped.has_duplicates:
            # Logger jitter of a few seconds (review F11): on the step grid.
            idx = snapped
            notes.append("timestamps_snapped_to_the_step_grid")
    code = np.asarray(idx.year * 100 + idx.month)
    dur, step = _durations(idx, code)
    kw = s.to_numpy(dtype=float)
    known = ~np.isnan(kw)
    neg = int((kw[known] < 0).sum())
    if neg:
        notes.append(f"negative_rows_read_as_export:{neg}")
        if neg > 0.5 * known.sum():
            notes.append("mostly_negative_rows_check_the_sign_convention")
    imp = np.clip(np.nan_to_num(kw, nan=0.0), 0.0, None)
    frame = pd.DataFrame({"m": code, "e": imp * dur, "dur": np.where(known, dur, 0.0),
                          "nan": ~known, "t0": idx.asi8, "t1": idx.asi8 + dur * 3.6e12})
    g = frame.groupby("m").agg(e=("e", "sum"), cov=("dur", "sum"), nan=("nan", "any"),
                               t0=("t0", "min"), t1=("t1", "max"))
    complete: list[int] = []
    for m, row in g.iterrows():
        y, mo = divmod(int(m), 100)
        start = pd.Timestamp(year=y, month=mo, day=1)
        end = start + pd.offsets.MonthBegin(1)
        if idx.tz is not None:
            # A DST change at midnight on the 1st (review F2): shift forward.
            start = start.tz_localize(idx.tz, nonexistent="shift_forward", ambiguous=True)
            end = end.tz_localize(idx.tz, nonexistent="shift_forward", ambiguous=True)
        hours = (end - start).total_seconds() / 3600.0
        tol = 0.5 * float(np.median(step[code == m]))
        ok = (not row["nan"] and row["cov"] >= hours - tol and row["t0"] <= start.value
              and row["t1"] >= end.value - tol * 3.6e12)
        (complete.append(int(m)) if ok else notes.append(f"month_incomplete:{y:04d}-{mo:02d}"))
    first, last = int(g.index.min()), int(g.index.max())
    cur = pd.Period(year=first // 100, month=first % 100, freq="M")
    while cur.year * 100 + cur.month < last:
        if cur.year * 100 + cur.month not in g.index:
            notes.append(f"month_missing:{cur.strftime('%Y-%m')}")
        cur += 1

    def key(m: int) -> str:
        return f"{m // 100:04d}-{m % 100:02d}"

    energy = {key(m): float(g.at[m, "e"]) for m in complete}
    for m in complete:
        if energy[key(m)] == 0.0 and neg:
            notes.append(f"month_without_import:{key(m)}")
    peaks: dict[str, float] = {}
    settle_h = SETTLEMENTS[settlement]
    coarse = float(np.max(step)) > settle_h * (1 + _STEP_TOL)
    ikey = interval_key(idx, {"15min": "15min", "30min": "30min", "h": "h"}[settlement])
    offset_h = (idx.asi8 - ikey) / 3.6e12
    crosses = bool((offset_h + dur > settle_h * (1 + _STEP_TOL)).any())
    if coarse:
        notes.append(f"peaks_not_established:meter_step_{float(np.max(step)):g}h_coarser_than_"
                     f"settlement_{settlement}")
    elif crosses:
        notes.append(f"peaks_not_established:meter_rows_cross_{settlement}_intervals")
    else:
        per = pd.DataFrame({"key": ikey, "e": imp * dur, "dur": dur, "m": code}) \
            .groupby("key").agg(e=("e", "sum"), dur=("dur", "sum"), m=("m", "first"))
        top = (per["e"] / per["dur"]).groupby(per["m"]).max()
        peaks = {key(m): float(top[m]) for m in complete if m in top.index}
    return {"meter_history_peaks_kw": peaks, "meter_history_energy_kwh": energy,
            "settlement": settlement, "notes": notes}
