"""
The Scenarios sheet (owner decision RW-SCEN-1): the base run beside every
sensitivity run of it, key KPIs per period and over the horizon, deltas
against the base, and native charts.

The runs are the base's sensitivity cases as `harness/sensitivity_comparison`
reads them: each case is opened only when the acting user may open it
(`project_registry.resolve_project`, the check opening the project makes; the
others are counted, never named), at most `MAX_CASES` of them, each read from
its own saved context (`workflow_tools._load_context`: its own network,
solver state and user series; the active context is not touched), and judged
comparable with the base by the same reasons the comparison gives
(`sensitivity_comparison._reasons`). A run that is not comparable shows its
reasons and no figures; a delta is a live formula only between two number
cells.

The route supplies the two router-owned reads per run: `result_df_for(state)`
(the result lookup over that run's solver state) and `ic_key(ctx, result_df)`
(the run's current investment-case staleness key, or `(None, why)`).
"""
from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass, field
from typing import Any

from openpyxl.utils import get_column_letter

from services.results_workbook import charts
from services.results_workbook.cells import BOLD, formula, heading, put, width
from services.results_workbook.overview import (
    CASE_ROWS, RUN_THE_CASE, OverviewFigures, overview_figures,
)

__all__ = ["TITLE", "Run", "Scenarios", "gather_experiments", "gather_scenarios", "write_scenarios"]

TITLE = "Scenarios"
HORIZON = "horizon"


@dataclass
class Run:
    """One run of the comparison; `runs[0]` is the base."""

    name: str
    project_id: str | None = None
    figs: OverviewFigures | None = None
    # `sensitivity_comparison._values(read)`: metric → (value, why, note).
    headline: dict | None = None
    ic_report: Any = None
    case_state: dict | None = None
    # `sensitivity_comparison` reason codes; empty when comparable with the base.
    reasons: list[str] = field(default_factory=list)
    # The run's own results, which the result-view sheets read (package 02):
    # its network, solver config, lost-load capture, result lookup over its own
    # solver state, and the settings its solve used. None when not comparable.
    network: Any = None
    cfg: Any = None
    lost_load: dict | None = None
    result_df: Callable | None = None
    solved_settings: Callable[[], dict | None] | None = None
    # The run's own saved context (`workflow_tools._load_context`), for the
    # views whose route reads a whole context (Load flow). None for the base,
    # whose context is the active one.
    context: Any = None


@dataclass
class Scenarios:
    runs: list[Run]                # the sensitivity runs, base excluded
    no_access: int = 0             # cases the acting user may not open
    over_cap: int = 0              # cases past `MAX_CASES`
    base_headline: dict | None = None
    # The open base differs from its saved run (the Compare-view rows read the saved one).
    base_unsaved: bool = False
    # The base's own results are current (else no run is compared with it).
    base_current: bool = True
    # Children named as experiment cases (`experiment_execution._name`): never
    # sensitivity runs; `gather_experiments` shows the acting user's own and
    # counts the others in `unlisted`, never naming them. Child id → name.
    experiment_children: dict[str, str] = field(default_factory=dict)
    unlisted: int = 0
    # The experiment store could not be read: no experiment case is shown.
    experiments_unreadable: bool = False

    @property
    def not_shown(self) -> int:
        return self.no_access + self.over_cap + self.unlisted


# The Scenarios sheet's own reasons for an experiment case, beside
# `sensitivity_comparison._REASONS` (the shared vocabulary, GK B (11) condition 3).
EXPERIMENT_REASONS = {
    "evidence_tampered": "its run evidence changed since its receipt",
    "case_failed": "its solve failed",
    "case_not_completed": "it has not completed",
    "case_unavailable": "its recorded result cannot be read back (see the experiment)",
}


def _cases(base_project_id: str) -> tuple[Any, list[tuple[Any, dict]], int, set[str]]:
    """(base, [(case, spec)] the acting user may open, by name, count of the
    others, the ids of the children named as experiment cases)."""
    from fastapi import HTTPException

    from harness import sensitivity_comparison as sc
    from harness.experiment_execution import is_case_project_name
    from services import chat_tools, project_registry, workflow_tools as wt
    with chat_tools._acting() as (db, user):
        base = sc._resolve(db, user, base_project_id)
        cases, no_access, experiment_children = [], 0, {}
        for child in project_registry.direct_children(db, base):
            if is_case_project_name(child.name):  # an experiment case: shown once, as one
                experiment_children[str(child.id)] = child.name
                continue
            try:  # access first: nothing of a child the user cannot open is read
                child = project_registry.resolve_project(db, user, str(child.id))
            except HTTPException:
                no_access += 1
                continue
            spec = wt._read_json(project_registry.project_dir(child) / wt.SENSITIVITY_CASE_FILE)
            if spec:
                cases.append((child, spec))
    return base, sorted(cases, key=lambda pair: pair[0].name.lower()), no_access, experiment_children


def _fill(run: Run, read: dict, *, result_df_for, ic_key, solved_settings_for, dispatch_ready) -> None:
    """A comparable run's figures from its own saved context (`read`, as
    `sensitivity_comparison._read` gives it); a run whose dispatch the routes
    would not serve gets `case_unsolved` and no figures."""
    from services.finance.investment_case_runner import staleness
    from services.finance.report import load_ic_report
    from harness import sensitivity_comparison as sc
    if not run.reasons and dispatch_ready is not None and not dispatch_ready(read["ctx"].network):
        run.reasons = ["case_unsolved"]
    if run.reasons:
        return
    ctx = read["ctx"]
    state = ctx.solver_state
    result_df = result_df_for(state)
    run.figs = overview_figures(ctx.network, state.get("solver_config"),
                                state.get("last_lost_load"), result_df)
    run.network, run.cfg, run.result_df = ctx.network, state.get("solver_config"), result_df
    run.context = ctx
    run.lost_load = state.get("last_lost_load")
    run.solved_settings = (solved_settings_for(state, ctx.network) if solved_settings_for
                           else lambda state=state: state.get("last_solve_settings"))
    run.headline = sc._values(read)
    run.ic_report = load_ic_report(state)
    if run.ic_report is not None:
        current, why = ic_key(ctx, result_df)
        run.case_state = staleness(run.ic_report.model_dump(mode="json"), current, reason=why)


def gather_scenarios(base_project_id: str, base_basis: dict, *,
                     result_df_for: Callable[[dict], Callable],
                     ic_key: Callable[[Any, Callable], tuple[tuple[str, dict] | None, str | None]],
                     solved_settings_for: Callable[[dict, Any], Callable[[], dict | None]] | None = None,
                     dispatch_ready: Callable[[Any], bool] | None = None,
                     ) -> Scenarios | None:
    """The base's sensitivity runs, or None when it has none the user can see
    or hide. `base_basis` is the open base's `fingerprint`, `mode`, `basis`
    (`experiment_comparison.basis_digests`) and `ok` (its results are current).
    `solved_settings_for(state, network)` and `dispatch_ready(network)` are the
    routes' own reads (`_solved_settings`, `_dispatch_ready`) over a run's
    context: a run whose dispatch the routes would not serve shows no figures."""
    from harness import sensitivity_comparison as sc
    base, cases, no_access, experiment_children = _cases(base_project_id)
    if not cases and not no_access and not experiment_children:
        return None
    shown, over_cap = cases[:sc.MAX_CASES], max(0, len(cases) - sc.MAX_CASES)
    runs = []
    for project, spec in shown:
        read = sc._read(project)
        reasons = ([] if base_basis["ok"] else ["baseline_unavailable"]) + sc._reasons(
            base_basis, read, spec, base_basis["fingerprint"])
        run = Run(name=project.name, project_id=str(project.id), reasons=reasons)
        _fill(run, read, result_df_for=result_df_for, ic_key=ic_key,
              solved_settings_for=solved_settings_for, dispatch_ready=dispatch_ready)
        runs.append(run)
    return _with_base(base, base_basis, Scenarios(
        runs=runs, no_access=no_access, over_cap=over_cap, experiment_children=experiment_children))


def _with_base(base, base_basis: dict, scen: Scenarios) -> Scenarios:
    from harness import sensitivity_comparison as sc
    base_read = sc._read(base)
    scen.base_headline = sc._values(base_read)
    scen.base_unsaved = base_read["fingerprint"] != base_basis["fingerprint"]
    scen.base_current = bool(base_basis["ok"])
    return scen


def _experiment_cases(base_project_id: str):
    """Under `experiments._LOCK`: (root, [(case record, ledger copy or None,
    ledger readable)]) of the acting user's own experiments on the base,
    in index order, each two-case plan as its two cases. Reads only the base's
    experiment store: no case project is opened here."""
    from fastapi import HTTPException

    from harness import experiment_execution as engine
    from harness.durable import network_guard
    from harness.experiment_two_case import select_case
    from services import experiments as ex
    root, project, owner = ex._authority()   # the open project, as the acting user may open it
    if project != str(base_project_id):
        return None
    out = []
    with network_guard(), ex._LOCK, ex._transaction(root):
        for experiment_id in ex._load_index(root)["order"]:
            try:  # the owner and project are `_read_record`'s own scope (GK B condition 1)
                record = ex._read_record(root, experiment_id, project, owner)
            except HTTPException as exc:
                if exc.status_code == 404:
                    continue
                raise
            if record["domain"] != "network":
                continue
            cases = ([select_case(record, c["id"]) for c in record["cases"]]
                     if len(record["cases"]) == 2 else [record])
            for case in cases:
                if not case.get("operations"):  # never armed: no ledger
                    out.append((case, None, True))
                    continue
                try:
                    out.append((case, engine._read(root, case), True))
                except (ValueError, KeyError, TypeError, OSError, HTTPException):
                    out.append((case, None, False))
    return root, out


def _unique(name: str, experiment_id: str, taken: set[str]) -> str:
    out = name if name not in taken else f"{name} ({experiment_id[:8]})"
    i = 2
    while out in taken:
        out = f"{name} ({experiment_id[:8]}, {i})"
        i += 1
    taken.add(out)
    return out


def _experiment_run(root, record, ledger, readable, base_basis, name, fill) -> Run | int:
    """One experiment case's run, outside `experiments._LOCK`, or 1 when the
    acting user may not open its project (counted, never named): access, then
    the receipt observed on the ledger copy, the baseline checks, the load and
    the post-load checks against the receipt."""
    from fastapi import HTTPException

    from harness import experiment_comparison as xc
    from harness import sensitivity_comparison as sc
    from services import chat_tools, project_registry
    prefix = [] if base_basis["ok"] else ["baseline_unavailable"]
    case_project_id = (ledger or {}).get("case_project_id")
    run = Run(name=name, project_id=case_project_id)
    if not readable or not xc.adapter_available(record) or ledger is None:
        run.reasons = prefix + ["case_unavailable" if not readable else "case_not_completed"]
        return run
    project = None
    if case_project_id:
        with chat_tools._acting() as (db, user):
            try:  # access first: nothing of a case the user cannot open is read
                project = project_registry.resolve_project(db, user, str(case_project_id))
            except HTTPException:
                return 1
    out = xc._start(record)
    try:
        data = xc._verify(root, record, ledger, out)
    except xc.BaselineEvidenceChanged:
        run.reasons = prefix + ["baseline_changed"]
        return run
    except (ValueError, KeyError, TypeError, OSError, RuntimeError, HTTPException):
        run.reasons = prefix + ["case_unavailable"]
        return run
    if data is None or project is None:
        status = out["status"]
        run.reasons = prefix + ["case_failed" if status == "failed" else
                                "case_unavailable" if status == "unavailable" else "case_not_completed"]
        return run
    baseline_fingerprint = record["preflight"]["input_fingerprint"]
    reasons = prefix + (["baseline_changed"] if baseline_fingerprint != base_basis["fingerprint"] else [])
    if reasons:
        run.reasons = reasons
        return run
    try:
        read = sc._read(project)
        xc._load_case(record, data, saved=read["ctx"])
    except xc.ReceiptMismatch:
        run.reasons = ["case_inputs_changed"]
        return run
    except xc.EvidenceChanged:
        run.reasons = ["evidence_tampered"]
        return run
    except (ValueError, KeyError, TypeError, OSError, RuntimeError, HTTPException):
        run.reasons = ["case_unavailable"]
        return run
    # The receipt fixed the case's inputs; the base's are checked above.
    run.reasons = sc._reasons(base_basis, read, {"input_fingerprint": read["fingerprint"],
                                                 "baseline_fingerprint": baseline_fingerprint},
                              base_basis["fingerprint"])
    fill(run, read)
    return run


def gather_experiments(base_project_id: str, base_basis: dict, scen: Scenarios | None, *,
                       result_df_for: Callable[[dict], Callable],
                       ic_key: Callable[[Any, Callable], tuple[tuple[str, dict] | None, str | None]],
                       solved_settings_for: Callable[[dict, Any], Callable[[], dict | None]] | None = None,
                       dispatch_ready: Callable[[Any], bool] | None = None,
                       ) -> Scenarios | None:
    """`scen` (the sensitivity runs, or None) with the acting user's own
    experiment cases on the base after them, in plan order: the cases left
    after the sensitivity runs take their `MAX_CASES` slots are counted before
    any is opened; each shown case is observed and loaded outside
    `experiments._LOCK` and checked against its receipt. Experiment-case
    children that are not one of these cases are counted, never named. None
    when there is nothing to show or count."""
    from fastapi import HTTPException

    from harness import experiment_execution as engine
    from harness import sensitivity_comparison as sc
    try:
        gathered = _experiment_cases(base_project_id)
        unreadable = False
    except HTTPException as exc:
        if exc.status_code == 404:
            raise
        gathered, unreadable = None, True
    root, cases = gathered if gathered else (None, [])
    mine = {str(ledger["case_project_id"]) for _c, ledger, _ok in cases
            if ledger and ledger.get("case_project_id")}
    # A case whose ledger cannot be read is shown as unavailable: its child,
    # known by the engine's name for it, is that case's, not another's.
    unread = {engine._name(c) for c, _ledger, ok in cases if not ok}
    children = scen.experiment_children if scen else {}
    unlisted = sum(1 for cid, name in children.items() if cid not in mine and name not in unread)
    if not cases and not unlisted and not unreadable:
        return scen
    if scen is None:
        from services import chat_tools
        with chat_tools._acting() as (db, user):
            row = sc._resolve(db, user, str(base_project_id))
        scen = _with_base(row, base_basis, Scenarios(runs=[]))
    scen.unlisted += unlisted
    scen.experiments_unreadable = unreadable
    slots = max(0, sc.MAX_CASES - len(scen.runs))
    shown, scen.over_cap = cases[:slots], scen.over_cap + max(0, len(cases) - slots)
    from services.pypsa_service import PyPSAService
    taken = {PyPSAService.get_loaded_project() or "Base", *(r.name for r in scen.runs)}

    def fill(run, read):
        _fill(run, read, result_df_for=result_df_for, ic_key=ic_key,
              solved_settings_for=solved_settings_for, dispatch_ready=dispatch_ready)
    for record, ledger, readable in shown:
        name = _unique(f"{record['title']}: {record['cases'][0]['id']}", record["experiment_id"], taken)
        run = _experiment_run(root, record, ledger, readable, base_basis, name, fill)
        if isinstance(run, Run):
            scen.runs.append(run)
        else:
            taken.discard(name)
            scen.no_access += run
    return scen


# ── the sheet ────────────────────────────────────────────────────────────

def _results_text(run: Run) -> str:
    from harness import sensitivity_comparison as sc
    if not run.reasons:
        return "current"
    known = [r for r in run.reasons if r in sc._REASONS]
    # The run's own name may hold ": " (an experiment case's "<title>: <case>").
    texts = [sc._not_comparable("", known).split(": ", 1)[1]] if known else []
    return "; ".join(texts + [EXPERIMENT_REASONS[r] for r in run.reasons if r not in sc._REASONS])


def _base_results_text(scen: Scenarios) -> str:
    from harness import sensitivity_comparison as sc
    return "current" if scen.base_current else (
        f"{sc._BASE_REASONS['case_stale']}; no run is compared with it")


def _case_text(run: Run) -> str:
    if run.reasons:
        return "not compared"
    if run.ic_report is None:
        return f"none: {RUN_THE_CASE}"
    st = run.case_state or {}
    if st.get("stale"):
        changed = ", ".join(st.get("changed") or []) or st.get("reason") or "?"
        return f"out of date (changed: {changed}): its figures are not shown"
    return "current"


def _case_value(run: Run, reader) -> float | None:
    if run.reasons or run.ic_report is None or (run.case_state or {}).get("stale"):
        return None
    try:
        v = reader(run.ic_report)
    except (AttributeError, KeyError, TypeError):
        return None
    return v if _is_num(v) else None


def _rows(base: Run, runs: list[Run], base_headline: dict | None) -> list[tuple[str, str, list]]:
    """(key, label, [value per run, base first]) in sheet order: every Overview line
    over the horizon then per period, LCOE by carrier, the investment case,
    then the Compare view's headline metrics."""
    from harness import sensitivity_comparison as sc
    every = [base, *runs]
    periods = base.figs.periods if base.figs else []
    out: list[tuple[str, str, list]] = []

    def line_value(run: Run, key: str, period: str):
        if run.figs is None:
            return None
        line = next((x for x in run.figs.lines if x.key == key), None)
        if line is None:
            return None
        return line.total if period == HORIZON else line.by_period.get(period)

    for period in [HORIZON, *periods]:
        for line in base.figs.lines if base.figs else []:
            out.append((f"{line.key}@{period}", f"{line.label.strip()} ({line.unit}), {period}",
                        [line_value(r, line.key, period) for r in every]))
    carriers = sorted({c for r in every if r.figs for c in r.figs.lcoe_by_carrier})
    for period in [HORIZON, *periods]:
        for carrier in carriers:
            vals = []
            for r in every:
                total, by = (r.figs.lcoe_by_carrier.get(carrier, (None, {})) if r.figs else (None, {}))
                vals.append(total if period == HORIZON else by.get(period))
            out.append((f"lcoe:{carrier}@{period}", f"LCOE, {carrier} (EUR/MWh), {period}", vals))
    for key, label, unit, reader in CASE_ROWS:
        out.append((f"case:{key}", f"{label} ({unit}), {HORIZON}", [_case_value(r, reader) for r in every]))
    for metric, (label, unit, *_rest) in sc.METRICS.items():
        vals = [None if (r is base and base_headline is None) else
                ((base_headline if r is base else r.headline) or {}).get(metric, (None,))[0]
                for r in every]
        out.append((f"compare:{metric}", f"{label} ({unit}), {HORIZON}, Compare view", vals))
    return out


def _is_num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _is_established(v) -> bool:
    """A number, or `True` for a formula written over two number cells."""
    return v is True or _is_num(v)


def write_scenarios(ws, base: Run, scen: Scenarios) -> dict:
    """Write the sheet; returns where things went (rows by key, columns)."""
    ws.title = TITLE
    runs = scen.runs
    every = [base, *runs]
    put(ws, 1, 1, "Scenarios: the base run, its sensitivity runs and its experiment cases")
    ws.cell(1, 1).font = BOLD
    put(ws, 2, 1, "Deltas are run − base, live formulas between two number cells; a run "
                  "that cannot be compared with the base shows why and no figures.")
    notes = []
    if scen.no_access:
        notes.append(f"{scen.no_access} scenario(s) or experiment case(s) not shown: no access")
    if scen.over_cap:
        from harness import sensitivity_comparison as sc
        notes.append(f"{scen.over_cap} scenario(s) or experiment case(s) not shown: the sheet "
                     f"compares up to {sc.MAX_CASES}, sensitivity runs by name, then "
                     "experiment cases in plan order")
    if scen.unlisted:
        notes.append(f"{scen.unlisted} experiment case project(s) not shown: not a case of "
                     "your own experiments on this base")
    if scen.experiments_unreadable:
        notes.append("Experiment cases not shown: the experiment store needs repair.")
    if scen.base_unsaved:
        notes.append("The open base has unsaved changes: its Compare-view rows read its saved run.")
    for i, text in enumerate(notes):
        put(ws, 3 + i, 1, text)

    r = 4 + len(notes)
    heading(ws, r, ["Run", "Results", "Investment case"])
    for i, run in enumerate(every):
        row = r + 1 + i
        put(ws, row, 1, f"{run.name} (base)" if run is base else run.name)
        put(ws, row, 2, _base_results_text(scen) if run is base else _results_text(run))
        put(ws, row, 3, _case_text(run))

    header = r + len(every) + 3
    first_val, last_val = 2, 1 + len(every)
    first_delta = last_val + 1
    heading(ws, header, ["KPI", *[run.name for run in every],
                         *[f"Δ {run.name} − base" for run in runs]])
    rows = {}
    row = header
    for key, label, vals in _rows(base, runs, scen.base_headline):
        row += 1
        rows[key] = row
        put(ws, row, 1, label)
        # System LCOE is a formula on the Overview, so the line has no value: here
        # too it is a live formula over the run's own cost and energy cells.
        line, _, period = key.partition("@")
        cost, energy = rows.get(f"system_cost@{period}"), rows.get(f"energy_served@{period}")
        for j, v in enumerate(vals):
            col = get_column_letter(first_val + j)
            if (line == "system_lcoe" and cost and energy and _is_num(ws[f"{col}{cost}"].value)
                    and _is_num(ws[f"{col}{energy}"].value) and ws[f"{col}{energy}"].value != 0):
                formula(ws, row, first_val + j, f"={col}{cost}*1000/{col}{energy}")
                vals[j] = True        # established: a delta may read it
            else:
                put(ws, row, first_val + j, v)
        base_ref = f"{get_column_letter(first_val)}{row}"
        for j, run in enumerate(runs):
            col = first_delta + j
            if not run.reasons and _is_established(vals[0]) and _is_established(vals[1 + j]):
                formula(ws, row, col, f"={get_column_letter(first_val + 1 + j)}{row}-{base_ref}")
            else:
                put(ws, row, col, None)

    made = []
    anchor = f"{get_column_letter(first_delta + len(runs) + 1)}{header}"
    cap = rows.get(f"capex_annualised@{HORIZON}")
    if cap is not None and rows.get(f"opex@{HORIZON}") == cap + 1:
        made.append(charts.bar(ws, title="Capex (annualised) and opex over the horizon, by run",
                               stacked=True, from_rows=True, data=(1, cap, last_val, cap + 1),
                               cats=(first_val, header, last_val, header), anchor=anchor,
                               y_title="M EUR"))
    lcoe = rows.get(f"system_lcoe@{HORIZON}")
    if lcoe is not None:
        made.append(charts.bar(ws, title="System LCOE over the horizon, by run", from_rows=True,
                               data=(1, lcoe, last_val, lcoe),
                               cats=(first_val, header, last_val, header),
                               anchor=f"{get_column_letter(first_delta + len(runs) + 1)}{header + 20}",
                               y_title="EUR/MWh"))
    width(ws, {"A": 60, "B": 22, "C": 40})
    return {"header": header, "rows": rows, "first_val": first_val,
            "first_delta": first_delta, "charts": len(made)}
