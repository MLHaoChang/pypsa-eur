"""
The finance sheets: the investment-case workbook (`finance/export_xlsx.py`)
inside the results workbook, registered for the `investment` result view.

`export_xlsx.build_workbook` stays the one writer of those sheets; its output
is copied in cell by cell (values and cell types, so a formula-looking text it
stored as text stays text) with number formats, each sheet prefixed "Case "
so it cannot collide with a result view's sheet. That writer emits values
only: no formula, defined name or chart refers to a sheet title, so the
rename breaks no reference (`write_case_sheets` checks it and refuses a
source that ever does). Without a stored report the builder writes nothing
and About says "run the investment case": export never runs it (RW-KPI-1).
"""
from __future__ import annotations

import io

from services.results_workbook.overview import RUN_THE_CASE
from services.results_workbook.registry import register

__all__ = ["CASE_PREFIX", "write_case_sheets"]

CASE_PREFIX = "Case "


def _copy(src, dst) -> None:
    for row in src.iter_rows():
        for c in row:
            if c.value is None:
                continue
            d = dst.cell(row=c.row, column=c.column)
            d.value = c.value
            d.data_type = c.data_type
            d.number_format = c.number_format
    for letter, dim in src.column_dimensions.items():
        dst.column_dimensions[letter].width = dim.width


def write_case_sheets(wb, report, *, project: str | None = None) -> list[str]:
    """Append the case workbook's sheets to `wb`; returns their titles."""
    from openpyxl import load_workbook

    from services.finance.export_xlsx import build_workbook

    src = load_workbook(io.BytesIO(build_workbook(report, project=project)))
    # GK B condition 8: copying cells would break any cross-sheet reference, and
    # openpyxl does not carry charts between workbooks. The source has none;
    # if it ever gains one, fail loudly rather than export a broken workbook.
    if src.defined_names or any(ws._charts or ws.defined_names or any(c.data_type == "f" for row in ws.iter_rows()
                                                  for c in row) for ws in src.worksheets):
        raise RuntimeError("export_xlsx now writes formulas, names or charts: "
                           "write_case_sheets must rewrite their sheet references")
    titles = []
    for ws in src.worksheets:
        dst = wb.create_sheet((CASE_PREFIX + ws.title)[:31])
        _copy(ws, dst)
        titles.append(dst.title)
    return titles


@register("investment", no_results=f"no results: {RUN_THE_CASE}")
def investment_sheets(wb, ctx) -> bool:
    if ctx.ic_report is None:
        return False
    ctx.sheets["investment"] = write_case_sheets(wb, ctx.ic_report, project=ctx.project)
    return True
