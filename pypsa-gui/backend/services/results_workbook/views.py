"""
The result-view sheets (package 02, `.scratch/results-workbook/issues/02-result-view-sheets.md`).

Each view's builder reads its data per run from the same functions its tab's
route calls (`views_network.py`) and hands `write_view` one `ViewData` per
run (`ctx.runs`, the base first). The sheet holds:

* a key-results block: one row per key figure and period (and the horizon),
  one column per run, then a live delta against the base for each other run
  where both cells are numbers (RW-SCEN-1);
* the view's full data: each table in long form with a Scenario column, every
  run's rows under one header, every snapshot of every investment period;
* 1–3 native charts (`CHARTS`, one table for every view), each drawn from the
  base run's rows of one of the sheet's own tables, over number cells only.

A run with no data for the view (not comparable with the base, or the view
has nothing for it) shows `not_established` in its key column and no rows.
"""
from __future__ import annotations

import math
from dataclasses import dataclass, field
from typing import Any, Literal

import pandas as pd
from openpyxl.utils import get_column_letter

from services.results_workbook import charts
from services.results_workbook.cells import BOLD, formula, heading, put, width

__all__ = ["CHARTS", "HORIZON", "Chart", "Key", "Table", "ViewData", "write_view"]

HORIZON = "horizon"


@dataclass
class Key:
    """One key figure: `by` maps a period label (or `HORIZON`) to its value."""

    label: str
    unit: str
    by: dict[str, float | None]


@dataclass
class Table:
    """A long table: the first `labels` columns of `frame` label each row
    (period, snapshot, carrier …), the rest hold the values."""

    name: str
    unit: str
    frame: pd.DataFrame
    labels: int = 1
    # A snapshot table (its last label column is the snapshot): how a coarser
    # resolution aggregates it (`budget.KINDS`) and each row's hours.
    time: str | None = None
    weights: Any = None
    # The coarser resolution it was written at, named (`budget.coarsen`).
    resolution: str | None = None


@dataclass
class ViewData:
    keys: list[Key] = field(default_factory=list)
    tables: list[Table] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class Chart:
    """A chart over the base run's rows of table `table`: categories from its
    label column `cats` (or the x values of a scatter), one series per value
    column whose base cells are all numbers."""

    kind: Literal["bar", "stacked", "line", "scatter"]
    title: str
    table: str
    y_title: str | None = None
    cats: int = -1
    # Only these value columns (default: every number column of the table).
    columns: tuple[str, ...] = ()


#: The charts of every view, chosen here in one place (spec item 5).
CHARTS: dict[str, tuple[Chart, ...]] = {
    "capex": (
        Chart("stacked", "Optimal capacity by carrier by period",
              "Optimal capacity by carrier and period", "MW"),
        Chart("stacked", "Capex by carrier by period", "Capex by carrier and period", "EUR"),
        Chart("stacked", "Opex by carrier by period", "Opex by carrier and period", "EUR"),
    ),
    "dispatch": (
        Chart("stacked", "Generation by carrier by period", "Energy by carrier and period", "MWh"),
        Chart("line", "Generator dispatch (base, every snapshot)", "Generator dispatch", "MW"),
    ),
    "flows": (
        Chart("stacked", "Energy by source by period", "Energy by source and period", "MWh"),
    ),
    "loadshape": (
        Chart("stacked", "Annual load served and supply by period", "Annual totals", "MWh"),
    ),
    "storage": (
        Chart("bar", "Equivalent full cycles by unit by period", "Cycles by unit and period",
              "cycles"),
        Chart("line", "State of charge (base, every snapshot)", "State of charge", "MWh"),
    ),
    "curtailment": (
        Chart("stacked", "Curtailed energy by plant by period",
              "Curtailed energy by plant and period", "MWh"),
        Chart("line", "Curtailment (base, every snapshot)", "Curtailment", "MW"),
    ),
    "lostload": (
        Chart("bar", "Lost load by bus by period", "Lost load by bus and period", "MWh"),
    ),
    "prices": (
        Chart("scatter", "Price duration curve by bus and period", "Price duration", "EUR/MWh"),
        Chart("bar", "Mean price by bus by period", "Mean price by bus and period", "EUR/MWh"),
    ),
    "economics": (
        Chart("bar", "Net profit by asset by period", "Net profit by asset and period", "EUR"),
        Chart("bar", "LCOE by asset", "Asset economics", "EUR/MWh", cats=0,
              columns=("lcoe_eur_per_mwh",)),
    ),
    "emissions": (
        Chart("stacked", "Emissions by carrier by period", "Emissions by carrier and period", "tCO2"),
    ),
    "loadflow": (
        Chart("bar", "Peak loading by branch by period", "Peak loading by branch and period", "%",
              cats=0),
        Chart("line", "Line flow (base, every snapshot)", "Line flow (bus0)", "MW"),
    ),
    "grid_code": (
        Chart("bar", "PCC reactive power at the critical hours", "PCC reactive power", "Mvar",
              columns=("q_final_mvar",)),
        Chart("bar", "DC capacity estimate by bus", "DC capacity estimate by bus", "MW", cats=0),
    ),
    "electrical_design": (
        Chart("bar", "Electrical capex by period", "Electrical cost by period", "EUR", cats=0,
              columns=("capex_eur",)),
        Chart("bar", "Transformer utilisation (worst hour)", "Transformers", "per unit",
              cats=0, columns=("utilisation",)),
    ),
    "island": (
        Chart("bar", "Cost of island capability by scenario", "Cost of island capability", "EUR",
              cats=0, columns=("cost of island capability",)),
    ),
    "grid_forming": (
        Chart("bar", "Capacity needed by case", "Capacity needed", None, cats=0,
              columns=("delta",)),
        Chart("bar", "Grid-forming settings: current and recommended", "Recommended settings",
              None, cats=2),
    ),
    "adequacy": (
        Chart("bar", "ENS cap and achieved ENS by period", "ENS target by period", "MWh",
              cats=0, columns=("cap_mwh", "achieved_ens_mwh")),
        Chart("bar", "Firm and required capacity at peak by period", "Reserve margin by period",
              "MW", cats=0, columns=("required_mw", "firm_mw")),
    ),
    # Ticket 30 (condition 6): one chart per criticality basis. A column is
    # charted only when every base row has a number in it, so no chart mixes
    # a per-horizon row with a per-year one.
    "fmea": (
        Chart("bar", "Criticality by failure mode, EUR per study horizon (ticket 30)",
              "Failure modes", "EUR per study horizon (ticket 30)", cats=0,
              columns=("criticality, EUR per study horizon (ticket 30)",)),
        Chart("bar", "Criticality by failure mode, EUR per year", "Failure modes",
              "EUR per year", cats=0, columns=("criticality, EUR per year",)),
        Chart("bar", "ΔEUE by failure mode", "Failure modes", "MWh", cats=0,
              columns=("ΔEUE, MWh",)),
    ),
}

_KEY_ROW = 4
#: More series than this make a chart unreadable; the cells keep them all.
_MAX_SERIES = 24


def _is_num(v: Any) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)


def _label(v: Any) -> Any:
    if isinstance(v, pd.Timestamp):
        return v.isoformat()
    if hasattr(v, "item"):  # numpy scalar
        v = v.item()
    return v


def _union(seqs) -> list:
    out: dict = {}
    for s in seqs:
        for x in s:
            out.setdefault(x, None)
    return list(out)


def _write_keys(ws, row: int, names: list[str], datas: list[ViewData | None]) -> int:
    labels = _union([[k.label for k in d.keys] for d in datas if d])
    if not labels:
        return row
    others = names[1:]
    heading(ws, row, ["Key result", "Unit", "Period", *names,
                      *[f"Δ {name} vs base" for name in others]])
    row += 1
    for label in labels:
        per_run = [next((k for k in d.keys if k.label == label), None) if d else None
                   for d in datas]
        unit = next(k.unit for k in per_run if k)
        for period in _union([k.by for k in per_run if k]):
            put(ws, row, 1, label)
            put(ws, row, 2, unit)
            put(ws, row, 3, period)
            for i, k in enumerate(per_run):
                put(ws, row, 4 + i, k.by.get(period) if k else None)
            base = ws.cell(row=row, column=4)
            for j in range(1, len(names)):
                other = ws.cell(row=row, column=4 + j)
                col = 3 + len(names) + j
                if _is_num(base.value) and _is_num(other.value):
                    formula(ws, row, col, f"={other.coordinate}-{base.coordinate}")
                else:
                    put(ws, row, col, None)
            row += 1
    return row + 1


def _write_table(ws, row: int, name: str, names: list[str],
                 tables: list[Table | None]) -> tuple[int, dict]:
    """Writes one table's rows for every run; returns the next free row and the
    base block's position for its charts."""
    base = next(t for t in tables if t is not None)
    label_cols = list(base.frame.columns[:base.labels])
    value_cols = _union([list(t.frame.columns[t.labels:]) for t in tables if t is not None])
    unit = "; ".join(x for x in (base.unit, base.resolution) if x)
    put(ws, row, 1, f"{name} ({unit})" if unit else name)
    ws.cell(row=row, column=1).font = BOLD
    row += 1
    header = row
    heading(ws, row, ["Scenario", *map(str, label_cols), *map(str, value_cols)])
    row += 1
    pos = {"header": header, "first": row, "last": row - 1, "labels": len(label_cols),
           "values": {}, "resolution": base.resolution}
    for i, (run_name, t) in enumerate(zip(names, tables)):
        if t is None:
            continue
        frame = t.frame.reindex(columns=label_cols + value_cols)
        for rec in frame.itertuples(index=False, name=None):
            put(ws, row, 1, run_name)
            for j, v in enumerate(rec, start=2):
                v = _label(v)
                if j - 2 < len(label_cols):
                    put(ws, row, j, v)
                elif _is_num(v):
                    ws.cell(row=row, column=j).value = float(v)
                else:
                    put(ws, row, j, None)
            row += 1
        if i == 0:
            pos["last"] = row - 1
            for j, col in enumerate(value_cols, start=2 + len(label_cols)):
                if t.frame.shape[0] and col in t.frame.columns and all(
                        _is_num(ws.cell(row=r, column=j).value)
                        for r in range(pos["first"], pos["last"] + 1)):
                    pos["values"][col] = j
    return row + 1, pos


def _charts(ws, view_id: str, positions: dict[str, dict], anchor_col: str) -> int:
    drawn = 0
    for spec in CHARTS.get(view_id, ()):
        pos = positions.get(spec.table)
        if not pos or not pos["values"] or pos["last"] < pos["first"]:
            continue
        cat_col = 2 + (spec.cats % pos["labels"])
        anchor = f"{anchor_col}{_KEY_ROW + 20 * drawn}"
        title = spec.title
        if pos["resolution"]:  # a coarsened table's chart names its resolution
            title = title.replace("every snapshot", pos["resolution"])
        h, r0, r1 = pos["header"], pos["first"], pos["last"]
        cols = sorted(j for col, j in pos["values"].items()
                      if not spec.columns or col in spec.columns)
        if not cols:
            continue
        if spec.kind == "scatter":
            charts.scatter(ws, title=title, x=(cat_col, r0, cat_col, r1),
                           ys=[(c, h, c, r1) for c in cols[:_MAX_SERIES]],
                           anchor=anchor, y_title=spec.y_title)
        else:
            # The leading contiguous run of number columns, one box.
            last = cols[0]
            while last + 1 in cols and last + 1 - cols[0] < _MAX_SERIES:
                last += 1
            box = (cols[0], h, last, r1)
            cats = (cat_col, r0, cat_col, r1)
            if spec.kind == "line":
                charts.line(ws, title=title, data=box, cats=cats, anchor=anchor,
                            y_title=spec.y_title)
            else:
                charts.bar(ws, title=title, data=box, cats=cats, anchor=anchor,
                           stacked=spec.kind == "stacked", y_title=spec.y_title)
        drawn += 1
    return drawn


def write_view(wb, ctx, view_id: str, title: str, datas: list[ViewData | None]) -> bool:
    """Writes the view's sheet; False (and no sheet) when no run has data."""
    if not any(d and (d.keys or d.tables) for d in datas):
        return False
    from services.results_workbook import budget
    names = [r.name for r in ctx.runs]
    key_rows = _KEY_ROW + 2 + sum(len(k.by) for d in datas if d for k in d.keys)
    datas, resolution, base_only = budget.fit(datas, key_rows, ctx.cells_left)
    ws = wb.create_sheet(title[:31])
    put(ws, 1, 1, title)
    ws.cell(row=1, column=1).font = BOLD
    notes = _union([d.notes for d in datas if d])
    said = []
    if resolution:
        said.append(f"resolution {resolution}")
        notes.append(f"The full data passes Excel's row limit or the workbook's cell budget, "
                     f"so every run's snapshot tables are at {resolution} resolution (each "
                     "table's title says how its cells aggregate; each bucket stays within "
                     "one investment period). The key results use every snapshot.")
    if base_only and len(datas) > 1:
        said.append("other scenarios at key results only")
        notes.append(f"Even then the sheet does not fit, so the tables hold {names[0]} only; "
                     "the other scenarios are at key-results level.")
    if said:
        ctx.resolution[view_id] = "; ".join(said)
    ctx.cells_left -= sum(t.frame.size for d in datas if d for t in d.tables)
    missing = [f"{name}: {why}" for name, why in zip(names, _why_missing(ctx, datas)) if why]
    put(ws, 2, 1, " ".join(notes + ([f"No figures for {'; '.join(missing)}."] if missing else [])))
    row = _write_keys(ws, _KEY_ROW, names, datas)
    positions: dict[str, dict] = {}
    for name in _union([[t.name for t in d.tables] for d in datas if d]):
        per_run = [next((t for t in d.tables if t.name == name), None) if d else None
                   for d in datas]
        row, pos = _write_table(ws, row, name, names, per_run)
        if per_run[0] is not None:
            positions[name] = pos
    widest = max(3 + 2 * len(names), max((p["labels"] + 1 + max(p["values"].values(), default=0)
                                          for p in positions.values()), default=0))
    _charts(ws, view_id, positions, get_column_letter(widest + 2))
    width(ws, {"A": 34, "B": 14, "C": 14})
    ctx.sheets[view_id] = [ws.title]
    return True


def _why_missing(ctx, datas: list[ViewData | None]) -> list[str | None]:
    out = []
    for run, d in zip(ctx.runs, datas):
        if d is not None:
            out.append(None)
        elif run.reasons:
            from services.results_workbook.scenarios import _results_text
            out.append("not comparable with the base (" + _results_text(run) + ")")
        else:
            out.append("no results for this view")
    return out
