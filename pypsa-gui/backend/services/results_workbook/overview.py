"""
The Overview sheet: the run's investment-grade KPIs (owner decision RW-KPI-1).

One column per investment period and a Horizon column, from the optimisation:

* capex: overnight (`cost_breakdown.capex_lifetime`, the present value of the
  upfront investment; horizon only) and annualised (Σ over carriers of
  `economics_by_carrier.capex_meur`, which carries fixed O&M as the LP's
  periodized cost does);
* opex (Σ `opex_meur`) and its parts: fixed O&M (`cost_breakdown.fom`, inside
  the annualised capex), variable (`gen_cost_meur`), storage charging,
  curtailment and lost load;
* fuel: PyPSA's marginal cost does not separate fuel from variable O&M, so the
  row is `not_established` rather than a guessed split;
* total system cost (`cost_breakdown.total`, which reconciles with the LP
  objective) with a reconciliation row against capex + opex, energy served
  (load less lost load, on the same weighting as the costs) and the
  system LCOE as a live formula over those two cells;
* LCOE per carrier and period (`economics_by_carrier.lcoe_eur_per_mwh`).

Then the investment-case block, read from the stored report
(`finance.report.load_ic_report`): NPV, IRR (project and equity, pre and post
tax), payback, DSCR, financial LCOE and debt size. Without a report every row
is `not_established` with "run the investment case"; the export never runs it.
With one, the project post-tax cash flow is written below with live discount
formulas on a rate cell (the WACC) and the NPV cell is `=CF0+NPV(rate,CF1:CFn)`
(`finance.metrics.npv`'s convention: year 0 undiscounted).

Charts (native, from this sheet's cells): capex vs opex stacked by period, LCOE
by carrier by period, and the yearly cash flow with its cumulative discounted
sum when the case exists.
"""
from __future__ import annotations

import math
from dataclasses import dataclass, field
from typing import Any

from openpyxl.utils import get_column_letter

from services.results_workbook import charts
from services.results_workbook.cells import formula, heading, put, width

__all__ = ["CASE_ROWS", "OverviewFigures", "Line", "RUN_THE_CASE", "TITLE",
           "overview_figures", "write_overview"]

TITLE = "Overview"
RUN_THE_CASE = "run the investment case"
_HEADER_ROW = 4


@dataclass
class Line:
    key: str
    label: str
    unit: str
    total: float | None
    by_period: dict[str, float | None]
    source: str
    note: str | None = None


@dataclass
class OverviewFigures:
    periods: list[str]
    lines: list[Line]
    # carrier → (total, {period: value}) of the LCOE, EUR/MWh.
    lcoe_by_carrier: dict[str, tuple[float | None, dict[str, float | None]]]
    notes: list[str] = field(default_factory=list)

    def line(self, key: str) -> Line:
        return next(x for x in self.lines if x.key == key)


# ── figures ──────────────────────────────────────────────────────────────

def _finite(v: Any) -> float | None:
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return f if math.isfinite(f) else None


def _econ_sum(by_carrier: dict, key: str, periods: list[str]) -> tuple[float | None, dict]:
    total = 0.0
    pp = {p: 0.0 for p in periods}
    for v in by_carrier.values():
        cell = v.get(key) or {}
        t = _finite(cell.get("total"))
        if t is None:
            return None, {p: None for p in periods}
        total += t
        for p in periods:
            x = _finite((cell.get("by_period") or {}).get(p, 0.0))
            pp[p] = None if x is None or pp[p] is None else pp[p] + x
    return total, pp


def _energy_served(n, lost_load: dict | None, result_df) -> tuple[float | None, dict[str, float]]:
    """MWh of load served: Σ load less the lost-load capture, on the COST
    columns' weighting (snapshot objective weighting × the period's years), so
    the system LCOE divides like by like (GK B plan condition 3). A capture
    whose lost energy is unknown makes the figure unknown, never "demand
    minus 0"."""
    import pandas as pd

    from services.compare.support import _per_period_groupby
    from services.period_utils import snapshot_weights

    sns = n.snapshots
    is_multi = isinstance(sns, pd.MultiIndex)
    w = snapshot_weights(n, "objective")
    p = result_df(n, "loads_t", "p")
    if p is None or getattr(p, "empty", True):
        return None, {}
    served = p.reindex(sns).fillna(0.0).astype(float).sum(axis=1)
    ll = (lost_load or {}).get("lost_load_t")
    if isinstance(lost_load, dict) and ll is None and lost_load.get("lost_load_total_mwh", 0.0) is None:
        return None, {}
    if ll is not None and not getattr(ll, "empty", True):
        served = served - ll.reindex(sns).fillna(0.0).astype(float).sum(axis=1)
    weighted = served.mul(w, axis=0)
    return float(weighted.sum()), _per_period_groupby(weighted, sns, is_multi)


def overview_figures(n, cfg, lost_load: dict | None, result_df) -> OverviewFigures:
    """The Overview's numbers, each from the service the Results tabs read."""
    from services.results.cost_breakdown import compute_cost_breakdown
    from services.results.economics_by_carrier import compute_economics_by_carrier

    econ = compute_economics_by_carrier(n, cfg, lost_load, result_df=result_df)
    cb = compute_cost_breakdown(n, cfg) or {}
    notes: list[str] = []
    by_carrier = econ.get("by_carrier") or {}
    if "error" in econ or not econ.get("available"):
        notes.append("The per-carrier economics could not be computed for this run; "
                     "its rows are not_established.")
        by_carrier = {}
    periods = sorted({p for v in by_carrier.values()
                      for p in ((v.get("capex_meur") or {}).get("by_period") or {})}, key=int)
    if not periods:
        periods = [str(r["period"]) for r in cb.get("by_period") or []]
    econ_ok = bool(by_carrier)

    def econ_line(key, label, field_, note=None):
        if not econ_ok:
            return Line(key, label, "M EUR", None, {p: None for p in periods},
                        "economics_by_carrier", note)
        t, pp = _econ_sum(by_carrier, field_, periods)
        return Line(key, label, "M EUR", t, pp, f"economics_by_carrier.{field_} (Σ carriers)", note)

    cb_pp = {str(r["period"]): r for r in cb.get("by_period") or []}

    def cb_line(key, label, field_, note=None):
        if not cb:
            return Line(key, label, "M EUR", None, {p: None for p in periods},
                        "cost_breakdown", note)
        t = _finite(cb.get(field_))
        pp = {p: (None if p not in cb_pp or _finite(cb_pp[p].get(field_)) is None
                  else cb_pp[p][field_] / 1e6) for p in periods}
        return Line(key, label, "M EUR", None if t is None else t / 1e6, pp,
                    f"cost_breakdown.{field_}", note)

    overnight = _finite(cb.get("capex_lifetime")) if cb else None
    energy_mwh, energy_pp = _energy_served(n, lost_load, result_df)
    lines = [
        Line("capex_overnight", "Capex, overnight (present value of the upfront investment)",
             "M EUR", None if overnight is None else overnight / 1e6, {p: None for p in periods},
             "cost_breakdown.capex_lifetime",
             "horizon only: the upfront cost is not split by period"
             if overnight is not None else
             "an upfront cost could not be resolved for at least one component class"),
        econ_line("capex_annualised", "Capex, annualised (incl. fixed O&M)", "capex_meur"),
        econ_line("opex", "Opex", "opex_meur",
                  "variable + storage charging + curtailment + lost load"),
        econ_line("opex_variable", "  of which variable opex (marginal cost)", "gen_cost_meur"),
        econ_line("opex_storage_charging", "  of which storage charging", "storage_charge_cost_meur"),
        econ_line("opex_curtailment", "  of which curtailment", "curtailment_cost_meur"),
        econ_line("opex_lost_load", "  of which lost load", "lost_load_cost_meur"),
        cb_line("fixed_om", "Fixed O&M (already inside annualised capex, not in opex)", "fom"),
        Line("fuel", "Fuel", "M EUR", None, {p: None for p in periods}, "not separable",
             "PyPSA's marginal cost does not separate fuel from variable O&M"),
        cb_line("system_cost", "Total system cost", "total",
                "annualised capex + opex as the LP objective counts them"),
        Line("energy_served", "Energy served", "GWh",
             None if energy_mwh is None else energy_mwh / 1e3,
             {p: (energy_pp[p] / 1e3 if p in energy_pp else None) for p in periods},
             "load less lost load, on the cost columns' weighting"),
        # The value is a formula over the two cells above (write_overview).
        Line("system_lcoe", "System LCOE", "EUR/MWh", None, {p: None for p in periods},
             "total system cost / energy served (live formula)"),
    ]
    lcoe: dict[str, tuple[float | None, dict[str, float | None]]] = {}
    for carrier in sorted(by_carrier):
        cell = by_carrier[carrier].get("lcoe_eur_per_mwh") or {}
        disp = _finite((by_carrier[carrier].get("dispatch_gwh") or {}).get("total"))
        if not disp:
            continue  # a carrier that served no energy has no LCOE row
        lcoe[carrier] = (_finite(cell.get("total")),
                         {p: _finite((cell.get("by_period") or {}).get(p)) for p in periods})
    return OverviewFigures(periods, lines, lcoe, notes)


# ── the investment-case block ────────────────────────────────────────────

def _payload(report, section: str) -> dict:
    st = report.sections.get(section) if report is not None else None
    return (st.payload or {}) if st is not None else {}


#: (key, label, unit, reader(report) -> value)
CASE_ROWS: tuple[tuple[str, str, str, Any], ...] = (
    ("npv_project_post_tax", "NPV, project post-tax at the WACC", "EUR",
     lambda r: r.npv_at_wacc),
    ("npv_project_pre_tax", "NPV, project pre-tax at the WACC", "EUR",
     lambda r: _payload(r, "project").get("project_pre_tax_npv")),
    ("irr_project_pre_tax", "IRR, project pre-tax", "per unit", lambda r: r.project_irr_pre_tax),
    ("irr_project_post_tax", "IRR, project post-tax", "per unit", lambda r: r.project_irr_post_tax),
    ("irr_equity_pre_tax", "IRR, equity pre-tax", "per unit",
     lambda r: _payload(r, "project").get("equity_pre_tax_irr")),
    ("irr_equity_post_tax", "IRR, equity post-tax", "per unit",
     lambda r: _payload(r, "project").get("equity_post_tax_irr")),
    ("payback_years", "Payback (equity post-tax cash)", "years",
     lambda r: _payload(r, "project").get("payback_years")),
    ("min_dscr", "DSCR, minimum", "ratio", lambda r: r.min_dscr),
    ("avg_dscr", "DSCR, average", "ratio", lambda r: r.avg_dscr),
    ("lcoe_financial", "LCOE, financial (finance-consistent)", "EUR/MWh",
     lambda r: r.lcoe_finance_consistent_eur_per_mwh),
    ("debt_size", "Debt size", "EUR", lambda r: _payload(r, "debt").get("amount")),
)


def _cash_flow(report) -> tuple[list[int], list[float], float] | None:
    """(years, project post-tax cash, WACC) when all are established."""
    p = _payload(report, "project")
    years = p.get("years") or []
    cash = (p.get("cash") or {}).get("project_post_tax")
    rate = _finite(p.get("wacc_nominal"))
    if not years or not cash or rate is None or len(cash) != len(years):
        return None
    vals = [_finite(v) for v in cash]
    if any(v is None for v in vals):
        return None
    return [int(y) for y in years], vals, rate


# ── writing ──────────────────────────────────────────────────────────────

def _num(ws, row: int, col: int, value: float | None) -> bool:
    """A number, or an EMPTY cell (the reason goes in the row's note): a
    formula or a SUM never reads `not_established` text or a 0 for it."""
    if value is None:
        return False
    ws.cell(row=row, column=col).value = value
    return True


def _is_num(ws, row: int, col: int) -> bool:
    v = ws.cell(row, col).value
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _missing(note: str, cols: list[str]) -> str:
    return f"{note}; not_established: {', '.join(cols)}" if cols else note


def write_overview(ws, figs: OverviewFigures, report, case_state: dict | None = None
                   ) -> dict[str, Any]:
    """Write the sheet; returns the cell map the tests and About read.
    `case_state` is `investment_case_runner.staleness(...)` for the stored
    report: a stale case's figures are not shown beside this run's."""
    ws.title = TITLE
    periods = figs.periods
    headers = periods + ["Horizon"]
    horizon_col = 1 + len(headers)
    unit_col, src_col = horizon_col + 1, horizon_col + 2
    heading(ws, 1, ["Overview: investment-grade KPIs over the simulation horizon"])
    put(ws, 2, 1, "Costs in M EUR, undiscounted: each period column is the period's cost "
                  "(snapshot objective weighting × the period's years); Horizon is the sum of "
                  "the periods. Energy uses the same weighting, so each LCOE divides like by "
                  "like. An empty cell is not_established: the note on its row says why, and "
                  "no number stands in for it.")
    heading(ws, _HEADER_ROW, ["Metric"] + headers + ["Unit", "Source / note"])
    rows: dict[str, int] = {}
    r = _HEADER_ROW + 1
    for ln in figs.lines:
        rows[ln.key] = r
        put(ws, r, 1, ln.label)
        values = [ln.by_period.get(p) for p in periods] + [ln.total]
        absent = [h for j, (h, v) in enumerate(zip(headers, values), start=2)
                  if not _num(ws, r, j, v)]
        put(ws, r, unit_col, ln.unit)
        note = ln.source if not ln.note else f"{ln.source}: {ln.note}"
        put(ws, r, src_col, note if ln.key == "system_lcoe" else _missing(note, absent))
        r += 1

    def live(row: int, expr_of, inputs: tuple[int, ...], why: str, positive: int | None = None):
        """A formula per column over `inputs` rows, only where every input is a
        number (and `positive` > 0); elsewhere the cell stays empty."""
        absent = []
        for j, h in enumerate(headers, start=2):
            col = get_column_letter(j)
            ok = all(_is_num(ws, i, j) for i in inputs) and (
                positive is None or ws.cell(positive, j).value > 0)
            if ok:
                formula(ws, row, j, expr_of(col))
            else:
                absent.append(h)
        if absent:
            put(ws, row, src_col, f"{ws.cell(row, src_col).value}; not_established "
                                  f"({why}): {', '.join(absent)}")

    cost, energy = rows["system_cost"], rows["energy_served"]
    live(rows["system_lcoe"], lambda c: f"={c}{cost}*1000/{c}{energy}", (cost, energy),
         "cost or energy served not established", positive=energy)
    # Reconciliation: what the component rows leave unexplained in the total.
    rec = r
    rows["reconciliation"] = rec
    cap, op = rows["capex_annualised"], rows["opex"]
    put(ws, rec, 1, "Reconciliation: total system cost − (annualised capex + opex)")
    put(ws, rec, unit_col, "M EUR")
    put(ws, rec, src_col, "residual between cost_breakdown.total (n.statistics, the LP's "
                          "costs) and economics_by_carrier's capex + opex: opex there also "
                          "counts storage charging at bus prices, the curtailment penalty and "
                          "lost load at VOLL; a non-zero residual is these, not a lost cost")
    live(rec, lambda c: f"={c}{cost}-({c}{cap}+{c}{op})", (cost, cap, op),
         "an input row is not established")
    r += 1
    for note in figs.notes:
        put(ws, r, 1, note)
        r += 1

    # LCOE by carrier.
    r += 1
    put(ws, r, 1, "LCOE by carrier (EUR/MWh; source economics_by_carrier.lcoe_eur_per_mwh)")
    r += 1
    lcoe_header = r
    heading(ws, r, ["Carrier"] + headers + ["Unit", "Note"])
    for carrier, (total, pp) in figs.lcoe_by_carrier.items():
        r += 1
        put(ws, r, 1, carrier)
        values = [pp.get(p) for p in periods] + [total]
        absent = [h for j, (h, v) in enumerate(zip(headers, values), start=2)
                  if not _num(ws, r, j, v)]
        put(ws, r, unit_col, "EUR/MWh")
        if absent:
            put(ws, r, src_col, _missing("no LCOE where the carrier served no energy", absent))
    lcoe_last = r

    # Investment case.
    stale = bool(report is not None and case_state and case_state.get("stale"))
    r += 2
    put(ws, r, 1, "Investment case")
    if stale:
        changed = ", ".join(case_state.get("changed") or []) or case_state.get("reason") or "?"
        put(ws, r, 2, f"investment case from an earlier run (changed since: {changed}); its "
                      f"figures are not shown beside this run's. {RUN_THE_CASE} again.")
    r += 1
    heading(ws, r, ["KPI", "Value", "Unit", "Source / note"])
    case_rows: dict[str, int] = {}
    src = f"investment case report {report.case_id}" if report is not None else None
    for key, label, unit, read in CASE_ROWS:
        r += 1
        case_rows[key] = r
        put(ws, r, 1, label)
        put(ws, r, 3, unit)
        if report is None:
            put(ws, r, 4, f"not_established: {RUN_THE_CASE}")
            continue
        if stale:
            put(ws, r, 4, f"not_established: the stored case is from an earlier run; "
                          f"{RUN_THE_CASE} again")
            continue
        value = _finite(read(report))
        _num(ws, r, 2, value)
        put(ws, r, 4, src if value is not None
            else f"{src}: not_established by the case (see Case About)")

    cash: dict[str, Any] = {}
    flow = _cash_flow(report) if report is not None and not stale else None
    if report is not None and not stale and flow is None:
        r += 2
        put(ws, r, 1, "Cash flow: the case's project post-tax cash or WACC is not "
                      "established, so no live NPV formula is written.")
    if flow is not None:
        years, vals, rate = flow
        r += 2
        put(ws, r, 1, "Discount rate (WACC, nominal)")
        put(ws, r, 2, rate)
        put(ws, r, 3, "edit this cell: the discount factors and the NPV follow it")
        rate_cell = f"$B${r}"
        r += 1
        hdr = r
        heading(ws, r, ["Year", "t", "Project post-tax cash (EUR)", "Discount factor",
                        "Discounted cash flow", "Cumulative discounted"])
        first = r + 1
        for i, (y, v) in enumerate(zip(years, vals)):
            r = first + i
            ws.cell(r, 1).value = y
            ws.cell(r, 2).value = i
            ws.cell(r, 3).value = v
            formula(ws, r, 4, f"=1/(1+{rate_cell})^B{r}")
            formula(ws, r, 5, f"=C{r}*D{r}")
            formula(ws, r, 6, f"=E{r}" if i == 0 else f"=F{r - 1}+E{r}")
        last = r
        # finance.metrics.npv: index 0 undiscounted, end-of-year; Excel's NPV
        # discounts its first value one period, so CF0 stays outside it.
        npv = f"=C{first}" + (f"+NPV({rate_cell},C{first + 1}:C{last})" if last > first else "")
        r += 2
        put(ws, r, 1, "NPV (live formula)")
        formula(ws, r, 2, npv)
        npv_row = r
        r += 1
        put(ws, r, 1, "Check: sum of discounted cash flows")
        formula(ws, r, 2, f"=SUM(E{first}:E{last})")
        r += 1
        put(ws, r, 1, "NPV in the stored report")
        _num(ws, r, 2, _finite(report.npv_at_wacc))
        put(ws, r, 3, "static: does not move when the rate or a cash flow is edited")
        formula(ws, case_rows["npv_project_post_tax"], 2, f"=B{npv_row}")
        put(ws, case_rows["npv_project_post_tax"], 4,
            f"{src}: live formula over the cash flow below")
        cash = {"header": hdr, "first": first, "last": last, "npv_row": npv_row,
                "rate_cell": rate_cell}

    # Charts, to the right of the tables.
    anchor_col = get_column_letter(src_col + 2)
    data_last = 1 + len(periods) if periods else horizon_col
    data_first_col = 2 if periods else horizon_col
    made = [charts.bar(ws, title="Capex (annualised) vs opex by period", stacked=True,
                       from_rows=True, data=(1, cap, data_last, op),
                       cats=(data_first_col, _HEADER_ROW, data_last, _HEADER_ROW),
                       anchor=f"{anchor_col}{_HEADER_ROW}", y_title="M EUR")]
    if lcoe_last > lcoe_header:
        made.append(charts.bar(ws, title="LCOE by carrier by period",
                               data=(data_first_col, lcoe_header, data_last, lcoe_last),
                               cats=(1, lcoe_header + 1, 1, lcoe_last),
                               anchor=f"{anchor_col}{_HEADER_ROW + 20}", y_title="EUR/MWh"))
    if cash:
        made.append(charts.bar_line(ws, title="Project post-tax cash flow and cumulative NPV",
                                    bars=(3, cash["header"], 3, cash["last"]),
                                    line_data=(6, cash["header"], 6, cash["last"]),
                                    cats=(1, cash["first"], 1, cash["last"]),
                                    anchor=f"{anchor_col}{_HEADER_ROW + 40}",
                                    y_title="EUR", line_title="cumulative discounted, EUR"))
    width(ws, {"A": 52, get_column_letter(src_col): 60})
    return {"rows": rows, "horizon_col": horizon_col, "periods": periods,
            "lcoe": (lcoe_header, lcoe_last), "case_rows": case_rows, "cash": cash,
            "charts": len(made), "case_stale": stale}
