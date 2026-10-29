"""
The results workbook: one Excel file of a run's results (spec
`.scratch/results-workbook/spec.md`; owner decisions RW-CHARTS-1, RW-TABS-1,
RW-KPI-1, RW-PLAN-1).

`build_results_workbook(inputs) -> bytes` writes:

* `About` — project, run, input fingerprint and freshness, workflow, cards
  and views included (each view's sheet, or "no results"), units, the
  `not_established` legend;
* `Scenarios` — when the base has sensitivity runs, its key KPIs beside
  theirs with deltas against the base (`scenarios.py`, RW-SCEN-1);
* `Overview` — the investment-grade KPIs per period and over the horizon, and
  their native charts (`overview.py`);
* one sheet (or set of sheets) per included result view that has a builder in
  `registry.SHEET_BUILDERS`: the finance sheets for `investment` here, the
  other views in package 02.

The route gathers `WorkbookInputs` (the live network and solver state); this
module never solves, never runs the investment case and never reads router
state, so a test drives it with a network and a stored report.
"""
from __future__ import annotations

import io
import logging
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import UTC, datetime
from typing import Any

from services.results_workbook import registry, scenarios
from services.results_workbook.cells import NOT_ESTABLISHED, heading, put, width
from services.results_workbook.overview import (
    RUN_THE_CASE, overview_figures, write_overview,
)

logger = logging.getLogger(__name__)

__all__ = ["MAX_BYTES", "WorkbookInputs", "WorkbookTooLarge", "build_results_workbook"]

#: The export cap the other exports apply (`chat_tools.export_to_excel`).
MAX_BYTES = 25 * 1024 * 1024


class WorkbookTooLarge(ValueError):
    def __init__(self, size: int):
        self.size = size
        super().__init__(f"The workbook is {size / 1024 / 1024:.1f} MB, over the 25 MB "
                         "export limit. Include fewer result views.")


@dataclass
class WorkbookInputs:
    network: Any
    cfg: Any
    lost_load: dict | None
    result_df: Callable
    views: list[str]
    ic_report: Any = None
    project: str | None = None
    scope: dict | None = None
    template: dict | None = None
    # `workflow_tools.solve_freshness()`; `stored` when the caller asked for the
    # stored run although its inputs changed since the solve.
    freshness: dict = field(default_factory=dict)
    stored: bool = False
    # `investment_case_runner.staleness(...)` of the stored report against the
    # current inputs: a stale case is labelled, its figures not shown.
    case_state: dict | None = None
    # The latest completed run's summary (`run_evidence.list_project_runs`), if any.
    run: dict | None = None
    # The base's sensitivity runs (`scenarios.gather_scenarios`), if it has any.
    scenarios: scenarios.Scenarios | None = None
    # The settings the base's last solve used (the route's `_solved_settings`).
    solved_settings: Callable[[], dict | None] | None = None


@dataclass
class WorkbookContext:
    """What a sheet builder reads; builders record their sheet titles in `sheets`."""

    inputs: WorkbookInputs
    sheets: dict[str, list[str]] = field(default_factory=dict)
    # Every run, the base first (`scenarios.Run`).
    runs: list[scenarios.Run] = field(default_factory=list)

    def __getattr__(self, name):
        return getattr(self.inputs, name)


def _freshness_text(inp: WorkbookInputs) -> str:
    f = inp.freshness or {}
    if f.get("results_current"):
        return "current: the inputs are unchanged since the solve"
    if inp.stored:
        return ("stored run: the inputs changed since the solve (or the solve recorded no "
                "fingerprint); these are the stored results, exported on request")
    return NOT_ESTABLISHED


def _case_text(rep, state: dict | None) -> str:
    if rep is None:
        return f"none stored: {RUN_THE_CASE} (the export never runs it)"
    if state and state.get("stale"):
        changed = ", ".join(state.get("changed") or []) or state.get("reason") or "?"
        return (f"report {rep.case_id} is from an earlier run (changed since: {changed}); "
                f"the Overview does not show its figures and its Case sheets describe that "
                f"earlier run. {RUN_THE_CASE} again")
    return f"report {rep.case_id}, current with this run's inputs"


def _about(ws, inp: WorkbookInputs, statuses: list[tuple[str, str, str]], figs) -> None:
    from services import flow_scope
    from services.result_views import view as result_view

    ws.title = "About"
    n = inp.network
    meta = getattr(n, "meta", None) or {}
    run = inp.run or {}
    scope = inp.scope or {}
    rep = inp.ic_report
    proj = {}
    if rep is not None and rep.sections.get("project") is not None:
        proj = rep.sections["project"].payload or {}
    cards = [c.label for c in flow_scope.CARDS if c.id in (scope.get("groups") or ())]
    rows: list[tuple[str, Any]] = [
        ("Workbook", "Results workbook: the run's results, one sheet per included view, "
                     "with native Excel charts drawn from the sheets' own cells"),
        ("Project", inp.project or "(unsaved)"),
        ("Generated at", datetime.now(UTC).isoformat(timespec="seconds")),
        ("Run", run.get("run_id") or "not recorded"),
        ("Run finished at", run.get("finished_at") or "not recorded"),
        ("Results", _freshness_text(inp)),
        ("Input fingerprint", meta.get("assistant_solve_input_fingerprint")),
        ("Workflow", flow_scope.preset_label(inp.template)),
        ("Template", (inp.template or {}).get("id") or "none (a network without a site template)"),
        ("Cards included", ", ".join(cards) if cards else "none switched on"),
        ("Investment periods", ", ".join(figs.periods) if figs.periods else "one (no periods)"),
        ("Currency", proj.get("currency") or "EUR (the cost inputs' currency)"),
        ("Money year", proj.get("currency_year") if proj.get("currency_year") is not None
         else "as in the cost inputs (not stated by the run)"),
        ("Units", "costs M EUR (Overview), EUR (case); energy GWh; LCOE EUR/MWh; "
                  "IRR per unit; DSCR ratio"),
        ("not_established", "the figure could not be computed from this run; no number "
                            "stands in for it"),
        ("Investment case", _case_text(rep, inp.case_state)),
    ]
    heading(ws, 1, ["Item", "Value"])
    r = 2
    for k, v in rows:
        put(ws, r, 1, k)
        put(ws, r, 2, v)
        r += 1
    r += 1
    heading(ws, r, ["Result view", "Card", "In this workbook"])
    for view_id, card, status in statuses:
        r += 1
        put(ws, r, 1, result_view(view_id).label)
        put(ws, r, 2, card)
        put(ws, r, 3, status)
    width(ws, {"A": 28, "B": 90, "C": 60})


def build_results_workbook(inp: WorkbookInputs) -> bytes:
    from openpyxl import Workbook

    from services import flow_scope
    from services.result_views import view as result_view

    builders = registry.load_builders()
    wb = Workbook()
    about = wb.active
    overview = wb.create_sheet()
    figs = overview_figures(inp.network, inp.cfg, inp.lost_load, inp.result_df)
    write_overview(overview, figs, inp.ic_report, inp.case_state)
    base = scenarios.Run(name=inp.project or "Base", figs=figs, ic_report=inp.ic_report,
                         case_state=inp.case_state, network=inp.network, cfg=inp.cfg,
                         lost_load=inp.lost_load, result_df=inp.result_df,
                         solved_settings=inp.solved_settings)
    if inp.scenarios is not None:
        scenarios.write_scenarios(wb.create_sheet(index=1), base, inp.scenarios)
    ctx = WorkbookContext(inp, runs=[base, *(inp.scenarios.runs if inp.scenarios else [])])
    card_label = {c.id: c.label for c in flow_scope.CARDS}
    statuses: list[tuple[str, str, str]] = []
    for view_id in inp.views:
        v = result_view(view_id)
        build = builders.get(view_id)
        if build is None:
            status = "not exported to Excel yet"
        else:
            before = set(wb.sheetnames)
            try:
                wrote = build(wb, ctx)
            except Exception as exc:  # noqa: BLE001 — an error is reported, never "no results"
                logger.exception("results workbook: the %s sheet failed", view_id)
                for title in set(wb.sheetnames) - before:
                    del wb[title]
                ctx.sheets.pop(view_id, None)
                status = f"error: the sheet could not be written ({type(exc).__name__})"
            else:
                status = ("sheets: " + ", ".join(ctx.sheets.get(view_id) or []) if wrote
                          else registry.NO_RESULTS.get(view_id, "no results for this run"))
        statuses.append((view_id, card_label.get(v.card, v.card or ""), status))
    _about(about, inp, statuses, figs)
    buf = io.BytesIO()
    wb.save(buf)
    blob = buf.getvalue()
    if len(blob) > MAX_BYTES:
        raise WorkbookTooLarge(len(blob))
    return blob
