"""
Map an uploaded equipment table (CSV or XLSX) onto several assets of one class, cell by cell.

Data-to-model slice 1 (owner AI-567/567b, inside LC-16). `read_table` parses an owned
upload; `plan` maps its columns onto native fields and says exactly what an apply would
write, or why each cell is left out. It writes nothing. The edit tool
`apply_equipment_mapping` (`harness/edit_card.py`, `routers/network_equipment_table.py`)
takes its write from `plan` through the native component update, then `stamp` records
each written field's source cell.

Rules (server verdicts; the model proposes a mapping, never a value):

- **Bounded, inert parse.** At most `MAX_ROWS` data rows, `MAX_COLUMNS` columns and
  `MAX_CELL_CHARS` characters per cell; more refuses (`table_too_large`) instead of
  truncating. Formulas are never evaluated: an XLSX formula cell is read from the value
  the workbook cached, flagged `formula`, and one with no cached value is left out
  (`formula_without_value`). Macros, links and other sheets are not read. Before openpyxl
  opens an XLSX, its zip is refused (`table_unsafe`) when it would inflate past
  `MAX_XLSX_INFLATED` or by more than `MAX_XLSX_RATIO`, or when an XML part declares a
  DOCTYPE or entity (defusedxml is not installed, so openpyxl's parser is not hardened).
- **Exact lineage.** Every cell names the upload id, the full sha256 of the bytes read,
  the sheet, the sheet row and column and its A1 reference, with the raw value and unit.
- **Exact entities.** A row's name matches one asset of the class exactly, or the row is
  left out (`component_not_found`, `duplicate_entity`).
- **Native fields.** Unknown, derived, read-only and non-numeric fields are left out with
  the field schema's own reason; values go through the edit card's value checks. A
  time-series field, or a static-or-series field the asset holds as a series, is left out
  (`time_series_not_supported_here`); a static-or-series field is written as static.
- **Units from the schema only.** A column's unit is the field's stored unit, or its display
  unit converted by the schema's `display_scale`. Any other unit, or none, leaves the
  column out (`unit_mismatch`, `unit_missing`): nothing is assumed.
- **Money stated.** A cost column needs the currency and money year it is in. A currency
  other than the field's leaves it out (`currency_mismatch`); a money year other than the
  one the project reads costs in needs the user's `money_year` confirmation, as a Library
  reference does, and the value is copied unadjusted.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
import zipfile
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation

from fastapi import HTTPException

MAX_ROWS = 500
MAX_COLUMNS = 40
MAX_CELL_CHARS = 200
# Changed cells one card shows; a larger mapping is split by the user (rows or columns).
MAX_CARD_CELLS = 200
# Left-out cells one card or preview lists; the rest are counted.
MAX_LISTED_GAPS = 50

MAX_XLSX_INFLATED = 100 * 1024 * 1024
MAX_XLSX_RATIO = 100

CSV_SHEET = "(csv)"
CONFIRMATIONS = ("money_year",)

_NUMBER = re.compile(r"^[+-]?(\d+(\.\d*)?|\.\d+)([eE][+-]?\d+)?$")


# Every refusal kind below, written as literals so `tests/test_tool_error_kind_manifest.py`
# derives them (`_refuse` takes its kind as a variable).
_ERROR_KINDS = (
    {"error_kind": "table_too_large"}, {"error_kind": "table_unsafe"}, {"error_kind": "unreadable_table"},
    {"error_kind": "sheet_not_found"}, {"error_kind": "source_changed"},
    {"error_kind": "unsupported_table_format"}, {"error_kind": "duplicate_header"},
    {"error_kind": "column_not_found"}, {"error_kind": "empty_mapping"}, {"error_kind": "mapping_too_large"},
    {"error_kind": "table_confirmation_needed"}, {"error_kind": "nothing_to_change"},
    {"error_kind": "edit_preview_stale"},
)


def _refuse(kind: str, message: str, status: int = 422, **extra):
    raise HTTPException(status_code=status, detail={"error_kind": kind, "message": message, **extra})


def _column_letter(index: int) -> str:
    """0 → A, 25 → Z, 26 → AA."""
    out = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        out = chr(65 + rem) + out
    return out


# ── Read ───────────────────────────────────────────────────────────────────


def _check_size(rows: int, columns: int) -> None:
    if rows > MAX_ROWS or columns > MAX_COLUMNS:
        _refuse("table_too_large", f"The table has {rows} data rows and {columns} columns; one mapping reads at "
                f"most {MAX_ROWS} rows and {MAX_COLUMNS} columns. Split the file.",
                limits={"rows": MAX_ROWS, "columns": MAX_COLUMNS})


def _cell_text(value) -> None:
    if isinstance(value, str) and len(value) > MAX_CELL_CHARS:
        _refuse("table_too_large", f"A cell holds {len(value)} characters; a table cell is read up to "
                f"{MAX_CELL_CHARS}.", limits={"cell_chars": MAX_CELL_CHARS})


def _read_csv(blob: bytes) -> tuple[list[str], list[list[dict]]]:
    try:
        text = blob.decode("utf-8-sig")
    except UnicodeDecodeError:
        _refuse("unreadable_table", "The CSV is not UTF-8 text.")
    grid = [row for row in csv.reader(io.StringIO(text))]
    while grid and not any(c.strip() for c in grid[-1]):
        grid.pop()
    if not grid:
        _refuse("unreadable_table", "The table is empty.")
    width = max(len(r) for r in grid)
    _check_size(len(grid) - 1, width)
    for row in grid:
        for value in row:
            _cell_text(value)
    cells = [[{"raw": (row[i] if i < len(row) else "")} for i in range(width)] for row in grid]
    return [c["raw"] for c in cells[0]], cells[1:]


_DTD = re.compile(rb"<!(DOCTYPE|ENTITY)", re.IGNORECASE)
_DECLARED = re.compile(rb"""<\?xml[^>]*encoding\s*=\s*["']([^"']+)["']""", re.IGNORECASE)
_WIDE = (b"\xff\xfe", b"\xfe\xff", b"\x00<", b"<\x00", b"\x00\x00", b"\x4c\x6f\xa7\x94")


def _check_xlsx(blob: bytes) -> None:
    """Refuse a zip bomb, or a DOCTYPE/entity in ANY member, before openpyxl parses anything.
    Every member is scanned whatever its name (openpyxl follows the package's own part
    names), and a member in a wide or declared non-UTF-8 encoding refuses, because the
    byte scan could not see a DOCTYPE in it."""
    try:
        archive = zipfile.ZipFile(io.BytesIO(blob))
        members = archive.infolist()
    except (zipfile.BadZipFile, ValueError):
        _refuse("unreadable_table", "The workbook could not be read as XLSX: not a zip archive.")
    inflated = sum(m.file_size for m in members)
    packed = sum(m.compress_size for m in members)
    if inflated > MAX_XLSX_INFLATED or inflated > MAX_XLSX_RATIO * max(packed, 1):
        _refuse("table_unsafe", f"The workbook inflates to {inflated} bytes from {packed}; a table is read up to "
                f"{MAX_XLSX_INFLATED} bytes and {MAX_XLSX_RATIO}x compression.",
                limits={"inflated_bytes": MAX_XLSX_INFLATED, "ratio": MAX_XLSX_RATIO})
    try:
        for m in members:
            tail, head = b"", True
            with archive.open(m) as part:
                while chunk := part.read(1 << 16):
                    if head:
                        declared = _DECLARED.match(chunk.lstrip(b"\xef\xbb\xbf \t\r\n"))
                        if chunk.startswith(_WIDE) or (declared and declared.group(1).lower().replace(b"-", b"")
                                                       not in (b"utf8", b"usascii", b"ascii")):
                            _refuse("table_unsafe", f"{m.filename} is not UTF-8 text; the workbook is not read.")
                        head = False
                    if _DTD.search(tail + chunk):
                        _refuse("table_unsafe", f"{m.filename} declares a DOCTYPE or entity; the workbook is not read.")
                    tail = chunk[-16:]
    except (zipfile.BadZipFile, zipfile.LargeZipFile, EOFError, OSError, NotImplementedError) as exc:
        _refuse("unreadable_table", f"The workbook could not be read as XLSX: {type(exc).__name__}.")


def _read_xlsx(blob: bytes, sheet: str | None) -> tuple[str, list[str], list[str], list[list[dict]]]:
    import openpyxl

    _check_xlsx(blob)
    try:
        # Two inert reads of the same bytes: cached values, and which cells hold a formula.
        values = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True, keep_links=False)
        formulas = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:  # noqa: BLE001 — any parser failure is one refusal
        _refuse("unreadable_table", f"The workbook could not be read as XLSX: {type(exc).__name__}.")
    sheets = list(values.sheetnames)
    name = sheet if sheet is not None else sheets[0]
    if name not in sheets:
        _refuse("sheet_not_found", f"No sheet {name!r}; the workbook has {sheets}.", sheets=sheets)
    try:
        # Read-only sheets parse lazily: a malformed sheet fails here, not at load.
        vrows = [list(r) for r in values[name].iter_rows(values_only=True)]
        frows = [list(r) for r in formulas[name].iter_rows(values_only=True)]
    except Exception as exc:  # noqa: BLE001 — any parser failure is one refusal
        _refuse("unreadable_table", f"Sheet {name!r} could not be read: {type(exc).__name__}.")
    finally:
        values.close()
        formulas.close()
    def blank(r: int) -> bool:
        return all(v is None for v in vrows[r]) and all(f is None for f in (frows[r] if r < len(frows) else ()))

    while vrows and blank(len(vrows) - 1):
        vrows.pop()
    if not vrows:
        _refuse("unreadable_table", f"Sheet {name!r} is empty.")
    width = max(len(r) for r in vrows)
    _check_size(len(vrows) - 1, width)
    grid = []
    for r, row in enumerate(vrows):
        out = []
        for c in range(width):
            value = row[c] if c < len(row) else None
            raw_formula = frows[r][c] if r < len(frows) and c < len(frows[r]) else None
            _cell_text(value)
            cell = {"raw": value}
            if isinstance(raw_formula, str) and raw_formula.startswith("="):
                cell["formula"] = True
            out.append(cell)
        grid.append(out)
    headers = ["" if c["raw"] is None else str(c["raw"]) for c in grid[0]]
    return name, sheets, headers, grid[1:]


def read_table(project: str, upload_id: str, sheet: str | None = None) -> dict:
    """The upload as `{source, headers, rows}`; each data row `{row, cells}` keyed by header.
    `source` = `{upload_id, filename, sha256, sheet, sheets}`; `sha256` is of the bytes read."""
    from services import upload_service

    meta = upload_service.get_upload_meta(project, upload_id)
    blob = upload_service.get_upload_bytes(project, upload_id)
    sha = hashlib.sha256(blob).hexdigest()
    if meta.sha256 and meta.sha256 != sha:
        _refuse("source_changed", "The upload's bytes no longer match its recorded digest.")
    lower = meta.filename.lower()
    if lower.endswith(".csv") or meta.mime in {"text/csv", "text/plain"}:
        name, sheets = CSV_SHEET, [CSV_SHEET]
        headers, rows = _read_csv(blob)
    elif lower.endswith((".xlsx", ".xlsm")):
        name, sheets, headers, rows = _read_xlsx(blob, sheet)
    else:
        _refuse("unsupported_table_format", f"{meta.filename!r} is not a CSV or XLSX table.")
    clean = [h.strip() for h in headers]
    seen = set()
    for h in clean:
        if h and h in seen:
            _refuse("duplicate_header", f"Two columns are headed {h!r}; headers must be unique.", header=h)
        seen.add(h)
    out_rows = []
    for i, row in enumerate(rows):
        cells = {}
        for c, header in enumerate(clean):
            if not header:
                continue
            cell = dict(row[c])
            cell.update(row=i + 2, column=_column_letter(c), cell=f"{_column_letter(c)}{i + 2}")
            cells[header] = cell
        out_rows.append({"row": i + 2, "cells": cells})
    return {"source": {"upload_id": upload_id, "filename": meta.filename, "sha256": sha,
                       "sheet": name, "sheets": sheets},
            "headers": [h for h in clean if h], "rows": out_rows}


# ── Map ────────────────────────────────────────────────────────────────────


def _number(raw) -> Decimal | None:
    if isinstance(raw, bool):
        return None
    if isinstance(raw, (int, float)):
        try:
            d = Decimal(repr(raw)) if isinstance(raw, float) else Decimal(raw)
        except InvalidOperation:
            return None
        return d if d.is_finite() else None
    if isinstance(raw, str) and _NUMBER.match(raw.strip()):
        return Decimal(raw.strip())
    return None


def _as_value(d: Decimal):
    return int(d) if d == d.to_integral_value() and abs(d) < 2**53 else float(d)


def _cost_currency(unit: str) -> str | None:
    m = re.match(r"^([A-Z]{3})/", unit or "")
    return m.group(1) if m else None


def _column_verdict(field_name: str, field: dict | None, unit: str | None, money: dict) -> tuple[dict | None, dict | None]:
    """`(plan, None)` for a mappable column, else `(None, reason)`."""
    if field is None:
        return None, {"kind": "unknown_field", "message": f"No field {field_name!r}. Read get_component_schema."}
    if str(field.get("native_type") or "") == "series":
        return None, {"kind": "time_series_not_supported_here", "message": f"{field_name} is a time series; a "
                      "table maps one static value per asset."}
    if field.get("derived"):
        return None, {"kind": "derived_field", "message": f"{field_name} is derived from other fields."}
    if not field.get("editable"):
        return None, {"kind": "read_only_field", "message": f"{field_name} is {field.get('editability_reason')}."}
    kind = str(field.get("native_type") or "")
    stored = field.get("unit_storage") or field.get("native_unit") or ""
    if kind and not (kind.startswith(("float", "int")) or kind == "static or series"):
        return None, {"kind": "unsupported_field_type", "message": f"{field_name} takes {kind}, not a number."}
    if not kind and not stored:
        return None, {"kind": "unsupported_field_type", "message": f"{field_name} states no numeric type."}
    accepted = list(dict.fromkeys(u for u in (stored, field.get("unit_display")) if u))
    if not unit:
        return None, {"kind": "unit_missing", "message": f"State the unit of {field_name}: {' or '.join(accepted)}.",
                      "accepted_units": accepted}
    scale = None
    if unit == stored:
        scale = Decimal(1)
    elif unit == field.get("unit_display") and field.get("display_scale"):
        scale = Decimal(repr(float(field["display_scale"])))
    if scale is None:
        return None, {"kind": "unit_mismatch", "message": f"{field_name} is in {' or '.join(accepted)}, not {unit!r}; "
                      "nothing is converted outside the field's own units.", "accepted_units": accepted}
    currency = _cost_currency(stored)
    out = {"field": field_name, "label": field.get("label", field_name), "unit": unit, "stored_unit": stored,
           "scale": scale, "currency": currency, "money_year": None}
    if currency:
        if not money.get("currency") or not money.get("money_year"):
            return None, {"kind": "money_basis_missing", "message": f"{field_name} is a cost: state the table's "
                          "currency and money year."}
        if money["currency"] != currency:
            return None, {"kind": "currency_mismatch", "message": f"{field_name} is stored in {currency}; the table "
                          f"is in {money['currency']}. Nothing is converted."}
        out["money_year"] = int(money["money_year"])
    return out, None


def plan(n, table: dict, component_class: str, name_column: str, columns: dict, *,
         money: dict | None = None, confirm=(), base_year: int | None = None, fields: dict | None = None) -> dict:
    """What an apply would write: `{targets: {asset: {field: value}}, cells, unmapped, needs,
    missing, source}`. Reads `n`; writes nothing. `fields` is the class's field schema
    (`edit_card._fields`)."""
    from harness import edit_card
    from services.library import reference_apply as R

    money = dict(money or {})
    if name_column not in table["headers"]:
        _refuse("column_not_found", f"No column {name_column!r}; the table has {table['headers']}.")
    if not isinstance(columns, dict) or not columns:
        _refuse("empty_mapping", "Map at least one column onto a field.")
    attr = edit_card._component_attr(component_class)
    frame = getattr(n, attr)
    series = getattr(n, f"{attr}_t", {})
    unmapped: list[dict] = []
    plans: dict[str, dict] = {}
    for header, spec in columns.items():
        if header not in table["headers"]:
            _refuse("column_not_found", f"No column {header!r}; the table has {table['headers']}.")
        spec = spec or {}
        verdict, reason = _column_verdict(spec.get("field", ""), (fields or {}).get(spec.get("field")),
                                          (spec.get("unit") or "").strip() or None, money)
        if reason:
            unmapped.append({"column": header, "field": spec.get("field"), **reason})
        else:
            plans[header] = verdict
    fields_seen: dict[str, str] = {}
    for header, p in list(plans.items()):
        if p["field"] in fields_seen:
            unmapped.append({"column": header, "field": p["field"], "kind": "duplicate_field",
                             "message": f"Columns {fields_seen[p['field']]!r} and {header!r} both map to {p['field']}."})
            del plans[header]
        else:
            fields_seen[p["field"]] = header

    names = [str(r["cells"][name_column]["raw"] or "").strip() for r in table["rows"]]
    counts = {nm: names.count(nm) for nm in names}
    src = table["source"]
    targets: dict[str, dict] = {}
    cells: list[dict] = []
    for row, nm in zip(table["rows"], names):
        ref = row["cells"][name_column]
        if not nm:
            if any(row["cells"][h]["raw"] not in (None, "") for h in plans):
                unmapped.append({"row": row["row"], "cell": ref["cell"], "kind": "name_missing",
                                 "message": f"Row {row['row']} names no asset."})
            continue
        if counts[nm] > 1:
            unmapped.append({"row": row["row"], "cell": ref["cell"], "name": nm, "kind": "duplicate_entity",
                             "message": f"{nm!r} appears on {counts[nm]} rows; none of them is mapped."})
            continue
        if nm not in frame.index:
            unmapped.append({"row": row["row"], "cell": ref["cell"], "name": nm, "kind": "component_not_found",
                             "message": f"No {component_class} named {nm!r}."})
            continue
        for header, p in plans.items():
            cell = row["cells"][header]
            lineage = {"upload_id": src["upload_id"], "filename": src["filename"], "sha256": src["sha256"],
                       "sheet": src["sheet"],
                       "row": cell["row"], "column": cell["column"], "cell": cell["cell"], "header": header,
                       "raw": cell["raw"], "raw_unit": p["unit"], **({"formula": True} if cell.get("formula") else {})}
            if cell["raw"] in (None, "") and cell.get("formula"):
                unmapped.append({"name": nm, "field": p["field"], "kind": "formula_without_value", "source": lineage,
                                 "message": f"{cell['cell']} is a formula with no saved value; it is not evaluated."})
                continue
            if cell["raw"] in (None, ""):
                unmapped.append({"name": nm, "field": p["field"], "kind": "empty_cell", "source": lineage,
                                 "message": f"{cell['cell']} is empty; {nm}.{p['field']} keeps its value."})
                continue
            held = series.get(p["field"]) if hasattr(series, "get") else None
            if held is not None and nm in getattr(held, "columns", ()):
                unmapped.append({"name": nm, "field": p["field"], "kind": "time_series_not_supported_here",
                                 "source": lineage, "message": f"{nm}.{p['field']} is held as a time series; "
                                 "a static value from the table would not apply. It is left as it is."})
                continue
            d = _number(cell["raw"])
            if d is None:
                unmapped.append({"name": nm, "field": p["field"], "kind": "unparsed_number", "source": lineage,
                                 "message": f"{cell['cell']} does not hold a plain number."})
                continue
            value = _as_value(d / p["scale"])
            try:
                edit_card._check_value(component_class, p["field"], value, (fields or {}).get(p["field"]))
                edit_card._check_rating(nm, p["field"], value, (fields or {}).get(p["field"]))
            except HTTPException as exc:
                detail = exc.detail if isinstance(exc.detail, dict) else {"message": str(exc.detail)}
                unmapped.append({"name": nm, "field": p["field"], "kind": detail.get("error_kind", "invalid_edit_value"),
                                 "source": lineage, "message": detail.get("message", "")})
                continue
            targets.setdefault(nm, {})[p["field"]] = value
            cells.append({"name": nm, "field": p["field"], "label": p["label"], "value": value,
                          "unit": p["stored_unit"], "currency": p["currency"], "money_year": p["money_year"],
                          "source": lineage})

    needs = []
    years = sorted({p["money_year"] for p in plans.values() if p["money_year"] is not None})
    for year in years:
        if base_year is not None and base_year != year:
            currency = next(p["currency"] for p in plans.values() if p["money_year"] == year)
            text = R.money_disclosure(year, currency, base_year)
            needs.append({"confirmation": "money_year", "reference_year": year, "read_as_year": base_year,
                          "disclosure": text, "text": f"{text} Compared with the investment case's base year "
                          f"({base_year}). Unconfirmed, nothing is applied."})
    notes = []
    if years and base_year is None:
        notes.append(f"No investment case to compare money years with: costs are {', '.join(map(str, years))} "
                     "money, as the table states.")
    return {"component_class": component_class, "source": src, "name_column": name_column,
            "targets": targets, "cells": cells, "unmapped": unmapped, "needs": needs,
            "missing": [x["confirmation"] for x in needs if x["confirmation"] not in set(confirm or ())],
            "notes": notes}


def plan_active(args: dict) -> dict:
    """`plan` for the active project and network, from the tools' arguments."""
    from harness import edit_card
    from services import chat_tools
    from services.library import reference_apply as R
    from services.pypsa_service import PyPSAService

    table = read_table(chat_tools._require_active_project(), args.get("upload_id"), args.get("sheet"))
    want = args.get("source_sha256")
    if want is not None and want != table["source"]["sha256"]:
        _refuse("source_changed", "The table's bytes differ from the ones previewed. Preview it again.",
                previewed=want, now=table["source"]["sha256"])
    ctx = PyPSAService.get_active_context()
    config = ctx.solver_state.get("solver_config")
    klass = args.get("component_class")
    return plan(ctx.network, table, klass, args.get("name_column"), args.get("columns") or {},
                money={"currency": args.get("currency"), "money_year": args.get("money_year")},
                confirm=args.get("confirm") or (), base_year=R.case_base_year(ctx.network, config),
                fields=edit_card._fields(klass))


def require_confirmed(p: dict) -> None:
    if p["missing"]:
        _refuse("table_confirmation_needed", "Confirm before applying: " + " ".join(
            x["text"] for x in p["needs"] if x["confirmation"] in p["missing"]), missing=p["missing"])


def require_card_size(todo: dict[str, dict]) -> None:
    count = sum(len(v) for v in todo.values())
    if count > MAX_CARD_CELLS:
        _refuse("mapping_too_large", f"The mapping changes {count} cells; one card shows at most "
                f"{MAX_CARD_CELLS}. Map fewer rows or columns per apply.", limits={"card_cells": MAX_CARD_CELLS})


def changes(n, p: dict) -> dict[str, dict]:
    """The targets whose value differs from the network's now; a repeat apply has none."""
    from harness import edit_card

    frame = getattr(n, edit_card._component_attr(p["component_class"]))
    out: dict[str, dict] = {}
    for name, attrs in p["targets"].items():
        for field, value in attrs.items():
            old = edit_card._scalar(frame.at[name, field]) if field in frame.columns else None
            if not edit_card._same(old, value):
                out.setdefault(name, {})[field] = value
    return out


def source_id(lineage: dict) -> str:
    """A written field's provenance source: `upload:<id>#<sha16> <sheet>!<cell>`."""
    return f"upload:{lineage['upload_id']}#{lineage['sha256'][:16]} {lineage['sheet']}!{lineage['cell']}"


def stamp(n, p: dict, written: dict[str, dict], *, approved_by: str | None) -> None:
    """Record each written field's source cell (and money basis for a cost)."""
    from services.asset_schema import provenance

    when = datetime.now(UTC).isoformat(timespec="seconds")
    for cell in p["cells"]:
        if cell["field"] not in written.get(cell["name"], {}):
            continue
        provenance.record(n, p["component_class"], cell["name"], cell["field"], status="customised",
                          source=source_id(cell["source"]),
                          currency=cell["currency"], currency_year=cell["money_year"],
                          approved_by=approved_by, approved_at=when)


def apply(p: dict, card: dict, *, approved_by: str | None) -> dict:
    """Write exactly the cells the card's copy preview accepted, each asset through the
    native component update, in table order, and stamp them.

    Under the network's mutation lock for the whole loop (`update_component` re-enters
    it per asset), so no other write lands between assets. Before the first write the
    live network must still hold every typed row's `old` and the plan must change
    exactly the card's typed cells, else `edit_preview_stale` and nothing is written.
    Not one transaction: a native refusal part-way stops, and the result names what was
    written and what was not (one undo reverts the written part)."""
    from harness import edit_card
    from services import chat_tools
    from services.pypsa_service import PyPSAService

    shown = {(r["name"], r["field"]): r for r in card["rows"] if r["role"] == "typed"}
    with PyPSAService.get_lock():
        n = PyPSAService.get_network()
        todo = changes(n, p)
        if not todo:
            _refuse("nothing_to_change", "Every mapped value already equals the network's; nothing to apply.")
        frame = getattr(n, edit_card._component_attr(p["component_class"]))
        cells = {(name, field) for name, attrs in todo.items() for field in attrs}

        def now(name, field):
            return edit_card._scalar(frame.at[name, field]) if field in frame.columns else None

        if cells != set(shown) or any(not edit_card._same(now(*k), shown[k]["old"]) for k in cells):
            _refuse("edit_preview_stale", "The values changed since the card was shown; nothing was applied. "
                    "Propose the edit again.", 409)
        written: dict[str, dict] = {}
        failed = None
        for name, attrs in todo.items():
            try:
                chat_tools.update_component(p["component_class"], name, dict(attrs))
            except HTTPException as exc:
                detail = exc.detail if isinstance(exc.detail, dict) else {"message": str(exc.detail)}
                failed = {"name": name, "error_kind": detail.get("error_kind", "edit_refused"),
                          "message": detail.get("message", "")}
                break
            written[name] = attrs
        stamp(PyPSAService.get_network(), p, written, approved_by=approved_by)
    out = {"component_class": p["component_class"], "source": p["source"], "written": written,
           "unmapped": p["unmapped"], "notes": p["notes"]}
    if failed:
        out.update(partial=True, failed=failed, not_attempted=[k for k in todo if k not in written and k != failed["name"]])
    return out
