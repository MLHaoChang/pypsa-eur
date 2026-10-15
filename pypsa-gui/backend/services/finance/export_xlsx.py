"""
The investment-case workbook (IC P4 plan WP4.6d; the `asset_results` pattern).

`build_workbook(report, *, project=None) -> bytes` writes:
- `About` — the case, the assumptions hash, every pack with its hash, the CFADS
  definition, the WACC gate, the counterfactual's provenance, the completeness
  of each section and every flag the sections carry;
- `Summary` — the headline figures;
- one sheet per report section (`IC_REPORT_SECTIONS`), its status, note and
  payload (scalars as rows; lists of records as tables);
- `CashflowLines` — every `CashflowLine` with its provenance, and a
  `not_established` row per year for each line whose amount is unknown.

An unknown value (`None`) is written as an explicit `not_established` cell,
never an empty one or a 0 (ADR-0001). A text value that a spreadsheet would
read as a formula (starting with = + - @, a tab or a carriage return — a
user-named asset, party or contract) is stored as TEXT, never a formula
(formula injection guard).
"""
from __future__ import annotations

import io
import json
from datetime import datetime, timezone
from typing import Any

from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from models.finance import (
    IC_EXPORT_KEYS, IC_REPORT_SECTIONS, InvestmentCaseReport, export_investment_case,
)
from services.finance.debt import CFADS_DEFINITION

NOT_ESTABLISHED = "not_established"
_FORMULA_LEADS = ("=", "+", "-", "@", "\t", "\r")



def _put(ws, row: int, col: int, value: Any) -> None:
    """One cell: None → `not_established`; a formula-looking string → text."""
    cell = ws.cell(row=row, column=col)
    if value is None:
        cell.value = NOT_ESTABLISHED
        return
    if isinstance(value, bool):
        cell.value = value
        return
    if isinstance(value, (int, float)):
        cell.value = value
        return
    if not isinstance(value, str):
        value = json.dumps(value, default=str, sort_keys=True)
    # C0 control characters are illegal in xlsx cells (openpyxl raises): strip
    # them rather than fail the export (WP4.6b review B4).
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    cell.value = value
    if value.startswith(_FORMULA_LEADS):
        cell.data_type = "s"                 # never a formula


def _rows(ws, rows: list[list[Any]], start: int = 1) -> int:
    r = start
    for row in rows:
        for c, v in enumerate(row, start=1):
            _put(ws, r, c, v)
        r += 1
    return r


def _flatten(prefix: str, obj: Any, out: list[list[Any]], tables: list[tuple[str, list]]) -> None:
    if isinstance(obj, dict):
        for k, v in obj.items():
            _flatten(f"{prefix}.{k}" if prefix else str(k), v, out, tables)
    elif isinstance(obj, list) and obj and all(isinstance(x, dict) for x in obj):
        tables.append((prefix, obj))
    elif isinstance(obj, list):
        out.append([prefix, *obj] if obj else [prefix, "(empty)"])
    else:
        out.append([prefix, obj])


def _flags(report: InvestmentCaseReport) -> list[str]:
    found: set[str] = set()
    for st in report.sections.values():
        payload = st.payload or {}
        for key in ("flags", "reasons"):
            v = payload.get(key)
            if isinstance(v, list):
                found.update(str(x) for x in v)
            elif isinstance(v, dict):
                for xs in v.values():
                    if isinstance(xs, list):
                        found.update(str(x) for x in xs)
    return sorted(found)


#: Counterfactual flags under which the incremental figures compare against a
#: site that could not run (review of #283): said beside the Summary figures.
BASELINE_CAVEATS = (
    ("counterfactual_exceeds_connection",
     "The comparison site without the investment draws more than the connection can carry in "
     "some hours, and no value of lost load is set to price the excess: it could not run as "
     "modelled, so the NPV, IRR and payback compare against an impossible baseline."),
    ("counterfactual_shed_at_voll",
     "The comparison site without the investment cannot import all of its load: the load above "
     "its connection is shed and valued at the value of lost load, so the NPV, IRR and payback "
     "count that shed as a cost the investment avoids."),
    ("counterfactual_buys",
     "The comparison site without the investment buys the energy of its non-electric load (grey "
     "hydrogen, say) at the stated reference price: the NPV, IRR and payback count that purchase as "
     "a cost the investment avoids."),
    ("co2_price_period",
     "A CO2 price is applied: the comparison counts the certificate cost of what it buys (and of "
     "grid imports, on both sides, when a grid intensity is set) as cash the site would otherwise "
     "pay. It is additional to any carbon cost already in the tariff or the wholesale price."),
    ("counterfactual_exceeds_supply",
     "The comparison site without the investment needs more than its grid supply can deliver in "
     "some hours: it could not run as modelled, so the NPV, IRR and payback compare against an "
     "impossible baseline."),
)


def _counterfactual_summary(block: Any) -> str | None:
    """The About row: the counterfactual's provenance in one line (the full
    block is on the Project sheet). None → `not_established`."""
    if not isinstance(block, dict):
        return None
    if not block.get("present"):
        return f"none — {block.get('basis')}"
    sources = ", ".join(block.get("sources") or []) or "(no lines)"
    reasons = block.get("reasons") or []
    state = "not established: " + ", ".join(reasons) if reasons else "established"
    return (f"{block.get('basis')}; {block.get('n_lines')} lines from {sources}; {state}")


def _avoided_co2_total(block: Any) -> Any:
    """The counterfactual's avoided CO2 cost over the horizon (undiscounted);
    "none" without a CO2 price; None (`not_established`) when not established."""
    if not isinstance(block, dict) or not block.get("present"):
        return "none"
    v = block.get("avoided_co2_cost")
    if v is None:
        return None if block.get("co2_priced") else "none"
    return round(float(sum(v)), 2)


def _currency_year(payload: dict) -> Any:
    """A stated year; "not stated" when the inputs leave it None; None
    (`not_established`) when the report carries no basis (a refusal)."""
    if "currency_year" not in payload:
        return None
    v = payload.get("currency_year")
    return "not stated" if v is None else v


def build_workbook(report: InvestmentCaseReport, *, project: str | None = None) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    about = wb.active
    about.title = "About"
    project_payload = (report.sections.get("project").payload or {}) \
        if report.sections.get("project") else {}
    rows: list[list[Any]] = [
        ["Case", report.case_id],
        ["Project", project or "(unsaved)"],
        ["Reference design", report.reference_design_id],
        ["Assumptions hash", report.assumptions_hash],
        ["Generated at", datetime.now(timezone.utc).isoformat(timespec="seconds")],
        ["CFADS", CFADS_DEFINITION],
        ["Shed cost", "excluded (an opportunity cost on both sides, disclosed)"
         if report.excludes_shed_cost else "included"],
        ["WACC vs discount rate consistent", report.gates.wacc_vs_discount_rate_consistent],
        ["Value-flow conservation", report.gates.conservation_ok],
        ["Counterfactual", _counterfactual_summary(project_payload.get("counterfactual"))],
        # Of which avoided CO2 cost (H2e): the horizon's sum, undiscounted; lenders ask.
        ["Of which avoided CO2 cost", _avoided_co2_total(project_payload.get("counterfactual"))],
        # The price basis (GS Q6): a currency year not stated says so, never a guess.
        ["Price basis", project_payload.get("price_basis")],
        ["Currency year", _currency_year(project_payload)],
        ["Basis", project_payload.get("basis_statement")],
        # The terminal value's method (IC S0b plan S6); its terms are on the project sheet.
        ["Terminal value method", (project_payload.get("terminal_value") or {}).get("method")],
    ]
    # Campus equipment bought as owner capex (IC G2 plan G-10), one row each.
    for e in project_payload.get("extra_assets") or []:
        rows.append(["Extra asset", f"{e.get('name')}: {e.get('kind')}, {e.get('quantity')} "
                     f"{e.get('basis')}, overnight {e.get('overnight_cost')}, built "
                     f"{e.get('build_year')}, source {e.get('source')} ({e.get('source_hash')})"])
    for jur, h in sorted(report.packs.items()):
        rows.append([f"Pack {jur}", h])
    for name in IC_REPORT_SECTIONS:
        rows.append([f"Section {name}", report.completeness.get(name)])
    for f in _flags(report):
        rows.append(["Flag", f])
    _rows(about, rows)

    summary = wb.create_sheet("Summary")
    # The export view (the gate flags live under `gates`; reading the dump's
    # top level wrote them as not established — P4 gate driver finding).
    raw = export_investment_case(report)
    _rows(summary, [["Figure", "Value"]] + [
        [k, raw.get(k)] for k in IC_EXPORT_KEYS
        if k not in ("completeness", "sections", "cashflow_lines", "gates", "packs", "pipeline")
        and not isinstance(raw.get(k), (dict, list))])
    flags = set(_flags(report))
    # A flag matches its key alone or with a `:<value>` suffix (counterfactual_shed_at_voll:<MWh>).
    _rows(summary, [["Caveat", text] for f, text in BASELINE_CAVEATS
                    if any(x == f or x.startswith(f + ":") for x in flags)],
          summary.max_row + 1)
    # What the incremental figures are compared with, and the baseline's notes
    # (owner decision 2026-10-15), next to the figures.
    be = project_payload.get("island_breakeven")
    if isinstance(be, dict) and be.get("text"):
        _rows(summary, [["Island capability break-even", be["text"]]], summary.max_row + 1)
    cf = project_payload.get("counterfactual")
    if isinstance(cf, dict) and cf.get("basis"):
        _rows(summary, [["Compared with", cf["basis"]]]
              + [["Note", note] for note in cf.get("notes") or []], summary.max_row + 1)

    for name in IC_REPORT_SECTIONS:
        ws = wb.create_sheet(name[:31])
        st = report.sections.get(name)
        status = report.completeness.get(name)
        r = _rows(ws, [["Status", status], ["Note", st.note if st else None]])
        if st is None or st.payload is None:
            continue
        flat: list[list[Any]] = []
        tables: list[tuple[str, list]] = []
        _flatten("", st.payload, flat, tables)
        r = _rows(ws, [[]] + flat, r)
        for title, records in tables:
            cols = sorted({k for rec in records for k in rec})
            r = _rows(ws, [[], [title], cols] + [[rec.get(c) for c in cols] for rec in records], r)

    cf = wb.create_sheet("CashflowLines")
    header = ["year", "participant", "counterparty", "value_stream", "tariff_item", "asset",
              "amount", "source", "mode", "pack_hash", "seed", "source_id", "contract_id",
              "period"]
    body = [[ln.year, ln.participant, ln.counterparty, ln.value_stream, ln.tariff_item, ln.asset,
             ln.amount, ln.provenance.source, ln.provenance.mode, ln.provenance.pack_hash,
             ln.provenance.seed, ln.provenance.source_id, ln.provenance.contract_id,
             ln.provenance.period] for ln in report.cashflow_lines]
    # A line whose amount is not established has no CashflowLine: one explicit
    # `not_established` row per year it covers, so a sum over this sheet can
    # never read as complete (C12; P4 gate assessor B2).
    proj = report.sections.get("project")
    missing = ((proj.payload or {}).get("lines_not_established") or []) if proj else []
    for m in missing:
        for y in m.get("years") or [None]:
            body.append([y, report.sections["project"].payload.get("owner"), m.get("counterparty"),
                         m.get("value_stream"), None, None, None, m.get("source"), None, None,
                         None, m.get("key"), None, None])
    _rows(cf, [header] + body)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


__all__ = ["build_workbook", "NOT_ESTABLISHED", "CFADS_DEFINITION"]
