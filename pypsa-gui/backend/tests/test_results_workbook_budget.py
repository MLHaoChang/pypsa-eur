"""
Results workbook 02, slice R4: the cell budget and time coarsening
(`services/results_workbook/budget.py`).

GK B condition 4 and its R4 conditions
(`.scratch/results-workbook/issues/02-result-view-sheets.md`): hand values with
non-unit hours; MW is the weighted mean and energy Σ p·w; the hours are the
generators weighting, never the objective; a missing cell drops out of both
sums and an empty bucket is blank; non-datetime snapshots go in named blocks;
no bucket crosses a period; one resolution per sheet, named in the note and
About; keys from full data; charts on coarsened number cells; fallback (b);
the MAX_BYTES backstop drops a whole sheet and says so in About.
"""
from __future__ import annotations

import io
import math
from types import SimpleNamespace

import numpy as np
import pandas as pd
import pypsa
import pytest
from openpyxl import Workbook, load_workbook

import services.results_workbook as RW
from services.results_workbook import budget
from services.results_workbook.views import Key, Table, ViewData, write_view
from services.results_workbook.views_network import _ts_table
from services.solver_service import SolverConfig

DAY = ["2030-01-01T00:00:00", "2030-01-01T12:00:00", "2030-01-01T18:00:00"]
W = [2.0, 1.0, 3.0]


def _table(values, *, snaps=DAY, weights=W, periods=None, time="mean") -> Table:
    labels = {"Snapshot": snaps}
    if periods is not None:
        labels = {"Period": periods, **labels}
    frame = pd.concat([pd.DataFrame(labels), pd.DataFrame({"g": values})], axis=1)
    return Table("Generator dispatch", "MW", frame, len(labels), time=time, weights=weights)


def _cell(t: Table, row: int = 0):
    return t.frame["g"].iloc[row]


# ── the arithmetic, on hand values ───────────────────────────────────────

def test_a_daily_mw_value_is_the_hour_weighted_mean():
    t, name = budget.coarsen(_table([10.0, 40.0, 20.0]), 0)
    assert name == "daily" and len(t.frame) == 1
    # (10·2 + 40·1 + 20·3) / 6 = 20; the unweighted mean would be 23.33.
    assert _cell(t) == pytest.approx(20.0)
    assert _cell(t) != pytest.approx(np.mean([10.0, 40.0, 20.0]))
    assert t.frame["Snapshot"].iloc[0] == "2030-01-01"
    assert t.resolution == "daily, time-weighted means"


def test_energy_is_the_sum_of_value_times_hours():
    t, _ = budget.coarsen(_table([10.0, 40.0, 20.0], time="sum"), 0)
    assert _cell(t) == pytest.approx(120.0)
    assert _cell(t) != pytest.approx(70.0)  # an unweighted sum fails


def test_prices_and_the_state_of_charge_say_how_they_are_averaged():
    assert budget.coarsen(_table([1.0, 2.0, 3.0], time="price"), 0)[0].resolution == \
        "daily, time-weighted means, not load-weighted"
    assert budget.coarsen(_table([1.0, 2.0, 3.0], time="state"), 0)[0].resolution == \
        "daily, weighted means of the state"


def test_a_missing_cell_drops_out_of_both_sums():
    t, _ = budget.coarsen(_table([10.0, math.nan, 20.0]), 0)
    # (10·2 + 20·3) / 5 = 16, not (10·2 + 20·3) / 6.
    assert _cell(t) == pytest.approx(16.0)


@pytest.mark.parametrize("time", ["mean", "sum"])
def test_an_empty_or_zero_hour_bucket_is_blank_never_zero(time):
    t, _ = budget.coarsen(_table([math.nan] * 3, time=time), 0)
    assert math.isnan(_cell(t))
    t, _ = budget.coarsen(_table([5.0, 6.0, 7.0], weights=[0.0, 0.0, 0.0]), 0)
    assert math.isnan(_cell(t))


def test_a_day_straddling_a_period_boundary_is_two_rows():
    t, _ = budget.coarsen(_table([10.0, 40.0, 20.0], periods=[2030, 2030, 2040]), 0)
    assert list(zip(t.frame["Period"], t.frame["Snapshot"])) == [
        (2030, "2030-01-01"), (2040, "2030-01-01")]
    assert list(t.frame["g"]) == pytest.approx([20.0, 20.0])  # (20+40)/3, 60/3


def test_weeks_and_months_never_cross_a_period():
    snaps = ["2030-01-30T00:00:00", "2030-02-02T00:00:00", "2030-02-02T00:00:00"]
    t, name = budget.coarsen(_table([1.0, 2.0, 3.0], snaps=snaps, periods=[2030, 2030, 2040]), 2)
    assert name == "monthly"
    assert list(zip(t.frame["Period"], t.frame["Snapshot"])) == [
        (2030, "2030-01-01"), (2030, "2030-02-01"), (2040, "2030-02-01")]


def test_integer_snapshots_go_in_named_blocks_per_period():
    snaps = list(range(30)) * 2
    t, name = budget.coarsen(_table([1.0] * 60, snaps=snaps, weights=[1.0] * 60,
                                    periods=[2030] * 30 + [2040] * 30), 0)
    assert name == "blocks of 24 snapshots"
    assert list(zip(t.frame["Period"], t.frame["Snapshot"])) == [
        (2030, "snapshots 1–24"), (2030, "snapshots 25–48"),
        (2040, "snapshots 1–24"), (2040, "snapshots 25–48")]


def test_the_hours_are_the_generators_weighting_never_the_objective():
    n = pypsa.Network()
    n.set_snapshots(pd.to_datetime(DAY))
    n.snapshot_weightings["generators"] = W
    n.snapshot_weightings["objective"] = [w * 10 for w in W]
    payload = {"index": DAY, "columns": ["g"], "data": [[10.0], [40.0], [20.0]]}
    t = _ts_table(SimpleNamespace(network=n), "Generator dispatch", "MW", payload)
    assert list(t.weights) == W
    assert _cell(budget.coarsen(t, 0)[0]) == pytest.approx(20.0)


# ── the sheet ────────────────────────────────────────────────────────────

def _ctx(names, cells_left):
    return SimpleNamespace(runs=[SimpleNamespace(name=n, reasons=[]) for n in names],
                           cells_left=cells_left, resolution={}, sheets={})


def _data(scale: float) -> ViewData:
    t = _table([10.0 * scale, 40.0 * scale, 20.0 * scale])
    return ViewData(keys=[Key("Generation", "MWh", {"horizon": 120.0 * scale})], tables=[t])


def _rows(ws):
    return list(ws.iter_rows(values_only=True))


def test_a_sheet_over_its_budget_is_daily_for_every_run_and_says_so():
    wb, ctx = Workbook(), _ctx(["Base", "High"], cells_left=10)
    assert write_view(wb, ctx, "dispatch", "Dispatch", [_data(1.0), _data(2.0)])
    ws = wb["Dispatch"]
    text = " ".join(str(c) for r in _rows(ws) for c in r if c)
    assert "at daily resolution" in text and ctx.resolution["dispatch"] == "resolution daily"
    assert "Generator dispatch (MW; daily, time-weighted means)" in text
    data = [r for r in _rows(ws) if r[0] in ("Base", "High") and r[1] == "2030-01-01"]
    assert [(r[0], r[2]) for r in data] == [("Base", pytest.approx(20.0)),
                                            ("High", pytest.approx(40.0))]
    # The keys are the full data's.
    key = next(r for r in _rows(ws) if r[0] == "Generation")
    assert key[3:5] == (120.0, 240.0)


def test_fallback_b_keeps_the_base_table_and_the_scenarios_keys():
    wb, ctx = Workbook(), _ctx(["Base", "High"], cells_left=1)
    write_view(wb, ctx, "dispatch", "Dispatch", [_data(1.0), _data(2.0)])
    ws = wb["Dispatch"]
    text = " ".join(str(c) for r in _rows(ws) for c in r if c)
    assert "the other scenarios are at key-results level" in text
    assert ctx.resolution["dispatch"] == ("resolution monthly; other scenarios at key "
                                          "results only")
    assert any(r[0] == "Base" and r[1] == "2030-01-01" for r in _rows(ws))
    assert not any(r[0] == "High" and r[1] == "2030-01-01" for r in _rows(ws))
    assert next(r for r in _rows(ws) if r[0] == "Generation")[3:5] == (120.0, 240.0)


def test_a_sheet_within_budget_is_untouched():
    wb, ctx = Workbook(), _ctx(["Base"], cells_left=budget.CELL_BUDGET)
    write_view(wb, ctx, "dispatch", "Dispatch", [_data(1.0)])
    assert "dispatch" not in ctx.resolution
    assert sum(1 for r in _rows(wb["Dispatch"]) if r[0] == "Base") == 3


# ── the whole workbook ───────────────────────────────────────────────────

@pytest.fixture(scope="module")
def solved():
    from tests.test_results_workbook_views import _network
    return _network()


def _build(n, views, **kw):
    from tests.test_results_workbook_views import _rdf
    return RW.build_results_workbook(RW.WorkbookInputs(
        network=n, cfg=SolverConfig(), lost_load=None, result_df=_rdf, views=list(views),
        project="Base", scenarios=None, solved_settings=lambda: {"voll_eur_per_mwh": 3000.0},
        **kw))


def _about(wb) -> dict:
    from services.result_views import view
    by_label = {r[0]: r[2] for r in wb["About"].iter_rows(values_only=True) if r and r[0]}
    return {v: by_label.get(view(v).label) for v in ("dispatch", "storage", "prices")}


VIEWS = ["dispatch", "storage", "prices"]


def test_a_coarsened_workbook_keeps_its_keys_names_each_resolution_and_charts_number_cells(
        solved, monkeypatch):
    from tests.test_results_workbook_views import TITLES, _cells
    full = load_workbook(io.BytesIO(_build(solved, VIEWS)))
    monkeypatch.setattr(budget, "CELL_BUDGET", 1)
    coarse = load_workbook(io.BytesIO(_build(solved, VIEWS)))
    about = _about(coarse)
    for view in VIEWS:
        ws, ref = coarse[TITLES[view]], full[TITLES[view]]
        assert "resolution" in str(about[view]), about[view]
        # Keys come from full data: the key block is unchanged.
        keys = lambda s: [r[:5] for r in s.iter_rows(min_row=4, max_row=12, values_only=True)]  # noqa: E731
        assert keys(ws) == keys(ref)
        assert ws.max_row < ref.max_row
        for chart in ws._charts:
            for s in chart.series:
                if s.val is not None:
                    sheet, cells = _cells(s.val.numRef.f)
                    assert sheet == ws.title
                    assert all(isinstance(ws[c].value, (int, float)) for c in cells)
    soc = " ".join(str(c) for r in coarse[TITLES["storage"]].iter_rows(values_only=True)
                   for c in r if c)
    assert "State of charge (MWh; monthly, weighted means of the state)" in soc


def test_over_max_bytes_drops_a_whole_sheet_and_says_so_in_about(solved, monkeypatch):
    small = len(_build(solved, []))
    big = len(_build(solved, ["dispatch"]))
    assert big > small
    monkeypatch.setattr(RW, "MAX_BYTES", (small + big) // 2)
    wb = load_workbook(io.BytesIO(_build(solved, ["dispatch"])))
    from tests.test_results_workbook_views import TITLES
    assert TITLES["dispatch"] not in wb.sheetnames
    assert str(_about(wb)["dispatch"]).startswith("error: over the 25 MB export limit")
