"""
The price basis of a finance case (GS Q6, IC U1 follow-up item d) and the
storage LCOS (owner decision 6) on hand cases.

Basis: `FinanceInputs.currency_year` (None = not stated) and `price_basis`
("nominal" default | "real"); the report's project payload echoes both and
states the basis in words; a real basis with a non-zero escalation or
inflation is flagged `real_basis_with_escalation:<class>`; on a real basis the
real LCOE / LCOS equal the nominal ones (the cash is already real).

LCOS = (PV of the storage asset's capex + replacements + its own O&M + the cost
of the energy charged) / PV of the energy discharged, at `wacc_nominal`, index
0 (the financial close) undiscounted — per storage asset and in total. The
charging cost comes from the adapter (`StorageYear`): the import part escalates
with `tariff`, the surplus (on-site, e.g. PV) part — the export revenue the
site forwent — with `export`; both and the discharge scale with the asset's
degradation. Unknown is None + a reason, never 0 (plan C12).
"""
from __future__ import annotations

import io
from datetime import date

import numpy as np
import pytest
from pydantic import ValidationError

from models.finance import FinanceInputs
from services.finance.case import (
    AssetFinance, FinanceCase, StorageYear, Template, TemplateLine,
)
from services.finance.engine import run_case
from services.finance.tax import DepreciationClass, TaxLayer, sl_half_year

LAYER = (TaxLayer(name="corp", rate=0.25,
                  depreciation=(DepreciationClass("all", 1.0, sl_half_year(10)),)),)
OM_KEY = "asset:fom:StorageUnit:bess"


def _fin(**over) -> FinanceInputs:
    kw = dict(currency="EUR", financial_close=date(2030, 1, 1),
              cod_by_asset={"pv": date(2031, 1, 1), "bess": date(2031, 1, 1)},
              capex_phasing=[1.0], contingency_share=0.0, analysis_years=10,
              escalation={"opex": 0.0, "tariff": 0.02, "export": 0.01, "capex": 0.0},
              degradation_by_asset={"pv": 0.0, "bess": 0.0},
              tax_losses="offset_other_income", financing_fee_tax="not_deducted",
              wacc_nominal=0.07, cost_of_equity=0.09, inflation=0.02)
    kw.update(over)
    return FinanceInputs(**kw)


def _case(fin=None, *, storage=None, assets=None, lines=None) -> FinanceCase:
    storage = {"bess": StorageYear(discharge_mwh=900.0, charge_mwh=1000.0,
                                   charge_import_cost=40_000.0, charge_surplus_cost=6_000.0,
                                   om_keys=(OM_KEY,))} if storage is None else storage
    lines = lines if lines is not None else (
        TemplateLine("bill", "energy_import", 150_000.0, "tariff"),
        TemplateLine(OM_KEY, "fom", -2_000.0, "opex", source="asset"))
    assets = assets or (AssetFinance("pv", "Generator", 100_000.0, 30.0, carrier="solar"),
                        AssetFinance("bess", "StorageUnit", 500_000.0, 15.0, carrier="battery"))
    return FinanceCase(inputs=fin or _fin(), owner="o", base_year=2031, cod=date(2031, 1, 1),
                       templates=(Template(first_year=2031, lines=tuple(lines),
                                           energy_mwh={"pv": 4_000.0}, storage=storage),),
                       assets=assets)


def _hand_lcos(*, capex=500_000.0, om=2_000.0, imp=40_000.0, sur=6_000.0, dis=900.0,
               r=0.07, g_t=0.02, g_e=0.01, deg=0.0, years=10, real_infl=None,
               replacement=None) -> float:
    num = capex
    den = 0.0
    for k in range(1, years + 1):                 # axis index k = operating year k
        f = (1.0 - deg) ** (k - 1)
        d = (1.0 + r) ** k
        charge = (imp * (1.0 + g_t) ** (k - 1) + sur * (1.0 + g_e) ** (k - 1)) * f
        num += (om + charge) / d
        if replacement and replacement[0] == k:
            num += replacement[1] / d
        rr = r if real_infl is None else (1.0 + r) / (1.0 + real_infl) - 1.0
        den += dis * f / (1.0 + rr) ** k
    return num / den


# ── the price basis (item d) ─────────────────────────────────────────────────


def test_the_basis_fields_default_to_nominal_and_not_stated():
    fin = _fin()
    assert fin.currency_year is None and fin.price_basis == "nominal"
    assert _fin(currency_year=2020, price_basis="real").price_basis == "real"
    with pytest.raises(ValidationError):
        _fin(price_basis="constant")
    with pytest.raises(ValidationError):
        _fin(currency_yaer=2020)                    # extra="forbid" still holds


def test_a_real_basis_with_escalation_or_inflation_is_flagged():
    r = run_case(_case(_fin(price_basis="real", currency_year=2020)), layers=LAYER)
    assert {"real_basis_with_escalation:tariff", "real_basis_with_escalation:export",
            "real_basis_with_escalation:inflation"} <= set(r.flags)
    assert "real_basis_with_escalation:opex" not in r.flags         # a 0 rate is real
    clean = _fin(price_basis="real", currency_year=2020, inflation=None,
                 escalation={c: 0.0 for c in ("opex", "tariff", "export", "capex")})
    r2 = run_case(_case(clean), layers=LAYER)
    assert not any(f.startswith("real_basis_with_escalation") for f in r2.flags)
    # The cash is already real: the real LCOE and LCOS are the nominal ones.
    m = r2.metrics
    assert m["lcoe_real_per_mwh"] == m["lcoe_nominal_per_mwh"] is not None
    assert m["lcos_real_per_mwh"] == m["lcos_nominal_per_mwh"] is not None
    assert "lcoe_real_equals_nominal:real_basis" in r2.flags
    # A nominal basis never carries the flag, whatever its rates.
    r3 = run_case(_case(), layers=LAYER)
    assert not any(f.startswith("real_basis_with_escalation") for f in r3.flags)


@pytest.mark.parametrize("over,words", [
    ({"price_basis": "real", "currency_year": 2020}, "real basis, 2020 EUR"),
    ({}, "nominal basis, EUR (currency year not stated)"),
    ({"currency": "USD", "currency_year": 2024}, "nominal basis, 2024 USD"),
])
def test_the_report_echoes_the_basis_and_states_it_in_words(over, words):
    from services.finance.report import assemble_finance_sections

    case = _case(_fin(**over))
    rep = assemble_finance_sections(run_case(case, layers=LAYER), case)
    p = rep.sections["project"].payload
    assert p["price_basis"] == case.inputs.price_basis
    assert p["currency_year"] == case.inputs.currency_year
    assert p["basis_statement"] == words


def test_the_workbook_about_sheet_carries_the_basis():
    from openpyxl import load_workbook

    from services.finance.export_xlsx import build_workbook
    from services.finance.report import assemble_finance_sections

    for over, basis, year in (({"price_basis": "real", "currency_year": 2020}, "real", 2020),
                              ({}, "nominal", "not stated")):
        case = _case(_fin(**over))
        rep = assemble_finance_sections(run_case(case, layers=LAYER), case)
        ws = load_workbook(io.BytesIO(build_workbook(rep)))["About"]
        rows = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, ws.max_row + 1)}
        assert rows["Price basis"] == basis
        assert rows["Currency year"] == year


# ── storage LCOS (owner decision 6) ──────────────────────────────────────────


def test_lcos_on_a_hand_case():
    """BESS: 500,000 capex at close; 2,000 O&M a year; it charges 1,000 MWh a
    year — 800 MWh imported for 40,000 (tariff +2 %/yr) and 200 MWh of PV
    surplus that would have earned 6,000 of export (+1 %/yr) — and discharges
    900 MWh; WACC 7 %, 10 years: 136.39 per MWh."""
    r = run_case(_case(), layers=LAYER)
    want = _hand_lcos()
    assert want == pytest.approx(136.3903, abs=1e-4)
    got = r.lcos["assets"]["bess"]
    assert got["lcos_nominal_per_mwh"] == pytest.approx(want, rel=1e-12)
    assert r.metrics["lcos_nominal_per_mwh"] == pytest.approx(want, rel=1e-12)
    d = (1.07 ** np.arange(11))
    assert got["pv_capex"] == pytest.approx(500_000.0)
    assert got["pv_om"] == pytest.approx(float(np.sum(2_000.0 / d[1:])))
    assert got["pv_discharge_mwh"] == pytest.approx(float(np.sum(900.0 / d[1:])))
    assert got["pv_charging_cost"] == pytest.approx(
        float(sum((40_000 * 1.02 ** (k - 1) + 6_000 * 1.01 ** (k - 1)) / 1.07 ** k
                  for k in range(1, 11))))
    # Real: the energy discounted at the real rate (1.07 / 1.02 − 1), as the LCOE.
    assert r.metrics["lcos_real_per_mwh"] == pytest.approx(_hand_lcos(real_infl=0.02), rel=1e-12)
    assert r.lcos["basis"].startswith("(PV of the storage asset's capex")
    assert r.lcos["reasons"] == []
    # The LCOE is not touched by the storage (generators only).
    no_storage = run_case(_case(storage={}), layers=LAYER)
    assert r.metrics["lcoe_nominal_per_mwh"] == no_storage.metrics["lcoe_nominal_per_mwh"]


def test_lcos_with_degradation_a_replacement_and_contingency():
    fin = _fin(degradation_by_asset={"pv": 0.0, "bess": 0.02}, contingency_share=0.1,
               replacement_capex=[(2036, "bess", 50_000.0), (2036, "pv", 9_999.0)])
    r = run_case(_case(fin), layers=LAYER)
    want = _hand_lcos(capex=550_000.0, deg=0.02, replacement=(6, 50_000.0))
    assert r.metrics["lcos_nominal_per_mwh"] == pytest.approx(want, rel=1e-12)


def test_lcos_in_total_over_two_storage_assets():
    storage = {"bess": StorageYear(900.0, 1000.0, 40_000.0, 6_000.0, om_keys=(OM_KEY,)),
               "bess2": StorageYear(450.0, 500.0, 20_000.0, 0.0)}
    assets = (AssetFinance("pv", "Generator", 100_000.0, 30.0),
              AssetFinance("bess", "StorageUnit", 500_000.0, 15.0),
              AssetFinance("bess2", "StorageUnit", 200_000.0, 15.0))
    fin = _fin(degradation_by_asset={"pv": 0.0, "bess": 0.0, "bess2": 0.0},
               cod_by_asset={})
    r = run_case(_case(fin, storage=storage, assets=assets), layers=LAYER)
    a, b = r.lcos["assets"]["bess"], r.lcos["assets"]["bess2"]
    assert b["lcos_nominal_per_mwh"] == pytest.approx(
        _hand_lcos(capex=200_000.0, om=0.0, imp=20_000.0, sur=0.0, dis=450.0), rel=1e-12)
    total = (a["pv_capex"] + a["pv_om"] + a["pv_charging_cost"] + b["pv_capex"] + b["pv_om"]
             + b["pv_charging_cost"]) / (a["pv_discharge_mwh"] + b["pv_discharge_mwh"])
    assert r.metrics["lcos_nominal_per_mwh"] == pytest.approx(total, rel=1e-12)


@pytest.mark.parametrize("make,reason", [
    (lambda: _case(_fin(wacc_nominal=None)), "wacc_nominal_missing"),
    (lambda: _case(_fin(degradation_by_asset={"pv": 0.0})), "degradation_missing:bess"),
    (lambda: _case(storage={"bess": StorageYear(900.0, 1000.0, None, 6_000.0)}),
     "charging_cost_not_established:bess"),
    (lambda: _case(storage={"bess": StorageYear(900.0, 1000.0, 40_000.0, None)}),
     "charging_cost_not_established:bess"),
    (lambda: _case(storage={"bess": StorageYear(0.0, 0.0, 0.0, 0.0)}), "no_discharge:bess"),
    (lambda: _case(_fin(contingency_share=None)), "capex_not_established"),
    (lambda: _case(lines=(TemplateLine("bill", "energy_import", 150_000.0, "tariff"),
                          TemplateLine(OM_KEY, "fom", None, "opex", source="asset"))),
     f"om_not_established:{OM_KEY}"),
    (lambda: _case(_fin(escalation={"opex": 0.0, "export": 0.0, "capex": 0.0})),
     "escalation_missing:tariff"),
])
def test_lcos_not_established_is_none_with_a_reason(make, reason):
    r = run_case(make(), layers=LAYER)
    assert r.metrics["lcos_nominal_per_mwh"] is None
    assert r.metrics["lcos_real_per_mwh"] is None
    assert reason in r.lcos["reasons"]
    assert f"lcos_not_established:{reason}" in r.flags


def test_the_real_lcos_without_inflation_is_none_with_its_reason():
    """Review B round 1: the reason is in `lcos.reasons` too (the headline reads
    it), while the nominal LCOS stays established."""
    r = run_case(_case(_fin(inflation=None)), layers=LAYER)
    assert r.metrics["lcos_nominal_per_mwh"] == pytest.approx(_hand_lcos(), rel=1e-12)
    assert r.metrics["lcos_real_per_mwh"] is None
    assert r.lcos["reasons"] == ["real_not_established:inflation_missing"]
    assert r.lcos["assets"]["bess"]["reasons"] == []
    assert "lcos_real_not_established:inflation_missing" in r.flags
    assert not any(f.startswith("lcos_not_established") for f in r.flags)


def test_no_storage_is_none_and_adds_no_flag():
    base = run_case(_case(storage={}), layers=LAYER)
    assert base.metrics["lcos_nominal_per_mwh"] is None
    assert base.lcos["reasons"] == ["no_storage"]
    assert not any(f.startswith("lcos") for f in base.flags)


def test_the_report_carries_the_lcos():
    from services.finance.report import assemble_finance_sections

    case = _case()
    r = run_case(case, layers=LAYER)
    p = assemble_finance_sections(r, case).sections["project"].payload
    assert p["lcos_nominal_per_mwh"] == pytest.approx(_hand_lcos(), rel=1e-12)
    assert p["lcos_real_per_mwh"] == pytest.approx(_hand_lcos(real_infl=0.02), rel=1e-12)
    block = p["lcos"]
    assert block["assets"]["bess"]["lcos_nominal_per_mwh"] == pytest.approx(_hand_lcos(),
                                                                          rel=1e-12)
    assert "charging" in block["charging_basis"]
    assert block["reasons"] == []
