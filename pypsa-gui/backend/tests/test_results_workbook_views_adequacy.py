"""
Results workbook 02, slice R3: the Adequacy and FMEA sheets
(`services/results_workbook/views_adequacy.py`).

GK B conditions (`.scratch/results-workbook/issues/02-result-view-sheets.md`):
1 the lifted `/adequacy` body changes nothing; 2 keys equal the route and each
scenario reads its own state and network; 3 a period-2 figure (Adequacy) or
the no-period note (FMEA); 5 no results vs an error; 6 FMEA's criticality and
severity are labelled by each row's own basis (per year for class A/B since
LC-30, per study horizon for a row saved before it, per year for C), the
horizon's length is given, and
no key compares two bases; 9 charts read number cells on their own sheet.
"""
from __future__ import annotations

import io
import json

import numpy as np
import pandas as pd
import pypsa
import pytest
from openpyxl import load_workbook

from models.adequacy import (
    AdequacyReport, CostBlock, EnergyBlock, InputsBlock, MetricsBlock, PeriodTarget,
    ReserveMarginBlock, ReserveMarginPeriod, SystemTarget, TargetBlock, VollBlock, ZoneTarget,
)
from services import results_workbook as RW
from services.adequacy.copt_endpoint import build_fmea_modes_payload
from services.adequacy.report import adequacy_payload
from services.pypsa_service import PyPSAService
from services.results_workbook import scenarios as S
from services.results_workbook.views import HORIZON
from services.results_workbook.views_adequacy import (
    CRITICALITY, CRITICALITY_PER_YEAR, PER_HORIZON, SEVERITY, SEVERITY_PER_EVENT,
)
from services.results_workbook.views_grid import NO_PERIODS
from services.solver_service import SolverConfig

ANNUALISED = "per year (annualised)"
LEGACY = f"{PER_HORIZON} (saved before annualising)"

TITLES = {"adequacy": "Adequacy", "fmea": "FMEA"}
P2 = "2040"
CFG = SolverConfig(voll=3000.0)


# ── fixtures ─────────────────────────────────────────────────────────────

def _network(p_nom: float = 90.0) -> pypsa.Network:
    """Two generators with outage data: COPT gives each a class-A mode."""
    n = pypsa.Network()
    n.set_snapshots(pd.date_range("2030-01-01", periods=2, freq="h"))
    n.add("Carrier", "gas")
    n.add("Bus", "b", carrier="AC", country="AA")
    n.add("Load", "l", bus="b", p_set=100.0)
    for name, p in (("g1", p_nom), ("g2", 40.0)):
        n.add("Generator", name, bus="b", carrier="gas", p_nom=p, marginal_cost=10.0,
              outage_rate_value=0.05, outage_rate_basis="EFORd", mttr_hours=50.0)
    return n


def _report(scale: float = 1.0) -> dict:
    """An AdequacyReport over two periods (2030, 2040) and one zone."""
    periods = [PeriodTarget(period="2030", cap_mwh=10.0, achieved_ens_mwh=4.0 * scale),
               PeriodTarget(period=P2, cap_mwh=20.0, achieved_ens_mwh=7.0 * scale,
                            binding=scale > 1)]
    ach = sum(p.achieved_ens_mwh for p in periods)
    margin = [ReserveMarginPeriod(period=str(y), peak_mw=100.0, required_mw=115.0,
                                  firm_mw=f * scale, margin_achieved=f * scale / 100.0 - 1.0,
                                  met=f * scale >= 115.0)
              for y, f in ((2030, 120.0), (2040, 118.0))]
    return AdequacyReport(
        engine="lp_proxy", fidelity="deterministic_scenario",
        target=TargetBlock(basis="energy",
                           system=SystemTarget(cap_mwh=30.0, achieved_ens_mwh=ach,
                                               achieved_shed_hours=3.0, by_period=periods),
                           zones=[ZoneTarget(zone="AA", cap_mwh=30.0, achieved_ens_mwh=ach)],
                           binding="system_cap", zone_field_populated=True,
                           energy_target_set=True),
        reserve_margin=ReserveMarginBlock(margin=0.15, horizon_wide=True, by_period=margin),
        metrics=MetricsBlock(ens_mwh=ach, shed_hours=3.0 * scale, time_basis="hours_per_horizon"),
        cost=CostBlock(total_system_cost_eur=5e6 * scale, period_basis="npv_multi_period"),
        inputs=InputsBlock(weather_years=["modelled"], voll=VollBlock(default_eur_per_mwh=3000.0),
                           assumptions_hash="h"),
        energy=EnergyBlock(involuntary_mwh=ach, demand_response_mwh=0.0),
    ).model_dump()


@pytest.fixture
def active(install_network):
    """The base: the active network and solver state, carrying a report."""
    n = _network()
    install_network(n)
    state = PyPSAService.get_solver_state()
    state.update({"status": "completed", "solver_config": CFG, "adequacy_report": _report()})
    return n


def _scenario(scale: float = 1.5, *, report: bool = True) -> S.Run:
    """A scenario run: its own network, config and saved context."""
    n = _network(p_nom=60.0)
    ctx = PyPSAService.build_context()
    ctx.network = n
    ctx.solver_state.update({"status": "completed", "solver_config": CFG,
                             **({"adequacy_report": _report(scale)} if report else {})})
    return S.Run(name="High load", project_id="scen", network=n, cfg=CFG, context=ctx,
                 result_df=lambda *a, **k: None, solved_settings=lambda: None)


def _build(n, views, *, scenarios=None, solved_settings=lambda: None) -> bytes:
    return RW.build_results_workbook(RW.WorkbookInputs(
        network=n, cfg=CFG, lost_load=None, result_df=lambda *a, **k: None, views=list(views),
        project="Base", scenarios=scenarios, solved_settings=solved_settings,
        project_id="base"))


def _wb(blob):
    return load_workbook(io.BytesIO(blob))


def _key(ws, label: str, period: str, col: int = 4):
    for row in ws.iter_rows(min_row=4):
        if row[0].value == label and str(row[2].value) == period:
            return row[col - 1].value
    raise AssertionError(f"no key {label!r} / {period!r} on {ws.title}")


def _about(wb) -> dict:
    from services.result_views import view
    by_label = {row[0]: row[2] for row in wb["About"].iter_rows(values_only=True) if row and row[0]}
    return {v: by_label.get(view(v).label) for v in TITLES}


def _text(ws) -> str:
    return " ".join(str(c) for row in ws.iter_rows(values_only=True) for c in row if c)


def _dump(x) -> str:
    return json.dumps(x, sort_keys=True)


# ── condition 1: the lifted /adequacy body changes nothing ────────────────

def test_adequacy_route_json_equals_the_lifted_function_in_every_state(
        client, install_network, session_ctx):
    from routers.results import _solved_settings as route_settings
    install_network(_network())
    assert client.get("/api/results/adequacy").status_code == 204
    ctx = session_ctx(client)       # the client's own context: what the route reads
    state = ctx.solver_state
    _solved_settings = lambda: route_settings(state, ctx.network)  # noqa: E731
    # Nothing solved: 204, and the function says None.
    assert adequacy_payload(state, _solved_settings) is None
    # Solved without a target: configured false.
    state["last_solve_settings"] = {"voll_eur_per_mwh": 3000.0}
    r = client.get("/api/results/adequacy")
    assert r.status_code == 200 and r.json() == {"configured": False, "engine": "lp_proxy"}
    assert _dump(r.json()) == _dump(adequacy_payload(state, _solved_settings))
    # A report: the route serves it with configured true, byte-equal.
    state["adequacy_report"] = _report()
    r = client.get("/api/results/adequacy")
    assert r.json() == {**_report(), "configured": True}
    assert r.content == client.get("/api/results/adequacy").content
    assert _dump(r.json()) == _dump(json.loads(json.dumps(adequacy_payload(state, _solved_settings))))


def test_fmea_route_json_is_the_function_the_sheet_reads(client, active, session_ctx):
    ctx = session_ctx(client)
    ctx.solver_state["solver_config"] = CFG
    r = client.get("/api/results/fmea_modes")
    assert r.status_code == 200
    payload = build_fmea_modes_payload(ctx.network, CFG, sweep_record=None,
                                       get_lock=PyPSAService.get_lock)
    assert payload["per_mode"][0]["criticality_eur_per_year"] > 0
    assert _dump(r.json()) == _dump(json.loads(json.dumps(payload)))


# ── conditions 2 and 3: Adequacy ─────────────────────────────────────────

def test_adequacy_keys_equal_the_route_per_period_and_each_scenario_reads_its_own_state(active):
    run = _scenario()
    ws = _wb(_build(active, ["adequacy"], scenarios=S.Scenarios(runs=[run])))[TITLES["adequacy"]]
    base = adequacy_payload(PyPSAService.get_solver_state(), lambda: None)
    mine = adequacy_payload(run.context.solver_state, lambda: None)
    by = {r["period"]: r for r in base["target"]["system"]["by_period"]}
    # Condition 3: the second period's figure is the report's.
    assert _key(ws, "Achieved ENS against the target", P2) == by[P2]["achieved_ens_mwh"] == 7.0
    assert _key(ws, "ENS cap (the target)", P2) == 20.0
    assert _key(ws, "Achieved ENS against the target", HORIZON) == pytest.approx(
        base["target"]["system"]["achieved_ens_mwh"])
    assert _key(ws, "Energy not served", HORIZON) == base["metrics"]["ens_mwh"]
    assert _key(ws, "System cost excluding shed, discounted over the plan", HORIZON) == base["cost"]["total_system_cost_eur"]
    assert _key(ws, "Firm capacity at peak", P2) == 118.0
    # MW is per period only: no horizon row.
    with pytest.raises(AssertionError):
        _key(ws, "Firm capacity at peak", HORIZON)
    # The scenario column is its own state: the base's figure would fail.
    assert mine["target"]["system"]["by_period"][1]["achieved_ens_mwh"] != 7.0
    assert _key(ws, "Achieved ENS against the target", P2, col=5) == pytest.approx(
        mine["target"]["system"]["by_period"][1]["achieved_ens_mwh"])
    assert _key(ws, "Shed hours", HORIZON, col=5) == pytest.approx(mine["metrics"]["shed_hours"])
    # Long tables keep every period of every run.
    rows = list(ws.iter_rows(values_only=True))
    assert {(r[0], r[1]) for r in rows if r[1] in ("2030", P2) and r[0] in ("Base", "High load")
            } >= {("Base", "2030"), ("Base", P2), ("High load", "2030"), ("High load", P2)}


def test_adequacy_solved_without_a_target_writes_not_modelled(active):
    state = PyPSAService.get_solver_state()
    state.pop("adequacy_report")
    wb = _wb(_build(active, ["adequacy"], solved_settings=lambda: {"voll_eur_per_mwh": 1.0}))
    assert "not modelled" in _text(wb[TITLES["adequacy"]])



def test_a_single_period_system_cost_is_per_year(active):
    report = PyPSAService.get_solver_state()["adequacy_report"]
    report["cost"]["period_basis"] = "single_period"
    ws = _wb(_build(active, ["adequacy"]))[TITLES["adequacy"]]
    assert _key(ws, "System cost excluding shed, per year", HORIZON) == report["cost"][
        "total_system_cost_eur"]
    assert "discounted over the plan" not in _text(ws)
    # Per year only on a full-year horizon: the sheet says how long this one is.
    assert "per year only when the horizon is a full year" in _text(ws)
    assert "Study horizon: 2 h of modelled time" in _text(ws)


def test_a_multi_period_cost_has_no_per_year_caveat_but_the_horizon_note(active):
    ws = _wb(_build(active, ["adequacy"]))[TITLES["adequacy"]]
    assert "per year only when" not in _text(ws)
    assert "Study horizon: 2 h of modelled time" in _text(ws)


def test_a_scenario_without_a_target_beside_a_base_with_one_says_not_modelled(active):
    run = _scenario(report=False)
    run.solved_settings = lambda: {"voll_eur_per_mwh": 1.0}
    ws = _wb(_build(active, ["adequacy"], scenarios=S.Scenarios(runs=[run])))[TITLES["adequacy"]]
    assert "High load: solved without a reliability target: not modelled" in _text(ws)

# ── conditions 2, 3 and 6: FMEA ──────────────────────────────────────────

def test_fmea_keys_equal_the_route_and_each_scenario_reads_its_own_network(active):
    run = _scenario()
    ws = _wb(_build(active, ["fmea"], scenarios=S.Scenarios(runs=[run])))[TITLES["fmea"]]
    base = build_fmea_modes_payload(active, CFG, sweep_record=None, get_lock=PyPSAService.get_lock)
    mine = build_fmea_modes_payload(run.network, CFG, sweep_record=None,
                                    get_lock=PyPSAService.get_lock)
    top, top_mine = base["per_mode"][0], mine["per_mode"][0]
    label = "Highest criticality, classes A/B"
    assert _key(ws, label, HORIZON, col=2) == f"EUR {ANNUALISED}"
    assert _key(ws, label, HORIZON) == pytest.approx(top["criticality_eur_per_year"]) and top[
        "criticality_eur_per_year"] > 0
    assert _key(ws, "Failure modes", HORIZON) == len(base["per_mode"])
    # The scenario column reads its own network: the base's would fail.
    assert top_mine["criticality_eur_per_year"] != pytest.approx(top["criticality_eur_per_year"])
    assert _key(ws, label, HORIZON, col=5) == pytest.approx(top_mine["criticality_eur_per_year"])
    # Condition 3: no period split, said in the note; every run's rows are listed.
    assert NO_PERIODS in _text(ws)
    rows = list(ws.iter_rows(values_only=True))
    assert sum(1 for r in rows if r[0] == "High load") == len(mine["per_mode"])


def test_fmea_class_a_rows_of_the_route_are_labelled_annualised_per_year(active):
    ws = _wb(_build(active, ["fmea"]))[TITLES["fmea"]]
    cells = [str(c) for row in ws.iter_rows(values_only=True) for c in row if isinstance(c, str)]
    assert CRITICALITY_PER_YEAR in cells and SEVERITY_PER_EVENT in cells
    assert PER_HORIZON == "per study horizon"
    assert not any("ticket 30" in c for c in cells)
    # No route field name reaches the sheet.
    assert not any("criticality_eur_per_year" in c or "severity_eur" in c for c in cells)
    # Every row is class A, annualised by LC-30: per year, and its per-horizon cells are empty.
    rows = _fmea_rows(ws)
    assert rows and all(r["criticality basis"] == ANNUALISED for r in rows)
    assert all(_number(r[CRITICALITY_PER_YEAR]) and not _number(r[CRITICALITY]) for r in rows)
    assert "Study horizon: 2 h of modelled time" in _text(ws)
    titles = [_chart_title(ch) for ch in ws._charts]
    assert any("Criticality" in t for t in titles)
    assert all("per year" in t for t in titles if "riticality" in t)
    for ch in ws._charts:
        for s in ch.series:
            name = ws[s.tx.strRef.f.split("!")[1].replace("$", "")].value
            if "riticality" in str(name) or "everity" in str(name):
                assert PER_HORIZON not in str(name)


def _mode(mode_id: str, cls: str, crit: float, basis: str | None = None) -> dict:
    m = {"mode_id": mode_id, "name": mode_id, "failure_class": cls,
         "occurrence_per_year": 2.0, "criticality_eur_per_year": crit,
         "severity_eur": crit / 2.0, "delta_eue_mwh": crit / 1000.0}
    return {**m, "criticality_basis": basis} if basis else m


def _fmea_sheet(active, monkeypatch, modes):
    monkeypatch.setattr("services.adequacy.copt_endpoint.build_fmea_modes_payload",
                        lambda *a, **k: {"per_mode": modes, "voll_eur_per_mwh": 1000.0})
    return _wb(_build(active, ["fmea"]))[TITLES["fmea"]]


def _number(v) -> bool:
    """A filled cell: the sheet writes an empty one as a "not_established" label."""
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _fmea_rows(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    h = next(i for i, r in enumerate(rows) if "criticality basis" in r)
    head = rows[h]
    return [dict(zip(head, r)) for r in rows[h + 1:] if r[0] == "Base"]


@pytest.mark.parametrize("cls,basis,crit_col,sev_col,says", [
    ("A", None, CRITICALITY, SEVERITY, LEGACY),
    ("B", "per_horizon_legacy", CRITICALITY, SEVERITY, LEGACY),
    ("A", "per_horizon", CRITICALITY, SEVERITY, PER_HORIZON),
    ("A", "annualised", CRITICALITY_PER_YEAR, SEVERITY_PER_EVENT, "per year (annualised)"),
    ("C", None, CRITICALITY_PER_YEAR, SEVERITY_PER_EVENT, "per year (frequency × severity)"),
    ("C", "per_event", CRITICALITY_PER_YEAR, SEVERITY_PER_EVENT,
     "per year (frequency × severity)"),
])
def test_each_fmea_row_is_labelled_by_its_own_basis(active, monkeypatch, cls, basis, crit_col,
                                                    sev_col, says):
    ws = _fmea_sheet(active, monkeypatch, [_mode("m1", cls, 400.0, basis)])
    (row,) = _fmea_rows(ws)
    assert row["criticality basis"] == says
    assert row[crit_col] == pytest.approx(400.0) and row[sev_col] == pytest.approx(200.0)
    # Its figures sit in no other basis's column.
    others = {CRITICALITY, CRITICALITY_PER_YEAR, SEVERITY, SEVERITY_PER_EVENT} - {crit_col, sev_col}
    assert not any(_number(row[c]) for c in others)
    family = "class C" if cls == "C" else "classes A/B"
    key = f"Highest criticality, {family}"
    assert _key(ws, key, HORIZON) == pytest.approx(400.0)
    assert _key(ws, key, HORIZON, col=2) == f"EUR {says}"


def test_mixed_bases_get_one_highest_key_each_and_no_chart_mixes_them(active, monkeypatch):
    ws = _fmea_sheet(active, monkeypatch, [
        _mode("c1", "C", 9000.0), _mode("a1", "A", 500.0), _mode("a2", "A", 700.0)])
    # The per-year C row does not outrank the per-horizon A rows' key.
    assert _key(ws, "Highest criticality, classes A/B", HORIZON) == pytest.approx(700.0)
    assert _key(ws, "Highest criticality, class C", HORIZON) == pytest.approx(9000.0)
    assert not any(str(r[0]).startswith("Highest criticality (")
                   for r in ws.iter_rows(min_row=4, values_only=True))
    # Neither criticality column is full, so neither is charted.
    assert not any("riticality" in _chart_title(ch) for ch in ws._charts)


def test_a_family_on_two_bases_has_no_highest(active, monkeypatch):
    ws = _fmea_sheet(active, monkeypatch, [
        _mode("a1", "A", 500.0, "annualised"), _mode("b1", "B", 700.0)])
    assert not _number(_key(ws, "Highest criticality, classes A/B", HORIZON))
    assert _key(ws, "Highest criticality, classes A/B", HORIZON, col=2) == \
        "EUR (rows on different bases)"


def test_per_year_rows_chart_on_the_per_year_chart(active, monkeypatch):
    ws = _fmea_sheet(active, monkeypatch, [_mode("c1", "C", 9000.0), _mode("c2", "C", 10.0)])
    titles = [_chart_title(ch) for ch in ws._charts]
    assert "Criticality by failure mode, EUR per year" in titles
    assert not any(PER_HORIZON in t for t in titles)


def _chart_title(chart) -> str:
    return "".join(r.t for p in chart.title.tx.rich.p for r in (p.r or []))


# ── condition 5: no results vs an error ──────────────────────────────────

def test_nothing_solved_and_no_failure_mode_are_no_results(install_network):
    n = pypsa.Network()
    n.set_snapshots(pd.date_range("2030-01-01", periods=2, freq="h"))
    n.add("Bus", "b")
    install_network(n)
    PyPSAService.get_solver_state().pop("adequacy_report", None)
    assert build_fmea_modes_payload(n, CFG, sweep_record=None, get_lock=PyPSAService.get_lock) is None
    wb = _wb(_build(n, ["adequacy", "fmea"]))
    about = _about(wb)
    assert not set(TITLES.values()) & set(wb.sheetnames)
    assert "no results" in str(about["adequacy"]) and "no results" in str(about["fmea"])


@pytest.mark.parametrize("view,target", [
    ("adequacy", "services.adequacy.report.adequacy_payload"),
    ("fmea", "services.adequacy.copt_endpoint.build_fmea_modes_payload")])
def test_a_failing_read_is_an_error_in_about_not_no_results(active, monkeypatch, view, target):
    def boom(*a, **k):
        raise ValueError("outage rate refused")   # the route's 422
    mod, name = target.rsplit(".", 1)
    monkeypatch.setattr(mod + "." + name, boom)
    wb = _wb(_build(active, [view]))
    assert TITLES[view] not in wb.sheetnames
    assert str(_about(wb)[view]).startswith("error")


# ── condition 9: charts ──────────────────────────────────────────────────

def _cells(ref: str) -> tuple[str, list[str]]:
    from openpyxl.utils import range_boundaries
    from openpyxl.utils.cell import get_column_letter as L
    sheet, rng = ref.rsplit("!", 1)
    c0, r0, c1, r1 = range_boundaries(rng.replace("$", ""))
    return sheet.strip("'"), [f"{L(c)}{r}" for c in range(c0, c1 + 1) for r in range(r0, r1 + 1)]


@pytest.mark.parametrize("view", list(TITLES))
def test_every_chart_reads_number_cells_on_its_own_sheet(active, view):
    ws = _wb(_build(active, [view], scenarios=S.Scenarios(runs=[_scenario()])))[TITLES[view]]
    assert 1 <= len(ws._charts) <= 3, view
    for chart in ws._charts:
        for s in chart.series:
            for ref in filter(None, [s.val.numRef.f if s.val else None]):
                sheet, cells = _cells(ref)
                assert sheet == ws.title, ref
                for c in cells:
                    v = ws[c].value
                    assert isinstance(v, (int, float)) and not isinstance(v, bool), (view, c, v)
                    assert np.isfinite(v)
            if s.cat is not None:
                assert _cells((s.cat.numRef or s.cat.strRef).f)[0] == ws.title
