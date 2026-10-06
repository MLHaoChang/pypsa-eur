"""
`gui-requirements.txt` must pin every third-party module the shipped backend
imports WITHOUT a guard.

**The bug this exists to prevent, measured in a shipped build.** A user clicked
"download template" in the packaged macOS app and got a 500. From
`~/Library/Application Support/PyPSA GUI/pypsa-gui.log`:

    File "routers/network.py", line 3568, in download_load_profile_template
    File "routers/network.py", line 441, in _xlsx_response
    File "pandas/io/excel/_openpyxl.py", line 57, in __init__
    ModuleNotFoundError: No module named 'openpyxl'

`openpyxl` is in `pixi.toml`, so every test and every dev run had it. The app
is built from a pip venv driven by `gui-requirements.txt` (D14), which did not
list it. Six modules were in that state; the build venv was the only place it
was observable and nothing was looking there.

**Why "unguarded" is the rule and not "imported".** A guarded import is a
deliberate optional — `time_aggregation_service.py` catches ImportError and
falls back to the full period, and pinning `tsam` would pull pyomo and
scikit-learn into the bundle to service a path that already works. An
UNGUARDED import is a promise that the module is there; when the promise is
broken the user gets a 500. Those are the ones this file makes non-negotiable.

**Why a static check rather than an endpoint test.** `test_desktop_downloads`
and friends run in the pixi environment, where all six modules are present. No
test executing against pixi can observe this defect — the difference lives in
the packaging manifest, so the check has to read the manifest.
"""
from __future__ import annotations

import ast
import importlib.metadata as md
import re
import sys
from pathlib import Path

import pytest

BACKEND = Path(__file__).resolve().parent.parent
REQUIREMENTS = BACKEND.parent / "gui-requirements.txt"

# `tests/` and `smoke/` are development tooling. The frozen app's entry point
# is `desktop/gui.py` and neither is reachable from it, so `pytest` and
# `requests` are not shipped dependencies.
NOT_SHIPPED = {"tests", "smoke"}

# Third-party modules the backend imports INSIDE a try/except that catches the
# import failure. Each is allowed to be absent from the bundle, and the value
# records what the user actually gets when it is — which is the thing worth
# reviewing. A guarded import missing from this map fails the test: adding an
# optional dependency should be a decision someone wrote down, not a default.
OPTIONAL_AT_RUNTIME = {
    "gridspine": (
        "gridspine_service catches ImportError and every action answers 503 "
        "`planning → dynamics not available in this build`. Since D5 gridspine "
        "is installed EDITABLE into every pixi environment; since increment 6 "
        "the frozen app freezes it too — from the repo root through the spec's "
        "`pathex` (pixi's editable install is an import-finder hook PyInstaller "
        "cannot follow), with its engines pinned in gui-requirements.txt and "
        "its YAML library and pandapower's case data collected. The guard stays "
        "as defence in depth: a build from a partial checkout would otherwise "
        "500 on the first study instead of saying what is missing. Verified by "
        "a Linux onedir freeze and a frozen-tree inspection; a macOS .app run "
        "is still owed."
    ),
    "tsam": (
        "time_aggregation_service catches ImportError and falls back to the "
        "full period. Correct result, slower solve — but the fallback is "
        "SILENT (log line only), so a user who asked for representative "
        "periods is never told they did not get them."
    ),
    "magic": (
        "uploads.py catches (ImportError, OSError) and trusts the client's "
        "DECLARED content-type. The wheel is ctypes bindings over a system "
        "libmagic this bundle does not ship, so pinning it alone would not "
        "change behaviour on a clean Mac."
    ),
    "anthropic": (
        "chat_service returns the typed error `sdk_not_installed` and the "
        "panel renders disabled. Pinned anyway — the panel is meant to work. "
        "The packaged app still ships no .env (check_bundle.py keeps it out), "
        "but since U-1 the key gate is no longer a dead end: the "
        "missing_api_key banner carries a field that writes "
        "<app-data>/user.env (services/app_secrets.py)."
    ),
    "pypdf": (
        "upload_service returns (None, False) — page count unknown. A "
        ">100-page PDF then reaches Anthropic's cap as a 415 instead of a "
        "clean local truncation banner. Pinned to keep the banner working."
    ),
}


def _normalise(name: str) -> str:
    """PEP 503: `python-magic`, `Python_Magic` and `python.magic` are one
    distribution."""
    return re.sub(r"[-_.]+", "-", name).strip().lower()


def _pinned_distributions() -> set[str]:
    pinned: set[str] = set()
    for raw in REQUIREMENTS.read_text(encoding="utf-8").splitlines():
        line = raw.split("#", 1)[0].strip()
        if not line or line.startswith("-"):
            continue
        name = re.split(r"[<>=!~\[;]", line, maxsplit=1)[0]
        if name:
            pinned.add(_normalise(name))
    return pinned


def _first_party() -> set[str]:
    names = {p.name for p in BACKEND.iterdir() if p.is_dir()}
    names |= {p.stem for p in BACKEND.glob("*.py")}
    return names


def _catches_import_error(handler: ast.ExceptHandler) -> bool:
    """Does this `except` clause swallow a failed import?"""
    caught = handler.type
    if caught is None:  # bare `except:`
        return True
    names = caught.elts if isinstance(caught, ast.Tuple) else [caught]
    for node in names:
        if isinstance(node, ast.Name) and node.id in {
            "ImportError", "ModuleNotFoundError", "Exception", "BaseException",
        }:
            return True
    return False


def _collect(node: ast.AST, guarded: bool, out: dict[str, set[tuple[str, bool]]],
             where: str) -> None:
    """
    Walk `node`, recording every import as (file, guarded).

    Hand-rolled rather than `ast.walk` because guardedness depends on ancestry:
    only the `body` of a try whose handler catches ImportError is protected.
    An import in the HANDLER of that try is not — `io.py:95` imports openpyxl
    inside `except AttributeError`, which catches nothing about openpyxl.
    """
    for child in ast.iter_child_nodes(node):
        if isinstance(child, ast.Try):
            protected = guarded or any(
                _catches_import_error(h) for h in child.handlers
            )
            for stmt in child.body:
                _collect(stmt, protected, out, where)
            for section in (child.handlers, child.orelse, child.finalbody):
                for stmt in section:
                    _collect(stmt, guarded, out, where)
            continue

        if isinstance(child, ast.Import):
            modules = [alias.name for alias in child.names]
        elif isinstance(child, ast.ImportFrom):
            # `level > 0` is relative — first-party by definition.
            modules = [child.module] if child.level == 0 and child.module else []
        else:
            modules = []

        for module in modules:
            out.setdefault(module.split(".")[0], set()).add((where, guarded))

        _collect(child, guarded, out, where)


def _shipped_imports() -> dict[str, set[tuple[str, bool]]]:
    """Third-party top-level imports in shipped backend code, each tagged with
    the file that imports it and whether that import is guarded."""
    first_party = _first_party()
    raw: dict[str, set[tuple[str, bool]]] = {}

    for path in sorted(BACKEND.rglob("*.py")):
        rel = path.relative_to(BACKEND)
        if NOT_SHIPPED & set(rel.parts):
            continue
        tree = ast.parse(path.read_text(encoding="utf-8"), filename=str(path))
        _collect(tree, False, raw, str(rel))

    return {
        module: sites
        for module, sites in raw.items()
        if module not in sys.stdlib_module_names
        and module not in first_party
        and not module.startswith("_")
    }


def _distribution_for(module: str) -> list[str] | None:
    return md.packages_distributions().get(module)


def test_every_unguarded_third_party_import_is_pinned_for_the_build():
    """
    An unguarded import is a promise the module is there. Break it and the
    user gets a 500 — which is exactly what shipped.
    """
    pinned = _pinned_distributions()
    problems: list[str] = []

    for module, sites in sorted(_shipped_imports().items()):
        unguarded = sorted(where for where, guarded in sites if not guarded)
        if not unguarded:
            continue
        providers = _distribution_for(module)
        if not providers:
            problems.append(
                f"{module!r} (imported unguarded by {', '.join(unguarded)}) is "
                "not installed here, so its distribution cannot be resolved"
            )
        elif not any(_normalise(p) in pinned for p in providers):
            problems.append(
                f"{module!r} (provided by {'/'.join(providers)}, imported "
                f"unguarded by {', '.join(unguarded)}) is missing from "
                "gui-requirements.txt"
            )

    assert not problems, (
        "These modules will NOT be in the frozen app, and nothing catches "
        "their absence — the feature 500s at runtime:\n  "
        + "\n  ".join(problems)
    )


def test_every_optional_dependency_has_a_recorded_consequence():
    """
    A guarded import may legitimately be left out of the bundle — but somebody
    has to have decided that, and written down what the user gets instead.
    A new guarded import fails here until that decision is recorded.
    """
    undocumented = sorted(
        module
        for module, sites in _shipped_imports().items()
        if all(guarded for _, guarded in sites)
        and module not in OPTIONAL_AT_RUNTIME
    )

    assert not undocumented, (
        "These imports are guarded, so the app survives without them — but "
        "the fallback behaviour is unrecorded. Add each to "
        "OPTIONAL_AT_RUNTIME describing what the user actually gets, then "
        "decide whether to pin it:\n  " + "\n  ".join(undocumented)
    )


def test_openpyxl_is_required_because_nothing_catches_its_absence():
    """
    The regression, named. `io.py:95` imports openpyxl inside an
    `except AttributeError` handler — which catches nothing about openpyxl —
    so the import is unguarded and the module must ship.

    This asserts the CLASSIFIER, not just the pin: a guard-detector that
    wrongly treated any try/except as protection would demote openpyxl to
    optional and silently retire the test above.
    """
    sites = _shipped_imports()["openpyxl"]

    assert any(not guarded for _, guarded in sites), (
        "openpyxl is now classified as guarded everywhere it is imported — if "
        "a real guard was added, confirm the fallback works before relaxing "
        "the pin, because 'download template' has no fallback"
    )
    assert "openpyxl" in _pinned_distributions()


def test_defusedxml_ships_because_openpyxl_only_hardens_itself_when_it_can_import_it():
    """
    openpyxl parses every uploaded workbook (`io.py`'s template upload, and
    pandas' Excel engine) and guards against entity-expansion ("billion
    laughs") and quadratic-blowup XML ONLY if `defusedxml` is importable —
    `openpyxl.xml.DEFUSEDXML` is decided at import by trying it. Nothing in
    this codebase imports defusedxml, so the import scan above cannot see the
    need, and the failure is silent: an unhardened parser still parses. The
    `test` environment happened to carry it transitively while the build venv
    and the default environment did not — so every test ran hardened and the
    app did not.
    """
    assert "defusedxml" in _pinned_distributions()


def test_openpyxl_parses_uploads_with_defusedxml_in_this_environment():
    import openpyxl.xml

    assert openpyxl.xml.DEFUSEDXML is True


def test_a_guard_that_does_not_catch_importerror_is_not_a_guard():
    """
    Mutation check on `_catches_import_error`. `except ValueError` must not
    count as protection, or every unguarded import inside an unrelated
    try/except would be waved through.
    """
    protects = ast.parse("try:\n import x\nexcept ImportError:\n pass").body[0]
    does_not = ast.parse("try:\n import x\nexcept ValueError:\n pass").body[0]
    bare = ast.parse("try:\n import x\nexcept:\n pass").body[0]
    tupled = ast.parse(
        "try:\n import x\nexcept (OSError, ImportError):\n pass"
    ).body[0]

    assert _catches_import_error(protects.handlers[0])
    assert not _catches_import_error(does_not.handlers[0])
    assert _catches_import_error(bare.handlers[0])
    assert _catches_import_error(tupled.handlers[0])


def test_requirements_parsing_ignores_comments_and_blank_lines():
    """
    `gui-requirements.txt` is mostly prose. A parser that treated comment lines
    as requirements would "pin" everything and this file would assert nothing.
    """
    pinned = _pinned_distributions()

    assert "pandas" in pinned
    assert "pypsa" in pinned
    assert not any(p.startswith("#") for p in pinned)
    assert "" not in pinned
    # The prose names modules it deliberately does NOT pin; they must not be
    # parsed as pins.
    assert "tsam" not in pinned


# --------------------------------------------------------------------------
# the planning → dynamics pipeline ships (increment 6)
# --------------------------------------------------------------------------

SPEC = BACKEND.parent / "pypsa-gui.spec"
BUILD_SCRIPT = BACKEND.parent / "build-macos.sh"


def test_the_planning_pipelines_engines_are_pinned_for_the_build():
    """gridspine imports pandapower, lightsim2grid and yaml at module scope.
    gridspine itself is frozen from source (see the spec), so it is not a pin;
    what it imports must be."""
    pinned = _pinned_distributions()
    for dist in ("pandapower", "lightsim2grid", "pyyaml"):
        assert dist in pinned, f"{dist} is missing from gui-requirements.txt"


def test_the_spec_freezes_gridspine_from_the_repo_root_with_its_data():
    """The spec cannot be imported (PyInstaller injects its globals), so its
    text is held to the four things that make gridspine work frozen: the repo
    root on pathex, the YAML library and pandapower's case data as datas, and
    both packages written out as real directories for their `__file__`-relative
    reads."""
    text = SPEC.read_text(encoding="utf-8")
    assert "pathex=[str(BACKEND), str(ROOT.parent)]" in text
    assert '"gridspine/templates/data"' in text
    assert 'collect_data_files("pandapower"' in text
    assert '"gridspine": "pyz+py"' in text and '"pandapower": "pyz+py"' in text
    # and the one exclusion the frozen probe found necessary: without it the
    # frozen app's pandapower looks for case39.json one directory too deep
    assert '"pandapower.__init__",' in text


def test_the_spec_ships_the_generic_defaults_pack_data():
    """IC U1 (a): the defaults pack's loader is a module (collected by import)
    but its versioned data files are not — the spec must list them, at the
    path the loader resolves (`services/library/defaults_pack/versions`)."""
    from services.library.defaults_pack.loader import VERSIONS_DIR

    text = SPEC.read_text(encoding="utf-8")
    assert '"services/library/defaults_pack/versions"' in text
    assert VERSIONS_DIR.parts[-4:] == ("services", "library", "defaults_pack", "versions")
    assert any(VERSIONS_DIR.iterdir())


def test_the_build_script_builds_the_project_templates_with_gridspine_reachable():
    """A fresh checkout ships no `network.nc` templates unless the build makes
    them, and the IEEE 39-bus one needs gridspine on the path to be built
    rather than skipped."""
    text = BUILD_SCRIPT.read_text(encoding="utf-8")
    assert "project_templates/_build.py" in text
    assert 'PYTHONPATH="$REPO"' in text


def test_the_bundle_check_expects_the_pipelines_data_files():
    sys.path.insert(0, str(BACKEND / "smoke"))
    try:
        import check_bundle
    finally:
        sys.path.pop(0)
    assert {"case39_units.yaml", "case39.json"} <= set(check_bundle.EXPECTED)


# --------------------------------------------------------------------------
# the decision study ships its report template and its assumptions library
# (plan 2026-09-28 guided investment study MVP-1, S9; gates S2, S4, S7, S8)
# --------------------------------------------------------------------------

_BACKEND_DATAS = re.compile(r'\(str\(BACKEND((?:\s*/\s*"[^"]+")+)\),\s*"([^"]*)"\)')


def _spec_backend_datas() -> dict[str, str]:
    """
    ``{source relative to backend: destination}`` for every ``datas`` entry
    the spec takes from ``BACKEND``. The spec cannot be imported (PyInstaller
    injects its globals), so its text is read.
    """
    out: dict[str, str] = {}
    for m in _BACKEND_DATAS.finditer(SPEC.read_text(encoding="utf-8")):
        out["/".join(re.findall(r'"([^"]+)"', m.group(1)))] = m.group(2)
    return out


def _check_bundle():
    sys.path.insert(0, str(BACKEND / "smoke"))
    try:
        import check_bundle
    finally:
        sys.path.pop(0)
    return check_bundle


def test_the_decision_studys_template_and_library_are_in_the_spec_and_the_bundle_check():
    """
    `services/study/render_html.py` reads `templates/decision_report.html.j2`
    and `services/study/library.py` reads `study_library/` (its CSVs, the
    finance YAML and `load_profiles/`), both `__file__`-relative. A bundle
    without them launches fine and 500s on the first report or the first
    study.
    """
    datas = _spec_backend_datas()
    assert datas.get("templates/decision_report.html.j2") == "templates", datas
    assert datas.get("study_library") == "study_library", datas
    expected = set(_check_bundle().EXPECTED)
    assert {"decision_report.html.j2", "study_library", "technology_costs.csv",
            "tariffs.csv", "finance_defaults.yaml", "load_profiles"} <= expected


def test_the_report_charts_agg_backend_is_a_declared_hidden_import():
    """
    `report_charts._figure` selects `Agg` by NAME (`matplotlib.use("Agg")`)
    and `Figure.savefig(format="png")` resolves the canvas through a string
    table, so static analysis sees no import of the Agg backend. PyInstaller's
    matplotlib hook happens to pick it up from the `use()` call today; the
    spec says so itself rather than relying on the hook's detection.
    """
    text = SPEC.read_text(encoding="utf-8")
    assert '"matplotlib.backends.backend_agg"' in text
    assert "matplotlib" in _pinned_distributions()


def test_the_decision_studys_data_resolves_in_a_frozen_layout(tmp_path):
    """
    The frozen layout, rebuilt from the spec: `pathex=[BACKEND, ...]` puts
    `services/study/*.py` at `_MEIPASS/services/study/`, and each `datas`
    entry lands at its destination (a file under the destination directory,
    a directory's contents AS the destination). The two modules are loaded
    from that tree, so their `__file__`-relative paths are the frozen ones —
    the `presets.json` and `matpower.jinja2` precedent — and must find the
    template and every library file there.
    """
    import importlib.util
    import shutil

    meipass = tmp_path / "_MEIPASS"
    for src, dest in _spec_backend_datas().items():
        source = BACKEND / src
        if not src.startswith(("templates", "study_library")):
            continue
        if source.is_dir():
            shutil.copytree(source, meipass / dest)
        else:
            (meipass / dest).mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, meipass / dest / source.name)
    frozen = meipass / "services" / "study"
    frozen.mkdir(parents=True)
    modules = {}
    for name in ("render_html", "library"):
        shutil.copy2(BACKEND / "services" / "study" / f"{name}.py", frozen / f"{name}.py")
        spec = importlib.util.spec_from_file_location(
            f"_frozen_probe_{name}", frozen / f"{name}.py")
        module = importlib.util.module_from_spec(spec)
        # dataclasses resolve their annotations through `sys.modules`
        sys.modules[spec.name] = module
        try:
            spec.loader.exec_module(module)
        finally:
            sys.modules.pop(spec.name, None)
        modules[name] = module

    render_html, library = modules["render_html"], modules["library"]
    assert render_html.TEMPLATES_DIR == meipass / "templates"
    template = render_html.environment().get_template("decision_report.html.j2")
    assert template.filename == str(meipass / "templates" / "decision_report.html.j2")

    assert library.LIBRARY_DIR == meipass / "study_library"
    loaded = library.load_library()
    assert loaded.tariffs and loaded.finance["library_version"] == library.LIBRARY_VERSION
    profiles = library.LIBRARY_DIR / "load_profiles"
    shipped = {p.name for p in (BACKEND / "study_library" / "load_profiles").iterdir()
               if p.is_file()}
    assert shipped and shipped <= {p.name for p in profiles.iterdir()}


def test_python_docx_ships_its_templates_and_renders_in_a_frozen_layout(tmp_path):
    """
    Gate S9 [S3]: `docx.Document()` opens `docx/templates/default.docx`
    `__file__`-relative, and `docx/parts/*.py` read `parts/../templates/*.xml`
    — the `pypsa/optimization/../data` shape, which needs the package written
    out as a real directory (`pyz+py`) as well as its data collected. The
    spec must say both, `check_bundle.EXPECTED` must name the template, and a
    report must render with `docx` imported from the rebuilt layout (the
    package copied the way `pyz+py` + `collect_data_files` lay it out).
    """
    import importlib.util
    import shutil
    import subprocess

    text = SPEC.read_text(encoding="utf-8")
    assert 'collect_data_files("docx")' in text
    assert '"docx": "pyz+py"' in text
    assert "default.docx" in set(_check_bundle().EXPECTED)

    meipass = tmp_path / "_MEIPASS"
    source = Path(importlib.util.find_spec("docx").origin).parent
    shutil.copytree(source, meipass / "docx",
                    ignore=shutil.ignore_patterns("__pycache__"))
    from tests.test_study_decision_report import _report

    report_json = tmp_path / "report.json"
    report_json.write_text(_report().model_dump_json(by_alias=True), encoding="utf-8")
    out = tmp_path / "report.docx"
    probe = (
        "import sys, pathlib\n"
        f"sys.path[:0] = [{str(meipass)!r}, {str(BACKEND)!r}]\n"
        "import docx\n"
        f"assert pathlib.Path(docx.__file__).is_relative_to({str(meipass)!r}), docx.__file__\n"
        "from models.study import DecisionReport\n"
        "from services.study.render_docx import render_docx\n"
        f"r = DecisionReport.model_validate_json(pathlib.Path({str(report_json)!r}).read_text())\n"
        "doc = docx.Document()\n"
        "doc.sections[0].header.paragraphs[0].text = 'header'   # parts/../templates\n"
        f"pathlib.Path({str(out)!r}).write_bytes(render_docx(r, charts={{}}))\n"
    )
    proc = subprocess.run([sys.executable, "-c", probe], capture_output=True, text=True,
                          cwd=str(tmp_path), timeout=300)
    assert proc.returncode == 0, proc.stderr[-2000:]
    import docx

    doc = docx.Document(str(out))
    assert any(t.rows[0].cells[0].text == "Assumption" for t in doc.tables)
def _guarded_gridspine_modules() -> set[str]:
    """Every `gridspine.*` module the backend's import guard names.

    Read out of the source rather than listed here, so the check cannot go
    stale the way a hand-maintained second copy of the list would — and so a
    module ADDED to the guard later is covered without anyone remembering to
    extend this file.
    """
    tree = ast.parse(
        (BACKEND / "services" / "gridspine_service.py").read_text(encoding="utf-8"),
        filename="gridspine_service.py",
    )
    found: set[str] = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.ImportFrom) and (node.module or "").startswith(
            "gridspine."
        ):
            found.add(node.module)
        elif isinstance(node, ast.Import):
            for alias in node.names:
                if alias.name.startswith("gridspine."):
                    found.add(alias.name)
    assert found, "no gridspine imports found — did gridspine_service move?"
    return found


def test_the_spec_names_every_gridspine_module_the_backend_guard_imports():
    """
    The spec states this rule itself: the gridspine entry modules are spelled
    out "so a future change to that guard cannot drop the package from the
    bundle without a build error". Nothing was enforcing it, and it was
    already broken once — `drivers.readback` reached the guard with the
    PowerFactory read-back and never reached the spec.

    Why it matters even though PyInstaller's analysis follows a try/except
    import today: `gridspine` is frozen from the repo root through `pathex`,
    not installed into the build venv, so nothing but that analysis puts it in
    the bundle. When a module is missed the app still LAUNCHES — the guard
    catches the ImportError — and every planning→dynamics action answers 503
    "not available in this build". That reads as a deliberate build variant
    rather than a packaging bug, which is why it survives a smoke test.
    """
    text = SPEC.read_text(encoding="utf-8")
    missing = sorted(
        module
        for module in _guarded_gridspine_modules()
        if f'"{module}"' not in text
    )
    assert not missing, (
        "pypsa-gui.spec's hiddenimports does not name "
        f"{missing} — the backend imports them inside the gridspine guard, so "
        "a bundle built without them answers 503 on every gridspine action"
    )
