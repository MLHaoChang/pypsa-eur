"""
WP11 (backend) — user templates over HTTP.

  POST /api/projects/{name}/uploads?kind=report_template         the template as an upload
  POST /api/projects/{name}/reports/{id}/template  {file_id|null}  bind / unbind
  GET  /api/projects/{name}/reports/{id}/template                  outline + stored plan
  POST /api/projects/{name}/reports/{id}/template/plan             the "mapping" job
  PUT  /api/projects/{name}/reports/{id}/template/plan             a user-edited plan
  POST /api/projects/{name}/reports/{id}/export                    honours the template
  POST /api/projects/{name}/reports/generate {template_file_id}    binds, defaults language

The tagged leg runs WP9's real `render_tagged` on the tagged fixture. The
untagged leg codes against WP10's pinned signatures through the accessors
`services.reports.templates` exposes (`_render_untagged`, `_default_mapping`,
`_propose_mapping`), replaced here with fakes so the routes are pinned
whether or not `template_untagged` is present in the checkout. The provider
for the plan job is the scripted `llm_fake` installed at the seam the job
router looks up (`routers.report_jobs._provider_for_profile`).
"""
from __future__ import annotations

import importlib.util
import io
import json
import time
import uuid as _uuid
from datetime import datetime, UTC
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

import main
from models.energy_hub import ReferenceDesignReport, SectionState, empty_section_map
from services.adequacy import eh_report as R
from services.llm_fake import FakeProvider
from services.llm_provider import LLMEvent
from services.reports import store, templates
from services.reports.docx_writer import DOCX_MIME, render_document_docx
from services.reports.generator import SectionFailure
from tests.conftest import attach_session

_BUILD = Path(__file__).resolve().parent / "fixtures" / "report_templates" / "_build.py"


# ── fixtures ────────────────────────────────────────────────────────────────

def _load_builder():
    spec = importlib.util.spec_from_file_location("report_template_fixtures", _BUILD)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def fixtures(tmp_path_factory) -> dict[str, bytes]:
    builder = _load_builder()
    out = tmp_path_factory.mktemp("report_templates")
    paths = builder.build_all(out)
    paths.update(builder.build_all(out, language="de"))
    return {name: path.read_bytes() for name, path in paths.items()}


@pytest.fixture
def same_org_other_user(_auth_db, seeded_identity):
    from db.models import OrgMembership, User
    from services.auth_service import hash_password

    _engine, session_local = _auth_db
    with session_local() as db:
        u = User(
            id=_uuid.uuid4(),
            email=f"colleague-{_uuid.uuid4().hex[:6]}@example.com",
            password_hash=hash_password("irrelevant"),
            status="active",
            is_super_admin=False,
            created_at=datetime.now(tz=UTC),
        )
        db.add(u)
        db.flush()
        db.add(OrgMembership(id=_uuid.uuid4(), user_id=u.id,
                             org_id=seeded_identity["org_id"], role="admin"))
        db.commit()
        uid = u.id
    with TestClient(main.app) as c:
        yield attach_session(c, session_local, uid)


def _store_eh_report(state: dict) -> None:
    sections = empty_section_map(default="skipped")
    sections["fmea_top"] = SectionState(status="ok", payload={
        "top": [{"rank": 1, "mode_id": "gen:g:forced_outage",
                 "component_class": "Generator", "name": "g",
                 "failure_class": "A", "occurrence_per_year": 1.0,
                 "occurrence_basis": "FOR", "severity_eur": 10.0,
                 "criticality_eur_per_year": 10.0, "delta_eue_mwh": 0.5,
                 "engine": "copt", "fidelity": "analytic_convolution"}],
        "classes_included": ["A"], "note": "Link-primary residual risk",
    })
    report = ReferenceDesignReport(
        archetype="strong_grid", pack_hash="p", assumptions_hash="a",
        sections=sections, ens_cap_permyriad=10.0, mc_lole_h=3.21,
    )
    R.store_eh_report(state, report)


def _upload(client, name: str, filename: str, data: bytes, *, kind: str | None = "report_template",
            mime: str = DOCX_MIME) -> dict:
    params = {} if kind is None else {"kind": kind}
    r = client.post(f"/api/projects/{name}/uploads", params=params,
                    files={"file": (filename, data, mime)})
    assert r.status_code == 200, r.text[:300]
    return r.json()


def _report(client, name: str, session_state=None, *, title: str = "Template report") -> str:
    if session_state is not None:
        _store_eh_report(session_state(client))
    r = client.post(f"/api/projects/{name}/reports",
                    json={"mode": "evidence_only", "title": title})
    assert r.status_code == 200, r.text[:300]
    return r.json()["report_id"]


def _bind(client, name: str, rid: str, file_id: str | None) -> dict:
    r = client.post(f"/api/projects/{name}/reports/{rid}/template", json={"file_id": file_id})
    assert r.status_code == 200, r.text[:400]
    return r.json()


def _docx(blob: bytes):
    return Document(io.BytesIO(blob))


def _all_text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    for section in doc.sections:
        parts.extend(p.text for p in section.header.paragraphs)
        parts.extend(p.text for p in section.footer.paragraphs)
    return "\n".join(parts)


def _turn(payload: dict | str) -> dict:
    text = payload if isinstance(payload, str) else json.dumps(payload)
    return {"events": [LLMEvent(type="text_delta", text=text)],
            "blocks": [{"type": "text", "text": text}]}


def _install_provider(monkeypatch, provider):
    import routers.report_jobs as jobs
    monkeypatch.setattr(jobs, "_provider_for_profile", lambda profile: (provider, None))
    return provider


def _poll(client, name: str, timeout: float = 10.0) -> dict:
    deadline = time.time() + timeout
    body = None
    while time.time() < deadline:
        r = client.get(f"/api/projects/{name}/reports/generate/status")
        assert r.status_code == 200, r.text
        body = r.json()
        if body.get("status") != "running":
            return body
        time.sleep(0.02)
    raise AssertionError(f"report job never finished: {body!r}")


def _heading_index(outline: dict, prefix: str) -> int:
    return next(h["index"] for h in outline["headings"] if h["text"].startswith(prefix))


def _plan_for(outline: dict, *, rename_text: str = "4 Residual failure modes (renamed)") -> dict:
    """A valid plan: keep 1–3, rename heading 4, drop heading 5 and the placeholder."""
    return {
        "entries": [
            {"heading_index": _heading_index(outline, "1 "), "action": "keep",
             "new_text": None, "section_ids": ["executive_summary"]},
            {"heading_index": _heading_index(outline, "2 "), "action": "keep",
             "new_text": None, "section_ids": []},
            {"heading_index": _heading_index(outline, "3 "), "action": "keep",
             "new_text": None, "section_ids": ["certification"]},
            {"heading_index": _heading_index(outline, "4 "), "action": "rename",
             "new_text": rename_text, "section_ids": ["fmea_top"]},
            {"heading_index": _heading_index(outline, "5 "), "action": "drop",
             "new_text": None, "section_ids": []},
        ],
        "inserted": [],
        "placeholders": {"[Client name]": "ACME"},
        "unmapped_sections": [],
        "notes": [],
    }


class _FakeUntagged:
    """
    WP10's surface as fakes: `render_untagged` writes the default document
    with a marker paragraph naming the plan it received; `default_mapping`
    returns an empty plan; `propose_mapping` consumes ONE provider turn (so
    the seam is exercised) and returns the plan that turn carries, or a
    `SectionFailure` when the turn is not a plan.
    """

    def __init__(self) -> None:
        self.rendered: list[dict] = []
        self.proposals: list[dict] = []

    def render_untagged(self, doc, template, plan, *, figure_bytes):
        plan_dict = plan.model_dump() if hasattr(plan, "model_dump") else dict(plan)
        self.rendered.append(plan_dict)
        marked = doc.model_copy(update={"title": f"UNTAGGED::{doc.title}"})
        return render_document_docx(marked, figure_bytes=figure_bytes)

    def default_mapping(self, outline, doc):
        return templates.MappingPlanBody(
            entries=[], inserted=[], placeholders={}, unmapped_sections=[],
            notes=["default_mapping (fake)"])

    def propose_mapping(self, provider, *, base_request, outline, doc, language):
        from services.llm_provider import LLMRequest

        request = LLMRequest(
            model=base_request["model"], max_tokens=base_request["max_tokens"],
            system_blocks=list(base_request.get("system_blocks") or []),
            tools=[], tools_stable=True,
            messages=[{"role": "user", "content": [{"type": "text", "text": "plan"}]}],
            history_stable_anchor=None)
        text = "".join(e.text for e in provider.stream(request) if e.type == "text_delta")
        self.proposals.append({"language": language, "n_headings": len(outline.headings)})
        try:
            payload = json.loads(text)
        except ValueError:
            return SectionFailure("no_json", raw_head=text[:40], repairs=1)
        return templates.MappingPlanBody.model_validate(payload)


@pytest.fixture
def fake_untagged(monkeypatch) -> _FakeUntagged:
    fake = _FakeUntagged()
    monkeypatch.setattr(templates, "_render_untagged", lambda: fake.render_untagged)
    monkeypatch.setattr(templates, "_default_mapping", lambda: fake.default_mapping)
    monkeypatch.setattr(templates, "_propose_mapping", lambda: fake.propose_mapping)
    monkeypatch.setattr(templates, "_plan_model", lambda: templates.MappingPlanBody)
    return fake


# ── the upload kind ─────────────────────────────────────────────────────────

def test_a_docx_uploaded_as_report_template_lists_under_that_kind(client, api_project, fixtures):
    name = api_project("tpl-upload")
    meta = _upload(client, name, "tagged_minimal.docx", fixtures["tagged_minimal.docx"])
    assert meta["kind"] == "report_template" and meta["mime"] == DOCX_MIME
    plain = _upload(client, name, "notes.csv", b"a,b\n1,2\n", kind=None, mime="text/csv")
    assert plain["kind"] == "user_upload"

    only = client.get(f"/api/projects/{name}/uploads", params={"kind": "report_template"})
    assert only.status_code == 200, only.text
    assert [u["file_id"] for u in only.json()] == [meta["file_id"]]
    everything = client.get(f"/api/projects/{name}/uploads").json()
    assert {u["file_id"] for u in everything} == {meta["file_id"], plain["file_id"]}
    users = client.get(f"/api/projects/{name}/uploads", params={"kind": "user_upload"}).json()
    assert [u["file_id"] for u in users] == [plain["file_id"]]


def test_an_unknown_upload_kind_is_400_unsupported_upload_kind(client, api_project, fixtures):
    name = api_project("tpl-upload-kind")
    r = client.post(f"/api/projects/{name}/uploads", params={"kind": "agent_export"},
                    files={"file": ("t.docx", fixtures["tagged_minimal.docx"], DOCX_MIME)})
    assert r.status_code == 400, r.text
    assert r.json()["detail"]["error_kind"] == "unsupported_upload_kind"
    r = client.get(f"/api/projects/{name}/uploads", params={"kind": "nope"})
    assert r.status_code == 400
    assert r.json()["detail"]["error_kind"] == "unsupported_upload_kind"
    assert client.get(f"/api/projects/{name}/uploads").json() == []


def test_the_same_bytes_uploaded_again_as_a_template_are_promoted(client, api_project, fixtures):
    """Dedup is on bytes; a template re-sent as a template must be listed as one."""
    name = api_project("tpl-upload-dedup")
    first = _upload(client, name, "a.docx", fixtures["tagged_minimal.docx"], kind=None)
    assert first["kind"] == "user_upload"
    again = _upload(client, name, "a.docx", fixtures["tagged_minimal.docx"])
    assert again["file_id"] == first["file_id"] and again["kind"] == "report_template"
    listed = client.get(f"/api/projects/{name}/uploads", params={"kind": "report_template"}).json()
    assert [u["file_id"] for u in listed] == [first["file_id"]]


# ── bind / unbind / GET ─────────────────────────────────────────────────────

def test_get_template_is_all_null_before_any_binding(client, api_project):
    name = api_project("tpl-get-null")
    rid = _report(client, name)
    r = client.get(f"/api/projects/{name}/reports/{rid}/template")
    assert r.status_code == 200, r.text
    assert r.json() == {"template_file_id": None, "mode": None, "language": None,
                        "outline": None, "plan": None}


def test_bind_a_tagged_template_returns_its_outline_and_stores_the_binding(
        client, api_project, project_storage_dir, fixtures):
    name = api_project("tpl-bind")
    rid = _report(client, name)
    meta = _upload(client, name, "tagged_minimal.docx", fixtures["tagged_minimal.docx"])

    body = _bind(client, name, rid, meta["file_id"])
    assert body["template_file_id"] == meta["file_id"]
    assert body["mode"] == "tagged" and body["language"] is None
    outline = body["outline"]
    assert outline["mode"] == "tagged"
    assert [t["kind"] for t in outline["tags"]][:3] == ["var", "var", "for"]

    doc = client.get(f"/api/projects/{name}/reports/{rid}").json()
    assert doc["template_file_id"] == meta["file_id"]
    assert doc["version"] == 2, "binding is a new version: a version file is never rewritten"
    v1 = client.get(f"/api/projects/{name}/reports/{rid}/versions/1").json()
    assert v1["template_file_id"] is None
    m = store.load_meta(project_storage_dir(name), rid)
    assert m.template_mode == "tagged" and m.template_language is None
    assert m.mapping_plan is None

    got = client.get(f"/api/projects/{name}/reports/{rid}/template").json()
    assert got["template_file_id"] == meta["file_id"] and got["mode"] == "tagged"
    assert got["outline"]["n_paragraphs"] == outline["n_paragraphs"] and got["plan"] is None
    listed = client.get(f"/api/projects/{name}/reports").json()
    assert listed[0]["template_mode"] == "tagged"

    # Binding the SAME file again is a no-op on the versions.
    again = _bind(client, name, rid, meta["file_id"])
    assert again["template_file_id"] == meta["file_id"]
    assert client.get(f"/api/projects/{name}/reports/{rid}").json()["version"] == 2


def test_bind_the_german_corporate_template_detects_de_and_untagged(
        client, api_project, fixtures):
    name = api_project("tpl-bind-de")
    rid = _report(client, name)
    meta = _upload(client, name, "corporate_de.docx", fixtures["corporate_untagged_de.docx"])
    body = _bind(client, name, rid, meta["file_id"])
    assert body["mode"] == "untagged" and body["language"] == "de"
    assert body["outline"]["has_toc"] is True and body["outline"]["body_start_index"] == 5
    assert any(h["text"] == "1 Zusammenfassung" for h in body["outline"]["headings"])


def test_unbind_clears_the_binding_the_mode_and_the_plan(
        client, api_project, project_storage_dir, fixtures, fake_untagged):
    name = api_project("tpl-unbind")
    rid = _report(client, name)
    meta = _upload(client, name, "corporate.docx", fixtures["corporate_untagged.docx"])
    bound = _bind(client, name, rid, meta["file_id"])
    plan = _plan_for(bound["outline"])
    put = client.put(f"/api/projects/{name}/reports/{rid}/template/plan", json=plan)
    assert put.status_code == 200, put.text[:300]
    assert store.load_meta(project_storage_dir(name), rid).mapping_plan is not None

    body = _bind(client, name, rid, None)
    assert body == {"template_file_id": None, "mode": None, "language": None, "outline": None}
    doc = client.get(f"/api/projects/{name}/reports/{rid}").json()
    assert doc["template_file_id"] is None and doc["version"] == 3
    m = store.load_meta(project_storage_dir(name), rid)
    assert m.template_mode is None and m.template_language is None and m.mapping_plan is None
    assert client.get(f"/api/projects/{name}/reports/{rid}/template").json()["plan"] is None
    # Unbinding an unbound report changes nothing.
    _bind(client, name, rid, None)
    assert client.get(f"/api/projects/{name}/reports/{rid}").json()["version"] == 3


def test_rebinding_a_different_template_clears_the_plan(
        client, api_project, project_storage_dir, fixtures, fake_untagged):
    name = api_project("tpl-rebind")
    rid = _report(client, name)
    en = _upload(client, name, "corporate.docx", fixtures["corporate_untagged.docx"])
    de = _upload(client, name, "corporate_de.docx", fixtures["corporate_untagged_de.docx"])
    bound = _bind(client, name, rid, en["file_id"])
    assert client.put(f"/api/projects/{name}/reports/{rid}/template/plan",
                      json=_plan_for(bound["outline"])).status_code == 200
    _bind(client, name, rid, de["file_id"])
    m = store.load_meta(project_storage_dir(name), rid)
    assert m.mapping_plan is None and m.template_language == "de"


def test_bind_errors(client, api_project, fixtures):
    name = api_project("tpl-bind-errors")
    rid = _report(client, name)
    tpl = _upload(client, name, "tagged_minimal.docx", fixtures["tagged_minimal.docx"])

    r = client.post(f"/api/projects/{name}/reports/ffffffffffffffff/template",
                    json={"file_id": tpl["file_id"]})
    assert r.status_code == 404 and r.json()["detail"]["error_kind"] == "report_not_found"
    r = client.post(f"/api/projects/{name}/reports/not-an-id/template",
                    json={"file_id": tpl["file_id"]})
    assert r.status_code == 400 and r.json()["detail"]["error_kind"] == "invalid_report_id"

    r = client.post(f"/api/projects/{name}/reports/{rid}/template", json={"file_id": "0" * 16})
    assert r.status_code == 404, r.text
    assert r.json()["detail"]["error_kind"] == "upload_not_found"

    csv = _upload(client, name, "notes.csv", b"a,b\n1,2\n", kind=None, mime="text/csv")
    r = client.post(f"/api/projects/{name}/reports/{rid}/template", json={"file_id": csv["file_id"]})
    assert r.status_code == 400, r.text
    assert r.json()["detail"]["error_kind"] == "template_not_a_template"

    # A docx-by-extension that is not a Word document: the reader refuses it.
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("[Content_Types].xml", "<x/>")
    bogus = _upload(client, name, "memo.docx", buf.getvalue())
    r = client.post(f"/api/projects/{name}/reports/{rid}/template", json={"file_id": bogus["file_id"]})
    assert r.status_code == 400, r.text
    assert r.json()["detail"]["error_kind"] == "template_unreadable"

    # A docx uploaded as a plain user upload still counts (MIME is enough).
    plain = _upload(client, name, "plain.docx", fixtures["corporate_untagged.docx"], kind=None)
    assert plain["kind"] == "user_upload"
    assert _bind(client, name, rid, plain["file_id"])["mode"] == "untagged"

    # Nothing bound survives in the document after the errors above except the last bind.
    assert client.get(f"/api/projects/{name}/reports/{rid}").json()["template_file_id"] == plain["file_id"]


def test_bind_is_refused_under_a_foreign_lock(
        client, api_project, fixtures, same_org_other_user):
    name = api_project("tpl-bind-lock")
    rid = _report(client, name)
    tpl = _upload(client, name, "tagged_minimal.docx", fixtures["tagged_minimal.docx"])
    assert client.post(f"/api/projects/{name}/lock").status_code == 200
    r = same_org_other_user.post(f"/api/projects/{name}/reports/{rid}/template",
                                 json={"file_id": tpl["file_id"]})
    assert r.status_code == 409, r.text[:200]
    assert r.json()["detail"]["error_kind"] == "project_locked"
    # Reads are not writes.
    assert same_org_other_user.get(f"/api/projects/{name}/reports/{rid}/template").status_code == 200
    # The holder binds.
    assert _bind(client, name, rid, tpl["file_id"])["mode"] == "tagged"


# ── export: tagged ──────────────────────────────────────────────────────────

def test_export_with_a_tagged_template_fills_the_tags(
        client, api_project, session_state, fixtures):
    name = api_project("tpl-export-tagged")
    rid = _report(client, name, session_state, title="ACME reference design")
    tpl = _upload(client, name, "tagged_minimal.docx", fixtures["tagged_minimal.docx"])
    _bind(client, name, rid, tpl["file_id"])

    r = client.post(f"/api/projects/{name}/reports/{rid}/export", json={"filename": "tagged-out"})
    assert r.status_code == 200, r.text[:400]
    meta = r.json()
    assert meta["kind"] == "agent_export" and meta["filename"] == "tagged-out.docx"
    blob = client.get(f"/api/projects/{name}/uploads/{meta['file_id']}/blob").content
    doc = _docx(blob)
    text = _all_text(doc)
    assert "ACME reference design" in text
    assert "{{" not in text and "{%" not in text
    # The looped table: the header row plus one row per fmea_top row, the
    # template row gone; `{{ row[0] }}` / `{{ row[1] }}` are the report
    # table's own first two cells.
    report = client.get(f"/api/projects/{name}/reports/{rid}").json()
    fmea_rows = report["tables"]["fmea_top"]["rows"]
    assert len(fmea_rows) == 1
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Loop", "Rank", "Component", "End"]
    assert len(table.rows) == 2
    assert table.rows[1].cells[1].text == fmea_rows[0][0]
    assert table.rows[1].cells[2].text == fmea_rows[0][1]
    # The footer carries the evidence hash.
    assert report["evidence_hash"] in text
    # The document is the template's layout, not the default writer's.
    assert "Generated by PyPSA Studio" not in text


def test_export_with_an_unfillable_tag_is_400_tagged_render_error(
        client, api_project, session_state):
    name = api_project("tpl-export-tagged-error")
    rid = _report(client, name, session_state)
    doc = Document()
    doc.add_paragraph("{{ fields.nope.text }}")
    buf = io.BytesIO()
    doc.save(buf)
    tpl = _upload(client, name, "bad.docx", buf.getvalue())
    assert _bind(client, name, rid, tpl["file_id"])["mode"] == "tagged"
    r = client.post(f"/api/projects/{name}/reports/{rid}/export", json={})
    assert r.status_code == 400, r.text[:400]
    detail = r.json()["detail"]
    assert detail["error_kind"] == "tagged_render_error"
    assert "nope" in detail["message"]


def test_export_without_a_template_is_the_default_writer(client, api_project, session_state):
    name = api_project("tpl-export-default")
    rid = _report(client, name, session_state)
    r = client.post(f"/api/projects/{name}/reports/{rid}/export", json={})
    assert r.status_code == 200, r.text[:300]
    blob = client.get(f"/api/projects/{name}/uploads/{r.json()['file_id']}/blob").content
    assert "Generated by PyPSA Studio" in _all_text(_docx(blob))


def test_export_with_a_deleted_template_upload_is_404_upload_not_found(
        client, api_project, session_state, fixtures):
    name = api_project("tpl-export-gone")
    rid = _report(client, name, session_state)
    tpl = _upload(client, name, "tagged_minimal.docx", fixtures["tagged_minimal.docx"])
    _bind(client, name, rid, tpl["file_id"])
    assert client.delete(f"/api/projects/{name}/uploads/{tpl['file_id']}").json()["deleted"] is True
    r = client.post(f"/api/projects/{name}/reports/{rid}/export", json={})
    assert r.status_code == 404, r.text[:300]
    assert r.json()["detail"]["error_kind"] == "upload_not_found"
    got = client.get(f"/api/projects/{name}/reports/{rid}/template")
    assert got.status_code == 404 and got.json()["detail"]["error_kind"] == "upload_not_found"


# ── the mapping job ─────────────────────────────────────────────────────────

def test_plan_job_runs_propose_mapping_through_the_seams_and_stores_the_plan(
        client, api_project, project_storage_dir, session_state, fixtures, fake_untagged,
        monkeypatch):
    name = api_project("tpl-plan-job")
    rid = _report(client, name, session_state)
    tpl = _upload(client, name, "corporate.docx", fixtures["corporate_untagged.docx"])
    bound = _bind(client, name, rid, tpl["file_id"])
    plan = _plan_for(bound["outline"])
    provider = _install_provider(monkeypatch, FakeProvider([_turn(plan)]))

    r = client.post(f"/api/projects/{name}/reports/{rid}/template/plan", json={"language": "en"})
    assert r.status_code == 200, r.text[:400]
    assert r.json() == {"status": "running", "report_id": rid}
    body = _poll(client, name)
    assert body["status"] == "done", body
    assert body["mode"] == "mapping" and body["report_id"] == rid
    assert body["prose_failures"] == [] and body["error"] is None
    assert body["profile_id"] and body["model"]
    assert "thread" not in body and "stop_event" not in body
    assert len(provider.requests) == 1
    assert fake_untagged.proposals == [{"language": "en", "n_headings": 7}]

    stored = store.load_meta(project_storage_dir(name), rid).mapping_plan
    assert stored is not None
    assert [e["action"] for e in stored["entries"]] == ["keep", "keep", "keep", "rename", "drop"]
    assert stored["entries"][3]["new_text"] == "4 Residual failure modes (renamed)"
    assert stored["placeholders"] == {"[Client name]": "ACME"}
    got = client.get(f"/api/projects/{name}/reports/{rid}/template").json()
    assert got["plan"] == stored and got["mode"] == "untagged"
    # A plan does not touch the document's versions.
    assert client.get(f"/api/projects/{name}/reports/{rid}").json()["version"] == 2
    # Abort after the fact is idempotent, as for a generate.
    ab = client.post(f"/api/projects/{name}/reports/generate/abort")
    assert ab.status_code == 200 and ab.json()["aborting"] is False


def test_plan_job_defaults_the_language_to_the_templates(
        client, api_project, session_state, fixtures, fake_untagged, monkeypatch):
    name = api_project("tpl-plan-lang")
    rid = _report(client, name, session_state)
    tpl = _upload(client, name, "corporate_de.docx", fixtures["corporate_untagged_de.docx"])
    bound = _bind(client, name, rid, tpl["file_id"])
    _install_provider(monkeypatch, FakeProvider([_turn(_plan_for(bound["outline"]))]))
    r = client.post(f"/api/projects/{name}/reports/{rid}/template/plan", json={})
    assert r.status_code == 200, r.text[:300]
    assert _poll(client, name)["status"] == "done"
    assert fake_untagged.proposals[-1]["language"] == "de"


def test_plan_job_falls_back_to_default_mapping_on_a_failure(
        client, api_project, project_storage_dir, session_state, fixtures, fake_untagged,
        monkeypatch):
    name = api_project("tpl-plan-fail")
    rid = _report(client, name, session_state)
    tpl = _upload(client, name, "corporate.docx", fixtures["corporate_untagged.docx"])
    _bind(client, name, rid, tpl["file_id"])
    _install_provider(monkeypatch, FakeProvider([_turn("not a plan")]))
    r = client.post(f"/api/projects/{name}/reports/{rid}/template/plan", json={})
    assert r.status_code == 200, r.text[:300]
    body = _poll(client, name)
    assert body["status"] == "done", body
    assert body["error"] is None
    assert [(f["section_id"], f["reason"]) for f in body["prose_failures"]] == [("mapping", "no_json")]
    assert body["repairs"] == 1
    stored = store.load_meta(project_storage_dir(name), rid).mapping_plan
    assert stored["entries"] == [] and "default_mapping (fake)" in stored["notes"]


def test_plan_job_errors(client, api_project, session_state, fixtures, fake_untagged, monkeypatch):
    name = api_project("tpl-plan-errors")
    rid = _report(client, name, session_state)
    _install_provider(monkeypatch, FakeProvider([]))

    r = client.post(f"/api/projects/{name}/reports/{rid}/template/plan", json={})
    assert r.status_code == 400, r.text
    assert r.json()["detail"]["error_kind"] == "no_template"

    tagged = _upload(client, name, "tagged_minimal.docx", fixtures["tagged_minimal.docx"])
    _bind(client, name, rid, tagged["file_id"])
    r = client.post(f"/api/projects/{name}/reports/{rid}/template/plan", json={})
    assert r.status_code == 400, r.text
    assert r.json()["detail"]["error_kind"] == "template_not_untagged"

    r = client.post(f"/api/projects/{name}/reports/ffffffffffffffff/template/plan", json={})
    assert r.status_code == 404 and r.json()["detail"]["error_kind"] == "report_not_found"

    # Provider 400s as generate: the seam's own kind.
    corporate = _upload(client, name, "corporate.docx", fixtures["corporate_untagged.docx"])
    _bind(client, name, rid, corporate["file_id"])
    import routers.report_jobs as jobs
    monkeypatch.setattr(jobs, "_provider_for_profile", lambda profile: (None, "missing_api_key"))
    r = client.post(f"/api/projects/{name}/reports/{rid}/template/plan", json={})
    assert r.status_code == 400 and r.json()["detail"]["error_kind"] == "missing_api_key"


def test_plan_job_shares_the_slot_with_generate(
        client, api_project, session_state, fixtures, fake_untagged, monkeypatch):
    import threading

    class _Holding(FakeProvider):
        def __init__(self, turns, hold):
            super().__init__(turns)
            self.hold = hold

        def stream(self, request):
            self.hold.wait(timeout=10)
            yield from super().stream(request)

    name = api_project("tpl-plan-inflight")
    rid = _report(client, name, session_state)
    tpl = _upload(client, name, "corporate.docx", fixtures["corporate_untagged.docx"])
    bound = _bind(client, name, rid, tpl["file_id"])
    hold = threading.Event()
    _install_provider(monkeypatch, _Holding([_turn(_plan_for(bound["outline"]))], hold))
    try:
        r = client.post(f"/api/projects/{name}/reports/{rid}/template/plan", json={})
        assert r.status_code == 200, r.text[:300]
        st = client.get(f"/api/projects/{name}/reports/generate/status").json()
        assert st["status"] == "running" and st["mode"] == "mapping"
        again = client.post(f"/api/projects/{name}/reports/{rid}/template/plan", json={})
        assert again.status_code == 409 and again.json()["detail"]["error_kind"] == "report_job_in_flight"
        gen = client.post(f"/api/projects/{name}/reports/generate", json={})
        assert gen.status_code == 409 and gen.json()["detail"]["error_kind"] == "report_job_in_flight"
    finally:
        hold.set()
    assert _poll(client, name)["status"] == "done"


# ── PUT plan ────────────────────────────────────────────────────────────────

def test_put_plan_stores_the_sanitised_plan_and_notes_what_it_dropped(
        client, api_project, project_storage_dir, fixtures, fake_untagged):
    name = api_project("tpl-put-plan")
    rid = _report(client, name)
    tpl = _upload(client, name, "corporate.docx", fixtures["corporate_untagged.docx"])
    bound = _bind(client, name, rid, tpl["file_id"])
    plan = _plan_for(bound["outline"])
    plan["entries"].append({"heading_index": 99, "action": "keep", "new_text": None,
                            "section_ids": []})
    plan["entries"].append({"heading_index": _heading_index(bound["outline"], "4.1"),
                            "action": "rename", "new_text": None, "section_ids": []})
    plan["inserted"].append({"after_heading_index": 42, "section_id": "frontier",
                             "heading": "Frontier"})
    plan["inserted"].append({"after_heading_index": _heading_index(bound["outline"], "3 "),
                             "section_id": "sizing", "heading": "Sizing"})
    plan["unmapped_sections"] = ["pipeline", "not_a_section"]

    r = client.put(f"/api/projects/{name}/reports/{rid}/template/plan", json=plan)
    assert r.status_code == 200, r.text[:400]
    stored = r.json()
    assert [e["heading_index"] for e in stored["entries"]] == [
        _heading_index(bound["outline"], p) for p in ("1 ", "2 ", "3 ", "4 ", "5 ")]
    assert stored["inserted"] == [{"after_heading_index": _heading_index(bound["outline"], "3 "),
                                   "section_id": "sizing", "heading": "Sizing"}]
    assert stored["unmapped_sections"] == ["pipeline"]
    notes = "\n".join(stored["notes"])
    assert "99" in notes and "42" in notes and "new_text" in notes and "not_a_section" in notes
    assert store.load_meta(project_storage_dir(name), rid).mapping_plan == stored
    assert client.get(f"/api/projects/{name}/reports/{rid}/template").json()["plan"] == stored


def test_put_plan_strict_refuses_a_plan_that_needs_sanitising(
        client, api_project, project_storage_dir, fixtures, fake_untagged):
    name = api_project("tpl-put-strict")
    rid = _report(client, name)
    tpl = _upload(client, name, "corporate.docx", fixtures["corporate_untagged.docx"])
    bound = _bind(client, name, rid, tpl["file_id"])
    plan = _plan_for(bound["outline"])
    plan["entries"].append({"heading_index": 99, "action": "keep", "new_text": None,
                            "section_ids": []})
    r = client.put(f"/api/projects/{name}/reports/{rid}/template/plan",
                   json={**plan, "strict": True})
    assert r.status_code == 400, r.text[:400]
    detail = r.json()["detail"]
    assert detail["error_kind"] == "invalid_mapping_plan"
    assert any("99" in n for n in detail["notes"])
    assert store.load_meta(project_storage_dir(name), rid).mapping_plan is None
    # A clean plan passes strict.
    ok = client.put(f"/api/projects/{name}/reports/{rid}/template/plan",
                    json={**_plan_for(bound["outline"]), "strict": True})
    assert ok.status_code == 200, ok.text[:300]
    assert "strict" not in ok.json()


def test_put_plan_errors(client, api_project, fixtures, fake_untagged):
    name = api_project("tpl-put-errors")
    rid = _report(client, name)
    r = client.put(f"/api/projects/{name}/reports/{rid}/template/plan", json={"entries": []})
    assert r.status_code == 400 and r.json()["detail"]["error_kind"] == "no_template"

    tagged = _upload(client, name, "tagged_minimal.docx", fixtures["tagged_minimal.docx"])
    _bind(client, name, rid, tagged["file_id"])
    r = client.put(f"/api/projects/{name}/reports/{rid}/template/plan", json={"entries": []})
    assert r.status_code == 400 and r.json()["detail"]["error_kind"] == "template_not_untagged"

    corporate = _upload(client, name, "corporate.docx", fixtures["corporate_untagged.docx"])
    _bind(client, name, rid, corporate["file_id"])
    r = client.put(f"/api/projects/{name}/reports/{rid}/template/plan",
                   json={"entries": [{"heading_index": "five", "action": "keep"}]})
    assert r.status_code == 400, r.text
    assert r.json()["detail"]["error_kind"] == "invalid_mapping_plan"
    r = client.put(f"/api/projects/{name}/reports/{rid}/template/plan",
                   json={"entries": [{"heading_index": 5, "action": "explode"}]})
    assert r.status_code == 400 and r.json()["detail"]["error_kind"] == "invalid_mapping_plan"
    r = client.put(f"/api/projects/{name}/reports/ffffffffffffffff/template/plan", json={})
    assert r.status_code == 404 and r.json()["detail"]["error_kind"] == "report_not_found"


# ── export: untagged ────────────────────────────────────────────────────────

def test_export_with_an_untagged_template_uses_the_stored_plan(
        client, api_project, session_state, fixtures, fake_untagged):
    name = api_project("tpl-export-untagged")
    rid = _report(client, name, session_state, title="Corporate report")
    tpl = _upload(client, name, "corporate.docx", fixtures["corporate_untagged.docx"])
    bound = _bind(client, name, rid, tpl["file_id"])

    # No plan yet: default_mapping is used.
    r = client.post(f"/api/projects/{name}/reports/{rid}/export", json={})
    assert r.status_code == 200, r.text[:300]
    blob = client.get(f"/api/projects/{name}/uploads/{r.json()['file_id']}/blob").content
    assert "UNTAGGED::Corporate report" in _all_text(_docx(blob))
    assert fake_untagged.rendered[-1]["notes"] == ["default_mapping (fake)"]

    plan = _plan_for(bound["outline"])
    assert client.put(f"/api/projects/{name}/reports/{rid}/template/plan", json=plan).status_code == 200
    r = client.post(f"/api/projects/{name}/reports/{rid}/export", json={})
    assert r.status_code == 200, r.text[:300]
    assert fake_untagged.rendered[-1]["entries"][3]["new_text"] == "4 Residual failure modes (renamed)"


def test_render_with_template_dispatches_on_the_mode(fixtures, fake_untagged):
    from models.report import Paragraph, ReportDocument, Section

    doc = ReportDocument(
        report_id="0123456789abcdef", version=1, title="T",
        created_at="2026-09-28T10:00:00+00:00", evidence_hash="a" * 64,
        mode="evidence_only",
        sections=[Section(section_id="executive_summary", heading="Summary", source="code",
                          status="ok", blocks=[Paragraph(md="Hello.")])],
        tables={"fmea_top": {"table_id": "fmea_top", "columns": ["a", "b"], "rows": [["1", "2"]]}},
    )
    data, mode = templates.render_with_template(doc, None, figure_bytes={})
    assert mode is None and "Generated by PyPSA Studio" in _all_text(_docx(data))
    data, mode = templates.render_with_template(doc, fixtures["tagged_minimal.docx"], figure_bytes={})
    assert mode == "tagged" and "Hello." in _all_text(_docx(data))
    data, mode = templates.render_with_template(doc, fixtures["corporate_untagged.docx"],
                                                figure_bytes={}, plan={"entries": []})
    assert mode == "untagged" and "UNTAGGED::T" in _all_text(_docx(data))
    assert fake_untagged.rendered[-1]["entries"] == []
    with pytest.raises(templates.TemplateReadError):
        templates.render_with_template(doc, b"not a docx", figure_bytes={})


# ── generate with a template ────────────────────────────────────────────────

def _draft(section_id: str, *paragraphs: str) -> dict:
    return _turn({"section_id": section_id, "paragraphs": list(paragraphs), "bullets": []})


def test_generate_with_a_template_binds_it_and_defaults_the_language(
        client, api_project, project_storage_dir, session_state, fixtures, monkeypatch):
    name = api_project("tpl-generate")
    _store_eh_report(session_state(client))
    de = _upload(client, name, "corporate_de.docx", fixtures["corporate_untagged_de.docx"])
    _install_provider(monkeypatch, FakeProvider([_draft("fmea_top", "Generator g dominiert.")]))

    r = client.post(f"/api/projects/{name}/reports/generate",
                    json={"sections": ["fmea_top"], "template_file_id": de["file_id"]})
    assert r.status_code == 200, r.text[:400]
    rid = r.json()["report_id"]
    body = _poll(client, name)
    assert body["status"] == "done", body
    doc = client.get(f"/api/projects/{name}/reports/{rid}").json()
    assert doc["language"] == "de" and doc["template_file_id"] == de["file_id"]
    m = store.load_meta(project_storage_dir(name), rid)
    assert m.template_mode == "untagged" and m.template_language == "de"
    got = client.get(f"/api/projects/{name}/reports/{rid}/template").json()
    assert got["template_file_id"] == de["file_id"] and got["mode"] == "untagged"
    assert got["language"] == "de" and got["plan"] is None

    # An explicit language wins over the template's.
    _install_provider(monkeypatch, FakeProvider([_draft("fmea_top", "Generator g dominates.")]))
    r = client.post(f"/api/projects/{name}/reports/generate",
                    json={"sections": ["fmea_top"], "template_file_id": de["file_id"],
                          "language": "en"})
    assert r.status_code == 200, r.text[:300]
    rid2 = r.json()["report_id"]
    assert _poll(client, name)["status"] == "done"
    assert client.get(f"/api/projects/{name}/reports/{rid2}").json()["language"] == "en"

    # A template without a detectable language → "en".
    tagged = _upload(client, name, "tagged_minimal.docx", fixtures["tagged_minimal.docx"])
    _install_provider(monkeypatch, FakeProvider([_draft("fmea_top", "Generator g dominates.")]))
    r = client.post(f"/api/projects/{name}/reports/generate",
                    json={"sections": ["fmea_top"], "template_file_id": tagged["file_id"]})
    assert r.status_code == 200, r.text[:300]
    rid3 = r.json()["report_id"]
    assert _poll(client, name)["status"] == "done"
    doc3 = client.get(f"/api/projects/{name}/reports/{rid3}").json()
    assert doc3["language"] == "en" and doc3["template_file_id"] == tagged["file_id"]
    assert store.load_meta(project_storage_dir(name), rid3).template_mode == "tagged"


def test_generate_without_a_template_is_unchanged(
        client, api_project, project_storage_dir, session_state, monkeypatch):
    name = api_project("tpl-generate-none")
    _store_eh_report(session_state(client))
    _install_provider(monkeypatch, FakeProvider([_draft("fmea_top", "Generator g dominates.")]))
    r = client.post(f"/api/projects/{name}/reports/generate", json={"sections": ["fmea_top"]})
    assert r.status_code == 200, r.text[:300]
    rid = r.json()["report_id"]
    assert _poll(client, name)["status"] == "done"
    doc = client.get(f"/api/projects/{name}/reports/{rid}").json()
    assert doc["language"] == "en" and doc["template_file_id"] is None
    assert store.load_meta(project_storage_dir(name), rid).template_mode is None


def test_generate_with_a_bad_template_is_refused_before_the_job_starts(
        client, api_project, session_state, monkeypatch):
    name = api_project("tpl-generate-bad")
    _store_eh_report(session_state(client))
    _install_provider(monkeypatch, FakeProvider([]))
    r = client.post(f"/api/projects/{name}/reports/generate",
                    json={"sections": ["fmea_top"], "template_file_id": "0" * 16})
    assert r.status_code == 404, r.text
    assert r.json()["detail"]["error_kind"] == "upload_not_found"
    csv = _upload(client, name, "notes.csv", b"a,b\n1,2\n", kind=None, mime="text/csv")
    r = client.post(f"/api/projects/{name}/reports/generate",
                    json={"sections": ["fmea_top"], "template_file_id": csv["file_id"]})
    assert r.status_code == 400 and r.json()["detail"]["error_kind"] == "template_not_a_template"
    assert client.get(f"/api/projects/{name}/reports/generate/status").status_code == 204


# ── sanitiser ───────────────────────────────────────────────────────────────

def test_sanitise_plan_never_raises_on_a_shape_valid_plan(fixtures):
    from models.report import ReportDocument, Section
    from services.reports.docx_reader import read_template

    outline = read_template(fixtures["corporate_untagged.docx"])
    doc = ReportDocument(
        report_id="0123456789abcdef", version=1, title="T",
        created_at="2026-09-28T10:00:00+00:00", evidence_hash="a" * 64,
        mode="evidence_only",
        sections=[Section(section_id="executive_summary", heading="Summary", source="code",
                          status="ok")])
    plan, notes = templates.sanitise_plan({
        "entries": [
            {"heading_index": 5, "action": "keep", "section_ids": ["executive_summary", "ghost"]},
            {"heading_index": 6, "action": "drop"},
        ],
        "inserted": [{"after_heading_index": 5, "section_id": "ghost", "heading": "G"}],
        "placeholders": {"[Client name]": "ACME"},
        "unmapped_sections": ["executive_summary"],
    }, outline, doc)
    assert [e["heading_index"] for e in plan["entries"]] == [5]
    assert plan["entries"][0]["section_ids"] == ["executive_summary"]
    assert plan["inserted"] == []
    assert plan["placeholders"] == {"[Client name]": "ACME"}
    assert len(notes) == 3 and plan["notes"] == notes
    with pytest.raises(templates.InvalidMappingPlan):
        templates.sanitise_plan({"entries": "nope"}, outline, doc)
    with pytest.raises(templates.InvalidMappingPlan):
        templates.sanitise_plan([], outline, doc)
