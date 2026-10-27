"""
The U2 WP9 gate's binding condition V2 (taken in WP10): the tornado POST's
centre pre-check (gate U2-WP8b Z3).

Gate note: docs/superpowers/notes/2026-10-12-u2-wp9-gate.md (S2, V2).

* (a) A failure inside the pre-check other than a `FindingsRefused` — a fork
  file a run that started meanwhile is replacing (probe P3: `OSError("NetCDF:
  HDF error")`) — answers 409 `study_running` when a run now holds the base
  context, else a typed error with `error_kind`; never an untyped 500.
* (b) The worker reuses the centre cases the pre-check valued: no option's
  centre network is valued twice per tornado.

V1 (the bill-note scan) is `test_decision_vocabulary_parity.py`'s.
"""
from __future__ import annotations

import threading

import pytest

# Money-year 01: the report tests' `_run` states the custom tariff in 2025 EUR.
from tests.test_study_report_routes import _run
from tests.test_study_tornado_routes import wait_tornado


@pytest.fixture
def studies_on(monkeypatch):
    from tests.study_s4_support import enable_studies

    yield from enable_studies(monkeypatch)


@pytest.fixture
def fake(monkeypatch):
    from services import solver_service
    from tests.test_study_runner import FakeSolver

    solver = FakeSolver()
    monkeypatch.setattr(solver_service, "run_simulation", solver)
    return solver


def _hdf_error(monkeypatch, on_read=None):
    """Probe P3: the pre-check's fork read fails as a fork being replaced does."""
    from services.study import runner as R

    def broken(row):
        if on_read is not None:
            on_read(row)
        raise OSError("NetCDF: HDF error")

    monkeypatch.setattr(R, "fork_network_from_disk", broken)


def test_a_precheck_failure_while_a_run_now_holds_the_base_is_409_study_running(
        client, api_project, studies_on, fake, monkeypatch, project_row):
    from services.study import runner as R

    name = "v2-race"
    sid = _run(client, api_project, name)
    stop = threading.Event()
    worker = threading.Thread(target=stop.wait, daemon=True)
    worker.start()

    def a_run_starts(_row):
        ctx = R.base_context(project_row(name))
        with ctx.solver_state_lock:
            ctx.solver_state[R.STUDY_KEY] = {"status": "running", "thread": worker}

    _hdf_error(monkeypatch, a_run_starts)
    try:
        r = client.post(f"/api/projects/{name}/studies/{sid}/findings/tornado", json={})
        assert r.status_code == 409, (r.status_code, r.text)
        assert r.json()["detail"]["error_kind"] == "study_running"
    finally:
        stop.set()
        worker.join(5)
        ctx = R.base_context(project_row(name))
        with ctx.solver_state_lock:
            ctx.solver_state.pop(R.STUDY_KEY, None)


def test_a_precheck_failure_without_a_run_is_a_typed_error(
        client, api_project, studies_on, fake, monkeypatch):
    name = "v2-typed"
    sid = _run(client, api_project, name)
    _hdf_error(monkeypatch)
    r = client.post(f"/api/projects/{name}/studies/{sid}/findings/tornado", json={})
    assert r.status_code == 500, (r.status_code, r.text)
    detail = r.json()["detail"]
    assert detail["error_kind"] == "findings_unreadable", detail
    assert "re-run the study" in detail["message"].lower()
    never = client.get(f"/api/projects/{name}/studies/{sid}/findings/tornado")
    assert never.status_code == 404 and never.json()["detail"]["error_kind"] == "tornado_never_run"


def test_the_worker_does_not_value_the_centre_a_second_time(
        client, api_project, studies_on, fake, monkeypatch):
    """
    V2 (b): the pre-check values every centre case (≈ the findings GET); the
    worker's `run_tornado` reuses them instead of valuing them again (the
    gate measured about 5 s of repeated engine work per tornado).
    """
    from services.study import engine_adapter as A

    name = "v2-reuse"
    sid = _run(client, api_project, name)
    real = A.option_case
    centre: list[tuple[str, int]] = []

    def counted(n, *a, **kw):
        if kw.get("centre", True):
            centre.append((kw["option_id"], id(n)))
        return real(n, *a, **kw)

    monkeypatch.setattr(A, "option_case", counted)
    r = client.post(f"/api/projects/{name}/studies/{sid}/findings/tornado", json={})
    assert r.status_code == 202, r.text
    rec = wait_tornado(client, name, sid)
    assert rec["status"] == "done", rec
    assert centre, "no centre case was valued"
    twice = sorted({c for c in centre if centre.count(c) > 1})
    assert twice == [], centre


# ── V3: recorded or fixed (gate U2-WP9 N2, N6, N8, G7) ───────────────────

def _stored_case_on_the_engine(vom: float, grant: float):
    """
    The S8 stored case as the engine would hand it to the formula view, with a
    non-zero variable O&M every operating year and a grant in year 1, and its
    cash series built from the engine's identity (operating = savings − FOM −
    VOM; cash = operating + terminal − CAPEX − replacements + grants).
    """
    import json
    import pathlib

    from models.study import InvestmentCase

    raw = json.loads((pathlib.Path(__file__).resolve().parent / "fixtures" / "pre_u2_study"
                      / "case.json").read_text(encoding="utf-8"))
    raw.pop("available", None)
    case = InvestmentCase.model_validate(raw)
    years = [y.model_copy(update={"opex_variable": 0.0 if y.year == 0 else vom}) for y in case.years]
    case = case.model_copy(update={"engine": "finance_engine", "years": years})
    n = len(years)
    op = [0.0 if y.year == 0 else y.savings - y.opex_fixed - y.opex_variable for y in years]
    grants = [grant if y.year == 1 else 0.0 for y in years]
    capex = [y.capex for y in years]
    repl = [y.replacements for y in years]
    term = [y.salvage for y in years]
    cash = [op[t] + term[t] - capex[t] - repl[t] + grants[t] for t in range(n)]
    rate = case.discount_rate
    series = {"engine": "finance_engine", "case_id": case.case_id, "rate": rate,
              "years": list(range(n)), "capex": capex, "replacements": repl, "operating": op,
              "terminal": term, "grants": grants, "cash": cash,
              "npv": sum(c / (1 + rate) ** t for t, c in enumerate(cash))}
    return case, series


def test_the_check_column_and_the_net_cash_read_variable_om_and_grants():
    """
    Gate U2-WP9 V3 (G7 survived: the golden fixture's VOM and grants are 0):
    with a non-zero variable O&M and a grant, the check column (savings − FOM
    − VOM − operating cash) is 0 in every year and the net cash flow is the
    engine's cash, so a formula that omits either goes red.
    """
    import openpyxl

    from services.study import workbook_view as WV
    from tests.xlsx_formula_eval import FormulaEvaluator

    case, series = _stored_case_on_the_engine(vom=1234.5, grant=50_000.0)
    wb = openpyxl.Workbook()
    ws = wb.active
    WV.engine_cash_flows(ws, case, series, title="CF")
    head = {ws.cell(4, j).value: j for j in range(1, ws.max_column + 1)}
    ev = FormulaEvaluator(wb)
    letter = {k: ws.cell(4, j).column_letter for k, j in head.items()}
    for t in series["years"]:
        r = 5 + t
        check = ev.value("CF", f"{letter['Check: savings - O&M - operating cash']}{r}")
        assert abs(check) <= 1e-6, (t, check)
        net = ev.value("CF", f"{letter['Net cash flow']}{r}")
        assert abs(net - series["cash"][t]) <= 1e-6, (t, net, series["cash"][t])
    assert any(v for v in series["grants"]) and case.years[1].opex_variable == 1234.5


def test_the_kpi_sheets_discount_rate_follows_the_rate_cell():
    """Gate U2-WP9 N2/V3: the KPI sheet's rate is the rate cell (live), not a static 0.07."""
    import openpyxl

    from services.study import workbook_view as WV
    from tests.xlsx_formula_eval import FormulaEvaluator

    case, series = _stored_case_on_the_engine(vom=0.0, grant=0.0)
    wb = openpyxl.Workbook()
    npv_row = WV.engine_cash_flows(wb.active, case, series, title="CF")
    WV.kpis(wb.create_sheet(), case, npv_row, title="K", cash_flows_title="CF")
    k = wb["K"]
    [row] = [r for r in k.iter_rows(min_row=2) if r[0].value == "discount_rate"]
    assert row[1].value == "='CF'!B1"
    assert row[4].value == WV.LIVE
    wb["CF"]["B1"].value = 0.10
    assert FormulaEvaluator(wb).value("K", row[1].coordinate) == pytest.approx(0.10)


def test_no_user_facing_sentence_says_library():
    """
    Gate U2-WP9 N8/V3: the word the owner retired is gone from every HELP
    sentence (backend and frontend mirror), the fidelity copy and the report
    appendix's labels; the stored KEY `library_version` stays (read-compat)
    and is shown as `defaults_version`.
    """
    import pathlib
    import re

    from services.study import report as R

    words = [s for s in list(R.HELP.values()) + [t for _p, t in R._PREFIX_HELP]
             if re.search(r"\blibrary\b", s, re.I)]
    assert words == []
    ts = (pathlib.Path(__file__).resolve().parents[2] / "frontend" / "src" / "utils"
          / "decisionVocabulary.ts").read_text(encoding="utf-8")
    fidelity = re.search(r"screening: '([^']*)'", ts).group(1)
    assert "library" not in fidelity.lower()


def test_the_provenance_shows_the_defaults_version_label():
    from services.study import render_docx, render_html
    from services.study import report_xlsx
    from tests.test_study_decision_report import _ledger, _report

    report = _report(ledger=_ledger())
    html = render_html.render_html(report, charts={})
    assert "<th>defaults_version</th>" in html and "<th>library_version</th>" not in html
    import io

    import docx
    import openpyxl

    doc = docx.Document(io.BytesIO(render_docx.render_docx(report, charts={})))
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    assert "defaults_version" in cells and "library_version" not in cells
    wb = openpyxl.load_workbook(io.BytesIO(report_xlsx.write_report_xlsx(report, {}, _ledger())))
    keys = [r[0].value for r in wb["Provenance"].iter_rows(min_row=2)]
    assert "defaults_version" in keys and "library_version" not in keys
