"""
Edge Investment Case P4 WP4.6b — the study runner and the report's finance
sections, without HTTP.

`start_investment_case` takes `build_case` injected (plan C1); these tests pass
the SAM oracle case S2 and, through the runner's test-only `layers` hook,
SAM's two tax layers, so the tax section and the post-tax headlines are
established and can be checked against `run_case` directly.
"""
from __future__ import annotations

import dataclasses
import threading
from datetime import date

import pytest
from fastapi import HTTPException

from models.finance import IC_REPORT_SECTIONS, InvestmentCaseReport
from services.finance.case import FinanceRefused, LpBasis
from services.finance.engine import run_case
from services.finance.investment_case_runner import (
    STAGES,
    assumptions_digest,
    finance_digest,
    public_record,
    staleness,
    start_investment_case,
)
from services.finance.report import assemble_finance_sections, refused_finance_report
from tests.fixtures.investment_case.sam import sam_case as S


class _Cfg:
    def __init__(self, finance):
        self.finance = finance
        self.commercial = None


def _run(case_fn, *, finance=None, layers=None, assumptions=None, stop_before=False):
    """Run the study synchronously; returns (record, stored report dict | None)."""
    finance = finance if finance is not None else S.to_finance_case("s2").inputs.model_dump(
        mode="json")
    state: dict = {"solver_config": _Cfg(finance)}
    published: dict = {}

    def publish(key, record, thread):
        published["key"] = key
        state[key] = record
        if stop_before:
            record["stop_event"].set()
        thread.start()
        thread.join(timeout=60)

    out = start_investment_case(
        None, build_case=case_fn, solver_state=state, state_update=state.update,
        publish_study=publish, layers=layers, assumptions=assumptions,
        case_hash=lambda case: "h" * 16)
    assert out["status"] == "running" and published["key"] == "investment_case"
    return state["investment_case"], state.get("investment_case_report")


def test_a_run_with_tax_layers_fills_the_five_sections():
    case = S.to_finance_case("s2")
    record, stored = _run(lambda: case, layers=S.sam_tax_layers("s2"))
    assert record["status"] == "done", record
    assert record["stages_done"] == list(STAGES)
    rep = InvestmentCaseReport.model_validate(stored)
    ref = run_case(case, layers=S.sam_tax_layers("s2")).metrics
    assert rep.project_irr_pre_tax == pytest.approx(ref["project_pre_tax_irr"])
    assert rep.project_irr_post_tax == pytest.approx(ref["project_post_tax_irr"])
    assert rep.npv_at_wacc == pytest.approx(ref["project_post_tax_npv"])
    assert rep.lcoe_finance_consistent_eur_per_mwh == pytest.approx(ref["lcoe_nominal_per_mwh"])
    assert (rep.min_dscr, rep.avg_dscr) == pytest.approx((ref["min_dscr"], ref["avg_dscr"]))
    assert (rep.llcr, rep.plcr) == pytest.approx((ref["llcr"], ref["plcr"]))
    proj = rep.sections["project"].payload
    assert proj["equity_post_tax_irr"] == pytest.approx(ref["equity_post_tax_irr"])
    assert proj["equity_post_tax_npv"] == pytest.approx(ref["equity_post_tax_npv"])
    assert proj["lifecycle_npv"] == pytest.approx(ref["lifecycle_npv"])
    assert proj["provenance"]["finance_case_hash"] == "h" * 16
    for name in ("project", "debt", "tax", "participants"):
        assert rep.completeness[name] == "ok", (name, rep.sections[name].note)
    for name in set(IC_REPORT_SECTIONS) - {"project", "debt", "tax", "participants", "gates"}:
        assert rep.completeness[name] == "skipped" and rep.sections[name].note
    # S2 carries no LP basis: the WACC gate cannot be decided (None, never False).
    assert rep.completeness["gates"] == "not_established"
    assert rep.gates.wacc_vs_discount_rate_consistent is None
    assert rep.excludes_shed_cost is True
    assert rep.ppa_price_for_target_irr_eur_per_mwh is None       # no solve_ppa asked


def test_without_a_pack_or_layers_the_tax_headlines_are_none_not_zero():
    record, stored = _run(lambda: S.to_finance_case("s2"))
    rep = InvestmentCaseReport.model_validate(stored)
    assert rep.completeness["tax"] == "not_established"
    assert "tax_pack_missing" in rep.sections["tax"].note
    assert rep.project_irr_post_tax is None and rep.npv_at_wacc is None
    assert rep.lcoe_finance_consistent_eur_per_mwh is None
    assert rep.project_irr_pre_tax is not None


def test_no_debt_leaves_the_cover_ratios_none():
    def case():
        c = S.to_finance_case("s2")
        return dataclasses.replace(c, inputs=c.inputs.model_copy(update={"debt": []}))
    _, stored = _run(case, layers=S.sam_tax_layers("s2"))
    rep = InvestmentCaseReport.model_validate(stored)
    assert (rep.min_dscr, rep.avg_dscr, rep.llcr, rep.plcr) == (None, None, None, None)


def test_the_wacc_gate_is_the_gates_block():
    def case(rate):
        c = S.to_finance_case("s2")
        return dataclasses.replace(c, lp_basis=LpBasis(discount_rate=rate))
    wacc = S.to_finance_case("s2").inputs.wacc_nominal
    _, ok = _run(lambda: case(wacc), layers=S.sam_tax_layers("s2"))
    _, off = _run(lambda: case(wacc + 0.01), layers=S.sam_tax_layers("s2"))
    ok, off = InvestmentCaseReport.model_validate(ok), InvestmentCaseReport.model_validate(off)
    assert ok.gates.wacc_vs_discount_rate_consistent is True and ok.completeness["gates"] == "ok"
    assert off.gates.wacc_vs_discount_rate_consistent is False
    assert off.sections["gates"].payload["legs"]["discount_rate"] == "differs"


def test_solve_for_ppa_fills_its_headline():
    case = S.to_finance_case("s1b", solve=True)
    _, stored = _run(lambda: case, finance=case.inputs.model_dump(mode="json"),
                     layers=S.sam_tax_layers("s1b"))
    rep = InvestmentCaseReport.model_validate(stored)
    ref = run_case(case, layers=S.sam_tax_layers("s1b")).metrics
    assert ref["solve_ppa_status"] == "ok"
    assert rep.ppa_price_for_target_irr_eur_per_mwh == pytest.approx(ref["solved_ppa_price"])


def test_cashflow_lines_carry_provenance_and_add_up():
    case = S.to_finance_case("s2")
    result = run_case(case, layers=S.sam_tax_layers("s2"))
    rep = assemble_finance_sections(result, case, assumptions_hash="a" * 16,
                                    packs={"us_federal": "p" * 16})
    lines = rep.cashflow_lines
    assert lines and all(ln.participant == "owner" and ln.provenance.mode == "pf" for ln in lines)
    years = [int(y) for y in result.tl.years]
    # The operating lines per year add up to the engine's operating net.
    for i, y in enumerate(years):
        op = sum(ln.amount for ln in lines if ln.year == y
                 and ln.value_stream in ("ppa_settlement", "fom"))
        assert op == pytest.approx(float(result.op.revenue[i] - result.op.costs[i]), abs=1e-6)
    ppa = [ln for ln in lines if ln.value_stream == "ppa_settlement"]
    assert ppa[0].provenance.contract_id == "ppa" and ppa[0].counterparty == "offtaker"
    tax = [ln for ln in lines if ln.value_stream == "corporate_tax"]
    assert tax and all(ln.provenance.pack_hash == "p" * 16 for ln in tax)
    assert sum(ln.amount for ln in tax) == pytest.approx(-float(result.tax.total_liability.sum()))
    capex = [ln for ln in lines if ln.value_stream == "capex"]
    assert sum(ln.amount for ln in capex) == pytest.approx(-float(result.op.capex.sum()))
    assert {"interest", "principal", "financing_fee"} <= {ln.value_stream for ln in lines}
    assert all(ln.amount != 0 for ln in lines)


def test_a_refusal_from_the_adapter_or_the_engine_is_a_status():
    def adapter():
        raise FinanceRefused("template_not_annual", "168 h")
    record, stored = _run(adapter)
    assert record["status"] == "refused" and record["error_code"] == "template_not_annual"
    rep = InvestmentCaseReport.model_validate(stored)
    assert rep.completeness["project"] == "not_established"
    assert rep.sections["project"].note.startswith("refused:template_not_annual")
    assert rep.project_irr_pre_tax is None

    def engine_refuses():
        c = S.to_finance_case("s2")
        return dataclasses.replace(c, inputs=c.inputs.model_copy(update={"analysis_years": None}))
    record, _ = _run(engine_refuses)
    assert record["status"] == "refused" and record["error_code"] == "analysis_years_missing"


def test_an_unexpected_error_is_a_failed_status():
    def boom():
        raise RuntimeError("adapter bug")
    record, stored = _run(boom)
    assert record["status"] == "failed" and record["error_code"] == "internal_error"
    assert "adapter bug" in record["error"] and stored is None


def test_a_stop_before_the_first_stage_stores_nothing():
    record, stored = _run(lambda: S.to_finance_case("s2"), stop_before=True)
    assert record["status"] == "aborted" and stored is None
    assert "thread" not in public_record(record) and "stop_event" not in public_record(record)


def test_the_pack_is_loaded_when_the_inputs_name_one():
    from services.finance.packs.base import load_pack

    def case():
        c = S.to_finance_case("s2")
        return dataclasses.replace(c, inputs=c.inputs.model_copy(
            update={"tax_pack_id": "us_federal"}))
    fin = case().inputs.model_dump(mode="json")
    record, stored = _run(case, finance=fin)
    pack = load_pack("us_federal", as_of=date(2030, 1, 1))
    assert record["packs"] == {"us_federal": pack.pack_hash}
    rep = InvestmentCaseReport.model_validate(stored)
    assert rep.packs == {"us_federal": pack.pack_hash}
    assert "tax_pack_missing" not in (rep.sections["tax"].note or "")


def test_synchronous_refusals():
    state = {"solver_config": _Cfg(None)}

    def call(body=None):
        return start_investment_case(body, build_case=lambda: None, solver_state=state,
                                     state_update=state.update,
                                     publish_study=lambda *a: None)
    with pytest.raises(HTTPException) as e:
        call()
    assert e.value.status_code == 422 and e.value.detail["code"] == "finance_inputs_missing"
    state["solver_config"] = _Cfg({"financial_close": "2030-01-01", "capex_phasing": [0.3]})
    with pytest.raises(HTTPException) as e:
        call()
    assert e.value.detail["code"] == "finance_inputs_invalid"
    assert ["capex_phasing"] in [x["loc"] for x in e.value.detail["errors"]] or \
        "capex_phasing" in e.value.detail["message"]
    state["solver_config"] = _Cfg({"financial_close": "2030-01-01", "tax_pack_id": "zz"})
    with pytest.raises(HTTPException) as e:
        call()
    assert e.value.detail["code"] == "tax_pack_not_found"
    with pytest.raises(HTTPException) as e:
        call({"nope": 1})
    assert e.value.status_code == 422


# ── digests and staleness ────────────────────────────────────────────────────


def test_assumptions_digest_moves_with_each_part():
    base = dict(finance={"a": 1}, commercial={"poc_link": "x", "value_flows": {"p": 1}},
                dispatch="d0", packs={"us_federal": "h1"})
    h0, parts = assumptions_digest(**base)
    assert set(parts) == {"finance", "value_flows", "commercial", "solver_config",
                          "dispatch", "packs"}
    assert assumptions_digest(**base)[0] == h0                       # deterministic
    for key, val in (("finance", {"a": 2}), ("dispatch", "d1"), ("packs", {"us_federal": "h2"}),
                     ("commercial", {"poc_link": "y", "value_flows": {"p": 1}}),
                     ("commercial", {"poc_link": "x", "value_flows": {"p": 2}})):
        assert assumptions_digest(**{**base, key: val})[0] != h0, key
    assert finance_digest(None) != finance_digest({})


def test_staleness_is_never_silently_current():
    h, parts = assumptions_digest(finance={"a": 1}, commercial=None, dispatch="d", packs={})
    stored = refused_finance_report("x", assumptions_hash=h).model_dump(mode="json")
    stored["sections"]["project"]["payload"] = {"provenance": {"inputs": parts}}
    assert staleness(None, (h, parts))["present"] is False
    assert staleness(stored, (h, parts))["stale"] is False
    s = staleness(stored, None, reason="network_busy")
    assert s["stale"] is True and s["reason"] == "network_busy"
    h2, parts2 = assumptions_digest(finance={"a": 2}, commercial=None, dispatch="d", packs={})
    s = staleness(stored, (h2, parts2))
    assert s["stale"] is True and s["changed"] == ["finance"]


def test_the_runner_record_is_published_under_the_state_lock():
    lock = threading.RLock()
    state: dict = {"solver_config": _Cfg(S.to_finance_case("s2").inputs.model_dump(mode="json"))}

    def publish(key, record, thread):
        state[key] = record
        thread.start()
        thread.join(timeout=60)
    start_investment_case(None, build_case=lambda: S.to_finance_case("s2"), solver_state=state,
                          state_update=state.update, publish_study=publish, state_lock=lock)
    assert state["investment_case"]["status"] == "done"


# ── WP4.6b review round 1 ────────────────────────────────────────────────────


@pytest.mark.parametrize("terminal", ["none", "book_value", "multiple_of_ebitda"])
def test_b2_each_years_lines_sum_to_the_post_tax_equity_cash(terminal):
    """With a counterfactual (its lines negated), a resolved terminal value
    and a levered tranche, Σ lines per year = the engine's post-tax equity
    cash — the report's drill-down reconciles to its headline cash."""
    from models.finance import DebtTranche, TerminalValueRule
    from services.finance.case import Template, TemplateLine
    from tests.test_finance_engine import LAYER, _case

    cf = (Template(first_year=2031, lines=(TemplateLine("bill", "energy_import", -100.0, "tariff",
                                                        source="counterfactual",
                                                        source_id="bill"),)),)
    tv = {"none": TerminalValueRule(), "book_value": TerminalValueRule(method="book_value"),
          "multiple_of_ebitda": TerminalValueRule(method="multiple_of_ebitda", value=5.0)}[terminal]
    loan = DebtTranche(kind="term_loan", amount=400.0, rate=0.05, tenor_years=5,
                       upfront_fee=0.01, dsra_months=0, grace_years=0, commitment_fee=0.0)
    case = _case(cf=cf, debt=[loan], fin_over={"terminal_value": tv, "analysis_years": 8})
    result = run_case(case, layers=LAYER)
    assert result.cash["equity_post_tax"] is not None, result.reasons
    rep = assemble_finance_sections(result, case, assumptions_hash="a" * 16, packs={})
    years = [int(y) for y in result.tl.years]
    for i, y in enumerate(years):
        total = sum(ln.amount for ln in rep.cashflow_lines if ln.year == y)
        assert total == pytest.approx(float(result.cash["equity_post_tax"][i]), abs=1e-6), y
    cf_lines = [ln for ln in rep.cashflow_lines
                if ln.provenance.source.startswith("counterfactual:")]
    assert cf_lines and all(ln.amount == pytest.approx(100.0) for ln in cf_lines)
    if terminal != "none":
        tv_lines = [ln for ln in rep.cashflow_lines if ln.value_stream == "terminal_value"]
        assert tv_lines and tv_lines[-1].year == years[-1]
    # The payload's terminal status is the engine's resolution, not the cash
    # stage's deferral (WP4.7 review round 2).
    assert rep.sections["project"].payload["operating_status"]["terminal"] == "ok"


def test_b6_the_counterfactual_block_and_one_cfads_definition():
    from services.finance.case import Template, TemplateLine
    from services.finance.debt import CFADS_DEFINITION
    from services.finance.export_xlsx import CFADS_DEFINITION as XLSX_CFADS
    from services.finance.export_xlsx import _counterfactual_summary
    from tests.test_finance_engine import LAYER, _case

    assert XLSX_CFADS is CFADS_DEFINITION and "replacement capex" in CFADS_DEFINITION
    cf = (Template(first_year=2031, lines=(TemplateLine("bill", "energy_import", -100.0, "tariff",
                                                        source="counterfactual",
                                                        source_id="bill"),)),)
    case = _case(cf=cf)
    rep = assemble_finance_sections(run_case(case, layers=LAYER), case,
                                    assumptions_hash="a" * 16, packs={})
    payload = rep.sections["project"].payload
    assert payload["cfads_definition"] == CFADS_DEFINITION
    assert payload["cost_of_equity"] == 0.09 and payload["wacc_nominal"] == 0.07
    block = payload["counterfactual"]
    assert block["present"] is True and block["n_lines"] == 1
    assert block["sources"] == ["counterfactual:bill"] and block["reasons"] == []
    assert "1 lines from counterfactual:bill; established" in _counterfactual_summary(block)
    # Without one: present False with its basis — never read as not established.
    plain = _case()
    block = assemble_finance_sections(run_case(plain, layers=LAYER), plain,
                                      assumptions_hash="a" * 16,
                                      packs={}).sections["project"].payload["counterfactual"]
    assert block["present"] is False and "incremental = total" in block["basis"]
    assert _counterfactual_summary(block).startswith("none — ")
    assert _counterfactual_summary(None) is None


def test_b1_the_solver_config_is_part_of_the_staleness_key():
    base = dict(finance={"a": 1}, commercial=None, dispatch="d0", packs={},
                solver={"discount_rate": 0.07, "finance": {"a": 1}, "commercial": None})
    h0, parts = assumptions_digest(**base)
    assert set(parts) == {"finance", "value_flows", "commercial", "solver_config", "dispatch",
                          "packs"}
    for field, val in (("discount_rate", 0.05), ("inflation_rate", 0.02),
                       ("auto_discount_periods", True), ("voll", 5000.0),
                       ("dsr_price_eur_per_mwh", 300.0), ("dsr_share_of_load", 0.1),
                       ("dsr_buses", ["b1"]), ("investment_periods", [2030, 2040])):
        h, p = assumptions_digest(**{**base, "solver": {**base["solver"], field: val}})
        assert h != h0 and p["solver_config"] != parts["solver_config"], field
    # A save that round-trips an int through a float field is not a change.
    same = assumptions_digest(**{**base, "solver": {**base["solver"], "mip_time_limit_s": 0}})
    assert assumptions_digest(**{**base, "solver": {**base["solver"],
                                                     "mip_time_limit_s": 0.0}})[0] == same[0]
    # finance / commercial have their own parts: not double-counted here.
    h, p = assumptions_digest(**{**base, "solver": {**base["solver"], "finance": {"a": 9}}})
    assert p["solver_config"] == parts["solver_config"]


def test_b5_an_abort_during_a_refusing_build_stores_nothing():
    state_box: dict = {}

    def adapter():
        state_box["record"]["stop_event"].set()          # the user aborts mid-build
        raise FinanceRefused("template_not_annual", "168 h")

    finance = S.to_finance_case("s2").inputs.model_dump(mode="json")
    state: dict = {"solver_config": _Cfg(finance), "investment_case_report": {"prior": True}}

    def publish(key, record, thread):
        state[key] = record
        state_box["record"] = record
        thread.start()
        thread.join(timeout=60)

    start_investment_case(None, build_case=adapter, solver_state=state,
                          state_update=state.update, publish_study=publish)
    assert state["investment_case"]["status"] == "aborted"
    assert state["investment_case_report"] == {"prior": True}     # the prior report kept

    def engine_refuses():
        c = S.to_finance_case("s2")
        state_box["record"]["stop_event"].set()
        return dataclasses.replace(c, inputs=c.inputs.model_copy(update={"analysis_years": None}))
    start_investment_case(None, build_case=engine_refuses, solver_state=state,
                          state_update=state.update, publish_study=publish)
    assert state["investment_case"]["status"] == "aborted"
    assert state["investment_case_report"] == {"prior": True}


def test_b3_b4_unknown_finance_keys_and_bad_case_ids_are_refused():
    from pydantic import ValidationError

    from models.finance import FinanceInputs
    from services.finance.investment_case_runner import InvestmentCaseRequest

    fin = S.to_finance_case("s2").inputs.model_dump(mode="json")
    with pytest.raises(ValidationError):
        FinanceInputs.model_validate({**fin, "terminalvalue": {"method": "fixed", "value": 5}})
    with pytest.raises(ValidationError):
        FinanceInputs.model_validate({**fin, "terminal_value": {"method": "none", "valu": 1}})
    with pytest.raises(ValidationError):
        FinanceInputs.model_validate({**fin, "debt": [{**fin["debt"][0], "tenor": 5}]})
    state = {"solver_config": _Cfg(fin)}
    with pytest.raises(HTTPException) as e:
        start_investment_case({"case_id": "bad\x07id"}, build_case=lambda: None,
                              solver_state=state, state_update=state.update,
                              publish_study=lambda *a: None)
    assert e.value.status_code == 422
    assert InvestmentCaseRequest(case_id="Base case v2.1-final").case_id


def test_b4_the_workbook_strips_control_characters():
    import io

    from openpyxl import load_workbook

    from services.finance.export_xlsx import build_workbook

    rep = refused_finance_report("x\x07y", "detail\x01", case_id="case",
                                 assumptions_hash="a" * 16)
    rep = rep.model_copy(update={"case_id": "c\x07ase"})
    wb = load_workbook(io.BytesIO(build_workbook(rep, project="p\x0bq")))
    about = {r[0].value: r[1].value for r in wb["About"].iter_rows(max_col=2)}
    assert about["Case"] == "case" and about["Project"] == "pq"
