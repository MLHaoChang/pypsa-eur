"""
S7, unit level: the decision report's prose guard, its disclosures, the
stale rule, and the three renderers on a report built from constructed
findings (plan S7 "Acceptance", "Mutation", "TDD evidence"; review v1 S6,
N4, N5, N7; the S5 and S6 gate carries). The route-level report of a run
is `tests/test_study_report_routes.py`.
"""
from __future__ import annotations

import io
import re
from datetime import UTC, datetime

import pytest

from models.study import (
    BaselineResult,
    DecisionStudy,
    Figure,
    Findings,
    FindingsHashes,
    OptionResult,
    ProseParagraph,
    ReportSection,
)
from services.study import findings as F
from services.study import library as L
from services.study import questions as Q
from services.study import report as R
from tests.study_s4_support import INTAKE
from tests.test_study_findings import _att, _rob

LIB = L.load_library()


def _study(name="Site battery", **kw) -> DecisionStudy:
    now = datetime(2026, 9, 29, tzinfo=UTC)
    return DecisionStudy(study_id="a" * 32, name=name, question_id="bess_at_site",
                         base_project="b" * 32, intake=kw.pop("intake", INTAKE),
                         created_at=now, updated_at=now, **kw)


def _findings(atts, rob, *, streams=None, options=("bess_1h", "bess_2h", "bess_4h"),
              hashes=None) -> Findings:
    v = F.verdict(atts, rob, streams=streams, fidelity="full_study",
                  expected=[o for o in options if Q.max_hours(Q.option(Q.BESS_AT_SITE, o))])
    opts = [OptionResult(option_id=o, label=o, solve_status="ok", system_cost=1.0, bill=1.0,
                         fidelity="full_study",
                         delta_vs_baseline={"npv": None, "payback": None, "capex": None,
                                            "co2": None,
                                            "unavailable": {k: "computed_by_the_pro_forma_s5"
                                                            for k in ("npv", "payback",
                                                                      "capex", "co2")}})
            for o in ("none", *options)]
    return Findings(options=opts, options_status="ok",
                    hashes=hashes or FindingsHashes(ledger_hash="l" * 16,
                                                    option_network_hashes={"f1": "h" * 16}),
                    baseline=BaselineResult(project_ref="p", solve_status="ok", bill=1.0),
                    verdict=v, robustness=rob or F.Robustness(), battery_attribution=atts,
                    value_streams=streams or [],
                    value_streams_status="ok" if streams else "not_established",
                    value_streams_option=v.option_id if streams else None,
                    honesty_notes=("perfect_foresight_dispatch", *F.BY_CONSTRUCTION))


def _report(name="Site battery", atts=None, rob="default", study=None, ledger=None, **kw):
    atts = atts if atts is not None else [_att("bess_1h", 1e5), _att("bess_2h", 3e5),
                                          _att("bess_4h", 2e5)]
    rob = _rob("bess_2h", (4e5, 2e5), (3.5e5, 2.5e5)) if rob == "default" else rob
    study = study or _study(name)
    ledger = ledger or L.seed_ledger(Q.BESS_AT_SITE, INTAKE, LIB)
    tariff = LIB.tariffs["de_industrial_illustrative"]
    return R.build_decision_report(R.ReportInputs(
        study=study, question=Q.BESS_AT_SITE, findings=_findings(atts, rob, **kw),
        cases={}, ledger=ledger, tariff=tariff, option_forks={"f1": "bess_2h"},
        generated_at=datetime(2026, 9, 29, tzinfo=UTC)))


# ── validate_prose (review v1 N4) ────────────────────────────────────────

def _with_prose(report, *texts, section="executive_summary"):
    s = report.sections[section]
    sections = {**report.sections, section: s.model_copy(update={
        "prose": [ProseParagraph(text=t) for t in texts]})}
    return report.model_copy(update={"sections": sections})


def test_validate_prose_rejects_a_hand_typed_number_and_names_where():
    report = _with_prose(_report(), "The battery NPV is {{battery_npv}}.",
                         "The battery is worth 4.1 M over its life.", section="economics")
    with pytest.raises(R.ProseError) as err:
        R.validate_prose(report)
    assert (err.value.section, err.value.paragraph) == ("economics", 1)
    assert "4.1" in str(err.value) and "economics" in str(err.value)


def test_validate_prose_accepts_the_allowed_units():
    report = _with_prose(_report(), "It avoids CO2 and H2, keeps N-1 security and runs 24/7.")
    out = R.validate_prose(report)
    assert out["executive_summary"] == ["It avoids CO2 and H2, keeps N-1 security and runs 24/7."]
    for bad in ("It runs 25/7.", "N-2 security.", "A CO22 cap."):
        with pytest.raises(R.ProseError):
            R.validate_prose(_with_prose(_report(), bad))


def test_an_unresolved_reference_is_refused():
    with pytest.raises(R.ProseError, match="unresolved reference"):
        R.validate_prose(_with_prose(_report(), "The NPV is {{no_such_fact}}."))


def test_a_null_fact_renders_not_established_rather_than_failing():
    report = _report()
    facts = {**report.facts, "battery_npv": Figure(
        key="battery_npv", label="Battery NPV", value=None, unit="EUR",
        engine="cash_flow_expander", fidelity="full_study",
        unavailable="currency_year_unknown")}
    report = _with_prose(report.model_copy(update={"facts": facts}),
                         "The battery NPV is {{battery_npv}}.")
    assert R.validate_prose(report)["executive_summary"] == [
        "The battery NPV is not established."]


def test_every_template_paragraph_of_a_built_report_passes_and_resolves():
    for atts, rob in (
            (None, "default"),                                   # recommended
            ([_att("bess_2h", 3e4)], _rob("bess_2h", (5e4, -2e4))),  # marginal
            ([_att(o, -5.0) for o in ("bess_1h", "bess_2h", "bess_4h")], None),
            ([_att("bess_2h", 3e5)], None)):                      # not established
        report = _report(atts=atts, rob=rob)
        rendered = R.validate_prose(report)
        assert set(rendered) == {sid for sid, _t in R.SECTIONS}
        assert not any("{{" in p for ps in rendered.values() for p in ps)


def test_every_help_sentence_is_digit_free():
    """
    DECISION (module docstring): disclosures render outside validate_prose, yet
    the study's own sentences obey the digit rule anyway.
    """
    for code, text in [*R.HELP.items(), *((p, t) for p, t in R._PREFIX_HELP)]:
        assert not re.search(r"(?<![A-Za-z])\d", text), code


# ── facts and the verdict (the KPIs are the findings') ───────────────────

def test_the_verdict_facts_are_the_findings_facts_unchanged():
    report = _report()
    v = _findings([_att("bess_1h", 1e5), _att("bess_2h", 3e5), _att("bess_4h", 2e5)],
                  _rob("bess_2h", (4e5, 2e5), (3.5e5, 2.5e5))).verdict
    for fig in v.headline_kpis:
        assert report.facts[fig.key] == fig
    assert report.sections["executive_summary"].payload["headline_kpis"] == [
        f.key for f in v.headline_kpis]
    assert report.ai_paragraphs == []


def test_positive_at_every_bound_is_not_claimed_at_the_tolerance_edge():
    """Gate S6 [N1]: a bound in [-tol, 0] is break-even, and the report says so."""
    report = _report(atts=[_att("bess_2h", 3e5)], rob=_rob("bess_2h", (4e5, -0.5)),
                     options=("bess_2h",))
    text = " ".join(R.validate_prose(report)["executive_summary"])
    assert "positive at every" not in text and "no tornado bound turns it negative" in text
    assert "break-even" in text
    assert "break_even_at_a_tornado_bound" in {d.code for d in report.required_disclosures}
    plain = _report()
    assert "break_even_at_a_tornado_bound" not in {d.code for d in plain.required_disclosures}


def test_a_bess_pv_verdict_shows_the_battery_only_best_and_the_total_npv():
    """Gate S6 [S4] carry."""
    pv = _att("bess_pv_2h", 5e5).model_copy(update={"option_npv": 1.2e6, "reference_npv": 7e5})
    report = _report(atts=[_att("bess_1h", 1e5), _att("bess_2h", 2e5), _att("bess_4h", 1.5e5),
                           pv],
                     rob=_rob("bess_pv_2h", (6e5, 4e5)),
                     options=("bess_1h", "bess_2h", "bess_4h", "bess_pv_2h"))
    assert report.facts["option_total_npv"].value == 1.2e6
    assert report.facts["battery_only_best_npv"].value == 2e5
    assert report.facts["battery_only_best_max_hours"].value == 2.0
    text = " ".join(R.validate_prose(report)["executive_summary"])
    assert "PV included" in text and "battery-only option" in text
    codes = {d.code for d in report.required_disclosures}
    assert "battery_value_against_pv_only_reference" in codes


def test_zero_sized_battery_only_options_are_stated():
    zero = [a.model_copy(update={"status": "skipped", "battery_npv": None,
                                 "battery_p_nom_mw": 0.0, "battery_payback_simple": None,
                                 "notes": ("size_zero_no_investment",),
                                 "unavailable": {k: "size_zero_no_investment" for k in (
                                     "battery_npv", "reference_npv",
                                     "battery_payback_simple")}})
            for a in (_att("bess_1h", 0.0), _att("bess_2h", 0.0), _att("bess_4h", 0.0))]
    report = _report(atts=zero, rob=None)
    codes = {d.code for d in report.required_disclosures}
    assert "battery_only_options_sized_to_zero" in codes
    # the default intake is a sector profile: the understatement is stated too
    assert "synthetic_load_understates_peak_shaving" in codes


# ── disclosures and codes ────────────────────────────────────────────────

def test_every_code_shown_has_a_sentence_and_the_tariffs_are_its_own():
    report = _report()
    tariff = LIB.tariffs["de_industrial_illustrative"]
    by_code = {d.code: d for d in report.required_disclosures}
    for code in tariff.honesty_notes:
        assert by_code[code].text == tariff.honesty_help[code]
        assert by_code[code].source == "tariff"
    assert all(d.source != "fallback" for d in report.required_disclosures)
    for code in report.honesty_notes:
        assert report.honesty_help.get(code) or code.startswith("technology_costs"), code
    assert not [c for c, t in report.honesty_help.items() if t == R._FALLBACK]


def test_the_bess_disclosure_set_is_its_own_not_the_adequacy_one():
    from services.adequacy import study_report as adequacy

    report = _report()
    texts = " ".join(d.text for d in report.required_disclosures)
    assert adequacy._ENGINE_NAMING not in texts
    assert R.DISCLOSURE_RULES and all(code in R.HELP for _tag, code in R.DISCLOSURE_RULES)


# ── stale (review v1 S6; gate S1/S4/S5 carries) ──────────────────────────

def _ledger():
    return L.seed_ledger(Q.BESS_AT_SITE, INTAKE, LIB)


def _hashed_report():
    from services.study import packs, run_hashes

    ledger = _ledger()
    hashes = FindingsHashes(ledger_hash=packs.ledger_hash(ledger),
                            intake_hash=run_hashes.intake_hash(INTAKE),
                            option_network_hashes={"f1": "h" * 16})
    return _report(hashes=hashes), ledger


def test_nothing_changed_is_not_stale():
    report, ledger = _hashed_report()
    assert R.stale_reasons(report, _study(), ledger, lambda u: "h" * 16) == []


def test_a_ledger_edit_after_assembly_is_stale():
    report, ledger = _hashed_report()
    rows = [r.model_copy(update={"value": 150.0, "status": "customised", "provenance": "user"})
            if r.key == "battery_storage_eur_per_kwh" else r for r in ledger.rows]
    edited = ledger.model_copy(update={"rows": rows})
    assert R.stale_reasons(report, _study(), edited, lambda u: "h" * 16) == [
        "ledger_changed_since_findings"]


def test_an_option_fork_edit_after_assembly_is_stale():
    report, ledger = _hashed_report()
    assert R.stale_reasons(report, _study(), ledger, lambda u: "x" * 16) == [
        "fork_changed_since_findings:bess_2h"]
    assert R.stale_reasons(report, _study(), ledger, lambda u: None) == [
        "fork_changed_since_findings:bess_2h"]


def test_an_intake_edit_after_assembly_is_stale():
    """Gate S7 [S4]: the intake is in the one rule."""
    report, ledger = _hashed_report()
    edited = _study().model_copy(update={"intake": {**INTAKE, "load": {
        **INTAKE["load"], "annual_mwh": 9000.0}}})
    assert R.stale_reasons(report, edited, ledger, lambda u: "h" * 16) == [
        "intake_changed_since_findings"]


def test_a_study_marked_stale_carries_its_reasons():
    report, ledger = _hashed_report()
    study = _study(stale=True, stale_reasons=["copied_record_findings_computed_on_the_origin_forks"])
    assert R.stale_reasons(report, study, ledger, lambda u: "h" * 16) == [
        "copied_record_findings_computed_on_the_origin_forks"]


# ── renderers ────────────────────────────────────────────────────────────

XSS = "<script>alert(1)</script>"


def test_html_escapes_a_hostile_study_name():
    from services.study.render_html import render_html

    html = render_html(_report(name=XSS), charts={})
    assert XSS not in html
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html


def test_html_renders_disclosures_before_the_first_number():
    from services.study.render_html import render_html

    report = _report()
    html = render_html(report)
    end = html.index("</section>", html.index('class="disclosures"'))
    first_fact = html.index('class="fact"')
    assert end < first_fact
    for fig in report.sections["executive_summary"].facts.values():
        value = R.format_fact(fig)
        assert value in html and html.index(value) > end, value
    assert "data:image/png;base64," in html


def test_html_shows_the_stale_banner():
    from services.study.render_html import render_html

    html = render_html(_report(), charts={}, stale=True,
                       stale_reasons=["ledger_changed_since_findings"])
    assert "This report is stale" in html
    assert html.index("This report is stale") < html.index('class="disclosures"')


def test_docx_opens_and_has_the_assumptions_table_with_edits_marked():
    import docx

    from services.study.render_docx import render_docx

    report = _report(name=XSS)
    rows = report.sections["assumptions"].payload["rows"]
    rows[0]["edited"] = True
    doc = docx.Document(io.BytesIO(render_docx(report)))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "Assumptions table" in headings and "Provenance" in headings
    assert headings.index("Read this first") < headings.index(
        "1. Executive summary and recommendation (ok)")
    tables = [t for t in doc.tables if t.rows[0].cells[0].text == "Assumption"]
    assert len(tables) == 1
    table = tables[0]
    assert len(table.rows) == len(rows) + 1
    assert table.rows[1].cells[-1].text == "edited"
    assert doc.paragraphs[0].text == XSS            # text, never markup
    assert doc.inline_shapes, "the charts are embedded"


def test_report_workbook_keeps_hostile_text_as_text():
    import openpyxl

    from services.study.report_xlsx import write_report_xlsx

    report = _report(name="=HYPERLINK(\"http://x\",\"y\")")
    wb = openpyxl.load_workbook(io.BytesIO(write_report_xlsx(report, {}, _ledger())))
    assert wb.sheetnames[:5] == ["Verdict", "Value streams", "Tornado", "Assumptions",
                                 "Provenance"]
    cell = wb["Verdict"]["B2"]
    assert cell.value.startswith("=HYPERLINK") and cell.data_type == "s"
    edited = [c.value for c in wb["Assumptions"][1]]
    assert edited[-1] == "edited"


def test_the_charts_are_pngs_of_the_report_data():
    from services.study import report_charts

    report = _report()
    charts = report_charts.render_all(report)
    assert set(charts) == {"tornado"}   # no streams, no case in this constructed report
    assert charts["tornado"][:8] == b"\x89PNG\r\n\x1a\n"


def test_a_report_section_without_payload_renders():
    from services.study.render_html import render_html

    report = _report()
    sections = {**report.sections, "robustness": ReportSection(status="skipped")}
    assert "Robustness" in render_html(report.model_copy(update={"sections": sections}),
                                       charts={})


# ── the S4/S6 carries: scope, size bound, default tariff, pending ────────

def test_a_recommendation_among_judged_options_says_so():
    report = _report(atts=[_att("bess_1h", 1e5), _att("bess_2h", 3e5)])   # bess_4h unjudged
    text = " ".join(R.validate_prose(report)["executive_summary"])
    assert "not judged" in text and "among the battery options" in text
    assert "options_not_all_judged" in {d.code for d in report.required_disclosures}


def test_size_bound_default_tariff_and_pending_drivers_are_disclosed():
    from models.study import InvestmentCase

    atts = [_att("bess_1h", 1e5), _att("bess_2h", 3e5), _att("bess_4h", 2e5)]
    rob = _rob("bess_2h", (4e5, 2e5)).model_copy(update={
        "status": "not_established", "pending": ["discount_rate"], "note": "tornado_aborted",
        "skipped": {"energy_price_level": "energy_price_level_no_effect_single_band"}})
    f = _findings(atts, _rob("bess_2h", (4e5, 2e5)))
    f = f.model_copy(update={
        "verdict": f.verdict.model_copy(update={
            "disclosures": (*f.verdict.disclosures, "size_at_upper_bound")}),
        "robustness": rob, "completeness": {}})
    case = InvestmentCase(case_id="c", study_id="a" * 32, option_id="bess_2h",
                          horizon_years=25, discount_rate=0.07,
                          honesty_notes=("tariff_not_chosen_library_default",))
    report = R.build_decision_report(R.ReportInputs(
        study=_study(), question=Q.BESS_AT_SITE, findings=f, cases={"bess_2h": case},
        ledger=_ledger(), tariff=LIB.tariffs["de_industrial_illustrative"]))
    codes = {d.code for d in report.required_disclosures}
    assert {"size_at_upper_bound", "tariff_not_chosen_library_default",
            "robustness_pending"} <= codes
    assert report.sections["robustness"].payload["pending"] == ["discount_rate"]
    assert report.sections["economics"].status == "not_established"
    assert any("discount_rate" in g for g in report.evidence_gaps)
    assert report.honesty_help["energy_price_level_no_effect_single_band"]


# ── gate S7 (BC-S7-2, BC-S7-3, N1, N2, N3, N4, N6, N9) ───────────────────

def test_the_verdicts_disclosures_are_carried_verbatim():
    """BC-S7-2: a battery NPV is never shown without the by-construction pair."""
    report = _report()                      # recommended, no case: economics skipped
    assert report.sections["economics"].status != "ok"
    codes = [d.code for d in report.required_disclosures]
    v = _findings([_att("bess_1h", 1e5), _att("bess_2h", 3e5), _att("bess_4h", 2e5)],
                  _rob("bess_2h", (4e5, 2e5), (3.5e5, 2.5e5))).verdict
    assert set(v.disclosures) <= set(codes)
    assert set(F.BY_CONSTRUCTION) <= set(codes)


def _money_headers(html: str) -> list[str]:
    return re.findall(r"<th>([^<]*(?:NPV|CAPEX|Savings|Salvage|cash flow|Swing|Annual value|"
                      r"Replacements|O&amp;M|discounted)[^<]*)</th>", html)


def test_money_table_headers_state_currency_and_year_in_html_and_docx():
    """BC-S7-3 (unit): the options and tornado tables of a constructed report."""
    import docx

    from services.study.render_docx import render_docx
    from services.study.render_html import render_html

    report = _report()
    html = render_html(report, charts={})
    heads = _money_headers(html)
    assert heads, "no money header found"
    for h in heads:
        assert "EUR" in h and "2020" in h, h
    doc = docx.Document(io.BytesIO(render_docx(report, charts={})))
    dheads = [c.text for t in doc.tables for c in t.rows[0].cells
              if re.search(r"NPV|Swing|Annual value|cash flow|CAPEX", c.text)]
    assert dheads
    for h in dheads:
        assert "EUR" in h and "2020" in h, h


def test_tornado_bounds_carry_the_drivers_unit():
    """BC-S7-3: "149.7" alone is not a value; the row's unit goes beside it."""
    from services.study.render_html import render_html

    rob = _rob("bess_2h", (4e5, 2e5))
    rob = rob.model_copy(update={"tornado": [rob.tornado[0].model_copy(
        update={"unit": "EUR/kWh", "low_value": 149.7, "high_value": 278.1})]})
    html = render_html(_report(rob=rob), charts={})
    assert "149.7 EUR/kWh" in html and "278.1 EUR/kWh" in html


def test_chart_money_axes_carry_the_currency_year(monkeypatch):
    """BC-S7-3: the axes' unit labels come with the year."""
    from services.study import report_charts

    seen = {}
    monkeypatch.setattr(report_charts, "tornado",
                        lambda rows, centre, **kw: seen.setdefault("tornado", kw) and None)
    report_charts.render_all(_report())
    assert "2020" in seen["tornado"]["unit"] and "EUR" in seen["tornado"]["unit"]


def test_a_control_character_in_the_name_does_not_break_docx_or_xlsx():
    """Gate [N9]: an untyped 500 on the user's own download."""
    import docx
    import openpyxl

    from services.study.render_docx import render_docx
    from services.study.report_xlsx import write_report_xlsx

    report = _report(name="Site\x0bA\x01B")
    doc = docx.Document(io.BytesIO(render_docx(report, charts={})))
    assert doc.paragraphs[0].text.startswith("Site")
    wb = openpyxl.load_workbook(io.BytesIO(write_report_xlsx(report, {}, _ledger())))
    assert wb["Verdict"]["B2"].value.startswith("Site")


def test_the_zero_size_sentence_blames_sector_profiles_only_on_a_synthetic_load():
    """Gate [N3]."""
    zero = [a.model_copy(update={"status": "skipped", "battery_npv": None,
                                 "battery_p_nom_mw": 0.0, "battery_payback_simple": None,
                                 "notes": ("size_zero_no_investment",),
                                 "unavailable": {k: "size_zero_no_investment" for k in (
                                     "battery_npv", "reference_npv",
                                     "battery_payback_simple")}})
            for a in (_att("bess_1h", 0.0), _att("bess_2h", 0.0), _att("bess_4h", 0.0))]
    uploaded = {**INTAKE, "load": {"source": "upload", "upload_id": "u1"}}
    now = datetime(2026, 9, 29, tzinfo=UTC)
    study = DecisionStudy(study_id="a" * 32, name="Metered", question_id="bess_at_site",
                          base_project="b" * 32, intake=uploaded, created_at=now,
                          updated_at=now)
    report = _report(atts=zero, rob=None, study=study)
    texts = " ".join(d.text for d in report.required_disclosures)
    codes = {d.code for d in report.required_disclosures}
    assert "battery_only_options_sized_to_zero" in codes
    assert "sector profile" not in texts and "synthetic_load_understates_peak_shaving" not in codes
    synthetic = " ".join(d.text for d in _report(atts=zero, rob=None).required_disclosures)
    assert "sector profiles" in synthetic


def test_the_demand_charge_foresight_note_shows_whenever_that_stream_is_not_zero():
    """Gate [N4]: not only when it is the main caveat."""
    from models.study import ValueStream

    streams = [ValueStream(key=k, label=k, annual_value=v, share=None, engine="bill_calculator",
                           unavailable={"share": "zero_savings"})
               for k, v in (("demand_charge_reduction", 1e3), ("energy_shift", 5e4))]
    report = _report(streams=streams)
    assert report.sections["executive_summary"].payload["main_caveat"] != \
        "demand_charge_perfect_foresight"
    assert "demand_charge_perfect_foresight" in {d.code for d in report.required_disclosures}
    none = [s.model_copy(update={"annual_value": 0.0}) if s.key == "demand_charge_reduction"
            else s for s in streams]
    assert "demand_charge_perfect_foresight" not in {
        d.code for d in _report(streams=none).required_disclosures}


def test_engine_labels_name_the_engine_that_made_the_figure():
    """Gate [N6]: the tolerance is a method constant, not a ledger row."""
    report = _report(atts=[_att("bess_2h", 3e5)], rob=_rob("bess_2h", (4e5, -0.5)),
                     options=("bess_2h",))
    assert report.facts["npv_tolerance_eur"].engine == "method_constant"


def test_validate_prose_rejects_a_digit_glued_to_a_unit_or_currency():
    """Gate [N1]: "EUR4M" slipped past the letter lookbehind."""
    for bad in ("It is worth EUR4M.", "A k5 bonus.", "About MW3 of power.", "USD2bn."):
        with pytest.raises(R.ProseError):
            R.validate_prose(_with_prose(_report(), bad))
    ok = "It avoids CO2 and H2, keeps N-1 security and runs 24/7."
    assert R.validate_prose(_with_prose(_report(), ok))["executive_summary"] == [ok]


def test_tariff_sentences_are_labelled_as_the_tariffs_own_in_read_this_first():
    """Gate [N2]: a supplied tariff's sentence is its author's text."""
    import docx

    from services.study.render_docx import render_docx
    from services.study.render_html import render_html

    html = render_html(_report(), charts={})
    block = html[html.index('class="disclosures"'):html.index("</section>")]
    assert "From the tariff" in block
    doc = docx.Document(io.BytesIO(render_docx(_report(), charts={})))
    assert any(p.text.startswith("From the tariff") for p in doc.paragraphs)


def test_an_unchanged_intake_resaved_by_a_browser_keeps_its_hash():
    """
    Gate S7 re-verification [S-R1]: JSON from a browser turns `4000.0` into
    `4000`, so an unchanged intake re-saved from the UI must hash the same,
    or the report reads stale and findings refuse a run nothing changed.
    Integral floats and ints are one number; a real change still differs,
    and a boolean is not a number.
    """
    from services.study import run_hashes as H

    stored = {"site": {"connection_mw": 2.0}, "load": {"annual_mwh": 4000.0},
              "pv": {"enabled": True}, "tariff_id": "de_seed", "extra": [1.0, 2.5]}
    resaved = {"site": {"connection_mw": 2}, "load": {"annual_mwh": 4000},
               "pv": {"enabled": True}, "tariff_id": "de_seed", "extra": [1, 2.5]}
    assert H.intake_hash(stored) == H.intake_hash(resaved)
    assert H.intake_hash(stored) != H.intake_hash({**resaved, "load": {"annual_mwh": 4000.5}})
    assert H.intake_hash({"pv": {"enabled": True}}) != H.intake_hash({"pv": {"enabled": 1}})


# ── gate S9 BC-S9-1: a verdict that cites "the drivers listed" lists them ──

def _marginal_on_the_demand_charge():
    from models.study import Robustness, TornadoRow

    rows = [TornadoRow(key="demand_charge_price", label="Demand charge on peak import",
                       unit="EUR/MW/month", centre_value=9000.0, low_value=6300.0,
                       high_value=11700.0, npv_low=-2.2e4, npv_high=5.6e5, swing=5.8e5,
                       evaluation="redispatch"),
            TornadoRow(key="discount_rate", label="Real discount rate", unit="per unit",
                       centre_value=0.07, low_value=0.049, high_value=0.091,
                       npv_low=4.4e5, npv_high=1.4e5, swing=3.0e5, evaluation="rate_only")]
    rob = Robustness(status="ok", tornado=rows, option_id="bess_1h", npv_centre=2.7e5)
    report = _report(atts=[_att("bess_1h", 2.7e5), _att("bess_2h", 0.0, p=0.0, status="skipped"),
                           _att("bess_4h", 0.0, p=0.0, status="skipped")], rob=rob)
    es = report.sections["executive_summary"]
    assert es.payload["verdict_class"] == "marginal", es.payload
    assert es.payload["drivers"] == ["demand_charge_price"]
    return report


def _section_text(html: str, sid: str) -> str:
    start = html.index(f'id="{sid}"')
    return html[start:html.index("</section>", start)]


def test_a_marginal_verdict_lists_its_drivers_by_label_in_html_docx_and_xlsx():
    """
    The marginal sentence ends "... on the drivers listed." (gate S9 [S1]): the
    executive summary must then list them, by the label the tornado table
    uses, in all three formats — not only in the section payload.
    """
    import docx
    import openpyxl

    from services.study.render_docx import render_docx
    from services.study.render_html import render_html
    from services.study.report_xlsx import write_report_xlsx

    report = _marginal_on_the_demand_charge()
    label = "Demand charge on peak import"
    summary = _section_text(render_html(report, charts={}), "executive_summary")
    assert "on the drivers listed" in summary
    assert label in summary and "Real discount rate" not in summary

    doc = docx.Document(io.BytesIO(render_docx(report, charts={})))
    texts = [p.text for p in doc.paragraphs]
    first = next(i for i, t in enumerate(texts) if t.startswith("1. Executive summary"))
    second = next(i for i, t in enumerate(texts) if t.startswith("2. "))
    assert any(label in t for t in texts[first:second]), texts[first:second]

    wb = openpyxl.load_workbook(io.BytesIO(write_report_xlsx(report, {}, _ledger())))
    cells = [c.value for row in wb["Verdict"].iter_rows() for c in row if c.value]
    assert "verdict_driver" in cells and label in cells



# ── gate S9 non-binding fixes: -0, the tested range, the maturity advice ──

def test_a_negative_zero_size_renders_as_zero():
    """An LP's -0.0 MW reads as a negative size to a novice (gate S9 [N4])."""
    from services.study.render_html import num

    assert num(-0.0, 3) == num(0.0, 3) == "0"
    assert num(-0.0) == "0"
    assert num(-4e-7, 3) == "0" and num(-4e-7) == "0"   # solver noise below display
    assert num(-0.25, 3) == "-0.25"                     # a real negative stays


def _edited_storage_ledger():
    from services.study import ledger as LG

    return LG.apply_user_row(_ledger(), "battery_storage_eur_per_kwh", 450.0,
                             unit="EUR/kWh", changed_by="test")


def test_the_assumptions_show_the_range_the_tornado_tested_and_mark_it_recentred():
    """
    An edited value outside the library band is tested on a range re-centred
    on it (`findings.bounds_for`); the Assumptions table must show THAT range,
    marked, not the library's (gate S9 [S4]) — in the payload, the HTML, the
    DOCX and the report and case workbooks.
    """
    import docx
    import openpyxl

    from services.study import findings as F
    from services.study.proforma_xlsx import _assumptions
    from services.study.render_docx import render_docx
    from services.study.render_html import num, render_html
    from services.study.report_xlsx import write_report_xlsx

    ledger = _edited_storage_ledger()
    row = next(r for r in ledger.rows if r.key == "battery_storage_eur_per_kwh")
    low, high, notes = F.bounds_for(row)
    assert notes == ("range_recentred_on_user_value",) and low < 450.0 < high
    report = _report(ledger=ledger)
    rows = {r["key"]: r for r in report.sections["assumptions"].payload["rows"]}
    got = rows["battery_storage_eur_per_kwh"]
    assert (got["tested_low"], got["tested_high"]) == (low, high)
    assert got["range_note"] == "range_recentred_on_user_value"
    inside = rows["battery_inverter_eur_per_kw"]          # a default: its own band
    assert (inside["tested_low"], inside["range_note"]) == (inside["range_low"], None)

    shown = f"{num(low, 4)} to {num(high, 4)}"
    html = render_html(report, charts={})
    line = next(tr for tr in html.split("<tr") if "battery_storage_eur_per_kwh" in tr
                and 'class="num"' in tr)
    assert shown in line and "re-centred" in line, line
    doc = docx.Document(io.BytesIO(render_docx(report, charts={})))
    table = next(t for t in doc.tables if t.rows[0].cells[0].text == "Assumption")
    cells = next(r.cells for r in table.rows if r.cells[1].text == "battery_storage_eur_per_kwh")
    assert shown in cells[4].text and "re-centred" in cells[4].text

    for wb in (openpyxl.load_workbook(io.BytesIO(write_report_xlsx(report, {}, ledger))),
               _assumptions_book(_assumptions, ledger)):
        ws = wb["Assumptions"]
        head = [c.value for c in ws[1]]
        line = next([c.value for c in r] for r in ws.iter_rows(min_row=2)
                    if r[0].value == "battery_storage_eur_per_kwh")
        assert line[head.index("tested_low")] == pytest.approx(low, rel=1e-12)
        assert line[head.index("tested_high")] == pytest.approx(high, rel=1e-12)
        assert line[head.index("tested_range_note")] == "range_recentred_on_user_value"


def _assumptions_book(writer, ledger):
    import openpyxl

    wb = openpyxl.Workbook()
    writer(wb.active, ledger)
    return wb


def _maturity(*reasons):
    from models.study import StudyMaturity

    return StudyMaturity.model_validate({"status": "ok", "class": "screening",
                                         "reasons": list(reasons)})


def _advice(report) -> str:
    return " ".join(p.text for p in report.sections["limitations"].prose)


def test_the_maturity_advice_names_only_what_still_holds_the_badge():
    """
    Gate S9 [N5]: with a metered load and the equipment quotes supplied, the
    advice must not ask for them; it names what holds the badge (here the
    illustrative tariff and the defaults still in use).
    """
    upload = {**INTAKE, "load": {"source": "upload", "upload_id": "u1"}}
    held = _study(intake=upload, maturity=_maturity(
        "discount_rate: default (library)",
        "tariff: tariff (illustrative, library)"))
    text = _advice(_report(study=held))
    assert "metered load" not in text and "equipment costs" not in text, text
    assert "tariff" in text and "defaults" in text and "weather" in text

    synthetic = _study(maturity=_maturity(
        "battery_storage_eur_per_kwh: default (library)",
        "load: sector_profile (upload metered load to raise maturity)"))
    text = _advice(_report(study=synthetic))
    assert "metered load" in text and "defaults" in text and "tariff" not in text, text


def test_the_report_hides_a_zero_stream_the_tariff_has_no_item_for():
    """
    Gate U2-S1 (visible now), both ways: a zero stream with no tariff item
    leaves the drivers payload (the waterfall, the docx and html tables); a
    zero stream the tariff has an item for stays.
    """
    from models.study import ValueStream

    def vs(key, value, itemised):
        return ValueStream(key=key, label=key, annual_value=value, share=None,
                           engine="bill_calculator", itemised=itemised,
                           unavailable={"share": "zero_savings"})

    streams = [vs("energy_shift", 5e4, True), vs("taxes_levies", 0.0, False),
               vs("fixed", 0.0, True)]
    payload = _report(streams=streams).sections["drivers"].payload
    assert [s["key"] for s in payload["value_streams"]] == ["energy_shift", "fixed"]

