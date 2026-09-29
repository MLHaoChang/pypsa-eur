"""
Build the report-template fixtures for WP8 (``services/reports/docx_reader``)
with python-docx — no ``.docx`` binary is committed; the tests call
``build_all(tmp_path)`` and a human regenerates them with::

    cd pypsa-gui/backend && python tests/fixtures/report_templates/_build.py OUT_DIR

``build_all(out_dir)`` writes ``tagged_minimal.docx``, ``corporate_untagged.docx``
and ``with_textbox.docx``; ``build_all(out_dir, language="de")`` writes
``corporate_untagged_de.docx`` (the corporate template with German headings).
See README.md next to this file for what each fixture contains.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

TAGGED = "tagged_minimal.docx"
CORPORATE = "corporate_untagged.docx"
CORPORATE_DE = "corporate_untagged_de.docx"
WITH_TEXTBOX = "with_textbox.docx"

#: The corporate template's words per language. The German variant proves
#: ``detect_language`` on real headings rather than on a word list.
_TEXT: dict[str, dict[str, object]] = {
    "en": {
        "title": "[Client name] — Energy Hub Reference Design",
        "prepared": "Prepared for [Client name]",
        "date": "<Date>",
        "headings": [
            "1 Executive Summary",
            "2 Introduction",
            "3 Availability Target",
            "4 Residual Failure Modes",
            "5 Lorem ipsum",
        ],
        "subheading": "4.1 Critical components",
        "body": "Body text of the section. XXX",
        "table_header": ["Component", "Failure mode", "Criticality"],
        "header": "ACME Energy Consulting",
        "footer": "Confidential",
    },
    "de": {
        "title": "[Kundenname] — Referenzdesign Energy Hub",
        "prepared": "Erstellt für [Kundenname]",
        "date": "<Datum>",
        "headings": [
            "1 Zusammenfassung",
            "2 Einleitung",
            "3 Verfügbarkeitsziel",
            "4 Verbleibende Ausfallarten",
            "5 Lorem ipsum",
        ],
        "subheading": "4.1 Kritische Komponenten",
        "body": "Text des Abschnitts. XXX",
        "table_header": ["Komponente", "Ausfallart", "Kritikalität"],
        "header": "ACME Energy Consulting",
        "footer": "Vertraulich",
    },
}


def _toc_paragraph(doc) -> None:
    """A ``w:fldSimple`` TOC field, the shape Word writes for a simple TOC."""
    p = doc.add_paragraph()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), r'TOC \o "1-3" \h \z \u')
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "Table of contents (update field)"
    r.append(t)
    fld.append(r)
    p._p.append(fld)


def build_tagged(path: Path) -> Path:
    doc = Document()
    doc.add_heading("{{ meta.title }}", 0)
    doc.add_heading("Summary", 1)
    p = doc.add_paragraph()
    # Word splits a tag across runs when the user edits it piecemeal; the
    # reader must read it whole.
    p.add_run("{{ fields.")
    p.add_run("executive_summary")
    p.add_run(".text }}")
    table = doc.add_table(rows=2, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells,
                          ["Loop", "Rank", "Component", "End"]):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, [
        "{% for row in tables.fmea_top.rows %}",
        "{{ row[0] }}",
        "{{ row[1] }}",
        "{% endfor %}",
    ]):
        cell.text = text
    doc.add_paragraph("Evidence hash in the footer.")
    footer = doc.sections[0].footer
    footer.paragraphs[0].text = "Evidence {{ meta.evidence_hash }}"
    doc.save(str(path))
    return path


def _build_corporate_document(language: str):
    words = _TEXT[language]
    doc = Document()
    doc.add_heading(str(words["title"]), 0)
    doc.add_paragraph(str(words["prepared"]))
    doc.add_paragraph(str(words["date"]))
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    _toc_paragraph(doc)
    for i, heading in enumerate(words["headings"]):  # type: ignore[arg-type]
        doc.add_heading(heading, 1)
        doc.add_paragraph(str(words["body"]))
        if i == 3:
            doc.add_heading(str(words["subheading"]), 2)
            table = doc.add_table(rows=2, cols=3)
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells,
                                  words["table_header"]):  # type: ignore[arg-type]
                cell.text = text
            for cell in table.rows[1].cells:
                cell.text = "…"
    section = doc.sections[0]
    section.header.paragraphs[0].text = str(words["header"])
    section.footer.paragraphs[0].text = str(words["footer"])
    return doc


def build_corporate(path: Path, language: str = "en") -> Path:
    _build_corporate_document(language).save(str(path))
    return path


def _add_textbox(doc, text: str) -> None:
    """A VML text box (``w:pict/v:shape/v:textbox/w:txbxContent``) in a run."""
    p = doc.add_paragraph()
    r = p.add_run()
    pict = parse_xml(
        f'<w:pict {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml">'
        '<v:shape id="tb1" type="#_x0000_t202" style="width:200pt;height:60pt">'
        "<v:textbox><w:txbxContent><w:p><w:r><w:t>"
        f"{text}"
        "</w:t></w:r></w:p></w:txbxContent></v:textbox>"
        "</v:shape></w:pict>"
    )
    r._r.append(pict)


def build_with_textbox(path: Path) -> Path:
    doc = _build_corporate_document("en")
    _add_textbox(doc, "Text in a floating box")
    doc.save(str(path))
    return path


def build_all(out_dir: Path | str, language: str = "en") -> dict[str, Path]:
    """
    Write the fixtures into ``out_dir`` and return ``{name: path}``.

    ``language="en"`` (default) writes the three English fixtures;
    ``language="de"`` writes the German corporate variant only.
    """
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    if language == "de":
        return {CORPORATE_DE: build_corporate(out / CORPORATE_DE, "de")}
    if language != "en":
        raise ValueError(f"unknown fixture language {language!r}")
    return {
        TAGGED: build_tagged(out / TAGGED),
        CORPORATE: build_corporate(out / CORPORATE, "en"),
        WITH_TEXTBOX: build_with_textbox(out / WITH_TEXTBOX),
    }


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("usage: _build.py OUT_DIR")
    written = build_all(sys.argv[1])
    written.update(build_all(sys.argv[1], language="de"))
    for name, path in written.items():
        print(f"{name}\t{path}")
