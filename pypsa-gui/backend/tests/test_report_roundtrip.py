"""
WP12 — the round trip: an edited Word copy of a report read back into
`ReportDocument` blocks and merged as a new version
(docs/superpowers/plans/2026-09-28-llm-report-generation-increment-1.md,
"Increment 3 — round trip").

Every input is built the way the feature meets it: a `ReportDocument`
rendered by the one writer (`render_document_docx`), then edited with
python-docx the way Word would (a paragraph changed, a bullet added, a table
cell changed, a comment on a run, a tracked insertion and deletion written
as `w:ins` / `w:del`).
"""
from __future__ import annotations

import copy
import io
import struct
import zlib

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from models.report import (
    Bullets,
    Callout,
    Field,
    Figure,
    FigureRef,
    Paragraph,
    ReportDocument,
    Section,
    SectionAudit,
    Table,
    TableRef,
)
from services.reports.docx_reader import TemplateReadError
from services.reports.docx_writer import render_document_docx
from services.reports.roundtrip import (
    RoundTripResult,
    RoundTripSection,
    merge_round_trip,
    read_edited_docx,
)


# ── fixtures ─────────────────────────────────────────────────────────────


def _png_1x1() -> bytes:
    def chunk(kind: bytes, body: bytes) -> bytes:
        return (struct.pack(">I", len(body)) + kind + body
                + struct.pack(">I", zlib.crc32(kind + body) & 0xFFFFFFFF))

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 0, 0, 0, 0)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr)
            + chunk(b"IDAT", zlib.compress(b"\x00\x80")) + chunk(b"IEND", b""))


_PNG = _png_1x1()
_FIGURES = {"fmea_pareto": _PNG}


def _document(*, unverified: list[str] | None = None) -> ReportDocument:
    """The writer test's shape: every block type, three sections, a table, a figure."""
    return ReportDocument(
        report_id="0123456789abcdef", version=2, title="Client study report",
        created_at="2026-09-28T10:00:00+00:00", evidence_hash="c" * 64,
        profile_id="local-llama", model="llama-3.1-8b", mode="generated",
        sections=[
            Section(section_id="executive_summary", heading="Executive summary",
                    source="code", status="ok", blocks=[
                        Callout(kind="gap", text="the demand profile is frozen"),
                        Callout(kind="disclosure", text="excludes load-shedding cost"),
                        TableRef(table_id="headline"),
                    ]),
            Section(section_id="fmea_top", heading="Residual failure modes",
                    source="llm", status="ok", note="prose checked by hand",
                    audit=SectionAudit(unverified=unverified or [],
                                       verified=[{"text": "100,000 €/yr",
                                                  "path": "/fmea/top/0/crit"}]),
                    blocks=[
                        Paragraph(md="The **top mode** is *backup* at `100,000 €/yr` "
                                     "([source](https://example.org/x))."),
                        Paragraph(md="First chunk.\n\nSecond chunk of the same block."),
                        Bullets(items=["first point", "second point"]),
                        TableRef(table_id="fmea_top", caption="Ranked modes"),
                        FigureRef(figure_id="fmea_pareto", caption="Pareto"),
                        Field(key="Engine", value="copt"),
                        TableRef(table_id="does_not_exist"),
                    ]),
            Section(section_id="certification", heading="Certification",
                    source="code", status="skipped", note="budget exhausted",
                    blocks=[Callout(kind="not_established",
                                    text="this section was not run in this study "
                                         "(budget exhausted).")]),
        ],
        tables={
            "headline": Table(table_id="headline", columns=["Item", "Value"],
                              rows=[["MC LOLE (h/yr)", "not established"],
                                    ["Cost at target", "1,234,568 €"]],
                              caption="Headline results"),
            "fmea_top": Table(table_id="fmea_top",
                              columns=["Rank", "Name", "Criticality"],
                              rows=[["1", "backup", "100,000 €/yr"],
                                    ["2", "import", "25,000 €/yr"],
                                    ["3", "spare", "not established"]],
                              source_path="/fmea/top"),
        },
        figures={"fmea_pareto": Figure(figure_id="fmea_pareto",
                                       png_file="figures/fmea_pareto.png",
                                       caption="criticality by mode")},
    )


def _rendered(doc: ReportDocument | None = None) -> bytes:
    return render_document_docx(doc or _document(), figure_bytes=_FIGURES)


def _open(data: bytes):
    return Document(io.BytesIO(data))


def _save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _paragraph_containing(doc, needle: str):
    for p in doc.paragraphs:
        if needle in p.text:
            return p
    raise AssertionError(f"no paragraph contains {needle!r}")


def _strip_bookmarks(doc) -> None:
    body = doc.element.body
    for el in list(body.iter(qn("w:bookmarkStart"))) + list(body.iter(qn("w:bookmarkEnd"))):
        el.getparent().remove(el)


def _run_element(text: str, *, deleted: bool = False):
    r = OxmlElement("w:r")
    t = OxmlElement("w:delText" if deleted else "w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def _tracked(kind: str, text: str, change_id: str):
    el = OxmlElement(f"w:{kind}")
    el.set(qn("w:id"), change_id)
    el.set(qn("w:author"), "Reviewer")
    el.set(qn("w:date"), "2026-09-29T10:00:00Z")
    el.append(_run_element(text, deleted=(kind == "del")))
    return el


def _section(result: RoundTripResult, section_id: str) -> RoundTripSection:
    for s in result.sections:
        if s.section_id == section_id:
            return s
    raise AssertionError(f"section {section_id!r} not in the result: "
                         f"{[s.section_id for s in result.sections]}")


def _by_id(doc: ReportDocument, section_id: str) -> Section:
    return next(s for s in doc.sections if s.section_id == section_id)


# ── an unedited copy ─────────────────────────────────────────────────────


def test_unedited_render_round_trips_unchanged():
    base = _document()
    result = read_edited_docx(_rendered(base), base)
    assert [s.section_id for s in result.sections] == [
        "executive_summary", "fmea_top", "certification"]
    assert all(s.changed is False for s in result.sections), [
        (s.section_id, s.blocks) for s in result.sections if s.changed]
    assert result.unmatched == []
    assert result.comments_global == []
    assert result.accepted_tracked_changes == 0
    assert result.edited_tables == {}
    for s in result.sections:
        assert s.blocks == _by_id(base, s.section_id).blocks


def test_merge_of_an_unedited_copy_is_byte_identical_except_version_and_created_at():
    base = _document()
    merged = merge_round_trip(base, read_edited_docx(_rendered(base), base))
    assert merged.version == base.version + 1
    assert merged.created_at != base.created_at
    assert merged.created_at.endswith("+00:00")
    left = base.model_dump(exclude={"version", "created_at"})
    right = merged.model_dump(exclude={"version", "created_at"})
    assert left == right
    assert merged.model_dump_json(exclude={"version", "created_at"}) == \
        base.model_dump_json(exclude={"version", "created_at"})


# ── edits ────────────────────────────────────────────────────────────────


def test_edited_paragraph_marks_the_section_changed_and_keeps_inline_markdown():
    base = _document()
    doc = _open(_rendered(base))
    p = _paragraph_containing(doc, "top mode")
    assert p.runs[0].text == "The "
    p.runs[0].text = "The edited "
    result = read_edited_docx(_save(doc), base)
    sec = _section(result, "fmea_top")
    assert sec.changed is True
    first = sec.blocks[0]
    assert isinstance(first, Paragraph)
    assert first.md == ("The edited **top mode** is *backup* at `100,000 €/yr` "
                        "([source](https://example.org/x)).")
    # The other sections did not move.
    assert _section(result, "executive_summary").changed is False
    assert _section(result, "certification").changed is False
    merged = merge_round_trip(base, result)
    sec_m = _by_id(merged, "fmea_top")
    assert sec_m.source == "user_edit"
    assert sec_m.blocks[0].md == first.md
    assert sec_m.audit.unverified == [] and sec_m.audit.verified == []
    assert sec_m.status == "ok"
    # Untouched sections are byte-identical.
    assert _by_id(merged, "executive_summary") == _by_id(base, "executive_summary")
    assert _by_id(merged, "certification") == _by_id(base, "certification")


def test_added_bullet_is_merged_into_the_bullets_block():
    base = _document()
    doc = _open(_rendered(base))
    last = _paragraph_containing(doc, "second point")
    assert last.style.name == "List Bullet"
    new_p = copy.deepcopy(last._p)
    last._p.addnext(new_p)
    added = doc.paragraphs[[p._p for p in doc.paragraphs].index(new_p)]
    for r in added.runs[1:]:
        r._r.getparent().remove(r._r)
    added.runs[0].text = "third point"
    result = read_edited_docx(_save(doc), base)
    sec = _section(result, "fmea_top")
    assert sec.changed is True
    bullets = [b for b in sec.blocks if isinstance(b, Bullets)]
    assert len(bullets) == 1
    assert bullets[0].items == ["first point", "second point", "third point"]
    # The paragraphs around it are untouched: the base's two-chunk paragraph
    # comes back as the two chunks the writer wrote.
    assert [b.md for b in sec.blocks if isinstance(b, Paragraph)][1:] == [
        "First chunk.", "Second chunk of the same block."]


def test_edited_table_cell_becomes_an_edited_table_in_the_result():
    base = _document()
    doc = _open(_rendered(base))
    table = next(t for t in doc.tables if t.cell(0, 0).text == "Rank")
    table.cell(1, 1).text = "backup (edited)"
    result = read_edited_docx(_save(doc), base)
    sec = _section(result, "fmea_top")
    assert sec.changed is True
    new_id = "fmea_top_edited_v3"
    refs = [b for b in sec.blocks if isinstance(b, TableRef)]
    assert refs[0].table_id == new_id
    assert refs[0].caption == "Ranked modes"
    assert refs[1].table_id == "does_not_exist"  # the missing-table line survives
    assert list(result.edited_tables) == [new_id]
    edited = result.edited_tables[new_id]
    assert edited.columns == ["Rank", "Name", "Criticality"]
    assert edited.rows[0] == ["1", "backup (edited)", "100,000 €/yr"]
    assert edited.rows[1:] == [["2", "import", "25,000 €/yr"],
                               ["3", "spare", "not established"]]
    # The other section's table is unchanged and keeps its original ref.
    exec_blocks = _section(result, "executive_summary").blocks
    assert exec_blocks[-1] == TableRef(table_id="headline")
    merged = merge_round_trip(base, result)
    assert new_id in merged.tables and merged.tables[new_id] == edited
    assert "fmea_top" in merged.tables  # the original is kept for earlier versions
    assert [b.table_id for b in _by_id(merged, "fmea_top").blocks
            if isinstance(b, TableRef)][0] == new_id


# ── comments ─────────────────────────────────────────────────────────────


def test_comment_on_a_run_becomes_the_sections_pending_instruction():
    base = _document()
    doc = _open(_rendered(base))
    p = _paragraph_containing(doc, "first point")
    doc.add_comment(runs=p.runs[0], text="Please expand on the second mode",
                    author="Reviewer")
    result = read_edited_docx(_save(doc), base)
    sec = _section(result, "fmea_top")
    assert sec.comments == ["Please expand on the second mode"]
    assert sec.changed is False  # a comment is not an edit
    assert result.comments_global == []
    merged = merge_round_trip(base, result)
    m = _by_id(merged, "fmea_top")
    assert m.pending_instruction == "Please expand on the second mode"
    assert m.comments == ["Please expand on the second mode"]
    assert m.source == "llm" and m.blocks == _by_id(base, "fmea_top").blocks
    assert _by_id(merged, "executive_summary").pending_instruction is None


def test_two_comments_join_into_one_instruction():
    base = _document()
    doc = _open(_rendered(base))
    p = _paragraph_containing(doc, "first point")
    q = _paragraph_containing(doc, "second point")
    doc.add_comment(runs=p.runs[0], text="one", author="A")
    doc.add_comment(runs=q.runs[0], text="two", author="B")
    merged = merge_round_trip(base, read_edited_docx(_save(doc), base))
    assert _by_id(merged, "fmea_top").comments == ["one", "two"]
    assert _by_id(merged, "fmea_top").pending_instruction == "one; two"


def test_comment_before_the_first_section_is_global():
    base = _document()
    doc = _open(_rendered(base))
    title = doc.paragraphs[0]
    assert title.style.name == "Title"
    doc.add_comment(runs=title.runs[0], text="Rename the report", author="R")
    result = read_edited_docx(_save(doc), base)
    assert result.comments_global == ["Rename the report"]
    assert all(s.comments == [] for s in result.sections)


# ── tracked changes ──────────────────────────────────────────────────────


def test_tracked_insertion_is_kept_and_tracked_deletion_dropped():
    base = _document()
    doc = _open(_rendered(base))
    p = _paragraph_containing(doc, "First chunk.")
    p._p.append(_tracked("ins", " Inserted by review.", "901"))
    q = _paragraph_containing(doc, "Second chunk")
    q._p.append(_tracked("del", " Deleted by review.", "902"))
    result = read_edited_docx(_save(doc), base)
    assert result.accepted_tracked_changes == 2
    sec = _section(result, "fmea_top")
    mds = [b.md for b in sec.blocks if isinstance(b, Paragraph)]
    assert "First chunk. Inserted by review." in mds
    assert "Second chunk of the same block." in mds
    assert not any("Deleted by review" in md for md in mds)
    assert sec.changed is True


# ── section matching ─────────────────────────────────────────────────────


def test_bookmarks_stripped_still_matches_sections_by_heading_text():
    base = _document()
    doc = _open(_rendered(base))
    _strip_bookmarks(doc)
    data = _save(doc)
    assert b"sec:" not in _open(data).element.body.xml.encode()
    result = read_edited_docx(data, base)
    assert [s.section_id for s in result.sections] == [
        "executive_summary", "fmea_top", "certification"]
    assert all(not s.changed for s in result.sections)
    assert result.unmatched == []


def test_renumbered_and_recased_headings_still_match():
    base = _document()
    doc = _open(_rendered(base))
    _strip_bookmarks(doc)
    for p in doc.paragraphs:
        if p.text == "Residual failure modes":
            p.runs[0].text = "2. RESIDUAL FAILURE MODES"
        elif p.text == "Executive summary":
            p.runs[0].text = "1 Executive Summary"
    result = read_edited_docx(_save(doc), base)
    assert [s.section_id for s in result.sections] == [
        "executive_summary", "fmea_top", "certification"]


def test_headings_that_match_nothing_are_unmatched():
    base = _document()
    doc = _open(_rendered(base))
    _strip_bookmarks(doc)
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading"):
            p.runs[0].text = "Chapter " + p.text[:1]
    result = read_edited_docx(_save(doc), base)
    assert all(s.section_id is None for s in result.sections)
    assert result.unmatched == ["Chapter E", "Chapter R", "Chapter C"]
    # The unmatched chunks carry their content so the route can show it.
    assert any(isinstance(b, Paragraph) and "top mode" in b.md
               for b in result.sections[1].blocks)


def test_numbers_to_check_appendix_is_not_a_section():
    base = _document(unverified=["4.0 h/yr"])
    data = _rendered(base)
    text = "\n".join(p.text for p in _open(data).paragraphs)
    assert "Numbers to check" in text
    result = read_edited_docx(data, base)
    assert [s.section_id for s in result.sections] == [
        "executive_summary", "fmea_top", "certification"]
    assert result.unmatched == []
    assert not any("4.0 h/yr" in b.md for s in result.sections
                   for b in s.blocks if isinstance(b, Paragraph))
    # Without bookmarks the appendix is still recognised by its heading.
    doc = _open(data)
    _strip_bookmarks(doc)
    result = read_edited_docx(_save(doc), base)
    assert [s.section_id for s in result.sections] == [
        "executive_summary", "fmea_top", "certification"]
    assert result.unmatched == []


def test_a_new_heading_the_user_added_is_reported_as_unmatched_content():
    base = _document()
    doc = _open(_rendered(base))
    doc.add_paragraph("Reviewer's remarks", style="Heading 1")
    doc.add_paragraph("Some free text the reviewer added.")
    result = read_edited_docx(_save(doc), base)
    assert result.unmatched == ["Reviewer's remarks"]
    extra = [s for s in result.sections if s.section_id is None]
    assert len(extra) == 1
    assert extra[0].blocks == [Paragraph(md="Some free text the reviewer added.")]
    assert _section(result, "certification").changed is False


# ── the merge ────────────────────────────────────────────────────────────


def test_deleted_section_is_kept_in_the_merge_with_a_note():
    base = _document()
    doc = _open(_rendered(base))
    body = doc.element.body
    # Delete the Certification heading and everything after it up to sectPr.
    start = _paragraph_containing(doc, "Certification")._p
    el = start
    while el is not None and not el.tag.endswith("}sectPr"):
        nxt = el.getnext()
        body.remove(el)
        el = nxt
    result = read_edited_docx(_save(doc), base)
    assert [s.section_id for s in result.sections] == ["executive_summary", "fmea_top"]
    merged = merge_round_trip(base, result)
    assert [s.section_id for s in merged.sections] == [
        "executive_summary", "fmea_top", "certification"]
    kept = _by_id(merged, "certification")
    assert kept.blocks == _by_id(base, "certification").blocks
    assert kept.source == "code"
    assert "removed from the edited copy; kept from version 2" in (kept.note or "")
    assert "budget exhausted" in kept.note


def test_merge_bumps_version_keeps_mode_and_adds_edited_tables():
    base = _document()
    result = RoundTripResult(
        sections=[RoundTripSection(section_id="executive_summary",
                                   heading="Executive summary",
                                   blocks=[Paragraph(md="Rewritten by hand.")],
                                   changed=True, comments=["tighten"])],
        unmatched=[], comments_global=[], accepted_tracked_changes=0,
        edited_tables={"headline_edited_v3": Table(
            table_id="headline_edited_v3", columns=["Item", "Value"],
            rows=[["x", "y"]])},
    )
    merged = merge_round_trip(base, result)
    assert merged.version == 3 and merged.mode == "generated"
    assert merged.report_id == base.report_id
    assert merged.profile_id == base.profile_id and merged.model == base.model
    assert merged.template_file_id == base.template_file_id
    assert "headline_edited_v3" in merged.tables and "headline" in merged.tables
    s = _by_id(merged, "executive_summary")
    assert s.blocks == [Paragraph(md="Rewritten by hand.")]
    assert s.source == "user_edit" and s.pending_instruction == "tighten"
    assert s.comments == ["tighten"]
    # The other two were absent from the result: kept with the note.
    for sid in ("fmea_top", "certification"):
        assert "removed from the edited copy" in (_by_id(merged, sid).note or "")
    # The base was not mutated.
    assert base.version == 2 and _by_id(base, "executive_summary").source == "code"
    assert "headline_edited_v3" not in base.tables


def test_result_sections_with_no_id_do_not_enter_the_merge():
    base = _document()
    result = RoundTripResult(
        sections=[RoundTripSection(section_id=None, heading="Chapter X",
                                   blocks=[Paragraph(md="stray")], changed=True,
                                   comments=[])],
        unmatched=["Chapter X"], comments_global=[], accepted_tracked_changes=0)
    merged = merge_round_trip(base, result)
    assert [s.section_id for s in merged.sections] == [
        s.section_id for s in base.sections]
    assert not any(b == Paragraph(md="stray") for s in merged.sections for b in s.blocks)


# ── errors ───────────────────────────────────────────────────────────────


def test_a_png_is_refused_with_the_readers_error():
    with pytest.raises(TemplateReadError):
        read_edited_docx(_PNG, _document())
    with pytest.raises(TemplateReadError):
        read_edited_docx(b"", _document())
    assert issubclass(TemplateReadError, ValueError)


# ── the real document shape ──────────────────────────────────────────────


def test_assemblers_evidence_only_document_round_trips_unchanged():
    """
    The real shape: 14 sections with status sentences, stage notes and
    callouts the writer renders from `status` / `note`, not from blocks —
    none of it may come back as an edit.
    """
    from services.reports.assemble import evidence_only_document
    from services.reports.evidence import collect_evidence

    evidence = collect_evidence(study_report=None, eh_report=None, worksheet=None)
    base = evidence_only_document(evidence, title="Evidence only",
                                  report_id="0123456789abcdef", figure_pngs={})
    assert len(base.sections) >= 10
    result = read_edited_docx(render_document_docx(base, figure_bytes={}), base)
    assert [s.section_id for s in result.sections] == [s.section_id for s in base.sections]
    assert [s.section_id for s in result.sections if s.changed] == []
    assert result.unmatched == [] and result.edited_tables == {}
    merged = merge_round_trip(base, result)
    assert merged.model_dump(exclude={"version", "created_at"}) == \
        base.model_dump(exclude={"version", "created_at"})


def test_edited_field_value_comes_back_as_a_field():
    base = _document()
    doc = _open(_rendered(base))
    p = _paragraph_containing(doc, "Engine:")
    assert p.runs[0].bold and p.runs[1].text == "copt"
    p.runs[1].text = "mc"
    sec = _section(read_edited_docx(_save(doc), base), "fmea_top")
    assert sec.changed is True
    assert [b for b in sec.blocks if isinstance(b, Field)] == [Field(key="Engine", value="mc")]
