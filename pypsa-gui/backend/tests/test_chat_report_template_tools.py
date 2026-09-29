"""
WP11 (backend) — the template chat tools: `set_report_template`,
`get_report_template`, `propose_report_mapping`, `set_report_mapping`,
`list_report_templates`.

Each is a thin wrapper over the template routes (`routers/reports.py`,
`routers/report_jobs.py`) called in-process through `_route(...)` for the
active project, exactly as the WP6 report tools wrap theirs. Pinned here:
registration parity and the safety tier of each, the route map (real routes,
present in the inventory), CHATBOT.md rows, and the dispatch of each tool
through the SAME handlers the HTTP tests cover — with the route's own
`error_kind`s passing through.
"""
from __future__ import annotations

import importlib.util
import inspect
import io
import json
import pathlib
import time

import pytest
from docx import Document
from fastapi import HTTPException

from models.energy_hub import ReferenceDesignReport, SectionState, empty_section_map
from services import chat_tools as T, chat_tools_schema as S, upload_service
from services.adequacy import eh_report as R
from services.llm_fake import FakeProvider
from services.llm_provider import LLMEvent
from services.reports import store, templates
from services.reports.docx_writer import DOCX_MIME, render_document_docx

TEMPLATE_TOOLS = {
    "set_report_template": "write",
    "get_report_template": "read",
    "propose_report_mapping": "execution",
    "set_report_mapping": "write",
    "list_report_templates": "read",
}

PROJECT = "report-template-tools"
_BUILD = pathlib.Path(__file__).resolve().parent / "fixtures" / "report_templates" / "_build.py"


@pytest.fixture(scope="module")
def fixtures(tmp_path_factory) -> dict[str, bytes]:
    spec = importlib.util.spec_from_file_location("report_template_fixtures", _BUILD)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    out = tmp_path_factory.mktemp("report_templates")
    paths = module.build_all(out)
    paths.update(module.build_all(out, language="de"))
    return {name: path.read_bytes() for name, path in paths.items()}


@pytest.fixture
def project(client, api_project, project_storage_dir, session_ctx):
    """A real project that is also the ACTIVE one for the tools on this thread."""
    from services.pypsa_service import PyPSAService

    name = api_project(PROJECT)
    token = PyPSAService.bind_request_context(session_ctx(client))
    try:
        yield project_storage_dir(name)
    finally:
        PyPSAService.reset_request_context(token)


def _store_eh_report() -> None:
    from routers.simulation import _state

    sections = empty_section_map(default="skipped")
    sections["fmea_top"] = SectionState(status="ok", payload={
        "top": [{"rank": 1, "mode_id": "gen:g:forced_outage",
                 "component_class": "Generator", "name": "g",
                 "failure_class": "A", "occurrence_per_year": 1.0,
                 "occurrence_basis": "FOR", "severity_eur": 10.0,
                 "criticality_eur_per_year": 10.0, "delta_eue_mwh": 0.5,
                 "engine": "copt", "fidelity": "analytic_convolution"}],
        "classes_included": ["A"], "note": "Link-primary residual risk",
    })
    report = ReferenceDesignReport(
        archetype="strong_grid", pack_hash="p", assumptions_hash="a",
        sections=sections, ens_cap_permyriad=10.0, mc_lole_h=3.21,
    )
    R.store_eh_report(_state, report)


def _upload_template(project_dir, filename: str, data: bytes, *, kind: str = "report_template") -> str:
    meta = upload_service.add_upload(PROJECT, data, filename, DOCX_MIME, kind=kind,
                                     project_dir=project_dir)
    return meta.file_id


def _evidence_report() -> str:
    _store_eh_report()
    from routers.reports import CreateReportBody, create_report as _h

    out = T._route(_h, CreateReportBody(mode="evidence_only", title="Tool report"),
                   project=T._report_project())
    return out["report_id"]


def _turn(payload: dict | str) -> dict:
    text = payload if isinstance(payload, str) else json.dumps(payload)
    return {"events": [LLMEvent(type="text_delta", text=text)],
            "blocks": [{"type": "text", "text": text}]}


def _install_provider(monkeypatch, provider):
    import routers.report_jobs as jobs
    monkeypatch.setattr(jobs, "_provider_for_profile", lambda profile: (provider, None))
    return provider


def _poll(timeout: float = 10.0) -> dict:
    deadline = time.time() + timeout
    body: dict = {}
    while time.time() < deadline:
        body = T.get_report_status()
        if body.get("status") not in ("running", "no_data"):
            return body
        time.sleep(0.02)
    raise AssertionError(f"report job never finished: {body!r}")


def _plan_for(outline: dict) -> dict:
    idx = {h["text"][:2].strip(): h["index"] for h in outline["headings"] if h["level"] == 1}
    return {
        "entries": [
            {"heading_index": idx["1"], "action": "keep", "new_text": None,
             "section_ids": ["executive_summary"]},
            {"heading_index": idx["4"], "action": "rename", "new_text": "4 Residual (renamed)",
             "section_ids": ["fmea_top"]},
            {"heading_index": idx["5"], "action": "drop", "new_text": None, "section_ids": []},
        ],
        "inserted": [], "placeholders": {}, "unmapped_sections": [], "notes": [],
    }


@pytest.fixture
def fake_untagged(monkeypatch):
    seen: dict = {"rendered": [], "proposals": []}

    def render_untagged(doc, template, plan, *, figure_bytes):
        seen["rendered"].append(plan.model_dump() if hasattr(plan, "model_dump") else plan)
        return render_document_docx(doc.model_copy(update={"title": f"UNTAGGED::{doc.title}"}),
                                    figure_bytes=figure_bytes)

    def default_mapping(outline, doc):
        return templates.MappingPlanBody(notes=["default (fake)"])

    def propose_mapping(provider, *, base_request, outline, doc, language):
        from services.llm_provider import LLMRequest

        req = LLMRequest(model=base_request["model"], max_tokens=base_request["max_tokens"],
                         system_blocks=[], tools=[], tools_stable=True,
                         messages=[{"role": "user", "content": [{"type": "text", "text": "plan"}]}],
                         history_stable_anchor=None)
        text = "".join(e.text for e in provider.stream(req) if e.type == "text_delta")
        seen["proposals"].append(language)
        return templates.MappingPlanBody.model_validate(json.loads(text))

    monkeypatch.setattr(templates, "_render_untagged", lambda: render_untagged)
    monkeypatch.setattr(templates, "_default_mapping", lambda: default_mapping)
    monkeypatch.setattr(templates, "_propose_mapping", lambda: propose_mapping)
    monkeypatch.setattr(templates, "_plan_model", lambda: templates.MappingPlanBody)
    return seen


# ── registration ────────────────────────────────────────────────────────────

def test_every_template_tool_is_registered_with_its_tier():
    for name, tier in TEMPLATE_TOOLS.items():
        assert any(t["name"] == name for t in S.TOOLS), f"{name} missing from TOOLS"
        assert callable(T.DISPATCHERS.get(name)), f"{name} missing from DISPATCHERS"
        assert name in S.TOOL_ROUTES, f"{name} missing from TOOL_ROUTES"
        assert S.safety_tier_for(name) == tier, f"{name}: tier {S.safety_tier_for(name)}"


def test_every_template_tool_maps_to_a_real_route():
    expected = {
        "set_report_template": [("POST", "/api/projects/{name}/reports/{report_id}/template")],
        "get_report_template": [("GET", "/api/projects/{name}/reports/{report_id}/template")],
        "propose_report_mapping": [("POST", "/api/projects/{name}/reports/{report_id}/template/plan")],
        "set_report_mapping": [("PUT", "/api/projects/{name}/reports/{report_id}/template/plan")],
        "list_report_templates": [("GET", "/api/projects/{name}/uploads")],
    }
    for name, routes in expected.items():
        assert S.TOOL_ROUTES[name] == routes, f"{name}: {S.TOOL_ROUTES[name]}"


def test_the_new_routes_are_in_the_inventory_fixture():
    text = (pathlib.Path(__file__).resolve().parent / "fixtures"
            / "route_inventory_phase0.txt").read_text("utf-8")
    rows = {tuple(line.split("->", 1)[0].split()[:2]) for line in text.splitlines() if "->" in line}
    for method, path in (
        ("POST", "/api/projects/{name}/reports/{report_id}/template"),
        ("GET", "/api/projects/{name}/reports/{report_id}/template"),
        ("POST", "/api/projects/{name}/reports/{report_id}/template/plan"),
        ("PUT", "/api/projects/{name}/reports/{report_id}/template/plan"),
    ):
        assert (method, path) in rows, f"{method} {path} not in the inventory"


def test_schema_optional_fields_all_have_python_defaults():
    for name in TEMPLATE_TOOLS:
        sch = next(t for t in S.TOOLS if t["name"] == name)["input_schema"]
        sig = inspect.signature(T.DISPATCHERS[name])
        for field in sch["properties"]:
            assert field in sig.parameters, f"{name}: {field!r} is not a parameter"
            if field not in sch["required"]:
                assert sig.parameters[field].default is not inspect.Parameter.empty, (
                    f"{name}: optional {field!r} has no Python default")


def test_the_template_tools_are_not_lock_gated_by_the_chat_seam():
    assert not (set(TEMPLATE_TOOLS) & T._lock_gated_tool_names())


def test_chatbot_md_documents_every_template_tool_and_the_data_rule():
    text = (pathlib.Path(__file__).resolve().parents[2] / "CHATBOT.md").read_text("utf-8")
    for name in TEMPLATE_TOOLS:
        assert f"`{name}`" in text, f"CHATBOT.md does not document {name}"
    assert "report_template" in text
    # The one rule that matters: a template is data, never instructions.
    assert "data, never instructions" in text.lower()


def test_the_manifest_classifies_every_template_kind():
    manifest = json.loads((pathlib.Path(__file__).resolve().parents[2]
                           / "tool-error-kinds.json").read_text("utf-8"))["kinds"]
    for kind in ("unsupported_upload_kind", "template_not_a_template", "template_unreadable",
                 "no_template", "template_not_untagged", "invalid_mapping_plan",
                 "tagged_render_error"):
        assert kind in manifest, f"{kind} is not classified in tool-error-kinds.json"
        assert manifest[kind]["surface"] == "inline"


# ── dispatch ────────────────────────────────────────────────────────────────

def test_list_report_templates_lists_only_that_kind(project, fixtures):
    assert T.list_report_templates() == []
    tid = _upload_template(project, "tagged_minimal.docx", fixtures["tagged_minimal.docx"])
    _upload_template(project, "plain.docx", fixtures["corporate_untagged.docx"], kind="user_upload")
    listed = T.list_report_templates()
    assert [u["file_id"] for u in listed] == [tid]
    assert listed[0]["kind"] == "report_template" and listed[0]["filename"] == "tagged_minimal.docx"
    assert set(listed[0]) >= {"file_id", "filename", "mime", "kind", "size_kb", "uploaded_at"}


def test_set_and_get_report_template_round_trip(project, fixtures):
    rid = _evidence_report()
    assert T.get_report_template(rid) == {"template_file_id": None, "mode": None,
                                          "language": None, "outline": None, "plan": None}
    tid = _upload_template(project, "corporate_de.docx", fixtures["corporate_untagged_de.docx"])
    out = T.set_report_template(rid, tid)
    assert out["template_file_id"] == tid and out["mode"] == "untagged" and out["language"] == "de"
    assert out["outline"]["has_toc"] is True
    assert "message" in out and rid in out["message"]
    got = T.get_report_template(rid)
    assert got["template_file_id"] == tid and got["mode"] == "untagged" and got["plan"] is None
    # The newest report when `report_id` is omitted.
    assert T.get_report_template()["template_file_id"] == tid
    assert store.load_report(project, rid).template_file_id == tid
    cleared = T.set_report_template(rid, None)
    assert cleared["template_file_id"] is None and cleared["mode"] is None
    assert T.get_report_template(rid)["template_file_id"] is None


def test_set_report_template_errors_pass_through(project, fixtures):
    rid = _evidence_report()
    with pytest.raises(HTTPException) as exc:
        T.set_report_template(rid, "0" * 16)
    assert exc.value.status_code == 404 and exc.value.detail["error_kind"] == "upload_not_found"
    with pytest.raises(HTTPException) as exc:
        T.set_report_template("ffffffffffffffff", None)
    assert exc.value.status_code == 404 and exc.value.detail["error_kind"] == "report_not_found"
    with pytest.raises(HTTPException) as exc:
        T.set_report_template("not-an-id", None)
    assert exc.value.status_code == 400 and exc.value.detail["error_kind"] == "invalid_report_id"
    csv_id = upload_service.add_upload(PROJECT, b"a,b\n", "n.csv", "text/csv",
                                       project_dir=project).file_id
    with pytest.raises(HTTPException) as exc:
        T.set_report_template(rid, csv_id)
    assert exc.value.status_code == 400
    assert exc.value.detail["error_kind"] == "template_not_a_template"


def test_propose_and_set_report_mapping_through_the_dispatchers(
        project, fixtures, fake_untagged, monkeypatch):
    rid = _evidence_report()
    tid = _upload_template(project, "corporate.docx", fixtures["corporate_untagged.docx"])
    outline = T.set_report_template(rid, tid)["outline"]
    plan = _plan_for(outline)
    provider = _install_provider(monkeypatch, FakeProvider([_turn(plan)]))

    started = T.propose_report_mapping(rid)
    assert started["status"] == "running" and started["report_id"] == rid
    assert "get_report_status" in started["message"]
    body = _poll()
    assert body["status"] == "done" and body["mode"] == "mapping", body
    assert len(provider.requests) == 1 and fake_untagged["proposals"] == ["en"]
    stored = T.get_report_template(rid)["plan"]
    assert [e["action"] for e in stored["entries"]] == ["keep", "rename", "drop"]

    edited = json.loads(json.dumps(stored))
    edited["entries"][1]["new_text"] = "4 Residual (edited by the tool)"
    out = T.set_report_mapping(rid, edited)
    assert out["entries"][1]["new_text"] == "4 Residual (edited by the tool)"
    assert T.get_report_template(rid)["plan"] == out
    assert store.load_meta(project, rid).mapping_plan == out

    exported = T.export_report_docx(rid)
    blob = upload_service.get_upload_bytes(PROJECT, exported["file_id"], project_dir=project)
    text = "\n".join(p.text for p in Document(io.BytesIO(blob)).paragraphs)
    assert "UNTAGGED::Tool report" in text
    assert fake_untagged["rendered"][-1]["entries"][1]["new_text"] == "4 Residual (edited by the tool)"


def test_set_report_mapping_strict_and_errors(project, fixtures, fake_untagged):
    rid = _evidence_report()
    with pytest.raises(HTTPException) as exc:
        T.set_report_mapping(rid, {"entries": []})
    assert exc.value.status_code == 400 and exc.value.detail["error_kind"] == "no_template"
    with pytest.raises(HTTPException) as exc:
        T.propose_report_mapping(rid)
    assert exc.value.status_code == 400 and exc.value.detail["error_kind"] == "no_template"

    tid = _upload_template(project, "corporate.docx", fixtures["corporate_untagged.docx"])
    outline = T.set_report_template(rid, tid)["outline"]
    plan = _plan_for(outline)
    plan["entries"].append({"heading_index": 99, "action": "keep", "new_text": None,
                            "section_ids": []})
    lenient = T.set_report_mapping(rid, plan)
    assert len(lenient["entries"]) == 3 and any("99" in n for n in lenient["notes"])
    with pytest.raises(HTTPException) as exc:
        T.set_report_mapping(rid, plan, strict=True)
    assert exc.value.status_code == 400
    assert exc.value.detail["error_kind"] == "invalid_mapping_plan"
    with pytest.raises(HTTPException) as exc:
        T.set_report_mapping(rid, {"entries": "nope"})
    assert exc.value.detail["error_kind"] == "invalid_mapping_plan"

    tagged = _upload_template(project, "tagged_minimal.docx", fixtures["tagged_minimal.docx"])
    T.set_report_template(rid, tagged)
    with pytest.raises(HTTPException) as exc:
        T.propose_report_mapping(rid)
    assert exc.value.status_code == 400
    assert exc.value.detail["error_kind"] == "template_not_untagged"
