"""
IC G2 (plan `docs/superpowers/plans/2026-10-07-ic-g1-g2-campus-equipment.md` §4 WP-G2; G-4 … G-11): campus-chosen
equipment as owner capex in the same investment case, `build_finance_case(..., extra_assets=())`.

- G-4. `ExtraOwnerAsset(name, kind, basis, quantity, parts, build_year, source, source_hash)` validates itself
  (`extra_asset_invalid:<name>:<field>`); its parts are read by attribute (an `asset_schema` `UpfrontPart` as is)
  and `services.finance` never loads the solve stack.
- G-5. Each extra asset is one `AssetFinance` after the network assets: `campus:<kind>`, no carrier, overnight =
  Σ upfront per unit × quantity; excluded from incentives (flagged); its tax class from
  `depreciation_class_by_asset`; the replacement and terminal-value rules apply.
- G-6. Its fixed O&M is one `extra_asset_fom:<name>` template line per period, unscaled, never in the
  counterfactual. G-7. COD by year against the case COD. G-8. Never in the counterfactual.
- G-9. Refusals and the may-double-count flag. G-10. The report block, the xlsx rows and the hash.

The site (`test_finance_case_storage_and_meters._site`): 1 MW load, 2 MW PV (800 k/MW, 25 y), a 1 MW / 2 h
BESS (600 k, 15 y), one representative day weighted to a year. The extras: a 132/33 kV 40 MVA transformer
(1.8 MEUR, 40 y, FOM 1.5 %), 2.5 km of a 33 kV cable (90 kEUR/km, 40 y, FOM 0.5 %), a 5 Mvar capacitor bank
(90 kEUR, 25 y, FOM 2 %) and a 5 Mvar STATCOM (450 kEUR, 20 y, FOM 3 %), the defaults pack's campus rows.
"""
from __future__ import annotations

import copy
import dataclasses
import io
from datetime import date

import pytest

from models.finance import ESCALATION_CLASSES, FinanceInputs
from services.finance.case import FinanceRefused

COD = date(2030, 1, 1)
HASH = "0123456789abcdef"


def _part(per_unit, life, fom, *, derived=False):
    from services.asset_schema.access import UpfrontPart

    return UpfrontPart("investment", per_unit, life, fom, derived_from_capital_cost=derived)


def _ext(name="trf", kind="transformer", basis="lump", quantity=1, parts=None, build_year=2030,
         source="campus study s1", source_hash=HASH):
    from services.finance.case import ExtraOwnerAsset

    return ExtraOwnerAsset(name=name, kind=kind, basis=basis, quantity=quantity,
                           parts=parts if parts is not None else (_part(1_800_000.0, 40.0, 0.015),),
                           build_year=build_year, source=source, source_hash=source_hash)


TRF = dict(name="trf", kind="transformer", basis="lump", quantity=1)
CABLE = dict(name="feeder_cable", kind="cable", basis="per_km", quantity=2.5,
             parts="cable")
CAP = dict(name="cap", kind="capacitor_bank", basis="lump", quantity=1, parts="cap")
STATCOM = dict(name="sc", kind="statcom", basis="lump", quantity=1, parts="statcom")
PARTS = {"cable": (90_000.0, 40.0, 0.005), "cap": (90_000.0, 25.0, 0.02),
         "statcom": (450_000.0, 20.0, 0.03)}


def _mk(spec, **over):
    kw = dict(spec, **over)
    if isinstance(kw.get("parts"), str):
        kw["parts"] = (_part(*PARTS[kw["parts"]]),)
    return _ext(**kw)


def _fin(**over) -> FinanceInputs:
    kw = dict(financial_close=date(2029, 1, 1), cod_by_asset={"pv": COD, "bess": COD},
              analysis_years=15, contingency_share=0.0,
              escalation={c: 0.02 for c in ESCALATION_CLASSES},
              degradation_by_asset={"pv": 0.0, "bess": 0.0}, tax_losses="offset_other_income",
              financing_fee_tax="not_deducted", wacc_nominal=0.07, cost_of_equity=0.10,
              inflation=0.02)
    kw.update(over)
    return FinanceInputs(**kw)


def _build(n, cfg, fin=None, extras=(), **kw):
    import routers.results as R
    from services.results.finance_case import build_finance_case

    return build_finance_case(n, cfg, fin or _fin(), result_df=R._result_df,
                              extra_assets=tuple(extras), **kw)


def _refused(*a, **kw) -> FinanceRefused:
    with pytest.raises(FinanceRefused) as exc:
        _build(*a, **kw)
    return exc.value


@pytest.fixture
def solved(reset_backend):
    from tests.test_finance_case_storage_and_meters import _commercial, _single_owner, _site
    from tests.test_value_flow_reconciliation import _solve

    n = _site()
    return _solve(n, _commercial(_single_owner(n)))


def _lines(t) -> dict:
    return {ln.key: ln for ln in t.lines}


# ── G-4: the record validates itself; no solve stack ─────────────────────────


@pytest.mark.parametrize("over,field", [
    ({"name": ""}, "name"),
    ({"name": "trf:a"}, "name"),                     # a ':' would make the codes ambiguous
    ({"kind": "busbar"}, "kind"),
    ({"basis": "per_MW"}, "basis"),
    ({"quantity": 0}, "quantity"),
    ({"quantity": float("inf")}, "quantity"),
    ({"quantity": float("nan")}, "quantity"),
    ({"quantity": 1.5}, "quantity"),                                    # lump: whole units
    ({"basis": "per_bay", "kind": "switchgear", "quantity": 2.5}, "quantity"),
    ({"parts": ()}, "parts"),
    ({"build_year": 1899}, "build_year"),
    ({"build_year": 2201}, "build_year"),
    ({"source_hash": "abc"}, "source_hash"),
    ({"source_hash": "0123456789abcdeg"}, "source_hash"),
    ({"parts": ("lifetime", None)}, "lifetime"),
    ({"parts": ("lifetime", float("nan"))}, "lifetime"),
    ({"parts": ("lifetime", float("inf"))}, "lifetime"),
    ({"parts": ("lifetime", 0.5)}, "lifetime"),
    ({"parts": ("upfront_per_unit", -1.0)}, "upfront_per_unit"),
    ({"parts": ("upfront_per_unit", float("inf"))}, "upfront_per_unit"),
    ({"parts": ("fom_share", 1.5)}, "fom_share"),
    ({"parts": ("fom_share", None)}, "fom_share"),
])
def test_an_invalid_extra_asset_names_its_field(over, field):
    if isinstance(over.get("parts"), tuple) and over["parts"] and isinstance(over["parts"][0], str):
        attr, v = over["parts"]
        p = {"per_unit": 1_800_000.0, "life": 40.0, "fom": 0.015}
        p[{"lifetime": "life", "upfront_per_unit": "per_unit", "fom_share": "fom"}[attr]] = v
        over = {"parts": (_part(p["per_unit"], p["life"], p["fom"]),)}
    name = over.get("name", "trf")
    with pytest.raises(ValueError, match=f"extra_asset_invalid:{name}:{field}"):
        _ext(**over)


def test_a_cable_takes_fractional_km_and_a_real_upfront_part_is_accepted():
    from services.asset_schema.access import UpfrontPart

    e = _mk(CABLE)
    assert e.quantity == 2.5 and isinstance(e.parts[0], UpfrontPart)


def test_a_part_that_is_not_a_dataclass_or_misses_a_field_is_invalid():
    import types
    from dataclasses import dataclass

    ns = types.SimpleNamespace(name="investment", upfront_per_unit=1.0, lifetime=40.0, fom_share=0.0,
                               derived_from_capital_cost=False)

    @dataclass(frozen=True)
    class NoFom:
        name: str
        upfront_per_unit: float
        lifetime: float

    for bad in (ns, NoFom("investment", 1.0, 40.0), ("investment", 1.0, 40.0, 0.0)):
        with pytest.raises(ValueError, match="extra_asset_invalid:trf:parts"):
            _ext(parts=(bad,))


def test_importing_the_finance_case_never_loads_the_solve_stack_or_asset_schema():
    import subprocess
    import sys
    from pathlib import Path

    code = ("import sys, services.finance.case, services.finance.engine\n"
            "bad = sorted(m for m in sys.modules if m.startswith(('services.solver', "
            "'services.asset_schema', 'services.results', 'routers')))\n"
            "print(','.join(bad))\n")
    res = subprocess.run([sys.executable, "-c", code], cwd=Path(__file__).resolve().parents[1],
                         capture_output=True, text=True, timeout=120)
    assert res.returncode == 0, res.stderr[-2000:]
    assert res.stdout.strip() == ""


# ── G-5, G-6, G-8: into the case ─────────────────────────────────────────────


@pytest.mark.live_solve
def test_extra_assets_are_owner_capex_with_their_fom_and_no_counterfactual(solved):
    from services.finance.engine import run_case

    n, cfg = solved
    base = _build(n, cfg)
    case = _build(n, cfg, extras=(_ext(), _mk(CABLE)))
    names = [a.name for a in case.assets]
    assert names == [a.name for a in base.assets] + ["trf", "feeder_cable"]
    trf, cable = case.assets[-2:]
    assert (trf.component, trf.carrier, trf.lifetime_years) == ("campus:transformer", None, 40.0)
    assert trf.overnight_cost == pytest.approx(1_800_000.0)
    assert (cable.component, cable.carrier) == ("campus:cable", None)
    assert cable.overnight_cost == pytest.approx(2.5 * 90_000.0)
    assert [(p.name, p.overnight_cost, p.lifetime_years, p.fom_share) for p in cable.parts] == \
        [("investment", pytest.approx(225_000.0), 40.0, 0.005)]
    assert case.extra_assets[0].name == "trf" and base.extra_assets == ()
    assert "cod_from_build_year:trf" in case.flags
    # G-8: the counterfactual is the same with and without extras.
    assert case.counterfactual == base.counterfactual
    assert case.counterfactual_hash == base.counterfactual_hash
    # G-6: one FOM line per extra and template, −fom_share × overnight, opex, base-year money.
    (t,) = case.templates
    fom = _lines(t)["extra_asset_fom:trf"]
    assert (fom.stream, fom.esc_class, fom.source, fom.source_id, fom.money_year) == \
        ("fom", "opex", "extra_asset_fom", "trf", case.base_year)
    assert fom.amount == pytest.approx(-0.015 * 1_800_000.0)
    assert _lines(t)["extra_asset_fom:feeder_cable"].amount == pytest.approx(-0.005 * 225_000.0)
    assert not any(k.startswith("extra_asset_fom") for cf in case.counterfactual for k in _lines(cf))
    # Capex: the project capex includes them, × (1 + contingency).
    r0 = run_case(dataclasses.replace(base, inputs=_fin(contingency_share=0.1)))
    r = run_case(dataclasses.replace(case, inputs=_fin(contingency_share=0.1)))
    assert float(r.op.capex.sum()) == pytest.approx(float(r0.op.capex.sum())
                                                    + 1.1 * (1_800_000.0 + 225_000.0))
    tl = r.tl
    for y in (2030, 2035, 2044):
        assert r.op.lines["extra_asset_fom:trf"][tl.index(y)] == \
            pytest.approx(-27_000.0 * 1.02 ** (y - case.base_year))


@pytest.mark.live_solve
def test_replacements_and_the_terminal_value_include_the_extra_parts(solved):
    """A 25-y capacitor bank on a 30-y axis from COD 2030: re-bought in 2054 (its last service year),
    that purchase serves 2055–2079, 20 years past the horizon 2060."""
    from services.finance.engine import run_case
    from services.finance.report import assemble_finance_sections

    n, cfg = solved
    fin = _fin(analysis_years=30, replacement_rule="part_lifetimes",
               terminal_value={"method": "remaining_life_annuity"})
    case = _build(n, cfg, fin, extras=(_mk(CAP),))
    r = run_case(case)
    cap_items = [(x.year, x.part, v) for x, v in r.op.replacement_items if x.asset == "cap"]
    assert cap_items == [(2054, "investment", pytest.approx(90_000.0 * 1.02 ** 24))]
    (term,) = [t for t in r.op.terminal_terms if t.asset == "cap"]
    rate = case.lp_basis.discount_rate
    crf = rate / (1 - (1 + rate) ** -25)
    apf = (1 - (1 + rate) ** -20) / rate
    assert (term.start_year, term.remaining_years, term.asset_rate) == (2055, 20.0, rate)
    assert term.value == pytest.approx(90_000.0 * 1.02 ** 24 * crf * apf, rel=1e-12)
    rep = assemble_finance_sections(r, case)
    lines = [ln for ln in rep.cashflow_lines if ln.provenance.source_id == "cap:investment"]
    assert [ln.year for ln in lines] == [2054]
    tv = rep.sections["project"].payload["terminal_value"]
    assert "cap" in {t["asset"] for t in tv["terms"]}
    assert "never charged by the LP" in tv["rate_basis"]


@pytest.mark.live_solve
def test_under_fixed_a_short_extra_needs_a_replacement_entry(solved):
    from services.finance.engine import run_case

    n, cfg = solved
    fin = _fin(analysis_years=25, replacement_capex=[(2044, "bess", 600_000.0)])
    case = _build(n, cfg, fin, extras=(_mk(STATCOM),))
    with pytest.raises(FinanceRefused) as exc:
        run_case(case)
    assert exc.value.code == "asset_lifetime_short" and "sc" in exc.value.detail
    ok = dataclasses.replace(case, inputs=_fin(analysis_years=25, replacement_capex=[
        (2044, "bess", 600_000.0), (2049, "sc", 450_000.0)]))
    assert run_case(ok).op.status["capex"] == "ok"


# ── G-7: COD by year ─────────────────────────────────────────────────────────


@pytest.mark.live_solve
def test_an_extra_asset_cod_is_compared_by_year(solved):
    n, cfg = solved
    cod = date(2027, 7, 1)
    fin = _fin(financial_close=date(2027, 1, 1), cod_by_asset={"pv": cod, "bess": cod})
    case = _build(n, cfg, fin, extras=(_ext(build_year=2027),))
    assert case.cod == cod and "cod_from_build_year:trf" in case.flags
    assert _refused(n, cfg, fin, extras=(_ext(build_year=2026),)).code == "cod_mismatch"
    assert _refused(n, cfg, fin, extras=(_ext(build_year=2028),)).code == \
        "extra_asset_staged_build:trf"
    typed = _fin(financial_close=date(2027, 1, 1),
                 cod_by_asset={"pv": cod, "bess": cod, "trf": date(2027, 1, 1)})
    assert _refused(n, cfg, typed, extras=(_ext(build_year=2027),)).code == "cod_mismatch"
    same = _fin(financial_close=date(2027, 1, 1), cod_by_asset={"pv": cod, "bess": cod, "trf": cod})
    got = _build(n, cfg, same, extras=(_ext(build_year=2028),))      # a typed entry wins
    assert got.cod == cod and "cod_from_build_year:trf" not in got.flags


# ── G-6 on a two-period network with an annualised template ─────────────────


@pytest.mark.live_solve
def test_the_fom_line_is_unscaled_once_per_period_and_never_in_the_counterfactual(reset_backend):
    from tests.test_finance_case_adapter import OWNED, _fin as adapter_fin
    from tests.test_value_flow_reconciliation import _commercial, _network, _solve

    n, cfg = _solve(_network(multi=True), _commercial(OWNED), multi=True)
    case = _build(n, cfg, adapter_fin(annualise=True), extras=(_ext(),))
    assert any(f.startswith("template_annualised:") for f in case.flags)
    assert [t.first_year for t in case.templates] == [2030, 2040]
    for t, period in zip(case.templates, ("2030", "2040")):
        got = [ln for ln in t.lines if ln.key == "extra_asset_fom:trf"]
        assert len(got) == 1
        assert got[0].amount == pytest.approx(-27_000.0)            # not × 8760 / 168
        assert (got[0].period, got[0].money_year) == (period, 2030)
    assert not any(ln.key.startswith("extra_asset_fom") for t in case.counterfactual
                   for ln in t.lines)


# ── G-5: incentives and tax ──────────────────────────────────────────────────


@pytest.mark.live_solve
@pytest.mark.parametrize("classes", [["battery"], []], ids=["asset_classes", "all_assets"])
def test_incentives_exclude_the_extra_assets(solved, classes):
    from services.finance.engine import run_case

    n, cfg = solved
    grant = {"kind": "grant", "rate": 0.3, "grant_tax_treatment": "taxable",
             "eligibility": {"asset_classes": classes}}
    fin = _fin(incentives=[grant])
    base = run_case(_build(n, cfg, fin))
    r = run_case(_build(n, cfg, fin, extras=(_ext(),)))
    assert r.incentives.established(), r.incentives.reasons
    assert "incentive_excludes_extra_asset:trf" in r.flags
    assert [ln.assets for ln in r.incentives.lines] == [ln.assets for ln in base.incentives.lines]
    assert "trf" not in r.incentives.lines[0].assets
    assert float(r.incentives.grant.sum()) == pytest.approx(float(base.incentives.grant.sum()))


@pytest.mark.live_solve
def test_an_extra_asset_tax_class_comes_from_the_inputs(solved):
    from services.finance.engine import run_case
    from services.finance.packs.base import load_pack

    n, cfg = solved
    pack = load_pack("eu_de", as_of=date(2026, 6, 1))
    common = dict(hebesatz_pct=400.0, acquisition_date=date(2029, 6, 1))
    missing = run_case(_build(n, cfg, _fin(**common, depreciation_class_by_asset={
        "pv": "afa_20", "bess": "afa_10"}), extras=(_ext(),)), pack)
    assert "tax_input_missing:depreciation_class:trf" in missing.reasons["tax"]
    stated = run_case(_build(n, cfg, _fin(**common, depreciation_class_by_asset={
        "pv": "afa_20", "bess": "afa_10", "trf": "afa_20"}), extras=(_ext(),)), pack)
    assert not any("depreciation_class:trf" in x for x in stated.reasons.get("tax", []))
    assert stated.tax is not None, stated.reasons
    assert all(float(d.sum()) > 0 for d in stated.tax.depreciation.values())


# ── G-9: refusals and the double-count flag ──────────────────────────────────


@pytest.mark.live_solve
def test_the_adapter_refusals(solved):
    n, cfg = solved
    derived = _ext(parts=(_part(1_800_000.0, 40.0, 0.015, derived=True),))
    assert _refused(n, cfg, extras=(derived,)).code == "extra_asset_derived_upfront:trf"
    assert _refused(n, cfg, extras=(_ext(), _ext())).code == "extra_asset_duplicate:trf"
    assert _refused(n, cfg, extras=(_ext(name="bess"),)).code == \
        "extra_asset_duplicates_network_asset:bess"


@pytest.mark.live_solve
def test_a_transformer_or_cable_beside_an_owned_network_line_may_double_count(reset_backend):
    from tests.test_finance_case_storage_and_meters import _commercial, _single_owner, _site
    from tests.test_value_flow_reconciliation import _solve

    n = _site()
    n.add("Bus", "mv", carrier="AC")
    n.add("Line", "feeder", bus0="site", bus1="mv", s_nom=10.0, x=0.01, r=0.001)
    vf = _single_owner(n)
    vf = copy.deepcopy(vf)
    vf["asset_owners"].append({"asset_id": "feeder", "component": "Line", "owner": "site"})
    n, cfg = _solve(n, _commercial(vf))
    fin = _fin(cod_by_asset={"pv": COD, "bess": COD, "feeder": COD})
    case = _build(n, cfg, fin, extras=(_ext(), _mk(CABLE), _mk(CAP)))
    assert "feeder" in {a.name for a in case.assets}
    assert {"extra_asset_may_double_count:trf", "extra_asset_may_double_count:feeder_cable"} <= \
        set(case.flags)
    assert "extra_asset_may_double_count:cap" not in case.flags
    plain = _build(*_solve(_site(), _commercial(_single_owner(_site()))), extras=(_ext(),))
    assert not any(f.startswith("extra_asset_may_double_count") for f in plain.flags)


# ── G-10: report, workbook and hash ──────────────────────────────────────────


def _hash_fixture(**over):
    """A hand case whose hash is pinned to its value on master 546f2cb (before G2)."""
    from services.finance.case import (
        AssetFinance, AssetPart, FinanceCase, LpBasis, Template, TemplateLine,
    )

    kw = dict(
        inputs=FinanceInputs(financial_close=date(2029, 1, 1), analysis_years=10,
                             contingency_share=0.0,
                             escalation={"tariff": 0.02, "opex": 0.02, "capex": 0.0}),
        owner="site", base_year=2030, cod=date(2030, 1, 1),
        templates=(Template(2030, (TemplateLine("bill", "energy_import", -1000.0, "tariff"),),
                            money_year=2030),),
        assets=(AssetFinance("bess", "StorageUnit", 1_600_000.0, 15.0, "battery",
                             (AssetPart("power", 400_000.0, 10.0, 0.01),
                              AssetPart("energy", 1_200_000.0, 15.0, None))),),
        flags=("cod_from_build_year:bess",), lp_basis=LpBasis(discount_rate=0.07))
    kw.update(over)
    return FinanceCase(**kw)


def test_a_case_without_extras_keeps_its_master_hash_and_extras_enter_the_hash():
    from services.results.finance_case import finance_case_hash

    assert finance_case_hash(_hash_fixture()) == "64e7a34307ea1af3"        # master 546f2cb
    with_one = finance_case_hash(_hash_fixture(extra_assets=(_ext(),)))
    assert with_one != "64e7a34307ea1af3"
    other = finance_case_hash(_hash_fixture(extra_assets=(_ext(source_hash="fedcba9876543210"),)))
    assert other != with_one
    more = finance_case_hash(_hash_fixture(extra_assets=(_ext(quantity=2),)))
    assert more not in (with_one, other)


@pytest.mark.live_solve
def test_the_report_block_reads_the_case_assets_and_the_workbook_lists_them(solved):
    from openpyxl import load_workbook

    from services.finance.case import scale_capex
    from services.finance.engine import run_case
    from services.finance.export_xlsx import build_workbook
    from services.finance.report import assemble_finance_sections

    n, cfg = solved
    case = scale_capex(_build(n, cfg, _fin(currency_year=2026), extras=(_ext(), _mk(CABLE))), 1.1)
    rep = assemble_finance_sections(run_case(case), case)
    block = rep.sections["project"].payload["extra_assets"]
    by = {b["name"]: b for b in block}
    assert set(by) == {"trf", "feeder_cable"}
    trf = by["trf"]
    assert (trf["kind"], trf["basis"], trf["quantity"], trf["build_year"], trf["source"],
            trf["source_hash"]) == ("transformer", "lump", 1, 2030, "campus study s1", HASH)
    assert trf["overnight_cost"] == pytest.approx(1.1 * 1_800_000.0)      # scale_capex reflected
    assert trf["parts"] == [{"name": "investment", "upfront_per_unit_as_passed": 1_800_000.0,
                             "overnight_cost": pytest.approx(1.1 * 1_800_000.0),
                             "lifetime": 40.0, "fom_share": 0.015}]
    assert "within its build year 2030" in trf["cod"]
    assert "2026 EUR" in trf["money_year"]
    assert by["feeder_cable"]["overnight_cost"] == pytest.approx(1.1 * 225_000.0)
    ws = load_workbook(io.BytesIO(build_workbook(rep)))["About"]
    rows = [(ws.cell(i, 1).value, ws.cell(i, 2).value) for i in range(1, ws.max_row + 1)]
    extra = [v for k, v in rows if k == "Extra asset"]
    assert len(extra) == 2 and any("trf" in v and HASH in v for v in extra)
    # A case without extras carries an empty block (readers use .get on old payloads).
    plain = _build(n, cfg)
    assert assemble_finance_sections(run_case(plain), plain).sections["project"].payload[
        "extra_assets"] == []


@pytest.mark.live_solve
def test_the_report_block_shows_scaled_part_costs_and_a_typed_cod(solved):
    """Review notes 1, 2: after `scale_capex` each part's overnight cost is the scaled one from
    `case.assets` (the per-unit cost is labelled as passed); a typed `cod_by_asset` entry over a
    different build year is stated as typed, not "within its build year"."""
    from services.finance.case import scale_capex
    from services.finance.engine import run_case
    from services.finance.report import assemble_finance_sections

    n, cfg = solved
    cod = date(2027, 7, 1)
    fin = _fin(financial_close=date(2027, 1, 1), cod_by_asset={"pv": cod, "bess": cod, "trf": cod})
    case = scale_capex(_build(n, cfg, fin, extras=(_ext(build_year=2028), _mk(CABLE, build_year=2027))),
                       1.2)
    by = {b["name"]: b for b in assemble_finance_sections(run_case(case), case)
          .sections["project"].payload["extra_assets"]}
    (part,) = by["trf"]["parts"]
    assert part["upfront_per_unit_as_passed"] == 1_800_000.0
    assert part["overnight_cost"] == pytest.approx(1.2 * 1_800_000.0)
    assert by["trf"]["overnight_cost"] == pytest.approx(1.2 * 1_800_000.0)
    assert by["feeder_cable"]["parts"][0]["overnight_cost"] == pytest.approx(1.2 * 2.5 * 90_000.0)
    assert "typed in cod_by_asset" in by["trf"]["cod"] and "within" not in by["trf"]["cod"]
    assert "within its build year 2027" in by["feeder_cable"]["cod"]
