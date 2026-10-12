"""
QA: the guided decision study (BESS at a site), end to end over HTTP with a
REAL LP — plan docs/superpowers/plans/2026-09-28-guided-investment-study-mvp1-v2.md
§ S9 and its "Definition of done".

The journey a user takes with no canvas: a meter file and a seed tariff in,
a verdict, a pro forma and a report out.

* **The load is an upload** through the guided flow's real path (a draft's
  file travels in the intake as ``load.csv_text`` and creation writes it as an
  upload into the study's OWN base project, gate S8 BC-S8-5), in kW with
  timestamps. Its shape is the site golden fixture's evening spike
  (``tests/golden/site_fixture.py``): the shipped sector profiles size every
  battery-only option to zero (gate S5), which would make every CAPEX and FOM
  check vacuous.
* **The seed tariff** (``de_industrial_illustrative``); PV off, so four
  options (``none``, ``bess_1h``, ``bess_2h``, ``bess_4h``).
* **Two ledger rows edited**: a storage quote of 500 EUR/kWh and an inverter
  quote of 230 EUR/kW. On this load that sizes the 1-hour battery (about
  0.8 MW) and leaves the 2- and 4-hour ones at zero, and the demand-charge
  row's low bound turns the 1-hour battery's NPV negative with its size held
  fixed: the verdict is ``marginal``, reached through the API (plan
  "Definition of done": `marginal` demonstrably reachable).

Checked on the way (plan S9; gate carries S2, S4, S5, S6, S7, S8):

1. creating the study (M0) and its upload change no pre-existing project;
2. the run solves every option, charged to a ``decision_study`` campaign on
   the study's base context, not the session's;
3. the best option's case reconciles: CAPEX to ``packs.battery_upfront_eur_per_mw``
   x size from the ledger, which ``upfront_cost_series`` now reads too (U2 WP7
   C1: the pack writes the battery as its two upfront parts, so the S4 gate's
   24-35 % single-lifetime overstatement is gone; checked equal), and
   fixed O&M to ``asset_economics.fom_cost_eur`` and to ``cost_breakdown``'s
   StorageUnit FOM recomputed on the fork's own network;
4. the findings carry a verdict; the tornado re-dispatches the demand-charge
   row twice on throw-away forks and leaves none behind, nothing study-owned,
   and the option forks unchanged;
5. the report renders to DOCX, XLSX and HTML, and ``stale`` flips after a
   ledger edit;
6. an abort mid-run leaves every pre-existing project unchanged and removes
   the unsolved forks;
7. a user project resident at the cap is neither evicted nor saved, and
   nothing is left marked study-owned;
8. deleting a study cascades to its forks.

U2 WP10 (plan docs/superpowers/plans/2026-10-05-guided-study-u2-engine-rewire.md
WP10, §4): the study runs on the Investment Case engine only. ``_reconcile``
compiles the fork's config with ``compile`` (no GS tariff path is left) and
checks that every option fork carries the engine's ``commercial`` and
``finance`` blocks bound to the study's OWN export series, that the case's
engines are ``tariff_engine`` / ``finance_engine``, that the case's WACC gate
is consistent at the centre, and the §4.6 identity on the named option
(NPV = LP objective saving x AF(r, H)). The §4 judging rule closes the run:
the verdict class, its drivers, the named option and the skipped options
equal the WP0 record, and every recorded figure equals it within the
fixture's tolerances or is a recorded delta with a cause
(``EVIDENCE["u2_deltas"]``).

Runs through ``tests/run_qa_drivers.py`` (``pixi run gui-qa-drivers``).
Real solves: 4 (run) + 2 (tornado) + 1-2 (the aborted run). Exit 0 = every
step passed.
"""
from __future__ import annotations

import hashlib
import io
import os
import pathlib
import sys
import time

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

# THE load-bearing import, and it must come first: `qa_support` pins the
# sandbox before anything imports `main` or `settings`.
from tests import qa_support  # noqa: E402

import main  # noqa: E402
import pandas as pd  # noqa: E402
from routers import studies as studies_router  # noqa: E402
from services.pypsa_service import PyPSAService  # noqa: E402
from tests import conftest as harness  # noqa: E402

PASS = 0
FAIL = 0
SRC = "qa-ds-src"
BASE = "qa-ds-base"
ABORT_BASE = "qa-ds-abort"
USERS = [f"qa-ds-user{i}" for i in range(5)]
YEAR = 2025
# 500, not the 450 of the first S9 driver: at 450 the demand-charge low bound
# was only -21.6k EUR, and at 420 the verdict was `recommended` (gate S9 [S2]).
# At 500 the flip and the centre both clear the stated margins below.
STORAGE_QUOTE = 500.0      # EUR/kWh
#: `marginal` must hold with margin on both sides (gate S9 [S2]): the centre
#: battery NPV positive by at least this much, the flipping bound negative by at
#: least FLIP_MARGIN_EUR, and every other bound positive by at least that.
CENTRE_MARGIN_EUR = 100_000.0
FLIP_MARGIN_EUR = 25_000.0
INVERTER_QUOTE = 230.0     # EUR/kW
OPTIONS = ["none", "bess_1h", "bess_2h", "bess_4h"]
EVIDENCE: dict = {}


def _step(label: str, ok: bool, msg: str = "") -> bool:
    global PASS, FAIL
    if ok:
        PASS += 1
        print(f"  [PASS] {label}" + (f" — {msg}" if msg else ""))
    else:
        FAIL += 1
        print(f"  [FAIL] {label}" + (f" — {msg}" if msg else ""))
    return ok


def _enable_studies() -> None:
    """
    The study routes refuse in multi-user mode until OPEN-ITEMS 1 is closed
    (BC-6), and the qa sandbox is multi-user; the suite enables them the same
    way (`tests/study_s4_support.enable_studies`).
    """
    os.environ["PYPSAGUI_DECISION_STUDIES"] = "1"
    dep = studies_router.require_decision_studies_enabled
    main.app.dependency_overrides[dep] = lambda: None
    studies_router.require_decision_studies_enabled = lambda: None


def _meter_csv() -> str:
    """The site fixture's evening-spike shape, as a kW meter export."""
    idx = pd.date_range(f"{YEAR}-01-01", periods=8760, freq="h")
    lines = ["timestamp,site load (kW)"]
    for t in idx:
        kw = 600.0 + (800.0 if (t.weekday() < 5 and t.hour == 17) else 0.0)
        lines.append(f"{t:%Y-%m-%d %H:%M},{kw:.1f}")
    return "\n".join(lines) + "\n"


def _intake() -> dict:
    return {
        "site": {"zone": "DE", "connection_mw": 2.0, "year": YEAR, "latitude": 51.0},
        "tariff": {"tariff_id": "de_industrial_illustrative"},
        "load": {"source": "upload", "csv_text": _meter_csv(), "filename": "meter.csv"},
        "pv": {"enabled": False},
    }


def _dir_hash(path: pathlib.Path) -> str:
    h = hashlib.sha256()
    for f in sorted(p for p in pathlib.Path(path).rglob("*") if p.is_file()):
        h.update(str(f.relative_to(path)).encode())
        h.update(f.read_bytes())
    return h.hexdigest()


def _all_projects() -> dict[str, pathlib.Path]:
    from sqlalchemy import select

    from db.models import Project
    from services import project_registry

    with qa_support.db_session() as db:
        return {p.name: project_registry.project_dir(p)
                for p in db.scalars(select(Project)).all()}


def _hashes(exclude=()) -> dict[str, str]:
    return {n: _dir_hash(d) for n, d in _all_projects().items()
            if d.exists() and not any(n == e or n.startswith(f"{e}-") for e in exclude)}


def _changed(before: dict[str, str]) -> list[str]:
    now = _all_projects()
    return sorted(n for n, h in before.items()
                  if n not in now or _dir_hash(now[n]) != h)


def _key(name: str) -> str:
    row = qa_support.project_row(name)
    return f"{row.org_id}:{row.id}"


def _poll(path: str, *, timeout: float = 900.0, until=None) -> dict:
    c = qa_support.client()
    t0 = time.time()
    body: dict = {}
    while time.time() - t0 < timeout:
        r = c.get(path)
        if r.status_code == 200:
            body = r.json()
            if until is not None and until(body):
                return body
            if until is None and body.get("status") != "running":
                return body
        time.sleep(0.2)
    return {"status": "timeout", **body}


def _study_url(base: str, sid: str, tail: str = "") -> str:
    return f"/api/projects/{base}/studies/{sid}{tail}"


# ── 0. user projects at the resident cap ─────────────────────────────────

def _user_projects_at_the_cap() -> dict:
    c = qa_support.client()
    for name in [*USERS, SRC]:
        qa_support.install_network(harness.build_network(), name=name)
        r = c.post(f"/api/projects/{name}", params={"force": True, "rebind": True})
        if r.status_code != 200:
            _step(f"user project {name} saved", False, r.text[:200])
    edited = PyPSAService.get_context(_key(USERS[1]))
    ok = _step("a user project is resident at the cap with an unsaved edit",
               edited is not None and PyPSAService.RESIDENT_CAP == 5,
               f"cap={PyPSAService.RESIDENT_CAP}")
    if ok:
        edited.network.add("Bus", "qa_unsaved_edit_bus")
    return {"edited": edited, "resident": set(PyPSAService._contexts)}


# ── 1. create: M0 and the upload write no user project ───────────────────

def _create(before: dict[str, str], base: str) -> str | None:
    c = qa_support.client()
    r = c.post(f"/api/projects/{SRC}/studies/", json={
        "question_id": "bess_at_site", "name": f"Site battery ({base})",
        "project_name": base, "intake": _intake()})
    if not _step(f"the study {base} is created (M0)", r.status_code == 201,
                 f"HTTP {r.status_code}" + ("" if r.status_code == 201 else f" {r.text[:300]}")):
        return None
    body = r.json()
    load = body["intake"]["load"]
    _step("the draft's meter file became an upload in the study's own project",
          bool(load.get("upload_id")) and "csv_text" not in load
          and (qa_support.project_dir(base) / "uploads" / load["upload_id"]).is_dir(),
          f"upload_id={load.get('upload_id')}")
    changed = _changed(before)
    _step("creating the study and its upload changed no pre-existing project",
          changed == [], f"changed={changed}")
    return body["study_id"]


# ── 2. ledger, run ───────────────────────────────────────────────────────

def _edit_ledger(sid: str) -> dict | None:
    c = qa_support.client()
    seed = c.get(_study_url(BASE, sid, "/ledger")).json()
    _step("the seeded ledger is the library's", seed["ledger"] and all(
        r["provenance"] == "library" for r in seed["ledger"]["rows"]
        if r["key"] in ("battery_storage_eur_per_kwh", "battery_inverter_eur_per_kw")))
    r = c.put(_study_url(BASE, sid, "/ledger"), json={"rows": [
        {"key": "battery_storage_eur_per_kwh", "value": STORAGE_QUOTE, "unit": "EUR/kWh",
         "source": "vendor quote (QA driver)"},
        {"key": "battery_inverter_eur_per_kw", "value": INVERTER_QUOTE, "unit": "EUR/kW",
         "source": "vendor quote (QA driver)"}]})
    if not _step("two ledger rows edited", r.status_code == 200, r.text[:300]):
        return None
    rows = {row["key"]: row for row in r.json()["ledger"]["rows"]}
    _step("the edited rows are the user's",
          all(rows[k]["provenance"] == "user" and rows[k]["status"] == "customised"
              for k in ("battery_storage_eur_per_kwh", "battery_inverter_eur_per_kw")),
          str({k: (rows[k]["provenance"], rows[k]["status"], rows[k]["value"])
               for k in ("battery_storage_eur_per_kwh", "battery_inverter_eur_per_kw")}))
    return r.json()["ledger"]


def _run(sid: str) -> dict | None:
    from services.adequacy import campaign
    from services.study import runner as R

    c = qa_support.client()
    t0 = time.monotonic()
    r = c.post(_study_url(BASE, sid, "/run"), json={})
    if not _step("the run starts", r.status_code == 202, f"HTTP {r.status_code} {r.text[:300]}"):
        return None
    _step("the run is charged one solve per option", r.json()["solves_charged"] == 4,
          str(r.json()["solves_charged"]))
    rec = _poll(_study_url(BASE, sid, "/run"))
    EVIDENCE["run_s"] = round(time.monotonic() - t0, 1)
    if not _step("the run finishes", rec.get("status") == "done",
                 f"status={rec.get('status')} error={rec.get('error')}"):
        return None
    _step("every option solved", sorted(rec["solved"]) == sorted(OPTIONS), str(rec["solved"]))
    camp = rec.get("campaign") or {}
    _step("the solves were charged to a decision_study campaign on the study's context",
          [(e["study"], e["solves_charged"]) for e in camp.get("entries") or []]
          == [("decision_study", 4)] and camp.get("active") is False, str(camp)[:300])
    fg = qa_support.session_context()
    _step("the session's own context has no campaign charged",
          R._on_ctx(fg, campaign.status)["active"] is False
          and fg.registry_key != _key(BASE))
    return rec


# ── 3. findings, tornado ─────────────────────────────────────────────────

def _option_fork_hashes(base: str) -> dict[str, str]:
    return {n: _dir_hash(d) for n, d in _all_projects().items() if n.startswith(f"{base}-opt-")}


def _tornado(sid: str) -> dict | None:
    c = qa_support.client()
    pre = c.get(_study_url(BASE, sid, "/findings")).json()
    _step("before the tornado the findings say why the verdict waits",
          pre.get("robustness", {}).get("note") == "tornado_not_run",
          str(pre.get("robustness", {}).get("note")))
    forks_before = _option_fork_hashes(BASE)
    t0 = time.monotonic()
    r = c.post(_study_url(BASE, sid, "/findings/tornado"), json={})
    if not _step("the tornado starts", r.status_code == 202, f"HTTP {r.status_code} {r.text[:300]}"):
        return None
    EVIDENCE["tornado_solves_estimated"] = r.json().get("solves_estimated")
    rec = _poll(_study_url(BASE, sid, "/findings/tornado"))
    EVIDENCE["tornado_s"] = round(time.monotonic() - t0, 1)
    if not _step("the tornado finishes", rec.get("status") == "done",
                 f"status={rec.get('status')} error={rec.get('error')}"):
        return None
    _step("the tornado re-dispatched the demand-charge row twice",
          rec.get("solves_charged") == 2 and len(rec.get("variants") or []) == 2,
          f"solves={rec.get('solves_charged')} variants={rec.get('variants')}")
    left = [n for n in _all_projects() if "-var-" in n]
    _step("no throw-away fork is left", left == [] and not rec.get("forks_left"), str(left))
    _step("nothing is left marked study-owned", PyPSAService._study_owned == set(),
          str(sorted(PyPSAService._study_owned)))
    _step("the option forks are unchanged by the tornado",
          _option_fork_hashes(BASE) == forks_before)
    f = c.get(_study_url(BASE, sid, "/findings")).json()
    v = f.get("verdict") or {}
    rob = f.get("robustness") or {}
    EVIDENCE["verdict"] = {"class": v.get("class"), "option": v.get("option_id"),
                           "drivers": v.get("drivers"), "sentence": v.get("sentence")}
    _step("the findings carry a verdict", f.get("available") is True and v.get("status") == "ok",
          f"available={f.get('available')} status={v.get('status')}")
    bars = [r for r in rob.get("tornado") or [] if r.get("swing") is not None]
    EVIDENCE["tornado"] = [(r["key"], r["low_value"], r["high_value"],
                            round(r["npv_low"]), round(r["npv_high"])) for r in bars]
    _step("the tornado has bars on at least two rows, the demand charge re-dispatched",
          rob.get("status") == "ok" and len(bars) >= 2 and any(
              r["key"] == "demand_charge_price" and r["evaluation"] == "redispatch"
              for r in bars), str([r["key"] for r in bars]))
    _step("marginal is reached through the API: the demand-charge low bound flips the sign",
          v.get("class") == "marginal" and v.get("drivers") == ["demand_charge_price"],
          f"class={v.get('class')} drivers={v.get('drivers')}")
    centre = rob.get("npv_centre")
    dc = next((r for r in bars if r["key"] == "demand_charge_price"), {})
    others = [min(r["npv_low"], r["npv_high"]) for r in bars if r["key"] != "demand_charge_price"]
    EVIDENCE["marginal_margins"] = {"npv_centre": round(centre or 0),
                                    "demand_charge_low_bound": round(dc.get("npv_low") or 0),
                                    "lowest_other_bound": round(min(others)) if others else None}
    _step(f"marginal with margin: centre NPV > {CENTRE_MARGIN_EUR:,.0f} EUR",
          centre is not None and centre > CENTRE_MARGIN_EUR, f"centre={centre}")
    _step(f"the demand-charge low bound is below -{FLIP_MARGIN_EUR:,.0f} EUR and every "
          f"other bound above +{FLIP_MARGIN_EUR:,.0f} EUR",
          dc.get("npv_low") is not None and dc["npv_low"] < -FLIP_MARGIN_EUR
          and dc.get("npv_high", 0) > FLIP_MARGIN_EUR
          and bool(others) and min(others) > FLIP_MARGIN_EUR,
          str(EVIDENCE["marginal_margins"]))
    atts = {a["option_id"]: a for a in f.get("battery_attribution") or []}
    _step("the zero-size options are judged by size (skipped), the best by its NPV",
          atts.get("bess_2h", {}).get("status") == "skipped"
          and atts.get("bess_4h", {}).get("status") == "skipped"
          and atts.get("bess_1h", {}).get("status") == "ok"
          and v.get("option_id") == "bess_1h",
          str({k: (a["status"], a.get("battery_p_nom_mw")) for k, a in atts.items()}))
    return f


# ── 4. the best option's case reconciles ─────────────────────────────────

def _reconcile(sid: str, ledger: dict, run: dict, findings: dict) -> None:
    import pypsa

    from models.study import AssumptionsLedger
    from services.results.cost_breakdown import compute_cost_breakdown
    from services.solver.periodized_costs import upfront_cost_series
    from services.solver_service import with_periodized_cost_defaults
    from services.study import compile as C
    from services.study import packs
    from services.study import questions as Q

    c = qa_support.client()
    oid = findings["verdict"]["option_id"]
    r = c.get(_study_url(BASE, sid, f"/options/{oid}/case"))
    if not _step(f"the {oid} case is available", r.status_code == 200
                 and r.json().get("available") is True, f"HTTP {r.status_code}"
                 + ("" if r.status_code == 200 else f" {r.text[:300]}")):
        return
    case = r.json()
    res = next(o for o in findings["options"] if o["option_id"] == oid)
    p = next(s["p_nom_opt"] for s in res["sizes"] if s["asset"] == packs.BATTERY_NAME)
    led = AssumptionsLedger.model_validate(ledger)
    hours = Q.max_hours(Q.option(Q.BESS_AT_SITE, oid))
    upfront = packs.battery_upfront_eur_per_mw(led, hours)["total"]
    capex = sum(y["capex"] for y in case["years"])
    EVIDENCE["capex"] = {"p_nom_mw": p, "upfront_eur_per_mw": upfront,
                         "case_capex_total": case["kpis"]["capex_total"], "sum_years_capex": capex}
    _step("CAPEX = packs.battery_upfront_eur_per_mw x size, from the ledger",
          abs(case["kpis"]["capex_total"] - upfront * p) <= 1e-6 * upfront * p
          and abs(capex - upfront * p) <= 1e-6 * upfront * p,
          f"{case['kpis']['capex_total']:.2f} vs {upfront * p:.2f}")
    # FOM: the run record's asset economics, and cost_breakdown on the fork
    econ = (run["details"][oid].get("asset_economics") or {}).get("storage_units") or []
    fom = next((row["fom_cost_eur"] for row in econ if row.get("name") == packs.BATTERY_NAME),
               None)
    horizon = case["horizon_years"]
    opex_fixed = sum(y["opex_fixed"] for y in case["years"])
    fork_dir = qa_support.project_dir(f"{BASE}-opt-{oid}")
    n = pypsa.Network(str(fork_dir / "network.nc"))
    # U2 WP6/WP10: the config the fork was solved with — the engine's
    # commercial block compiled from the run's ledger and bound to the study's
    # minted export series (`compile`, the one path; no GS tariff is left).
    compiled = C.commercial_from_ledger(run["intake"], led, _library(), n.snapshots,
                                        export_series=run.get("export_series"))
    cfg = C.solver_config(led, compiled)
    bd = compute_cost_breakdown(n, cfg)
    su_fom = next((float(row["fom"]) for row in bd.get("by_component") or []
                   if row.get("component") == "StorageUnit"), None)
    EVIDENCE["fom"] = {"asset_economics_fom_cost_eur": fom, "cost_breakdown_storageunit_fom":
                       su_fom, "sum_years_opex_fixed": opex_fixed, "horizon_years": horizon}
    _step("fixed O&M = asset_economics.fom_cost_eur x horizon",
          fom is not None and abs(opex_fixed - fom * horizon) <= 1e-6 * max(1.0, fom * horizon),
          f"{opex_fixed:.2f} vs {fom} x {horizon}")
    _step("and = cost_breakdown's StorageUnit FOM on the fork's own network",
          fom is not None and su_fom is not None and abs(su_fom - fom) <= 1e-6 * max(1.0, fom),
          f"{su_fom} vs {fom}")
    with with_periodized_cost_defaults(n, cfg):
        back = float(upfront_cost_series(n, "StorageUnit").get(packs.BATTERY_NAME, float("nan")))
    EVIDENCE["upfront_cost_series_eur_per_mw"] = back
    EVIDENCE["upfront_cost_series_overstatement_pct"] = round(100 * (back / upfront - 1), 1)
    # U2 WP7 C1: the battery is written as its two upfront parts, which
    # `upfront_cost_series` sums — the S4 gap (a single-lifetime back-
    # calculation, 881,618 EUR/MW here at WP0) is closed (u2_deltas.json).
    _step("upfront_cost_series equals the ledger upfront on this pack (C1 parts)",
          abs(back - upfront) <= 1e-9 * upfront, f"{back:.0f} vs ledger {upfront:.0f} EUR/MW")
    x = c.get(_study_url(BASE, sid, f"/options/{oid}/case.xlsx"))
    _step("the case workbook downloads", x.status_code == 200 and x.content[:2] == b"PK")
    _one_engine(sid, led, run, findings, case, compiled)
    _record_figures(run, findings, case)


def _one_engine(sid: str, led, run: dict, findings: dict, case: dict, compiled) -> None:
    """
    U2 WP10: the study is valued on the Investment Case engine only — the
    forks carry its inputs, bound to the study's own export series; the case
    names its engines; its WACC gate is consistent; the §4.6 identity holds.
    """
    import json
    import uuid

    import pypsa

    from routers.projects import _solver_config_from_dict
    from services.study import compile as C
    from services.study import engine_adapter as A
    from tests.golden.oracle import annuity_pv_factor

    oid = case["option_id"]
    ref = run.get("export_series") or {}
    base_uuid = str(qa_support.project_row(BASE).id)
    blocks, refs = {}, {}
    for o in OPTIONS:
        raw = json.loads((qa_support.project_dir(f"{BASE}-opt-{o}") / "solver_config.json")
                         .read_text(encoding="utf-8"))
        blocks[o] = (raw.get("commercial") is not None, raw.get("finance") is not None,
                     "demand_charge" in raw)
        refs[o] = (raw.get("commercial") or {}).get("export_price_ref")
    _step("every option fork carries the engine's commercial block and every invested "
          "option its finance block, and no fork the GS demand charge",
          all(b[0] and not b[2] for b in blocks.values())
          and all(blocks[o][1] for o in OPTIONS if o != "none"), str(blocks))
    _step("every fork's export series is the study's own (decision-study:<base>:<study>:export)",
          C.parse_export_series_name(ref.get("id")) is not None
          and C.parse_export_series_name(ref.get("id")) == (uuid.UUID(base_uuid), sid)
          and all(r is not None and r.get("id") == ref.get("id")
                  and r.get("version") == ref.get("version") for r in refs.values()),
          f"run={ref.get('id')} forks={sorted({(r or {}).get('id') for r in refs.values()})}")
    bills = {o: (d.get("bill") or {}).get("engine") for o, d in (run.get("details") or {}).items()}
    _step("the case's engines are tariff_engine and finance_engine",
          case.get("engine") == "finance_engine"
          and case["sources"]["bill_refs"] == ["tariff_engine:counterfactual",
                                               f"tariff_engine:{oid}"]
          and {"tariff_engine", "finance_engine"} <= set(case["provenance"]["engines"])
          and set(bills.values()) == {"tariff_engine"},
          f"case={case.get('engine')} refs={case['sources']['bill_refs']} bills={bills}")
    fork = qa_support.project_dir(f"{BASE}-opt-{oid}")
    n = pypsa.Network(str(fork / "network.nc"))
    cfg = _solver_config_from_dict(json.loads((fork / "solver_config.json").read_text()))
    bundle = A.option_case(n, cfg, led, compiled=compiled, option_id=oid, study_id=sid,
                           fidelity=case.get("fidelity"))
    gate = bundle.result.gate
    EVIDENCE["wacc_gate"] = {"consistent": gate["wacc_vs_discount_rate_consistent"],
                             "legs": gate["legs"]}
    _step("the case's WACC gate is consistent at the centre (discount rate and asset "
          "rates ok, inflation n/a)",
          gate["wacc_vs_discount_rate_consistent"] is True
          and gate["legs"] == {"discount_rate": "ok", "asset_rates": "ok", "inflation": "n/a"}
          and abs(bundle.view.kpis.npv - case["kpis"]["npv"]) <= 1e-6 * abs(case["kpis"]["npv"]),
          str(EVIDENCE["wacc_gate"]))
    cost = {o["option_id"]: o.get("system_cost") for o in findings["options"]}
    af = annuity_pv_factor(case["discount_rate"], case["horizon_years"])
    npv = case["kpis"]["npv"]
    want = (cost["none"] - cost[oid]) * af
    EVIDENCE["identity_4_6"] = {"option": oid, "npv": npv, "lp_saving_x_af": want,
                                "rel": abs(npv - want) / abs(npv)}
    _step(f"the §4.6 identity holds on {oid}: NPV = LP objective saving x AF(r, H)",
          abs(npv - want) <= 1e-6 * abs(npv), str(EVIDENCE["identity_4_6"]))


def _record_figures(run: dict, findings: dict, case: dict) -> None:
    """
    Evidence only, no check (U2 plan WP0): the full-precision figures the U2
    comparison reads, taken from payloads this driver already fetched — the
    run record, the findings and the named option's case. No request, no
    step, no exit-code change.
    """
    from services.study import packs

    rob = findings.get("robustness") or {}
    EVIDENCE["sizes"] = {
        o["option_id"]: {s["asset"]: {"p_nom_opt": s.get("p_nom_opt"),
                                      "e_nom_opt": s.get("e_nom_opt")}
                         for s in o.get("sizes") or []}
        for o in findings.get("options") or []}
    EVIDENCE["attribution"] = {
        a["option_id"]: {"status": a.get("status"), "battery_p_nom_mw": a.get("battery_p_nom_mw"),
                         "battery_npv": a.get("battery_npv"), "option_npv": a.get("option_npv")}
        for a in findings.get("battery_attribution") or []}
    EVIDENCE["npv_bounds"] = {
        "npv_centre": rob.get("npv_centre"), "option_id": rob.get("option_id"),
        "tornado": {r["key"]: {"evaluation": r.get("evaluation"), "low_value": r["low_value"],
                               "high_value": r["high_value"], "npv_low": r.get("npv_low"),
                               "npv_high": r.get("npv_high")}
                    for r in rob.get("tornado") or []}}
    k = case.get("kpis") or {}
    econ = (run["details"][case["option_id"]].get("asset_economics") or {}).get(
        "storage_units") or []
    lcos_incl = next((row.get("lcos_eur_per_mwh") for row in econ
                      if row.get("name") == packs.BATTERY_NAME), None)
    # Gate U2-WP7 X5 (c): the finance engine's LCOS includes the charging
    # energy, so it is recorded under its own key, never the charging-free
    # one (the WP0 freeze's `lcos_excl_charging`; recorded delta with its
    # `post_figure`). A pro forma case (a fork written before WP7) keeps it.
    lcos_key = ("lcos_finance_engine" if case.get("engine") == "finance_engine"
                else "lcos_excl_charging")
    # U2 WP8 part B (§5.4): the case says `terminal_value_eur` and
    # `levelised_cost`; the evidence key follows (the WP0 freeze's
    # `salvage_eur`, moved by its recorded delta's `post_figure`).
    EVIDENCE["case_kpis"] = {
        "option_id": case["option_id"], "npv": k.get("npv"), "irr": k.get("irr"),
        "payback_simple": k.get("payback_simple"),
        "payback_discounted": k.get("payback_discounted"), "capex_total": k.get("capex_total"),
        "terminal_value_eur": k.get("terminal_value_eur"), lcos_key: k.get("levelised_cost"),
        "lcos_incl_charging": lcos_incl,
        "ledger_hash": (case.get("provenance") or {}).get("ledger_hash")}
    EVIDENCE["bills"] = {
        oid: {"total": (d.get("bill") or {}).get("total"),
              "annual_bill": (d.get("bill") or {}).get("annual_bill"),
              "by_component": (d.get("bill") or {}).get("by_component")}
        for oid, d in sorted((run.get("details") or {}).items())}


def _library():
    # U2 WP8 (gate C4): the study seeds and compiles from the pinned pack.
    from services.study import library as L

    return L.load_defaults()


# ── 5. the report ────────────────────────────────────────────────────────

def _report(sid: str) -> None:
    import docx
    import openpyxl

    c = qa_support.client()
    r = c.post(_study_url(BASE, sid, "/report"))
    if not _step("the report assembles", r.status_code == 200, f"HTTP {r.status_code}"
                 + ("" if r.status_code == 200 else f" {r.text[:300]}")):
        return
    body = r.json()
    _step("the report is available and fresh", body.get("available") is True
          and body.get("stale") is False, f"stale={body.get('stale')}")
    codes = [d["code"] for d in body.get("required_disclosures") or []]
    EVIDENCE["report_disclosures"] = codes
    _step("the report discloses the by-construction NPV and reads the upload as measured",
          "npv_nonnegative_at_optimum_by_construction" in codes
          and "demand_charge_perfect_foresight" in codes
          and "synthetic_load_understates_peak_shaving" not in codes, f"{len(codes)} codes")
    h = c.get(_study_url(BASE, sid, "/report.html"))
    _step("the report renders to HTML (attachment, CSP sandbox, charts inline)",
          h.status_code == 200 and h.headers.get("content-security-policy") == "sandbox"
          and h.text.count("data:image/png;base64,") >= 3, f"HTTP {h.status_code}")
    d = c.get(_study_url(BASE, sid, "/report.docx"))
    ok = d.status_code == 200
    if ok:
        doc = docx.Document(io.BytesIO(d.content))
        ok = any(t.rows[0].cells[0].text == "Assumption" for t in doc.tables)
    _step("the report renders to DOCX with the Assumptions table", ok, f"HTTP {d.status_code}")
    if ok:
        texts = [p.text for p in doc.paragraphs]
        first = next(i for i, t in enumerate(texts) if t.startswith("1. Executive summary"))
        second = next(i for i, t in enumerate(texts) if t.startswith("2. "))
        summary = texts[first:second]
        hs = h.text.index('id="executive_summary"')
        _step("the marginal verdict's drivers are listed by label in the HTML and DOCX summary",
              any("Demand charge on peak import" in t for t in summary)
              and "Demand charge on peak import" in h.text[hs:h.text.index("</section>", hs)],
              " | ".join(summary)[:300])
    x = c.get(_study_url(BASE, sid, "/report.xlsx"))
    ok = x.status_code == 200
    sheets: list[str] = []
    if ok:
        sheets = openpyxl.load_workbook(io.BytesIO(x.content)).sheetnames
        ok = {"Verdict", "Tornado", "Assumptions", "Provenance", "Cash flows bess_1h"} <= set(sheets)
    _step("the report renders to XLSX", ok, str(sheets))
    r = c.put(_study_url(BASE, sid, "/ledger"), json={"rows": [
        {"key": "battery_storage_eur_per_kwh", "value": STORAGE_QUOTE + 10.0,
         "unit": "EUR/kWh"}]})
    after = c.get(_study_url(BASE, sid, "/report")).json()
    _step("stale flips after a ledger edit",
          r.status_code == 200 and after.get("stale") is True
          and after.get("stale_reasons") == ["ledger_changed_since_findings"],
          f"stale={after.get('stale')} reasons={after.get('stale_reasons')}")
    case = c.get(_study_url(BASE, sid, "/options/bess_1h/case"))
    _step("and the case refuses on the same rule", case.status_code == 409, str(case.status_code))


# ── 6. abort mid-run ─────────────────────────────────────────────────────

def _abort_mid_run() -> str | None:
    c = qa_support.client()
    before = _hashes()
    sid = _create(before, ABORT_BASE)
    if sid is None:
        return None
    before = _hashes(exclude=[ABORT_BASE])
    r = c.post(_study_url(ABORT_BASE, sid, "/run"), json={})
    if not _step("the second run starts", r.status_code == 202,
                 f"HTTP {r.status_code}" + ("" if r.status_code == 202 else f" {r.text[:300]}")):
        return sid
    live = _poll(_study_url(ABORT_BASE, sid, "/run"), timeout=600,
                 until=lambda b: b.get("status") != "running" or len(b.get("solved") or []) >= 1)
    a = c.post(_study_url(ABORT_BASE, sid, "/run/abort"))
    _step("the abort is accepted mid-run", a.status_code == 200 and a.json().get("aborting") is True,
          f"solved at abort={live.get('solved')} answer={a.text[:200]}")
    rec = _poll(_study_url(ABORT_BASE, sid, "/run"), timeout=600)
    _step("the aborted run says so and names the options it did not reach",
          rec.get("status") == "aborted" and rec.get("pending"),
          f"status={rec.get('status')} solved={rec.get('solved')} pending={rec.get('pending')}")
    EVIDENCE["abort"] = {"solved": rec.get("solved"), "pending": rec.get("pending")}
    names = _all_projects()
    unsolved = [o for o in OPTIONS if o not in (rec.get("solved") or [])]
    left = [o for o in unsolved if f"{ABORT_BASE}-opt-{o}" in names]
    kept = [o for o in rec.get("solved") or [] if f"{ABORT_BASE}-opt-{o}" in names]
    _step("the unsolved forks are removed, the solved ones kept",
          left == [] and kept == list(rec.get("solved") or []), f"left={left} kept={kept}")
    study = c.get(_study_url(ABORT_BASE, sid)).json()
    findings = qa_support.project_dir(ABORT_BASE) / study["findings_ref"]
    import json

    faux = json.loads(findings.read_text())
    _step("the aborted run's options are not established",
          faux["options_status"] == "not_established" and faux["pending_options"],
          str(faux["pending_options"]))
    changed = _changed(before)
    _step("the aborted run changed no pre-existing project", changed == [], str(changed))
    _step("nothing is left marked study-owned after the abort",
          PyPSAService._study_owned == set(), str(sorted(PyPSAService._study_owned)))
    return sid


# ── 7. the cap invariant, 8. delete ──────────────────────────────────────

def _cap_invariant(cap: dict, user_before: dict[str, str]) -> None:
    import pypsa

    edited = cap["edited"]
    still = PyPSAService.get_context(_key(USERS[1]))
    _step("the user project resident at the cap was not evicted",
          still is edited and "qa_unsaved_edit_bus" in still.network.buses.index)
    on_disk = pypsa.Network(str(qa_support.project_dir(USERS[1]) / "network.nc"))
    _step("and its unsaved edit was not saved", "qa_unsaved_edit_bus" not in on_disk.buses.index)
    _step("no user context was evicted", cap["resident"] <= set(PyPSAService._contexts),
          str(sorted(cap["resident"] - set(PyPSAService._contexts))))
    changed = _changed(user_before)
    _step("no user project changed on disk over the whole journey", changed == [], str(changed))
    _step("neither study base is left loaded",
          _key(BASE) not in PyPSAService._contexts and _key(ABORT_BASE) not in PyPSAService._contexts)


def _delete(base: str, sid: str) -> None:
    c = qa_support.client()
    forks = [n for n in _all_projects() if n.startswith(f"{base}-opt-")]
    dirs = [qa_support.project_dir(n) for n in forks]
    r = c.delete(_study_url(base, sid))
    names = _all_projects()
    _step(f"deleting the study cascades to its {len(forks)} fork(s)",
          r.status_code == 204 and not any(n in names for n in forks)
          and not any(d.exists() for d in dirs) and base in names,
          f"HTTP {r.status_code} left={[n for n in forks if n in names]}")
    _step("the study record is gone", c.get(_study_url(base, sid)).status_code == 404)


def _judge() -> None:
    """
    Plan §4 judging rule (U2 WP10): the verdict class, its drivers, the named
    option and the skipped options equal the WP0 record; every recorded
    figure equals WP0 within the fixture's tolerances, or is a recorded delta
    (`tests/fixtures/u2_deltas.json`) with a cause. `EVIDENCE["u2_deltas"]`
    lists each recorded delta and any figure outside tolerance as
    `{figure, pre, post, abs, rel, cause}` (cause None: unrecorded).
    """
    import json

    from tests import test_u2_pre_numbers_recorded as PRE
    from tests import u2_record_pre_numbers as REC

    frozen = json.loads(REC.FIXTURE.read_text(encoding="utf-8"))
    tol = frozen["tolerances"]
    try:
        got = PRE._json_round_trip(REC.driver_section({**EVIDENCE, "checks": None}))
    except KeyError as exc:
        _step("the §4 judging rule can read the run's figures", False, f"missing {exc}")
        return
    want = {k: v for k, v in frozen["driver"].items() if k != "checks"}
    v, w = got["verdict"], want["verdict"]
    _step("§4 judging rule: the verdict class, drivers, named option and skipped options "
          "equal WP0", v == w, f"got={v} WP0={w}")
    rows = [r for r in json.loads(PRE.DELTAS.read_text(encoding="utf-8"))["deltas"]
            if r.get("test") == PRE.DRIVER_TEST_ID]
    moved = PRE._with_recorded_deltas(want, "driver", PRE.DRIVER_TEST_ID, tol, rows)
    got.pop("checks", None)
    unrecorded = PRE._diff(moved, got, tol, "driver")
    out = [{k: r.get(k) for k in ("figure", "post_figure", "pre", "post", "abs", "rel", "cause")}
           for r in rows]
    out += [{"figure": line, "pre": None, "post": None, "abs": None, "rel": None,
             "cause": None} for line in unrecorded]
    EVIDENCE["u2_deltas"] = out
    _step("§4 judging rule: every figure equals WP0 within tolerance or is a recorded delta "
          "with a cause", unrecorded == [] and all(r.get("cause") for r in rows),
          "; ".join(unrecorded)[:400])


def main_() -> int:
    t0 = time.monotonic()
    _enable_studies()
    print("QA: the guided decision study, over HTTP with a real LP\n")
    cap = _user_projects_at_the_cap()
    user_before = _hashes()
    sid = _create(user_before, BASE)
    if sid is None:
        return 1
    ledger = _edit_ledger(sid)
    run = _run(sid) if ledger is not None else None
    findings = _tornado(sid) if run is not None else None
    if findings is not None:
        _reconcile(sid, ledger, run, findings)
        _report(sid)
        _judge()
    abort_sid = _abort_mid_run()
    _cap_invariant(cap, user_before)
    _delete(BASE, sid)
    if abort_sid is not None:
        _delete(ABORT_BASE, abort_sid)
    EVIDENCE["total_s"] = round(time.monotonic() - t0, 1)
    print("\nEvidence:")
    for k, v in EVIDENCE.items():
        print(f"  {k}: {v}")
    print(f"\n{PASS} passed, {FAIL} failed ({EVIDENCE['total_s']} s)")
    # U2 plan WP0: `QA_DECISION_STUDY_EVIDENCE_OUT=<path>` dumps EVIDENCE as
    # JSON (read by `tests/u2_record_pre_numbers.py`); nothing else changes.
    out = os.environ.get("QA_DECISION_STUDY_EVIDENCE_OUT")
    if out:
        import json

        pathlib.Path(out).write_text(json.dumps(
            {**EVIDENCE, "checks": {"passed": PASS, "failed": FAIL}}, indent=1, default=str),
            encoding="utf-8")
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main_())
