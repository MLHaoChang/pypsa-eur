"""
The results workbook, package 01 (`.scratch/results-workbook/issues/01-engine-overview-finance.md`):
the Overview's figures equal the services the Results tabs read, the
investment-case block equals the stored report (or is `not_established`
without one, nothing run), the charts are native and read the sheet's own
cells, a stale run is refused, formula-like text stays text, the 25 MB cap
holds, and the workbook opens in openpyxl and (when installed) LibreOffice.
"""
from __future__ import annotations

import io
import shutil
import subprocess

import pandas as pd
import pypsa
import pytest
from openpyxl import load_workbook

from models.finance import IC_REPORT_SECTIONS, IcSectionState, InvestmentCaseReport
from routers import results as results_router
from services import results_workbook as RW
from services.finance.metrics import npv
from services.finance.report import IC_REPORT_STORE_KEY
from services.results.asset_economics import compute_asset_economics
from services.results.cost_breakdown import compute_cost_breakdown
from services.results.economics_by_carrier import compute_economics_by_carrier
from services.results_workbook import registry
from services.results_workbook.overview import CASE_ROWS, RUN_THE_CASE
from services.solver_service import SolverConfig

URL = "/api/results/workbook.xlsx"
EVIL = '=HYPERLINK("http://x","y")'
PERIODS = ["2030", "2040"]


def _network() -> pypsa.Network:
    n = pypsa.Network()
    hours = pd.date_range("2030-01-01", periods=4, freq="h")
    mi = pd.MultiIndex.from_product([[2030, 2040], hours], names=["period", "timestep"])
    n.set_snapshots(mi)
    n.investment_periods = [2030, 2040]
    n.snapshot_weightings.loc[:, :] = 2190.0  # four hours stand for a year
    n.add("Carrier", ["gas", "solar", EVIL, "AC"])
    n.add("Bus", "A", carrier="AC")
    n.add("Load", "L", bus="A", p_set=[50.0, 60.0, 70.0, 55.0] * 2)
    n.add("Generator", "gas", bus="A", carrier="gas", p_nom_extendable=True,
          capital_cost=60_000.0, marginal_cost=45.0)
    n.add("Generator", "solar", bus="A", carrier="solar", p_nom_extendable=True,
          capital_cost=25_000.0, marginal_cost=0.5, p_max_pu=[0.0, 0.6, 0.9, 0.3] * 2)
    n.add("Generator", "evil", bus="A", carrier=EVIL, p_nom=5.0, capital_cost=1_000.0,
          marginal_cost=20.0)
    n.optimize(solver_name="highs", multi_investment_periods=True)
    return n


@pytest.fixture(scope="module")
def solved():
    return _network()


def _rdf(n, accessor, attr, source="lopf"):
    return getattr(getattr(n, accessor), attr, None)


def _report() -> InvestmentCaseReport:
    years = [2029, 2030, 2031, 2032, 2033]
    cash = [-1_000_000.0, 260_000.0, 270_000.0, 280_000.0, 400_000.0]
    rate = 0.07
    sections = {s: IcSectionState(status="skipped", note="not in this fixture")
                for s in IC_REPORT_SECTIONS}
    sections["project"] = IcSectionState(status="ok", payload={
        "years": years, "cash": {"project_post_tax": cash}, "wacc_nominal": rate,
        "project_pre_tax_npv": 123_456.0, "equity_pre_tax_irr": 0.15,
        "equity_post_tax_irr": 0.13, "payback_years": 3.4, "currency": "EUR",
        "currency_year": 2025})
    sections["debt"] = IcSectionState(status="ok", payload={"amount": 650_000.0})
    return InvestmentCaseReport(
        case_id="case-rw", assumptions_hash="0123456789abcdef", npv_at_wacc=npv(rate, cash),
        project_irr_pre_tax=0.11, project_irr_post_tax=0.095, min_dscr=1.25, avg_dscr=1.6,
        lcoe_finance_consistent_eur_per_mwh=71.5, sections=sections,
        completeness={s: sections[s].status for s in IC_REPORT_SECTIONS})


def _build(n, *, report=None, views=(), **kw) -> bytes:
    return RW.build_results_workbook(RW.WorkbookInputs(
        network=n, cfg=SolverConfig(), lost_load=None, result_df=_rdf, views=list(views),
        ic_report=report, **kw))


def _wb(blob):
    return load_workbook(io.BytesIO(blob))


def _row(ws, label):
    for row in ws.iter_rows():
        if row[0].value == label:
            return row
    raise AssertionError(f"no row {label!r}")


def _approx(a, b):
    return a == pytest.approx(b, rel=1e-9, abs=1e-12)


# ── the Overview's figures ───────────────────────────────────────────────

def test_overview_matches_the_services_per_period_and_over_the_horizon(solved):
    n = solved
    econ = compute_economics_by_carrier(n, SolverConfig(), None, result_df=_rdf)["by_carrier"]
    cb = compute_cost_breakdown(n, SolverConfig())
    ws = _wb(_build(n))["Overview"]
    assert [c.value for c in ws[4][:4]] == ["Metric", *PERIODS, "Horizon"]

    def check(label, field_):
        row = _row(ws, label)
        for j, p in enumerate(PERIODS, start=1):
            assert _approx(row[j].value, sum(v[field_]["by_period"].get(p, 0.0) for v in econ.values()))
        assert _approx(row[3].value, sum(v[field_]["total"] for v in econ.values()))

    check("Capex, annualised (incl. fixed O&M)", "capex_meur")
    check("Opex", "opex_meur")
    check("  of which variable opex (marginal cost)", "gen_cost_meur")
    total = _row(ws, "Total system cost")
    cb_pp = {str(r["period"]): r["total"] for r in cb["by_period"]}
    for j, p in enumerate(PERIODS, start=1):
        assert _approx(total[j].value, cb_pp[p] / 1e6)
    assert _approx(total[3].value, cb["total"] / 1e6)
    assert _approx(_row(ws, "Capex, overnight (present value of the upfront investment)")[3].value,
                   cb["capex_lifetime"] / 1e6)

    # Energy served: load × the energy weighting (no lost load in this fixture).
    w = n.snapshot_weightings.generators * n.investment_period_weightings.years.reindex(
        n.snapshots.get_level_values(0)).to_numpy()
    load = n.loads_t.p.sum(axis=1) * w
    served = _row(ws, "Energy served")
    for j, p in enumerate(PERIODS, start=1):
        assert _approx(served[j].value, load.xs(int(p), level=0).sum() / 1e3)
    assert _approx(served[3].value, load.sum() / 1e3)

    # The per-asset economics (Σ VOM == variable opex: generators only here).
    ae = compute_asset_economics(n, SolverConfig(), result_df=_rdf)
    vom = sum(a["vom_cost_eur"] for a in ae["generators"])
    assert _approx(_row(ws, "  of which variable opex (marginal cost)")[3].value, vom / 1e6)

    # System LCOE is a live formula over the cost and energy cells.
    lcoe = _row(ws, "System LCOE")
    c, e = total[3].coordinate, served[3].coordinate
    assert lcoe[3].value == f"={c}*1000/{e}"

    # Per-carrier LCOE.
    for carrier in ("gas", "solar"):
        row = _row(ws, carrier)
        cell = econ[carrier]["lcoe_eur_per_mwh"]
        for j, p in enumerate(PERIODS, start=1):
            assert _approx(row[j].value, cell["by_period"][p])
        assert _approx(row[3].value, cell["total"])

    # Fuel is not separable in PyPSA: stated, not invented.
    assert _row(ws, "Fuel")[3].value == RW.NOT_ESTABLISHED


# ── the investment case ──────────────────────────────────────────────────

def test_without_a_case_every_kpi_is_not_established_and_nothing_is_run(solved, monkeypatch):
    import services.finance.engine as engine

    monkeypatch.setattr(engine, "run_case", lambda *a, **k: pytest.fail("export ran the case"))
    wb = _wb(_build(solved, views=["investment"]))
    ws = wb["Overview"]
    for _key, label, _unit, _read in CASE_ROWS:
        row = _row(ws, label)
        assert row[1].value == RW.NOT_ESTABLISHED and row[3].value == RUN_THE_CASE
    assert not [s for s in wb.sheetnames if s.startswith("Case ")]
    assert len(ws._charts) == 2
    about = {r[0].value: r[2].value for r in wb["About"].iter_rows() if r[0].value}
    assert about["Investment case"] == f"no results: {RUN_THE_CASE}"


def test_with_a_case_the_kpis_equal_the_report_and_the_npv_is_live(solved):
    rep = _report()
    wb = _wb(_build(solved, report=rep, views=["investment"]))
    ws = wb["Overview"]
    rows = {key: _row(ws, label) for key, label, _u, _r in CASE_ROWS}
    for key, _label, _unit, read in CASE_ROWS:
        if key == "npv_project_post_tax":
            continue
        assert rows[key][1].value == read(rep), key
    # NPV: the KPI cell reads the live formula `=CF0+NPV(rate,CF1:CFn)`.
    npv_cell = rows["npv_project_post_tax"][1].value
    assert npv_cell.startswith("=B")
    npv_row = int(npv_cell[2:])
    expr = ws.cell(npv_row, 2).value
    assert expr.startswith("=C") and "+NPV($B$" in expr
    rate_row = int(expr.split("$B$")[1].split(",")[0])
    assert ws.cell(rate_row, 2).value == 0.07
    first = int(expr[2:].split("+")[0])
    cash = [ws.cell(first + i, 3).value for i in range(5)]
    assert cash == rep.sections["project"].payload["cash"]["project_post_tax"]
    # Evaluating the formula by hand gives the report's NPV.
    assert _approx(cash[0] + sum(c / 1.07 ** (i + 1) for i, c in enumerate(cash[1:])),
                   rep.npv_at_wacc)
    assert _approx(_row(ws, "NPV in the stored report")[1].value, rep.npv_at_wacc)
    # Finance sheets come from export_xlsx, prefixed.
    case = [s for s in wb.sheetnames if s.startswith("Case ")]
    assert {"Case About", "Case Summary", "Case CashflowLines"} <= set(case)
    assert {f"Case {s}"[:31] for s in IC_REPORT_SECTIONS} <= set(case)
    assert len(ws._charts) == 3


def test_the_charts_reference_the_overviews_own_cells(solved):
    ws = _wb(_build(solved, report=_report()))["Overview"]
    refs = []
    for chart in ws._charts:
        for plot in [chart, *getattr(chart, "_charts", [])]:
            for s in plot.series:
                refs.append(s.val.numRef.f)
                if s.cat is not None:
                    refs.append((s.cat.numRef or s.cat.strRef).f)
    assert refs and all(r.startswith("'Overview'!") for r in refs), refs
    titles = [c.title.tx.rich.p[0].r[0].t for c in ws._charts]
    assert titles == ["Capex (annualised) vs opex by period", "LCOE by carrier by period",
                      "Project post-tax cash flow and cumulative NPV"]


# ── text, size and views ─────────────────────────────────────────────────

def test_formula_like_text_is_stored_as_text(solved):
    wb = _wb(_build(solved, project=EVIL))
    row = _row(wb["Overview"], EVIL.lower())  # economics_by_carrier lower-cases carriers
    assert row[0].data_type == "s"
    about = _row(wb["About"], "Project")
    assert about[1].value == EVIL and about[1].data_type == "s"


def test_the_25_mb_cap_holds(solved, monkeypatch):
    monkeypatch.setattr(RW, "MAX_BYTES", 1000)
    with pytest.raises(RW.WorkbookTooLarge):
        _build(solved)


def test_views_resolve_from_the_selected_cards_plus_extras():
    scope = {"groups": ["commercial"], "options": {"commercial": ["investment_case"]}}
    assert registry.resolve_views(scope) == ["investment"]
    assert registry.resolve_views(scope, ["prices", "", "prices"]) == ["prices", "investment"]
    assert registry.resolve_views(None, all_=True) == [v.view_id for v in registry.workbook_views()]
    with pytest.raises(registry.UnknownView):
        registry.resolve_views(scope, ["nope"])
    with pytest.raises(ValueError):
        registry.register("summary")


def test_round_trip_and_libreoffice_recalculates(solved, tmp_path):
    rep = _report()
    blob = _build(solved, report=rep, views=["investment"])
    wb = _wb(blob)
    assert wb.sheetnames[:2] == ["About", "Overview"]
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        pytest.skip("LibreOffice is not installed here")
    src = tmp_path / "in" / "wb.xlsx"
    src.parent.mkdir()
    src.write_bytes(blob)
    out = tmp_path / "out"
    subprocess.run([soffice, "--headless", "--norestore", f"-env:UserInstallation=file://{tmp_path}/lo",
                    "--convert-to", "xlsx", "--outdir", str(out), str(src)],
                   check=True, capture_output=True, timeout=180)
    calc = load_workbook(out / "wb.xlsx", data_only=True)["Overview"]
    npv_row = _row(calc, "NPV (live formula)")
    assert _approx(npv_row[1].value, rep.npv_at_wacc)
    total, served = _row(calc, "Total system cost"), _row(calc, "Energy served")
    assert _approx(_row(calc, "System LCOE")[3].value, total[3].value * 1000 / served[3].value)


# ── the route ─────────────────────────────────────────────────────────────

def test_route_204_before_a_solve(client, install_network):
    n = pypsa.Network()
    n.add("Bus", "A")
    install_network(n)
    assert client.get(URL).status_code == 204


def test_route_refuses_a_stale_run_and_exports_the_stored_one_on_request(client, install_network):
    install_network(_network())
    r = client.get(URL)  # no solve fingerprint recorded: not current
    assert r.status_code == 409 and r.json()["detail"]["error_kind"] == "results_workbook_stale"
    assert "no input fingerprint" in r.json()["detail"]["message"]
    with pytest.MonkeyPatch.context() as mp:
        mp.setattr(results_router, "_solver_in_flight", lambda: True)
        r = client.get(URL, params={"stored": "true"})
        assert r.status_code == 409 and r.json()["detail"]["error_kind"] == "results_workbook_stale"
    r = client.get(URL, params={"stored": "true"})
    assert r.status_code == 200
    about = {row[0].value: row[1].value for row in _wb(r.content)["About"].iter_rows()}
    assert about["Results"].startswith("stored run")


def test_route_exports_a_current_run_with_the_stored_case(client, install_network, monkeypatch):
    import services.workflow_tools as wt

    install_network(_network())
    monkeypatch.setattr(wt, "solve_freshness",
                        lambda ctx=None: {"inputs_changed_since_solve": False, "results_current": True})
    results_router._state[IC_REPORT_STORE_KEY] = _report().model_dump(mode="json")
    try:
        r = client.get(URL, params={"views": "investment,prices"})
        assert r.status_code == 200
        wb = _wb(r.content)
        assert "Case Summary" in wb.sheetnames
        assert client.get(URL, params={"views": "bogus"}).json()["detail"]["error_kind"] \
            == "unknown_result_view"
        monkeypatch.setattr(RW, "MAX_BYTES", 1000)
        r = client.get(URL)
        assert r.status_code == 413 and r.json()["detail"]["error_kind"] == "file_too_large"
    finally:
        results_router._state.pop(IC_REPORT_STORE_KEY, None)
