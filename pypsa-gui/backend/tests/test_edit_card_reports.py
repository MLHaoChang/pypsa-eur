"""Old -> new edit cards for the report tools (Q3-CARDS-1): set_report_template,
set_report_mapping and import_edited_report.

Each card shows the stored value beside the value the route would store; the receipt reads the
stores back and verifies new == stored; a preview writes nothing; a caller without access is
refused before any old value is read.
"""
from __future__ import annotations

import io

import pytest
from docx import Document
from fastapi import HTTPException

import harness.edit_card_reports  # noqa: F401  (registers the adapters; the integrator imports it too)
from harness import edit_card
from services import chat_service, chat_tools as T
from services.reports import store
from tests.study_s4_support import dir_hash
from tests.test_chat_report_roundtrip_tools import _docx, _upload as _upload_copy
from tests.test_chat_report_template_tools import (  # noqa: F401  (fixtures)
    PROJECT, _evidence_report, _plan_for, _upload_template, fake_untagged, fixtures, project)
from tests.test_edit_card import _dispatch


@pytest.fixture
def session(project):
    chat = chat_service.ChatSession()
    T.set_chat_session(chat)
    yield chat
    T.set_chat_session(None)


def _cells(card):
    return [(r["field"], r["old"], r["new"]) for r in card["rows"] if r["role"] == "typed"]


def _refusal(tool, args):
    with pytest.raises(HTTPException) as exc:
        edit_card.preview(tool, args)
    return exc.value


def _edited(data: bytes, suffix: str = " (checked by the user)") -> bytes:
    """The exported copy with one body paragraph edited in Word."""
    doc = Document(io.BytesIO(data))
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("Not established") and paragraph.runs:
            paragraph.runs[-1].text += suffix
            break
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ── set_report_template ─────────────────────────────────────────────────────

def test_template_bind_shows_default_to_the_file_then_type_and_language(project, fixtures):
    rid = _evidence_report()
    tid = _upload_template(project, "corporate_de.docx", fixtures["corporate_untagged_de.docx"])
    card = edit_card.preview("set_report_template", {"report_id": rid, "file_id": tid})
    assert card["complete"] and card["kind"] == "settings_edit" and card["section"]["id"] == "results"
    cells = {f: (o, n) for f, o, n in _cells(card)}
    assert cells["template_file_id"] == (None, "corporate_de.docx")
    assert cells["template_mode"] == (None, "untagged")
    assert cells["template_language"][0] is None
    assert "Template Built-in default → corporate_de.docx" in card["says"]


def test_template_swap_clears_the_stored_plan_and_the_receipt_verifies(project, fixtures, session, fake_untagged):
    rid = _evidence_report()
    first = _upload_template(project, "first.docx", fixtures["corporate_untagged.docx"])
    second = _upload_template(project, "second.docx", fixtures["tagged_minimal.docx"])
    T.set_report_template(rid, first)
    T.set_report_mapping(rid, _plan_for(T.get_report_template(rid)["outline"]))
    frames, cards, result = _dispatch(session, "set_report_template", {"report_id": rid, "file_id": second})
    rows = {r["field"]: r for r in cards[0]["edit_card"]["rows"]}
    assert (rows["template_file_id"]["old"], rows["template_file_id"]["new"]) == ("first.docx", "second.docx")
    assert (rows["template_mode"]["old"], rows["template_mode"]["new"]) == ("untagged", "tagged")
    assert rows["mapping_plan"]["old"] == "3 heading entries" and rows["mapping_plan"]["new"] is None
    assert rows["mapping_plan"]["words"]["new"] == "cleared"
    assert result["_edit_receipt"]["outcome"] == "verified"
    assert store.load_meta(project, rid).mapping_plan is None


def test_template_unbind_and_noop_refusals(project, fixtures):
    rid = _evidence_report()
    assert _refusal("set_report_template", {"report_id": rid}).detail["error_kind"] == "empty_edit"
    tid = _upload_template(project, "t.docx", fixtures["tagged_minimal.docx"])
    T.set_report_template(rid, tid)
    assert _refusal("set_report_template", {"report_id": rid, "file_id": tid}).detail["error_kind"] == "empty_edit"
    card = edit_card.preview("set_report_template", {"report_id": rid})
    assert ("template_file_id", "t.docx", None) in _cells(card)
    assert card["rows"][0]["words"]["new"] == "Built-in default"
    refused = _refusal("set_report_template", {"report_id": rid, "file_id": "no-such-file"})
    assert refused.status_code == 404  # the route's own upload_not_found


def test_template_not_a_template_is_the_routes_refusal(project):
    from services import upload_service
    from services.reports.docx_writer import DOCX_MIME

    rid = _evidence_report()
    meta = upload_service.add_upload(PROJECT, b"x", "notes.txt", "text/plain", kind="user_upload", project_dir=project)
    assert _refusal("set_report_template", {"report_id": rid, "file_id": meta.file_id}).detail["error_kind"] \
        == "template_not_a_template"


# ── set_report_mapping ──────────────────────────────────────────────────────

def _untagged_report(project, fixtures):
    rid = _evidence_report()
    T.set_report_template(rid, _upload_template(project, "u.docx", fixtures["corporate_untagged.docx"]))
    return rid, T.get_report_template(rid)["outline"]


def test_mapping_shows_only_what_differs_in_the_panels_words(project, fixtures, session, fake_untagged):
    rid, outline = _untagged_report(project, fixtures)
    plan = _plan_for(outline)
    frames, cards, result = _dispatch(session, "set_report_mapping", {"report_id": rid, "plan": plan})
    rows = cards[0]["edit_card"]["rows"]
    typed = {r["field"]: r for r in rows if r["role"] == "typed"}
    renamed = next(r for r in typed.values() if r["field"].endswith(".new_text"))
    assert (renamed["old"], renamed["new"]) == (None, "4 Residual (renamed)")
    assert renamed["label"].startswith("Template heading ") and "· New text" in renamed["label"]
    action = [r for r in typed.values() if r["field"].endswith(".action")]
    assert {(r["old"], r["new"]) for r in action} == {("keep", "rename"), ("keep", "drop")}
    sections = next(r for r in typed.values() if r["field"].endswith(".section_ids"))
    assert sections["old"] is None and sections["new"]
    assert not any(r["field"].endswith(".action") and r["new"] == "keep" for r in typed.values())
    assert result["_edit_receipt"]["outcome"] == "verified"
    # the same plan again: nothing differs
    assert _refusal("set_report_mapping", {"report_id": rid, "plan": plan}).detail["error_kind"] == "empty_edit"


def test_mapping_dropped_entries_are_noted_and_strict_is_the_routes_refusal(project, fixtures, fake_untagged):
    rid, outline = _untagged_report(project, fixtures)
    plan = _plan_for(outline)
    plan["entries"].append({"heading_index": 9999, "action": "keep", "new_text": None, "section_ids": []})
    card = edit_card.preview("set_report_mapping", {"report_id": rid, "plan": plan})
    notes = [r["note"] for r in card["rows"] if r["role"] == "note"]
    assert any("9999" in n for n in notes)
    refused = _refusal("set_report_mapping", {"report_id": rid, "plan": plan, "strict": True})
    assert refused.detail["error_kind"] == "invalid_mapping_plan"
    assert _refusal("set_report_mapping", {"report_id": rid, "plan": "nope"}).detail["error_kind"] == "invalid_edit_value"


def test_mapping_untagged_only_and_needs_a_template(project, fixtures, fake_untagged):
    rid = _evidence_report()
    assert _refusal("set_report_mapping", {"report_id": rid, "plan": {}}).detail["error_kind"] == "no_template"
    T.set_report_template(rid, _upload_template(project, "tag.docx", fixtures["tagged_minimal.docx"]))
    assert _refusal("set_report_mapping", {"report_id": rid, "plan": {}}).detail["error_kind"] == "template_not_untagged"


def test_mapping_caps_rows_and_hides_secret_like_placeholders(project, fixtures, fake_untagged):
    rid, outline = _untagged_report(project, fixtures)
    plan = {"entries": [], "inserted": [], "unmapped_sections": [], "notes": [],
            "placeholders": {"api_key": "hunter2", **{f"p{i}": f"v{i}" for i in range(edit_card.STUDY_LEAF_CAP + 2)}}}
    card = edit_card.preview("set_report_mapping", {"report_id": rid, "plan": plan})
    typed = [r for r in card["rows"] if r["role"] == "typed"]
    assert len(typed) == edit_card.STUDY_LEAF_CAP
    assert "hunter2" not in repr(card)
    secret = next(r for r in typed if r["field"] == "placeholders.api_key")
    assert (secret["old"], secret["new"]) == (None, "set")
    assert any(r["role"] == "note" and "more" in r["note"] for r in card["rows"])
    # a rotated secret is not "unchanged", and still never appears
    T.set_report_mapping(rid, plan)
    rotated = edit_card.preview("set_report_mapping", {"report_id": rid, "plan": {**plan, "placeholders": {
        **plan["placeholders"], "api_key": "hunter3"}}})
    assert _cells(rotated) == [("placeholders.api_key", "set", "set (changed)")]
    assert "hunter" not in repr(rotated)
    args = {"report_id": rid, "plan": {**plan, "placeholders": {**plan["placeholders"], "api_key": "hunter3"}}}
    T.set_report_mapping(rid, args["plan"])
    assert edit_card.ADAPTERS["set_report_mapping"].readback(args, rotated["rows"]) == [r["new"] for r in rotated["rows"]]


# ── import_edited_report ────────────────────────────────────────────────────

def test_import_shows_the_version_edited_section_and_template_then_verifies(project, session):
    rid = _evidence_report()
    copy = _upload_copy(project, _edited(_docx(project, rid)), filename="edited.docx")
    frames, cards, result = _dispatch(session, "import_edited_report", {"report_id": rid, "file_id": copy})
    rows = {r["field"]: r for r in cards[0]["edit_card"]["rows"] if r["role"] == "typed"}
    assert (rows["version"]["old"], rows["version"]["new"]) == (1, 2)
    section = next(r for f, r in rows.items() if f.startswith("sections."))
    assert section["old"] != section["new"] and section["new"].startswith("your edit · ")
    assert section["label"].startswith("Section · ")
    assert (rows["template_file_id"]["old"], rows["template_file_id"]["new"]) == (None, "edited.docx")
    assert rows["template_file_id"]["words"]["old"] == "Built-in default"
    # the merged result is large enough for the chat to spill it, so the receipt is read from the stores
    assert edit_card.receipt(cards[0]["edit_card"], {"report_id": rid, "file_id": copy})["outcome"] == "verified"
    assert store.load_report(project, rid).version == 2


def test_import_without_binding_and_a_copy_that_changes_nothing(project):
    rid = _evidence_report()
    data = _docx(project, rid)
    copy = _upload_copy(project, data, filename="same.docx")
    # An unedited copy only binds itself as the template; with binding off nothing changes.
    assert _refusal("import_edited_report", {"report_id": rid, "file_id": copy,
                                             "bind_as_template": False}).detail["error_kind"] == "empty_edit"
    card = edit_card.preview("import_edited_report", {"report_id": rid, "file_id": copy})
    assert [f for f, _, _ in _cells(card)] == ["version", "template_file_id"]
    edited = _upload_copy(project, _edited(data), filename="e.docx")
    card = edit_card.preview("import_edited_report", {"report_id": rid, "file_id": edited, "bind_as_template": False})
    assert "template_file_id" not in {f for f, _, _ in _cells(card)}


def test_import_refuses_what_the_route_refuses(project):
    rid = _evidence_report()
    junk = _upload_copy(project, b"not a word document", filename="junk.docx")
    assert _refusal("import_edited_report", {"report_id": rid, "file_id": junk}).detail["error_kind"] \
        == "roundtrip_unreadable"
    other = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Something else entirely")
    doc.save(other)
    stranger = _upload_copy(project, other.getvalue(), filename="other.docx")
    assert _refusal("import_edited_report", {"report_id": rid, "file_id": stranger}).detail["error_kind"] \
        == "roundtrip_not_a_report"
    assert _refusal("import_edited_report", {"report_id": rid, "file_id": "nope"}).status_code == 404
    assert _refusal("import_edited_report", {"report_id": "no-report", "file_id": junk}).status_code in (400, 404)


# ── shared conditions: nothing written, access first, a running job ─────────

def test_previews_write_nothing(project, fixtures, fake_untagged):
    rid, outline = _untagged_report(project, fixtures)
    other = _upload_template(project, "other.docx", fixtures["corporate_untagged_de.docx"])
    copy = _upload_copy(project, _edited(_docx(project, rid)))
    before = dir_hash(project)
    edit_card.preview("set_report_template", {"report_id": rid, "file_id": other})
    edit_card.preview("set_report_mapping", {"report_id": rid, "plan": _plan_for(outline)})
    edit_card.preview("import_edited_report", {"report_id": rid, "file_id": copy})
    assert dir_hash(project) == before


def test_an_outsider_gets_the_handlers_refusal_not_old_values(project, fixtures, second_identity):
    rid = _evidence_report()
    tid = _upload_template(project, "t.docx", fixtures["tagged_minimal.docx"])
    copy = _upload_copy(project, _edited(_docx(project, rid)))
    T.set_acting_user(second_identity["user_id"])  # after the owner's setup
    for tool, args in (("set_report_template", {"report_id": rid, "file_id": tid}),
                       ("set_report_mapping", {"report_id": rid, "plan": {}}),
                       ("import_edited_report", {"report_id": rid, "file_id": copy})):
        exc = _refusal(tool, args)
        assert exc.status_code in (403, 404), tool
        assert "rows" not in repr(exc.detail) and "old" not in exc.detail, tool


def test_a_running_report_job_refuses_like_the_route(project, fixtures, fake_untagged, monkeypatch):
    from routers import reports

    rid, outline = _untagged_report(project, fixtures)
    copy = _upload_copy(project, _edited(_docx(project, rid)))
    monkeypatch.setattr(reports, "_job_slot_busy", lambda: True)
    for tool, args in (("set_report_mapping", {"report_id": rid, "plan": _plan_for(outline)}),
                       ("import_edited_report", {"report_id": rid, "file_id": copy})):
        assert _refusal(tool, args).detail["error_kind"] == "report_job_in_flight", tool
        with pytest.raises(HTTPException) as real:  # the handler says the same
            getattr(T, tool)(**args)
        assert real.value.detail["error_kind"] == "report_job_in_flight", tool
