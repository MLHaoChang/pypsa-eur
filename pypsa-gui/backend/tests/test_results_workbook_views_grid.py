"""
The results workbook, package 02 slice R2 (`.scratch/results-workbook/issues/02-result-view-sheets.md`):
the Load Flow, Grid code & connection, Electrical design, Island & ride-through
and Grid-forming sheets. Each view's key figures equal its tab's route payload,
each scenario column comes from that run's own context or project, every
per-period view shows the second period, the project views resolve access
first, a view with nothing to show is "no results" and a failing one is an
error, and every chart reads number cells on its own sheet.
"""
from __future__ import annotations

import io
import json
from types import SimpleNamespace

import numpy as np
import pandas as pd
import pypsa
import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from services import results_workbook as RW
from services import run_evidence
from services.loadflow_summary import get_loadflow_view
from services.pypsa_service import PyPSAService
from services.results_workbook import scenarios as S
from services.results_workbook.views import HORIZON
from services.results_workbook.views_grid import NO_PERIODS
from services.solver_service import SolverConfig

PROJECT = "44444444-4444-4444-4444-444444444444"
SCEN = "55555555-5555-5555-5555-555555555555"
TITLES = {"loadflow": "Load Flow", "grid_code": "Grid code & connection",
          "electrical_design": "Electrical design", "island": "Island & ride-through",
          "grid_forming": "Grid-forming"}
P2 = "2040"


# ── fixtures ─────────────────────────────────────────────────────────────

def _lf_network(scale: float = 1.0) -> pypsa.Network:
    """Two periods of two hours; L1's peak is in 2040, T1's in 2030."""
    n = pypsa.Network()
    hours = pd.date_range("2030-01-01", periods=2, freq="h")
    n.set_snapshots(pd.MultiIndex.from_product([[2030, 2040], hours], names=["period", "timestep"]))
    n.investment_periods = [2030, 2040]
    n.snapshot_weightings.loc[:, :] = 3.0
    n.investment_period_weightings.loc[:, "years"] = 10.0
    for bus in ("A", "B", "C"):
        n.add("Bus", bus)
    n.add("Line", "L1", bus0="A", bus1="B", x=0.1, r=0.01, s_nom=100.0)
    n.add("Line", "L2", bus0="A", bus1="B", x=0.1, r=0.01, s_nom=100.0)
    n.add("Transformer", "T1", bus0="B", bus1="C", s_nom=50.0, x=0.1, r=0.01)
    idx = n.snapshots
    n.lines_t.p0 = pd.DataFrame({"L1": [50.0, 60.0, 80.0 * scale, 105.0 * scale],
                                 "L2": [10.0, 20.0, 30.0, 40.0]}, index=idx)
    n.lines_t.p1 = -0.99 * n.lines_t.p0
    n.transformers_t.p0 = pd.DataFrame({"T1": [45.0 * scale, 10.0, 20.0, 25.0]}, index=idx)
    n.transformers_t.p1 = -0.99 * n.transformers_t.p0
    n._objective = 0.0
    return n


def _solved_state(extra: dict | None = None) -> dict:
    return {"status": "completed", "condition": "optimal", "objective": 1000.0,
            "ac_pf_results": None, "ac_pf_convergence": None, "ac_pf_convergence_list": None,
            "ac_pf_converged_count": None, "ac_pf_total_snapshots": None, **(extra or {})}


@pytest.fixture
def active_lf(install_network, monkeypatch):
    """The base: the active network, solved, in a saved project."""
    from services import workflow_tools
    n = _lf_network()
    install_network(n)
    PyPSAService.get_solver_state().update(_solved_state())
    workflow_tools._record_solve_inputs(n, PyPSAService.get_solver_state()["solver_config"])
    monkeypatch.setattr(run_evidence, "_authority", lambda: (None, PROJECT, "owner"))
    return n


def _scenario_context(n: pypsa.Network):
    """A scenario's own saved context, as `workflow_tools._load_context` builds it."""
    from services import workflow_tools
    ctx = PyPSAService.build_context()
    ctx.network, ctx.project_uuid, ctx.loaded_project = n, SCEN, "High load"
    ctx.solver_state.update(_solved_state({"solver_config": SolverConfig()}))
    n.meta["assistant_solve_exact_input_fingerprint"] = workflow_tools._network_input_fingerprint(
        ctx, exact=True)
    return ctx


def _rdf(n, accessor, attr, source="lopf"):
    return getattr(getattr(n, accessor), attr, None)


def _run(n, name, **kw):
    return S.Run(name=name, network=n, cfg=SolverConfig(), result_df=_rdf,
                 solved_settings=lambda: None, **kw)


def _build(n, views, *, scenarios=None, project_for=None, project_id="base") -> bytes:
    return RW.build_results_workbook(RW.WorkbookInputs(
        network=n, cfg=SolverConfig(), lost_load=None, result_df=_rdf, views=list(views),
        project="Base", scenarios=scenarios, solved_settings=lambda: None,
        project_id=project_id, project_for=project_for))


def _wb(blob):
    return load_workbook(io.BytesIO(blob))


def _key(ws, label: str, period: str, col: int = 4):
    for row in ws.iter_rows(min_row=4):
        if row[0].value == label and str(row[2].value) == period:
            return row[col - 1].value
    raise AssertionError(f"no key {label!r} / {period!r} on {ws.title}")


def _about(wb) -> dict:
    from services.result_views import view
    by_label = {row[0]: row[2] for row in wb["About"].iter_rows(values_only=True) if row and row[0]}
    return {v: by_label.get(view(v).label) for v in TITLES}


# ── condition 1: the lifted load-flow read changes nothing ───────────────

def test_loadflow_route_json_is_unchanged_and_a_context_read_of_the_active_run_equals_it(
        client, active_lf, monkeypatch):
    plain = get_loadflow_view()
    assert plain["status"] == "ok"
    # Passing the active context explicitly reads exactly what the default reads.
    active = PyPSAService.get_active_context()
    monkeypatch.setattr(active, "project_uuid", PROJECT, raising=False)
    dump = lambda x: json.dumps(x, sort_keys=True)  # noqa: E731
    assert dump(get_loadflow_view(context=active)) == dump(plain)
    windows = {(a, b): get_loadflow_view("lopf", a, b) for a, b in ((0, 1), (2, 3))}
    for (a, b), payload in windows.items():
        assert dump(get_loadflow_view("lopf", a, b, context=active)) == dump(payload)
    # The route serves the same JSON (its body is unchanged: no context passed).
    assert dump(client.get("/api/results/loadflow-view").json()) == dump(plain)
    for (a, b), payload in windows.items():
        assert dump(client.get("/api/results/loadflow-view",
                               params={"window_from": a, "window_to": b}).json()) == dump(payload)


# ── condition 2 and 3: load flow ─────────────────────────────────────────

def test_loadflow_keys_equal_the_route_and_each_scenario_reads_its_own_context(active_lf):
    other = _lf_network(scale=1.5)
    scen = S.Scenarios(runs=[_run(other, "High load", project_id=SCEN,
                                  context=_scenario_context(other))])
    wb = _wb(_build(active_lf, ["loadflow"], scenarios=scen))
    ws = wb[TITLES["loadflow"]]
    # The route's own function (`/loadflow-view` calls it as is).
    get = lambda **p: get_loadflow_view("lopf", p.get("window_from"), p.get("window_to"))  # noqa: E731
    whole, p2 = get(), get(window_from=2, window_to=3)
    assert _key(ws, "Branches over their limit", HORIZON) == whole["counts"]["over_limit"]
    assert _key(ws, "Highest peak loading", HORIZON) == pytest.approx(whole["rows"][0]["util_pct"])
    # The second period is the tab's window over its snapshots.
    assert _key(ws, "Highest peak loading", P2) == pytest.approx(p2["rows"][0]["util_pct"])
    assert _key(ws, "Weighted hours in the window", P2) == pytest.approx(p2["window"]["weighted_hours"])
    assert _key(ws, "Highest peak loading", "2030") == pytest.approx(
        get(window_from=0, window_to=1)["rows"][0]["util_pct"])
    # The scenario column is that run's own context: a base read would fail.
    mine = get_loadflow_view(context=scen.runs[0].context)
    assert mine["rows"][0]["util_pct"] != pytest.approx(whole["rows"][0]["util_pct"])
    assert _key(ws, "Highest peak loading", HORIZON, col=5) == pytest.approx(mine["rows"][0]["util_pct"])
    assert _key(ws, "Branches over their limit", HORIZON, col=5) == mine["counts"]["over_limit"]
    # Long tables: every period, both runs; the line flows are the /lines series.
    rows = list(ws.iter_rows(values_only=True))
    loading = [r for r in rows if r[0] in ("Base", "High load") and r[1] == P2]
    assert {r[0] for r in loading} == {"Base", "High load"}
    l1 = float(active_lf.lines_t.p0["L1"].iloc[3])  # what /lines serves for 2040's last hour
    assert any(r[0] == "Base" and r[1] == 2040 and r[3] == l1 for r in rows)



def test_a_branch_list_the_tab_cuts_short_says_so(active_lf, monkeypatch):
    from services import loadflow_summary as lf

    def capped(*a, **k):
        out = get_loadflow_view(*a, **k)
        if out.get("status") == "ok":
            out = {**out, "rows": out["rows"][:1], "rows_truncated": True}
        return out
    monkeypatch.setattr(lf, "get_loadflow_view", capped)
    ws = _wb(_build(active_lf, ["loadflow"]))[TITLES["loadflow"]]
    notes = " ".join(str(r[0]) for r in ws.iter_rows(values_only=True) if r[0])
    ranked = get_loadflow_view("lopf")["counts"]["ranked"]
    assert f"ranks at most 1 of {ranked} branches" in notes

def test_a_loadflow_the_tab_cannot_show_is_no_results(install_network, monkeypatch):
    n = pypsa.Network()
    n.set_snapshots(pd.date_range("2030-01-01", periods=2, freq="h"))
    n.add("Bus", "A")
    install_network(n)
    monkeypatch.setattr(run_evidence, "_authority", lambda: (None, PROJECT, "owner"))
    assert get_loadflow_view()["status"] == "unavailable"
    wb = _wb(_build(n, ["loadflow"]))
    assert "Load Flow" not in wb.sheetnames and "no results" in str(_about(wb)["loadflow"])


# ── the project views: payloads per project ──────────────────────────────

def _hub_state(scale: float = 1.0, results: bool = True) -> dict:
    res = {
        "selection": [{"period": 2030, "hour": 5, "reasons": ["peak_import"]},
                      {"period": 2040, "hour": 17, "reasons": ["min_pf", "peak_import"]}],
        # Columns as a real campus run writes them (gridspine campus_study).
        "transformers": [{"group": "GRID_IMPORT", "units": 1, "unit_rating_mva": 80.0,
                          "max_s_intact_mva": 46.6 * scale, "required_unit_mva": 55.9 * scale,
                          "recommended_unit_mva": 63.0, "adequate": True,
                          "utilisation": 0.7 * scale, "worst_period": 2040, "worst_hour": 0}],
        "compensation": [{"direction": "capacitive", "required_mvar": 4.0 * scale,
                          "recommended_mvar": 5.0, "worst_period": 2030, "worst_hour": 5,
                          "margin": 0.2}],
        "short_circuit": [{"bus": "mv", "vn_kv": 20.0, "ikss_max_ka": 12.5 * scale,
                           "ip_max_ka": 30.0, "ikss_min_ka": 11.0, "rated_ka": None,
                           "adequate": None, "detail": "no switchgear rating given",
                           "period": 2040}],
        "compliance": [
            {"check": "pcc_reactive", "status_as_is": "fail", "status_with_measures": "pass",
             "value": 12.0 * scale, "unit": "Mvar", "limit": 10.0, "worst_period": 2040,
             "worst_hour": 17, "detail": "over the band", "clause": "RfG 21", "source": "code"},
            {"check": "transformer_loading", "status_as_is": "pass", "status_with_measures": "pass",
             "value": 81.0 * scale, "unit": "%", "limit": 100.0, "worst_period": 2030,
             "worst_hour": 5, "detail": "", "clause": None, "source": None}],
        "reactive": [{"period": 2030, "hour": 5, "converged": True, "q0_mvar": 9.0,
                      "q_final_mvar": 8.0 * scale, "compliant_without": True},
                     {"period": 2040, "hour": 17, "converged": True, "q0_mvar": 14.0,
                      "q_final_mvar": 12.0 * scale, "compliant_without": False}],
        "ibr_q_curve": None, "ibr_q_curve_not_computed": None,
        "requirement": {"q_limit_mvar": 10.0, "p_ref_mw": 60.0 * scale, "clause": "RfG 21",
                        "p_ref_from": "the connection agreement"},
        "reported_requirements": [{"id": "frt", "requirement": "fault ride-through",
                                   "status": "needs an RMS study"}],
        "investment": [{"need": "pcc_reactive", "action": "buy", "status": "chosen", "invest_period": 2030,
                        "capex_eur": 100_000.0 * scale, "annualised_eur_per_a": 9_000.0 * scale},
                       {"need": "transformer", "action": "add", "status": "chosen", "invest_period": 2040,
                        "capex_eur": 250_000.0 * scale, "annualised_eur_per_a": 21_000.0 * scale}],
        "cost": [{"period": 2030, "capex_eur": 100_000.0 * scale, "capex_replacement_eur": 0.0,
                  "annualised_eur_per_a": 9_000.0 * scale},
                 {"period": 2040, "capex_eur": 250_000.0 * scale, "capex_replacement_eur": 0.0,
                  "annualised_eur_per_a": 21_000.0 * scale}],
        "compliance_invested": [{"check": "pcc_reactive", "status_as_is": "fail",
                                 "status_with_measures": "pass", "value_with_measures": 9.0}],
        "history": [], "unresolved": [], "taps": None, "tap_hours_moved": None, "scope": None,
        "cost_basis": None, "harmonic_scan": None, "voltage_stability": None,
    }
    return {"results": res if results else None, "stale": False, "grid_code_changed": False}


def _island_report(scale: float = 1.0) -> dict:
    return {"caveat": "Screening, not a certificate.", "missing": [],
            "tables": [
                {"id": "requirements", "group": "inputs", "title": "Island requirements",
                 "columns": ["field", "value"], "rows": [["rocof_hz_s", 1.0]], "assumed_rows": []},
                {"id": "costs", "group": "study", "title": "Cost of island capability",
                 "columns": ["scenario", "status", "objective", "cost of island capability"],
                 "rows": [["baseline (no island rows)", "ok", 1.0e6, 0.0],
                          ["seamless", "ok", 1.0e6 + 5.0e4 * scale, 5.0e4 * scale]],
                 "assumed_rows": [1]},
                {"id": "island_hours", "group": "hours", "title": "Island critical hours",
                 "columns": ["period", "hour", "island reasons"],
                 "rows": [[2030, 5, "peak_import"], [2040, 17, "min_pf"]], "assumed_rows": []}]}


def _gfm(scale: float = 1.0, state: str = "ready") -> dict:
    if state == "not_run":
        return {"state": "not_run", "reason": "Run the campus study first", "recommendation": None}
    return {"state": state, "reason": None, "recommendation": {
        "label": "heuristic screening", "status": "not_achievable", "simulations": 32,
        "settings": {"bess_new_2030": {"model": "gfm_droop",
                                       "current": {"droop_pct": 4.0, "tau_f_s": 0.1},
                                       "recommended": {"droop_pct": 2.0 * scale, "tau_f_s": 0.5}}},
        "cases": [{"key": "2030 h5 seamless E1", "ok_before": False, "binding_before": "rocof",
                   "ok_after": scale > 1, "binding_after": "rocof"},
                  {"key": "2040 h17 seamless E1", "ok_before": True, "binding_before": None,
                   "ok_after": True, "binding_after": None}],
        "capacity": [{"case": "2030 h5 seamless E1", "kind": "inertia", "delta": 40.0 * scale,
                      "unit": "MW·s", "binding": "rocof", "derivation": "f0·ΔP/RoCoF",
                      "feasible": True, "note": ""}]}}


@pytest.fixture
def projects(monkeypatch):
    """Two hub projects (base, scenario) and a planning study; the services
    their tabs' routes call answer per project, and record what they read."""
    from services import campus_electrical_service as ce
    from services import campus_island_report, campus_island_tuning
    from services import gridspine_service as gs
    rows = {"base": SimpleNamespace(name="base", project_kind="capacity_expansion"),
            SCEN: SimpleNamespace(name="scen", project_kind="capacity_expansion"),
            "study": SimpleNamespace(name="study", project_kind="planning_dynamics")}
    scale = {"base": 1.0, "scen": 2.0}
    read: list[tuple[str, str]] = []

    def spy(fn_name, make):
        def fn(project, *a, **k):
            read.append((fn_name, project.name))
            return make(project)
        return fn
    monkeypatch.setattr(ce, "get_state", spy("get_state", lambda p: _hub_state(scale[p.name])))
    monkeypatch.setattr(ce, "get_island", spy("get_island", lambda p: {
        "config": {}, "study": {}, "study_at": "2030-01-01T00:00:00+00:00"}))
    monkeypatch.setattr(campus_island_report, "island_report",
                        spy("island_report", lambda p: _island_report(scale[p.name])))
    monkeypatch.setattr(campus_island_tuning, "get_settings",
                        spy("get_settings", lambda p: _gfm(scale[p.name])))
    monkeypatch.setattr(gs, "get_capacity", spy("get_capacity", lambda p: {"rows": [
        {"bus": "b1", "hour": 3, "kind": "load", "capacity_mw": None, "dc_estimate_mw": 120.0,
         "binding_kind": "thermal_n1", "binding_element": "l1", "binding_contingency": "l2",
         "binding_preexisting": False, "method": "dc"},
        {"bus": "b2", "hour": 3, "kind": "load", "capacity_mw": 80.0, "dc_estimate_mw": 95.0,
         "binding_kind": "v_low_n1", "binding_element": "b2", "binding_contingency": "l1",
         "binding_preexisting": False, "method": "ac"}], "hours": [3], "buses": ["b1", "b2"]}))
    monkeypatch.setattr(gs, "get_connection", spy("get_connection", lambda p: {"rows": [
        {"assessment_id": "a1", "hour": 3, "bus": "b1", "check": "voltage", "status": "fail",
         "value": 0.93, "unit": "pu", "limit": 0.95, "detail": "", "clause": "", "source": "code"}]}))
    calls: list = []

    def project_for(run):
        calls.append(run.project_id)
        return rows.get(run.project_id)
    return SimpleNamespace(rows=rows, read=read, project_for=project_for, calls=calls,
                           scale=scale)


@pytest.fixture(scope="module")
def net():
    return _lf_network()


def _two_runs(net, *, scen_id=SCEN):
    return S.Scenarios(runs=[_run(_lf_network(), "High load", project_id=scen_id)])


@pytest.mark.parametrize("view", ["grid_code", "electrical_design", "island", "grid_forming"])
def test_project_views_read_each_runs_own_project_after_resolving_it(net, projects, view):
    wb = _wb(_build(net, [view], scenarios=_two_runs(net), project_for=projects.project_for))
    ws = wb[TITLES[view]]
    assert projects.calls == ["base", SCEN]
    # Every file read is of a project resolved first, and each run reads its own.
    assert {name for _, name in projects.read} == {"base", "scen"}
    first = {"grid_code": ("Checks fail as is", 1, 1), "electrical_design": None,
             "island": ("Cost of island capability — seamless", 5.0e4, 1.0e5),
             "grid_forming": ("Capacity needed — inertia", 40.0, 80.0)}[view]
    if first:
        label, base, scen = first
        assert _key(ws, label, HORIZON) == pytest.approx(base)
        assert _key(ws, label, HORIZON, col=5) == pytest.approx(scen)


def test_grid_code_keys_equal_the_campus_payload(net, projects):
    ws = _wb(_build(net, ["grid_code"], scenarios=_two_runs(net),
                    project_for=projects.project_for))[TITLES["grid_code"]]
    res = _hub_state(1.0)["results"]
    assert _key(ws, "P_ref", HORIZON) == res["requirement"]["p_ref_mw"]
    assert _key(ws, "P_ref", HORIZON, col=5) == _hub_state(2.0)["results"]["requirement"]["p_ref_mw"]
    assert _key(ws, "Checks fail with measures", HORIZON) == 0
    assert NO_PERIODS in ws["A2"].value  # no per-period split: said, not faked


def test_grid_code_of_a_planning_study_reads_capacity_and_connection(net, projects):
    ws = _wb(_build(net, ["grid_code"], project_id="study",
                    project_for=projects.project_for))[TITLES["grid_code"]]
    assert _key(ws, "Lowest DC capacity estimate — load", HORIZON) == 95.0
    assert _key(ws, "Lowest AC connection capacity — load", HORIZON) == 80.0
    assert _key(ws, "Connection checks — fail", HORIZON) == 1
    assert {f for f, _ in projects.read} == {"get_capacity", "get_connection"}


def test_electrical_design_costs_per_period_with_the_second_period(net, projects):
    ws = _wb(_build(net, ["electrical_design"], scenarios=_two_runs(net),
                    project_for=projects.project_for))[TITLES["electrical_design"]]
    cost = _hub_state(1.0)["results"]["cost"]
    assert _key(ws, "Electrical capex", P2) == cost[1]["capex_eur"]
    assert _key(ws, "Electrical capex", "2030") == cost[0]["capex_eur"]
    assert _key(ws, "Electrical capex", HORIZON) == cost[0]["capex_eur"] + cost[1]["capex_eur"]
    assert _key(ws, "Electrical cost, annualised", P2, col=5) == _hub_state(2.0)["results"]["cost"][1][
        "annualised_eur_per_a"]
    tables = [r[1] for r in ws.iter_rows(values_only=True) if r[0] == "Base"]
    assert 2040 in tables  # the long tables carry the second period's rows



def test_electrical_design_sums_only_whole_capex_and_never_annualised_cost(net, projects,
                                                                          monkeypatch):
    from services import campus_electrical_service as ce
    ws = _wb(_build(net, ["electrical_design"], project_for=projects.project_for))[
        TITLES["electrical_design"]]
    assert _key(ws, "Electrical cost, annualised", HORIZON) == "not_established"
    first_only = _hub_state(1.0)
    first_only["results"]["cost"] = first_only["results"]["cost"][:1]
    monkeypatch.setattr(ce, "get_state", lambda p: first_only)
    ws = _wb(_build(net, ["electrical_design"], project_for=projects.project_for))[
        TITLES["electrical_design"]]
    assert _key(ws, "Electrical capex", "2030") == first_only["results"]["cost"][0]["capex_eur"]
    assert _key(ws, "Electrical capex", HORIZON) == "not_established"  # 2040 unknown


def test_grid_forming_counts_cases_no_capacity_fixes(net, projects, monkeypatch):
    from services import campus_island_tuning
    got = _gfm()
    cap = got["recommendation"]["capacity"]
    cap.append({**cap[0], "case": "2040 h2 seamless E1", "delta": float("inf")})
    monkeypatch.setattr(campus_island_tuning, "get_settings", lambda p: got)
    wb = _wb(_build(net, ["grid_forming"], project_for=projects.project_for))
    ws = wb[TITLES["grid_forming"]]
    assert _key(ws, "Capacity needed — inertia", HORIZON) == 40.0
    assert _key(ws, "Cases no capacity fixes — inertia", HORIZON) == 1
    rows = list(ws.iter_rows(values_only=True))
    assert any("not achievable by capacity" in r for r in rows)
    assert ws._charts  # the capacity chart keeps its all-number column


def _gfm_capacity(rows):
    got = _gfm()
    got["recommendation"]["capacity"] = [
        {"case": c, "kind": k, "delta": v, "unit": u, "binding": "rocof", "derivation": "",
         "feasible": True, "note": ""} for c, k, v, u in rows]
    return got


def test_grid_forming_capacity_is_the_largest_delta_of_each_kind(net, projects, monkeypatch):
    from services import campus_island_tuning
    monkeypatch.setattr(campus_island_tuning, "get_settings", lambda p: _gfm_capacity([
        ("2030 h5", "inertia", 40.0, "MW·s"), ("2040 h2", "inertia", 25.0, "MW·s"),
        ("2030 h5", "headroom_up", 3.0, "MW")]))
    ws = _wb(_build(net, ["grid_forming"], project_for=projects.project_for))[
        TITLES["grid_forming"]]
    assert _key(ws, "Capacity needed — inertia", HORIZON) == 40.0
    assert _key(ws, "Capacity needed — headroom_up", HORIZON) == 3.0


def test_grid_forming_charts_capacity_only_for_a_single_kind(net, projects, monkeypatch):
    from services import campus_island_tuning
    titles = lambda ws: [c.title.tx.rich.p[0].r[0].t for c in ws._charts]  # noqa: E731
    monkeypatch.setattr(campus_island_tuning, "get_settings", lambda p: _gfm_capacity([
        ("2030 h5", "inertia", 40.0, "MW·s"), ("2030 h5", "headroom_up", 3.0, "MW")]))
    ws = _wb(_build(net, ["grid_forming"], project_for=projects.project_for))[
        TITLES["grid_forming"]]
    assert "Capacity needed by case" not in titles(ws)  # MW·s and MW share no axis
    assert any("not charted" in str(r[0]) for r in ws.iter_rows(values_only=True))
    monkeypatch.setattr(campus_island_tuning, "get_settings", lambda p: _gfm_capacity([
        ("2030 h5", "inertia", 40.0, "MW·s"), ("2040 h2", "inertia", 25.0, "MW·s")]))
    ws = _wb(_build(net, ["grid_forming"], project_for=projects.project_for))[
        TITLES["grid_forming"]]
    assert "Capacity needed by case" in titles(ws)


@pytest.mark.parametrize("view", ["grid_code", "island", "grid_forming"])
def test_a_view_without_a_period_split_keeps_the_second_periods_rows(net, projects, view):
    """GK B condition 3, ruling A: no per-period keys, but no first-period-only table."""
    ws = _wb(_build(net, [view], project_for=projects.project_for))[TITLES[view]]
    assert NO_PERIODS in ws["A2"].value
    base = [r for r in ws.iter_rows(values_only=True) if r[0] == "Base"]
    assert any(v == 2040 or str(v).startswith("2040 ") for r in base for v in r[1:])

def test_electrical_design_of_a_planning_study_is_no_results(net, projects):
    wb = _wb(_build(net, ["electrical_design"], project_id="study",
                    project_for=projects.project_for))
    assert "Electrical design" not in wb.sheetnames
    assert "no results" in str(_about(wb)["electrical_design"])
    assert projects.read == []  # the campus study is not read for a study project


def test_a_run_the_user_cannot_open_reads_nothing_of_its_project(net, projects):
    wb = _wb(_build(net, ["grid_forming"], scenarios=_two_runs(net, scen_id="hidden"),
                    project_for=projects.project_for))
    ws = wb[TITLES["grid_forming"]]
    assert projects.calls == ["base", "hidden"]
    assert [name for _, name in projects.read] == ["base"]
    assert _key(ws, "Capacity needed — inertia", HORIZON, col=5) == "not_established"


def test_without_project_access_the_project_views_have_no_data(net, projects):
    wb = _wb(_build(net, ["island"], project_for=None))
    assert "no results" in str(_about(wb)["island"]) and projects.read == []


# ── condition 5: no results vs error ─────────────────────────────────────

def test_grid_forming_not_run_is_no_results(net, projects, monkeypatch):
    from services import campus_island_tuning
    monkeypatch.setattr(campus_island_tuning, "get_settings", lambda p: _gfm(state="not_run"))
    wb = _wb(_build(net, ["grid_forming"], project_for=projects.project_for))
    assert "Grid-forming" not in wb.sheetnames
    assert "no results" in str(_about(wb)["grid_forming"])


def test_a_campus_study_not_run_is_no_results(net, projects, monkeypatch):
    from services import campus_electrical_service as ce
    monkeypatch.setattr(ce, "get_state", lambda p: _hub_state(results=False))
    wb = _wb(_build(net, ["grid_code", "electrical_design"], project_for=projects.project_for))
    about = _about(wb)
    assert "no results" in str(about["grid_code"]) and "no results" in str(about["electrical_design"])


def test_a_failing_project_view_is_an_error_in_about_not_no_results(net, projects, monkeypatch):
    from services import campus_electrical_service as ce
    monkeypatch.setattr(ce, "get_state", lambda p: (_ for _ in ()).throw(RuntimeError("boom")))
    wb = _wb(_build(net, ["electrical_design", "grid_forming"], project_for=projects.project_for))
    about = _about(wb)
    assert "error" in str(about["electrical_design"]) and "no results" not in str(about["electrical_design"])
    assert "Electrical design" not in wb.sheetnames and "Grid-forming" in wb.sheetnames


def test_a_planning_study_before_its_screen_is_no_results_but_other_errors_raise(net, projects,
                                                                                 monkeypatch):
    from services import gridspine_service as gs
    monkeypatch.setattr(gs, "get_connection", lambda p, **k: {"rows": []})
    monkeypatch.setattr(gs, "get_capacity", lambda p, **k: (_ for _ in ()).throw(
        HTTPException(404, "no screened run")))
    wb = _wb(_build(net, ["grid_code"], project_id="study", project_for=projects.project_for))
    assert "no results" in str(_about(wb)["grid_code"])
    monkeypatch.setattr(gs, "get_capacity", lambda p, **k: (_ for _ in ()).throw(
        HTTPException(500, "broken")))
    wb = _wb(_build(net, ["grid_code"], project_id="study", project_for=projects.project_for))
    assert "error" in str(_about(wb)["grid_code"])


# ── condition 8: the route resolves each run's project for the user ──────

def test_the_routes_project_reader_resolves_access_first(monkeypatch):
    import contextlib

    from routers import results as results_router
    from services import chat_tools, project_registry
    seen = []

    @contextlib.contextmanager
    def acting():
        yield "db", "user"

    def resolve(db, user, pid):
        seen.append((db, user, pid))
        if pid == "hidden":
            raise HTTPException(404, "Project not found")
        return SimpleNamespace(id=pid)
    monkeypatch.setattr(chat_tools, "_acting", acting)
    monkeypatch.setattr(project_registry, "resolve_project", resolve)
    read = results_router._workbook_project
    assert read(SimpleNamespace(project_id="p1")).id == "p1"
    assert read(SimpleNamespace(project_id="hidden")) is None
    assert read(SimpleNamespace(project_id=None)) is None
    assert seen == [("db", "user", "p1"), ("db", "user", "hidden")]


# ── condition 9: charts ──────────────────────────────────────────────────

def _cells(ref: str) -> tuple[str, list[str]]:
    from openpyxl.utils import range_boundaries
    from openpyxl.utils.cell import get_column_letter as L
    sheet, rng = ref.rsplit("!", 1)
    c0, r0, c1, r1 = range_boundaries(rng.replace("$", ""))
    return sheet.strip("'"), [f"{L(c)}{r}" for c in range(c0, c1 + 1) for r in range(r0, r1 + 1)]


@pytest.mark.parametrize("view,project_id", [
    ("loadflow", "base"), ("grid_code", "base"), ("grid_code", "study"),
    ("electrical_design", "base"), ("island", "base"), ("grid_forming", "base")])
def test_every_chart_reads_number_cells_on_its_own_sheet(active_lf, projects, view, project_id):
    blob = _build(active_lf, [view], project_id=project_id, project_for=projects.project_for)
    ws = _wb(blob)[TITLES[view]]
    assert 1 <= len(ws._charts) <= 3, view
    for chart in ws._charts:
        for s in chart.series:
            for ref in filter(None, [s.val.numRef.f if s.val else None]):
                sheet, cells = _cells(ref)
                assert sheet == ws.title, ref
                for c in cells:
                    v = ws[c].value
                    assert isinstance(v, (int, float)) and not isinstance(v, bool), (view, c, v)
                    assert np.isfinite(v)
            if s.cat is not None:
                assert _cells((s.cat.numRef or s.cat.strRef).f)[0] == ws.title
