"""
The XLSX pro forma (plan S5; review v1 S9 "Excel NPV semantics").

The workbook's cash-flow sheet carries live formulas — discount factor,
discounted and cumulative cash flow, and the NPV pinned as
``=CF0 + NPV(rate, CF1:CFn)`` — so a reader can audit and change it. Excel's
``NPV`` discounts its FIRST value one period; a range that starts at CF0
understates the NPV by a factor (1 + r). `tests/xlsx_formula_eval.py`
implements that semantics (checked here against Excel's documented examples)
and evaluates the workbook's cells against the case.
"""
from __future__ import annotations

import io

import openpyxl
import pytest

from tests.golden import site_fixture as sf
from tests.xlsx_formula_eval import FormulaEvaluator, excel_npv


# ── the evaluator against Excel's documentation ──────────────────────────

def test_the_evaluator_matches_excels_documented_npv_examples():
    """
    Microsoft's NPV function page, both examples:
    * ``=NPV(10%, -10000, 3000, 4200, 6800)`` is 1,188.44 (the -10,000 is
      paid at the END of the first period, so it is discounted too);
    * ``=NPV(8%, B3:B7) + B2`` with an initial -40,000 and 8,000, 9,200,
      10,000, 12,000, 14,500 is 1,922.06, and with a -9,000 loss in year six
      ``=NPV(8%, B3:B7, -9000) + B2`` is -3,749.47.
    """
    assert excel_npv(0.10, [-10000, 3000, 4200, 6800]) == pytest.approx(1188.44, abs=0.005)
    flows = [8000, 9200, 10000, 12000, 14500]
    assert excel_npv(0.08, flows) - 40000 == pytest.approx(1922.06, abs=0.005)
    assert excel_npv(0.08, flows + [-9000]) - 40000 == pytest.approx(-3749.47, abs=0.005)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    for i, v in enumerate([-40000] + flows, start=2):
        ws[f"B{i}"] = v
    ws["A1"] = 0.08
    ws["C1"] = "=NPV(A1,B3:B7)+B2"
    ws["C2"] = "=NPV(A1,B3:B7,-9000)+B2"
    ws["C3"] = "=NPV(10%,-10000,3000,4200,6800)".replace("10%", "0.1")
    ev = FormulaEvaluator(wb)
    assert ev.value("S", "C1") == pytest.approx(1922.06, abs=0.005)
    assert ev.value("S", "C2") == pytest.approx(-3749.47, abs=0.005)
    assert ev.value("S", "C3") == pytest.approx(1188.44, abs=0.005)


def test_the_evaluator_refuses_what_it_does_not_implement():
    wb = openpyxl.Workbook()
    wb.active.title = "S"
    wb["S"]["A1"] = "=IRR(B1:B3)"
    with pytest.raises(ValueError):
        FormulaEvaluator(wb).value("S", "A1")


# ── the workbook ─────────────────────────────────────────────────────────

def _case_and_ledger(option="bess_2h"):
    from services.study import proforma as P

    n, cfg = sf.solve_site_option(option)
    nb, _ = sf.solve_site_option("none")
    ledger = sf.site_ledger()
    case = P.build_investment_case(n, cfg, nb, ledger, None, option,
                                   study_id=sf.SITE_STUDY_ID, tariff=sf.site_tariff(ledger),
                                   fidelity="full_study")
    return case, ledger


@pytest.fixture(scope="module")
def book():
    from services.study.proforma_xlsx import write_proforma_xlsx

    case, ledger = _case_and_ledger()
    blob = write_proforma_xlsx(case, ledger)
    assert isinstance(blob, bytes) and blob[:2] == b"PK"
    return case, ledger, openpyxl.load_workbook(io.BytesIO(blob))


def _find(ws, label):
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == label:
                return cell
    raise AssertionError(f"{label!r} not on {ws.title}")


def test_the_sheets(book):
    _case, _ledger, wb = book
    assert wb.sheetnames == ["Cash flows", "KPIs", "Assumptions", "Provenance"]


def test_the_npv_cell_is_cf0_plus_npv_of_cf1_to_cfn_and_evaluates_to_the_case(book):
    case, _ledger, wb = book
    ws = wb["Cash flows"]
    npv_label = _find(ws, "NPV")
    formula = ws.cell(npv_label.row, npv_label.column + 1).value
    first = _find(ws, "Year").row + 1          # the year-0 row
    last = first + case.horizon_years
    net_col = _find(ws, "Net cash flow").column_letter
    ev = FormulaEvaluator(wb)
    got = ev.value("Cash flows", ws.cell(npv_label.row, npv_label.column + 1).coordinate)
    assert got == pytest.approx(case.kpis.npv, abs=1.0)
    assert formula == f"={net_col}{first}+NPV($B$1,{net_col}{first + 1}:{net_col}{last})"
    # …and the KPI sheet's NPV is that cell, not a copy of the number.
    kpi = _find(wb["KPIs"], "npv")
    assert ev.value("KPIs", f"B{kpi.row}") == pytest.approx(case.kpis.npv, abs=1.0)


def test_every_year_row_evaluates_to_the_case(book):
    case, _ledger, wb = book
    ws = wb["Cash flows"]
    ev = FormulaEvaluator(wb)
    first = _find(ws, "Year").row + 1
    cols = {name: _find(ws, name).column_letter for name in (
        "Net cash flow", "Discount factor", "Discounted cash flow", "Cumulative discounted",
        "Market revenue at duals (excluded)")}
    for y in case.years:
        r = first + y.year
        assert ws[f"{cols['Discount factor']}{r}"].value.startswith("=")
        assert ws[f"{cols['Discounted cash flow']}{r}"].value.startswith("=")
        assert ev.value("Cash flows", f"{cols['Net cash flow']}{r}") == pytest.approx(
            y.net_cash_flow, abs=1e-6)
        assert ev.value("Cash flows", f"{cols['Discount factor']}{r}") == pytest.approx(
            (1 + case.discount_rate) ** -y.year, rel=1e-12)
        assert ev.value("Cash flows", f"{cols['Discounted cash flow']}{r}") == pytest.approx(
            y.discounted_cash_flow, abs=1e-6)
        assert ev.value("Cash flows", f"{cols['Cumulative discounted']}{r}") == pytest.approx(
            y.cumulative_discounted, abs=1e-6)
    # The excluded column is not referenced by the net cash flow formula.
    mr = cols["Market revenue at duals (excluded)"]
    assert all(mr not in str(ws[f"{cols['Net cash flow']}{first + y.year}"].value)
               for y in case.years)


def test_assumptions_are_the_ledger_and_provenance_is_complete(book):
    from services.study import packs

    case, ledger, wb = book
    ws = wb["Assumptions"]
    keys = [c.value for c in ws["A"][1:] if c.value]
    assert keys == [r.key for r in ledger.rows]
    prov = {r[0].value: r[1].value for r in wb["Provenance"].iter_rows() if r[0].value}
    assert prov["basis"] == "real, pre-tax, without subsidy"
    assert prov["currency_year"] == 2020
    assert prov["fidelity"] == "full_study"
    assert prov["library_version"] == ledger.ledger_version
    assert prov["ledger_hash"] == packs.ledger_hash(ledger)
    for engine in ("cash_flow_expander", "bill_calculator", "lp_duals", "ledger"):
        assert engine in prov["engines"]
    assert prov["salvage_basis"] == "annuity_pv"


def test_a_ledger_text_that_looks_like_a_formula_stays_text():
    from services.study.proforma_xlsx import write_proforma_xlsx

    case, ledger = _case_and_ledger()
    rows = [r.model_copy(update={"source": '=HYPERLINK("http://x","y")'})
            if r.key == "discount_rate" else r for r in ledger.rows]
    wb = openpyxl.load_workbook(io.BytesIO(write_proforma_xlsx(
        case, ledger.model_copy(update={"rows": rows}))))
    cell = _find(wb["Assumptions"], '=HYPERLINK("http://x","y")')
    assert cell.data_type == "s"


def test_a_not_established_case_writes_no_npv():
    from services.study import tariff as T
    from services.study import proforma as P
    from services.study.proforma_xlsx import write_proforma_xlsx

    n, cfg = sf.solve_site_option("bess_2h")
    nb, _ = sf.solve_site_option("none")
    ledger = sf.site_ledger()
    tariff = sf.site_tariff(ledger)
    base = T.BillCalculator().bill(nb.links_t.p0["grid_import"], nb.links_t.p0["grid_export"],
                                   tariff, nb.snapshot_weightings)
    null = T.BillCalculator().bill(None, None, tariff, n.snapshot_weightings)
    case = P.build_investment_case(n, cfg, nb, ledger, {"baseline": base, "option": null},
                                   "bess_2h", study_id=sf.SITE_STUDY_ID, tariff=tariff)
    wb = openpyxl.load_workbook(io.BytesIO(write_proforma_xlsx(case, ledger)))
    kpi = _find(wb["KPIs"], "status")
    assert wb["KPIs"].cell(kpi.row, 2).value == "not_established"
    with pytest.raises(AssertionError):
        _find(wb["Cash flows"], "NPV")
