"""
U2 WP8 part B — plan §5.4 (the vocabulary rename) with read-compat, the C9
row of the §2 call-site table.

The case's KPIs are renamed: `CaseKpis.lcos` → `levelised_cost` (with its
`levelised_cost_basis`) and `salvage_eur` → `terminal_value_eur`. A study
stored before U2 holds the old names, and `_Model` forbids unknown keys, so
`CaseKpis` maps them on read (`_legacy_names`); nothing writes them again.

The engine literals: the Investment Case engine's figures are
`tariff_engine` / `finance_engine` (WP7), and every default now names them.
GS's own two engines keep their labels where they made the figure —
`cash_flow_expander` (the pro forma) and `bill_calculator` (GS's bill
calculator): the findings' fallback until WP10, and every figure a study
stored before U2 holds. Read-compat never relabels those as the engine's
(plan WP8 part B, decision B1): a pre-U2 case IS the pro forma's.

The stored pre-U2 payloads are `tests/fixtures/pre_u2_study/*.json`: the S8
route dumps (`b6566fa6`, the frontend's fixtures before this rename), kept
verbatim. Mutation: remove `CaseKpis._legacy_names` → the stored-case tests
and the stored report route go red.
"""
from __future__ import annotations

import copy
import io
import json
import pathlib
import re

import pytest
from pydantic import ValidationError

from models.study import (
    BatteryAttribution,
    Bill,
    CaseKpis,
    DecisionReport,
    Findings,
    InvestmentCase,
)

PRE = pathlib.Path(__file__).resolve().parent / "fixtures" / "pre_u2_study"
OLD_KPIS = {"lcos", "salvage_eur"}
NEW_KPIS = {"levelised_cost", "levelised_cost_basis", "terminal_value_eur"}


def _pre(name: str) -> dict:
    raw = json.loads((PRE / name).read_text(encoding="utf-8"))
    raw.pop("available", None)        # the route's payload flag, not the model's
    return raw


def test_the_fixture_is_a_pre_u2_payload():
    """The guard on the oracle: the stored case really carries the old names."""
    raw = _pre("case.json")
    assert OLD_KPIS <= set(raw["kpis"]) and not NEW_KPIS & set(raw["kpis"])
    assert raw["engine"] == "cash_flow_expander"
    assert raw["sources"]["bill_refs"] == ["bill_calculator:none", "bill_calculator:bess_1h"]


def test_a_stored_pre_u2_case_loads_under_the_new_kpi_names():
    raw = _pre("case.json")
    case = InvestmentCase.model_validate(raw)
    k = case.kpis
    assert k.levelised_cost == raw["kpis"]["lcos"]
    # Both producers before U2 (the pro forma, and the WP7 view) were real-terms.
    assert k.levelised_cost_basis == "real"
    assert k.terminal_value_eur == raw["kpis"]["salvage_eur"]
    assert k.terminal_value_eur == pytest.approx(case.years[-1].salvage, rel=1e-12)
    # Honest labels (decision B1): the pro forma and GS's calculator made it.
    assert case.engine == "cash_flow_expander"
    assert case.sources.bill_refs == raw["sources"]["bill_refs"]
    assert case.provenance.engines == raw["provenance"]["engines"]
    assert {s.engine for s in case.value_streams} == {"bill_calculator"}


def test_a_stored_flag_moves_with_its_figure():
    """A null stored KPI keeps its reason under the new name (ADR-0001)."""
    raw = _pre("case.json")
    raw["kpis"] = {**raw["kpis"], "lcos": None, "salvage_eur": None,
                   "unavailable": {**raw["kpis"]["unavailable"], "lcos": "no_discharge",
                                   "salvage_eur": "salvage_not_computed"}}
    last = raw["years"][-1]
    raw["years"][-1] = {**last, "salvage": None,
                        "unavailable": {**last["unavailable"], "salvage": "salvage_not_computed"}}
    raw["salvage_basis"] = None
    k = InvestmentCase.model_validate(raw).kpis
    assert k.levelised_cost is None and k.levelised_cost_basis is None
    assert k.unavailable["levelised_cost"] == "no_discharge"
    assert k.terminal_value_eur is None
    assert k.unavailable["terminal_value_eur"] == "salvage_not_computed"
    assert not OLD_KPIS & set(k.unavailable)


def test_it_is_written_back_in_the_new_names_only():
    case = InvestmentCase.model_validate(_pre("case.json"))
    out = case.model_dump(mode="json")
    assert NEW_KPIS <= set(out["kpis"]) and not OLD_KPIS & set(out["kpis"])
    assert not OLD_KPIS & set(out["kpis"]["unavailable"])
    assert InvestmentCase.model_validate(out) == case


def test_an_old_and_a_new_name_together_are_refused():
    """A payload naming both is ambiguous: never silently one of them."""
    raw = _pre("case.json")["kpis"]
    with pytest.raises(ValidationError, match="both lcos and levelised_cost"):
        CaseKpis.model_validate({**raw, "levelised_cost": 1.0, "levelised_cost_basis": "real"})
    with pytest.raises(ValidationError, match="both salvage_eur and terminal_value_eur"):
        CaseKpis.model_validate({**raw, "terminal_value_eur": 1.0})


def test_a_levelised_cost_states_its_basis():
    k = CaseKpis.model_validate(_pre("case.json")["kpis"]).model_dump(mode="json")
    with pytest.raises(ValidationError, match="needs its levelised_cost_basis"):
        CaseKpis.model_validate({**k, "levelised_cost_basis": None})
    with pytest.raises(ValidationError, match="null levelised_cost has no levelised_cost_basis"):
        CaseKpis.model_validate({**k, "levelised_cost": None, "levelised_cost_basis": "real",
                                 "unavailable": {**k["unavailable"],
                                                 "levelised_cost": "no_discharge"}})


@pytest.mark.parametrize("name", ["findings.json", "findings_pv.json"])
def test_stored_pre_u2_findings_load_with_their_own_labels(name):
    raw = _pre(name)
    f = Findings.model_validate(raw)
    assert {a.engine for a in f.battery_attribution} == {"cash_flow_expander"}
    assert {s.engine for s in f.value_streams} <= {"bill_calculator"}
    assert f.model_dump(mode="json", by_alias=True)["verdict"] == raw["verdict"]


def test_a_stored_pre_u2_report_and_bill_load():
    report = DecisionReport.model_validate(_pre("report.json"))
    assert "cash_flow_expander" in {fig.engine for fig in report.facts.values()}
    bill = Bill.model_validate(_pre("preview_upload.json")["bill"]["bill"])
    assert bill.engine == "bill_calculator"
    assert bill.by_component.taxes_levies == 0.0          # owner decision 7, read as 0.0


def test_new_writes_name_the_engines():
    """Every default names the Investment Case engine; the old KPI names are gone."""
    assert NEW_KPIS <= set(CaseKpis.model_fields) and not OLD_KPIS & set(CaseKpis.model_fields)
    assert InvestmentCase.model_fields["engine"].default == "finance_engine"
    assert BatteryAttribution.model_fields["engine"].default == "finance_engine"
    assert Bill.model_fields["engine"].default == "tariff_engine"


def test_the_gs_fallbacks_state_their_own_label():
    """
    Decision B1: GS's calculator and pro forma name themselves explicitly,
    so the new defaults never label their output as the engine's.
    """
    import inspect

    from services.study import proforma, tariff

    assert 'engine="bill_calculator"' in inspect.getsource(tariff.BillCalculator)
    src = inspect.getsource(proforma.build_investment_case)
    calls = re.findall(r"InvestmentCase\(\s*\*\*(?:\{\*\*)?common\b", src)
    assert calls and len(calls) == src.count("InvestmentCase(")
    assert re.search(r"common = dict\([^)]*engine=PROFORMA_ENGINE", src, re.S)
    assert proforma.PROFORMA_ENGINE == "cash_flow_expander"
    assert proforma.ENGINES[0] == proforma.PROFORMA_ENGINE


# ── the stored report record, through the routes ────────────────────────

def _legacy(record: dict) -> dict:
    """A report record as a pre-U2 build wrote it (the old KPI names and labels)."""
    old = copy.deepcopy(record)
    for case in old["cases"].values():
        k = case.get("kpis")
        if k is not None:
            k.pop("levelised_cost_basis")
            k["lcos"] = k.pop("levelised_cost")
            k["salvage_eur"] = k.pop("terminal_value_eur")
            un = k["unavailable"]
            for new, was in (("levelised_cost", "lcos"), ("terminal_value_eur", "salvage_eur")):
                if new in un:
                    un[was] = un.pop(new)
        case["engine"] = "cash_flow_expander"
        case["sources"]["bill_refs"] = [
            "bill_calculator:none", f"bill_calculator:{case['option_id']}"]
        for s in case["value_streams"]:
            s["engine"] = "bill_calculator"
    for fig in old["report"]["facts"].values():
        fig["engine"] = {"finance_engine": "cash_flow_expander",
                         "tariff_engine": "bill_calculator"}.get(fig["engine"], fig["engine"])
    return old


def test_a_stored_pre_u2_report_record_reads_through_the_routes(
        client, api_project, studies_on, fake, project_storage_dir):
    from routers.studies import REPORT_AUX
    from services.study import store
    from tests.test_study_report_routes import _full, _post

    name = "pre-u2-record"
    sid = _full(client, api_project, name)
    _post(client, name, sid)
    pdir = project_storage_dir(name)
    record = store.load_aux(pdir, sid, REPORT_AUX)
    assert record["cases"], "the assembled report stored no case"
    old = _legacy(record)
    store.save_aux(pdir, sid, REPORT_AUX, old)

    r = client.get(f"/api/projects/{name}/studies/{sid}/report")
    assert r.status_code == 200, r.text
    x = client.get(f"/api/projects/{name}/studies/{sid}/report.xlsx")
    assert x.status_code == 200, x.text

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(x.content))
    cells = {c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row
             if isinstance(c.value, str)}
    assert {"levelised_cost", "terminal_value_eur"} <= cells
    assert not OLD_KPIS & cells


@pytest.fixture
def studies_on(monkeypatch):
    from tests.study_s4_support import enable_studies

    yield from enable_studies(monkeypatch)


@pytest.fixture
def fake(monkeypatch):
    from services import solver_service
    from tests.test_study_runner import FakeSolver

    solver = FakeSolver()
    monkeypatch.setattr(solver_service, "run_simulation", solver)
    return solver
