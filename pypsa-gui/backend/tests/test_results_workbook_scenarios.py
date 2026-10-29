"""
The results workbook's Scenarios sheet (owner decision RW-SCEN-1): the base
and its sensitivity runs side by side, values equal to the services the
sensitivity comparison and the Overview read, deltas live only between two
number cells, each run read from its own context (no global bleed), and a
case the user may not open counted, never shown.

The base (tests.test_workflow_tools.baseline) serves 5 MW for 2 h from G1 at
10 EUR/MWh; a case at c EUR/MWh has opex = 10 c EUR (the hand oracle).
"""
# ruff: noqa: F811 -- imported pytest fixtures
from __future__ import annotations

import functools
import io
import re
import uuid

import pandas as pd
import pytest
from openpyxl import load_workbook

from harness import sensitivity_comparison as sc
from routers import results as results_router
from services import chat_tools, results_workbook as RW
from services import workflow_tools as wt
from services.finance.export_xlsx import NOT_ESTABLISHED
from services.results_workbook import scenarios as S
from services.results_workbook.overview import CASE_ROWS, overview_figures
from services.solver_service import SolverConfig
from tests.test_results_workbook import _report, _rdf, solved  # noqa: F401
from tests.test_sensitivity_comparison import edit_saved, opex, solve, sweep
from tests.test_workflow_tools import baseline, helpers_state  # noqa: F401

URL = "/api/results/workbook.xlsx"


def _export(client) -> bytes:
    r = client.get(URL)
    assert r.status_code == 200, r.text[:300]
    return r.content


def _sheet(blob):
    wb = load_workbook(io.BytesIO(blob))
    return wb, wb[S.TITLE] if S.TITLE in wb.sheetnames else None


def _header(ws) -> tuple[int, dict[str, int]]:
    for row in range(1, ws.max_row + 1):
        if ws.cell(row, 1).value == "KPI":
            return row, {ws.cell(row, c).value: c for c in range(2, ws.max_column + 1)
                         if ws.cell(row, c).value}
    raise AssertionError("no KPI header")


def _kpi_row(ws, label: str) -> int:
    hits = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == label]
    assert len(hits) == 1, (label, hits)
    return hits[0]


def _texts(ws) -> list[str]:
    return [str(ws.cell(r, c).value) for r in range(1, ws.max_row + 1)
            for c in range(1, ws.max_column + 1) if ws.cell(r, c).value is not None]


def _is_num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _refs(expr: str) -> list[str]:
    return [f"{c}{r}" for c, r in re.findall(r"\$?([A-Z]{1,3})\$?(\d+)", expr)]


def _value(ws, ref: str) -> float:
    """A number cell, or a formula over number cells evaluated by hand."""
    v = ws[ref].value
    if _is_num(v):
        return v
    assert isinstance(v, str) and v.startswith("="), (ref, v)
    expr = v[1:]
    for r in sorted(_refs(expr), key=len, reverse=True):
        expr = expr.replace(r, repr(_value(ws, r)))
    assert re.fullmatch(r"[-+*/(). 0-9e]+", expr), expr
    return eval(expr)  # noqa: S307 -- arithmetic over numbers only, checked above


def _assert_deltas_read_two_numbers(ws):
    """Every formula on the sheet reads two cells, each a number or a formula
    over numbers (never text or an empty cell)."""
    seen = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                refs = _refs(cell.value)
                assert len(refs) == 2, cell.value
                for ref in refs:
                    assert _is_num(_value(ws, ref)), (cell.coordinate, cell.value)
                seen += 1
    return seen


@pytest.fixture
def swept(baseline):
    _ctx, base_id, _root = baseline
    solve(base_id)
    return base_id, sweep(base_id, 5, 20)


def test_the_scenarios_sheet_equals_the_comparison_and_overview_services(client, swept):
    base_id, ids = swept
    wb, ws = _sheet(_export(client))
    assert wb.sheetnames[:3] == ["About", S.TITLE, "Overview"]
    header, cols = _header(ws)
    names = [n for n in cols if not str(n).startswith("Δ")]
    base_name = names[0]
    assert names[1:] == ["Gas 20", "Gas 5"]     # by name, base first

    # The Compare-view rows equal compare_sensitivities, value for value.
    report = chat_tools.DISPATCHERS["compare_sensitivities"](baseline_project_id=base_id)
    by_case = {c["name"]: c for c in report["cases"]}
    for metric, (label, unit, *_r) in sc.METRICS.items():
        row = _kpi_row(ws, f"{label} ({unit}), horizon, Compare view")
        want = report["base"]["values"][metric].get("value")
        got = ws.cell(row, cols[base_name]).value
        assert (got == NOT_ESTABLISHED) if want is None else got == pytest.approx(want, rel=1e-12)
        for name in ("Gas 20", "Gas 5"):
            want = by_case[name]["values"][metric].get("value")
            got = ws.cell(row, cols[name]).value
            assert (got == NOT_ESTABLISHED) if want is None else got == pytest.approx(want, rel=1e-12)
    opex_row = _kpi_row(ws, f"{sc.METRICS['opex_meur'][0]} (M€), horizon, Compare view")
    assert ws.cell(opex_row, cols["Gas 20"]).value == pytest.approx(opex(20), rel=1e-12)
    assert ws.cell(opex_row, cols["Gas 5"]).value == pytest.approx(opex(5), rel=1e-12)
    # A live delta: run − base.
    d = ws.cell(opex_row, cols["Δ Gas 20 − base"]).value
    assert d == f"={ws.cell(opex_row, cols['Gas 20']).column_letter}{opex_row}-" \
                f"{ws.cell(opex_row, cols[base_name]).column_letter}{opex_row}"

    # The Overview rows equal overview_figures on each run's own saved context.
    with chat_tools._acting() as (db, user):
        from services import project_registry
        ctx = wt._load_context(project_registry.resolve_project(db, user, ids["Gas 20"]))
    figs = overview_figures(ctx.network, ctx.solver_state.get("solver_config"),
                            ctx.solver_state.get("last_lost_load"),
                            functools.partial(results_router._result_df, state=ctx.solver_state))
    for line in figs.lines:
        if line.key == "system_lcoe":   # a formula here, checked by hand below
            continue
        row = _kpi_row(ws, f"{line.label.strip()} ({line.unit}), horizon")
        got = ws.cell(row, cols["Gas 20"]).value
        assert (got == NOT_ESTABLISHED) if line.total is None else got == pytest.approx(line.total, rel=1e-9)
    # Opex in the Overview's own unit agrees with the hand oracle too.
    assert ws.cell(_kpi_row(ws, "Opex (M EUR), horizon"), cols["Gas 5"]).value \
        == pytest.approx(opex(5), rel=1e-9)

    # System LCOE: a live formula over each run's own cost and energy cells; by
    # hand, no capex and 10 MWh served, so it is the run's gas price.
    lcoe = _kpi_row(ws, "System LCOE (EUR/MWh), horizon")
    cost = _kpi_row(ws, "Total system cost (M EUR), horizon")
    energy = _kpi_row(ws, "Energy served (GWh), horizon")
    for name, price in ((base_name, 10), ("Gas 20", 20), ("Gas 5", 5)):
        col = ws.cell(lcoe, cols[name]).column_letter
        assert ws.cell(lcoe, cols[name]).value == f"={col}{cost}*1000/{col}{energy}"
        assert _value(ws, f"{col}{lcoe}") == pytest.approx(price, rel=1e-9)
    assert _value(ws, ws.cell(lcoe, cols["Δ Gas 5 − base"]).coordinate) == pytest.approx(-5, rel=1e-9)

    # No case: every investment-case row and its delta is not_established, no formula.
    for _key, label, unit, _reader in CASE_ROWS:
        row = _kpi_row(ws, f"{label} ({unit}), horizon")
        assert {ws.cell(row, c).value for c in cols.values()} == {NOT_ESTABLISHED}
    assert _assert_deltas_read_two_numbers(ws) > 0
    assert len(ws._charts) == 2


def test_a_project_without_sensitivity_runs_has_no_scenarios_sheet(client, baseline):
    _ctx, base_id, _root = baseline
    solve(base_id)
    wb, ws = _sheet(_export(client))
    assert ws is None and wb.sheetnames[:2] == ["About", "Overview"]


def test_a_case_of_another_org_is_counted_never_shown(client, swept, second_identity):
    """Per-run ACL: the check opening the project makes (resolve_project)."""
    from db.models import Project
    from db.session import SessionLocal
    _base_id, ids = swept
    with SessionLocal() as db:
        row = db.get(Project, uuid.UUID(ids["Gas 5"]))
        row.org_id = uuid.UUID(str(second_identity["org_id"]))
        db.commit()
    _wb, ws = _sheet(_export(client))
    texts = _texts(ws)
    assert "1 scenario(s) not shown: no access" in texts
    assert not any("Gas 5" in t for t in texts)
    _header_row, cols = _header(ws)
    assert "Gas 20" in cols


def test_each_run_reads_its_own_context_and_the_active_one_is_untouched(client, swept, monkeypatch):
    """No global bleed: a case with its own uploaded series (user_ts) is read
    into its own context; the active context and the module-global store are
    unchanged by the export."""
    from services.pypsa_service import PyPSAService
    from services.user_timeseries import _user_ts
    _base_id, ids = swept

    def profile(ctx):
        series = pd.Series([3.0, 3.0], index=ctx.network.snapshots, name="L1")
        ctx.network.loads_t.p_set["L1"] = series
        ctx.user_ts[("loads", "p_set", "L1")] = series
    edit_saved(ids["Gas 20"], profile)       # not re-solved: the case reads out of date

    loaded = {}
    real = wt._load_context

    def spy(project):
        ctx = real(project)
        loaded[project.name] = ctx
        return ctx
    monkeypatch.setattr(wt, "_load_context", spy)

    active = PyPSAService.get_active_context()
    before = (id(active), id(active.network), dict(_user_ts),
              set(active.solver_state.get("lopf_results") or {}), wt._network_input_fingerprint(active))
    assert not before[2]                     # the base has no uploaded series
    _wb, ws = _sheet(_export(client))
    after = PyPSAService.get_active_context()
    assert (id(after), id(after.network)) == before[:2]
    assert dict(_user_ts) == before[2]
    assert set(after.solver_state.get("lopf_results") or {}) == before[3]
    assert wt._network_input_fingerprint(after) == before[4]
    # The case's series went into its own context only.
    case = loaded["Gas 20"]
    assert case is not after and case.user_ts is not after.user_ts
    assert list(case.user_ts) == [("loads", "p_set", "L1")]
    assert list(loaded["Gas 5"].user_ts) == []
    run_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Gas 20")
    assert "out of date" in ws.cell(run_row, 2).value


def test_a_stale_case_shows_its_reason_and_no_figures_or_deltas(client, swept):
    _base_id, ids = swept

    def reprice(ctx):
        ctx.network.generators.loc["G1", "marginal_cost"] = 99
    edit_saved(ids["Gas 5"], reprice)
    _wb, ws = _sheet(_export(client))
    run_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Gas 5")
    assert "out of date" in ws.cell(run_row, 2).value
    _h, cols = _header(ws)
    opex_row = _kpi_row(ws, "Opex (M EUR), horizon")
    assert ws.cell(opex_row, cols["Gas 5"]).value == NOT_ESTABLISHED
    assert ws.cell(opex_row, cols["Δ Gas 5 − base"]).value == NOT_ESTABLISHED
    assert isinstance(ws.cell(opex_row, cols["Δ Gas 20 − base"]).value, str) \
        and ws.cell(opex_row, cols["Δ Gas 20 − base"]).value.startswith("=")
    _assert_deltas_read_two_numbers(ws)


# ── the sheet from given runs (no project store) ─────────────────────────

def _blob(n, scen, report=None, case_state=None) -> bytes:
    return RW.build_results_workbook(RW.WorkbookInputs(
        network=n, cfg=SolverConfig(), lost_load=None, result_df=_rdf, views=[], ic_report=report,
        case_state=case_state, project="Base", scenarios=scen))


def _run(n, name, **kw) -> S.Run:
    return S.Run(name=name, figs=overview_figures(n, SolverConfig(), None, _rdf), **kw)


def test_per_run_case_staleness_and_a_not_established_delta(solved):
    """The base has a current case; one run's case is current, another's is
    out of date (labelled, figures not shown), a third has none: its deltas
    over the case rows are not_established, not a formula."""
    rep = _report()
    current = {"present": True, "stale": False, "changed": []}
    stale = {"present": True, "stale": True, "changed": ["finance"], "reason": "inputs_changed"}
    scen = S.Scenarios(runs=[_run(solved, "A", ic_report=rep, case_state=current),
                             _run(solved, "B", ic_report=rep, case_state=stale),
                             _run(solved, "C")])
    ws = _sheet(_blob(solved, scen, rep, current))[1]
    runs = {ws.cell(r, 1).value: ws.cell(r, 3).value for r in range(1, ws.max_row + 1)}
    assert runs["A"] == "current" and runs["C"] == "none: run the investment case"
    assert runs["B"].startswith("out of date (changed: finance)")
    _h, cols = _header(ws)
    row = _kpi_row(ws, "NPV, project post-tax at the WACC (EUR), horizon")
    assert ws.cell(row, cols["A"]).value == pytest.approx(rep.npv_at_wacc)
    assert ws.cell(row, cols["B"]).value == NOT_ESTABLISHED
    assert ws.cell(row, cols["Δ A − base"]).value.startswith("=")
    assert ws.cell(row, cols["Δ B − base"]).value == NOT_ESTABLISHED
    assert ws.cell(row, cols["Δ C − base"]).value == NOT_ESTABLISHED
    _assert_deltas_read_two_numbers(ws)


def test_not_comparable_reasons_are_written_and_hidden_runs_counted(solved):
    scen = S.Scenarios(runs=[S.Run(name="H", reasons=["baseline_unavailable", "horizon_differs"])],
                       no_access=2, over_cap=1, base_current=False)
    ws = _sheet(_blob(solved, scen))[1]
    texts = _texts(ws)
    base_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Base (base)")
    assert ws.cell(base_row, 2).value.startswith(sc._BASE_REASONS["case_stale"])
    assert f"1 scenario(s) not shown: the sheet compares up to {sc.MAX_CASES}, by name" in texts
    assert "2 scenario(s) not shown: no access" in texts
    assert any(t.startswith("1 scenario(s) not shown") for t in texts)
    assert any(sc._REASONS["horizon_differs"] in t for t in texts)
    _assert_deltas_read_two_numbers(ws)
    _h, cols = _header(ws)
    delta = cols["Δ H − base"]
    assert not any(ws.cell(r, delta).data_type == "f" for r in range(1, ws.max_row + 1))


def test_formula_like_run_names_stay_text(solved):
    scen = S.Scenarios(runs=[_run(solved, '=HYPERLINK("http://x","y")'), _run(solved, "+cmd|' /C calc'!A0")])
    ws = _sheet(_blob(solved, scen))[1]
    for name in ('=HYPERLINK("http://x","y")', "+cmd|' /C calc'!A0"):
        cells = [c for row in ws.iter_rows() for c in row if c.value == name]
        assert cells and all(c.data_type == "s" for c in cells)
    _assert_deltas_read_two_numbers(ws)


def test_system_lcoe_is_a_per_run_formula_per_period_and_the_charts_read_numbers(solved):
    """Two investment periods: each run's System LCOE, per period and over the
    horizon, is cost × 1000 / energy of its own cells (hand value from the
    figures); every chart series reads number or formula cells only."""
    from tests.test_results_workbook import _refs as range_refs
    figs = overview_figures(solved, SolverConfig(), None, _rdf)
    scen = S.Scenarios(runs=[_run(solved, "A"), _run(solved, "B")])
    ws = _sheet(_blob(solved, scen))[1]
    _h, cols = _header(ws)
    cost, energy = figs.line("system_cost"), figs.line("energy_served")
    assert figs.periods
    for period in ["horizon", *figs.periods]:
        c = cost.total if period == "horizon" else cost.by_period[period]
        e = energy.total if period == "horizon" else energy.by_period[period]
        row = _kpi_row(ws, f"System LCOE (EUR/MWh), {period}")
        for name in ("Base", "A", "B"):
            cell = ws.cell(row, cols[name])
            assert cell.data_type == "f", (period, name, cell.value)
            assert _value(ws, cell.coordinate) == pytest.approx(c * 1000 / e, rel=1e-9)
        assert _value(ws, ws.cell(row, cols["Δ A − base"]).coordinate) == pytest.approx(0, abs=1e-9)
    series = []
    for chart in ws._charts:
        for s in chart.series:
            series += range_refs(s.val.numRef.f.split("!")[1])
    assert series
    for ref in series:
        assert ws[ref].data_type == "f" or _is_num(ws[ref].value), (ref, ws[ref].value)


# ── per-view scenario columns (package 02, R1) ───────────────────────────

def _view_keys(ws, label: str) -> dict[str, object]:
    """Run name → the key's horizon value, from the view sheet's key block."""
    head = next(r for r in ws.iter_rows(min_row=4) if r[0].value == "Key result")
    names = [c.value for c in head[3:] if c.value and not str(c.value).startswith("Δ")]
    row = next(r for r in ws.iter_rows(min_row=head[0].row + 1)
               if r[0].value == label and r[2].value == "horizon")
    return {name: row[3 + i].value for i, name in enumerate(names)}


def test_each_view_column_reads_that_cases_own_saved_run(client, swept):
    """Through the route: each case's Capacity Expansion column equals the cost
    breakdown of its own saved network and config, not the base's."""
    from services import project_registry
    from services.results.cost_breakdown import compute_cost_breakdown
    _base_id, ids = swept
    r = client.get(URL, params={"views": "capex"})
    assert r.status_code == 200, r.text
    values = _view_keys(load_workbook(io.BytesIO(r.content))["Capacity Expansion"], "Opex")
    assert set(ids) <= set(values)
    with chat_tools._acting() as (db, user):
        for name, pid in ids.items():
            ctx = wt._load_context(project_registry.resolve_project(db, user, pid))
            expected = compute_cost_breakdown(ctx.network, ctx.solver_state.get("solver_config"))
            assert values[name] == pytest.approx(expected["opex"], rel=1e-9), name
    # The runs differ, so a column read from the base would fail.
    assert len({round(v, 3) for v in values.values()}) == len(values)


def test_a_case_whose_dispatch_the_routes_would_not_serve_has_no_figures(client, swept, monkeypatch):
    from services.pypsa_service import PyPSAService
    active = PyPSAService.get_network()
    real = results_router._dispatch_ready
    monkeypatch.setattr(results_router, "_dispatch_ready",
                        lambda n: real(n) if n is active else False)
    r = client.get(URL, params={"views": "dispatch"})
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Dispatch"]
    assert "not solved yet" in ws["A2"].value
    label = next(c.value for c in ws["A"] if str(c.value).startswith("Generation — "))
    values = _view_keys(ws, label)
    assert [v for k, v in values.items() if k != next(iter(values))] == \
        [NOT_ESTABLISHED] * (len(values) - 1)
