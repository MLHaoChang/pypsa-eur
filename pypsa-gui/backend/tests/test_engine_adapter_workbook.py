"""
U2 WP9 TARGET — `engine_adapter.case_workbook`: IC's finance workbook
(`export_xlsx.build_workbook` over `assemble_finance_sections`) with the
guided sheets appended — `Cash flows (formulas)` (the live-formula view over
the engine's cash series, owner decision 9), `Assumptions` (the ledger with
`engine_path` and the tested range) and `Provenance` (pack version and hash,
compiled digest, export series ref, engines, basis, currency year).

Plan: docs/superpowers/plans/2026-10-05-guided-study-u2-engine-rewire.md WP2
(port of `test_proforma_xlsx.py`), WP9. The evaluator's own two tests are
ported to `tests/test_xlsx_formula_eval.py` (gate U2-S1 C8 b).
"""
from __future__ import annotations

import io

import numpy as np
import openpyxl
import pytest

from tests.xlsx_formula_eval import FormulaEvaluator

SHEET = "Cash flows (formulas)"
#: The engine's project pre-tax cash, column by column (`finance.engine`:
#: incremental operating cash + terminal value − CAPEX − replacements + grants).
ENGINE_COLUMNS = ("CAPEX", "Replacements", "Operating cash (engine)", "Terminal value",
                  "Grants")


def _A():
    from services.study import engine_adapter as A

    return A


def _bundle_and_ledger(option="bess_2h"):
    from tests import test_engine_adapter_case_golden as G

    n, cfg, compiled, ledger = G._solved_ic(option)
    return G._bundle(option), ledger


@pytest.fixture(scope="module")
def book():
    bundle, ledger = _bundle_and_ledger()
    blob = _A().case_workbook(bundle, ledger)
    assert isinstance(blob, bytes) and blob[:2] == b"PK"
    return bundle, ledger, openpyxl.load_workbook(io.BytesIO(blob))


def _find(ws, label):
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == label:
                return cell
    raise AssertionError(f"{label!r} not on {ws.title}")


def _first(ws):
    return _find(ws, "Year").row + 1          # the year-0 row


def test_the_engines_sheets_come_first_and_the_guided_sheets_are_appended(book):
    _b, _l, wb = book
    assert wb.sheetnames[-3:] == [SHEET, "Assumptions", "Provenance"]
    assert len(wb.sheetnames) > 3
    # IC's own workbook, unchanged, ahead of the guided sheets (decision 9).
    assert wb.sheetnames[:2] == ["About", "Summary"] and "CashflowLines" in wb.sheetnames
    assert wb.sheetnames.index("CashflowLines") < wb.sheetnames.index("KPIs") \
        < wb.sheetnames.index(SHEET)


def test_the_npv_cell_is_cf0_plus_npv_of_cf1_to_cfn_and_evaluates_to_the_case(book):
    bundle, _l, wb = book
    ws = wb[SHEET]
    npv_label = _find(ws, "NPV")
    formula = ws.cell(npv_label.row, npv_label.column + 1).value
    first = _first(ws)
    last = first + bundle.view.horizon_years
    net = _find(ws, "Net cash flow").column_letter
    assert formula == f"={net}{first}+NPV($B$1,{net}{first + 1}:{net}{last})"
    got = FormulaEvaluator(wb).value(SHEET, ws.cell(npv_label.row, npv_label.column + 1).coordinate)
    # The task's bar: the evaluated formula is the ENGINE's NPV to 1e-6.
    assert got == pytest.approx(bundle.result.metrics["project_pre_tax_npv"], abs=1e-6)
    assert got == pytest.approx(bundle.view.kpis.npv, abs=1e-6)
    # …and the KPI sheet's NPV is that cell (live), not a copy of the number.
    kpi = _find(wb["KPIs"], "npv")
    assert wb["KPIs"][f"B{kpi.row}"].value == f"='{SHEET}'!B{npv_label.row}"
    assert FormulaEvaluator(wb).value("KPIs", f"B{kpi.row}") == pytest.approx(
        bundle.view.kpis.npv, abs=1e-6)


def test_the_rate_cell_is_the_cases_rate_and_moving_it_moves_the_npv(book):
    bundle, _l, wb = book
    ws = wb[SHEET]
    assert ws["B1"].value == bundle.case.inputs.wacc_nominal == bundle.view.discount_rate
    npv = _find(ws, "NPV")
    coord = ws.cell(npv.row, npv.column + 1).coordinate
    ws["B1"].value = 0.10
    try:
        cash = bundle.result.cash["project_pre_tax"]
        want = float(np.sum(cash / 1.10 ** np.arange(len(cash))))
        assert FormulaEvaluator(wb).value(SHEET, coord) == pytest.approx(want, abs=1e-6)
    finally:                    # the workbook is the module's: put the rate back
        ws["B1"].value = bundle.view.discount_rate


def test_every_year_row_evaluates_to_the_engines_cash(book):
    bundle, _l, wb = book
    ws = wb[SHEET]
    ev = FormulaEvaluator(wb)
    first = _first(ws)
    col = {name: _find(ws, name).column_letter for name in (
        "Net cash flow", "Discount factor", "Discounted cash flow", "Cumulative discounted",
        "Market revenue at duals (excluded)", "Check: savings - O&M - operating cash")}
    cash = bundle.result.cash["project_pre_tax"]
    for y in bundle.view.years:
        r = first + y.year
        for name in ("Net cash flow", "Discount factor", "Discounted cash flow",
                     "Cumulative discounted"):
            assert str(ws[f"{col[name]}{r}"].value).startswith("="), (name, r)
        assert ev.value(SHEET, f"{col['Net cash flow']}{r}") == pytest.approx(
            float(cash[y.year]), abs=1e-6)
        assert ev.value(SHEET, f"{col['Net cash flow']}{r}") == pytest.approx(
            y.net_cash_flow, abs=1e-6)
        assert ev.value(SHEET, f"{col['Discount factor']}{r}") == pytest.approx(
            (1 + bundle.view.discount_rate) ** -y.year, rel=1e-12)
        assert ev.value(SHEET, f"{col['Discounted cash flow']}{r}") == pytest.approx(
            y.discounted_cash_flow, abs=1e-6)
        assert ev.value(SHEET, f"{col['Cumulative discounted']}{r}") == pytest.approx(
            y.cumulative_discounted, abs=1e-6)
        # the guided decomposition reconciles with the engine's operating cash
        assert abs(ev.value(SHEET, f"{col['Check: savings - O&M - operating cash']}{r}")) \
            <= 1e-6 * max(1.0, abs(y.savings))
    # The excluded column is read by no formula.
    mr = col["Market revenue at duals (excluded)"]
    assert all(mr not in str(c.value) for row in ws.iter_rows() for c in row
               if isinstance(c.value, str) and c.value.startswith("="))


def test_the_money_columns_are_the_engines_own_series(book):
    """A view over the engine's cash series, not the guided view's copy."""
    bundle, _l, wb = book
    ws = wb[SHEET]
    first = _first(ws)
    r = bundle.result
    series = {"CAPEX": r.op.capex, "Replacements": r.op.replacement,
              "Operating cash (engine)": r.op_incremental["net"], "Terminal value": r.terminal,
              "Grants": r.incentives.grant}
    for name in ENGINE_COLUMNS:
        c = _find(ws, name).column_letter
        got = [ws[f"{c}{first + t}"].value for t in range(len(series[name]))]
        # openpyxl writes a float with 16 significant digits (`%.16g`).
        assert got == pytest.approx([float(v) for v in series[name]], rel=1e-15, abs=0), name
    net = _find(ws, "Net cash flow").column_letter
    L = {name: _find(ws, name).column_letter for name in ENGINE_COLUMNS}
    assert ws[f"{net}{first}"].value == (
        f"={L['Operating cash (engine)']}{first}+{L['Terminal value']}{first}"
        f"-{L['CAPEX']}{first}-{L['Replacements']}{first}+{L['Grants']}{first}")


def test_assumptions_are_the_ledger_with_engine_paths_and_provenance_is_complete(book):
    from services.study import packs

    bundle, ledger, wb = book
    ws = wb["Assumptions"]
    header = [c.value for c in ws[1]]
    assert "engine_path" in header
    assert [c.value for c in ws["A"][1:] if c.value] == [r.key for r in ledger.rows]
    paths = {ws.cell(i, 1).value: ws.cell(i, header.index("engine_path") + 1).value
             for i in range(2, ws.max_row + 1)}
    assert paths == {r.key: r.engine_path for r in ledger.rows}
    for col in ("tested_low", "tested_high", "tested_range_note"):
        assert col in header
    prov = {r[0].value: r[1].value for r in wb["Provenance"].iter_rows() if r[0].value}
    # Gate U2-WP9 N8: shown as the defaults version (the stored key stays).
    assert prov["defaults_version"] == ledger.ledger_version and "library_version" not in prov
    assert prov["ledger_hash"] == packs.ledger_hash(ledger)
    assert prov["compiled_digest"] == bundle.compiled.digest
    assert prov["finance_digest"] == bundle.finance.digest
    assert {"tariff_engine", "finance_engine"} <= set(prov["engines"].split(", "))
    assert prov["engine"] == "finance_engine"
    assert prov["basis"] == "real, pre-tax, without subsidy"
    assert prov["currency_year"] == 2020
    stamp = bundle.compiled.tariff_meta["pack_stamp"]
    _pack_id, _, rest = stamp.partition("@")
    version, _, digest = rest.partition(":sha256:")
    assert (prov["pack_version"], prov["pack_hash"]) == (version, digest)
    assert ledger.ledger_version.endswith(f"{version} {digest[:12]}")
    ref = bundle.compiled.config.export_price_ref
    assert prov["export_series"] == f"{ref.id} v{ref.version} sha256:{ref.hash}"
    assert prov["terminal_value_method"] == bundle.case.inputs.terminal_value.method


def test_a_ledger_text_that_looks_like_a_formula_stays_text():
    bundle, ledger = _bundle_and_ledger()
    rows = [r.model_copy(update={"source": '=HYPERLINK("http://x","y")'})
            if r.key == "discount_rate" else r for r in ledger.rows]
    wb = openpyxl.load_workbook(io.BytesIO(_A().case_workbook(
        bundle, ledger.model_copy(update={"rows": rows}))))
    assert _find(wb["Assumptions"], '=HYPERLINK("http://x","y")').data_type == "s"


def test_a_not_established_case_writes_no_npv():
    """
    ADR-0001: a case the guided face does not establish (here the option's
    bill is missing) shows no NPV — not on the formula sheet, and not as the
    engine's headline either (IC's sheets are withheld, the reasons stated).
    """
    from tests import test_engine_adapter_case_golden as G

    bundle = G._bundle("bess_2h", bills={"option": None})
    assert bundle.view.status == "not_established"
    wb = openpyxl.load_workbook(io.BytesIO(_A().case_workbook(bundle, G._ledger())))
    with pytest.raises(AssertionError):
        _find(wb[SHEET], "NPV")
    assert "Summary" not in wb.sheetnames and "CashflowLines" not in wb.sheetnames
    kpi = _find(wb["KPIs"], "status")
    assert wb["KPIs"].cell(kpi.row, 2).value == "not_established"
    notes = {c.value for c in wb[SHEET]["B"] if c.value}
    assert "bill_unavailable_option" in notes
