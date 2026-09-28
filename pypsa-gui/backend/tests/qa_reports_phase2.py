"""
QA: phase 2 of the study-report plan, end to end — a real Energy Hub study
over HTTP, then the GENERATED report through the chat tools on the scripted
fake provider: one section repaired, one section the model cannot write, a
planted wrong number flagged by the audit, one section regenerated under an
instruction into a second version, a run aborted mid-way that still leaves a
partial version, the ``.docx`` re-opened with its "Numbers to check"
appendix, and the in-flight refusal.

Plan: docs/superpowers/plans/2026-09-28-llm-report-generation-increment-1.md
(phase 2 gate: WP3 generator + job + number audit, WP6 chat tools). The unit
and route tests cover each piece on hand-built evidence; this driver pins the
JOURNEY on the numbers the study engines actually produced, through the
tools a chat turn dispatches (``generate_report`` → ``get_report_status`` →
``get_report`` / ``get_report_table`` → ``regenerate_report_section`` →
``abort_report_generation`` → ``export_report_docx``), each called under the
signed-in session's ``ProjectContext`` exactly as ``chat_service`` calls it.

The provider is ``services.llm_fake.FakeProvider`` installed at the seam the
generate route looks up (``routers.report_jobs._provider_for_profile``), so
no key and no network are involved. Optionally, when
``PYPSA_GUI_TEST_LIVE_OPENAI_PROFILE`` names an openai-wire profile (the same
gate as ``docs/superpowers/runbooks/local-openai-wire-probe.md``), ONE real
generate runs on it and only the job's completion is asserted: a small local
model failing a section is an honest result the note records, not a failure.

Runs in CI through ``tests/run_qa_drivers.py``. Exit 0 = every step passed.
"""
from __future__ import annotations

import io
import json
import os
import pathlib
import sys
import threading
import time

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

# THE load-bearing import, and it must come first: `qa_support` pins the
# sandbox before anything imports `main` or `settings`.
from tests import qa_support          # noqa: E402

from tests.eh_stage_fixtures import VOLL, certifiable_weak_network  # noqa: E402

PROJECT = "qa_reports_phase2"
BUDGET_SOLVES = 8
STAGES = ["apply_pack", "ens_solve", "fmea_top", "assemble"]
WRONG_NUMBER_SENTENCE = "the plan sheds 99.9 hours per year"
JOB_TIMEOUT = 120.0

PASS = 0
FAIL = 0


def _step(label: str, ok: bool, msg: str = "") -> None:
    global PASS, FAIL
    if ok:
        PASS += 1
        print(f"  [PASS] {label}" + (f" — {msg}" if msg else ""))
    else:
        FAIL += 1
        print(f"  [FAIL] {label}" + (f" — {msg}" if msg else ""))


def _poll_http(path: str, *, timeout: float = 600.0) -> dict:
    c = qa_support.client()
    t0 = time.time()
    body: dict = {}
    while time.time() - t0 < timeout:
        r = c.get(path)
        if r.status_code == 204:
            return {"status": "no-content"}
        if r.status_code != 200:
            return {"status": f"http-{r.status_code}", "error": r.text[:300]}
        body = r.json()
        if body.get("status") != "running":
            return body
        time.sleep(0.25)
    return {"status": "timeout", **body}


def _as_session(fn, *args, **kwargs):
    """
    Call an in-process chat tool the way a chat turn does: under the signed-in
    session's ``ProjectContext`` (the active project is per SESSION) and with
    the signed-in user as the acting identity (the report tools call the
    project routers, which authorize per caller; `main` binds that once).
    """
    from services.pypsa_service import PyPSAService

    token = PyPSAService.bind_request_context(qa_support.session_context())
    try:
        return fn(*args, **kwargs)
    finally:
        PyPSAService.reset_request_context(token)


def _poll_job(*, timeout: float = JOB_TIMEOUT) -> dict:
    """``get_report_status`` until the job is no longer running."""
    from services import chat_tools

    t0 = time.time()
    body: dict = {}
    while time.time() - t0 < timeout:
        body = _as_session(chat_tools.get_report_status)
        if body.get("status") not in ("running", "no_data"):
            return body
        time.sleep(0.1)
    return {"status": "timeout", **body}


def _docx_text(blob: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(blob))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def _turn(payload: dict | str) -> dict:
    from services.llm_provider import LLMEvent

    text = payload if isinstance(payload, str) else json.dumps(payload)
    return {"events": [LLMEvent(type="text_delta", text=text)],
            "blocks": [{"type": "text", "text": text}]}


def _draft(section_id: str, *paragraphs: str) -> dict:
    return _turn({"section_id": section_id, "paragraphs": list(paragraphs), "bullets": []})


class _Holding:
    """
    A fake provider that answers its FIRST turn at once and waits on an event
    before every later one — so an abort lands between sections, after one
    section was written, and the partial version has something in it.
    """

    name = "fake"

    def __init__(self, turns: list[dict], hold: threading.Event) -> None:
        from services.llm_fake import FakeProvider

        self._inner = FakeProvider(turns)
        self.hold = hold
        self.requests = self._inner.requests

    def stream(self, request):
        if self.requests:
            self.hold.wait(timeout=30)
        yield from self._inner.stream(request)


class _Seam:
    """Install a provider at the generate route's seam for the whole driver."""

    def __init__(self) -> None:
        import routers.report_jobs as jobs

        self.jobs = jobs
        self.orig_provider = jobs._provider_for_profile
        self.orig_resolve = jobs._resolve_profile

    def install(self, provider) -> None:
        self.jobs._provider_for_profile = lambda profile: (provider, None)
        try:
            self.orig_resolve()
        except Exception as exc:                                  # noqa: BLE001
            # No active profile in this sandbox: the fake needs only an id and
            # a model to name on the report.
            from types import SimpleNamespace

            print(f"  ! active profile unresolvable ({type(exc).__name__}); using a stand-in")
            self.jobs._resolve_profile = lambda: SimpleNamespace(
                id="qa-fake-profile", model="qa-fake-model", max_output_tokens=None)

    def live(self, profile_id: str) -> None:
        from services import llm_config

        self.jobs._provider_for_profile = self.orig_provider
        self.jobs._resolve_profile = lambda: llm_config.resolve_profile(profile_id)

    def restore(self) -> None:
        self.jobs._provider_for_profile = self.orig_provider
        self.jobs._resolve_profile = self.orig_resolve


# ── the journey ───────────────────────────────────────────────────────────

def section_1_real_study() -> dict | None:
    print("\n[1] a real strong_grid study on the certifiable weak network")
    c = qa_support.client()
    # Install UNBOUND, then save: the first save claims the name AND binds the
    # session context's storage_dir (the lesson qa_reports_phase0 records).
    qa_support.install_network(certifiable_weak_network())
    qa_support.save_project(PROJECT)

    r = c.put("/api/simulation/solver_config",
              json={"solver_name": "highs", "voll": VOLL})
    _step("the solver config takes the VoLL", r.status_code == 200, r.text[:200])
    r = c.post("/api/results/eh_study", json={
        "archetype": "strong_grid", "budget_solves": BUDGET_SOLVES, "stages": STAGES})
    _step("the study starts", r.status_code == 200, r.text[:200])
    if r.status_code != 200:
        return None
    study = _poll_http("/api/results/eh_study")
    _step("the study finishes", study.get("status") == "done",
          f"status={study.get('status')} error={str(study.get('error'))[:200]}")
    if study.get("status") != "done":
        return None
    r = c.get("/api/results/eh_reference_design")
    _step("the report is persisted for GET", r.status_code == 200, r.text[:200])
    if r.status_code != 200:
        return None
    rep = r.json()
    comp = rep.get("completeness") or {}
    fmea = (rep.get("sections") or {}).get("fmea_top") or {}
    top = (fmea.get("payload") or {}).get("top") or []
    _step("fmea_top is established with at least one ranked mode",
          comp.get("fmea_top") == "ok" and len(top) >= 1,
          f"status={comp.get('fmea_top')} rows={len(top)}")
    ok = [k for k in comp if comp[k] == "ok"]
    _step("at least three EH sections are established (summary + repair + failure "
          "need distinct targets)", len(ok) >= 3, f"ok={ok}")
    return rep


def _plan_script(rep: dict) -> tuple[list[str], list[dict], dict]:
    """
    The target sections in DOCUMENT order and the fake's script for them.

    The job writes sections in document order (`REPORT_SECTIONS` order after
    the executive summary), so the script must be scripted in that order —
    the correction WP3 recorded.
    """
    from models.energy_hub import REPORT_SECTIONS
    from services.reports.assemble import EXECUTIVE_SUMMARY_ID

    comp = rep.get("completeness") or {}
    ok_sections = [s for s in REPORT_SECTIONS if comp.get(s) == "ok"]
    others = [s for s in ok_sections if s != "fmea_top"]
    repair_section = others[0]
    garbage_section = others[1] if len(others) > 1 else "fmea_top"
    top = ((rep.get("sections") or {}).get("fmea_top") or {}).get("payload", {}).get("top") or []
    right_value = float(top[0].get("criticality_eur_per_year") or 0.0)
    right_number = f"{right_value:.1f}"

    targets = [EXECUTIVE_SUMMARY_ID] + ok_sections
    script: list[dict] = []
    for sid in targets:
        if sid == EXECUTIVE_SUMMARY_ID:
            script.append(_draft(
                sid,
                f"The dominant residual failure mode carries a criticality of "
                f"{right_number} per year, and {WRONG_NUMBER_SENTENCE}."))
        elif sid == repair_section:
            script.append(_turn("Here is the section you asked for, in prose."))
            script.append(_draft(sid, f"The {sid} section, established by the study."))
        elif sid == garbage_section:
            script.append(_turn("not json at all"))
            script.append(_turn("{still: not json"))
        else:
            script.append(_draft(sid, f"The {sid} section, established by the study."))
    return targets, script, {
        "repair": repair_section, "garbage": garbage_section,
        "right_number": right_number, "ok": ok_sections,
    }


def section_2_generate(rep: dict, seam: _Seam) -> tuple[str, dict] | None:
    print("\n[2] generate_report on the fake provider → status → done")
    from services import chat_tools
    from services.llm_fake import FakeProvider
    from services.reports.assemble import EXECUTIVE_SUMMARY_ID

    targets, script, plan = _plan_script(rep)
    print(f"      targets={targets} repair={plan['repair']} garbage={plan['garbage']} "
          f"right_number={plan['right_number']}")
    provider = FakeProvider(script)
    seam.install(provider)

    st = _as_session(chat_tools.get_report_status)
    _step("get_report_status is no_data before any run",
          st.get("status") == "no_data" and "generate_report" in str(st.get("message")),
          str(st)[:160])
    started = _as_session(chat_tools.generate_report, title="QA phase-2 client report",
                          sections=targets)
    rid = started.get("report_id") or ""
    _step("generate_report answers running with a report id",
          started.get("status") == "running" and len(rid) == 16
          and "get_report_status" in str(started.get("message")), str(started)[:200])
    if len(rid) != 16:
        return None
    body = _poll_job()
    _step("the job reaches done", body.get("status") == "done",
          f"status={body.get('status')} error={body.get('error')}")
    if body.get("status") != "done":
        return None
    _step("the record carries version 1, the profile and the model, no thread",
          body.get("version") == 1 and body.get("profile_id") and body.get("model")
          and "thread" not in body and "stop_event" not in body, str(body)[:200])
    # The record counts repair TURNS: the repaired section's one, plus the
    # garbage section's failed attempt (WP3's own failure test pins 2).
    _step("two repair turns were counted (one succeeded, the garbage section's failed)",
          body.get("repairs") == 2, f"repairs={body.get('repairs')}")
    failures = [f.get("section_id") for f in body.get("prose_failures") or []]
    _step("the garbage section is the one prose failure", failures == [plan["garbage"]],
          f"prose_failures={failures}")
    _step("the fake saw one request per section plus the two extra turns",
          len(provider.requests) == len(targets) + 2,
          f"requests={len(provider.requests)} targets={len(targets)}")
    _step("no request carried tools",
          all(r.tools == [] for r in provider.requests))

    doc = _as_session(chat_tools.get_report)
    size = len(json.dumps(doc, default=str))
    _step("get_report (default: the newest report) fits the 4000-char result cap",
          size < 4000, f"size={size} outline={doc.get('outline')}")
    _step("it is the generated report, version 1",
          doc.get("report_id") == rid and doc.get("version") == 1
          and doc.get("mode") == "generated" and doc.get("title") == "QA phase-2 client report",
          str({k: doc.get(k) for k in ("report_id", "version", "mode", "title")}))
    by_id = {s.get("section_id"): s for s in doc.get("sections") or []}
    bad_source = [sid for sid in targets
                  if by_id.get(sid, {}).get("source") != ("code" if sid == plan["garbage"] else "llm")]
    _step("every target section is llm-sourced except the garbage one",
          not bad_source, f"wrong={bad_source}")
    garbage = _as_session(chat_tools.get_report, rid, section_id=plan["garbage"])
    gsec = (garbage.get("sections") or [{}])[0]
    _step("the garbage section keeps its status and says 'prose not established'",
          gsec.get("source") == "code" and gsec.get("status") == "ok"
          and str(gsec.get("note", "")).startswith("prose not established")
          and body["profile_id"] in str(gsec.get("note")),
          f"note={str(gsec.get('note'))[:160]}")
    summary = _as_session(chat_tools.get_report, rid, section_id=EXECUTIVE_SUMMARY_ID)
    ssec = (summary.get("sections") or [{}])[0]
    audit = ssec.get("audit") or {}
    unverified = [str(u) for u in audit.get("unverified") or []]
    verified = [str(v) for v in audit.get("verified") or []]
    _step("the planted wrong number is in audit.unverified",
          any("99.9" in u for u in unverified), f"unverified={unverified}")
    _step("the right number copied from the evidence is in audit.verified",
          any(plan["right_number"] in v for v in verified), f"verified={verified}")
    _step("the summary's prose is intact in the per-section read",
          any(b.get("type") == "paragraph" and WRONG_NUMBER_SENTENCE in b.get("md", "")
              for b in ssec.get("blocks") or []))

    listed = _as_session(chat_tools.list_reports)
    meta = next((m for m in listed if m.get("report_id") == rid), {})
    generation = meta.get("generation") or {}
    _step("list_reports shows it as generated with the job's generation summary",
          meta.get("mode") == "generated" and meta.get("latest_version") == 1
          and generation.get("repairs") == 2
          and [f.get("section_id") for f in generation.get("prose_failures") or []]
          == [plan["garbage"]],
          f"mode={meta.get('mode')} latest_version={meta.get('latest_version')} "
          f"generation={str(generation)[:160]}")

    top = ((rep.get("sections") or {}).get("fmea_top") or {}).get("payload", {}).get("top") or []
    page = _as_session(chat_tools.get_report_table, rid, "fmea_top", limit=1)
    _step("get_report_table pages the fmea_top rows",
          page.get("total_count") == len(top) and page.get("returned") == 1
          and page.get("has_more") == (len(top) > 1)
          and isinstance(page.get("columns"), list) and page.get("items"),
          f"total={page.get('total_count')} returned={page.get('returned')} "
          f"has_more={page.get('has_more')}")
    rest = _as_session(chat_tools.get_report_table, rid, "fmea_top", offset=1)
    _step("the second page is the rest",
          rest.get("returned") == max(len(top) - 1, 0) and rest.get("has_more") is False)
    return rid, plan


def section_3_regenerate(rid: str, plan: dict, seam: _Seam) -> None:
    print("\n[3] regenerate_report_section(fmea_top, 'one paragraph') → version 2")
    from services import chat_tools
    from services.llm_fake import FakeProvider

    provider = FakeProvider([_draft("fmea_top", "One paragraph on the residual modes.")])
    seam.install(provider)
    started = _as_session(chat_tools.regenerate_report_section, rid, "fmea_top",
                          instruction="one paragraph")
    _step("the rewrite starts", started.get("status") == "running"
          and started.get("report_id") == rid, str(started)[:200])
    body = _poll_job()
    _step("the job reaches done with version 2",
          body.get("status") == "done" and body.get("version") == 2,
          f"status={body.get('status')} version={body.get('version')}")
    _step("the instruction travelled inside the fence of the one request",
          len(provider.requests) == 1
          and "one paragraph" in json.dumps(provider.requests[0].messages, default=str),
          f"requests={len(provider.requests)}")
    c = qa_support.client()
    v1 = c.get(f"/api/projects/{PROJECT}/reports/{rid}/versions/1").json()
    v2 = c.get(f"/api/projects/{PROJECT}/reports/{rid}/versions/2").json()
    s1 = {s["section_id"]: json.dumps(s, sort_keys=True) for s in v1.get("sections") or []}
    s2 = {s["section_id"]: json.dumps(s, sort_keys=True) for s in v2.get("sections") or []}
    changed = [sid for sid in s1 if s1[sid] != s2.get(sid)]
    _step("every other section is byte-identical between the versions",
          changed == ["fmea_top"], f"changed={changed}")
    fmea2 = next((s for s in v2.get("sections") or [] if s["section_id"] == "fmea_top"), {})
    _step("fmea_top carries the new paragraph before its table and figure",
          fmea2.get("source") == "llm"
          and (fmea2.get("blocks") or [{}])[0].get("md") == "One paragraph on the residual modes."
          and [b["type"] for b in fmea2.get("blocks") or []][1:3] == ["table_ref", "figure_ref"],
          str([b.get("type") for b in fmea2.get("blocks") or []]))
    _step("the document's tables are the software's, unchanged by the rewrite",
          v1.get("tables") == v2.get("tables"))
    latest = _as_session(chat_tools.get_report, rid)
    _step("get_report reads version 2 by default", latest.get("version") == 2)


def section_4_abort_and_in_flight(rep: dict, seam: _Seam) -> None:
    print("\n[4] a generate held on an Event: second generate → 409; abort → partial version")
    from fastapi import HTTPException

    from services import chat_tools

    targets, script, _plan = _plan_script(rep)
    hold = threading.Event()
    provider = _Holding(script, hold)
    seam.install(provider)
    rid = ""
    try:
        started = _as_session(chat_tools.generate_report, sections=targets)
        rid = started.get("report_id") or ""
        _step("the held generate starts", started.get("status") == "running" and len(rid) == 16)
        try:
            _as_session(chat_tools.generate_report, sections=targets)
            _step("a second generate_report is refused with report_job_in_flight", False,
                  "no exception")
        except HTTPException as exc:
            detail = exc.detail if isinstance(exc.detail, dict) else {}
            _step("a second generate_report is refused with report_job_in_flight",
                  exc.status_code == 409 and detail.get("error_kind") == "report_job_in_flight",
                  f"HTTP {exc.status_code} {detail}")
        st = _as_session(chat_tools.get_report_status)
        _step("the status is running for that report",
              st.get("status") == "running" and st.get("report_id") == rid, str(st)[:160])
        ab = _as_session(chat_tools.abort_report_generation)
        _step("abort_report_generation answers aborting", ab == {"status": "running", "aborting": True},
              str(ab))
    finally:
        hold.set()
    body = _poll_job()
    _step("the job ends aborted", body.get("status") == "aborted",
          f"status={body.get('status')} error={body.get('error')}")
    ab2 = _as_session(chat_tools.abort_report_generation)
    _step("abort is idempotent once finished", ab2.get("aborting") is False, str(ab2))
    if len(rid) != 16:
        return
    c = qa_support.client()
    r = c.get(f"/api/projects/{PROJECT}/reports/{rid}")
    _step("a partial version 1 was saved", r.status_code == 200 and r.json().get("version") == 1,
          f"HTTP {r.status_code}")
    if r.status_code != 200:
        return
    doc = r.json()
    by_id = {s["section_id"]: s for s in doc.get("sections") or []}
    written = [sid for sid in targets if by_id.get(sid, {}).get("source") == "llm"]
    rest = [sid for sid in targets if sid not in written]
    _step("the sections written before the abort are llm, the rest code and noted aborted",
          len(written) >= 1 and rest
          and all("aborted" in str(by_id[sid].get("note")) for sid in rest),
          f"written={written} rest={rest}")
    _step("list_reports records the abort in the generation summary",
          any(m.get("report_id") == rid and (m.get("generation") or {}).get("aborted") is True
              for m in _as_session(chat_tools.list_reports)))


def section_5_export(rid: str, plan: dict) -> None:
    print("\n[5] export_report_docx → chip → blob → Word with the 'Numbers to check' appendix")
    from services import chat_tools
    from services.reports.docx_writer import APPENDIX_HEADING, SECTION_TITLES

    meta = _as_session(chat_tools.export_report_docx, rid, filename="qa-phase2-report")
    _step("the tool returns an agent_export chip named as asked",
          meta.get("kind") == "agent_export" and meta.get("filename") == "qa-phase2-report.docx"
          and meta.get("report_id") == rid and "file strip" in str(meta.get("message")),
          str(meta)[:200])
    c = qa_support.client()
    r = c.get(f"/api/projects/{PROJECT}/uploads/{meta.get('file_id')}/blob")
    _step("the blob route serves it with the Word MIME",
          r.status_code == 200 and r.headers.get("content-type", "").startswith(
              "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
          f"HTTP {r.status_code} {r.headers.get('content-type')}")
    if r.status_code != 200:
        return
    text = _docx_text(r.content)
    _step("the document re-opens with python-docx", bool(text))
    _step("it carries the regenerated paragraph", "One paragraph on the residual modes." in text)
    _step(f"it carries the '{APPENDIX_HEADING}' appendix with the planted number",
          APPENDIX_HEADING in text and "99.9" in text.split(APPENDIX_HEADING, 1)[-1],
          f"appendix present={APPENDIX_HEADING in text}")
    # The writer states a note only for a section whose STATUS is not ok; the
    # garbage section's evidence is ok (its table is there), so its "prose not
    # established" note is the viewer's and the status record's, not Word's.
    _step("the garbage section is in the document under its title, with its evidence",
          SECTION_TITLES[plan["garbage"]] in text, f"title={SECTION_TITLES[plan['garbage']]!r}")
    r = c.get(f"/api/projects/{PROJECT}/uploads")
    kinds = {u.get("file_id"): u.get("kind") for u in (r.json() if r.status_code == 200 else [])}
    _step("the uploads list shows the chip as an agent export",
          kinds.get(meta.get("file_id")) == "agent_export", str(kinds)[:200])


def section_6_live_probe(seam: _Seam) -> None:
    profile_id = os.environ.get("PYPSA_GUI_TEST_LIVE_OPENAI_PROFILE")
    if not profile_id:
        print("\n[6] live probe SKIPPED — set PYPSA_GUI_TEST_LIVE_OPENAI_PROFILE (and _MODEL) "
              "to run one real generate on a local openai-wire model; this means the wire "
              "is UNPROBED here, not that it passed")
        return
    print(f"\n[6] ONE real generate on the live openai-wire profile {profile_id!r}")
    from services import chat_tools, llm_config

    try:
        llm_config.resolve_profile(profile_id)
    except llm_config.ProfileNotConfiguredError:
        model = os.environ.get("PYPSA_GUI_TEST_LIVE_OPENAI_MODEL")
        if not model:
            _step("the live profile exists or PYPSA_GUI_TEST_LIVE_OPENAI_MODEL lets it be created",
                  False, f"profile {profile_id!r} unknown and _MODEL unset")
            return
        llm_config.save_profiles([llm_config.LLMProfile(
            id=profile_id, label="Live probe endpoint",
            preset=os.environ.get("PYPSA_GUI_TEST_LIVE_OPENAI_PRESET", "ollama"),
            wire="openai", base_url=os.environ.get("PYPSA_GUI_TEST_LIVE_OPENAI_BASE_URL"),
            model=model, tools=False, vision=False,
            auth=os.environ.get("PYPSA_GUI_TEST_LIVE_OPENAI_AUTH", "none"),
            fallback_model=None, max_output_tokens=None,
        )], profile_id)
    seam.live(profile_id)
    started = _as_session(chat_tools.generate_report, title="QA phase-2 live probe")
    rid = started.get("report_id") or ""
    _step("the live generate starts", started.get("status") == "running" and len(rid) == 16,
          str(started)[:200])
    if len(rid) != 16:
        return
    body = _poll_job(timeout=900.0)
    _step("the live job reaches done or failed (a model that cannot write a section is "
          "recorded, not retried)", body.get("status") in ("done", "failed"),
          f"status={body.get('status')} error={str(body.get('error'))[:200]} "
          f"failures={[f.get('section_id') for f in body.get('prose_failures') or []]}")
    if body.get("status") != "done":
        return
    listed = _as_session(chat_tools.list_reports)
    meta = next((m for m in listed if m.get("report_id") == rid), {})
    targets = (meta.get("generation") or {}).get("sections") or []
    c = qa_support.client()
    doc = c.get(f"/api/projects/{PROJECT}/reports/{rid}").json()
    by_id = {s["section_id"]: s for s in doc.get("sections") or []}
    bad = [sid for sid in targets
           if not (by_id.get(sid, {}).get("source") == "llm"
                   or str(by_id.get(sid, {}).get("note", "")).startswith("prose not established"))]
    _step("every target section is llm or carries a 'prose not established' note",
          not bad, f"bad={bad} model={body.get('model')}")


def main() -> int:
    print("=" * 60)
    print("QA: reports phase 2 — EH study → generated report on the fake provider")
    print("=" * 60)
    crashed = False
    seam: _Seam | None = None
    try:
        qa_support.reset_backend()
        qa_support.delete_project(PROJECT)
        from services import chat_tools

        chat_tools.set_acting_user(str(qa_support.user().id))
        seam = _Seam()
        rep = section_1_real_study()
        if rep is not None:
            created = section_2_generate(rep, seam)
            if created is not None:
                rid, plan = created
                section_3_regenerate(rid, plan, seam)
                section_4_abort_and_in_flight(rep, seam)
                section_5_export(rid, plan)
            section_6_live_probe(seam)
    except Exception as exc:                                     # noqa: BLE001
        crashed = True
        import traceback
        print(f"\n  [FAIL] the journey raised {type(exc).__name__}: {exc}")
        traceback.print_exc()
    finally:
        try:
            if seam is not None:
                seam.restore()
            from services import chat_tools

            chat_tools.set_acting_user(None)
            qa_support.delete_project(PROJECT)
            qa_support.reset_backend()
        except Exception as exc:                                 # noqa: BLE001
            print(f"  ! cleanup raised {type(exc).__name__}: {exc}")

    total = PASS + FAIL + (1 if crashed else 0)
    print("\n" + "=" * 60)
    print(f"Total: {total}")
    print(f"Pass:  {PASS}")
    print(f"Fail:  {FAIL + (1 if crashed else 0)}")
    print("=" * 60)
    return 1 if (FAIL or crashed) else 0


if __name__ == "__main__":
    sys.exit(main())
