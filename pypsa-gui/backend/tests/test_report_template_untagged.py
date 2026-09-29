"""
WP10 — ``services/reports/template_untagged``: the mapping plan for an
untagged (corporate) template and the body rebuild.

Plan: docs/superpowers/plans/2026-09-28-llm-report-generation-increment-1.md
(Increment 2, WP10). The corporate fixture is BUILT from
``tests/fixtures/report_templates/_build.py`` (see its README): body
positions 0 Title, 1 "Prepared for [Client name]", 2 "<Date>", 3 page
break, 4 TOC field, 5/7/9/11 Heading 1 ("1 Executive Summary",
"2 Introduction", "3 Availability Target", "4 Residual Failure Modes"),
13 Heading 2 "4.1 Critical components", 14 table, 15 Heading 1
"5 Lorem ipsum"; ``body_start_index == 5``; header "ACME Energy
Consulting", footer "Confidential".
"""
from __future__ import annotations

import copy
import importlib.util
import io
import json
import struct
import zlib
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from models.report import (
    Bullets,
    Callout,
    Field,
    Figure,
    FigureRef,
    Paragraph,
    ReportDocument,
    Section,
    SectionAudit,
    Table,
    TableRef,
)
from services.chat_service import _UNTRUSTED_CLOSE, _UNTRUSTED_OPEN
from services.llm_fake import FakeProvider
from services.llm_provider import LLMEvent
from services.reports import prompts
from services.reports.docx_reader import merged_paragraph_text, read_template
from services.reports.docx_writer import APPENDIX_HEADING
from services.reports.generator import SectionFailure
from services.reports.template_untagged import (
    MappingEntry,
    MappingPlan,
    default_mapping,
    propose_mapping,
    render_untagged,
)

_BUILD = Path(__file__).resolve().parent / "fixtures" / "report_templates" / "_build.py"
_BASE = {"model": "test-model", "max_tokens": 1024,
         "system_blocks": prompts.system_blocks()}


def _load_builder():
    spec = importlib.util.spec_from_file_location("report_template_fixtures_wp10", _BUILD)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def corporate(tmp_path_factory) -> bytes:
    builder = _load_builder()
    out = tmp_path_factory.mktemp("report_templates")
    return builder.build_all(out)[builder.CORPORATE].read_bytes()


def _png_1x1() -> bytes:
    def chunk(kind: bytes, body: bytes) -> bytes:
        return (struct.pack(">I", len(body)) + kind + body
                + struct.pack(">I", zlib.crc32(kind + body) & 0xFFFFFFFF))

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 0, 0, 0, 0)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr)
            + chunk(b"IDAT", zlib.compress(b"\x00\x80")) + chunk(b"IEND", b""))


_PNG = _png_1x1()


def _document(*, unverified: list[str] | None = None) -> ReportDocument:
    """
    Five sections: the executive summary and ``target`` match template
    headings, ``fmea_top`` too (with every block type), ``certification``
    is skipped, ``cost`` is not established and matches no heading.
    """
    return ReportDocument(
        report_id="0123456789abcdef", version=2, title="Client study report",
        created_at="2026-09-28T10:00:00+00:00", evidence_hash="c" * 64,
        profile_id="local-llama", model="llama-3.1-8b", mode="generated",
        sections=[
            Section(section_id="executive_summary", heading="Executive summary",
                    source="code", status="ok", blocks=[
                        Callout(kind="gap", text="the demand profile is frozen"),
                        Callout(kind="disclosure", text="excludes load-shedding cost"),
                        TableRef(table_id="headline"),
                    ]),
            Section(section_id="target",
                    heading="Availability target and achieved adequacy",
                    source="llm", status="ok",
                    blocks=[Paragraph(md="The target is met.")]),
            Section(section_id="certification", heading="Certification",
                    source="code", status="skipped", note="budget exhausted",
                    blocks=[Callout(kind="not_established",
                                    text="this section was not run in this study "
                                         "(budget exhausted).")]),
            Section(section_id="cost", heading="Cost at target",
                    source="code", status="not_established",
                    note="no cost stage ran", blocks=[]),
            Section(section_id="fmea_top", heading="Residual failure modes",
                    source="llm", status="ok",
                    audit=SectionAudit(unverified=unverified or []),
                    blocks=[
                        Paragraph(md="The **top mode** is *backup* at `100,000 €/yr` "
                                     "([source](https://example.org/x))."),
                        Bullets(items=["first point", "second point"]),
                        TableRef(table_id="fmea_top", caption="Ranked modes"),
                        FigureRef(figure_id="fmea_pareto", caption="Pareto"),
                        Field(key="Engine", value="copt"),
                        TableRef(table_id="does_not_exist"),
                    ]),
        ],
        tables={
            "headline": Table(table_id="headline", columns=["Item", "Value"],
                              rows=[["MC LOLE (h/yr)", "not established"],
                                    ["Cost at target", "1,234,568 €"]],
                              caption="Headline results"),
            "fmea_top": Table(table_id="fmea_top",
                              columns=["Rank", "Name", "Criticality"],
                              rows=[["1", "backup", "100,000 €/yr"],
                                    ["2", "import", "25,000 €/yr"],
                                    ["3", "spare", "not established"]]),
        },
        figures={"fmea_pareto": Figure(figure_id="fmea_pareto",
                                       png_file="figures/fmea_pareto.png",
                                       caption="criticality by mode")},
    )


def _body_children(data: bytes) -> list:
    d = Document(io.BytesIO(data))
    return [c for c in d.element.body.iterchildren()
            if isinstance(c.tag, str) and not c.tag.endswith("}sectPr")]


def _child_texts(children) -> list[str]:
    from docx.oxml.ns import qn as _qn
    out = []
    for c in children:
        if c.tag == _qn("w:p"):
            out.append(merged_paragraph_text(c))
        else:
            out.append("\n".join(merged_paragraph_text(p) for p in c.iter(_qn("w:p"))))
    return out


def _bookmarks(data: bytes) -> list[str]:
    d = Document(io.BytesIO(data))
    return [el.get(qn("w:name")) for el in d.element.body.iter(qn("w:bookmarkStart"))]


def _paragraph_styles(data: bytes) -> list[tuple[str, str]]:
    d = Document(io.BytesIO(data))
    return [(p.style.name if p.style is not None else "", merged_paragraph_text(p))
            for p in d.paragraphs]


def _entry(plan: MappingPlan, index: int) -> MappingEntry:
    found = [e for e in plan.entries if e.heading_index == index]
    assert len(found) == 1, (index, plan.entries)
    return found[0]


def _turn(text: str) -> dict:
    return {"events": [LLMEvent(type="text_delta", text=text)],
            "blocks": [{"type": "text", "text": text}]}


# ── default_mapping ──────────────────────────────────────────────────────


def test_default_mapping_on_the_corporate_fixture(corporate):
    outline = read_template(corporate)
    doc = _document()
    plan = default_mapping(outline, doc)
    assert isinstance(plan, MappingPlan)
    # "1 Executive Summary" → executive_summary (number and case ignored).
    e5 = _entry(plan, 5)
    assert e5.action == "keep" and e5.section_ids == ["executive_summary"]
    # "2 Introduction" is real content the report has nothing for: kept, empty.
    e7 = _entry(plan, 7)
    assert e7.action == "keep" and e7.section_ids == []
    assert any("Introduction" in n for n in plan.notes)
    # "3 Availability Target" → target; "4 Residual Failure Modes" → fmea_top.
    assert _entry(plan, 9).section_ids == ["target"]
    assert _entry(plan, 11).section_ids == ["fmea_top"]
    # "5 Lorem ipsum" is a placeholder chapter.
    assert _entry(plan, 15).action == "drop"
    # The Heading 2 "4.1 Critical components" is kept without content.
    assert _entry(plan, 13).action == "keep" and _entry(plan, 13).section_ids == []
    # Headings before the body start (the Title) are not in the plan.
    assert all(e.heading_index >= 5 for e in plan.entries)
    # cost (not established, unmatched) is inserted after the last kept
    # top-level heading — the omission is the finding, not a reason to drop.
    assert [i["section_id"] for i in plan.inserted] == ["cost"]
    assert plan.inserted[0]["after_heading_index"] == 11
    assert plan.inserted[0]["heading"] == "Cost at target"
    # certification is skipped: listed, not inserted.
    assert plan.unmapped_sections == ["certification"]
    # Every section is covered exactly once or listed.
    covered = [s for e in plan.entries for s in e.section_ids]
    covered += [i["section_id"] for i in plan.inserted]
    assert sorted(covered + plan.unmapped_sections) == sorted(
        s.section_id for s in doc.sections)
    assert len(covered) == len(set(covered))


def test_default_mapping_fills_the_cover_placeholders(corporate):
    outline = read_template(corporate)
    plan = default_mapping(outline, _document())
    assert plan.placeholders["[Client name]"] == "Client"
    assert plan.placeholders["<Date>"] == "2026-09-28"
    assert "XXX" not in plan.placeholders
    assert any("[Client name]" in n for n in plan.notes)
    # XXX only occurs in body paragraphs the rebuild removes: said, not filled.
    assert any("XXX" in n and "rebuild" in n for n in plan.notes)


def test_default_mapping_renames_a_near_match_and_uses_synonyms(corporate):
    outline = read_template(corporate)
    outline = outline.model_copy(deep=True)
    outline.headings[1].text = "1 Summary"                      # synonym
    outline.headings[3].text = "3 Availability targets"         # near match
    outline.headings[4].text = "4 Residual failure mode"        # near match
    doc = _document()
    plan = default_mapping(outline, doc)
    assert _entry(plan, 5).section_ids == ["executive_summary"]
    assert _entry(plan, 9).section_ids == ["target"]
    assert _entry(plan, 11).section_ids == ["fmea_top"]
    for index in (5, 9, 11):
        assert _entry(plan, index).action in ("keep", "rename")


def test_default_mapping_drops_bracket_placeholders_and_samples(corporate):
    outline = read_template(corporate).model_copy(deep=True)
    outline.headings[2].text = "2 [Chapter title]"
    outline.headings[3].text = "3 Sample chapter"
    plan = default_mapping(outline, _document())
    assert _entry(plan, 7).action == "drop"
    assert _entry(plan, 9).action == "drop"
    # target lost its heading → inserted, never lost.
    assert "target" in [i["section_id"] for i in plan.inserted]


def test_default_mapping_without_a_body_start_uses_every_heading(corporate):
    outline = read_template(corporate).model_copy(deep=True)
    outline.body_start_index = None
    outline.has_toc = False
    plan = default_mapping(outline, _document())
    assert {e.heading_index for e in plan.entries} >= {5, 7, 9, 11, 13, 15}


def test_default_mapping_with_no_kept_heading_inserts_at_the_body_start(corporate):
    outline = read_template(corporate).model_copy(deep=True)
    for h in outline.headings:
        h.text = "Lorem ipsum"
    plan = default_mapping(outline, _document())
    assert all(e.action == "drop" for e in plan.entries)
    inserted = [i["section_id"] for i in plan.inserted]
    assert inserted == ["executive_summary", "target", "cost", "fmea_top"]
    assert all(i["after_heading_index"] == -1 for i in plan.inserted)
    assert plan.unmapped_sections == ["certification"]


# ── render_untagged ──────────────────────────────────────────────────────


@pytest.fixture(scope="module")
def rendered(corporate) -> tuple[bytes, MappingPlan]:
    outline = read_template(corporate)
    doc = _document(unverified=["100,000 €/yr"])
    plan = default_mapping(outline, doc)
    plan.placeholders = {}
    data = render_untagged(doc, corporate, plan, figure_bytes={"fmea_pareto": _PNG})
    return data, plan


def test_render_reopens_and_keeps_the_cover_and_toc_byte_identical(corporate, rendered):
    data, _ = rendered
    assert data[:2] == b"PK"
    before = _body_children(corporate)
    after = _body_children(data)
    from lxml import etree
    for i in range(5):
        assert etree.tostring(before[i]) == etree.tostring(after[i]), i
    assert _child_texts(after[:5]) == _child_texts(before[:5])


def test_render_rebuilds_the_body_with_the_templates_heading_style(rendered):
    data, _ = rendered
    styles = _paragraph_styles(data)
    body = [(s, t) for s, t in styles if t.strip()][3:]   # after cover + TOC
    headings = [(s, t) for s, t in body if s.startswith("Heading")]
    assert headings[0] == ("Heading 1", "1 Executive Summary")
    texts = [t for _, t in headings]
    assert texts[:4] == ["1 Executive Summary", "2 Introduction",
                         "3 Availability Target", "4 Residual Failure Modes"]
    assert "4.1 Critical components" in texts
    assert "Cost at target" in texts
    assert "5 Lorem ipsum" not in texts
    # The Heading 2 kept its level; the inserted section is top level.
    assert ("Heading 2", "4.1 Critical components") in headings
    assert ("Heading 1", "Cost at target") in headings
    # The inserted section comes after the whole chapter 4 (after 4.1).
    assert texts.index("Cost at target") > texts.index("4.1 Critical components")
    # Template sample text is gone; report content is in.
    all_text = "\n".join(t for _, t in styles)
    assert "Body text of the section" not in all_text
    assert "first point" in all_text and "second point" in all_text
    assert "the demand profile is frozen" in all_text
    assert "copt" in all_text
    assert "Not established" in all_text and "no cost stage ran" in all_text
    # A skipped section listed in unmapped_sections is not written.
    assert "budget exhausted" not in all_text


def test_render_writes_bookmarks_for_every_covered_section(rendered):
    data, plan = rendered
    names = _bookmarks(data)
    assert "sec:fmea_top" in names
    assert "sec:executive_summary" in names and "sec:target" in names
    assert "sec:cost" in names
    assert "sec:certification" not in names
    assert "sec:numbers_to_check" in names
    # Unique bookmark ids.
    d = Document(io.BytesIO(data))
    ids = [el.get(qn("w:id")) for el in d.element.body.iter(qn("w:bookmarkStart"))]
    assert len(ids) == len(set(ids))


def test_render_sets_update_fields_and_keeps_header_and_footer(rendered):
    data, _ = rendered
    d = Document(io.BytesIO(data))
    flags = d.settings.element.findall(qn("w:updateFields"))
    assert len(flags) == 1 and flags[0].get(qn("w:val")) == "true"
    assert d.sections[0].footer.paragraphs[0].text == "Confidential"
    assert d.sections[0].header.paragraphs[0].text == "ACME Energy Consulting"


def test_render_appends_the_numbers_to_check_appendix(rendered, corporate):
    data, _ = rendered
    styles = _paragraph_styles(data)
    assert ("Heading 1", APPENDIX_HEADING) in styles
    assert any("100,000 €/yr" in t for _, t in styles)
    # And not when nothing is flagged.
    doc = _document()
    plan = default_mapping(read_template(corporate), doc)
    clean = render_untagged(doc, corporate, plan, figure_bytes={})
    assert APPENDIX_HEADING not in [t for _, t in _paragraph_styles(clean)]


def test_render_replaces_placeholders_on_the_cover_and_in_the_footer(corporate):
    outline = read_template(corporate)
    doc = _document()
    plan = default_mapping(outline, doc)
    plan.placeholders = {"[Client name]": "ACME Grid Co", "<Date>": "2026-09-28",
                         "Confidential": "Confidential — ACME Grid Co"}
    data = render_untagged(doc, corporate, plan, figure_bytes={})
    texts = _child_texts(_body_children(data))
    assert texts[0] == "ACME Grid Co — Energy Hub Reference Design"
    assert texts[1] == "Prepared for ACME Grid Co"
    assert texts[2] == "2026-09-28"
    d = Document(io.BytesIO(data))
    assert d.sections[0].footer.paragraphs[0].text == "Confidential — ACME Grid Co"
    # The Title paragraph kept its style.
    assert _paragraph_styles(data)[0][0] == "Title"


def test_render_replaces_a_placeholder_split_across_runs(corporate):
    d = Document(io.BytesIO(corporate))
    title = d.paragraphs[0]
    for r in title.runs:
        r.text = ""
    title.runs[0].text = "[Client"
    title.add_run(" name] — Split")
    buf = io.BytesIO()
    d.save(buf)
    template = buf.getvalue()
    doc = _document()
    plan = default_mapping(read_template(template), doc)
    plan.placeholders = {"[Client name]": "Split Co"}
    data = render_untagged(doc, template, plan, figure_bytes={})
    assert _child_texts(_body_children(data))[0] == "Split Co — Split"


def test_render_uses_the_templates_own_heading_style_name(corporate):
    from docx.enum.style import WD_STYLE_TYPE
    d = Document(io.BytesIO(corporate))
    chapter = d.styles.add_style("Chapter", WD_STYLE_TYPE.PARAGRAPH)
    chapter.base_style = d.styles["Heading 1"]
    for p in d.paragraphs:
        if p.style.name == "Heading 1":
            p.style = chapter
    buf = io.BytesIO()
    d.save(buf)
    template = buf.getvalue()
    outline = read_template(template)
    assert outline.body_start_index == 5
    assert outline.headings[1].style == "Chapter"
    doc = _document()
    plan = default_mapping(outline, doc)
    data = render_untagged(doc, template, plan, figure_bytes={})
    headings = [(s, t) for s, t in _paragraph_styles(data) if s in ("Chapter", "Heading 1", "Heading 2")]
    assert headings[0] == ("Chapter", "1 Executive Summary")
    # The inserted section takes the same top-level style.
    assert ("Chapter", "Cost at target") in headings
    # The appendix too.
    assert not any(s == "Heading 1" for s, _ in headings)


def test_render_puts_a_second_section_on_one_heading_under_a_subheading(corporate):
    outline = read_template(corporate)
    doc = _document()
    plan = default_mapping(outline, doc)
    e = _entry(plan, 11)
    e.section_ids = ["fmea_top", "cost"]
    plan.inserted = []
    data = render_untagged(doc, corporate, plan, figure_bytes={})
    styles = _paragraph_styles(data)
    assert ("Heading 2", "Cost at target") in styles
    names = _bookmarks(data)
    assert "sec:fmea_top" in names and "sec:cost" in names


def test_render_honours_rename_and_drop(corporate):
    outline = read_template(corporate)
    doc = _document()
    plan = default_mapping(outline, doc)
    e = _entry(plan, 5)
    e.action = "rename"
    e.new_text = "1 Management summary"
    _entry(plan, 7).action = "drop"
    data = render_untagged(doc, corporate, plan, figure_bytes={})
    texts = [t for _, t in _paragraph_styles(data)]
    assert "1 Management summary" in texts
    assert "1 Executive Summary" not in texts
    assert "2 Introduction" not in texts


def test_render_never_loses_a_section_the_plan_forgot(corporate):
    outline = read_template(corporate)
    doc = _document()
    plan = default_mapping(outline, doc)
    plan.inserted = []          # cost is now neither mapped nor listed
    data = render_untagged(doc, corporate, plan, figure_bytes={})
    assert "sec:cost" in _bookmarks(data)


def test_render_without_toc_does_not_add_update_fields(corporate):
    d = Document(io.BytesIO(corporate))
    body = d.element.body
    children = [c for c in body.iterchildren() if isinstance(c.tag, str)
                and not c.tag.endswith("}sectPr")]
    body.remove(children[4])            # the TOC field paragraph
    buf = io.BytesIO()
    d.save(buf)
    template = buf.getvalue()
    outline = read_template(template)
    assert not outline.has_toc and outline.body_start_index == 4
    doc = _document()
    plan = default_mapping(outline, doc)
    plan.placeholders = {}
    data = render_untagged(doc, template, plan, figure_bytes={})
    out = Document(io.BytesIO(data))
    assert out.settings.element.find(qn("w:updateFields")) is None
    assert _child_texts(_body_children(data))[:4] == _child_texts(children[:4])


# ── propose_mapping ──────────────────────────────────────────────────────


def _model_plan() -> dict:
    return {
        "entries": [
            {"heading_index": 5, "action": "keep", "new_text": None,
             "section_ids": ["executive_summary"]},
            {"heading_index": 7, "action": "keep", "new_text": None, "section_ids": []},
            {"heading_index": 9, "action": "rename", "new_text": "3 Availability target",
             "section_ids": ["target"]},
            {"heading_index": 11, "action": "keep", "new_text": None,
             "section_ids": ["fmea_top"]},
            {"heading_index": 13, "action": "drop", "new_text": None, "section_ids": []},
            {"heading_index": 15, "action": "drop", "new_text": None, "section_ids": []},
            {"heading_index": 99, "action": "keep", "new_text": None,
             "section_ids": ["cost"]},
        ],
        "inserted": [],
        "placeholders": {"[Client name]": "ACME Grid Co", "<Date>": "2026-09-28"},
        "unmapped_sections": ["certification"],
        "notes": ["from the model"],
    }


def test_propose_mapping_returns_the_models_plan_sanitised(corporate):
    outline = read_template(corporate)
    doc = _document()
    provider = FakeProvider([_turn("```json\n" + json.dumps(_model_plan()) + "\n```")])
    plan = propose_mapping(provider, base_request=_BASE, outline=outline, doc=doc,
                           language="en")
    assert isinstance(plan, MappingPlan), plan
    assert len(provider.requests) == 1
    indices = [e.heading_index for e in plan.entries]
    assert 99 not in indices
    assert any("99" in n for n in plan.notes)
    assert _entry(plan, 9).action == "rename"
    assert _entry(plan, 9).new_text == "3 Availability target"
    assert _entry(plan, 13).action == "drop"
    # cost sat only on the dropped entry: appended by the default rule, noted.
    assert [i["section_id"] for i in plan.inserted] == ["cost"]
    assert plan.inserted[0]["after_heading_index"] == 11
    assert any("cost" in n for n in plan.notes)
    assert plan.unmapped_sections == ["certification"]
    assert plan.placeholders == {"[Client name]": "ACME Grid Co", "<Date>": "2026-09-28"}
    assert "from the model" in plan.notes
    # The request went out with no tools and the mapping message in the fence.
    req = provider.requests[0]
    assert req.tools == [] and req.tools_stable is True
    body = req.messages[0]["content"]
    text = body if isinstance(body, str) else body[0]["text"]
    assert _UNTRUSTED_OPEN in text and "1 Executive Summary" in text


def test_propose_mapping_drops_unknown_sections_duplicates_and_bad_inserts(corporate):
    outline = read_template(corporate)
    doc = _document()
    raw = _model_plan()
    raw["entries"] = raw["entries"][:6]
    raw["entries"][0]["section_ids"] = ["executive_summary", "nope"]
    raw["entries"][1]["section_ids"] = ["executive_summary"]        # duplicate
    raw["entries"].append({"heading_index": 11, "action": "drop",
                           "new_text": None, "section_ids": []})   # duplicate index
    raw["inserted"] = [
        {"after_heading_index": 15, "section_id": "cost", "heading": "Cost"},   # dropped anchor
        {"after_heading_index": 9, "section_id": "ghost", "heading": "Ghost"},
    ]
    raw["placeholders"] = {"[Client name]": "X", "[Nothing]": "Y"}
    raw["unmapped_sections"] = ["target", "ghost"]
    provider = FakeProvider([_turn(json.dumps(raw))])
    plan = propose_mapping(provider, base_request=_BASE, outline=outline, doc=doc,
                           language="en")
    assert isinstance(plan, MappingPlan), plan
    assert _entry(plan, 5).section_ids == ["executive_summary"]
    assert _entry(plan, 7).section_ids == []
    assert _entry(plan, 11).action == "keep" and _entry(plan, 11).section_ids == ["fmea_top"]
    assert [i["section_id"] for i in plan.inserted] == ["cost"]
    assert plan.inserted[0]["after_heading_index"] == 11
    assert plan.placeholders == {"[Client name]": "X"}
    assert plan.unmapped_sections == ["certification"]
    lowered = "\n".join(plan.notes).lower()
    for word in ("nope", "ghost", "[nothing]", "duplicate", "target"):
        assert word in lowered, word
    covered = [s for e in plan.entries for s in e.section_ids]
    covered += [i["section_id"] for i in plan.inserted]
    assert sorted(covered + plan.unmapped_sections) == sorted(
        s.section_id for s in doc.sections)


def test_propose_mapping_rename_without_text_becomes_keep(corporate):
    outline = read_template(corporate)
    doc = _document()
    raw = _model_plan()
    raw["entries"] = raw["entries"][:6]
    raw["entries"][2]["new_text"] = "  "
    provider = FakeProvider([_turn(json.dumps(raw))])
    plan = propose_mapping(provider, base_request=_BASE, outline=outline, doc=doc,
                           language="en")
    assert isinstance(plan, MappingPlan)
    assert _entry(plan, 9).action == "keep" and _entry(plan, 9).new_text is None


def test_propose_mapping_garbage_twice_is_a_section_failure(corporate):
    outline = read_template(corporate)
    provider = FakeProvider([_turn("nonsense " * 20), _turn("still nonsense")])
    out = propose_mapping(provider, base_request=_BASE, outline=outline,
                          doc=_document(), language="en")
    assert isinstance(out, SectionFailure), out
    assert out.reason and out.repairs == 1
    assert len(provider.requests) == 2


def test_propose_mapping_repairs_once_then_accepts(corporate):
    outline = read_template(corporate)
    doc = _document()
    good = copy.deepcopy(_model_plan())
    good["entries"] = good["entries"][:6]
    provider = FakeProvider([_turn("here is my plan"), _turn(json.dumps(good))])
    plan = propose_mapping(provider, base_request=_BASE, outline=outline, doc=doc,
                          language="en")
    assert isinstance(plan, MappingPlan), plan
    assert len(provider.requests) == 2
    roles = [m["role"] for m in provider.requests[1].messages]
    assert roles == ["user", "assistant", "user"]


def test_propose_mapping_plan_renders(corporate):
    outline = read_template(corporate)
    doc = _document()
    provider = FakeProvider([_turn(json.dumps(_model_plan()))])
    plan = propose_mapping(provider, base_request=_BASE, outline=outline, doc=doc,
                           language="en")
    assert isinstance(plan, MappingPlan)
    data = render_untagged(doc, corporate, plan, figure_bytes={})
    texts = [t for _, t in _paragraph_styles(data)]
    assert "3 Availability target" in texts and "4.1 Critical components" not in texts
    assert texts[0] == "ACME Grid Co — Energy Hub Reference Design"
    assert "sec:cost" in _bookmarks(data)


# ── the mapping user message ─────────────────────────────────────────────


def test_mapping_user_message_carries_fence_outline_and_shape(corporate):
    outline = read_template(corporate)
    doc = _document()
    msg = prompts.mapping_user_message(outline, doc, language="de")
    assert msg.count(_UNTRUSTED_OPEN) == 1 and msg.count(_UNTRUSTED_CLOSE) == 1
    inside = msg[msg.index(_UNTRUSTED_OPEN):msg.index(_UNTRUSTED_CLOSE)]
    for text in ("1 Executive Summary", "5 Lorem ipsum", "4.1 Critical components",
                 "Heading 2", "[Client name]", "fmea_top", "certification",
                 "skipped", "Residual failure modes"):
        assert text in inside, text
    assert '"body_start_index":5' in inside.replace(" ", "")
    assert '"has_toc":true' in inside.replace(" ", "")
    assert prompts.MAPPING_SHAPE in msg
    assert '"heading_index"' in prompts.MAPPING_SHAPE
    assert '"inserted"' in prompts.MAPPING_SHAPE
    assert "de" in msg
    assert msg.index(_UNTRUSTED_CLOSE) < msg.index(prompts.MAPPING_SHAPE)
