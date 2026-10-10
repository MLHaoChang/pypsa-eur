"""
Phase 4 end-to-end QA: the single-owner finance engine (Edge Investment Case
P4 plan, "Phase 4 e2e QA gate"; oracles S1 … S3f and F1–F6 in "Fixtures and
oracles"; conventions C1–C14).

Every expected value below is computed by this driver with stdlib arithmetic
from the oracle's own definition (or read from SAM's committed outputs), never
from the product's code; the product is only the thing under test.

  A. The SAM "Single Owner" oracle cases through `services.finance.engine.
     run_case` (the reviewed `sam_case` mapping), to tolerance T: per-year
     arrays |ours − SAM| ≤ max(1, 1e-6·|SAM|), IRR ≤ 1e-5 absolute, NPV / debt
     size / solved price ≤ 1e-6 relative, DSCR ≤ 1e-6 absolute. S1 (all
     equity; project = equity), S1b (solve-for-PPA for 11 % in year 20 → SAM's
     price), S2 (DSCR sculpting, DSRA, reserve interest), S3 (gearing, ITC,
     salvage) and S3f (the gearing-with-fee deviation sized exactly:
     SAM D − ours = 0.36·f·TIC; everything downstream at SAM's debt to T).
  B. The other recorded deviations, each sized: S1l (the remaining basis
     written off: the last year differs by the state + federal shield on
     (1 − 24.5/39)·TIC, every other year to T), S2t (salvage in SAM's sculpting
     basis: SAM D − ours = salvage / 1.3 / 1.07^25), S3d (the DSRA in SAM's
     gearing base: `capex` gives 0.6·TIC, `total_uses` reproduces SAM).
  C. F1 by hand: IDC on two construction years 60/40, one loan at 8 % geared
     60 %, commitment and upfront fees, the annuity from COD.
  D. F2 by hand: the `eu_de` pack (GewSt with the §8 Nr. 1 add-back and its
     own €1m + 60 % pool, KSt on the 2028–2032 path with SolZ, AfA 20 years),
     six years with two loss years.
  E. F3 by hand: the `us_federal` pack (MACRS-7 with 100 % bonus, §163(j) at
     30 % of EBITDA with carryforward, NOL 80 %), five years.
  F. F4 by hand: a base year two years before COD, a PPA on its own 2.5 %
     indexation ending after 8 operating years, class escalation, degradation.
  G. F5 by hand: fixed / EBITDA-multiple / book-value terminal values and an
     escalated replacement.
  H. F6 by hand: the counterfactual on the toy site (1 MW load, 2 MW PV, TOU +
     demand charge, supply at 50 $/MWh) through the config routes and the
     adapter: bill, demand and commodity on both sides, the incremental cash
     557,700 $/yr, the C5 energy term S = 503,700, year 2, the LCOE 176.18
     $/MWh (WP4.6a review B1), and two review variants: + a per-kWh levy (S
     547,500 — B7) and + a non-owner heat pump (incremental 725,600, its draw
     disclosed — B2).
  I. The integration fixture (`build_edge_hourly_year`) through the routes:
     PUT the commercial config, the value flows and the finance inputs
     (If-Match; a stale tag 412s) → solve → `/results/value_flows` closes →
     POST `/results/investment_case` → poll → `/report` and `detail=full` →
     `export.xlsx` read back with openpyxl (every number equal, every None
     `not_established`). The equity IRR is finite and equals the adapter
     test's pin 0.36208 (the runner's test-only `layers` hook carries the same
     25 % SL-10 layer); each year's Σ cashflow lines = `cash.equity_post_tax`;
     year 1's template = the ledger's owner net; the year-1 incremental EBITDA
     identity with the bill rated by hand; the counterfactual block in the
     payload and the About sheet's counterfactual and CFADS rows; the LCOE
     against its definition with the BESS vom as the only own cost (B1); and
     the production path (`us_federal` pack, no hook) finite and reconciled
     (printed as I/pack, after J and K, which read the hooked run's report).
  J. Staleness through the routes: a finance edit → `stale`, changed
     ["finance"]; the revert → current; a solver-config edit (discount_rate) →
     changed ["solver_config"]; the revert → current.
  K. The chat tools on the stored report: `get_investment_case` (headlines,
     current, under the cap), `explain_cashflow` (the stream PVs reconcile to
     the equity NPV) and `solve_ppa_price` on a solvable case (the S1b oracle
     through the router's adapter seam → SAM's price; nothing stored changes).
  L. Refusals and C12: the 7-day fixture refused `template_not_annual:168`
     through the routes (headlines None, the xlsx `not_established`) and, with
     `annualise`, flagged `template_annualised:52.14`; a None ledger line (carried
     as not established in the payload, the counterparty totals and the xlsx) →
     operating `not_established`; a missing escalation class → not
     established through the routes; every None renders as None / "not
     established" / `not_established`, never 0.
  M. IC S0b by hand (plan `2026-10-06-ic-s0b-replacements-terminal.md` S2–S6):
     a two-part battery (power 400 k · 10 y, energy 1.2 M · 15 y) on 25
     operating years from COD 2030 under `part_lifetimes` — the power part
     re-bought in 2039 and 2049, the energy part in 2044, each escalated by
     capex 2 % from the base year; the `remaining_life_annuity` terms (cost ×
     crf × annuity-PV factor, 5 years left on each); the identity at 7 %
     with no escalation and no tax: PV(capex + replacements − TV) =
     2,199,084.18 and the project NPV = (saving − LP annuity) × af(7 %, 25);
     `scale_capex` moves parts and total; a fixed entry for a `part_lifetimes`
     asset refused `replacement_rule_conflict:bess`.
  N. IC G1, G2 by hand (plan `2026-10-07-ic-g1-g2-campus-equipment.md`): the
     defaults pack 2026-10-07 prices a 132/33 kV 40 MVA transformer at
     1,800,000 EUR/unit (lump, 40 y, FOM 1.5 %) in 2026 EUR and keeps every
     2026-10-05 row; that transformer and a 25-y capacitor bank (90 k) as
     extra owner assets beside a 1 MEUR battery on 30 years under
     `part_lifetimes` + `remaining_life_annuity`: the capex with 10 %
     contingency, the FOM lines escalated by opex, the bank re-bought in 2054
     and its remaining 20 years valued; a grant on every asset excludes the
     extras; a case without extras hashes as it did before G2.

(The P1–P3 drivers — `qa_commercial_lp.py`, `qa_billing_contracts.py`,
`qa_value_flows.py` — run beside this one in `tests/run_qa_drivers.py`.)

Run: python tests/qa_investment_case.py   (from pypsa-gui/backend; also run by
tests/run_qa_drivers.py).
"""
from __future__ import annotations

import pathlib
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

try:
    sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
except (AttributeError, ValueError):
    pass

import contextlib  # noqa: E402
import copy  # noqa: E402
import dataclasses  # noqa: E402
import io  # noqa: E402
import math  # noqa: E402
import queue  # noqa: E402
import threading  # noqa: E402
import time  # noqa: E402
from datetime import date, datetime, timedelta  # noqa: E402

from tests import qa_support  # noqa: E402  (ordering is the point)

from services.solver_service import run_simulation  # noqa: E402

PASS = 0
FAIL = 0
CENT = 0.005
PREFIX = "_qa_ic_"
PROJECTS: list[str] = []
STASH: dict = {}                       # scenario I → J, K, L
FIN_URL = "/api/simulation/finance"
CFG_URL = "/api/simulation/solver_config"
VF_URL = "/api/simulation/commercial/value_flows"
STUDY_URL = "/api/results/investment_case"
REPORT_URL = "/api/results/investment_case/report"
XLSX_URL = "/api/results/investment_case/export.xlsx"
NE = "not_established"
PIN_EQUITY_IRR = 0.36207517            # tests/test_finance_case_adapter.py (HiGHS, 1e-5)


def _step(label: str, ok: bool, msg: str = "") -> None:
    global PASS, FAIL
    if ok:
        PASS += 1
    else:
        FAIL += 1
    print(f"  [{'PASS' if ok else 'FAIL'}] {label}" + (f" — {msg}" if msg else ""))


def _cent(a, b) -> bool:
    return a is not None and b is not None and abs(a - b) < CENT


def _rel(a, b, tol) -> bool:
    return a is not None and b is not None and abs(a - b) <= tol * max(abs(b), 1e-300)


def _why(r) -> str:
    return "" if 200 <= r.status_code < 300 else f"{r.status_code} {r.text[:300]}"


def _project(tag: str) -> str:
    name = f"{PREFIX}{tag}"
    if name not in PROJECTS:
        PROJECTS.append(name)
    return name


def _T(ours, sam) -> tuple[bool, str]:
    """Tolerance T on a per-year array: |ours − SAM| ≤ max(1, 1e-6·|SAM|)."""
    if ours is None:
        return False, "ours is None"
    ours, sam = [float(x) for x in ours], [None if x is None else float(x) for x in sam]
    if len(ours) != len(sam):
        return False, f"length {len(ours)} vs {len(sam)}"
    worst, at = 0.0, None
    for i, (a, b) in enumerate(zip(ours, sam)):
        if b is None:
            continue
        d = abs(a - b)
        if d > max(1.0, 1e-6 * abs(b)):
            return False, f"year {i}: {a:.6f} vs SAM {b:.6f}"
        if d > worst:
            worst, at = d, i
    return True, f"max |Δ| {worst:.3g} (year {at})"


def _npv(rate: float, cash) -> float:
    """End-of-year discounting, index 0 undiscounted (SAM, plan C9)."""
    return sum(float(c) / (1.0 + rate) ** i for i, c in enumerate(cash))


def _irr(cash) -> float | None:
    """The stdlib IRR: the root of NPV closest to 0 above −0.99, by a scan and
    bisection (plan C9's rule)."""
    grid = [-0.99 + 0.001 * i for i in range(int((3.0 + 0.99) / 0.001))]
    best = None
    prev_r, prev_v = grid[0], _npv(grid[0], cash)
    for r in grid[1:]:
        v = _npv(r, cash)
        if (prev_v <= 0) != (v <= 0):
            lo, hi, flo = prev_r, r, prev_v
            for _ in range(200):
                mid = 0.5 * (lo + hi)
                fm = _npv(mid, cash)
                if (fm <= 0) == (flo <= 0):
                    lo, flo = mid, fm
                else:
                    hi = mid
            root = 0.5 * (lo + hi)
            if best is None or abs(root) < abs(best):
                best = root
        prev_r, prev_v = r, v
    return best


@contextlib.contextmanager
def _bound():
    """Bind the signed-in client's project context to this thread, so a direct
    call (the adapter, a chat tool) sees what the client's requests see."""
    from services.pypsa_service import PyPSAService

    ctx = qa_support.session_context()
    tok = PyPSAService.bind_request_context(ctx)
    try:
        yield ctx
    finally:
        PyPSAService.reset_request_context(tok)


# ── A / B: the SAM oracle ──────────────────────────────────────────────────


def _sam_run(name: str, *, solve: bool = False, debt_amount: float | None = None):
    from services.finance.engine import run_case
    from tests.fixtures.investment_case.sam import sam_case as S

    case = S.to_finance_case(name, solve=solve)
    if debt_amount is not None:
        t = case.inputs.debt[0].model_copy(update={"gearing": None, "amount": debt_amount})
        case = dataclasses.replace(case, inputs=case.inputs.model_copy(update={"debt": [t]}))
    return run_case(case, layers=S.sam_tax_layers(name))


def _parity(name: str, r, label: str | None = None) -> None:
    from tests.fixtures.investment_case.sam import sam_case as S

    label = label or name.upper()
    e = S.sam_expected(name)
    a, sc = e["arrays"], e["scalars"]
    _step(f"{label}: every section established", all(v == "ok" for v in r.sections.values()),
          f"{r.sections} {r.reasons}")
    op = r.op
    revenue = None if op.revenue is None or r.terminal is None else \
        [x + y for x, y in zip(op.revenue, r.terminal)]
    checks = [
        ("energy (MWh)", sum(op.energy_mwh.values()), a["cf_energy_net"]),
        ("revenue (+ salvage, SAM's total revenue)", revenue, a["cf_total_revenue"]),
        ("operating expenses", op.costs, a["cf_operating_expenses"]),
        ("EBITDA", r.cash["ebitda"], a["cf_ebitda"]),
        ("federal depreciation", r.tax.depreciation["federal"], a["cf_feddepr_total"]),
        ("state depreciation", r.tax.depreciation["state"], a["cf_stadepr_total"]),
        ("federal tax", None if r.tax is None else [-x for x in r.tax.liability["federal"]],
         a["cf_fedtax"]),
        ("state tax", None if r.tax is None else [-x for x in r.tax.liability["state"]],
         a["cf_statax"]),
        ("equity cash pre-tax", r.cash["equity_pre_tax"], a["cf_project_return_pretax"]),
        ("equity cash post-tax", r.cash["equity_post_tax"], a["cf_project_return_aftertax"]),
    ]
    for what, ours, sam in checks:
        ok, msg = _T(ours, sam)
        _step(f"{label}: {what} per year to T", ok, msg)
    m = r.metrics
    irr_ = m["equity_post_tax_irr"]
    _step(f"{label}: post-tax equity IRR = SAM's (≤ 1e-5)",
          irr_ is not None and abs(irr_ - sc["equity_irr_post_tax"]) <= 1e-5,
          f"{irr_} vs {sc['equity_irr_post_tax']}")
    pre = _irr(a["cf_project_return_pretax"])
    _step(f"{label}: pre-tax equity IRR = the stdlib IRR of SAM's pre-tax cash (≤ 1e-5)",
          m["equity_pre_tax_irr"] is not None and pre is not None
          and abs(m["equity_pre_tax_irr"] - pre) <= 1e-5, f"{m['equity_pre_tax_irr']} vs {pre}")
    rate = sc["equity_discount_rate"] / 100.0 if sc["equity_discount_rate"] > 1 else \
        sc["equity_discount_rate"]
    hand_npv = _npv(rate, a["cf_project_return_aftertax"])
    _step(f"{label}: post-tax equity NPV = SAM's (1e-6 rel) and = Σ SAM cash / (1+r)^t",
          _rel(m["equity_post_tax_npv"], sc["equity_npv_post_tax"], 1e-6)
          and _rel(hand_npv, sc["equity_npv_post_tax"], 1e-6),
          f"{m['equity_post_tax_npv']:.4f} vs {sc['equity_npv_post_tax']:.4f} (hand "
          f"{hand_npv:.4f} at {rate:.6f})")
    _step(f"{label}: LCOE nominal / real = SAM's (1e-6 rel)",
          _rel(m["lcoe_nominal_per_mwh"], sc["lcoe_nominal_per_mwh"], 1e-6)
          and _rel(m["lcoe_real_per_mwh"], sc["lcoe_real_per_mwh"], 1e-6),
          f"{m['lcoe_nominal_per_mwh']} / {m['lcoe_real_per_mwh']} vs "
          f"{sc['lcoe_nominal_per_mwh']} / {sc['lcoe_real_per_mwh']}")
    if sc["min_dscr"] is None:
        _step(f"{label}: no debt (debt 0, min DSCR None)",
              m["debt_size"] == 0.0 and m["min_dscr"] is None, f"{m['debt_size']} {m['min_dscr']}")
        return
    d = r.debt
    _step(f"{label}: debt size = SAM's (1e-6 rel)", _rel(m["debt_size"], sc["debt_size"], 1e-6),
          f"{m['debt_size']:.4f} vs {sc['debt_size']:.4f}")
    dsra = [x + y for x, y in zip(a["cf_funding_debtservice"], a["cf_disbursement_debtservice"])]
    for what, ours, sam in (("interest", d.interest, a["cf_debt_payment_interest"]),
                            ("principal", d.principal, a["cf_debt_payment_principal"]),
                            ("balance", d.tranches[0].balance, a["cf_debt_balance"]),
                            ("DSRA balance", d.dsra_balance, a["cf_reserve_debtservice"]),
                            ("DSRA movements", d.dsra_funding, dsra),
                            ("reserve interest", d.reserve_interest, a["cf_reserve_interest"])):
        ok, msg = _T(ours, sam)
        _step(f"{label}: debt {what} per year to T", ok, msg)
    ten = d.tranches[0].tranche.tenor_years
    ok, msg = _T(list(d.cfads[1:ten + 1]), a["cf_cash_for_ds"][1:ten + 1])
    _step(f"{label}: CFADS = SAM's cf_cash_for_ds over the tenor (to T)", ok, msg)
    bad = [(i, d.dscr[i], s) for i, s in enumerate(a["cf_pretax_dscr"])
           if s is not None and not abs(float(d.dscr[i]) - s) <= 1e-6]
    _step(f"{label}: DSCR per year (≤ 1e-6) and min DSCR",
          not bad and m["min_dscr"] is not None and abs(m["min_dscr"] - sc["min_dscr"]) <= 1e-6,
          f"{bad[:2]}; min {m['min_dscr']} vs {sc['min_dscr']}")


def scenario_a() -> None:
    print("\n[A] SAM Single Owner oracle cases through the engine, to tolerance T")
    from tests.fixtures.investment_case.sam import sam_case as S

    for name in ("s1", "s2", "s3"):
        r = _sam_run(name)
        _parity(name, r)
        if name == "s1":
            pp, ep = r.metrics["project_post_tax_irr"], r.metrics["equity_post_tax_irr"]
            _step("S1: all equity — unlevered project IRR = equity IRR",
                  pp is not None and ep is not None and abs(pp - ep) <= 1e-12, f"{pp} vs {ep}")
        if name == "s3":
            sc = S.sam_expected("s3")["scalars"]
            p = S.sam_params("s3")
            hand = p.itc_federal_percent * p.itc_base_share * p.installed_cost
            itc = r.incentives.itc
            _step("S3: the ITC in operating year 1 = SAM's itc_total = 30 % × eligible TIC",
                  _rel(float(sum(itc)), sc["itc_total"], 1e-6) and _rel(hand, sc["itc_total"], 1e-9)
                  and abs(float(itc[1]) - sc["itc_total"]) < 1.0,
                  f"{float(sum(itc)):.2f} vs {sc['itc_total']:.2f} (hand {hand:.2f})")
            debt = p.debt["percent"] * p.installed_cost
            _step("S3: fee 0, DSRA 0 → SAM's debt = 0.6 × TIC (by hand)",
                  _rel(sc["debt_size"], debt, 1e-9) and _rel(r.metrics["debt_size"], debt, 1e-12),
                  f"{r.metrics['debt_size']:.2f} vs {debt:.2f}")

    # S1b: the cashflows at SAM's solved price, then the solve itself.
    _parity("s1b", _sam_run("s1b"), "S1b (at SAM's price)")
    e = S.sam_expected("s1b")["scalars"]
    r = _sam_run("s1b", solve=True)
    m = r.metrics
    _step("S1b: solve-for-PPA status ok", m.get("solve_ppa_status") == "ok",
          f"{m.get('solve_ppa_status')}")
    _step("S1b: solved PPA price = SAM's ppa_price (1e-6 rel)",
          _rel(m.get("solved_ppa_price"), e["ppa_price_per_mwh"], 1e-6),
          f"{m.get('solved_ppa_price')} vs {e['ppa_price_per_mwh']} $/MWh")
    _step("S1b: the solved case's post-tax IRR in year 20 = the 11 % target (≤ 1e-5)",
          m.get("solved_equity_irr_at_target_year") is not None
          and abs(m["solved_equity_irr_at_target_year"] - 0.11) <= 1e-5
          and e["target_year"] == 20,
          f"{m.get('solved_equity_irr_at_target_year')} (SAM year {e['target_year']})")
    _step("S1b: the stored input price is not mutated by the solve",
          r.case.templates[0].lines[0].price == S.sam_params("s1b").ppa_price_per_mwh)

    # S3f: our gearing on TIC, and SAM's one-step fee term.
    p = S.sam_params("s3f")
    e = S.sam_expected("s3f")
    ours = _sam_run("s3f")
    tic, g, f = p.installed_cost, p.debt["percent"], p.debt["fee_share"]
    _step("S3f: our debt (gearing_base='capex') = 0.6 × TIC exactly",
          _rel(ours.metrics["debt_size"], g * tic, 1e-12) and g == 0.6,
          f"{ours.metrics['debt_size']:.4f} vs {g * tic:.4f}")
    dev = e["scalars"]["debt_size"] - ours.metrics["debt_size"]
    _step("S3f: SAM D − ours = g²·f·TIC = 0.36·f·TIC (the recorded deviation, 1e-9 rel)",
          _rel(dev, 0.36 * f * tic, 1e-9) and _rel(g * tic * (1 + g * f), e["scalars"]["debt_size"],
                                                   1e-9),
          f"{dev:.4f} vs {0.36 * f * tic:.4f} (f {f}, TIC {tic:,.0f}); "
          f"{[d['name'] for d in e['deviations']]}")
    _step("S3f: the deviation is recorded in the fixture",
          [d["name"] for d in e["deviations"]] == ["gearing_with_fee"])
    print(f"      (S3f equity IRR at our debt {ours.metrics['equity_post_tax_irr']:.6f} vs SAM "
          f"{e['scalars']['equity_irr_post_tax']:.6f}: the deviation's consequence, not a "
          "parity check)")
    _parity("s3f", _sam_run("s3f", debt_amount=e["scalars"]["debt_size"]), "S3f (at SAM's debt)")


def scenario_b() -> None:
    print("\n[B] Recorded deviations S1l, S2t, S3d, each sized")
    from tests.fixtures.investment_case.sam import sam_case as S

    # S1l — SL-39 over 25 years: SAM drops (1 − 24.5/39) of the basis, we write it off.
    r = _sam_run("s1l")
    e = S.sam_expected("s1l")
    p = S.sam_params("s1l")
    post = [float(x) for x in r.cash["equity_post_tax"]]
    sam = e["arrays"]["cf_project_return_aftertax"]
    ok, msg = _T(post[:-1], sam[:-1])
    _step("S1l: post-tax equity cash to T in every year but the last", ok, msg)
    left = (1 - 24.5 / 39) * p.installed_cost
    shield = left * (p.state_rate + p.federal_rate - p.federal_rate * p.state_rate)
    _step("S1l: last year − SAM = the state + federal shield on (1 − 24.5/39)·TIC",
          _rel(post[-1] - sam[-1], shield, 1e-9),
          f"{post[-1] - sam[-1]:.4f} vs {shield:.4f}")
    _step("S1l: recorded (`remaining_basis_written_off`) and flagged",
          e["deviations"][0]["name"] == "remaining_basis_written_off"
          and "remaining_basis_written_off:federal" in r.flags)

    # S2t — SAM sculpts on the salvage in the last year.
    from services.finance.cashflow import build_operating
    from services.finance.debt import build_debt
    from services.finance.timeline import build_timeline

    def debt_of(case):
        tl = build_timeline(case)
        return build_debt(case.inputs, build_operating(case, tl), tl)

    e = S.sam_expected("s2t")
    p = S.sam_params("s2t")
    d = debt_of(S.to_finance_case("s2t"))
    salvage = p.salvage_share * p.installed_cost
    hand = salvage / p.debt["dscr"] / (1 + p.debt["rate"]) ** p.analysis_years
    _step("S2t: SAM D − ours = salvage / DSCR / (1+r)^25 (1e-9 rel)",
          d.established() and _rel(e["scalars"]["debt_size"] - d.amount, hand, 1e-9),
          f"{e['scalars']['debt_size'] - d.amount:.4f} vs {hand:.4f} "
          f"(DSCR {p.debt['dscr']}, r {p.debt['rate']})")
    _step("S2t: recorded (`salvage_in_cfads`)", e["deviations"][0]["name"] == "salvage_in_cfads")

    # S3d — the DSRA inside SAM's gearing base.
    e = S.sam_expected("s3d")
    p = S.sam_params("s3d")
    case = S.to_finance_case("s3d")
    d = debt_of(case)
    _step("S3d: gearing_base='capex' → 0.6 × TIC (the recorded deviation)",
          _rel(d.amount, p.debt["percent"] * p.installed_cost, 1e-12),
          f"{d.amount:.2f}; SAM {e['scalars']['debt_size']:.2f}, Δ "
          f"{e['scalars']['debt_size'] - d.amount:.2f}")
    t = case.inputs.debt[0].model_copy(update={"gearing_base": "total_uses"})
    d2 = debt_of(dataclasses.replace(case, inputs=case.inputs.model_copy(update={"debt": [t]})))
    ok, msg = _T(d2.interest, e["arrays"]["cf_debt_payment_interest"])
    _step("S3d: gearing_base='total_uses' reproduces SAM's debt (1e-6 rel) and interest (T)",
          _rel(d2.amount, e["scalars"]["debt_size"], 1e-6) and ok,
          f"{d2.amount:.4f} vs {e['scalars']['debt_size']:.4f}; {msg}")
    _step("S3d: recorded (`dsra_in_gearing_base`)",
          e["deviations"][0]["name"] == "dsra_in_gearing_base")


# ── C–G: hand oracles F1–F5 ────────────────────────────────────────────────


def scenario_c() -> None:
    print("\n[C] F1 by hand: IDC on 60/40 construction draws, one loan at 8 %")
    from tests.test_finance_debt import _case, _run, _t

    d, tl = _run(_case([_t(gearing=0.6, commitment_fee=0.005, upfront_fee=0.01)]))
    capex, g, r, ten = 1_000_000.0, 0.6, 0.08, 10
    D = g * capex                                  # sized at COD, IDC inside (WP4.2a)
    # Draws at the construction points: 0.6·P at year 0 compounds one year; 0.4·P none.
    P = D / (0.6 * (1 + r) + 0.4)
    idc = 0.6 * P * r
    _step("F1: the axis is two construction years then COD",
          (tl.y0, tl.cod_year) == (2031, 2033), f"{tl.y0} → {tl.cod_year}")
    _step("F1: cash drawn P = D / (0.6·1.08 + 0.4), draws 60/40",
          _rel(d.tranches[0].principal_drawn, P, 1e-12)
          and _rel(float(d.draws[0]), 0.6 * P, 1e-12) and _rel(float(d.draws[1]), 0.4 * P, 1e-12),
          f"{d.tranches[0].principal_drawn:.4f} vs {P:.4f}")
    _step("F1: IDC = 0.6·P·8 % (one year to the COD point) and D − P in total",
          _rel(float(d.idc[0]), idc, 1e-12) and _rel(d.idc_total, D - P, 1e-9),
          f"{d.idc_total:.4f} vs {D - P:.4f}")
    _step("F1: fees = 1 % of D at close + 0.5 % on the undrawn 40 % over year 0",
          _rel(float(d.fees[0]), 0.01 * D, 1e-12) and _rel(float(d.fees[1]), 0.005 * 0.4 * P, 1e-12),
          f"{float(d.fees[0]):.2f}, {float(d.fees[1]):.2f}")
    uses = capex + (D - P) + 0.01 * D + 0.005 * 0.4 * P
    _step("F1: total uses = capex + IDC + fees", _rel(d.uses["total"], uses, 1e-12),
          f"{d.uses['total']:.2f} vs {uses:.2f}")
    pay = D * r / (1 - (1 + r) ** -ten)
    svc = [float(x) for x in d.service[2:12]]
    _step("F1: annuity service D·r/(1 − 1.08^−10) for 10 years from COD, balance → 0",
          all(_rel(x, pay, 1e-12) for x in svc) and abs(float(d.tranches[0].balance[11])) < 1e-6,
          f"{svc[0]:.4f} vs {pay:.4f}")


def _pool(bases, allowance: float, share) -> list[float]:
    """Loss carryforward by hand: losses join the pool; a positive base is
    offset by at most allowance + share × (base − allowance) (share per year
    when a list), FIFO."""
    pool, out = 0.0, []
    for i, b in enumerate(bases):
        s = share[i] if isinstance(share, list) else share
        if b <= 0:
            pool += -b
            out.append(0.0)
            continue
        cap = min(b, allowance + s * max(0.0, b - allowance)) if allowance else s * b
        use = min(pool, cap)
        pool -= use
        out.append(b - use)
    return out


def scenario_d() -> None:
    print("\n[D] F2 by hand: the eu_de pack (GewSt + KSt + SolZ, two loss years)")
    import numpy as np

    from models.finance import FinanceInputs
    from services.finance.packs.base import load_pack
    from services.finance.tax import compute_tax
    from services.finance.tax_layers import OwnerAsset, resolve_tax_layers
    from services.finance.timeline import Timeline

    cod = date(2031, 1, 1)
    pack = load_pack("eu_de", as_of=date(2026, 1, 1))
    fin = FinanceInputs(financial_close=cod, tax_losses="carryforward", hebesatz_pct=400.0,
                        acquisition_date=date(2024, 6, 1))
    res = resolve_tax_layers(pack, fin, [OwnerAsset("pv", "solar", 2_000_000.0)], cod)
    _step("F2: the layers resolve (GewSt, then KSt), nothing missing",
          res.missing == [] and [ly.name for ly in res.layers] == ["gewst", "kst"],
          f"{res.missing} {[ly.name for ly in res.layers]}")
    if res.missing:
        return
    ebitda = [-3e6, -1e6, 2.5e6, 4e6, 4e6, 4e6]
    interest = 600_000.0
    t = compute_tax(Timeline(y0=2031, cod_year=2031, base_year=2031, analysis_years=6),
                    res.layers, ebitda=np.array(ebitda), basis=2_000_000.0,
                    interest=np.full(6, interest), losses="carryforward")
    afa = 2_000_000.0 / 20                         # PV 20 years, a January COD: a full year
    add_back = 0.25 * max(0.0, interest - 200_000.0)
    gew_base = [e - afa - interest + add_back for e in ebitda]
    gew = [0.035 * 4.0 * x for x in _pool(gew_base, 1e6, 0.6)]
    kst_base = [e - afa - interest for e in ebitda]   # GewSt is not deductible
    rate = {2028: 0.14, 2029: 0.13, 2030: 0.12, 2031: 0.11}
    kst = [rate.get(y, 0.10) * 1.055 * x
           for y, x in zip(range(2031, 2037), _pool(kst_base, 1e6, 0.6))]
    _step("F2: AfA 100 k€ a year", all(_cent(float(x), afa) for x in t.depreciation["kst"]),
          f"{list(t.depreciation['kst'])[:2]}")
    _step("F2: GewSt = 14 % of (EBITDA − AfA − interest + 25 % × (interest − 200k)) after its "
          "own €1m + 60 % pool", all(_cent(float(a), b) for a, b in zip(t.liability["gewst"], gew)),
          f"{[round(float(x)) for x in t.liability['gewst']]} vs {[round(x) for x in gew]}")
    _step("F2: KSt (11 % in 2031, 10 % on) × 1.055 SolZ after the €1m + 60 % pool",
          all(_cent(float(a), b) for a, b in zip(t.liability["kst"], kst)),
          f"{[round(float(x)) for x in t.liability['kst']]} vs {[round(x) for x in kst]}")
    _step("F2: below the Zinsschranke Freigrenze — no flag", t.flags == [], f"{t.flags}")


def scenario_e() -> None:
    print("\n[E] F3 by hand: the us_federal pack (MACRS-7 + 100 % bonus, §163(j), NOL 80 %)")
    import numpy as np

    from models.finance import FinanceInputs
    from services.finance.packs.base import load_pack
    from services.finance.tax import compute_tax
    from services.finance.tax_layers import OwnerAsset, resolve_tax_layers
    from services.finance.timeline import Timeline

    cod = date(2031, 1, 1)
    pack = load_pack("us_federal", as_of=date(2026, 1, 1))
    fin = FinanceInputs(financial_close=cod, tax_losses="carryforward", state_rate=0.0,
                        small_business_163j=False, acquisition_date=date(2030, 6, 1),
                        depreciation_class_by_asset={"bess": "macrs_7"})
    res = resolve_tax_layers(pack, fin, [OwnerAsset("bess", "battery", 1_000_000.0)], cod)
    _step("F3: one federal layer, nothing missing",
          res.missing == [] and [ly.name for ly in res.layers] == ["federal"], f"{res.missing}")
    if res.missing:
        return
    ebitda = [500e3, 600e3, 700e3, 800e3, 900e3]
    interest = [300e3, 300e3, 100e3, 100e3, 100e3]
    t = compute_tax(Timeline(y0=2031, cod_year=2031, base_year=2031, analysis_years=5),
                    res.layers, ebitda=np.array(ebitda), basis=1_000_000.0,
                    interest=np.array(interest), losses="carryforward")
    dep = [1_000_000.0, 0.0, 0.0, 0.0, 0.0]        # acquired after 2025-01-19: 100 % bonus
    allowed, carried = [], 0.0
    for e, i in zip(ebitda, interest):
        a = min(i + carried, 0.30 * e)
        carried = i + carried - a
        allowed.append(a)
    base = [e - d - a for e, d, a in zip(ebitda, dep, allowed)]
    taxable = _pool(base, 0.0, 0.8)
    _step("F3: depreciation 1 M$ in year 1 (bonus)",
          all(_cent(float(a), b) for a, b in zip(t.depreciation["federal"], dep)))
    _step("F3: taxable = EBITDA − depreciation − interest capped at 30 % with carryforward",
          all(_cent(float(a), b) for a, b in zip(t.taxable["federal"], base)),
          f"{[round(float(x)) for x in t.taxable['federal']]} vs {[round(x) for x in base]}")
    _step("F3: 21 % after the NOL at 80 % of taxable income",
          all(_cent(float(a), 0.21 * b) for a, b in zip(t.liability["federal"], taxable)),
          f"{[round(float(x), 2) for x in t.liability['federal']]} vs "
          f"{[round(0.21 * x, 2) for x in taxable]}")
    _step("F3: flagged interest_capped:federal", "interest_capped:federal" in t.flags,
          f"{t.flags}")


def scenario_f() -> None:
    print("\n[F] F4 by hand: escalation, indexation and tenor; base year before COD")
    from services.finance.case import CONTRACT_CLASS, TemplateLine
    from services.finance.cashflow import build_operating
    from services.finance.timeline import build_timeline
    from tests.test_finance_cashflow import _case

    ppa = TemplateLine("ppa", "ppa_settlement", 50_000.0, CONTRACT_CLASS, indexation=0.025,
                       tenor_years=8, contract_id="p1")
    bill = TemplateLine("bill", "energy_import", -20_000.0, "tariff")
    exp = TemplateLine("exp", "energy_export", 10_000.0, "export", degrades_with="pv")
    case = _case([ppa, bill, exp])
    tl = build_timeline(case)
    op = build_operating(case, tl)
    _step("F4: y0 2031, COD 2033, construction 2031–2032, base year 2031",
          (tl.y0, tl.cod_year, list(tl.construction_years), case.base_year) ==
          (2031, 2033, [2031, 2032], 2031))
    bad = []
    for i, y in enumerate(tl.years):
        k = y - 2033 + 1
        want = ((0.0, 0.0, 0.0) if y < 2033 else
                (50_000.0 * 1.025 ** (y - 2031) if k <= 8 else 0.0,
                 -20_000.0 * 1.03 ** (y - 2031),
                 10_000.0 * 1.01 ** (y - 2031) * 0.99 ** (k - 1)))
        got = (float(op.lines["ppa"][i]), float(op.lines["bill"][i]), float(op.lines["exp"][i]))
        if any(abs(a - b) > 1e-6 for a, b in zip(got, want)):
            bad.append((y, got, want))
    _step("F4: PPA on its own 2.5 % from the base year, stopping after 8 operating years; "
          "tariff at 3 %; export at 1 % degrading 1 %/yr — every year", not bad, f"{bad[:2]}")
    _step("F4: the tenor end is disclosed (contract_ends:p1:2040)",
          "contract_ends:p1:2040" in op.flags, f"{op.flags}")
    _step("F4: capex 1 M × 1.1 contingency, 60/40",
          _cent(float(op.capex[0]), 660_000.0) and _cent(float(op.capex[1]), 440_000.0))


def scenario_g() -> None:
    print("\n[G] F5 by hand: terminal values and replacement")
    from models.finance import TerminalValueRule
    from services.finance.case import TemplateLine
    from services.finance.cashflow import build_operating
    from services.finance.timeline import build_timeline
    from tests.test_finance_cashflow import _case

    om = TemplateLine("om", "fom", -10_000.0, "opex")
    rev = TemplateLine("rev", "energy_export", 100_000.0, "export")
    fixed = _case([om, rev], fin_over={"terminal_value": TerminalValueRule(method="fixed",
                                                                           value=50_000.0),
                                       "replacement_capex": [(2037, "pv", 200_000.0)]})
    tl = build_timeline(fixed)
    op = build_operating(fixed, tl)
    last = tl.years[-1]
    # Neither line names an asset: no degradation (C5), class escalation only.
    net_last = 100_000.0 * 1.01 ** (last - 2031) - 10_000.0 * 1.02 ** (last - 2031)
    _step("F5: fixed terminal 50 k in the last year only, not escalated, inside EBITDA",
          float(op.terminal[-1]) == 50_000.0 and float(sum(op.terminal[:-1])) == 0.0
          and _cent(float(op.ebitda[-1]), net_last + 50_000.0),
          f"EBITDA {float(op.ebitda[-1]):.2f} vs {net_last + 50_000.0:.2f}")
    _step("F5: replacement 200 k in 2037 escalated by capex 2 % from the base year",
          _cent(float(op.replacement[tl.index(2037)]), 200_000.0 * 1.02 ** (2037 - 2031)),
          f"{float(op.replacement[tl.index(2037)]):.2f}")
    mult = _case([om, rev], fin_over={"terminal_value": TerminalValueRule(
        method="multiple_of_ebitda", value=3.0)})
    op2 = build_operating(mult, build_timeline(mult))
    _step("F5: EBITDA-multiple terminal = 3 × the last year's operating net",
          _cent(float(op2.terminal[-1]), 3.0 * net_last), f"{float(op2.terminal[-1]):.2f}")
    book = _case([om, rev], fin_over={"terminal_value": TerminalValueRule(method="book_value")})
    op3 = build_operating(book, build_timeline(book))
    _step("F5: book value waits for the tax basis (not established in the operating pass)",
          op3.status["terminal"] == "not_established" and op3.ebitda is None)


def scenario_m() -> None:
    print("\n[M] IC S0b by hand: part replacements and the remaining-life terminal value")
    from models.finance import FinanceInputs
    from services.finance.case import (
        AssetFinance, AssetPart, FinanceCase, FinanceRefused, LpBasis, Template, TemplateLine,
        scale_capex,
    )
    from services.finance.engine import run_case
    from services.finance.tax import DepreciationClass, TaxLayer, sl_half_year

    def crf(r, life):
        return 1.0 / life if r == 0 else r / (1.0 - (1.0 + r) ** -life)

    def apf(r, y):
        return 0.0 if y <= 0 else (float(y) if r == 0 else (1.0 - (1.0 + r) ** -y) / r)

    no_tax = (TaxLayer(name="corp", rate=0.0,
                       depreciation=(DepreciationClass("all", 1.0, sl_half_year(10)),)),)
    bess = AssetFinance("bess", "StorageUnit", 1_600_000.0, 15.0, "battery",
                        (AssetPart("power", 400_000.0, 10.0, None),
                         AssetPart("energy", 1_200_000.0, 15.0, None)))
    saving = 300_000.0

    def case(capex_esc: float, **over) -> FinanceCase:
        kw = dict(financial_close=date(2029, 1, 1), cod_by_asset={}, analysis_years=25,
                  contingency_share=0.0, capex_phasing=[1.0],
                  escalation={"opex": 0.0, "tariff": 0.0, "capex": capex_esc},
                  tax_losses="offset_other_income", financing_fee_tax="not_deducted",
                  wacc_nominal=0.07, cost_of_equity=0.07, inflation=0.0,
                  replacement_rule="part_lifetimes",
                  terminal_value={"method": "remaining_life_annuity"})
        kw.update(over)
        return FinanceCase(
            inputs=FinanceInputs(**kw), owner="o", base_year=2030, cod=date(2030, 1, 1),
            templates=(Template(2030, (TemplateLine("bill", "energy_import", -700_000.0,
                                                    "tariff"),)),),
            counterfactual=(Template(2030, (TemplateLine("bill", "energy_import", -1_000_000.0,
                                                         "tariff"),)),),
            assets=(bess,), lp_basis=LpBasis(discount_rate=0.07))

    r = run_case(case(0.02), layers=no_tax)
    tl = r.tl
    want = {2039: 400_000.0 * 1.02 ** 9, 2044: 1_200_000.0 * 1.02 ** 14,
            2049: 400_000.0 * 1.02 ** 19}
    got = {int(y): float(r.op.replacement[tl.index(y)]) for y in tl.years
           if r.op.replacement[tl.index(y)] != 0.0}
    _step("M: part_lifetimes re-buys power in 2039 and 2049, energy in 2044 (the last service "
          "years), escalated by capex 2 % from 2030, no contingency",
          got.keys() == want.keys() and all(_cent(got[y], v) for y, v in want.items()), f"{got}")
    tv_power = want[2049] * crf(0.07, 10) * apf(0.07, 5)          # serves 2050–2059
    tv_energy = want[2044] * crf(0.07, 15) * apf(0.07, 5)         # serves 2045–2059
    _step("M: remaining_life_annuity TV = Σ base × crf(7 %, L) × af(7 %, 5) at 2054, inside EBITDA",
          r.op.terminal is not None and _cent(float(r.op.terminal[-1]), tv_power + tv_energy)
          and float(sum(r.op.terminal[:-1])) == 0.0,
          f"{None if r.op.terminal is None else float(r.op.terminal[-1]):.2f} vs "
          f"{tv_power + tv_energy:.2f}")
    flat = run_case(case(0.0), layers=no_tax)
    pv_costs = sum((float(flat.op.capex[i]) + float(flat.op.replacement[i])
                    - float(flat.op.terminal[i])) / 1.07 ** i for i in range(flat.tl.n))
    annuity = 400_000.0 * crf(0.07, 10) + 1_200_000.0 * crf(0.07, 15)
    _step("M: the identity — PV(capex + replacements − TV) = PV of the LP annuities 2,199,084.18",
          abs(pv_costs - annuity * apf(0.07, 25)) < CENT
          and abs(annuity * apf(0.07, 25) - 2_199_084.18) < CENT, f"{pv_costs:.4f}")
    npv_want = (saving - annuity) * apf(0.07, 25)
    _step("M: the project NPV = (LP saving − annuity) × af(7 %, 25) (GS BY_CONSTRUCTION)",
          _rel(flat.metrics["project_pre_tax_npv"], npv_want, 1e-9),
          f"{flat.metrics['project_pre_tax_npv']} vs {npv_want}")
    s = scale_capex(case(0.0), 1.1).assets[0]
    _step("M: scale_capex moves every part and the total together",
          _cent(s.overnight_cost, 1_760_000.0)
          and [round(p.overnight_cost, 2) for p in s.parts] == [440_000.0, 1_320_000.0])
    try:
        run_case(case(0.0, replacement_capex=[(2040, "bess", 1.0)]), layers=no_tax)
        refused = None
    except FinanceRefused as exc:
        refused = exc.code
    _step("M: a fixed entry for a part_lifetimes asset is refused (double counting)",
          refused == "replacement_rule_conflict:bess", f"{refused}")


def scenario_n() -> None:
    print("\n[N] IC G1, G2 by hand: campus equipment in the pack and as extra owner assets")
    from models.finance import FinanceInputs
    from services.asset_schema.access import UpfrontPart
    from services.finance.case import (
        AssetFinance, ExtraOwnerAsset, FinanceCase, LpBasis, Template, TemplateLine,
    )
    from services.finance.engine import run_case
    from services.library.defaults_pack.loader import load_defaults_pack
    from services.results.finance_case import finance_case_hash

    new, old = load_defaults_pack("2026-10-07"), load_defaults_pack("2026-10-05")
    (tr,) = new.cost_parts("transformer.tr_132_33_40")
    _step("N: pack 2026-10-07 — transformer.tr_132_33_40 is one lump part, 1,800,000 EUR/unit, "
          "40 y, FOM 0.015, 2026 EUR, illustrative",
          (tr.part, tr.basis, tr.overnight.value, tr.overnight.unit, tr.lifetime.value,
           tr.fom_share.value, tr.overnight.currency_year, tr.overnight.illustrative) ==
          ("investment", "lump", 1_800_000.0, "EUR/unit", 40.0, 0.015, 2026, True))
    _step("N: pack 2026-10-07 keeps every 2026-10-05 row unchanged; 2026-10-05's hash is unchanged",
          [v.model_dump() for v in new.cost_values[:len(old.cost_values)]] ==
          [v.model_dump() for v in old.cost_values]
          and old.hash == "962a0184861c1c594fdf4029ed5e7d55e241ab1aadfad5d52412df76ad2a4f78")

    def ext(name, kind, cost, life, fom):
        return ExtraOwnerAsset(name=name, kind=kind, basis="lump", quantity=1,
                               parts=(UpfrontPart("investment", cost, life, fom),),
                               build_year=2030, source="qa campus study",
                               source_hash="0123456789abcdef")

    extras = (ext("trf", "transformer", 1_800_000.0, 40.0, 0.015),
              ext("cap", "capacitor_bank", 90_000.0, 25.0, 0.02))
    bess = AssetFinance("bess", "StorageUnit", 1_000_000.0, 40.0, "battery")
    fom = tuple(TemplateLine(f"extra_asset_fom:{e.name}", "fom",
                             -e.parts[0].fom_share * e.parts[0].upfront_per_unit, "opex",
                             source="extra_asset_fom", source_id=e.name, money_year=2030)
                for e in extras)

    def case(**over):
        kw = dict(financial_close=date(2029, 1, 1), cod_by_asset={}, analysis_years=30,
                  contingency_share=0.1, capex_phasing=[1.0],
                  escalation={"opex": 0.02, "tariff": 0.0, "capex": 0.02},
                  tax_losses="offset_other_income", financing_fee_tax="not_deducted",
                  wacc_nominal=0.07, cost_of_equity=0.07, inflation=0.0,
                  replacement_rule="part_lifetimes",
                  terminal_value={"method": "remaining_life_annuity"},
                  incentives=[{"kind": "grant", "rate": 0.3, "grant_tax_treatment": "taxable"}])
        kw.update(over)
        return FinanceCase(
            inputs=FinanceInputs(**kw), owner="o", base_year=2030, cod=date(2030, 1, 1),
            templates=(Template(2030, (TemplateLine("bill", "energy_import", 500_000.0,
                                                    "tariff"),) + fom),),
            assets=(bess,) + tuple(e.asset_finance() for e in extras), extra_assets=extras,
            lp_basis=LpBasis(discount_rate=0.07))

    r = run_case(case())
    tl = r.tl
    _step("N: capex = (1 MEUR battery + 1.8 MEUR transformer + 90 k bank) × 1.1 contingency",
          r.op.capex is not None and _cent(float(r.op.capex.sum()), 2_890_000.0 * 1.1),
          f"{None if r.op.capex is None else float(r.op.capex.sum())}")
    _step("N: the transformer's FOM = 0.015 × 1.8 M a year, escalated by opex 2 % from 2030",
          all(_cent(float(r.op.lines["extra_asset_fom:trf"][tl.index(y)]),
                    -27_000.0 * 1.02 ** (y - 2030)) for y in (2030, 2045, 2059)))
    items = [(x.year, v) for x, v in r.op.replacement_items if x.asset == "cap"]
    _step("N: the 25-y bank is re-bought in 2054 (its last service year) at 90 k × 1.02^24",
          len(items) == 1 and items[0][0] == 2054 and _cent(items[0][1], 90_000.0 * 1.02 ** 24),
          f"{items}")
    crf = 0.07 / (1 - 1.07 ** -25)
    apf = (1 - 1.07 ** -20) / 0.07
    term = [t for t in r.op.terminal_terms if t.asset == "cap"]
    _step("N: its purchase serves 2055–2079: 20 years left, valued 90 k × 1.02^24 × crf × af(7 %, 20)",
          len(term) == 1 and term[0].remaining_years == 20.0
          and _cent(term[0].value, 90_000.0 * 1.02 ** 24 * crf * apf),
          f"{[(t.remaining_years, t.value) for t in term]}")
    _step("N: a grant on every asset excludes the extras (disclosed); its basis is the battery's",
          r.incentives.established() and "incentive_excludes_extra_asset:trf" in r.flags
          and r.incentives.lines[0].assets == ("bess",)
          and _cent(float(r.incentives.grant.sum()), 0.3 * 1_000_000.0 * 1.1),
          f"{float(r.incentives.grant.sum())}")
    plain = dataclasses.replace(case(), assets=(bess,), extra_assets=())
    _step("N: the case hash covers the extras and omits an empty extra_assets (pre-G2 hash shape)",
          finance_case_hash(case()) != finance_case_hash(plain)
          and "extra_assets" not in _canon_keys(plain) and "extra_assets" in _canon_keys(case()))


def _canon_keys(case) -> list[str]:
    from services.results.finance_case import _canon

    return sorted(_canon(case))


# ── through the routes ─────────────────────────────────────────────────────


def _export_price_ref(client, tag: str, n) -> dict | None:
    price = n.links_t["ic_export_price"]["export"]
    r = client.post("/api/library/series", json={
        "name": f"qa_ic_px_{tag}", "timestamps": [t.isoformat() for t in n.snapshots],
        "values": [float(v) for v in price], "meta": {"source": "qa driver"}})
    return r.json() if r.status_code in (200, 201) else None


def _setup(label: str, tag: str, n, commercial: dict, fin: dict | None):
    """Install → commercial (config route) → value flows (If-Match) → finance
    (If-Match) → solve on the session → `/results/value_flows`. Returns the
    session context or None (already a FAIL)."""
    client = qa_support.client()
    qa_support.install_network(n, name=_project(tag))
    commercial = copy.deepcopy(commercial)
    vf = commercial.pop("value_flows")
    if commercial.get("export_link") and "ic_export_price" in n.links_t:
        ref = _export_price_ref(client, tag, n)
        _step(f"{label}: export price stored in the Library", ref is not None)
        commercial["export_price_ref"] = ref
    r = client.put(CFG_URL, json={"commercial": commercial})
    _step(f"{label}: config route accepts the commercial config", r.status_code == 200, _why(r))
    tag_ = client.get(VF_URL).json()["digest"]
    r = client.put(VF_URL, json={"value_flows": vf}, headers={"If-Match": tag_})
    _step(f"{label}: value-flows route stores the config (If-Match)", r.status_code == 200,
          _why(r))
    if fin is not None:
        d0 = client.get(FIN_URL).json()["digest"]
        r = client.put(FIN_URL, json={"finance": fin}, headers={"If-Match": d0})
        _step(f"{label}: finance route stores the inputs (If-Match)",
              r.status_code == 200 and r.json()["status"] == "ok", _why(r))
    ctx = qa_support.session_context()
    cfg = ctx.solver_state["solver_config"]
    status, cond = run_simulation(cfg, ctx.network, ctx.mutation_lock, threading.Event(),
                                  queue.SimpleQueue(), state_update=lambda **kw: None)
    _step(f"{label}: solve is optimal", status in ("ok", "optimal"), f"{status}/{cond}")
    if status not in ("ok", "optimal"):
        return None
    r = client.get("/api/results/value_flows")
    body = r.json() if r.status_code == 200 else {}
    _step(f"{label}: /results/value_flows closes (status ok, conservation ok)",
          body.get("status") == "ok" and body.get("conservation_ok") is True,
          _why(r) or f"{body.get('status')} {body.get('conservation_ok')} {body.get('flags')}")
    STASH[f"vf:{tag}"] = body
    return ctx


def _run_study(label: str, *, layers=None, timeout: float = 120.0) -> dict:
    """POST the study (optionally through the runner's test-only `layers` hook)
    and poll it to the end."""
    import services.finance.investment_case_runner as RUN

    client = qa_support.client()
    orig = RUN.start_investment_case
    if layers is not None:
        def hooked(*a, **kw):
            kw.setdefault("layers", layers)
            return orig(*a, **kw)
        RUN.start_investment_case = hooked
    try:
        r = client.post(STUDY_URL, json={})
    finally:
        RUN.start_investment_case = orig
    _step(f"{label}: POST /results/investment_case starts the run",
          r.status_code == 200 and r.json().get("status") == "running", _why(r))
    deadline, body = time.time() + timeout, {}
    while time.time() < deadline:
        g = client.get(STUDY_URL)
        if g.status_code == 200:
            body = g.json()
            if body.get("status") != "running":
                break
        time.sleep(0.05)
    return body


def _layer():
    from services.finance.tax import DepreciationClass, TaxLayer, sl_half_year

    return (TaxLayer(name="corp", rate=0.25,
                     depreciation=(DepreciationClass("all", 1.0, sl_half_year(10)),)),)


def _fin_json(**over) -> dict:
    from tests.test_finance_case_adapter import _fin

    return _fin(**over).model_dump(mode="json")


# ── H: F6 ──────────────────────────────────────────────────────────────────


def _f6_hand(*, levy: float = 0.0, hp_mw: float = 0.0) -> dict:
    """The F6 toy site by hand, hour by hour over 2030: load 1 MW (+0.5 MW
    10–14 h) + an always-on heat pump's electricity, PV 2 MW 10–14 h, TOU 60 /
    180 $/MWh (night 0–6 h), 9 $/kW-month demand, supply 50 $/MWh, a per-kWh
    levy. The PV serves the site first (no export Link)."""
    t0 = datetime(2030, 1, 1)
    cf_e = act_e = cf_c = act_c = cf_l = act_l = gen = 0.0
    cf_pk: dict[int, float] = {}
    act_pk: dict[int, float] = {}
    for h in range(8760):
        ts = t0 + timedelta(hours=h)
        mid = 10 <= ts.hour < 14
        load = 1.0 + 0.5 * mid + hp_mw
        pv = 2.0 if mid else 0.0
        used = min(pv, load)
        imp = load - used
        rate = 60.0 if ts.hour < 6 else 180.0
        cf_e += load * rate
        act_e += imp * rate
        cf_c += 50.0 * load
        act_c += 50.0 * imp
        cf_l += load * 1000.0 * levy
        act_l += imp * 1000.0 * levy
        gen += used
        cf_pk[ts.month] = max(cf_pk.get(ts.month, 0.0), load * 1000.0)
        act_pk[ts.month] = max(act_pk.get(ts.month, 0.0), imp * 1000.0)
    cf_d = 9.0 * sum(cf_pk.values())
    act_d = 9.0 * sum(act_pk.values())
    inc = (cf_e + cf_d + cf_c + cf_l) - (act_e + act_d + act_c + act_l)
    return {"cf_energy": cf_e, "act_energy": act_e, "cf_demand": cf_d, "act_demand": act_d,
            "cf_commodity": cf_c, "act_commodity": act_c, "incremental": inc,
            "S": (cf_e + cf_c + cf_l) - (act_e + act_c + act_l), "gen": gen}


def _f6_lcoe(h: dict, fin: dict, capex: float) -> float:
    """LCOE = (PV value − PV post-tax equity) / PV energy at the cost of
    equity (SAM's definition generalised, WP4.5): the PV has no own costs, so
    value − equity = capex + tax; tax 25 % on (incremental − SL-10
    half-year depreciation); escalation 2 %, degradation 0.5 % on S only."""
    r, esc, dg = fin["cost_of_equity"], 0.02, 0.005
    sl = [0.05] + [0.10] * 9 + [0.05]
    pv_tax = pv_e = 0.0
    for k in range(1, fin["analysis_years"] + 1):
        f = (1 - dg) ** (k - 1)
        net = (h["incremental"] - h["S"] + h["S"] * f) * (1 + esc) ** (k - 1)
        dep = capex * sl[k - 1] if k - 1 < len(sl) else 0.0
        pv_tax += 0.25 * (net - dep) / (1 + r) ** k
        pv_e += h["gen"] * f / (1 + r) ** k
    return (capex + pv_tax) / pv_e


def _f6_case(label: str, tag: str, *, items=None, hp: bool = False):
    import routers.results as R
    from services.results.finance_case import build_finance_case
    from tests.test_finance_case_adapter import COD, F6_VF, _f6_network, _fin
    from tests.test_value_flow_reconciliation import DEMAND, TOU

    n = _f6_network()
    if hp:
        n.add("Bus", "heat", carrier="heat")
        n.add("Link", "hp", bus0="site", bus1="heat", p_nom=1.0, efficiency=3.0)
        n.add("Load", "heat_load", bus="heat", p_set=3.0)
    com = {"poc_link": "import", "value_flows": copy.deepcopy(F6_VF),
           "import_tariff": {"id": "t", "name": "t", "jurisdiction": "US",
                             "valid_from": "2029-01-01", "items": items or [TOU, DEMAND]}}
    fin = _fin(cod_by_asset={"pv": COD})
    ctx = _setup(label, tag, n, com, fin.model_dump(mode="json"))
    if ctx is None:
        return None, fin
    with _bound() as c:
        cfg = c.solver_state["solver_config"]
        case = build_finance_case(c.network, cfg, fin, result_df=R._result_df)
    return case, fin


def scenario_h() -> None:
    print("\n[H] F6 by hand: the counterfactual on the toy site (config routes → adapter)")
    from services.finance.engine import run_case
    from tests.test_value_flow_reconciliation import DEMAND, LEVY, TOU

    h = _f6_hand()
    _step("H: the hand working gives incremental 557,700 and S 503,700",
          abs(h["incremental"] - 557_700.0) < 1e-6 and abs(h["S"] - 503_700.0) < 1e-6,
          f"{h['incremental']:,.2f} / {h['S']:,.2f}")
    case, fin = _f6_case("H", "f6")
    if case is None:
        return
    (t,), (cf,) = case.templates, case.counterfactual
    act = {ln.key: ln for ln in t.lines}
    cfl = {ln.key: ln for ln in cf.lines}

    def amt(d, k):
        return None if k not in d or d[k].amount is None else d[k].amount

    _step("H: bill energy, actual −2,880 $/day and counterfactual −3,960 $/day (by hand)",
          _cent(amt(act, "bill:energy"), -h["act_energy"])
          and _cent(amt(cfl, "bill:energy"), -h["cf_energy"]),
          f"{amt(act, 'bill:energy')} / {amt(cfl, 'bill:energy')}")
    _step("H: demand charge 1,000 vs 1,500 kW × 9 $ × 12 months",
          _cent(amt(act, "bill:demand"), -h["act_demand"])
          and _cent(amt(cfl, "bill:demand"), -h["cf_demand"]),
          f"{amt(act, 'bill:demand')} / {amt(cfl, 'bill:demand')}")
    comm_act = [ln for ln in t.lines if ln.source == "asset" and ln.amount is not None
                and "grid_supply" in (ln.source_id or "")]
    comm_cf = [ln for ln in cf.lines if ln.source == "counterfactual"
               and ln.source_id == "commodity" and ln.amount is not None]
    _step("H: commodity 50 $ × 20 MWh/day actual, × 26 MWh/day counterfactual",
          len(comm_act) == 1 and _cent(comm_act[0].amount, -h["act_commodity"])
          and comm_cf and _cent(sum(ln.amount for ln in comm_cf), -h["cf_commodity"]),
          f"{[(ln.key, ln.amount) for ln in comm_act]} / {[(ln.key, ln.amount) for ln in comm_cf]}")
    _step("H: the C5 term S = the energy-volume saving 503,700 (degrading with pv), paired",
          _cent(amt(act, "bill_degradation:pv"), h["S"])
          and _cent(amt(act, "bill_degradation_base:pv"), -h["S"])
          and act["bill_degradation:pv"].degrades_with == "pv",
          f"{amt(act, 'bill_degradation:pv')} / {amt(act, 'bill_degradation_base:pv')}")
    _step("H: PV energy 1.5 MW × 4 h × 365 = 2,190 MWh; overnight 2 M$",
          abs(t.energy_mwh.get("pv", 0.0) - h["gen"]) < 1e-6
          and _cent(case.assets[0].overnight_cost, 2_000_000.0),
          f"{t.energy_mwh} {case.assets[0].overnight_cost}")
    r = run_case(case, layers=_layer())
    net = r.op_incremental["net"]
    _step("H: year-1 incremental cash = 557,700 (COD year)",
          net is not None and _cent(float(net[1]), h["incremental"]),
          f"{None if net is None else float(net[1]):,.4f}")
    y2 = h["incremental"] * 1.02 - h["S"] * 0.005 * 1.02
    _step("H: year 2 = 557,700 × 1.02 − S × 0.5 % × 1.02 (only S degrades)",
          net is not None and _cent(float(net[2]), y2), f"{None if net is None else float(net[2]):,.4f} vs {y2:,.4f}")
    want = _f6_lcoe(h, fin.model_dump(mode="json"), 2_000_000.0)
    got = r.metrics["lcoe_nominal_per_mwh"]
    _step("H: LCOE = 176.18 $/MWh by hand [WP4.6a review B1]",
          got is not None and abs(got - want) <= 1e-6 * want and abs(want - 176.18) < 0.005,
          f"engine {got} vs hand {want:.4f}")

    # B7: a per-kWh levy on import is an energy-volume saving too.
    h7 = _f6_hand(levy=0.02)
    case7, _ = _f6_case("H/levy", "f6_levy", items=[TOU, DEMAND, LEVY])
    if case7 is not None:
        r7 = run_case(case7, layers=_layer())
        n7 = r7.op_incremental["net"]
        _step("H/levy: incremental = 557,700 + the levy saving 43,800 = 601,500",
              n7 is not None and _cent(float(n7[1]), h7["incremental"])
              and abs(h7["incremental"] - 601_500.0) < 1e-6,
              f"{None if n7 is None else float(n7[1]):,.2f} vs {h7['incremental']:,.2f}")
        a7 = {ln.key: ln.amount for ln in case7.templates[0].lines}
        _step("H/levy: S includes the per-kWh levy: 547,500 [WP4.6a review B7]",
              _cent(a7.get("bill_degradation:pv"), h7["S"]) and abs(h7["S"] - 547_500.0) < 1e-6,
              f"{a7.get('bill_degradation:pv')} vs {h7['S']:,.2f}")

    # B2(a): a non-owner heat pump's electricity is site load on both sides.
    hh = _f6_hand(hp_mw=1.0)
    case2, _ = _f6_case("H/heat pump", "f6_hp", hp=True)
    if case2 is not None:
        r2 = run_case(case2, layers=_layer())
        n2 = r2.op_incremental["net"]
        flagged = any(f.startswith("counterfactual_includes_conversion_load:")
                      for f in case2.flags)
        v = None if n2 is None else float(n2[1])
        # The adapter includes the draw (WP4.6a round 1, B2): the number, disclosed
        # (the "or not established" alternative dropped after round 2).
        _step("H/heat pump: incremental = 725,600 by hand, the draw disclosed "
              "[WP4.6a review B2]",
              v is not None and _cent(v, hh["incremental"])
              and abs(hh["incremental"] - 725_600.0) < 1e-6 and flagged,
              f"{v} vs {hh['incremental']:,.2f}; flags "
              f"{[f for f in case2.flags if 'counterfactual' in f]}")


# ── I: the integration fixture through the routes ─────────────────────────


def _bill_by_hand(meter_mw) -> float:
    """The integration tariff (TOU 0.06 / 0.18 $/kWh night 0–6 h, 9 $/kW-month
    demand, 150 $/month) on an hourly 2030 meter, by hand."""
    t0 = datetime(2030, 1, 1)
    energy, peaks = 0.0, {}
    for h, mw in enumerate(meter_mw):
        ts = t0 + timedelta(hours=h)
        kw = max(0.0, float(mw)) * 1000.0
        energy += kw * (0.06 if ts.hour < 6 else 0.18)
        peaks[ts.month] = max(peaks.get(ts.month, 0.0), kw)
    return energy + 9.0 * sum(peaks.values()) + 150.0 * 12


def _xlsx_checks(label: str, content: bytes, export: dict, full: dict) -> None:
    import openpyxl

    from services.finance.debt import CFADS_DEFINITION

    wb = openpyxl.load_workbook(io.BytesIO(content))
    _step(f"{label}: workbook sheets (About, Summary, the sections, CashflowLines)",
          wb.sheetnames[:2] == ["About", "Summary"] and "CashflowLines" in wb.sheetnames
          and "project" in wb.sheetnames, f"{wb.sheetnames}")
    summary = {r[0].value: r[1].value for r in wb["Summary"].iter_rows(min_row=2)}
    gates = ("wacc_vs_discount_rate_consistent", "conservation_ok")
    off = {k: (summary.get(k), v) for k, v in export.items()
           if not isinstance(v, (dict, list)) and k in summary and k not in gates
           and not _xl_eq(summary[k], v)}
    _step(f"{label}: Summary — every number equal, every None `not_established`",
          not off and len(summary) >= 10, f"{off}")
    off = {k: (summary.get(k), export.get(k)) for k in gates
           if k in summary and not _xl_eq(summary[k], export.get(k))}
    _step(f"{label}: Summary — the gate rows equal the export view (a known gate is never "
          "written not_established)", not off, f"{off}")
    proj = {r[0].value: [c.value for c in r[1:]] for r in wb["project"].iter_rows() if r[0].value}
    payload = full["sections"]["project"]["payload"]
    off = {}
    for k, v in payload.items():
        if isinstance(v, (int, float)) and not isinstance(v, bool) or v is None:
            got = (proj.get(k) or [None])[0]
            if not _xl_eq(got, v):
                off[k] = (got, v)
    eq = payload["cash"]["equity_post_tax"]
    row = proj.get("cash.equity_post_tax") or []
    eq_ok = eq is not None and len(row) >= len(eq) and all(_xl_eq(a, b) for a, b in zip(row, eq))
    _step(f"{label}: project sheet — scalars and the equity cash row read back equal",
          not off and eq_ok, f"{off} {row[:3]} vs {None if eq is None else eq[:3]}")
    rows = list(wb["CashflowLines"].iter_rows(values_only=True))
    lines = full["cashflow_lines"]
    bad = []
    for ln, rw in zip(lines, rows[1:]):
        pv = ln["provenance"]
        want = [ln["year"], ln["participant"], ln["counterparty"], ln["value_stream"],
                ln.get("tariff_item"), ln.get("asset"), ln["amount"], pv.get("source"),
                pv.get("mode"), pv.get("pack_hash"), pv.get("seed"), pv.get("source_id"),
                pv.get("contract_id"), pv.get("period")]
        if len(rw) < len(want) or not all(_xl_eq(a, b) for a, b in zip(rw, want)):
            bad.append((rw, want))
    _step(f"{label}: CashflowLines — {len(lines)} lines read back equal, None → not_established",
          len(rows) - 1 == len(lines) and not bad and len(lines) > 0, f"{bad[:1]}")
    about = {r[0]: r[1] for r in wb["About"].iter_rows(values_only=True) if r[0]}
    cfb = payload.get("counterfactual") or {}
    _step(f"{label}: About — the counterfactual row (its basis) and the CFADS definition",
          isinstance(about.get("Counterfactual"), str)
          and about["Counterfactual"].startswith(str(cfb.get("basis")))
          and about.get("CFADS") == CFADS_DEFINITION == payload.get("cfads_definition"),
          f"{about.get('Counterfactual')!r:.120} | {about.get('CFADS')!r:.80}")
    _step(f"{label}: About — the assumptions hash", about.get("Assumptions hash") ==
          full["assumptions_hash"])
    # P4 gate assessor 3c: the conservation gate comes from the adapter's P3
    # ledger check, so a case built from a ledger has it established.
    _step(f"{label}: About / Summary — value-flow conservation is the ledger's (True here), "
          "not a constant not_established",
          export.get("conservation_ok") is True and about.get("Value-flow conservation") is True,
          f"export {export.get('conservation_ok')!r}, About {about.get('Value-flow conservation')!r}")


def _xl_eq(got, want) -> bool:
    """A cell read back: None ↔ `not_established`; a number equal to the
    spreadsheet's precision (openpyxl writes 16 significant digits, so the
    17th — one ulp — may move: relative 1e-15); anything else exactly."""
    if want is None:
        return got == NE
    if isinstance(want, bool) or isinstance(got, bool):
        return got is want
    if isinstance(want, (int, float)):
        return isinstance(got, (int, float)) and abs(got - want) <= 1e-15 * max(1.0, abs(want))
    return got == want


def _reconciles(full: dict) -> tuple[bool, str]:
    """Each year's Σ cashflow lines = `cash.equity_post_tax`."""
    p = full["sections"]["project"]["payload"]
    eq = p["cash"]["equity_post_tax"]
    if eq is None:
        return False, "equity_post_tax is None"
    by_year: dict[int, float] = {}
    for ln in full["cashflow_lines"]:
        by_year[ln["year"]] = by_year.get(ln["year"], 0.0) + ln["amount"]
    off = [(y, by_year.get(y, 0.0), e) for y, e in zip(p["years"], eq)
           if abs(by_year.get(y, 0.0) - e) > CENT]
    return not off, f"{len(p['years'])} years; {off[:2]}"


def scenario_i() -> None:
    print("\n[I] The integration fixture through the routes (PUT → solve → value flows → "
          "study → report → xlsx)")
    import numpy as np

    from tests.test_finance_case_adapter import _hourly_year

    client = qa_support.client()
    n, com = _hourly_year()
    fin = _fin_json()
    ctx = _setup("I", "hourly", n, com, fin)
    if ctx is None:
        return
    r = client.put(FIN_URL, json={"finance": {**fin, "analysis_years": 10}},
                   headers={"If-Match": "0" * 32})
    _step("I: a stale If-Match is refused (412 finance_changed), nothing stored",
          r.status_code == 412 and r.json()["detail"]["code"] == "finance_changed"
          and client.get(FIN_URL).json()["finance"] == fin, _why(r) or str(r.status_code))
    n = ctx.network
    w = n.snapshot_weightings.objective.to_numpy(dtype=float)
    _step("I: the BESS cycles and the site exports (the identity has an export term)",
          float((w * n.storage_units_t.p_dispatch["bess"]).sum()) > 0
          and float((w * n.links_t.p0["export"]).sum()) > 0)

    body = _run_study("I", layers=_layer())
    _step("I: the run finishes `done` through every stage",
          body.get("status") == "done" and body.get("stages_done") ==
          ["build_case", "load_pack", "run_engine", "assemble", "store"],
          f"{body.get('status')} {body.get('error_code')} {body.get('error')}")
    _step("I: the stored report is present and current",
          (body.get("report") or {}).get("present") is True
          and body["report"].get("stale") is False, f"{body.get('report')}")
    rx = client.get(REPORT_URL)
    rf = client.get(REPORT_URL, params={"detail": "full"})
    _step("I: GET …/report and detail=full answer 200",
          rx.status_code == 200 and rf.status_code == 200, _why(rx) or _why(rf))
    if rf.status_code != 200:
        return
    export, full = rx.json(), rf.json()
    STASH["full"], STASH["fin"] = full, fin
    p = full["sections"]["project"]["payload"]
    irr = p.get("equity_post_tax_irr")
    _step("I: the post-tax equity IRR is finite and = the pin 0.36208 (≤ 1e-5)",
          irr is not None and math.isfinite(irr) and abs(irr - PIN_EQUITY_IRR) <= 1e-5,
          f"{irr}")
    _step("I: the export view carries the same headlines as the payload",
          export["project_irr_post_tax"] == p["project_post_tax_irr"]
          and export["lcoe_finance_consistent_eur_per_mwh"] == p["lcoe_nominal_per_mwh"]
          and set(export["completeness"]) == set(full["completeness"]))
    ok, msg = _reconciles(full)
    _step("I: each year's Σ cashflow lines = cash.equity_post_tax", ok, msg)
    cf = p.get("counterfactual") or {}
    _step("I: the counterfactual block is in the project payload (present, basis, sources, "
          "established)",
          cf.get("present") is True and cf.get("basis") and cf.get("n_lines", 0) > 0
          and any("commodity" in s for s in cf.get("sources", []))
          and any(s.startswith("bill:") for s in cf.get("sources", []))
          and cf.get("lines_not_established") == [] and p.get("has_counterfactual") is True,
          f"{cf}")
    _step("I: the counterfactual's lines are in the cashflow lines, negated",
          any(str(ln["provenance"]["source"]).startswith("counterfactual:")
              for ln in full["cashflow_lines"]))

    # Year 1's template = the P3 ledger's owner net (the routes' own ledger).
    cod = p["cod_year"]
    finance_src = {"capex", "replacement_capex", "terminal_value", "itc", "ptc", "grant", "debt",
                   "debt_draws", "dsra", "reserve_interest"}
    op_lines = [ln for ln in full["cashflow_lines"] if ln["year"] == cod
                and ln["provenance"]["source"] not in finance_src
                and not str(ln["provenance"]["source"]).startswith(("counterfactual:", "tax:"))]
    template = sum(ln["amount"] for ln in op_lines)
    from services.commercial.lp_bindings import same_party
    led = STASH.get("vf:hourly") or {}
    owner_net = 0.0
    for ln in (led.get("periods") or {}).get("_", {}).get("lines", []):
        if ln.get("basis", "cash") != "cash" or ln["amount"] is None:
            continue
        if same_party(ln["payee"], "site"):
            owner_net += ln["amount"]
        elif same_party(ln["payer"], "site"):
            owner_net -= ln["amount"]
    _step("I: COD-year operating lines = the /results/value_flows owner net (to the cent)",
          _cent(template, owner_net), f"{template:,.4f} vs {owner_net:,.4f}")

    # The year-1 incremental EBITDA identity, the bill rated by hand.
    load = n.loads_t.p_set["site_load"].to_numpy(dtype=float)
    imp = np.clip(n.links_t.p0["import"].to_numpy(dtype=float), 0.0, None)
    export_mw = n.links_t.p0["export"].to_numpy(dtype=float)
    price = n.links_t["ic_export_price"]["export"].to_numpy(dtype=float)
    disch = n.storage_units_t.p_dispatch["bess"].to_numpy(dtype=float)
    vom = 0.5 * float((w * disch).sum())
    identity = ((_bill_by_hand(load) + 60.0 * float((w * load).sum()))
                - (_bill_by_hand(imp) + 60.0 * float((w * imp).sum()))
                + float((w * export_mw * price).sum()) - vom)
    inc = p["incremental_net"]
    _step("I: year-1 incremental EBITDA = (cf bill + commodity) − (bill + commodity) + export − "
          "BESS vom, the bills by hand (to the cent)",
          inc is not None and _cent(inc[1], identity) and identity > 0,
          f"{None if inc is None else inc[1]:,.4f} vs {identity:,.4f}")

    # LCOE against its definition: the BESS vom is the only own cost (B1).
    coe = fin["cost_of_equity"]
    gen = float((w * np.clip(n.generators_t.p["pv"].to_numpy(dtype=float), 0, None)).sum())
    eq = p["cash"]["equity_post_tax"]
    num = de = 0.0
    for i, y in enumerate(p["years"]):
        k = y - cod + 1
        d = (1 + coe) ** i
        own = vom * 1.02 ** (k - 1) if k >= 1 else 0.0
        num += ((inc[i] or 0.0) + own - eq[i]) / d
        de += (gen * 0.995 ** (k - 1) if k >= 1 else 0.0) / d
    want = num / de
    got = p.get("lcoe_nominal_per_mwh")
    _step("I: LCOE = (PV(incremental + own costs) − PV(equity)) / PV(energy) ≈ 143.23 "
          "[WP4.6a review B1]",
          got is not None and abs(got - want) <= 1e-6 * abs(want), f"report {got} vs {want:.4f}")

    x = client.get(XLSX_URL)
    _step("I: GET …/export.xlsx is a workbook",
          x.status_code == 200 and "spreadsheetml" in x.headers.get("content-type", ""), _why(x))
    if x.status_code == 200:
        _xlsx_checks("I/xlsx", x.content, export, full)


def scenario_i_pack() -> None:
    print("\n[I/pack] The production path: the us_federal pack, no test hook")
    if "full" not in STASH:
        _step("I/pack: the integration run exists", False)
        return
    client = qa_support.client()
    fin = {**STASH["fin"], "tax_pack_id": "us_federal", "state_rate": 0.0,
           "acquisition_date": "2029-06-01", "small_business_163j": True,
           "depreciation_class_by_asset": {"pv": "macrs_5", "bess": "macrs_7"}}
    d0 = client.get(FIN_URL).json()["digest"]
    r = client.put(FIN_URL, json={"finance": fin}, headers={"If-Match": d0})
    _step("I/pack: finance route stores the pack inputs", r.status_code == 200, _why(r))
    body = _run_study("I/pack")
    _step("I/pack: the run finishes `done`", body.get("status") == "done",
          f"{body.get('status')} {body.get('error_code')} {body.get('error')}")
    full = client.get(REPORT_URL, params={"detail": "full"}).json()
    p = full["sections"]["project"]["payload"]
    irr = p.get("equity_post_tax_irr")
    _step("I/pack: tax established and the post-tax equity IRR finite",
          full["completeness"].get("tax") == "ok" and irr is not None and math.isfinite(irr),
          f"tax {full['completeness'].get('tax')} ({full['sections']['tax'].get('note')}); "
          f"IRR {irr}")
    ok, msg = _reconciles(full)
    _step("I/pack: each year's Σ cashflow lines = cash.equity_post_tax", ok, msg)
    _step("I/pack: the report names the pack and its hash",
          set(full.get("packs") or {}) == {"us_federal"})
    d0 = client.get(FIN_URL).json()["digest"]
    client.put(FIN_URL, json={"finance": STASH["fin"]}, headers={"If-Match": d0})


# ── J: staleness ───────────────────────────────────────────────────────────


def _report_state() -> dict:
    return (qa_support.client().get(STUDY_URL).json() or {}).get("report") or {}


def scenario_j() -> None:
    print("\n[J] Staleness after a finance edit and a solver-config edit")
    if "fin" not in STASH:
        _step("J: the integration run exists", False)
        return
    client = qa_support.client()
    fin = STASH["fin"]
    d0 = client.get(FIN_URL).json()["digest"]
    r = client.put(FIN_URL, json={"finance": {**fin, "analysis_years": 12}},
                   headers={"If-Match": d0})
    st = _report_state()
    _step("J: a finance edit → stale, changed == ['finance']",
          r.status_code == 200 and st.get("stale") is True and st.get("changed") == ["finance"],
          _why(r) or f"{st}")
    d1 = client.get(FIN_URL).json()["digest"]
    client.put(FIN_URL, json={"finance": fin}, headers={"If-Match": d1})
    st = _report_state()
    _step("J: reverting the finance inputs → current again", st.get("stale") is False, f"{st}")
    orig = client.get(CFG_URL).json()["discount_rate"]
    r = client.put(CFG_URL, json={"discount_rate": orig + 0.01})
    st = _report_state()
    _step("J: a solver-config edit (discount_rate) → stale, changed == ['solver_config']",
          r.status_code == 200 and st.get("stale") is True
          and st.get("changed") == ["solver_config"], _why(r) or f"{st}")
    client.put(CFG_URL, json={"discount_rate": orig})
    st = _report_state()
    _step("J: reverting discount_rate → current again", st.get("stale") is False, f"{st}")


# ── K: chat tools ──────────────────────────────────────────────────────────


def scenario_k() -> None:
    print("\n[K] Chat tools on the stored report")
    if "full" not in STASH:
        _step("K: the integration run exists", False)
        return
    import json

    import routers.results as R
    from services import chat_tools
    from tests.fixtures.investment_case.sam import sam_case as S

    full = STASH["full"]
    p = full["sections"]["project"]["payload"]
    with _bound():
        out = chat_tools.DISPATCHERS["get_investment_case"]()
        h = out.get("headlines") or {}
        _step("K: get_investment_case — ok, present, current, under the 4,000-char cap",
              out.get("status") == "ok" and out["report"]["present"] is True
              and out["report"]["stale"] is False and len(json.dumps(out, default=str)) < 4000,
              f"{out.get('report')} {len(json.dumps(out, default=str))}")
        _step("K: get_investment_case — the headline IRR is the payload's (6 dp)",
              isinstance(h.get("equity_post_tax_irr"), float)
              and abs(h["equity_post_tax_irr"] - p["equity_post_tax_irr"]) < 1e-6
              and out.get("cashflow_lines") == len(full["cashflow_lines"]),
              f"{h.get('equity_post_tax_irr')} vs {p['equity_post_tax_irr']}")
        ex = chat_tools.DISPATCHERS["explain_cashflow"]()
        eq = ex.get("equity_irr") or {}
        _step("K: explain_cashflow — reconciles True at the cost of equity",
              ex.get("status") == "ok" and eq.get("reconciles") is True
              and ex["rate"]["basis"] == "cost_of_equity"
              and abs(ex["rate"]["value"] - STASH["fin"]["cost_of_equity"]) < 1e-9,
              f"{ex.get('rate')} {eq.get('status')} {eq.get('max_yearly_residual')}")
        _step("K: explain_cashflow — the stream PVs sum to the post-tax equity NPV",
              eq.get("sum_of_stream_pvs") is not None and p["equity_post_tax_npv"] is not None
              and abs(eq["sum_of_stream_pvs"] - p["equity_post_tax_npv"]) < 0.05,
              f"{eq.get('sum_of_stream_pvs')} vs {p['equity_post_tax_npv']}")
        _step("K: explain_cashflow — the avoided supply cost is one stream",
              sum(1 for x in eq.get("by_stream", []) if x["stream"] == chat_tools._IC_AVOIDED) == 1,
              f"{[x['stream'] for x in eq.get('by_stream', [])]}")

        # solve_ppa_price on a solvable case: the S1b oracle through the router seam.
        st = qa_support.session_context().solver_state
        fin_before = copy.deepcopy(st["solver_config"].finance)
        rep_before = copy.deepcopy(st.get("investment_case_report"))
        case = S.to_finance_case("s1b", solve=True)
        sp = case.inputs.solve_ppa
        old_build, old_layers = R._ic_build_case, chat_tools._ic_solve_layers
        R._ic_build_case = lambda n, cfg, *, owner, lost_load: (lambda: case)
        chat_tools._ic_solve_layers = lambda c: S.sam_tax_layers("s1b")
        try:
            sol = chat_tools.DISPATCHERS["solve_ppa_price"](
                target_irr=sp.target_irr, target_year=sp.target_year)
        except Exception as exc:  # noqa: BLE001
            sol = {"status": f"{type(exc).__name__}: {getattr(exc, 'detail', exc)}"}
        finally:
            R._ic_build_case, chat_tools._ic_solve_layers = old_build, old_layers
        want = S.sam_expected("s1b")["scalars"]["ppa_price_per_mwh"]
        _step("K: solve_ppa_price — SAM's S1b price (4 dp) at the 11 % target in year 20",
              sol.get("status") == "ok" and abs(sol["solved_ppa_price_per_mwh"] - want) < 1e-4
              and abs(sol["equity_post_tax_irr_at_target_year"] - sp.target_irr) < 1e-6,
              f"{sol.get('solved_ppa_price_per_mwh')} vs {want}; {sol.get('status')}")
        _step("K: solve_ppa_price — nothing stored changed",
              sol.get("stored_inputs_changed") is False
              and st["solver_config"].finance == fin_before
              and st.get("investment_case_report") == rep_before
              and case.inputs.solve_ppa == sp)


# ── L: refusals and C12 ────────────────────────────────────────────────────


def _no_zero_for_none(label: str, export: dict, full: dict, keys) -> None:
    p = full["sections"]["project"]["payload"]
    _step(f"{label}: headlines None in the export view, never 0",
          all(export.get(k) is None for k in keys), f"{ {k: export.get(k) for k in keys} }")
    _step(f"{label}: payload returns None, never 0",
          all(p.get(k) is None for k in ("equity_post_tax_irr", "equity_post_tax_npv",
                                           "lifecycle_npv", "lcoe_nominal_per_mwh")),
          f"{ {k: p.get(k) for k in ('equity_post_tax_irr', 'lcoe_nominal_per_mwh')} }")


def scenario_l() -> None:
    print("\n[L] Refusals and C12 (unknown is None + a reason)")
    import json

    import openpyxl

    from services import chat_tools

    client = qa_support.client()
    keys = ("project_irr_pre_tax", "project_irr_post_tax", "npv_at_wacc",
            "lcoe_finance_consistent_eur_per_mwh")

    # L1: a None ledger line → operating not established (the real adapter).
    if "fin" in STASH:
        import routers.results as R
        import services.results.value_flows as VF
        from models.finance import FinanceInputs
        from services.finance.engine import run_case
        from services.finance.export_xlsx import build_workbook
        from services.finance.report import assemble_finance_sections
        from services.results.finance_case import build_finance_case

        orig = VF.value_flow_ledger
        nulled: list = []

        def corrupt(*a, **kw):
            got = orig(*a, **kw)
            if got is None:
                return got
            inputs, vf, ledger, res = got
            ledger = copy.deepcopy(ledger)
            for ln in ledger.periods["_"]:
                if ln.source == "bill" and ln.tariff_item == "energy" and ln.amount:
                    ln.amount = None
                    nulled.append(ln.source_id)
                    break
            return inputs, vf, ledger, res

        VF.value_flow_ledger = corrupt
        try:
            with _bound() as c:
                cfg = c.solver_state["solver_config"]
                case = build_finance_case(c.network, cfg, FinanceInputs.model_validate(STASH["fin"]),
                                          result_df=R._result_df)
        finally:
            VF.value_flow_ledger = orig
        r = run_case(case, layers=_layer())
        reasons = (r.op.reasons or {}).get("operating", [])
        _step("L: a None ledger line (bill energy) → a None template line → operating "
              "not_established, named",
              bool(nulled) and r.sections["operating"] == "not_established"
              and any("line_not_established:bill:energy" in x for x in reasons),
              f"{nulled} {r.sections} {reasons[:3]}")
        _step("L: …and every return built on it is None",
              r.cash["equity_post_tax"] is None and r.metrics["equity_post_tax_irr"] is None
              and r.metrics["lcoe_nominal_per_mwh"] is None and r.metrics["lifecycle_npv"] is None)
        rep = assemble_finance_sections(r, case, assumptions_hash="a" * 16, packs={})
        wb = openpyxl.load_workbook(io.BytesIO(build_workbook(rep)))
        summary = {x[0].value: x[1].value for x in wb["Summary"].iter_rows(min_row=2)}
        _step("L: …the report says not_established and the xlsx writes not_established, never 0",
              rep.completeness["project"] == "not_established"
              and all(summary.get(k) == NE for k in keys), f"{ {k: summary.get(k) for k in keys} }")
        # P4 gate assessor B2: the unknown line is CARRIED, never dropped — listed in
        # the payload, the retailer's total unknown (not the rest summed), and a
        # not_established row per year on the CashflowLines sheet.
        pl = rep.sections["project"].payload
        miss = [m for m in pl.get("lines_not_established") or [] if m["key"] == "bill:energy"]
        part = rep.sections["participants"].payload
        cp = next((m["counterparty"] for m in miss), None)
        ws = wb["CashflowLines"]
        head = [c.value for c in next(ws.iter_rows(max_row=1))]
        rows = [dict(zip(head, (c.value for c in r_))) for r_ in ws.iter_rows(min_row=2)]
        ne_rows = [x for x in rows if x["source_id"] == "bill:energy" and x["amount"] == NE]
        _step("L: …the unknown line is carried: listed in the payload with its years, its "
              "counterparty's total not established, a not_established row per year in the xlsx "
              "[P4 gate assessor B2]",
              len(miss) == 1 and len(miss[0]["years"]) == r.tl.analysis_years
              and cp is not None and part["operating_cash_by_counterparty"][cp] is None
              and cp in part["counterparties_not_established"]
              and len(ne_rows) == r.tl.analysis_years,
              f"{miss[:1]} cp={cp} rows={len(ne_rows)}")

        # L2: a missing escalation class, through the routes.
        fin = {**STASH["fin"], "escalation": {k: v for k, v in STASH["fin"]["escalation"].items()
                                               if k != "tariff"}}
        d0 = client.get(FIN_URL).json()["digest"]
        rr = client.put(FIN_URL, json={"finance": fin}, headers={"If-Match": d0})
        _step("L: the finance route accepts escalation without `tariff` (unknown until typed)",
              rr.status_code == 200, _why(rr))
        body = _run_study("L/escalation", layers=_layer())
        export = client.get(REPORT_URL).json()
        full = client.get(REPORT_URL, params={"detail": "full"}).json()
        note = full["sections"]["project"].get("note") or ""
        _step("L: a missing escalation class → project not_established, the reason names it",
              body.get("status") == "done" and full["completeness"]["project"] == "not_established"
              and "escalation_missing:tariff" in note, f"{body.get('status')} {note[:200]}")
        _no_zero_for_none("L/escalation", export, full, keys)
        with _bound():
            out = chat_tools.DISPATCHERS["get_investment_case"]()
            ex = chat_tools.DISPATCHERS["explain_cashflow"]()
        h = out.get("headlines") or {}
        _step("L/escalation: chat headlines read 'not established', never 0 or null",
              h.get("equity_post_tax_irr") == "not established"
              and h.get("lcoe_nominal_per_mwh") == "not established"
              and all(v is not None and v != 0 for k, v in h.items() if k != "debt_amount"),
              f"{h}")
        nums = json.dumps(ex.get("equity_irr"))
        _step("L/escalation: explain_cashflow says not_established for the IRR",
              (ex.get("equity_irr") or {}).get("status") == "not_established", nums[:200])
        x = client.get(XLSX_URL)
        wb = openpyxl.load_workbook(io.BytesIO(x.content))
        summary = {c[0].value: c[1].value for c in wb["Summary"].iter_rows(min_row=2)}
        _step("L/escalation: the xlsx writes not_established for every unknown headline",
              all(summary.get(k) == NE for k in keys), f"{ {k: summary.get(k) for k in keys} }")
        d1 = client.get(FIN_URL).json()["digest"]
        client.put(FIN_URL, json={"finance": STASH["fin"]}, headers={"If-Match": d1})

    # L3: the 7-day fixture through the routes.
    from tests.test_finance_case_adapter import OWNED, _edge7
    from tests.test_value_flow_reconciliation import _commercial

    n7 = _edge7()
    # Typed overnight costs and lifetimes (C6: never back-calculated from the
    # annuity), so the annualised case can be a number.
    for frame, name, cost in (("generators", "pv", 700_000.0), ("storage_units", "bess", 1e6)):
        df = getattr(n7, frame)
        df.loc[name, "overnight_cost"] = cost
        df.loc[name, "lifetime"] = 30.0
    ctx = _setup("L/7-day", "edge7", n7, _commercial(copy.deepcopy(OWNED), contracts=[]),
                 _fin_json())
    if ctx is None:
        return
    body = _run_study("L/7-day", layers=_layer())
    _step("L/7-day: refused template_not_annual:168 (a status, never a 500)",
          body.get("status") == "refused" and body.get("error_code") == "template_not_annual:168",
          f"{body.get('status')} {body.get('error_code')} {body.get('error')}")
    export = client.get(REPORT_URL).json()
    full = client.get(REPORT_URL, params={"detail": "full"}).json()
    _step("L/7-day: the refusal report — project not_established with the code",
          full["completeness"]["project"] == "not_established"
          and "template_not_annual:168" in (full["sections"]["project"].get("note") or ""),
          f"{full['sections']['project'].get('note')}")
    _step("L/7-day: every headline None, never 0",
          all(export.get(k) is None for k in keys + ("min_dscr", "llcr")),
          f"{ {k: export.get(k) for k in keys} }")
    x = client.get(XLSX_URL)
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    summary = {c[0].value: c[1].value for c in wb["Summary"].iter_rows(min_row=2)}
    _step("L/7-day: the xlsx writes not_established", all(summary.get(k) == NE for k in keys),
          f"{ {k: summary.get(k) for k in keys} }")
    d0 = client.get(FIN_URL).json()["digest"]
    client.put(FIN_URL, json={"finance": _fin_json(annualise=True)}, headers={"If-Match": d0})
    body = _run_study("L/7-day annualised", layers=_layer())
    full = client.get(REPORT_URL, params={"detail": "full"}).json()
    flags = full["sections"]["project"]["payload"].get("flags") or []
    irr = full["sections"]["project"]["payload"].get("equity_post_tax_irr")
    _step("L/7-day: with annualise — done, a number, flagged template_annualised:52.14",
          body.get("status") == "done" and "template_annualised:52.14" in flags
          and irr is not None and math.isfinite(irr),
          f"{body.get('status')} IRR {irr} {[f for f in flags if 'annualis' in f]}; "
          f"{(full['sections']['project'].get('note') or '')[:200]}")
    # The week is one January week: a monthly demand charge annualises by
    # 12 / (months present) = 12, not by 8760 / 168 (WP4.6a review B5).
    led = (STASH.get("vf:edge7") or {}).get("periods", {}).get("_", {}).get("lines", [])
    week = sum(ln["amount"] for ln in led if ln["source"] == "bill"
               and ln.get("tariff_item") == "demand" and ln["payer"] == "site"
               and ln["amount"] is not None)
    cod = full["sections"]["project"]["payload"].get("cod_year")
    got = sum(ln["amount"] for ln in full.get("cashflow_lines") or []
              if ln["year"] == cod and ln["provenance"]["source"] == "bill"
              and ln.get("tariff_item") == "demand")
    _step("L/7-day: annualised, the demand charge = 12 × the week's one monthly charge "
          "[WP4.6a review B5]", week > 0 and _cent(got, -12.0 * week),
          f"{got:,.2f} vs {-12.0 * week:,.2f} (× 8760/168 would be {-week * 8760 / 168:,.2f})")


def main() -> int:
    started = time.monotonic()
    for fn in (scenario_a, scenario_b, scenario_c, scenario_d, scenario_e, scenario_f,
               scenario_g, scenario_h, scenario_i, scenario_j, scenario_k, scenario_i_pack,
               scenario_l, scenario_m, scenario_n):
        try:
            fn()
        except Exception as exc:  # noqa: BLE001 — a crashed scenario is a failure
            import traceback
            traceback.print_exc()
            _step(f"{fn.__name__} ran without crashing", False, f"{type(exc).__name__}: {exc}")
    try:
        qa_support.delete_project(*PROJECTS)
    except Exception:  # noqa: BLE001
        pass
    print()
    print("=" * 60)
    print(f"Total: {PASS + FAIL}  Pass: {PASS}  Fail: {FAIL}  "
          f"({time.monotonic() - started:.0f} s)")
    print("=" * 60)
    return 0 if FAIL == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
