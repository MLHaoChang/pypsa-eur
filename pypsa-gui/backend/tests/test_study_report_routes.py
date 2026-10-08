"""
S7: the decision report routes on a run with the fake solver of
`tests/test_study_runner.py` (the queue, the forks, the campaign, the mesh
and the reads are real; only `run_simulation` is replaced) — plan S7
"Acceptance"; gates S1 (in-flight check per handler, the copied record,
payload-level `available`), S5 (one hash rule with the case's 409).
"""
from __future__ import annotations

import io
import shutil
import threading

import pytest

from tests.study_s4_support import create_pack_study, enable_studies, wait_run
from tests.test_study_runner import FakeSolver
from tests.test_study_tornado_routes import NO_PV_TOU, _run, wait_tornado

XSS = "<script>alert(1)</script>"


@pytest.fixture
def studies_on(monkeypatch):
    yield from enable_studies(monkeypatch)


@pytest.fixture
def fake(monkeypatch):
    from services import solver_service

    solver = FakeSolver()
    monkeypatch.setattr(solver_service, "run_simulation", solver)
    return solver


def _full(client, api_project, name, **kw):
    sid = _run(client, api_project, name, **kw)
    assert client.post(f"/api/projects/{name}/studies/{sid}/findings/tornado",
                       json={}).status_code == 202
    assert wait_tornado(client, name, sid)["status"] == "done"
    return sid


def _post(client, name, sid):
    r = client.post(f"/api/projects/{name}/studies/{sid}/report")
    assert r.status_code == 200, r.text
    return r.json()


# ── the acceptance flow ──────────────────────────────────────────────────

def test_a_report_of_a_run_renders_and_every_verdict_kpi_is_the_findings(
        client, api_project, studies_on, fake):
    from services.study import report as R

    name = "rep-flow"
    sid = _full(client, api_project, name)
    findings = client.get(f"/api/projects/{name}/studies/{sid}/findings").json()
    body = _post(client, name, sid)
    assert body["available"] is True and body["stale"] is False
    assert body["ai_paragraphs"] == []
    assert [s for s in body["sections"]] == [sid_ for sid_, _t in R.SECTIONS]
    assert body["sections"]["reliability"]["status"] == "skipped"
    v = findings["verdict"]
    assert v["status"] == "ok"
    for kpi in v["headline_kpis"]:
        assert body["facts"][kpi["key"]] == kpi
    # stored: GET answers the same report
    again = client.get(f"/api/projects/{name}/studies/{sid}/report").json()
    assert again["facts"] == body["facts"] and again["available"] is True
    study = client.get(f"/api/projects/{name}/studies/{sid}").json()
    assert study["report_ref"] == f"studies/{sid}.report.json"

    # economics carries both LCOS figures, named, and the duals disclosure;
    # U2 WP8: the run's forks are the engine's, so the case's LCOS is the
    # finance engine's (charging included)
    econ = body["sections"]["economics"]
    assert econ["status"] == "ok"
    assert econ["facts"]["case_npv"]["engine"] == "finance_engine"
    assert {"lcos_finance_engine", "lcos_incl_charging", "market_revenue_at_duals",
            "case_irr", "case_payback_discounted"} <= set(econ["facts"])
    # the golden registration (coverage.py `decision_report`): every economics
    # fact is the case route's KPI for the same option, copied
    oid = econ["payload"]["option_id"]
    case = client.get(f"/api/projects/{name}/studies/{sid}/options/{oid}/case").json()
    k = case["kpis"]
    for fact, value in (("case_npv", k["npv"]), ("case_irr", k["irr"]),
                        ("case_payback_simple", k["payback_simple"]),
                        ("case_payback_discounted", k["payback_discounted"]),
                        ("case_capex_total", k["capex_total"]),
                        ("lcos_finance_engine", k["levelised_cost"]),
                        ("market_revenue_at_duals",
                         case["market_revenue_at_duals"]["annual_value"])):
        assert econ["facts"][fact]["value"] == value, fact
    assert econ["facts"]["case_npv"]["currency_year"] == case["currency_year"]
    codes = [d["code"] for d in body["required_disclosures"]]
    for code in ("market_revenue_at_duals_zero_profit_at_optimum",
                 "npv_nonnegative_at_optimum_by_construction",
                 "irr_and_discounted_payback_bounded_at_optimum_by_construction",
                 "synthetic_load_understates_peak_shaving", "tariff_illustrative"):
        assert code in codes, code
    assert body["sections"]["drivers"]["payload"]["value_streams_basis"] == "baseline"

    html = client.get(f"/api/projects/{name}/studies/{sid}/report.html")
    assert html.status_code == 200
    assert html.headers["content-security-policy"] == "sandbox"
    assert html.headers["content-disposition"].startswith("attachment;")
    assert html.headers["content-type"].startswith("text/html")
    text = html.text
    disclosures_end = text.index("</section>", text.index('class="disclosures"'))
    for kpi in v["headline_kpis"]:
        from models.study import Figure

        shown = R.format_fact(Figure.model_validate(kpi))
        assert shown in text and text.index(shown) > disclosures_end, shown
    assert text.count("data:image/png;base64,") >= 3

    import docx

    d = client.get(f"/api/projects/{name}/studies/{sid}/report.docx")
    assert d.status_code == 200 and d.headers["content-security-policy"] == "sandbox"
    doc = docx.Document(io.BytesIO(d.content))
    assert any(t.rows[0].cells[0].text == "Assumption" for t in doc.tables)
    all_text = "\n".join([p.text for p in doc.paragraphs]
                         + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    for kpi in v["headline_kpis"]:
        from models.study import Figure

        assert R.format_fact(Figure.model_validate(kpi)) in all_text

    import openpyxl

    x = client.get(f"/api/projects/{name}/studies/{sid}/report.xlsx")
    assert x.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    assert {"Verdict", "Value streams", "Tornado", "Assumptions", "Provenance",
            "Cash flows bess_2h", "KPIs bess_2h"} <= set(wb.sheetnames)
    kpis = wb["KPIs bess_2h"]
    cells = {r[0].value: r[4].value for r in kpis.iter_rows() if r[0].value}
    assert cells["npv"].startswith("live formula")
    assert cells["irr"].startswith("static") and cells["payback_discounted"].startswith("static")


def test_a_study_named_script_renders_escaped(client, api_project, studies_on, fake):
    api_project("rep-xss-src")
    r = create_pack_study(client, "rep-xss-src", "rep-xss", intake=NO_PV_TOU, name=XSS)
    assert r.status_code == 201, r.text
    sid = r.json()["study_id"]
    assert client.post(f"/api/projects/rep-xss/studies/{sid}/run", json={}).status_code == 202
    assert wait_run(client, "rep-xss", sid)["status"] == "done"
    _post(client, "rep-xss", sid)
    html = client.get(f"/api/projects/rep-xss/studies/{sid}/report.html").text
    assert XSS not in html and "&lt;script&gt;" in html


# ── stale: one hash rule with the case route's 409 ───────────────────────

def test_a_ledger_edit_after_assembly_sets_stale(client, api_project, studies_on, fake):
    name = "rep-led"
    sid = _run(client, api_project, name)
    assert _post(client, name, sid)["stale"] is False
    r = client.put(f"/api/projects/{name}/studies/{sid}/ledger", json={"rows": [
        {"key": "battery_storage_eur_per_kwh", "value": 150.0, "unit": "EUR/kWh"}]})
    assert r.status_code == 200, r.text
    body = client.get(f"/api/projects/{name}/studies/{sid}/report").json()
    assert body["stale"] is True
    assert body["stale_reasons"] == ["ledger_changed_since_findings"]
    # the case route refuses on the same rule
    case = client.get(f"/api/projects/{name}/studies/{sid}/options/bess_2h/case")
    assert case.status_code == 409
    html = client.get(f"/api/projects/{name}/studies/{sid}/report.html").text
    assert "This report is stale" in html


def test_an_option_fork_edit_after_assembly_sets_stale(
        client, api_project, studies_on, fake, project_storage_dir):
    import pypsa

    name = "rep-fork"
    sid = _run(client, api_project, name)
    assert _post(client, name, sid)["stale"] is False
    path = project_storage_dir(f"{name}-opt-bess_2h") / "network.nc"
    n = pypsa.Network(str(path))
    n.storage_units.loc["battery", "p_nom_opt"] = 5.0
    n.export_to_netcdf(str(path))
    body = client.get(f"/api/projects/{name}/studies/{sid}/report").json()
    assert body["stale"] is True
    assert body["stale_reasons"] == ["fork_changed_since_findings:bess_2h"]
    case = client.get(f"/api/projects/{name}/studies/{sid}/options/bess_2h/case")
    assert case.status_code == 409 and case.json()["detail"]["error_kind"] == "fork_changed_since_run"


def test_a_copied_records_report_is_stale_when_adopted(
        client, api_project, studies_on, fake, project_storage_dir):
    name = "rep-copy"
    sid = _run(client, api_project, name)
    _post(client, name, sid)
    other = api_project("rep-copy-dest")
    shutil.copytree(project_storage_dir(name) / "studies",
                    project_storage_dir(other) / "studies", dirs_exist_ok=True)
    body = client.get(f"/api/projects/{other}/studies/{sid}/report").json()
    assert body["stale"] is True
    assert "copied_record_findings_computed_on_the_origin_forks" in body["stale_reasons"]
    assert any(r.startswith("fork_changed_since_findings") for r in body["stale_reasons"])


# ── refusals, access, in-flight ──────────────────────────────────────────

def test_report_refusals_are_typed(client, api_project, studies_on, fake, other_org_client):
    api_project("rep-404-src")
    sid = create_pack_study(client, "rep-404-src", "rep-404",
                            intake=NO_PV_TOU).json()["study_id"]
    r = client.post(f"/api/projects/rep-404/studies/{sid}/report")
    assert r.status_code == 404 and r.json()["detail"]["error_kind"] == "study_never_run"
    for url in ("report", "report.html", "report.docx", "report.xlsx"):
        r = client.get(f"/api/projects/rep-404/studies/{sid}/{url}")
        assert r.status_code == 404, url
        assert r.json()["detail"]["error_kind"] == "report_never_assembled"
    for method, url in (("post", "report"), ("get", "report"), ("get", "report.html"),
                        ("get", "report.docx"), ("get", "report.xlsx")):
        r = getattr(other_org_client, method)(f"/api/projects/rep-404/studies/{sid}/{url}")
        assert r.status_code == 404, (url, r.status_code)
    assert fake.calls == []


def test_a_ledger_edited_after_the_run_refuses_the_assembly(client, api_project, studies_on, fake):
    name = "rep-refuse"
    sid = _run(client, api_project, name)
    client.put(f"/api/projects/{name}/studies/{sid}/ledger", json={"rows": [
        {"key": "battery_storage_eur_per_kwh", "value": 150.0, "unit": "EUR/kWh"}]})
    r = client.post(f"/api/projects/{name}/studies/{sid}/report")
    assert r.status_code == 409 and r.json()["detail"]["error_kind"] == "ledger_changed_since_run"


def test_an_unestablished_verdict_is_reported_not_refused(client, api_project, studies_on, fake):
    name = "rep-pre"
    sid = _run(client, api_project, name)          # no tornado: verdict not established
    body = _post(client, name, sid)
    assert body["available"] is False
    assert body["sections"]["executive_summary"]["status"] == "not_established"
    assert any("tornado" in g for g in body["evidence_gaps"])
    assert client.get(f"/api/projects/{name}/studies/{sid}/report.html").status_code == 200


def test_the_assembly_refuses_while_a_run_replaces_the_forks(
        client, api_project, studies_on, monkeypatch):
    """The POST reads the option networks: its own in-flight check (gate S1)."""
    from services import solver_service

    release = threading.Event()
    monkeypatch.setattr(solver_service, "run_simulation", FakeSolver())
    sid = _run(client, api_project, "rep-live")
    monkeypatch.setattr(solver_service, "run_simulation",
                        FakeSolver(on_call=lambda k: release.wait(30)))
    try:
        assert client.post(f"/api/projects/rep-live/studies/{sid}/run", json={}).status_code == 202
        r = client.post(f"/api/projects/rep-live/studies/{sid}/report")
        assert r.status_code == 409 and r.json()["detail"]["error_kind"] == "study_running"
    finally:
        client.post(f"/api/projects/rep-live/studies/{sid}/run/abort")
        release.set()
    wait_run(client, "rep-live", sid)


def test_every_report_handler_declares_access_and_refuses_when_disabled():
    import inspect

    from routers import studies as S
    from routers.deps import ProjectAccessDep

    handlers = (S.assemble_report, S.get_report, S.get_report_html, S.get_report_docx,
                S.get_report_xlsx)
    for fn in handlers:
        params = inspect.signature(fn).parameters
        assert params["project"].default is ProjectAccessDep, fn.__name__
        assert "_refuse_unless_enabled()" in inspect.getsource(fn), fn.__name__
    src = inspect.getsource(S.assemble_report)
    assert "_check_lock(" in src and "_refuse_while_a_run_runs(" in src


# ── gate S7 (BC-S7-1, BC-S7-3, S4 intake) ────────────────────────────────

class BoundSolver(FakeSolver):
    """The fake, sizing every battery AT its upper bound (a size limit binds)."""

    def __call__(self, config, n, *a, **k):
        out = super().__call__(config, n, *a, **k)
        if len(n.storage_units):
            n.storage_units["p_nom_opt"] = n.storage_units["p_nom_max"]
        return out


def _econ_text(body) -> str:
    return " ".join(p["text"] for p in body["sections"]["economics"]["prose"])


def test_zero_profit_is_stated_only_at_an_interior_optimum(
        client, api_project, studies_on, monkeypatch):
    """BC-S7-1: at a size limit the revenue at duals exceeds the annualised cost."""
    from services import solver_service

    monkeypatch.setattr(solver_service, "run_simulation", FakeSolver())
    sid = _full(client, api_project, "rep-int")
    body = _post(client, "rep-int", sid)
    codes = {d["code"] for d in body["required_disclosures"]}
    assert "market_revenue_at_duals_zero_profit_at_optimum" in codes
    assert "market_revenue_at_duals_exceeds_cost_at_size_limit" not in codes
    assert "zero profit by construction" in _econ_text(body)

    monkeypatch.setattr(solver_service, "run_simulation", BoundSolver())
    sid = _full(client, api_project, "rep-bound")
    body = _post(client, "rep-bound", sid)
    assert body["sections"]["economics"]["status"] == "ok"
    codes = {d["code"] for d in body["required_disclosures"]}
    assert "size_at_upper_bound" in codes
    assert "market_revenue_at_duals_zero_profit_at_optimum" not in codes
    assert "market_revenue_at_duals_exceeds_cost_at_size_limit" in codes
    text = _econ_text(body)
    assert "zero profit" not in text and "exceeds" in text
    html = client.get(f"/api/projects/rep-bound/studies/{sid}/report.html").text
    assert "zero profit by construction" not in html


def test_every_money_table_states_currency_and_year_in_html_and_docx(
        client, api_project, studies_on, fake):
    """BC-S7-3: options, cash flow, value streams and tornado, both formats."""
    import re

    import docx

    name = "rep-units"
    sid = _full(client, api_project, name)
    _post(client, name, sid)
    year = str(client.get(f"/api/projects/{name}/studies/{sid}/report").json()["currency_year"])
    html = client.get(f"/api/projects/{name}/studies/{sid}/report.html").text
    money = re.compile(r"NPV|CAPEX|Savings|Salvage|cash flow|Swing|Annual value|Replacements|"
                       r"O&amp;M|discounted")
    heads = [h for h in re.findall(r"<th>([^<]*)</th>", html) if money.search(h)]
    for label in ("Battery NPV", "Net cash flow", "Annual value", "Swing"):
        assert any(h.startswith(label) for h in heads), label
    for h in heads:
        assert "EUR" in h and year in h, h
    doc = docx.Document(io.BytesIO(
        client.get(f"/api/projects/{name}/studies/{sid}/report.docx").content))
    dheads = [c.text for t in doc.tables for c in t.rows[0].cells
              if money.search(c.text.replace("&", "&amp;"))]
    for label in ("Battery NPV", "Net cash flow", "Annual value", "Swing"):
        assert any(h.startswith(label) for h in dheads), label
    for h in dheads:
        assert "EUR" in h and year in h, h


def test_an_intake_edit_after_the_run_is_seen_by_every_reader(
        client, api_project, studies_on, fake):
    """Gate [S4]: one rule — findings and case 409, the report stale."""
    name = "rep-intake"
    sid = _run(client, api_project, name)
    assert _post(client, name, sid)["stale"] is False
    new = {**NO_PV_TOU, "load": {**NO_PV_TOU["load"], "annual_mwh": 9000.0}}
    assert client.patch(f"/api/projects/{name}/studies/{sid}",
                        json={"intake": new}).status_code == 200
    body = client.get(f"/api/projects/{name}/studies/{sid}/report").json()
    assert body["stale"] is True and "intake_changed_since_findings" in body["stale_reasons"]
    for method, url in (("get", "findings"), ("get", "options/bess_2h/case"),
                        ("get", "options/bess_2h/case.xlsx"), ("post", "report")):
        r = getattr(client, method)(f"/api/projects/{name}/studies/{sid}/{url}")
        assert r.status_code == 409, (url, r.status_code)
        assert r.json()["detail"]["error_kind"] == "intake_changed_since_run", url


# ── gate F1 BC-F1-1: the export-cycling preflight flags reach the study ──

def _tou_with_export(price: float) -> dict:
    """The gate's probe: the TOU seed tariff, supplied with a custom export price."""
    from services.study.library import load_defaults

    t = load_defaults().tariffs["tou_reference_illustrative"].model_dump(mode="json")
    t.update(tariff_id="tou_custom_export", source="my contract",
             export={"price_per_mwh": price, "series_ref": None, "cap_mw": None})
    return t


_CYCLING = "tariff_export_exceeds_import"


def test_an_export_credit_that_cycles_is_disclosed_by_the_run_findings_verdict_and_report(
        client, api_project, studies_on, fake):
    """
    125 EUR/MWh of export credit against an off-peak band of 90: the LP pays
    itself to cycle energy through the connection. The study path used to keep
    only preflight ERRORS, so this warning never reached the study.
    """
    from services.study import report as Rep

    name = "rep-cycling"
    intake = {**NO_PV_TOU, "tariff": {"custom": _tou_with_export(125.0)}}
    sid = _full(client, api_project, name, intake=intake)
    run = client.get(f"/api/projects/{name}/studies/{sid}/run").json()
    assert run["preflight_flags"] == {o: [_CYCLING] for o in run["solved"]}, run
    tornado = client.get(f"/api/projects/{name}/studies/{sid}/findings/tornado").json()
    assert tornado["preflight_flags"] == [_CYCLING], tornado
    findings = client.get(f"/api/projects/{name}/studies/{sid}/findings").json()
    assert _CYCLING in findings["honesty_notes"]
    assert _CYCLING in findings["verdict"]["disclosures"], findings["verdict"]
    body = _post(client, name, sid)
    shown = {d["code"]: d["text"] for d in body["required_disclosures"]}
    assert shown.get(_CYCLING) == Rep.HELP[_CYCLING]


def test_a_tariff_whose_export_stays_below_import_discloses_no_cycling(
        client, api_project, studies_on, fake):
    name = "rep-no-cycling"
    intake = {**NO_PV_TOU, "tariff": {"custom": _tou_with_export(30.0)}}
    sid = _full(client, api_project, name, intake=intake)
    run = client.get(f"/api/projects/{name}/studies/{sid}/run").json()
    assert run["preflight_flags"] == {}
    tornado = client.get(f"/api/projects/{name}/studies/{sid}/findings/tornado").json()
    assert tornado["preflight_flags"] == []
    findings = client.get(f"/api/projects/{name}/studies/{sid}/findings").json()
    codes = set(findings["honesty_notes"]) | set(findings["verdict"]["disclosures"])
    assert not {c for c in codes if c.startswith("tariff_export_exceeds")}
