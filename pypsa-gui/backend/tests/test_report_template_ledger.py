"""Template assumptions in the main report pipeline (outsourced task L).

Pipeline 2 remains out of scope: see the pack-base reachability test below
and the studies route's run-before-report rule.
"""
from services.reports import assemble, evidence as E


def _collect(**sources):
    return E.collect_evidence(study_report=None, eh_report=None, worksheet=None, **sources)


def _document(ev):
    return assemble.evidence_only_document(ev, title="Template report", report_id="ledger", figure_pngs={})


def test_template_rows_are_an_uncategorised_assumptions_appendix_not_measurements():
    ev = _collect(template_ledger={"ledger": [
        {"item": "discount_rate", "label": "Discount rate", "value": 0, "unit": "%",
         "source": "assumed", "note": "2020 money, unchanged"},
        {"item": "process", "value": "NH3", "source": "entered", "note": "user entry"},
    ]})
    table = ev.tables["template_ledger"]
    assert table.columns == ["Assumption", "Value", "Unit", "Source", "Note"]
    assert table.rows == [["Discount rate", "0", "%", "assumed", "2020 money, unchanged"],
                          ["process", "NH3", "", "entered", "user entry"]]
    doc = _document(ev)
    assert [s.section_id for s in doc.sections][-2:] == ["template_ledger", "pipeline"]
    section = doc.sections[-2]
    assert section.heading == "Template assumptions" and section.category is None
    assert ev.required_disclosures[0] == E.NOTHING_MEASURED_DISCLOSURE


def test_absent_ledger_is_document_only_and_keeps_the_master_evidence_hash():
    # Captured before implementation, on unmodified report modules at master
    # c9508d4c40fefab39ece792326eb2f8f4b96a9e8 (plan-only a1c9e5b).
    original_hash = "0123bc4df2c312a435863ed3b377cff469b4363ef5c48c429db9a9f2aee7590c"
    base = _collect()
    assert E.evidence_hash(base) == original_hash
    assert _collect(template_ledger=None).model_dump() == base.model_dump()
    assert E.evidence_hash(_collect(template_ledger=None)) == original_hash
    assert "template_ledger" not in base.sections and "template_ledger" not in base.tables
    section = next(s for s in _document(base).sections if s.section_id == "template_ledger")
    assert section.status == "not_established"
    assert "no template ledger for this project" in " ".join(b.text for b in section.blocks)


def test_unresolved_growth_rows_are_named_and_listed_first_without_inventing_rows(tmp_path):
    import json
    from tests.test_report_grid_readiness import T

    planning = T.planning_spec("eh_datacenter_staged", {
        "start_year": 2030, "end_year": 2045, "dates": [2038, 2042],
    })
    T.write_sidecars(tmp_path, "eh_datacenter_staged", planning=planning)
    meta = json.loads((tmp_path / "eh_template.json").read_text())
    ev = _collect(template_ledger=meta)
    rows = ev.tables["template_ledger"].rows
    assert [r[0] for r in rows[:3]] == ["IT load 2030 (placeholder)", "IT load 2038 (placeholder)", "IT load 2042 (placeholder)"]
    assert len(rows) == len(meta["ledger"])
    assert {r[0] for r in rows} == {r.get("label") or r["item"] for r in meta["ledger"]}


def test_missing_duplicate_and_invalid_metadata_is_disclosed_without_discarding_valid_rows():
    ev = _collect(template_ledger={"ledger": [
        {"item": "placeholder dc_load 2038", "label": "Old placeholder", "value": 12, "source": "assumed"},
        {"item": "rate", "value": 0.07}, {"item": "rate", "value": 0.08}, None,
    ], "growth_inputs": [{"kind": "dc_load", "year": 2042, "set": False}]})
    assert len(ev.tables["template_ledger"].rows) == 3
    assert ev.tables["template_ledger"].rows[0][:2] == ["Old placeholder", "12"]
    disclosures = " ".join(ev.required_disclosures)
    assert "placeholder dc_load 2038" in disclosures and "dc_load 2042" in disclosures
    assert "duplicate" in disclosures.lower() and "invalid" in disclosures.lower()


def test_docx_keeps_assumed_sources_and_money_years_and_routes_a_comment_to_the_appendix():
    import io
    from docx import Document
    from services.reports.docx_writer import render_document_docx
    from services.reports.roundtrip import read_edited_docx

    ev = _collect(template_ledger={"ledger": [{"item": "capex", "value": 2020,
        "source": "assumed", "note": "EUR in 2020 money; no escalation"}]})
    document = _document(ev)
    word = Document(io.BytesIO(render_document_docx(document, figure_bytes={})))
    table = next(t for t in word.tables if any(c.text == "capex" for row in t.rows for c in row.cells))
    assert [c.text for c in table.rows[1].cells] == ["capex", "2020", "", "assumed", "EUR in 2020 money; no escalation"]
    heading = next(p for p in word.paragraphs if p.text == "Template assumptions")
    word.add_comment(runs=heading.runs[0], text="Confirm these assumptions", author="Reviewer")
    buf = io.BytesIO()
    word.save(buf)
    roundtrip = read_edited_docx(buf.getvalue(), document)
    section = next(s for s in roundtrip.sections if s.section_id == "template_ledger")
    assert section.comments == ["Confirm these assumptions"]
    assert any(getattr(b, "table_id", None) for b in section.blocks)


def test_create_report_reads_the_projects_template_sidecar(client, api_project, project_storage_dir):
    import json
    name = api_project("report-ledger")
    directory = project_storage_dir(name)
    (directory / "eh_template.json").write_text(json.dumps({"ledger": [
        {"item": "original", "value": 12, "source": "assumed", "note": "Do not reprice"},
    ]}))
    response = client.post(f"/api/projects/{name}/reports", json={"mode": "evidence_only", "title": "Ledger"})
    assert response.status_code == 200, response.text
    document = response.json()["document"]
    assert document["tables"]["template_ledger"]["rows"][0][:4] == ["original", "12", "", "assumed"]


def test_template_metadata_alone_cannot_admit_a_generated_report(client, api_project, project_storage_dir):
    import json
    import pytest
    from fastapi import HTTPException
    from routers.deps import AuthorizedProject
    from routers.report_jobs import _evidence_for

    name = api_project("metadata-only")
    directory = project_storage_dir(name)
    (directory / "eh_template.json").write_text(json.dumps({"ledger": [
        {"item": "reference", "value": 12, "source": "assumed"},
    ]}))
    project = AuthorizedProject(name=name, directory=directory, uuid=name, org_id="local", registry_key=name)
    # Directly exercise the brief's required report-job evidence seam without
    # resolving a provider or starting a model call.
    with pytest.raises(HTTPException) as refused:
        _evidence_for(project)
    assert refused.value.status_code == 400 and refused.value.detail["error_kind"] == "no_evidence"


def test_unreadable_or_non_object_sidecars_build_an_explicit_absent_ledger_report(client, api_project, project_storage_dir):
    name = api_project("broken-ledger")
    directory = project_storage_dir(name)
    for raw in ("{", "[]", "null"):
        (directory / "eh_template.json").write_text(raw)
        response = client.post(f"/api/projects/{name}/reports", json={"mode": "evidence_only", "title": "Missing"})
        assert response.status_code == 200, response.text
        section = next(s for s in response.json()["document"]["sections"] if s["section_id"] == "template_ledger")
        assert section["status"] == "not_established"
        assert section["blocks"][0]["text"] == "no template ledger for this project"


def test_non_planning_template_lists_its_original_ledger_without_growth_placeholders(tmp_path):
    import json
    from tests.test_report_grid_readiness import T

    T.write_sidecars(tmp_path, "eh_datacenter")
    meta = json.loads((tmp_path / "eh_template.json").read_text())
    table = _collect(template_ledger=meta).tables["template_ledger"]
    assert len(table.rows) == len(meta["ledger"])
    assert not any("(placeholder)" in row[0] for row in table.rows)


def test_a_recognised_decision_question_uses_a_pack_base_without_the_eh_sidecar(
        client, api_project, project_storage_dir, monkeypatch):
    import json
    from tests.study_s4_support import create_pack_study, enable_studies

    origin = api_project("eh-study-origin")
    (project_storage_dir(origin) / "eh_template.json").write_text(json.dumps({"ledger": [{"item": "origin", "value": 12}]}))
    enabled = enable_studies(monkeypatch)
    next(enabled)
    try:
        response = create_pack_study(client, origin, "ledger-pack-base")
        assert response.status_code == 201, response.text
        base = project_storage_dir("ledger-pack-base")
        assert (base / "network.nc").is_file()
        assert not (base / "eh_template.json").exists()
        assert (project_storage_dir(origin) / "eh_template.json").is_file()
    finally:
        enabled.close()


def test_report_jobs_keep_the_template_appendix_when_reportable_worksheet_evidence_exists(
        client, api_project, project_storage_dir):
    import json
    from routers.deps import AuthorizedProject
    from routers.report_jobs import _evidence_for
    from services.adequacy.worksheet import save_worksheet
    from tests.test_adequacy_worksheet import _expert_row

    name = api_project("ledger-with-evidence")
    directory = project_storage_dir(name)
    (directory / "eh_template.json").write_text(json.dumps({"ledger": [
        {"item": "reference", "value": 12, "source": "assumed"},
    ]}))
    save_worksheet(directory, manual_rows=[_expert_row()], overlays={})
    project = AuthorizedProject(name=name, directory=directory, uuid=name, org_id="local", registry_key=name)
    evidence, _figures = _evidence_for(project)
    assert evidence.tables["template_ledger"].rows[0][:4] == ["reference", "12", "", "assumed"]
    assert evidence.sections[E.WORKSHEET_SECTION].status == "ok"
