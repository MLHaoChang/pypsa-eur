"""
Owner decision EXPORT-INJECT-1: the CSV bundle and the Excel export escape a
string cell that a spreadsheet would run as a formula.

Found by LC-32 (#920, LC-D20): an asset named "=HYPERLINK(1)" went into the
CSV unescaped and into the workbook as a live formula. OWASP's CSV-injection
rule: a string cell starting with = + - @ tab or CR gets a leading single
quote; a number (a negative one too) stays a number. The GUI's own imports
strip that quote again, so an export round-trips.
"""

from __future__ import annotations

import csv
import io
import zipfile

import openpyxl
import pandas as pd
import pypsa
import pytest

from routers.io import _escape_formula, _unescape_formula

BAD = "=HYPERLINK(1)"


def _network() -> pypsa.Network:
    n = pypsa.Network()
    n.set_snapshots(pd.date_range("2025-01-01", periods=2, freq="h"))
    n.add("Carrier", "AC")
    n.add("Carrier", "@wind")
    n.add("Bus", BAD, v_nom=110.0)
    n.add("Bus", "B2", v_nom=110.0)
    n.add(
        "Generator",
        "+G1",
        bus=BAD,
        carrier="@wind",
        p_nom=10.0,
        p_min_pu=-0.5,
        marginal_cost=-3.0,
    )
    n.add("Line", "-L1", bus0=BAD, bus1="B2", x=0.1, s_nom=50.0)
    return n


@pytest.mark.parametrize(
    "raw", ["=1+1", "+1+cmd", "-2+3+cmd", "@SUM(A1)", "\t=1", "\r=1", "=", "-"]
)
def test_a_formula_trigger_gets_a_quote(raw):
    assert _escape_formula(raw) == "'" + raw
    assert _unescape_formula(_escape_formula(raw)) == raw


@pytest.mark.parametrize(
    "raw", ["Bus 0", "a=b", "'quoted", "", "-1.5", "+3", "-inf", "1e-3", "-1e5"]
)
def test_plain_text_and_numbers_are_left_alone(raw):
    assert _escape_formula(raw) == raw
    assert _unescape_formula(raw) == raw


def test_non_strings_are_left_alone():
    for v in (-1, -0.5, None, True):
        assert _escape_formula(v) is v


def _cells(rows):
    return [c for row in rows for c in row]


def test_csv_bundle_escapes_names_and_keeps_numbers(client, install_network):
    install_network(_network())
    r = client.get("/api/io/export/csv")
    assert r.status_code == 200, r.text
    zf = zipfile.ZipFile(io.BytesIO(r.content))
    for member in zf.namelist():
        rows = list(csv.reader(io.StringIO(zf.read(member).decode())))
        for cell in _cells(rows):
            assert _escape_formula(cell) == cell, (member, cell)
    buses = list(csv.reader(io.StringIO(zf.read("buses.csv").decode())))
    assert "'" + BAD in _cells(buses)
    gens = list(csv.DictReader(io.StringIO(zf.read("generators.csv").decode())))
    assert gens[0]["name"] == "'+G1"
    assert gens[0]["bus"] == "'" + BAD
    assert gens[0]["carrier"] == "'@wind"
    assert float(gens[0]["p_min_pu"]) == -0.5
    assert float(gens[0]["marginal_cost"]) == -3.0


def test_excel_export_writes_no_formula(client, install_network):
    install_network(_network())
    r = client.get("/api/io/export/excel")
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    values = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                assert cell.data_type != "f", (ws.title, cell.coordinate, cell.value)
                values.append(cell.value)
                if isinstance(cell.value, str):
                    assert _escape_formula(cell.value) == cell.value, cell.value
    assert "'" + BAD in values
    assert "'+G1" in values and "'-L1" in values and "'@wind" in values
    gens = wb["generators"]
    header = [c.value for c in gens[1]]
    row = [c.value for c in gens[2]]
    assert row[header.index("p_min_pu")] == -0.5
    assert row[header.index("marginal_cost")] == -3.0


@pytest.mark.parametrize(
    "fmt,mime",
    [
        ("csv", "application/zip"),
        ("excel", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    ],
)
def test_an_export_imports_back_unchanged(
    client, install_network, session_ctx, fmt, mime
):
    install_network(_network())
    data = client.get(f"/api/io/export/{fmt}").content
    r = client.post(f"/api/io/import/{fmt}", files={"file": (f"x.{fmt}", data, mime)})
    assert r.status_code == 200, r.text
    n = session_ctx(client).network
    assert BAD in n.buses.index
    assert list(n.generators.index) == ["+G1"]
    assert n.generators.at["+G1", "bus"] == BAD
    assert n.generators.at["+G1", "carrier"] == "@wind"
    assert n.generators.at["+G1", "p_min_pu"] == -0.5
    assert list(n.lines.index) == ["-L1"]
    assert n.lines.at["-L1", "bus0"] == BAD
    assert "@wind" in n.carriers.index


def _with_cached_formula(
    xlsx: bytes, sheet: str, column: str, formula: str, cached: str
) -> bytes:
    """Put `formula` in the first data row of `column`, with the cached value
    Excel would have stored beside it (openpyxl writes none)."""
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    ws = wb[sheet]
    header = [c.value for c in ws[1]]
    ws.cell(row=2, column=header.index(column) + 1).value = "=" + formula
    buf = io.BytesIO()
    wb.save(buf)
    src = zipfile.ZipFile(buf)
    out = io.BytesIO()
    patched = 0
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item)
            old = f"<f>{formula}</f><v></v>".encode()
            patched += data.count(old)
            dst.writestr(
                item, data.replace(old, f"<f>{formula}</f><v>{cached}</v>".encode())
            )
    assert patched == 1
    return out.getvalue()


def test_an_excel_formula_imports_as_its_cached_value(
    client, install_network, session_ctx
):
    """GK B on #925: a workbook saved by Excel keeps a formula and its cached
    value; the import reads the value, as PyPSA's own reader does."""
    install_network(_network())
    data = _with_cached_formula(
        client.get("/api/io/export/excel").content, "generators", "p_nom", "10*2", "20"
    )
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    r = client.post("/api/io/import/excel", files={"file": ("x.xlsx", data, mime)})
    assert r.status_code == 200, r.text
    n = session_ctx(client).network
    assert n.generators.at["+G1", "p_nom"] == 20
    assert BAD in n.buses.index


def test_a_legacy_sheet_workbook_strips_only_the_escape_quote(client, session_ctx):
    """GK B on #925: a workbook without a `network` sheet goes through the
    legacy per-sheet reader, which strips the quote the same way."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    buses = wb.create_sheet("buses")
    buses.append(["name", "v_nom"])
    buses.append(["B1", 110.0])
    gens = wb.create_sheet("generators")
    gens.append(["name", "bus", "p_nom"])
    gens.append(["'" + BAD, "B1", 5.0])
    gens.append(["'-5", "B1", 6.0])
    buf = io.BytesIO()
    wb.save(buf)
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    r = client.post(
        "/api/io/import/excel", files={"file": ("x.xlsx", buf.getvalue(), mime)}
    )
    assert r.status_code == 200, r.text
    n = session_ctx(client).network
    assert sorted(n.generators.index) == sorted([BAD, "'-5"])
    assert n.generators.at[BAD, "p_nom"] == 5.0


@pytest.mark.parametrize(
    "raw",
    ["=1+1", "+1+cmd", "-2+3+cmd", "@SUM(A1)", "\t=1", "\r=1", "=", "-",
     "-5", "+3.2", "plain", "", 7, -3.5, None],
)
def test_chat_csv_cell_escapes_like_the_io_export(raw):
    from services.chat_tools import _inert_csv_cell

    assert _inert_csv_cell(raw) == _escape_formula(raw)
