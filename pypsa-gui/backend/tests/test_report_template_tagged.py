"""
WP9 of docs/superpowers/plans/2026-09-28-llm-report-generation-increment-1.md —
tagged rendering: Jinja2 tags inside a user's Word template are filled from
a ``ReportDocument``.

Every template is built here with python-docx into ``tmp_path``; no binary
fixture is committed. The fixture document mirrors the shape of
``tests/test_report_docx_writer.py::_document`` (copied, not imported).
"""
from __future__ import annotations

import io
import struct
import zlib
from collections.abc import Callable
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from models.report import (
    Bullets,
    Callout,
    Field,
    Figure,
    FigureRef,
    Paragraph,
    ReportDocument,
    Section,
    SectionAudit,
    Table,
    TableRef,
)
from services.reports.template_tagged import TaggedRenderError, render_tagged


# ── fixtures ─────────────────────────────────────────────────────────────


def _png_1x1() -> bytes:
    def chunk(kind: bytes, body: bytes) -> bytes:
        return (struct.pack(">I", len(body)) + kind + body
                + struct.pack(">I", zlib.crc32(kind + body) & 0xFFFFFFFF))

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 0, 0, 0, 0)
    idat = zlib.compress(b"\x00\x80")
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr)
            + chunk(b"IDAT", idat) + chunk(b"IEND", b""))


_PNG = _png_1x1()
_FIGURES = {"fmea_pareto": _PNG}


def _document() -> ReportDocument:
    """Three sections, every block type, a table and a figure."""
    return ReportDocument(
        report_id="0123456789abcdef", version=2, title="Client study report",
        created_at="2026-09-28T10:00:00+00:00", evidence_hash="c" * 64,
        profile_id="local-llama", model="llama-3.1-8b", mode="generated",
        sections=[
            Section(section_id="executive_summary", heading="Executive summary",
                    source="code", status="ok", blocks=[
                        Callout(kind="gap", text="the demand profile is frozen"),
                        Callout(kind="disclosure", text="excludes load-shedding cost"),
                        Field(key="Cost at target", value="1,234,568 €"),
                        TableRef(table_id="headline"),
                    ]),
            Section(section_id="fmea_top", heading="Residual failure modes",
                    source="llm", status="ok",
                    audit=SectionAudit(unverified=["42 h"]),
                    blocks=[
                        Paragraph(md="The **top mode** is *backup* at `100,000 €/yr`."),
                        Paragraph(md="A second paragraph."),
                        Bullets(items=["first point", "second point"]),
                        TableRef(table_id="fmea_top", caption="Ranked modes"),
                        FigureRef(figure_id="fmea_pareto", caption="Pareto"),
                        Field(key="Engine", value="copt"),
                    ]),
            Section(section_id="certification", heading="Certification",
                    source="code", status="skipped", note="budget exhausted",
                    blocks=[Callout(kind="not_established",
                                    text="this section was not run in this study "
                                         "(budget exhausted).")]),
        ],
        tables={
            "headline": Table(table_id="headline", columns=["Item", "Value"],
                              rows=[["MC LOLE (h/yr)", "not established"],
                                    ["Cost at target", "1,234,568 €"]],
                              caption="Headline results"),
            "fmea_top": Table(table_id="fmea_top",
                              columns=["Rank", "Name", "Criticality"],
                              rows=[["1", "backup", "100,000 €/yr"],
                                    ["2", "import", "25,000 €/yr"],
                                    ["3", "spare", "not established"]]),
        },
        figures={"fmea_pareto": Figure(figure_id="fmea_pareto",
                                       png_file="figures/fmea_pareto.png",
                                       caption="criticality by mode")},
    )


def _template(tmp_path: Path, build: Callable[[object], None]) -> bytes:
    """A .docx built by ``build(doc)``, saved under ``tmp_path`` and read back."""
    doc = Document()
    build(doc)
    path = tmp_path / "template.docx"
    doc.save(str(path))
    return path.read_bytes()


def _render(tmp_path: Path, build: Callable[[object], None]) -> Document:
    data = render_tagged(_document(), _template(tmp_path, build),
                         figure_bytes=_FIGURES)
    return Document(io.BytesIO(data))


def _text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


# ── acceptance ───────────────────────────────────────────────────────────


def test_meta_title_renders_in_a_title_paragraph(tmp_path):
    doc = _render(tmp_path, lambda d: d.add_heading("{{ meta.title }}", level=0))
    title = doc.paragraphs[0]
    assert title.text == "Client study report"
    assert title.style.name == "Title"


def test_a_tag_split_across_three_runs_renders_with_the_first_runs_formatting(tmp_path):
    def build(d):
        p = d.add_paragraph()
        p.add_run("{{ me").bold = True
        p.add_run("ta.title")
        p.add_run(" }}")

    doc = _render(tmp_path, build)
    p = doc.paragraphs[0]
    assert p.text == "Client study report"
    runs = [r for r in p.runs if r.text]
    assert len(runs) == 1 and runs[0].bold is True


def test_row_loop_yields_one_row_per_item_and_drops_the_template_row(tmp_path):
    def build(d):
        t = d.add_table(rows=2, cols=3)
        for cell, text in zip(t.rows[0].cells, ["Rank", "Name", "Criticality"]):
            cell.text = text
        cells = t.rows[1].cells
        cells[0].text = "{% for row in tables.fmea_top.rows %}{{ row[0] }}"
        cells[1].text = "{{ row[1] }}"
        cells[2].text = "{{ row[2] }}{% endfor %}"

    doc = _render(tmp_path, build)
    table = doc.tables[0]
    assert len(table.rows) == 4  # header + 3 data rows
    assert [c.text for c in table.rows[1].cells] == ["1", "backup", "100,000 €/yr"]
    assert [c.text for c in table.rows[3].cells] == ["3", "spare", "not established"]
    assert "{%" not in _text(doc) and "{{" not in _text(doc)


def test_row_loop_with_endfor_in_a_following_row(tmp_path):
    def build(d):
        t = d.add_table(rows=3, cols=2)
        t.rows[0].cells[0].text = "{% for row in tables.headline.rows %}{{ row[0] }}"
        t.rows[0].cells[1].text = "{{ row[1] }}"
        t.rows[1].cells[0].text = "{{ loop.index }}"
        t.rows[2].cells[0].text = "{% endfor %}"

    doc = _render(tmp_path, build)
    table = doc.tables[0]
    assert len(table.rows) == 4  # 2 items × 2 template rows, endfor row gone
    assert [c.text for c in table.rows[0].cells] == ["MC LOLE (h/yr)", "not established"]
    assert table.rows[1].cells[0].text == "1"
    assert table.rows[3].cells[0].text == "2"


def test_figure_tag_alone_in_a_paragraph_becomes_one_inline_picture(tmp_path):
    def build(d):
        d.add_paragraph("{{ figures.fmea_pareto }}")
        d.add_paragraph("See {{ figures.fmea_pareto }} above.")

    doc = _render(tmp_path, build)
    assert len(doc.inline_shapes) == 1
    assert doc.paragraphs[0].text == ""
    assert doc.paragraphs[1].text == "See criticality by mode above."


def test_fields_text_becomes_paragraphs_with_bold_runs(tmp_path):
    def build(d):
        d.add_paragraph("{{ fields.fmea_top.text }}", style="List Number")
        d.add_paragraph("after")

    doc = _render(tmp_path, build)
    texts = [p.text for p in doc.paragraphs]
    assert texts == ["The top mode is backup at 100,000 €/yr.",
                     "A second paragraph.", "after"]
    first = doc.paragraphs[0]
    assert first.style.name == "List Number"
    assert doc.paragraphs[1].style.name == "List Number"
    assert [r.text for r in first.runs if r.bold] == ["top mode"]
    assert [r.text for r in first.runs if r.italic] == ["backup"]
    assert [r.text for r in first.runs if r.font.name == "Consolas"] == ["100,000 €/yr"]


def test_unknown_field_raises_naming_it_and_listing_the_known_ids(tmp_path):
    with pytest.raises(TaggedRenderError) as exc:
        _render(tmp_path, lambda d: d.add_paragraph("{{ fields.nope.text }}"))
    message = str(exc.value)
    assert "nope" in message
    assert "fmea_top" in message and "executive_summary" in message
    assert "{{ fields.nope.text }}" in message
    assert "meta" in message and "tables" in message


def test_unknown_table_lists_the_known_table_ids(tmp_path):
    with pytest.raises(TaggedRenderError) as exc:
        _render(tmp_path, lambda d: d.add_paragraph("{{ tables.missing.rows }}"))
    assert "missing" in str(exc.value) and "headline" in str(exc.value)


def test_for_without_endfor_is_a_syntax_error(tmp_path):
    with pytest.raises(TaggedRenderError) as exc:
        _render(tmp_path, lambda d: d.add_paragraph(
            "{% for row in tables.fmea_top.rows %}{{ row[0] }}"))
    assert "endfor" in str(exc.value)


def test_row_loop_without_endfor_is_an_error_not_a_crash(tmp_path):
    def build(d):
        t = d.add_table(rows=1, cols=1)
        t.rows[0].cells[0].text = "{% for row in tables.fmea_top.rows %}{{ row[0] }}"

    with pytest.raises(TaggedRenderError) as exc:
        _render(tmp_path, build)
    assert "endfor" in str(exc.value)


def test_footer_tag_renders(tmp_path):
    def build(d):
        footer = d.sections[0].footer
        footer.is_linked_to_previous = False
        footer.paragraphs[0].text = "Evidence {{ meta.evidence_hash }}"
        d.add_paragraph("body")

    doc = _render(tmp_path, build)
    assert doc.sections[0].footer.paragraphs[0].text == "Evidence " + "c" * 64


def test_template_without_tags_renders_its_text_unchanged(tmp_path):
    def build(d):
        d.add_heading("Corporate report", level=1)
        d.add_paragraph("Nothing to fill here, not even {curly} braces.")
        t = d.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text = "a"
        t.rows[0].cells[1].text = "b"

    template = _template(tmp_path, build)
    before = _text(Document(io.BytesIO(template)))
    after = render_tagged(_document(), template, figure_bytes=_FIGURES)
    assert _text(Document(io.BytesIO(after))) == before


def test_if_hides_the_paragraph_on_a_skipped_section(tmp_path):
    def build(d):
        d.add_paragraph('{% if fields.certification.status == "ok" %}'
                        "Certified.{% endif %}")
        d.add_paragraph('{% if fields.fmea_top.status == "ok" %}'
                        "FMEA established.{% endif %}")

    doc = _render(tmp_path, build)
    assert [p.text for p in doc.paragraphs] == ["FMEA established."]


def test_context_exposes_headline_fields_section_fields_and_callouts(tmp_path):
    def build(d):
        d.add_paragraph("{{ headline['Cost at target'] }}")
        d.add_paragraph("{{ fields.executive_summary['Cost at target'] }}")
        d.add_paragraph("{{ fields.fmea_top.Engine }}")
        d.add_paragraph("{{ fields.fmea_top.heading }} / {{ fields.fmea_top.status }}")
        d.add_paragraph("{{ fields.fmea_top.bullets | join(', ') }}")
        d.add_paragraph("{{ fields.certification.note }}")
        d.add_paragraph("{{ disclosures | join('; ') }} | {{ gaps | join('; ') }}")
        d.add_paragraph("{{ unverified.fmea_top | join(', ') }}")
        d.add_paragraph("{{ tables.fmea_top.columns | join('/') }} "
                        "{{ tables.headline.caption }}")
        d.add_paragraph("{{ meta.version }} {{ meta.mode }} {{ meta.language }} "
                        "{{ meta.profile_id }} {{ meta.model }} {{ meta.created_at }}")

    doc = _render(tmp_path, build)
    assert [p.text for p in doc.paragraphs] == [
        "1,234,568 €",
        "1,234,568 €",
        "copt",
        "Residual failure modes / ok",
        "first point, second point",
        "budget exhausted",
        "excludes load-shedding cost | the demand profile is frozen",
        "42 h",
        "Rank/Name/Criticality Headline results",
        "2 generated en local-llama llama-3.1-8b 2026-09-28T10:00:00+00:00",
    ]


def test_a_figure_without_bytes_states_it_was_not_produced(tmp_path):
    template = _template(tmp_path, lambda d: d.add_paragraph("{{ figures.fmea_pareto }}"))
    data = render_tagged(_document(), template, figure_bytes={})
    doc = Document(io.BytesIO(data))
    assert len(doc.inline_shapes) == 0
    assert "not produced" in doc.paragraphs[0].text


def test_settings_and_section_properties_are_untouched(tmp_path):
    template = _template(tmp_path, lambda d: d.add_paragraph("{{ meta.title }}"))
    before = Document(io.BytesIO(template))
    after = Document(io.BytesIO(render_tagged(_document(), template,
                                              figure_bytes=_FIGURES)))
    assert after.settings.element.xml == before.settings.element.xml
    assert (after.element.body.find(qn("w:sectPr")).xml
            == before.element.body.find(qn("w:sectPr")).xml)
