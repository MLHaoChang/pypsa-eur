"""The report's investment-case and island (control) chapters (E2E demo plan
2026-10-14 D8, data center first; owner: docx finance and island chapters).

Pure, as the rest of the evidence: the route passes the stored investment
case report and the campus island report plus GFM settings; each becomes an
evidence section with code-rendered tables, its numbers in the audit's facts
(IRRs as percent, so "8.5 %" verifies), and a chapter in the evidence-only
document. A chapter is present only when its source was passed: a report of
a project without an investment case or a campus keeps its evidence hash.
"""
import pytest

from services.reports import assemble, chapters, evidence as E
from tests.test_chat_tools_investment_case import _report as _ic_report


def _island_report():
    return {
        "caveat": "screening, not a dynamic study",
        "missing": ["cost of islanding study"],
        "tables": [
            {"id": "requirements", "group": "inputs", "title": "Island requirements", "model": "sidecar",
             "caveat": "c", "columns": ["field", "value", "source"],
             "rows": [["bridge_s", 1.0, "assumed"], ["rocof_window_ms", 500.0, "assumed"]],
             "assumed_rows": [0, 1]},
            {"id": "freq_summary", "group": "frequency", "title": "Frequency check (I6)", "model": "RMS",
             "caveat": "c", "columns": ["scenario", "event", "verdict", "worst_nadir_hz"],
             "rows": [["seamless", "E2", "fail", 47.3]], "assumed_rows": []},
            {"id": "freq_events", "group": "frequency", "title": "Events", "model": "RMS",
             "caveat": "c", "columns": ["x"], "rows": [[1]], "assumed_rows": []},
        ],
    }


def _gfm(status="not_achievable", state="ready", applied=None):
    rec = {"label": "heuristic screening", "status": status, "achievable": status != "not_achievable",
           "simulations": 32, "max_simulations": 400, "margin": 0.1,
           "settings": {"BESS1": {"model": "gfm_droop",
                                  "current": {"droop_pct": 4.0, "tau_f_s": 0.1, "i_max_pu": 1.2,
                                              "H_equivalent_s": 0.4},
                                  "recommended": {"droop_pct": 2.5, "tau_f_s": 0.6, "i_max_pu": 1.2,
                                                  "H_equivalent_s": 1.2}}},
           "capacity": [{"case": "2030 h5 seamless E2", "kind": "inertia", "delta": 40.0, "unit": "MW·s",
                         "binding": "transition RoCoF", "derivation": "d", "feasible": True, "note": ""}]
           if status == "not_achievable" else [],
           "cases": [], "comparison": [], "ledger": ["droop 1-10 % (assumed)"], "notes": [], "excluded": []}
    if applied:
        rec["applied"] = applied
    return {"state": state, "reason": "", "recommendation": rec}


def _collect(**kw):
    return E.collect_evidence(study_report=None, eh_report=None, worksheet=None, **kw)


# ── presence and the hash ───────────────────────────────────────────────────


def test_without_the_sources_there_is_no_chapter_and_the_hash_holds():
    base = E.collect_evidence(study_report=None, eh_report=None, worksheet=None)
    assert chapters.INVESTMENT not in base.sections and chapters.ISLAND not in base.sections
    assert E.evidence_hash(_collect(ic_report=None, island=None)) == E.evidence_hash(base)


def test_a_passed_but_empty_source_is_a_not_established_chapter():
    ev = _collect(ic_report={}, island={"report": None, "gfm": None})
    assert ev.sections[chapters.INVESTMENT].status == "not_established"
    assert ev.sections[chapters.ISLAND].status == "not_established"
    doc = assemble.evidence_only_document(ev, title="t", report_id="r", figure_pngs={})
    ids = [s.section_id for s in doc.sections]
    assert chapters.INVESTMENT in ids and chapters.ISLAND in ids
    assert any(chapters.INVESTMENT in line or "Investment case" in line for line in ev.not_established)


# ── the investment case ─────────────────────────────────────────────────────


def test_the_investment_chapter_carries_the_headline_figures_and_the_basis():
    rep = _ic_report()
    p = rep["sections"]["project"]["payload"]
    p["counterfactual"] = {"present": True, "counterfactual_basis": "site_today",
                           "basis": "compared with the site as it stands today",
                           "notes": ["pv has no typed cost; treated as a new investment."]}
    p["island_breakeven"] = {"eur_per_year": 693523.0, "needs_outage_duration": True,
                             "text": "The island capability must be worth at least EUR 693,523 a year …"}
    ev = _collect(ic_report=rep, ic_stale=False)
    s = ev.sections[chapters.INVESTMENT]
    assert s.status == "ok" and s.kind == "investment"
    pay = s.payload
    irr = rep.get("project_irr_post_tax") or p.get("project_post_tax_irr")
    if irr is not None:
        assert pay["project_irr_post_tax_pct"] == pytest.approx(100 * irr)
    assert pay["compared_with"] == "compared with the site as it stands today"
    assert pay["island_breakeven_eur_per_year"] == 693523.0
    table = ev.tables["investment_case"]
    rows = dict(table.rows)
    assert rows["Compared with"].startswith("compared with the site")
    assert rows["Island capability break-even"].startswith("The island capability must be worth")
    assert table.source_path.startswith(f"/sections/{chapters.INVESTMENT}")
    # the number audit can verify an IRR quoted in percent
    facts = E.flatten_numbers(ev)
    assert any(f.unit == "%" and f.path.endswith("_pct") for f in facts) or irr is None


def test_a_stale_investment_case_says_so():
    ev = _collect(ic_report=_ic_report(), ic_stale=True)
    s = ev.sections[chapters.INVESTMENT]
    assert s.status == "ok" and "inputs changed" in (s.note or "")
    assert any("Investment case" in d and "inputs changed" in d for d in ev.required_disclosures)


# ── the island (control) chapter ────────────────────────────────────────────


def test_the_island_chapter_shows_the_checks_and_the_gfm_settings():
    ev = _collect(island={"report": _island_report(), "gfm": _gfm()})
    s = ev.sections[chapters.ISLAND]
    assert s.status == "ok" and s.kind == "island"
    # the summary tables, not every hourly table
    assert "island_requirements" in ev.tables and "island_freq_summary" in ev.tables
    assert "island_freq_events" not in ev.tables
    fs = ev.tables["island_freq_summary"]
    assert fs.columns == ["scenario", "event", "verdict", "worst_nadir_hz"] and fs.rows[0][2] == "fail"
    gfm = ev.tables["island_gfm_settings"]
    assert gfm.columns == ["Unit", "Control", "Setting", "Current", "Recommended"]
    assert ["BESS1", "gfm_droop", "droop_pct", "4", "2.5"] in gfm.rows
    assert "island_gfm_capacity" in ev.tables
    disc = " ".join(ev.required_disclosures)
    assert "heuristic screening" in disc and "capacity is needed" in disc
    assert "screening, not a dynamic study" in disc


def test_an_applied_recommendation_asks_for_the_re_solve():
    gfm = _gfm(status="recommended", state="stale")
    gfm["applied"] = {"at": 1, "changes": [{"unit": "BESS1"}], "hub_resolve_needed": True}
    ev = _collect(island={"report": _island_report(), "gfm": gfm})
    assert any("Re-solve the hub" in d for d in ev.required_disclosures)


def test_the_re_solve_disclosure_outlives_the_recommendation():
    """#310 review: a campus run removes the recommendation; the applied
    record still asks for the re-solve, even with no island table."""
    gfm = {"state": "not_run", "reason": "r", "recommendation": None,
           "applied": {"at": 1, "changes": [], "hub_resolve_needed": True}}
    ev = _collect(island={"report": None, "gfm": gfm})
    assert ev.sections[chapters.ISLAND].status == "ok"
    assert any("Re-solve the hub" in d for d in ev.required_disclosures)


def test_the_island_chapter_lands_after_the_reliability_chapters_and_the_investment_first():
    ev = _collect(ic_report=_ic_report(), island={"report": _island_report(), "gfm": _gfm()})
    doc = assemble.evidence_only_document(ev, title="t", report_id="r", figure_pngs={})
    ids = [s.section_id for s in doc.sections]
    assert ids[0] == "executive_summary"
    assert ids[1] == chapters.INVESTMENT                     # D8: summary → investment → …
    assert ids.index(chapters.ISLAND) > ids.index("multi_energy")
    assert ids[-1] == "pipeline"
    inv = next(s for s in doc.sections if s.section_id == chapters.INVESTMENT)
    assert inv.source == "code" and inv.status == "ok"
    assert any(getattr(b, "table_id", None) == "investment_case" for b in inv.blocks)


def test_the_new_chapters_are_titled_for_the_docx_and_the_prompt():
    from services.reports import docx_writer

    for sid in (chapters.INVESTMENT, chapters.ISLAND):
        assert sid in E.SECTION_TITLES and sid in docx_writer.SECTION_TITLES


# ── through the routes ──────────────────────────────────────────────────────


def _post(client, name):
    r = client.post(f"/api/projects/{name}/reports", json={"mode": "evidence_only", "title": "Demo"})
    assert r.status_code == 200, r.text[:400]
    return r.json()


def test_the_route_adds_the_investment_chapter_when_a_case_is_stored(
        client, api_project, session_state):
    name = api_project("rep-ic")
    session_state(client)["investment_case_report"] = _ic_report()
    body = _post(client, name)
    by_id = {s["section_id"]: s for s in body["document"]["sections"]}
    inv = by_id[chapters.INVESTMENT]
    assert inv["status"] == "ok"
    assert any(b["type"] == "table_ref" and b["table_id"] == "investment_case" for b in inv["blocks"])
    assert [s["section_id"] for s in body["document"]["sections"]][1] == chapters.INVESTMENT
    assert chapters.ISLAND not in by_id                     # no island data on disk
    # the docx carries the chapter's heading
    import io

    import docx

    r = client.post(f"/api/projects/{name}/reports/{body['report_id']}/export", json={})
    assert r.status_code == 200, r.text[:300]
    blob = client.get(f"/api/projects/{name}/uploads/{r.json()['file_id']}/blob")
    heads = [p.text for p in docx.Document(io.BytesIO(blob.content)).paragraphs
             if p.style.name.startswith("Heading")]
    assert "Investment case" in heads


def test_the_route_adds_the_island_chapter_when_the_project_has_island_data(
        client, api_project, project_storage_dir):
    import json

    from tests.test_campus_island_service import _config

    name = api_project("rep-island")
    (project_storage_dir(name) / "island_config.json").write_text(json.dumps(_config()))
    body = _post(client, name)
    by_id = {s["section_id"]: s for s in body["document"]["sections"]}
    isl = by_id[chapters.ISLAND]
    assert isl["status"] == "ok"
    assert any(b["type"] == "table_ref" and b["table_id"] == "island_requirements" for b in isl["blocks"])
    assert any(b["type"] == "callout" and "cost of islanding study" in b["text"] for b in isl["blocks"])
    assert chapters.INVESTMENT not in by_id


def test_a_project_without_either_keeps_its_report(client, api_project):
    name = api_project("rep-none")
    ids = [s["section_id"] for s in _post(client, name)["document"]["sections"]]
    assert chapters.INVESTMENT not in ids and chapters.ISLAND not in ids


# ── #307 review nits ─────────────────────────────────────────────────────────


def test_without_a_counterfactual_the_chapter_says_so_in_its_own_words():
    """#307 review: a case with no counterfactual reads "compared with nothing:
    incremental = total", never "the incremental figures are no counterfactual …"."""
    rep = _ic_report()
    rep["sections"]["project"]["payload"]["counterfactual"] = {
        "present": False, "basis": "no counterfactual in the case (the adapter builds one only "
                                   "when the owner is the site party): incremental = total"}
    ev = _collect(ic_report=rep)
    assert ev.sections[chapters.INVESTMENT].payload["compared_with"] is None
    assert dict(ev.tables["investment_case"].rows)["Compared with"].startswith("nothing")
    assert not any("the incremental figures are no counterfactual" in d
                   for d in ev.required_disclosures)
    assert any("no counterfactual" in d and "incremental = total" in d
               for d in ev.required_disclosures)


def test_an_incomplete_search_shows_its_best_settings_so_far_never_a_recommendation():
    """#307 review, #303's never-a-verdict rule: an incomplete search's column
    reads "Best found so far"."""
    ev = _collect(island={"report": _island_report(), "gfm": _gfm(status="incomplete")})
    assert ev.tables["island_gfm_settings"].columns[-1] == "Best found so far"
    ev = _collect(island={"report": _island_report(), "gfm": _gfm(status="recommended")})
    assert ev.tables["island_gfm_settings"].columns[-1] == "Recommended"


def test_an_island_chapter_without_an_apply_record_carries_no_applied_key():
    """#310 review: the key is added only with a record, so a report of an
    island with no recommendation keeps its evidence hash."""
    ev = _collect(island={"report": _island_report(), "gfm": {"state": "not_run", "reason": "r",
                                                               "recommendation": None}})
    assert "applied" not in ev.sections[chapters.ISLAND].payload["gfm"]


def test_island_constraints_off_says_the_hub_does_not_read_them():
    gfm = {"state": "ready", "reason": "", "recommendation": None,
           "applied": {"at": 1, "changes": [], "hub_resolve_needed": False}}
    ev = _collect(island={"report": _island_report(), "gfm": gfm})
    assert any("does not read them" in d for d in ev.required_disclosures)
    assert not any("Re-solve the hub" in d for d in ev.required_disclosures)
