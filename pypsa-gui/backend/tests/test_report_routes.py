"""
WP1 — the report routes under `/api/projects/{name}/reports`, and the bundle
transitions `reports/` must survive.

Authorization is `ProjectAccessDep`, so another org's project answers 404
(never 403 — the existence oracle `routers/deps.py` documents). A delete is a
write, so a live FOREIGN edit lock refuses it the way the worksheet PUT is
refused (`_check_project_lock`, check-only). Every path is joined from
`AuthorizedProject.directory` and a validated id; the id regex is the same
shape as upload ids, so `..` and a wrong length are 400 `invalid_report_id`.

Bundle inclusion: `reports/` joined `_BUNDLE_DIRS`, so it travels with
save-as, save-a-copy, scenario copy, snapshot create + restore and bundle
export + import — the seven transitions `docs/CHATBOT_UPLOADS_WORKFLOW.md`
lists for `uploads/`.
"""
from __future__ import annotations

import io
import uuid as _uuid
import zipfile
from datetime import datetime, UTC

import pytest
from fastapi.testclient import TestClient

import main
from models.report import Paragraph, ReportDocument, Section
from services.reports import store
from tests.conftest import attach_session, build_network

_PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 16


def _doc(report_id: str, title: str = "Report") -> ReportDocument:
    return ReportDocument(
        report_id=report_id, version=1, title=title,
        created_at="2026-09-28T10:00:00+00:00", evidence_hash="a" * 64,
        profile_id=None, model=None, mode="evidence_only",
        sections=[Section(section_id="s", heading="S", source="code", status="ok",
                          blocks=[Paragraph(md="v1 text")])],
        tables={}, figures={"fig": {"figure_id": "fig", "png_file": "fig.png"}},
    )


def _seed(project_dir, report_id: str = "0123456789abcdef", *, title="Report",
          versions: int = 1, figure: bool = True) -> str:
    store.create_report(project_dir, _doc(report_id, title))
    for _ in range(versions - 1):
        store.save_version(project_dir, _doc(report_id, title))
    if figure:
        store.write_figure(project_dir, report_id, "fig", _PNG)
    return report_id


@pytest.fixture
def same_org_other_user(_auth_db, seeded_identity):
    """
    A second user in the SAME org — the realistic intruder for a lock test
    (the same fixture `test_worksheet_foreign_lock.py` uses).
    """
    from db.models import OrgMembership, User
    from services.auth_service import hash_password

    _engine, session_local = _auth_db
    with session_local() as db:
        u = User(
            id=_uuid.uuid4(),
            email=f"colleague-{_uuid.uuid4().hex[:6]}@example.com",
            password_hash=hash_password("irrelevant"),
            status="active",
            is_super_admin=False,
            created_at=datetime.now(tz=UTC),
        )
        db.add(u)
        db.flush()
        db.add(OrgMembership(id=_uuid.uuid4(), user_id=u.id,
                             org_id=seeded_identity["org_id"], role="admin"))
        db.commit()
        uid = u.id
    with TestClient(main.app) as c:
        yield attach_session(c, session_local, uid)


# ── reads ───────────────────────────────────────────────────────────────────

def test_list_is_empty_for_a_project_without_reports(client, api_project):
    name = api_project("rep-empty")
    r = client.get(f"/api/projects/{name}/reports")
    assert r.status_code == 200, r.text
    assert r.json() == []


def test_list_get_latest_and_a_specific_version(client, api_project, project_storage_dir):
    name = api_project("rep-read")
    rid = _seed(project_storage_dir(name), versions=2)

    listed = client.get(f"/api/projects/{name}/reports")
    assert listed.status_code == 200, listed.text
    assert [m["report_id"] for m in listed.json()] == [rid]
    assert listed.json()[0]["latest_version"] == 2

    latest = client.get(f"/api/projects/{name}/reports/{rid}")
    assert latest.status_code == 200, latest.text
    assert latest.json()["version"] == 2
    assert latest.json()["sections"][0]["blocks"][0] == {"type": "paragraph", "md": "v1 text"}

    v1 = client.get(f"/api/projects/{name}/reports/{rid}", params={"version": 1})
    assert v1.status_code == 200, v1.text
    assert v1.json()["version"] == 1
    v1_path = client.get(f"/api/projects/{name}/reports/{rid}/versions/1")
    assert v1_path.status_code == 200, v1_path.text
    assert v1_path.json() == v1.json()

    gone = client.get(f"/api/projects/{name}/reports/{rid}", params={"version": 7})
    assert gone.status_code == 404
    assert gone.json()["detail"]["error_kind"] == "report_not_found"


def test_an_unknown_report_is_404_report_not_found(client, api_project):
    name = api_project("rep-missing")
    r = client.get(f"/api/projects/{name}/reports/ffffffffffffffff")
    assert r.status_code == 404
    assert r.json()["detail"]["error_kind"] == "report_not_found"


# `..` is sent percent-encoded: httpx collapses a literal `..` segment before
# the request leaves the client (so `/reports/..` would simply hit the list
# route), while `%2e%2e` survives the client and is unquoted server-side into
# the path parameter — which is what an attacker's raw socket sends. A slash
# cannot reach the handler at all (the converter is `[^/]+`); the store test
# covers that shape directly, since the chat path has no converter.
@pytest.mark.parametrize("bad", [
    "%2e%2e", "0123456789abcde", "0123456789abcdef0", "0123456789ABCDEF", "not-an-id",
])
def test_a_malformed_id_is_400_invalid_report_id(client, api_project, bad):
    name = api_project("rep-badid")
    for url in (
        f"/api/projects/{name}/reports/{bad}",
        f"/api/projects/{name}/reports/{bad}/figures/fig",
    ):
        r = client.get(url)
        assert r.status_code == 400, (url, r.status_code, r.text[:200])
        assert r.json()["detail"]["error_kind"] == "invalid_report_id"
    r = client.delete(f"/api/projects/{name}/reports/{bad}")
    assert r.status_code == 400, r.text[:200]
    assert r.json()["detail"]["error_kind"] == "invalid_report_id"


def test_another_orgs_project_is_404_never_403(client, other_org_client, api_project,
                                               project_storage_dir):
    name = api_project("rep-tenant")
    rid = _seed(project_storage_dir(name))
    probes = [
        ("GET", "/api/projects/{p}/reports"),
        ("GET", f"/api/projects/{{p}}/reports/{rid}"),
        ("GET", f"/api/projects/{{p}}/reports/{rid}/figures/fig"),
        ("DELETE", f"/api/projects/{{p}}/reports/{rid}"),
    ]
    for method, template in probes:
        hit = other_org_client.request(method, template.format(p=name))
        missing = other_org_client.request(method, template.format(p="NoSuchProject"))
        assert hit.status_code == 404, (method, template, hit.status_code, hit.text[:200])
        assert missing.status_code == 404
        assert hit.content == missing.content, "the 404 body must not leak existence"
    # The owner still sees everything; the intruder's DELETE did nothing.
    assert client.get(f"/api/projects/{name}/reports/{rid}").status_code == 200


# ── figures ─────────────────────────────────────────────────────────────────

def test_figure_route_serves_png_inline(client, api_project, project_storage_dir):
    name = api_project("rep-fig")
    rid = _seed(project_storage_dir(name))
    r = client.get(f"/api/projects/{name}/reports/{rid}/figures/fig")
    assert r.status_code == 200, r.text[:200]
    assert r.headers["content-type"].startswith("image/png")
    assert r.headers["content-disposition"].startswith("inline")
    assert r.content == _PNG


@pytest.mark.parametrize("figure_id", ["nope", "%2e%2e", "fig.png", "a%20b"])
def test_figure_route_404s_on_an_unknown_or_malformed_figure_id(
        client, api_project, project_storage_dir, figure_id):
    name = api_project("rep-fig-404")
    rid = _seed(project_storage_dir(name))
    r = client.get(f"/api/projects/{name}/reports/{rid}/figures/{figure_id}")
    assert r.status_code == 404, (figure_id, r.status_code, r.text[:200])
    assert r.json()["detail"]["error_kind"] == "figure_not_found"


# ── delete + edit lock ──────────────────────────────────────────────────────

def test_delete_is_refused_under_a_foreign_lock_the_way_the_worksheet_is(
        client, api_project, project_storage_dir, same_org_other_user):
    name = api_project("rep-lock")
    rid = _seed(project_storage_dir(name))
    assert client.post(f"/api/projects/{name}/lock").status_code == 200

    # Control: the worksheet PUT is refused for the same intruder.
    ws = same_org_other_user.put(f"/api/projects/{name}/worksheet",
                                 json={"manual_rows": [], "overlays": {}})
    assert ws.status_code == 409 and ws.json()["detail"]["error_kind"] == "project_locked"

    r = same_org_other_user.delete(f"/api/projects/{name}/reports/{rid}")
    assert r.status_code == 409, r.text[:200]
    assert r.json()["detail"]["error_kind"] == "project_locked"
    assert (project_storage_dir(name) / "reports" / rid / "v1.json").exists()
    # Reads are not writes: the intruder may still list and fetch.
    assert same_org_other_user.get(f"/api/projects/{name}/reports/{rid}").status_code == 200


def test_the_holder_can_delete_and_a_free_project_stays_free(
        client, api_project, project_storage_dir, same_org_other_user):
    name = api_project("rep-del")
    d = project_storage_dir(name)
    rid_a = _seed(d, "aaaaaaaaaaaaaaaa")
    rid_b = _seed(d, "bbbbbbbbbbbbbbbb")
    assert client.post(f"/api/projects/{name}/lock").status_code == 200

    r = client.delete(f"/api/projects/{name}/reports/{rid_a}")
    assert r.status_code == 200, r.text[:200]
    assert r.json() == {"deleted": True, "report_id": rid_a}
    assert not (d / "reports" / rid_a).exists()
    assert (d / "reports" / rid_b / "v1.json").exists()
    assert client.delete(f"/api/projects/{name}/reports/{rid_a}").status_code == 404

    # Check-only: on a FREE project a colleague's delete succeeds and does not
    # claim the lock.
    assert client.delete(f"/api/projects/{name}/lock").status_code in (200, 204)
    r = same_org_other_user.delete(f"/api/projects/{name}/reports/{rid_b}")
    assert r.status_code == 200, r.text[:200]
    assert client.post(f"/api/projects/{name}/lock").status_code == 200


# ── bundle transitions ──────────────────────────────────────────────────────

def test_reports_is_a_bundle_dir_like_uploads():
    from routers.projects import _BUNDLE_DIRS
    assert "uploads" in _BUNDLE_DIRS and "reports" in _BUNDLE_DIRS


def test_save_as_carries_reports(client, install_network, project_storage_dir):
    install_network(build_network(), name="RepA")
    assert client.post("/api/projects/RepA", params={"force": True, "rebind": True}
                       ).status_code == 200
    rid = _seed(project_storage_dir("RepA"))

    r = client.post("/api/projects/RepB", params={"rebind": "true"})
    assert r.status_code == 200, r.text[:200]
    assert (project_storage_dir("RepB") / "reports" / rid / "v1.json").exists()
    assert (project_storage_dir("RepB") / "reports" / rid / "figures" / "fig.png").exists()
    assert client.get(f"/api/projects/RepB/reports/{rid}").status_code == 200


def test_save_a_copy_carries_reports(client, install_network, project_storage_dir):
    install_network(build_network(), name="CopyA")
    assert client.post("/api/projects/CopyA", params={"force": True, "rebind": True}
                       ).status_code == 200
    rid = _seed(project_storage_dir("CopyA"))

    r = client.post("/api/projects/CopyOfCopyA")
    assert r.status_code == 200, r.text[:200]
    assert (project_storage_dir("CopyOfCopyA") / "reports" / rid / "v1.json").exists()
    # A copy, not a move: the source keeps its report.
    assert (project_storage_dir("CopyA") / "reports" / rid / "v1.json").exists()


def test_a_scenario_fork_inherits_reports(client, api_project, project_storage_dir):
    name = api_project("rep-parent")
    rid = _seed(project_storage_dir(name))
    r = client.post(f"/api/projects/{name}/scenarios",
                    json={"name": "rep-child", "description": "x"})
    assert r.status_code in (200, 201), r.text[:200]
    child = r.json().get("name") or "rep-child"
    assert (project_storage_dir(child) / "reports" / rid / "v1.json").exists()
    assert client.get(f"/api/projects/{child}/reports/{rid}").status_code == 200


def test_snapshot_carries_reports_and_restore_replaces_them(
        client, api_project, project_storage_dir):
    name = api_project("rep-snap")
    d = project_storage_dir(name)
    rid_old = _seed(d, "aaaaaaaaaaaaaaaa", title="in the snapshot")

    r = client.post(f"/api/projects/{name}/snapshots", json={"label": "v1"})
    assert r.status_code in (200, 201), r.text[:200]
    snap_id = r.json()["id"]
    snap_dirs = [p for p in (d / "snapshots").iterdir() if p.is_dir()]
    assert any((p / "reports" / rid_old / "v1.json").exists() for p in snap_dirs), \
        sorted(p.name for p in snap_dirs)

    # Diverge after the snapshot: drop the old report, add a new one.
    assert client.delete(f"/api/projects/{name}/reports/{rid_old}").status_code == 200
    rid_new = _seed(d, "bbbbbbbbbbbbbbbb", title="after the snapshot")
    assert [m["report_id"] for m in client.get(f"/api/projects/{name}/reports").json()] \
        == [rid_new]

    r = client.post(f"/api/projects/{name}/snapshots/{snap_id}/restore")
    assert r.status_code == 200, r.text[:200]
    ids = [m["report_id"] for m in client.get(f"/api/projects/{name}/reports").json()]
    assert ids == [rid_old], ids
    assert not (d / "reports" / rid_new).exists()
    assert (d / "reports" / rid_old / "figures" / "fig.png").read_bytes() == _PNG


def test_bundle_export_carries_reports_and_import_restores_them(
        client, api_project, project_storage_dir):
    name = api_project("rep-bundle")
    rid = _seed(project_storage_dir(name), versions=2)

    r = client.get(f"/api/projects/{name}/bundle")
    assert r.status_code == 200
    names = set(zipfile.ZipFile(io.BytesIO(r.content)).namelist())
    assert {f"reports/{rid}/meta.json", f"reports/{rid}/v1.json",
            f"reports/{rid}/v2.json", f"reports/{rid}/figures/fig.png"} <= names, names

    r = client.post("/api/projects/import_bundle?name=rep_bundle_imported",
                    files={"file": ("rep.pypsaproj.zip", r.content, "application/zip")})
    assert r.status_code in (200, 201), r.text[:200]
    imported = r.json().get("imported") or r.json().get("name") or "rep_bundle_imported"
    got = client.get(f"/api/projects/{imported}/reports/{rid}")
    assert got.status_code == 200, got.text[:200]
    assert got.json()["version"] == 2
    fig = client.get(f"/api/projects/{imported}/reports/{rid}/figures/fig")
    assert fig.status_code == 200 and fig.content == _PNG


# ═════════════════════════════════════════════════════════════════════════
# WP5 — POST /{name}/reports (evidence_only) and POST …/export
# ═════════════════════════════════════════════════════════════════════════

from models.energy_hub import (  # noqa: E402
    REPORT_SECTIONS,
    ReferenceDesignReport,
    SectionState,
    empty_section_map,
)
from services.adequacy import eh_report as R  # noqa: E402
from services.reports.docx_writer import DOCX_MIME  # noqa: E402


def _store_eh_report(state: dict) -> None:
    """The same hand-built report `test_chat_report_export_tools._store_report` uses."""
    sections = empty_section_map(default="skipped")
    sections["fmea_top"] = SectionState(status="ok", payload={
        "top": [{"rank": 1, "mode_id": "gen:g:forced_outage",
                 "component_class": "Generator", "name": "g",
                 "failure_class": "A", "occurrence_per_year": 1.0,
                 "occurrence_basis": "FOR", "severity_eur": 10.0,
                 "criticality_eur_per_year": 10.0, "delta_eue_mwh": 0.5,
                 "engine": "copt", "fidelity": "analytic_convolution"}],
        "classes_included": ["A"], "note": "Link-primary residual risk",
    })
    report = ReferenceDesignReport(
        archetype="strong_grid", pack_hash="p", assumptions_hash="a",
        sections=sections, ens_cap_permyriad=10.0,
    )
    R.store_eh_report(state, report)


def test_post_evidence_only_builds_v1_and_the_figure_files(
        client, api_project, project_storage_dir, session_state):
    name = api_project("rep-post")
    _store_eh_report(session_state(client))

    r = client.post(f"/api/projects/{name}/reports",
                    json={"mode": "evidence_only", "title": "Client report"})
    assert r.status_code == 200, r.text[:400]
    body = r.json()
    rid = body["report_id"]
    assert body["mode"] == "evidence_only"
    assert body["title"] == "Client report"
    assert body["latest_version"] == 1
    doc = body["document"]
    assert doc["version"] == 1 and doc["mode"] == "evidence_only"
    assert doc["evidence_hash"] == body["evidence_hash"]

    by_id = {s["section_id"]: s for s in doc["sections"]}
    for section in REPORT_SECTIONS:
        assert section in by_id, section
    fmea = by_id["fmea_top"]
    assert fmea["status"] == "ok"
    assert any(b["type"] == "table_ref" and b["table_id"] == "fmea_top"
               for b in fmea["blocks"])
    assert any(b["type"] == "figure_ref" and b["figure_id"] == "fmea_pareto"
               for b in fmea["blocks"])
    assert by_id["certification"]["status"] == "skipped"
    assert any(b["type"] == "callout" and b["kind"] == "not_established"
               for b in by_id["certification"]["blocks"])
    # ADR-0001 through the wire: the headline MC LOLE is null → not established.
    headline = doc["tables"]["headline"]
    lole = next(row for row in headline["rows"] if row[0].startswith("MC LOLE"))
    assert lole[1] == "not established"

    d = project_storage_dir(name)
    assert (d / "reports" / rid / "v1.json").is_file()
    assert (d / "reports" / rid / "meta.json").is_file()
    assert (d / "reports" / rid / "figures" / "fmea_pareto.png").is_file()
    assert doc["figures"]["fmea_pareto"]["png_file"] == "figures/fmea_pareto.png"
    fig = client.get(f"/api/projects/{name}/reports/{rid}/figures/fmea_pareto")
    assert fig.status_code == 200 and fig.content[:8] == b"\x89PNG\r\n\x1a\n"
    listed = client.get(f"/api/projects/{name}/reports").json()
    assert [m["report_id"] for m in listed] == [rid]
    assert client.get(f"/api/projects/{name}/reports/{rid}").json()["version"] == 1


def test_post_evidence_only_without_any_study_states_every_section(
        client, api_project):
    name = api_project("rep-post-empty")
    r = client.post(f"/api/projects/{name}/reports", json={"mode": "evidence_only"})
    assert r.status_code == 200, r.text[:400]
    doc = r.json()["document"]
    by_id = {s["section_id"]: s for s in doc["sections"]}
    for section in REPORT_SECTIONS:
        assert by_id[section]["status"] == "not_established"
        assert any(b["type"] == "callout" and b["kind"] == "not_established"
                   for b in by_id[section]["blocks"])
    assert doc["figures"] == {}
    assert r.json()["title"], "a default title is given when none is sent"


def test_post_with_another_mode_is_400_report_mode_not_supported(client, api_project):
    name = api_project("rep-post-mode")
    r = client.post(f"/api/projects/{name}/reports", json={"mode": "generated"})
    assert r.status_code == 400, r.text[:200]
    assert r.json()["detail"]["error_kind"] == "report_mode_not_supported"
    assert "phase 2" in r.json()["detail"]["message"]
    assert client.get(f"/api/projects/{name}/reports").json() == []


def test_post_is_refused_under_a_foreign_lock(
        client, api_project, same_org_other_user):
    name = api_project("rep-post-lock")
    assert client.post(f"/api/projects/{name}/lock").status_code == 200
    r = same_org_other_user.post(f"/api/projects/{name}/reports",
                                 json={"mode": "evidence_only"})
    assert r.status_code == 409, r.text[:200]
    assert r.json()["detail"]["error_kind"] == "project_locked"
    assert client.get(f"/api/projects/{name}/reports").json() == []


def test_export_saves_an_agent_export_chip_the_blob_route_serves(
        client, api_project, project_storage_dir, session_state):
    name = api_project("rep-export")
    _store_eh_report(session_state(client))
    rid = client.post(f"/api/projects/{name}/reports",
                      json={"mode": "evidence_only"}).json()["report_id"]

    r = client.post(f"/api/projects/{name}/reports/{rid}/export",
                    json={"filename": "client-report"})
    assert r.status_code == 200, r.text[:400]
    meta = r.json()
    assert meta["kind"] == "agent_export"
    assert meta["mime"] == DOCX_MIME
    assert meta["filename"] == "client-report.docx"
    assert (project_storage_dir(name) / "uploads" / meta["file_id"] / "blob").is_file()

    blob = client.get(f"/api/projects/{name}/uploads/{meta['file_id']}/blob")
    assert blob.status_code == 200
    assert blob.headers["content-type"].startswith(DOCX_MIME)
    from docx import Document
    doc = Document(io.BytesIO(blob.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    assert "Residual failure modes" in text
    assert "Link-primary residual risk" in text
    assert "g" in cells and "not established" in cells
    assert len(doc.inline_shapes) == 1
    from docx.oxml.ns import qn
    marks = [el.get(qn("w:name")) for el in doc.element.body.iter(qn("w:bookmarkStart"))]
    assert "sec:fmea_top" in marks
    # A default filename is derived when none is given, and it is a .docx.
    r2 = client.post(f"/api/projects/{name}/reports/{rid}/export", json={})
    assert r2.status_code == 200, r2.text[:200]
    assert r2.json()["filename"].endswith(".docx")
    kinds = {u["file_id"]: u["kind"]
             for u in client.get(f"/api/projects/{name}/uploads").json()}
    assert kinds.get(meta["file_id"]) == "agent_export"


def test_export_of_an_unknown_report_or_version_is_404(client, api_project):
    name = api_project("rep-export-404")
    r = client.post(f"/api/projects/{name}/reports/ffffffffffffffff/export", json={})
    assert r.status_code == 404, r.text[:200]
    assert r.json()["detail"]["error_kind"] == "report_not_found"
    rid = client.post(f"/api/projects/{name}/reports",
                      json={"mode": "evidence_only"}).json()["report_id"]
    r = client.post(f"/api/projects/{name}/reports/{rid}/export", json={"version": 9})
    assert r.status_code == 404
    assert r.json()["detail"]["error_kind"] == "report_not_found"
    bad = client.post(f"/api/projects/{name}/reports/not-an-id/export", json={})
    assert bad.status_code == 400
    assert bad.json()["detail"]["error_kind"] == "invalid_report_id"


# ═════════════════════════════════════════════════════════════════════════
# Follow-up — GET /{name}/reports/evidence_hash (the viewer's staleness badge)
# ═════════════════════════════════════════════════════════════════════════
#
# The badge "evidence changed since vN" compares a document's `evidence_hash`
# against the CURRENT session evidence, not against the newest evidence-only
# report. The route hashes exactly what the evidence-only POST would store,
# never answers 204 (the empty evidence has a hash too) and is declared before
# `/{report_id}` so its literal segment is not parsed as a report id.


def test_evidence_hash_route_never_204s_and_is_not_a_report_id(client, api_project):
    name = api_project("rep-ehash-empty")
    r = client.get(f"/api/projects/{name}/reports/evidence_hash")
    assert r.status_code == 200, r.text[:300]
    body = r.json()
    assert body.get("detail", {}).get("error_kind") != "invalid_report_id"
    assert isinstance(body["evidence_hash"], str) and len(body["evidence_hash"]) == 64
    # No EH report stored and no worksheet rows: every EH section is stated as
    # not established (the installed network alone may still yield the COPT
    # screening, so `sections_ok` is bounded, not zero).
    assert 0 <= body["sections_ok"] <= body["sections_total"]
    assert body["sections_total"] >= len(REPORT_SECTIONS)
    assert body["sections_ok"] < len(REPORT_SECTIONS)
    # The same evidence hashes the same way twice.
    assert client.get(f"/api/projects/{name}/reports/evidence_hash").json() == body


def test_evidence_hash_equals_what_the_evidence_only_post_stores(
        client, api_project, session_state):
    name = api_project("rep-ehash-post")
    _store_eh_report(session_state(client))
    before = client.get(f"/api/projects/{name}/reports/evidence_hash").json()
    assert before["sections_ok"] >= 1
    assert before["sections_total"] >= before["sections_ok"]

    created = client.post(f"/api/projects/{name}/reports",
                          json={"mode": "evidence_only"}).json()
    assert created["evidence_hash"] == before["evidence_hash"]
    assert created["document"]["evidence_hash"] == before["evidence_hash"]
    # Storing a report does not change the session evidence.
    after = client.get(f"/api/projects/{name}/reports/evidence_hash").json()
    assert after == before


def test_evidence_hash_changes_when_the_session_eh_report_changes(
        client, api_project, session_state):
    name = api_project("rep-ehash-change")
    empty = client.get(f"/api/projects/{name}/reports/evidence_hash").json()
    _store_eh_report(session_state(client))
    seeded = client.get(f"/api/projects/{name}/reports/evidence_hash").json()
    assert seeded["evidence_hash"] != empty["evidence_hash"]
    assert seeded["sections_ok"] > empty["sections_ok"]

    # One number moves → the hash moves with it (the collector's canonical form).
    sections = empty_section_map(default="skipped")
    sections["fmea_top"] = SectionState(status="ok", payload={
        "top": [{"rank": 1, "mode_id": "gen:g:forced_outage",
                 "component_class": "Generator", "name": "g",
                 "failure_class": "A", "occurrence_per_year": 2.0,
                 "occurrence_basis": "FOR", "severity_eur": 10.0,
                 "criticality_eur_per_year": 20.0, "delta_eue_mwh": 0.5,
                 "engine": "copt", "fidelity": "analytic_convolution"}],
        "classes_included": ["A"], "note": "Link-primary residual risk",
    })
    R.store_eh_report(session_state(client), ReferenceDesignReport(
        archetype="strong_grid", pack_hash="p", assumptions_hash="a",
        sections=sections, ens_cap_permyriad=10.0,
    ))
    moved = client.get(f"/api/projects/{name}/reports/evidence_hash").json()
    assert moved["evidence_hash"] != seeded["evidence_hash"]
    assert moved["sections_ok"] == seeded["sections_ok"]


def test_evidence_hash_is_project_scoped_404_for_another_org(
        client, other_org_client, api_project):
    name = api_project("rep-ehash-tenant")
    assert client.get(f"/api/projects/{name}/reports/evidence_hash").status_code == 200
    hit = other_org_client.get(f"/api/projects/{name}/reports/evidence_hash")
    missing = other_org_client.get("/api/projects/NoSuchProject/reports/evidence_hash")
    assert hit.status_code == 404, hit.text[:200]
    assert missing.status_code == 404
    assert hit.content == missing.content, "the 404 body must not leak existence"
