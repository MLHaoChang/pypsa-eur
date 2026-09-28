"""
QA: the report spike (WP0), end to end — a real Energy Hub study over HTTP,
then the chat tool that turns its stored ``ReferenceDesignReport`` into a
Word document, then the document fetched back through the uploads blob
route and re-opened.

Plan: docs/superpowers/plans/2026-09-28-llm-report-generation-increment-1.md
(phase 0 gate). ``tests/test_chat_report_export_tools.py`` covers the tool
with a hand-built report; this driver pins the JOURNEY: the numbers in the
document are the ones the study engines produced, the chip the panel would
render is really an ``agent_export`` upload the blob route serves with the
Word MIME, and a project with no study answers with the named error rather
than an empty file.

Runs in CI through ``tests/run_qa_drivers.py``. Exit 0 = every step passed.
"""
from __future__ import annotations

import io
import pathlib
import sys
import time

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

# THE load-bearing import, and it must come first: `qa_support` pins the
# sandbox before anything imports `main` or `settings`.
from tests import qa_support          # noqa: E402

from tests.eh_stage_fixtures import VOLL, certifiable_weak_network  # noqa: E402

PROJECT = "qa_reports_phase0"
BUDGET_SOLVES = 8
STAGES = ["apply_pack", "ens_solve", "fmea_top", "assemble"]

PASS = 0
FAIL = 0


def _step(label: str, ok: bool, msg: str = "") -> None:
    global PASS, FAIL
    if ok:
        PASS += 1
        print(f"  [PASS] {label}" + (f" — {msg}" if msg else ""))
    else:
        FAIL += 1
        print(f"  [FAIL] {label}" + (f" — {msg}" if msg else ""))


def _poll(path: str, *, timeout: float = 600.0) -> dict:
    c = qa_support.client()
    t0 = time.time()
    body: dict = {}
    while time.time() - t0 < timeout:
        r = c.get(path)
        if r.status_code == 204:
            return {"status": "no-content"}
        if r.status_code != 200:
            return {"status": f"http-{r.status_code}", "error": r.text[:300]}
        body = r.json()
        if body.get("status") != "running":
            return body
        time.sleep(0.25)
    return {"status": "timeout", **body}


def _as_session(fn, *args, **kwargs):
    """
    Call an in-process chat tool the way a chat turn does: under the signed-in
    session's ``ProjectContext``. The active project is per SESSION since the
    tenancy migration, so a bare call reads the process foreground and sees
    no project at all (``no_active_project``) — the driver's first defect.
    """
    from services.pypsa_service import PyPSAService

    token = PyPSAService.bind_request_context(qa_support.session_context())
    try:
        return fn(*args, **kwargs)
    finally:
        PyPSAService.reset_request_context(token)


def _docx_text(blob: bytes) -> tuple[str, int]:
    from docx import Document

    doc = Document(io.BytesIO(blob))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts), len(doc.inline_shapes)


# ── the journey ───────────────────────────────────────────────────────────

def section_1_no_study_is_a_named_error() -> None:
    print("\n[1] a project with no EH study")
    from fastapi import HTTPException

    from services import chat_tools

    # Install UNBOUND, then save: the first save claims the name AND binds the
    # session context's storage_dir (org-scoped). Installing under the name
    # first would pre-claim it, the save would not rebind, and every in-process
    # export would fall back to the legacy flat directory the routes no longer
    # read — a 404 on the chip's download.
    qa_support.install_network(certifiable_weak_network())
    qa_support.save_project(PROJECT)
    try:
        _as_session(chat_tools.export_eh_report_docx)
        _step("the tool refuses with eh_report_not_found", False, "no exception")
    except HTTPException as exc:
        detail = exc.detail if isinstance(exc.detail, dict) else {}
        _step("the tool refuses with eh_report_not_found",
              exc.status_code == 404 and detail.get("error_kind") == "eh_report_not_found",
              f"HTTP {exc.status_code} {detail}")
        _step("the refusal names the remedy", "run_eh_study" in str(detail.get("message")))


def section_2_real_study_to_docx() -> None:
    print("\n[2] a real strong_grid study → export_eh_report_docx → blob → Word")
    c = qa_support.client()
    r = c.put("/api/simulation/solver_config",
              json={"solver_name": "highs", "voll": VOLL})
    _step("the solver config takes the VoLL", r.status_code == 200, r.text[:200])

    r = c.post("/api/results/eh_study", json={
        "archetype": "strong_grid", "budget_solves": BUDGET_SOLVES, "stages": STAGES})
    _step("the study starts", r.status_code == 200, r.text[:200])
    if r.status_code != 200:
        return
    study = _poll("/api/results/eh_study")
    _step("the study finishes", study.get("status") == "done",
          f"status={study.get('status')} error={str(study.get('error'))[:200]}")
    if study.get("status") != "done":
        return
    r = c.get("/api/results/eh_reference_design")
    _step("the report is persisted for GET", r.status_code == 200, r.text[:200])
    if r.status_code != 200:
        return
    rep = r.json()
    comp = rep.get("completeness") or {}
    fmea = (rep.get("sections") or {}).get("fmea_top") or {}
    top = (fmea.get("payload") or {}).get("top") or []
    _step("fmea_top is established with at least one ranked mode",
          comp.get("fmea_top") == "ok" and len(top) >= 1,
          f"status={comp.get('fmea_top')} rows={len(top)} note={str(fmea.get('note'))[:120]}")

    from services import chat_tools

    meta = _as_session(chat_tools.export_eh_report_docx, filename="qa-phase0-report")
    _step("the tool returns an agent_export chip", meta.get("kind") == "agent_export"
          and meta.get("filename") == "qa-phase0-report.docx", str(meta)[:200])

    r = c.get(f"/api/projects/{PROJECT}/uploads/{meta['file_id']}/blob")
    _step("the blob route serves it", r.status_code == 200, f"HTTP {r.status_code}")
    _step("with the Word MIME",
          r.headers.get("content-type", "").startswith(
              "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
          r.headers.get("content-type", ""))
    if r.status_code != 200:
        return
    text, pictures = _docx_text(r.content)
    _step("the document re-opens with python-docx", bool(text))
    _step("it carries the study's top failure mode by name",
          str(top[0].get("name")) in text if top else False,
          f"name={top[0].get('name') if top else None}")
    _step("it carries the ranking's criticality as a grouped number",
          f"{float(top[0]['criticality_eur_per_year']):,.0f}" in text if top else False)
    _step("the skipped stages are stated, not omitted",
          all(f"This section was" in text for _ in [0])
          and text.count("This section was") >= 3,
          f"statements={text.count('This section was')}")
    _step("the Link-primary disclosure is present", "Link-primary" in text)
    _step("the criticality figure is embedded", pictures == 1, f"pictures={pictures}")
    # ADR-0001 at the document level: the study skipped mc_certify, so the
    # headline MC LOLE is null — it must read "not established", never 0.
    # (A legitimate zero — a stage that charged no solves — is fine.)
    _step("a null headline figure reads 'not established', not 0",
          rep.get("mc_lole_h") is None
          and "MC LOLE (h/yr)\nnot established" in text,
          f"mc_lole_h={rep.get('mc_lole_h')}")

    r = c.get(f"/api/projects/{PROJECT}/uploads")
    kinds = {u.get("file_id"): u.get("kind") for u in (r.json() if r.status_code == 200 else [])}
    _step("the uploads list shows the chip as an agent export",
          kinds.get(meta["file_id"]) == "agent_export", str(kinds)[:200])


def main() -> int:
    print("=" * 60)
    print("QA: report spike — EH study → .docx export chip")
    print("=" * 60)
    crashed = False
    try:
        qa_support.reset_backend()
        qa_support.delete_project(PROJECT)
        section_1_no_study_is_a_named_error()
        section_2_real_study_to_docx()
    except Exception as exc:                                     # noqa: BLE001
        crashed = True
        import traceback
        print(f"\n  [FAIL] the journey raised {type(exc).__name__}: {exc}")
        traceback.print_exc()
    finally:
        try:
            qa_support.delete_project(PROJECT)
            qa_support.reset_backend()
        except Exception as exc:                                 # noqa: BLE001
            print(f"  ! cleanup raised {type(exc).__name__}: {exc}")

    total = PASS + FAIL + (1 if crashed else 0)
    print("\n" + "=" * 60)
    print(f"Total: {total}")
    print(f"Pass:  {PASS}")
    print(f"Fail:  {FAIL + (1 if crashed else 0)}")
    print("=" * 60)
    return 1 if (FAIL or crashed) else 0


if __name__ == "__main__":
    sys.exit(main())
