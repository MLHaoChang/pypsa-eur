"""
The decision report from REAL LP solves (plan S9; gate S7 carry "at least one
live-LP report"; gate S8 carry). Every S7 route test runs the S4 fake solver,
so the `bess_pv`, zero-size and size-bound disclosures had only ever been
rendered from constructed findings. Here they come from HiGHS solves of the
site golden fixture's production pack:

* the site fixture's evening-spike load and PV (`tests/golden/site_fixture.py`);
* a user storage quote of 400 EUR/kWh, at which the 4-hour battery sizes to
  ZERO (a real `size_zero_no_investment`), and
* a sizing limit of 1.0 x the connection (2 MW), at which the PV of
  `bess_pv_2h` stops AT ITS BOUND (a real `size_at_upper_bound:pv`), while its
  2-hour battery is interior (0.63 MW) and is valued against a PV-only
  reference solved for real.

The pipeline is the production one, in process: `packs.build_site_network`,
the commercial config compiled and bound as the runner does (U2 WP6), the
LP through `run_simulation` (IC's commercial chain prices the PoC and carries
the demand charge), `runner._read_option` (the bill, the asset
economics, the shared sizing classifier), `findings.run_tornado`,
`findings.assemble_findings`, `findings.centre_attributions`,
`report.build_decision_report` and the three renderers. Only the disk reads
of `assemble_decision_report` (`load_inputs`, `context_from_disk`) are
replaced by the same objects built in memory; the route level is the S7
tests' and the QA driver's (`tests/qa_decision_study.py`). The tornado's
drivers are the no-solve rows, so the whole module is four LP solves.
"""
from __future__ import annotations

import html
import io
from datetime import UTC, datetime

import pytest

from models.study import BaselineResult, Bill, DecisionStudy
from services.study import findings as F
from services.study import ledger as LG
from services.study import packs
from services.study import questions as Q
from services.study import report as REP
from services.study import runner as RN
from services.study import run_hashes, store
from tests.golden import site_fixture as sf

OPTIONS = ("none", "bess_4h", "bess_pv_2h")
SID = "9" * 32
# the tornado's drivers: the no-solve rows (the dispatch rows are the
# driver's and the S6 LP tests')
QUESTION = Q.BESS_AT_SITE.model_copy(update={"key_drivers": [
    "battery_storage_eur_per_kwh", "battery_inverter_eur_per_kw", "discount_rate"]})


def _lp(net, cfg, _variant_id=None):
    from tests.u2_targets import ic_solve

    return ic_solve(net, cfg)


@pytest.fixture(scope="module")
def live(tmp_path_factory):
    intake = sf.site_intake()
    library = sf.site_library()
    ledger = sf.site_ledger()
    ledger = LG.apply_user_row(ledger, "battery_storage_eur_per_kwh", 400.0,
                               unit="EUR/kWh", changed_by="test")
    ledger = LG.apply_user_row(ledger, "sizing_limit_connection_multiple", 1.0,
                               unit="multiple of connection limit", changed_by="test")
    from tests.u2_targets import FAKE_REF, bound_option

    tariff = packs.effective_tariff(intake, ledger, library)
    nets, details, results = {}, {}, {}
    for oid in OPTIONS:
        n, compiled = bound_option(intake, ledger, oid, library)
        cfg = packs.option_solver_config(ledger, compiled)
        _lp(n, cfg)
        d = RN._read_option(n, cfg, Q.option(QUESTION, oid), "full_study",
                            tariff.currency_year, compiled=compiled)
        res = d.pop("result").model_copy(update={"project_ref": f"fork-{oid}"})
        nets[oid], details[oid], results[oid] = n, d, res
    # U2 WP8 (gate C6): the run record's bills are the engine's on each
    # solved meter, and the findings read them from there.
    bill = {oid: Bill.model_validate(details[oid]["bill"]) for oid in OPTIONS}
    ctx = F.TornadoContext(
        study_id=SID, question=QUESTION, intake=intake, ledger=ledger, library=library,
        tariff=tariff, baseline_network=nets["none"],
        baseline_bill=bill["none"], fidelity="full_study",
        options={oid: F.OptionInput(oid, nets[oid], bill[oid],
                                    details[oid]["asset_economics"],
                                    tuple(details[oid]["caveats"]))
                 for oid in OPTIONS if oid != "none"},
        export_series=FAKE_REF)
    ref_calls: list[str] = []
    outcome = F.run_tornado(ctx, lambda net, c, vid: ref_calls.append(vid) or _lp(net, c, vid))

    # what the run and the tornado would have written beside the study
    base_dir = tmp_path_factory.mktemp("live-base")
    hashes = {"ledger_hash": packs.ledger_hash(ledger),
              "intake_hash": run_hashes.intake_hash(intake),
              "option_network_hashes": {f"fork-{o}": "h" * 16 for o in OPTIONS}}
    faux = {"options": [results[o].model_dump(mode="json") for o in OPTIONS],
            "options_status": "ok", "pending_options": [], "hashes": hashes,
            "baseline": BaselineResult(project_ref="fork-none", solve_status="ok",
                                       bill=results["none"].bill).model_dump(mode="json"),
            "honesty_notes": []}
    store.save_aux(base_dir, SID, F.TORNADO_AUX, {
        "status": "done", "kind": "tornado", "ledger_hash": hashes["ledger_hash"],
        "option_network_hashes": hashes["option_network_hashes"],
        "robustness": outcome.robustness.model_dump(mode="json"),
        "attributions": [a.model_dump(mode="json") for a in outcome.attributions],
        "reference_bills": {k: b.model_dump(mode="json")
                            for k, b in outcome.reference_bills.items()}})
    now = datetime(2026, 9, 30, tzinfo=UTC)
    study = DecisionStudy(study_id=SID, name="Live site battery", question_id="bess_at_site",
                          base_project="b" * 32, intake=intake, ledger=ledger,
                          currency_year=2020, created_at=now, updated_at=now)
    inp = F.StudyInputs(study, QUESTION, {"details": details, "intake": intake}, faux, ledger,
                        library, tariff, rows={o: f"fork-{o}" for o in OPTIONS},
                        recorded={o: "h" * 16 for o in OPTIONS},
                        current={o: "h" * 16 for o in OPTIONS})
    findings = F.assemble_findings(study, base_dir, None, "b" * 32, loaded=(inp, ctx))
    _atts, cases = F.centre_attributions(ctx)
    report = REP.build_decision_report(REP.ReportInputs(
        study=study, question=QUESTION, findings=findings, cases=cases, ledger=ledger,
        tariff=tariff, details=details, option_forks={f"fork-{o}": o for o in OPTIONS},
        generated_at=now, run_intake=intake))
    return {"nets": nets, "details": details, "findings": findings, "report": report,
            "cases": cases, "ledger": ledger, "ref_calls": ref_calls, "outcome": outcome,
            "intake": intake, "library": library}


def test_the_real_solves_reach_the_three_states_the_report_must_disclose(live):
    sizes = {o: F.battery_size(n) for o, n in live["nets"].items() if o != "none"}
    assert F.is_zero_size(sizes["bess_4h"]), sizes
    assert sizes["bess_pv_2h"] > 0.5, sizes
    sizing = live["details"]["bess_pv_2h"]["sizing"]
    assert sizing["pv"]["binding_constraint"] == "at_upper_bound"
    assert sizing["battery"]["binding_constraint"] == "interior"
    assert live["details"]["bess_pv_2h"]["caveats"] == ["size_at_upper_bound:pv"]
    # the PV-only reference was solved, once
    assert live["ref_calls"] == ["ref-bess_pv_2h"]


def test_the_findings_judge_the_zero_size_battery_by_size_and_value_bess_pv_by_reference(live):
    f = live["findings"]
    atts = {a.option_id: a for a in f.battery_attribution}
    assert atts["bess_4h"].status == "skipped"
    assert atts["bess_pv_2h"].status == "ok"
    assert atts["bess_pv_2h"].method == "battery_removed_same_pv"
    assert "size_zero_no_investment" in f.honesty_notes
    assert f.verdict.status == "ok" and f.verdict.option_id == "bess_pv_2h"
    assert "size_at_upper_bound" in f.verdict.disclosures
    assert f.value_streams_basis == "pv_only_reference"
    assert f.completeness["battery_attribution"] == "ok"


def test_the_live_report_discloses_bess_pv_zero_size_and_the_size_bound(live):
    report = live["report"]
    codes = [d.code for d in report.required_disclosures]
    for code in ("battery_value_against_pv_only_reference",
                 "battery_only_options_sized_to_zero",
                 "size_at_upper_bound",
                 "market_revenue_at_duals_exceeds_cost_at_size_limit",
                 "value_streams_increment_over_pv_only",
                 "npv_nonnegative_at_optimum_by_construction"):
        assert code in codes, (code, codes)
    # at a size bound the zero-profit sentence must NOT be claimed (BC-S7-1)
    assert "market_revenue_at_duals_zero_profit_at_optimum" not in codes
    # the uploaded series is not a synthetic profile
    assert "sector_profiles_have_broad_peaks" not in codes
    # the prose guard passes on the live facts (it raises otherwise)
    assert REP.validate_prose(report)["executive_summary"]
    exec_facts = report.sections["executive_summary"].facts
    assert "option_total_npv" in exec_facts
    assert report.sections["executive_summary"].payload["option_id"] == "bess_pv_2h"


def test_the_live_report_renders_to_html_docx_and_xlsx_with_those_disclosures(live):
    import docx
    import openpyxl

    from services.study.render_docx import render_docx
    from services.study.render_html import render_html
    from services.study.report_xlsx import write_report_xlsx

    report = live["report"]
    wanted = [d for d in report.required_disclosures
              if d.code in ("battery_value_against_pv_only_reference",
                            "battery_only_options_sized_to_zero", "size_at_upper_bound")]
    assert len(wanted) == 3
    page = html.unescape(render_html(report))
    body = page[page.index("<body"):]
    for d in wanted:
        assert d.text in body, d.code
    assert body.count("data:image/png;base64,") >= 3

    doc = docx.Document(io.BytesIO(render_docx(report)))
    text = "\n".join([p.text for p in doc.paragraphs]
                     + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    for d in wanted:
        assert d.text in text, d.code

    wb = openpyxl.load_workbook(io.BytesIO(
        write_report_xlsx(report, live["cases"], live["ledger"])))
    assert {"Verdict", "Tornado", "Cash flows bess_pv_2h"} <= set(wb.sheetnames)


def test_the_findings_and_the_report_value_the_engine_solved_forks_on_the_engine(live):
    """
    Gate U2-WP7 X1 (mutation G4: `findings._on_engine` → False, every case
    back on the pro forma) and X4. The forks here are solved on the engine's
    commercial chain with the battery's two upfront parts (`engine_ready`),
    so every centre case and the PV-only reference case are the finance
    engine's. The report's economics section carries the engine's LCOS fact,
    its disclosure and the prose that explains it, and each money fact of the
    case names the engine that produced it (X4: never the pro forma's
    `cash_flow_expander` on an engine case; the HTML and XLSX print it). The
    case route's construction (`routers.studies._option_case` on an
    engine-ready fork with the run's export series) gives the report's LCOS.
    """
    from services.study import engine_adapter as A
    from tests.u2_targets import FAKE_REF

    out = live["outcome"]
    assert set(out.cases) == {"bess_4h", "bess_pv_2h"}
    assert set(out.reference_cases) == {"bess_pv_2h"}
    for case in (*out.cases.values(), *out.reference_cases.values(), *live["cases"].values()):
        assert case.engine == "finance_engine", case.case_id

    report = live["report"]
    econ = report.sections["economics"]
    oid = econ.payload["option_id"]
    assert oid == "bess_pv_2h"
    assert "lcos_finance_engine" in econ.facts and "lcos_excl_charging" not in econ.facts
    codes = [d.code for d in report.required_disclosures]
    assert "lcos_includes_charging_energy_cost" in codes
    assert "lcos_two_definitions" not in codes
    prose = " ".join(p.text for p in econ.prose)
    assert "{{lcos_finance_engine}}" in prose and "export income" in prose
    # X5 (a): imported charge at what the site paid, PV charge at the export
    # income given up, and why the optimisation's figure differs.
    [lcos_help] = [d.text for d in report.required_disclosures
                   if d.code == "lcos_includes_charging_energy_cost"]
    for words in ("price the site paid", "export income the site gave up",
                  "model's hourly price", "demand charge"):
        assert words in lcos_help and words in prose, words
    assert REP.validate_prose(report)["economics"]
    for key in ("case_npv", "case_irr", "case_payback_simple", "case_payback_discounted",
                "case_capex_total", "lcos_finance_engine"):
        assert econ.facts[key].engine == "finance_engine", key
        assert report.facts[key].engine == "finance_engine", key

    n = live["nets"][oid]
    ledger = live["ledger"]
    compiled = packs.option_commercial(live["intake"], ledger, live["library"], n.snapshots,
                                       export_series=FAKE_REF)
    assert A.engine_ready(n)
    route = A.option_case(n, packs.option_solver_config(ledger, compiled), ledger,
                          compiled=compiled, option_id=oid, study_id=SID,
                          fidelity="full_study",
                          asset_economics=live["details"][oid]["asset_economics"],
                          question=QUESTION, study_currency_year=2020).view
    assert route.engine == "finance_engine"
    assert route.kpis.lcos == pytest.approx(econ.facts["lcos_finance_engine"].value, rel=1e-12)


def test_every_bill_and_value_on_the_live_report_names_the_engine_that_made_it(live):
    """
    U2 WP8 (gate C6; the WP7 carry-forward of X4): the run's bills are the
    tariff engine's, established, and so are the streams built on them; the
    battery attributions, the verdict's money facts and the report's
    `option_total_npv` and `tornado_centre_npv` name the finance engine that
    valued the cases — never the pro forma's `cash_flow_expander`.
    """
    for oid, d in live["details"].items():
        assert d["bill"]["engine"] == "tariff_engine", oid
        assert d["bill"]["total"] is not None, (oid, d["bill"]["unavailable"])
    f, report = live["findings"], live["report"]
    assert f.value_streams and all(s.engine == "tariff_engine" for s in f.value_streams)
    assert report.facts["value_streams_total"].engine == "tariff_engine"
    for a in f.battery_attribution:
        assert a.engine == "finance_engine", a.option_id
    for key in ("battery_npv", "battery_payback_simple"):
        assert f.verdict.facts[key].engine == "finance_engine", key
    for key in ("option_total_npv", "tornado_centre_npv"):
        assert report.facts[key].engine == "finance_engine", key
    engines = report.sections["appendix"].payload["engines"]
    assert {"tariff_engine", "finance_engine"} <= set(engines)
    assert not {"bill_calculator", "cash_flow_expander"} & set(engines)
