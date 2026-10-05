"""
WP8 — ``services/reports/docx_reader``: the outline of a user's Word template.

Plan: docs/superpowers/plans/2026-09-28-llm-report-generation-increment-1.md
(Increment 2, WP8). The fixtures are BUILT here from
``tests/fixtures/report_templates/_build.py`` (python-docx) — no ``.docx``
binary is committed — so every assertion below is about a document whose
construction is visible in that script.

Body positions the assertions rely on (each top-level ``w:p``/``w:tbl`` is
one index step):

* ``tagged_minimal.docx``: 0 Title, 1 Heading 1, 2 split-tag paragraph,
  3 table, 4 closing paragraph.
* ``corporate_untagged.docx``: 0 Title, 1 "Prepared for", 2 "<Date>",
  3 page break, 4 TOC field, 5/7/9/11 Heading 1 with a body paragraph after
  each, 13 Heading 2, 14 table, 15 Heading 1, 16 body.
"""
from __future__ import annotations

import importlib.util
import io
import struct
import zipfile
import zlib
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from services.reports.docx_reader import (
    TemplateHeading,
    TemplateOutline,
    TemplateReadError,
    detect_language,
    merged_paragraph_text,
    read_template,
)

_BUILD = Path(__file__).resolve().parent / "fixtures" / "report_templates" / "_build.py"


def _load_builder():
    spec = importlib.util.spec_from_file_location("report_template_fixtures", _BUILD)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def fixtures(tmp_path_factory) -> dict[str, bytes]:
    builder = _load_builder()
    out = tmp_path_factory.mktemp("report_templates")
    paths = builder.build_all(out)
    paths.update(builder.build_all(out, language="de"))
    return {name: path.read_bytes() for name, path in paths.items()}


@pytest.fixture(scope="module")
def tagged(fixtures) -> TemplateOutline:
    return read_template(fixtures["tagged_minimal.docx"])


@pytest.fixture(scope="module")
def corporate(fixtures) -> TemplateOutline:
    return read_template(fixtures["corporate_untagged.docx"])


def _png_1x1() -> bytes:
    def chunk(kind: bytes, body: bytes) -> bytes:
        return (struct.pack(">I", len(body)) + kind + body
                + struct.pack(">I", zlib.crc32(kind + body) & 0xFFFFFFFF))

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 0, 0, 0, 0)
    idat = zlib.compress(b"\x00\x80")
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr)
            + chunk(b"IDAT", idat) + chunk(b"IEND", b""))


# ── tagged fixture ───────────────────────────────────────────────────────


def test_tagged_fixture_is_tagged_mode(tagged):
    assert tagged.mode == "tagged"
    assert tagged.has_toc is False
    assert tagged.unsupported == []


def test_tagged_every_tag_found_including_the_split_one(fixtures, tagged):
    # The fixture really splits the tag: three runs in the paragraph.
    doc = Document(io.BytesIO(fixtures["tagged_minimal.docx"]))
    assert [r.text for r in doc.paragraphs[2].runs] == [
        "{{ fields.", "executive_summary", ".text }}"]
    texts = [t.text for t in tagged.tags]
    assert texts == [
        "{{ meta.title }}",
        "{{ fields.executive_summary.text }}",
        "{% for row in tables.fmea_top.rows %}",
        "{{ row[0] }}",
        "{{ row[1] }}",
        "{% endfor %}",
        "{{ meta.evidence_hash }}",
    ]
    kinds = [t.kind for t in tagged.tags]
    assert kinds == ["var", "var", "for", "var", "var", "endfor", "var"]


def test_tagged_split_tag_has_its_paragraph_index(tagged):
    split = next(t for t in tagged.tags
                 if t.text == "{{ fields.executive_summary.text }}")
    assert split.paragraph_index == 2
    title = next(t for t in tagged.tags if t.text == "{{ meta.title }}")
    assert title.paragraph_index == 0


def test_tagged_for_endfor_pair_sit_on_the_table_index(tagged):
    loop = [t for t in tagged.tags if t.kind in ("for", "endfor")]
    assert [t.kind for t in loop] == ["for", "endfor"]
    assert {t.paragraph_index for t in loop} == {3}
    assert tagged.tables[0].index == 3


def test_tagged_footer_tag_is_found(tagged):
    footer_tag = next(t for t in tagged.tags if "evidence_hash" in t.text)
    assert footer_tag.kind == "var"
    assert footer_tag.paragraph_index == -1
    assert "{{ meta.evidence_hash }}" in tagged.footer_text


def test_tagged_indices_are_positional_across_paragraphs_and_tables(tagged):
    assert tagged.n_paragraphs == 4
    assert [h.index for h in tagged.headings] == [0, 1]
    assert tagged.tables[0].index == 3
    assert tagged.tables[0].n_rows == 2
    assert tagged.tables[0].n_cols == 4
    assert tagged.tables[0].header == ["Loop", "Rank", "Component", "End"]


def test_tagged_placeholders_do_not_pick_up_jinja_subscripts(tagged):
    # "{{ row[0] }}" contains "[0]" — that is a tag, not a "[…]" placeholder.
    assert tagged.placeholders == []


# ── untagged (corporate) fixture ─────────────────────────────────────────


def test_corporate_fixture_is_untagged_mode(corporate):
    assert corporate.mode == "untagged"
    assert corporate.tags == []


def test_corporate_headings_with_levels_and_style_names(corporate):
    got = [(h.index, h.level, h.text, h.style) for h in corporate.headings]
    assert got == [
        (0, 0, "[Client name] — Energy Hub Reference Design", "Title"),
        (5, 1, "1 Executive Summary", "Heading 1"),
        (7, 1, "2 Introduction", "Heading 1"),
        (9, 1, "3 Availability Target", "Heading 1"),
        (11, 1, "4 Residual Failure Modes", "Heading 1"),
        (13, 2, "4.1 Critical components", "Heading 2"),
        (15, 1, "5 Lorem ipsum", "Heading 1"),
    ]
    assert all(isinstance(h, TemplateHeading) for h in corporate.headings)


def test_corporate_has_toc_and_body_starts_after_it(corporate):
    assert corporate.has_toc is True
    assert corporate.body_start_index == 5
    flags = [(h.index, h.is_body_start) for h in corporate.headings]
    assert (5, True) in flags
    assert sum(1 for _, f in flags if f) == 1


def test_corporate_placeholders(corporate):
    got = [(p.text, p.paragraph_index) for p in corporate.placeholders]
    assert ("[Client name]", 0) in got
    assert ("[Client name]", 1) in got
    assert ("<Date>", 2) in got
    assert ("XXX", 6) in got
    assert all(len(p.text) <= 62 for p in corporate.placeholders)


def test_corporate_table_with_header_row(corporate):
    assert len(corporate.tables) == 1
    table = corporate.tables[0]
    assert table.index == 14
    assert table.n_rows == 2
    assert table.n_cols == 3
    assert table.header == ["Component", "Failure mode", "Criticality"]
    assert table.style == "Table Grid"


def test_corporate_header_and_footer_text(corporate):
    assert corporate.header_text == "ACME Energy Consulting"
    assert corporate.footer_text == "Confidential"


def test_corporate_paragraph_count_and_no_unsupported(corporate):
    assert corporate.n_paragraphs == 16
    assert corporate.unsupported == []


# ── unsupported content ──────────────────────────────────────────────────


def test_text_box_is_reported_unsupported(fixtures):
    outline = read_template(fixtures["with_textbox.docx"])
    assert "text box" in outline.unsupported
    assert outline.unsupported.count("text box") == 1
    # Everything else about the corporate template is still read.
    assert outline.mode == "untagged"
    assert outline.body_start_index == 5


def test_content_control_and_object_and_field_code_are_reported():
    doc = Document()
    doc.add_paragraph("before")
    sdt = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "controlled"
    r.append(t); p.append(r); content.append(p); sdt.append(content)
    doc.element.body.insert(1, sdt)
    obj_p = doc.add_paragraph()
    obj_p.add_run()._r.append(OxmlElement("w:object"))
    fld_p = doc.add_paragraph()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " AUTHOR ")
    fld_p._p.append(fld)
    page_p = doc.add_paragraph()
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), " PAGE ")
    page_p._p.append(page)
    buf = io.BytesIO(); doc.save(buf)
    outline = read_template(buf.getvalue())
    assert outline.unsupported == ["content control", "embedded object", "field code"]
    assert outline.has_toc is False


# ── language ─────────────────────────────────────────────────────────────


def test_detect_language_english(corporate):
    assert detect_language(corporate) == "en"
    assert corporate.language == "en"


def test_detect_language_german(fixtures):
    outline = read_template(fixtures["corporate_untagged_de.docx"])
    assert [h.text for h in outline.headings][1:3] == ["1 Zusammenfassung", "2 Einleitung"]
    assert detect_language(outline) == "de"
    assert outline.language == "de"


def test_detect_language_undecidable_is_none():
    empty = TemplateOutline(
        mode="untagged", language=None, headings=[], tags=[], placeholders=[],
        tables=[], header_text="", footer_text="", body_start_index=None,
        n_paragraphs=0, has_toc=False, unsupported=[])
    assert detect_language(empty) is None
    lorem = empty.model_copy(update={"headings": [
        TemplateHeading(index=0, level=1, text="Lorem ipsum", style="Heading 1",
                        is_body_start=True)]})
    assert detect_language(lorem) is None
    # One vote is not enough: "Summary" alone stays undecided.
    one = empty.model_copy(update={"headings": [
        TemplateHeading(index=0, level=1, text="Summary", style="Heading 1",
                        is_body_start=True)]})
    assert detect_language(one) is None


def test_tagged_fixture_language_is_none_or_en(tagged):
    # Jinja names carry no language; the one heading "Summary" is one vote.
    assert tagged.language in (None, "en")


# ── errors ───────────────────────────────────────────────────────────────


def test_png_bytes_raise_template_read_error():
    with pytest.raises(TemplateReadError) as info:
        read_template(_png_1x1())
    assert isinstance(info.value, ValueError)


def test_zip_that_is_not_a_docx_raises():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "not a word file")
    with pytest.raises(TemplateReadError):
        read_template(buf.getvalue())


def test_truncated_docx_raises(fixtures):
    data = fixtures["corporate_untagged.docx"]
    with pytest.raises(TemplateReadError):
        read_template(data[: len(data) // 2])


def test_empty_bytes_raise():
    with pytest.raises(TemplateReadError):
        read_template(b"")


# ── edge shapes ──────────────────────────────────────────────────────────


def test_empty_document_gives_an_untagged_outline():
    doc = Document()
    buf = io.BytesIO(); doc.save(buf)
    outline = read_template(buf.getvalue())
    assert outline.mode == "untagged"
    assert outline.headings == []
    assert outline.body_start_index is None
    assert outline.n_paragraphs == 0
    assert outline.tables == []
    assert outline.language is None


def test_body_start_falls_back_to_first_heading_without_toc():
    doc = Document()
    doc.add_paragraph("intro")
    doc.add_heading("Only a level 2", 2)
    doc.add_paragraph("text")
    buf = io.BytesIO(); doc.save(buf)
    outline = read_template(buf.getvalue())
    assert outline.has_toc is False
    assert outline.body_start_index == 1
    assert outline.headings[0].is_body_start is True


def test_outline_level_marks_a_heading_without_a_heading_style():
    doc = Document()
    p = doc.add_paragraph("Custom styled heading")
    lvl = OxmlElement("w:outlineLvl")
    lvl.set(qn("w:val"), "1")
    p._p.get_or_add_pPr().append(lvl)
    buf = io.BytesIO(); doc.save(buf)
    outline = read_template(buf.getvalue())
    assert [(h.level, h.text, h.style) for h in outline.headings] == [
        (2, "Custom styled heading", "Normal")]


def test_merged_paragraph_text_joins_runs_inside_hyperlink_wrappers():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("{{ meta.")
    link = OxmlElement("w:hyperlink")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "title"
    r.append(t)
    link.append(r)
    p._p.append(link)
    p.add_run(" }}")
    assert merged_paragraph_text(p) == "{{ meta.title }}"
    assert merged_paragraph_text(p._p) == "{{ meta.title }}"
    buf = io.BytesIO(); doc.save(buf)
    outline = read_template(buf.getvalue())
    assert [t.text for t in outline.tags] == ["{{ meta.title }}"]
    assert outline.mode == "tagged"


def test_if_endif_and_other_block_tags_are_classified():
    doc = Document()
    doc.add_paragraph("{% if fields.cost %}x{% endif %}{% set y = 1 %}")
    buf = io.BytesIO(); doc.save(buf)
    outline = read_template(buf.getvalue())
    assert [t.kind for t in outline.tags] == ["if", "endif", "other"]
    # if/endif alone do not make a template "tagged": nothing gets filled.
    assert outline.mode == "untagged"
