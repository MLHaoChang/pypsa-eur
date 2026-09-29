"""
WP6 — the report chat tools: `generate_report`, `get_report_status`,
`abort_report_generation`, `list_reports`, `get_report`, `get_report_table`,
`regenerate_report_section`, `export_report_docx`, `delete_report`.

Each is a thin wrapper over the WP1/WP3/WP5 routes (`routers/reports.py`,
`routers/report_jobs.py`) called in-process through `_route(...)` with the
active project, the way `run_eh_study` wraps `post_eh_study`. What these
tests pin beyond the route tests:

  * registration parity (TOOLS / DISPATCHERS / TOOL_ROUTES) and the safety
    tier of each — a generate on the read tier would spend minutes of model
    tokens with no confirmation card;
  * `get_report` stays under the chat harness's 4000-char result cap by
    construction: tables collapse to `{table_id, columns, n_rows, caption}`,
    figures to their id and caption, and a document whose prose alone would
    not fit degrades to a per-section outline with a hint to read one
    section at a time;
  * the route's own errors pass through with their `error_kind`, and the
    "no report yet" case names `generate_report` as the remedy.

The provider is the scripted `llm_fake` installed at the seam the route
looks up (`routers.report_jobs._provider_for_profile`), so no key and no
network are involved.
"""
from __future__ import annotations

import inspect
import io
import json
import threading
import time

import pytest
from docx import Document
from fastapi import HTTPException

from models.energy_hub import ReferenceDesignReport, SectionState, empty_section_map
from models.report import Paragraph, ReportDocument, Section, Table, TableRef
from services import chat_service, chat_tools as T, chat_tools_schema as S, upload_service
from services.adequacy import eh_report as R
from services.llm_fake import FakeProvider
from services.llm_provider import LLMEvent
from services.reports import store
from services.reports.assemble import EXECUTIVE_SUMMARY_ID

REPORT_TOOLS = {
    "generate_report": "execution",
    "get_report_status": "read",
    "abort_report_generation": "destructive",
    "list_reports": "read",
    "get_report": "read",
    "get_report_table": "read",
    "regenerate_report_section": "execution",
    "export_report_docx": "write",
    "delete_report": "destructive",
}

PROJECT = "report-tools"


# ── fixtures ────────────────────────────────────────────────────────────────

@pytest.fixture
def project(client, api_project, project_storage_dir, session_ctx):
    """
    A real project (row + org-scoped storage) that is also the ACTIVE one for
    the tools called from this thread.

    The active project is per SESSION: `api_project`'s save binds the client's
    session context to the project, and a tool called from the test thread
    without that binding reads the process foreground and sees no project at
    all (`no_active_project`). Binding the session's `ProjectContext` here is
    what a chat turn does before dispatching a tool (and what the QA drivers'
    `_as_session` does), so the tools, the job record and the EH report the
    tests store all resolve to the same context.
    """
    from services.pypsa_service import PyPSAService

    name = api_project(PROJECT)
    token = PyPSAService.bind_request_context(session_ctx(client))
    try:
        yield project_storage_dir(name)
    finally:
        PyPSAService.reset_request_context(token)


def _store_eh_report() -> None:
    """The EH report the tools read, in the state the test thread resolves to."""
    from routers.simulation import _state

    sections = empty_section_map(default="skipped")
    sections["fmea_top"] = SectionState(status="ok", payload={
        "top": [{"rank": 1, "mode_id": "gen:g:forced_outage",
                 "component_class": "Generator", "name": "g",
                 "failure_class": "A", "occurrence_per_year": 1.0,
                 "occurrence_basis": "FOR", "severity_eur": 10.0,
                 "criticality_eur_per_year": 10.0, "delta_eue_mwh": 0.5,
                 "engine": "copt", "fidelity": "analytic_convolution"}],
        "classes_included": ["A"], "note": "Link-primary residual risk",
    })
    report = ReferenceDesignReport(
        archetype="strong_grid", pack_hash="p", assumptions_hash="a",
        sections=sections, ens_cap_permyriad=10.0, mc_lole_h=3.21,
    )
    R.store_eh_report(_state, report)


def _turn(payload: dict | str) -> dict:
    text = payload if isinstance(payload, str) else json.dumps(payload)
    return {"events": [LLMEvent(type="text_delta", text=text)],
            "blocks": [{"type": "text", "text": text}]}


def _draft(section_id: str, *paragraphs: str) -> dict:
    return _turn({"section_id": section_id, "paragraphs": list(paragraphs),
                  "bullets": []})


def _script() -> list[dict]:
    return [_draft(EXECUTIVE_SUMMARY_ID, "The plan meets 3.2 h/yr."),
            _draft("fmea_top", "Generator g dominates.")]


def _install_provider(monkeypatch, provider):
    import routers.report_jobs as jobs
    monkeypatch.setattr(jobs, "_provider_for_profile", lambda profile: (provider, None))
    return provider


class _Holding(FakeProvider):
    def __init__(self, turns, hold: threading.Event):
        super().__init__(turns)
        self.hold = hold

    def stream(self, request):
        self.hold.wait(timeout=10)
        yield from super().stream(request)


def _poll(timeout: float = 10.0) -> dict:
    deadline = time.time() + timeout
    body: dict = {}
    while time.time() < deadline:
        body = T.get_report_status()
        if body.get("status") not in ("running", "no_data"):
            return body
        time.sleep(0.02)
    raise AssertionError(f"report job never finished: {body!r}")


def _big_doc(report_id: str = "0123456789abcdef", *, n_sections: int = 12,
             n_rows: int = 50, paragraph: str = "Short prose for section {i}.",
             mode: str = "generated", title: str = "Big report") -> ReportDocument:
    sections = [
        Section(section_id=f"s{i}", heading=f"Section {i}", source="llm", status="ok",
                blocks=[Paragraph(md=paragraph.format(i=i)),
                        TableRef(table_id="big" if i == 0 else f"t{i}", caption=f"T{i}")])
        for i in range(n_sections)
    ]
    tables = {"big": Table(table_id="big", columns=["rank", "mode", "value"],
                           rows=[[str(r), f"mode-{r}", f"{r}.0"] for r in range(n_rows)],
                           caption="The big one")}
    for i in range(1, n_sections):
        tables[f"t{i}"] = Table(table_id=f"t{i}", columns=["k", "v"], rows=[["a", "1"]])
    return ReportDocument(
        report_id=report_id, version=1, title=title,
        created_at="2026-09-28T10:00:00+00:00", evidence_hash="a" * 64,
        profile_id="fake-profile", model="fake-model", mode=mode,
        sections=sections, tables=tables, figures={},
    )


# ── registration ────────────────────────────────────────────────────────────

def test_every_report_tool_is_registered_with_its_tier():
    for name, tier in REPORT_TOOLS.items():
        assert any(t["name"] == name for t in S.TOOLS), f"{name} missing from TOOLS"
        assert callable(T.DISPATCHERS.get(name)), f"{name} missing from DISPATCHERS"
        assert name in S.TOOL_ROUTES, f"{name} missing from TOOL_ROUTES"
        assert S.safety_tier_for(name) == tier, f"{name}: tier {S.safety_tier_for(name)}"


def test_every_report_tool_maps_to_a_real_report_route():
    for name in REPORT_TOOLS:
        routes = S.TOOL_ROUTES[name]
        assert routes and all(isinstance(r, tuple) for r in routes), f"{name}: {routes}"
        assert all(path.startswith("/api/projects/{name}/reports") for _m, path in routes), (
            f"{name}: {routes}")


def test_schema_optional_fields_all_have_python_defaults():
    for name in REPORT_TOOLS:
        sch = next(t for t in S.TOOLS if t["name"] == name)["input_schema"]
        sig = inspect.signature(T.DISPATCHERS[name])
        for field in sch["properties"]:
            assert field in sig.parameters, f"{name}: {field!r} is not a parameter"
            if field not in sch["required"]:
                assert sig.parameters[field].default is not inspect.Parameter.empty, (
                    f"{name}: optional {field!r} has no Python default")


def test_the_report_tools_are_not_lock_gated_by_the_chat_seam():
    """`/api/projects/*` routes check the edit lock in the handler, as the docstring says."""
    assert not (set(REPORT_TOOLS) & T._lock_gated_tool_names())


def test_system_prompt_routes_a_document_to_generate_report_and_a_summary_to_build_study_report():
    guide = chat_service._ADEQUACY_GUIDE_CHAINING
    assert "generate_report" in guide and "build_study_report" in guide
    prompt = chat_service._build_system_prompt(chat_service.ChatSession())
    assert "generate_report" in prompt


def test_chatbot_md_documents_every_report_tool():
    import pathlib
    text = (pathlib.Path(__file__).resolve().parents[2] / "CHATBOT.md").read_text("utf-8")
    for name in REPORT_TOOLS:
        assert f"`{name}`" in text, f"CHATBOT.md does not document {name}"


# ── generate → status → done ────────────────────────────────────────────────

def test_status_before_any_run_is_no_data_and_abort_is_not_found(project):
    st = T.get_report_status()
    assert st["status"] == "no_data" and "generate_report" in st["message"]
    with pytest.raises(HTTPException) as exc:
        T.abort_report_generation()
    assert exc.value.status_code == 404
    assert exc.value.detail["error_kind"] == "report_job_not_found"


def test_generate_report_runs_to_done_through_the_dispatchers(project, monkeypatch):
    _store_eh_report()
    provider = _install_provider(monkeypatch, FakeProvider(_script()))

    started = T.DISPATCHERS["generate_report"](
        title="Client report", sections=[EXECUTIVE_SUMMARY_ID, "fmea_top"])
    assert started["status"] == "running"
    rid = started["report_id"]
    assert "get_report_status" in started["message"]

    body = _poll()
    assert body["status"] == "done", body
    assert body["report_id"] == rid and body["version"] == 1
    assert "thread" not in body and "stop_event" not in body
    assert body["repairs"] == 0 and body["prose_failures"] == []
    assert len(provider.requests) == 2 and provider.requests[0].tools == []

    listed = T.DISPATCHERS["list_reports"]()
    assert [m["report_id"] for m in listed] == [rid]
    assert listed[0]["mode"] == "generated" and listed[0]["latest_version"] == 1

    doc = T.DISPATCHERS["get_report"]()          # default: the newest report
    assert len(json.dumps(doc, default=str)) < 4000
    assert doc["report_id"] == rid and doc["version"] == 1
    assert doc["mode"] == "generated" and doc["title"] == "Client report"
    assert doc["profile_id"] == body["profile_id"] and doc["model"] == body["model"]
    assert "tables" not in doc and "figures" not in doc
    by_id = {s["section_id"]: s for s in doc["sections"]}
    # Sixteen sections (the fixture network adds an on-demand COPT surface)
    # do not fit in full: every code-only section is a row (no `blocks`), and
    # the prose sections are put back in full in document order while they
    # fit — the summary does, the FMEA section after it does not.
    assert doc["outline"] is True
    assert by_id["target"]["source"] == "code" and "blocks" not in by_id["target"]
    assert by_id["target"]["callouts"] == 1 and by_id["target"]["status"] == "skipped"
    assert "blocks" not in by_id["copt"] and by_id["copt"]["source"] == "code"
    summary = by_id[EXECUTIVE_SUMMARY_ID]
    assert summary["source"] == "llm" and "blocks" in summary
    assert summary["blocks"][0] == {"type": "paragraph", "md": "The plan meets 3.2 h/yr."}
    assert summary["audit"] == {"unverified": [], "verified": ["3.2 h/yr"]}
    assert by_id["fmea_top"]["source"] == "llm" and "blocks" not in by_id["fmea_top"]
    assert by_id["fmea_top"]["paragraphs"] == 1 and by_id["fmea_top"]["figures"] == ["fmea_pareto"]

    # One section at a time is always in full.
    one = T.get_report(rid, section_id="fmea_top")
    assert "outline" not in one and len(json.dumps(one, default=str)) < 4000
    fmea = one["sections"][0]
    assert fmea["source"] == "llm"
    assert fmea["blocks"][0] == {"type": "paragraph", "md": "Generator g dominates."}
    table = next(b for b in fmea["blocks"] if b["type"] == "table")
    assert table["table_id"] == "fmea_top" and table["n_rows"] == 1
    assert isinstance(table["columns"], list) and "rows" not in table
    figure = next(b for b in fmea["blocks"] if b["type"] == "figure")
    assert figure["figure_id"] == "fmea_pareto" and "png_file" not in figure
    assert T.get_report(rid, section_id=EXECUTIVE_SUMMARY_ID)["sections"] == [summary]


def test_generate_report_with_an_empty_section_list_means_the_default(project, monkeypatch):
    """A model that passes `sections: []` means "all of them", not a 422."""
    _store_eh_report()
    import routers.report_jobs as jobs
    seen: dict = {}

    def fake_start(project, **kwargs):
        seen.update(kwargs)
        return "0123456789abcdef"
    monkeypatch.setattr(jobs, "_start", fake_start)
    _install_provider(monkeypatch, FakeProvider([]))
    assert T.generate_report(sections=[])["report_id"] == "0123456789abcdef"
    assert seen["sections"] is None


def test_a_second_generate_while_running_is_in_flight_and_abort_ends_it(project, monkeypatch):
    _store_eh_report()
    hold = threading.Event()
    _install_provider(monkeypatch, _Holding(_script(), hold))
    try:
        rid = T.generate_report(sections=[EXECUTIVE_SUMMARY_ID, "fmea_top"])["report_id"]
        with pytest.raises(HTTPException) as exc:
            T.generate_report()
        assert exc.value.status_code == 409
        assert exc.value.detail["error_kind"] == "report_job_in_flight"
        st = T.get_report_status()
        assert st["status"] == "running" and st["report_id"] == rid
        assert T.abort_report_generation() == {"status": "running", "aborting": True}
    finally:
        hold.set()
    body = _poll()
    assert body["status"] == "aborted", body
    assert T.abort_report_generation()["aborting"] is False
    doc = T.get_report(rid)
    assert doc["report_id"] == rid and doc["version"] == 1


def test_no_evidence_and_seam_errors_pass_through(project, monkeypatch):
    import routers.report_jobs as jobs
    monkeypatch.setattr(jobs, "_current_study_report", lambda: None)
    _install_provider(monkeypatch, FakeProvider([]))
    with pytest.raises(HTTPException) as exc:
        T.generate_report()
    assert exc.value.status_code == 400
    assert exc.value.detail["error_kind"] == "no_evidence"

    monkeypatch.setattr(jobs, "_provider_for_profile", lambda profile: (None, "missing_api_key"))
    with pytest.raises(HTTPException) as exc:
        T.generate_report()
    assert exc.value.detail["error_kind"] == "missing_api_key"


# ── reading ─────────────────────────────────────────────────────────────────

def test_get_report_with_no_reports_names_generate_report(project):
    assert T.list_reports() == []
    with pytest.raises(HTTPException) as exc:
        T.get_report()
    assert exc.value.status_code == 404
    assert exc.value.detail["error_kind"] == "report_not_found"
    assert "generate_report" in exc.value.detail["message"]


@pytest.mark.parametrize("bad", ["../x", "0123456789ABCDEF", "short"])
def test_a_malformed_report_id_is_invalid_report_id(project, bad):
    with pytest.raises(HTTPException) as exc:
        T.get_report(bad)
    assert exc.value.status_code == 400
    assert exc.value.detail["error_kind"] == "invalid_report_id"


def test_an_unknown_report_or_version_is_report_not_found(project):
    store.create_report(project, _big_doc())
    with pytest.raises(HTTPException) as exc:
        T.get_report("ffffffffffffffff")
    assert exc.value.detail["error_kind"] == "report_not_found"
    with pytest.raises(HTTPException) as exc:
        T.get_report("0123456789abcdef", version=7)
    assert exc.value.detail["error_kind"] == "report_not_found"


def test_get_report_fits_the_cap_on_twelve_sections_and_a_fifty_row_table(project):
    store.create_report(project, _big_doc())
    doc = T.get_report("0123456789abcdef")
    assert len(json.dumps(doc, default=str)) < 4000
    assert len(doc["sections"]) == 12
    assert doc["sections"][3]["blocks"][0] == {"type": "paragraph",
                                               "md": "Short prose for section 3."}
    big = doc["sections"][0]["blocks"][1]
    assert big == {"type": "table", "table_id": "big", "columns": ["rank", "mode", "value"],
                   "n_rows": 50, "caption": "T0"}
    assert "get_report_table" in doc["message"]


def test_get_report_degrades_to_an_outline_when_the_prose_alone_would_not_fit(project):
    long = ("Section {i}: " + "x" * 900)
    store.create_report(project, _big_doc(paragraph=long))
    doc = T.get_report("0123456789abcdef")
    assert len(json.dumps(doc, default=str)) < 4000
    assert doc["outline"] is True and "section_id" in doc["message"]
    assert len(doc["sections"]) == 12
    # The first sections are put back in full while they fit; ~1 kB each
    # against a 3.9 kB budget means the fourth cannot be.
    assert "blocks" in doc["sections"][0]
    assert doc["sections"][0]["blocks"][0]["md"] == long.format(i=0)
    s3 = doc["sections"][3]
    assert s3["section_id"] == "s3" and s3["source"] == "llm"
    assert s3["paragraphs"] == 1 and s3["tables"] == ["t3"]
    assert "blocks" not in s3 and "heading" not in s3
    # One section at a time is always intact.
    one = T.get_report("0123456789abcdef", section_id="s3")
    assert len(json.dumps(one, default=str)) < 4000
    assert [s["section_id"] for s in one["sections"]] == ["s3"]
    assert one["sections"][0]["blocks"][0]["md"] == long.format(i=3)
    with pytest.raises(HTTPException) as exc:
        T.get_report("0123456789abcdef", section_id="nope")
    assert exc.value.status_code == 404
    assert exc.value.detail["error_kind"] == "report_section_not_found"


def test_get_report_table_pages_the_rows(project):
    store.create_report(project, _big_doc())
    page = T.DISPATCHERS["get_report_table"]("0123456789abcdef", "big", limit=20)
    assert page["table_id"] == "big" and page["columns"] == ["rank", "mode", "value"]
    assert page["caption"] == "The big one"
    assert page["total_count"] == 50 and page["returned"] == 20 and page["has_more"] is True
    assert page["items"][0] == ["0", "mode-0", "0.0"]
    last = T.get_report_table("0123456789abcdef", "big", offset=40, limit=20)
    assert last["returned"] == 10 and last["has_more"] is False
    assert last["items"][-1] == ["49", "mode-49", "49.0"]
    with pytest.raises(HTTPException) as exc:
        T.get_report_table("0123456789abcdef", "nope")
    assert exc.value.status_code == 404
    assert exc.value.detail["error_kind"] == "report_table_not_found"
    assert "big" in exc.value.detail["message"]


# ── regenerate ──────────────────────────────────────────────────────────────

def test_regenerate_section_saves_version_2_through_the_dispatchers(project, monkeypatch):
    _store_eh_report()
    _install_provider(monkeypatch, FakeProvider(_script()))
    rid = T.generate_report(sections=[EXECUTIVE_SUMMARY_ID, "fmea_top"])["report_id"]
    assert _poll()["status"] == "done"

    provider = _install_provider(monkeypatch, FakeProvider([
        _draft("fmea_top", "Generator g dominates, said shorter.")]))
    started = T.DISPATCHERS["regenerate_report_section"](rid, "fmea_top", instruction="shorter")
    assert started["status"] == "running" and started["report_id"] == rid
    assert "get_report_status" in started["message"]
    body = _poll()
    assert body["status"] == "done" and body["version"] == 2, body
    assert len(provider.requests) == 1
    v2 = T.get_report(rid, section_id="fmea_top")
    v1 = T.get_report(rid, version=1, section_id="fmea_top")
    assert v2["version"] == 2 and v1["version"] == 1
    assert v2["sections"][0]["blocks"][0]["md"] == "Generator g dominates, said shorter."
    assert v1["sections"][0]["blocks"][0]["md"] == "Generator g dominates."
    # The summary is untouched between the versions.
    s2 = T.get_report(rid, section_id=EXECUTIVE_SUMMARY_ID)["sections"]
    s1 = T.get_report(rid, version=1, section_id=EXECUTIVE_SUMMARY_ID)["sections"]
    assert s1 == s2

    with pytest.raises(HTTPException) as exc:
        T.regenerate_report_section(rid, "nope")
    assert exc.value.detail["error_kind"] == "report_section_not_found"


# ── export and delete ───────────────────────────────────────────────────────

def test_export_report_docx_returns_an_agent_export_chip(project):
    store.create_report(project, _big_doc())
    meta = T.DISPATCHERS["export_report_docx"]()       # default: the newest report
    assert meta["kind"] == "agent_export"
    assert meta["filename"] == "report_0123456789abcdef_v1.docx"
    assert meta["mime"].startswith("application/vnd.openxmlformats")
    assert meta["size"] > 0 and "file strip" in meta["message"]
    blob = upload_service.get_upload_bytes(PROJECT, meta["file_id"], project_dir=project)
    text = "\n".join(p.text for p in Document(io.BytesIO(blob)).paragraphs)
    assert "Big report" in text and "Section 3" in text

    # A byte-identical .docx is a dedup hit in the upload store (the first
    # chip comes back), so the named export is of a second report.
    store.create_report(project, _big_doc("fedcba9876543210", title="Second report"))
    named = T.export_report_docx("fedcba9876543210", filename="client-report")
    assert named["filename"] == "client-report.docx"
    assert named["report_id"] == "fedcba9876543210"

    with pytest.raises(HTTPException) as exc:
        T.export_report_docx("ffffffffffffffff")
    assert exc.value.detail["error_kind"] == "report_not_found"


def test_export_with_no_reports_names_generate_report(project):
    with pytest.raises(HTTPException) as exc:
        T.export_report_docx()
    assert exc.value.status_code == 404
    assert exc.value.detail["error_kind"] == "report_not_found"
    assert "generate_report" in exc.value.detail["message"]


def test_delete_report_removes_it(project):
    store.create_report(project, _big_doc())
    assert T.DISPATCHERS["delete_report"]("0123456789abcdef") == {
        "deleted": True, "report_id": "0123456789abcdef"}
    assert T.list_reports() == []
    with pytest.raises(HTTPException) as exc:
        T.delete_report("0123456789abcdef")
    assert exc.value.detail["error_kind"] == "report_not_found"
