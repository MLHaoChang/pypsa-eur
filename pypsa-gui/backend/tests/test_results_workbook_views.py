"""
The results workbook, package 02 slice R1 (`.scratch/results-workbook/issues/02-result-view-sheets.md`):
the network result views' sheets. Each view's key figures equal its tab's
route payload, every period is on every sheet, each scenario column comes
from that run's own network, every chart reads number cells on its own sheet,
a view with nothing to show is "no results" and a failing one is an error.
"""
from __future__ import annotations

import io
import re

import pandas as pd
import pypsa
import pytest
from openpyxl import load_workbook

from routers import results as results_router
from services import results_workbook as RW
from services.period_utils import snapshot_weights
from services.results.emissions import compute_emissions
from services.results_workbook import scenarios as S
from services.results_workbook.overview import overview_figures
from services.results_workbook.views import HORIZON
from services.solver_service import SolverConfig

URL = "/api/results/workbook.xlsx"
NETWORK_VIEWS = ["capex", "dispatch", "flows", "loadshape", "storage", "curtailment",
                 "lostload", "prices", "economics", "emissions"]
TITLES = {"capex": "Capacity Expansion", "dispatch": "Dispatch", "flows": "Energy flow",
          "loadshape": "Load shape", "storage": "Storage cycling", "curtailment": "Curtailment",
          "lostload": "Lost load", "prices": "Prices", "economics": "Economics",
          "emissions": "Emissions"}
P2 = "2040"  # the second period: every view must show it


def _network(load_scale: float = 1.0, gas_cost: float = 45.0) -> pypsa.Network:
    n = pypsa.Network()
    hours = pd.date_range("2030-01-01", periods=4, freq="h")
    n.set_snapshots(pd.MultiIndex.from_product([[2030, 2040], hours], names=["period", "timestep"]))
    n.investment_periods = [2030, 2040]
    # A non-unit weighting: an unweighted energy sum must not pass.
    n.snapshot_weightings.loc[:, :] = 2190.0
    n.investment_period_weightings.loc[:, "years"] = 10.0
    n.add("Carrier", ["gas", "solar", "AC"], co2_emissions=[0.2, 0.0, 0.0])
    n.add("Bus", "A", carrier="AC")
    n.add("Load", "L", bus="A", p_set=[x * load_scale for x in [50.0, 60.0, 70.0, 55.0, 52.0, 65.0, 72.0, 50.0]])
    n.add("Generator", "gas", bus="A", carrier="gas", p_nom_extendable=True,
          capital_cost=60_000.0, marginal_cost=gas_cost)
    # More solar than load at noon, so some is curtailed.
    n.add("Generator", "solar", bus="A", carrier="solar", p_nom=200.0, capital_cost=25_000.0,
          marginal_cost=0.5, p_max_pu=[0.0, 0.6, 0.9, 0.3, 0.1, 0.7, 0.8, 0.2])
    n.add("StorageUnit", "bat", bus="A", carrier="AC", p_nom=20.0, max_hours=2.0, marginal_cost=0.1)
    n.optimize(solver_name="highs", multi_investment_periods=True)
    return n


def _capture(n, mw: float = 3.0, voll: float = 3000.0) -> dict:
    ll = pd.DataFrame({"A": [0.0, mw, 0.0, 0.0, mw, 0.0, 0.0, 0.0]}, index=n.snapshots)
    w = snapshot_weights(n, "generators")
    total = float(ll["A"].mul(w).sum())
    return {"lost_load_t": ll, "lost_load_total_mwh": total, "lost_load_cost_eur": total * voll,
            "voll_eur_per_mwh": voll}


@pytest.fixture(scope="module")
def solved():
    return _network()


@pytest.fixture(scope="module")
def other():
    """A scenario whose results differ from the base's on every view."""
    return _network(load_scale=1.2, gas_cost=60.0)


def _rdf(n, accessor, attr, source="lopf"):
    return getattr(getattr(n, accessor), attr, None)


def _build(n, views=NETWORK_VIEWS, *, lost_load=None, scenarios=None, **kw) -> bytes:
    return RW.build_results_workbook(RW.WorkbookInputs(
        network=n, cfg=SolverConfig(), lost_load=lost_load, result_df=_rdf, views=list(views),
        project="Base", scenarios=scenarios,
        solved_settings=kw.pop("solved_settings", lambda: {"voll_eur_per_mwh": 3000.0}), **kw))


def _wb(blob):
    return load_workbook(io.BytesIO(blob))


def _key(ws, label: str, period: str, col: int = 4):
    for row in ws.iter_rows(min_row=4):
        if row[0].value == label and str(row[2].value) == period:
            return row[col - 1].value
    raise AssertionError(f"no key {label!r} / {period!r} on {ws.title}")


def _about(wb) -> dict:
    """View id → its About status."""
    by_label = {row[0]: row[2] for row in wb["About"].iter_rows(values_only=True) if row and row[0]}
    from services.result_views import view
    return {v: by_label.get(view(v).label) for v in NETWORK_VIEWS}


def _approx(a, b):
    return a == pytest.approx(b, rel=1e-9, abs=1e-9)


# ── the same numbers as the tab's route ──────────────────────────────────

def _period_rows(payload, period: int) -> list[int]:
    return [i for i, p in enumerate(payload["periods"]) if p == period]


def test_each_views_key_figures_equal_its_routes_payload(client, install_network, monkeypatch):
    import services.workflow_tools as wt

    n = _network()
    install_network(n)
    monkeypatch.setattr(wt, "solve_freshness",
                        lambda ctx=None: {"inputs_changed_since_solve": False, "results_current": True})
    try:
        r = client.get(URL, params={"views": ",".join(NETWORK_VIEWS)})
        assert r.status_code == 200, r.text
        wb = _wb(r.content)
        get = lambda path, **p: client.get(f"/api/results/{path}", params=p).json()  # noqa: E731
        w = snapshot_weights(n, "generators")
        w2 = [float(w.iloc[i]) for i in range(len(w)) if n.snapshots[i][0] == 2040]

        cb, ebc = get("cost_breakdown"), get("economics_by_carrier")
        ws = wb["Capacity Expansion"]
        assert _approx(_key(ws, "Capex (annualised)", HORIZON), cb["capex"])
        assert _approx(_key(ws, "Total system cost", HORIZON), cb["total"])
        cb2 = next(b for b in cb["by_period"] if b["period"] == 2040)
        assert _approx(_key(ws, "Capex (annualised)", P2), cb2["capex"])
        assert _approx(_key(ws, "Opex", P2), cb2["opex"])
        assert ebc["available"] and "Economics by carrier (MEUR, GWh, EUR/MWh)" in \
            [c.value for c in ws["A"]]

        gen = get("generators")
        rows, gas = _period_rows(gen, 2040), gen["columns"].index("gas")
        assert _approx(_key(wb["Dispatch"], "Generation — gas", P2),
                       sum(gen["data"][i][gas] * wi for i, wi in zip(rows, w2)))

        flows = get("energy_flows", period=2040)
        assert _approx(_key(wb["Energy flow"], "Generation", P2),
                       flows["conservation"]["generation_mwh"])

        shape = get("load_shape")
        a2 = next(a for a in shape["annual"] if a["period"] == 2040)
        assert _approx(_key(wb["Load shape"], "Load served", P2), a2["load_served"])

        # The tab's cycles: Σ |p| × w (per solved year) ÷ (2 × p_nom_opt × max_hours).
        sd = get("storage_dispatch")
        wy = n.snapshot_weightings.generators.loc[2040].tolist()
        tp = sum(abs(sd["data"][i][0]) * wi for i, wi in zip(_period_rows(sd, 2040), wy))
        assert tp > 0
        assert _approx(_key(wb["Storage cycling"], "Throughput — bat", P2), tp)
        assert _approx(_key(wb["Storage cycling"], "Cycles — bat", P2), tp / (2 * 20.0 * 2.0))

        curt, cause = get("curtailment"), get("curtailment/by_cause")
        assert _approx(_key(wb["Curtailment"], "Curtailed energy — solar", HORIZON),
                       cause["plants"]["solar"]["total_mwh"])
        sol = curt["columns"].index("solar")
        assert _key(wb["Curtailment"], "Curtailed energy — solar", P2) > 0
        assert _approx(_key(wb["Curtailment"], "Curtailed energy — solar", P2),
                       sum(curt["data"][i][sol] * wi for i, wi in zip(_period_rows(curt, 2040), w2)))
        prices = get("prices")
        p2 = [prices["data"][i][0] for i in _period_rows(prices, 2040)]
        assert _approx(_key(wb["Prices"], "Mean price — A", P2), sum(p2) / len(p2))
        assert _approx(_key(wb["Prices"], "Max price — A", P2), max(p2))

        ae = get("asset_economics")
        assets = [a for comp in ("generators", "storage_units", "stores", "links") for a in ae.get(comp) or []]
        assert _approx(_key(wb["Economics"], "Net profit (all assets)", P2),
                       sum(b["net_profit_eur"] for a in assets for b in a.get("by_period") or []
                           if b["period"] == 2040))

        em = get("emissions")
        assert _approx(_key(wb["Emissions"], "Total emissions", P2),
                       next(b["total_tCO2"] for b in em["by_period"] if b["period"] == 2040))
        assert _approx(_key(wb["Emissions"], "Total emissions", HORIZON), em["total_tCO2"])
        # A synthetic lost-load capture is not one the energy flows can balance,
        # so the Lost load sheet is read from a second export with it (the
        # network installed afresh, so the state proxy binds the same context).
        install_network(n)
        results_router._state["last_lost_load"] = _capture(n)
        r = client.get(URL, params={"views": "lostload"})
        assert r.status_code == 200, r.text
        wb_ll = _wb(r.content)
        ll = get("lost_load")
        assert ll["total_mwh"] is not None and ll["total_mwh"] > 0, ll
        assert _approx(_key(wb_ll["Lost load"], "Lost load", HORIZON), ll["total_mwh"])
        assert _approx(_key(wb_ll["Lost load"], "Lost-load cost", HORIZON), ll["total_cost_eur"])
        assert _approx(_key(wb_ll["Lost load"], "Shed hours (electrical buses)", P2),
                       ll["shed_hours"]["by_period"][P2])

    finally:
        results_router._state.pop("last_lost_load", None)


def test_derived_energy_is_weighted_not_a_plain_sum(solved):
    """Σ p·w with w = 2190 h × 10 years; Σ p alone (or w without the years) fails."""
    ws = _wb(_build(solved, ["dispatch"]))["Dispatch"]
    p = solved.generators_t.p.loc[2040, "gas"]
    assert _approx(_key(ws, "Generation — gas", P2), float(p.sum()) * 2190.0 * 10.0)


# ── every period ─────────────────────────────────────────────────────────

def test_every_view_shows_the_second_period_and_no_table_is_first_period_only(solved):
    wb = _wb(_build(solved))
    wb_ll = _wb(_build(solved, ["lostload"], lost_load=_capture(solved)))
    for view in NETWORK_VIEWS:
        ws = (wb_ll if view == "lostload" else wb)[TITLES[view]]
        periods_in_keys = {str(r[2].value) for r in ws.iter_rows(min_row=5) if r[1].value and r[2].value}
        assert P2 in periods_in_keys, view
        # Each long table with a Period column holds both periods.
        header = None
        seen: set = set()
        for row in ws.iter_rows(min_row=4, values_only=True):
            if row[0] == "Scenario":
                if header is not None and "Period" in header:
                    assert {"2030", "2040"} <= seen, (view, header, seen)
                header, seen = row, set()
            elif header is not None and row[0] is not None and "Period" in header:
                seen.add(str(row[header.index("Period")]))
        if header is not None and "Period" in header:
            assert {"2030", "2040"} <= seen, (view, header, seen)


# ── scenario columns ─────────────────────────────────────────────────────

def _run(n, name, **kw):
    return S.Run(name=name, figs=overview_figures(n, SolverConfig(), None, _rdf), network=n,
                 cfg=SolverConfig(), result_df=_rdf, solved_settings=lambda: None, **kw)


def test_each_scenario_column_comes_from_that_runs_own_network(solved, other):
    scen = S.Scenarios(runs=[_run(other, "High load"),
                             S.Run(name="Different horizon", reasons=["horizon_differs"])])
    wb = _wb(_build(solved, scenarios=scen))
    ws = wb["Emissions"]
    base, mine = compute_emissions(solved, "lopf", result_df=_rdf), compute_emissions(other, "lopf", result_df=_rdf)
    assert base["total_tCO2"] != pytest.approx(mine["total_tCO2"])  # so a base read would fail
    assert _approx(_key(ws, "Total emissions", HORIZON, col=4), base["total_tCO2"])
    assert _approx(_key(ws, "Total emissions", HORIZON, col=5), mine["total_tCO2"])
    assert _key(ws, "Total emissions", HORIZON, col=6) == "not_established"
    # The delta is a live formula; with no figures for a run there is none.
    assert re.fullmatch(r"=E(\d+)-D\1", _key(ws, "Total emissions", HORIZON, col=7))
    assert _key(ws, "Total emissions", HORIZON, col=8) == "not_established"
    assert "Different horizon: not comparable with the base (a different time horizon" in ws["A2"].value
    # The long tables carry the run's own rows under its name.
    disp = wb["Dispatch"]
    names = {r[0] for r in disp.iter_rows(min_row=5, values_only=True)}
    assert {"Base", "High load"} <= names and "Different horizon" not in names
    gas_other = [r for r in disp.iter_rows(min_row=5, values_only=True)
                 if r[0] == "High load" and r[1] == 2040 and r[2] is not None and len(r) > 3]
    assert gas_other


def test_every_scenario_view_reads_its_own_run(solved, other, monkeypatch):
    """Mutation: a builder fed the base's network for every run must fail this."""
    from services.results_workbook import views_network as VN
    seen = []
    real = VN._ts

    def spy(run, accessor, attr):
        seen.append((run.name, id(run.network)))
        return real(run, accessor, attr)
    monkeypatch.setattr(VN, "_ts", spy)
    _build(solved, ["dispatch", "storage"], scenarios=S.Scenarios(runs=[_run(other, "High load")]))
    assert ("Base", id(solved)) in seen and ("High load", id(other)) in seen
    assert all(nid == (id(solved) if name == "Base" else id(other)) for name, nid in seen)


# ── charts ───────────────────────────────────────────────────────────────

def _cells(ref: str) -> tuple[str, list[str]]:
    from openpyxl.utils import range_boundaries
    from openpyxl.utils.cell import get_column_letter as L
    sheet, rng = ref.rsplit("!", 1)
    c0, r0, c1, r1 = range_boundaries(rng.replace("$", ""))
    return sheet.strip("'"), [f"{L(c)}{r}" for c in range(c0, c1 + 1) for r in range(r0, r1 + 1)]


@pytest.mark.parametrize("view", NETWORK_VIEWS)
def test_every_chart_reads_number_cells_on_its_own_sheet(solved, view):
    ws = _wb(_build(solved, [view], lost_load=_capture(solved) if view == "lostload" else None)
             )[TITLES[view]]
    assert 1 <= len(ws._charts) <= 3, view
    for chart in ws._charts:
        for s in chart.series:
            for ref in filter(None, [s.val.numRef.f if s.val else None,
                                     s.xVal.numRef.f if getattr(s, "xVal", None) and s.xVal.numRef else None,
                                     s.yVal.numRef.f if getattr(s, "yVal", None) else None]):
                sheet, cells = _cells(ref)
                assert sheet == ws.title, ref
                for c in cells:
                    v = ws[c].value
                    assert isinstance(v, (int, float)) and not isinstance(v, bool), (view, c, v)
            if s.cat is not None:
                assert _cells((s.cat.numRef or s.cat.strRef).f)[0] == ws.title


# ── no results vs error ──────────────────────────────────────────────────

def test_a_view_with_nothing_to_show_is_no_results_and_has_no_sheet():
    n = pypsa.Network()
    n.set_snapshots(pd.date_range("2030-01-01", periods=2, freq="h"))
    n.add("Bus", "A")
    n.add("Load", "L", bus="A", p_set=[1.0, 2.0])
    n.add("Generator", "g", bus="A", p_nom=10.0, marginal_cost=1.0)
    n.optimize(solver_name="highs")
    wb = _wb(_build(n, ["storage", "dispatch"], solved_settings=lambda: None))
    assert "Storage cycling" not in wb.sheetnames and "Dispatch" in wb.sheetnames
    about = _about(wb)
    assert "no results" in str(about["storage"])


def test_a_failing_view_is_an_error_in_about_not_no_results(solved, monkeypatch):
    from services.results import emissions
    monkeypatch.setattr(emissions, "compute_emissions",
                        lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom")))
    wb = _wb(_build(solved, ["emissions", "storage"]))
    about = _about(wb)
    assert "error" in str(about["emissions"]) and "no results" not in str(about["emissions"])
    assert "Emissions" not in wb.sheetnames and "Storage cycling" in wb.sheetnames


def test_formula_like_names_stay_text():
    evil = '=HYPERLINK("http://x","y")'
    n = _network()
    n.generators.loc["gas", "carrier"] = evil
    ws = _wb(_build(n, ["dispatch"]))["Dispatch"]
    label = next(c for row in ws.iter_rows() for c in row if c.value == f"Generation — {evil}")
    assert label.data_type == "s"
    for row in ws.iter_rows():
        for c in row:
            if c.data_type == "f":
                assert re.fullmatch(r"=[A-Z]+\d+-[A-Z]+\d+", c.value), c.value


def test_a_period_total_with_an_unknown_part_is_not_established(solved, monkeypatch):
    """ADR-0001: no partial sum stands in for a total."""
    from services.results import asset_economics
    real = asset_economics.compute_asset_economics

    def blanked(*a, **k):
        out = real(*a, **k)
        out["generators"][0]["by_period"][1]["net_profit_eur"] = None
        return out
    monkeypatch.setattr(asset_economics, "compute_asset_economics", blanked)
    ws = _wb(_build(solved, ["economics"]))["Economics"]
    assert _key(ws, "Net profit (all assets)", P2) == "not_established"
    assert isinstance(_key(ws, "Net profit (all assets)", "2030"), (int, float))
