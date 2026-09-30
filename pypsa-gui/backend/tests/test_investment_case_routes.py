"""
Edge Investment Case P4 WP4.6b — the finance inputs route, the investment-case
study routes, persistence, staleness and the P3 hygiene items.

Plan: docs/superpowers/plans/2026-09-30-edge-investment-case-p4.md WP4.6b

The adapter (`services.results.finance_case.build_finance_case`) is replaced by
a fake `build_case` built from the SAM oracle case S2
(`tests.fixtures.investment_case.sam.sam_case.to_finance_case("s2")`): the
router's closure factory `_ic_build_case` is the seam (plan C1 — the runner
takes `build_case` injected). S2 has no `tax_pack_id`, so the HTTP report has
the tax section not established; the runner tests run S2 with SAM's layers.
"""
from __future__ import annotations

import io
import json
import threading
import time
import zipfile
from contextlib import contextmanager

import pytest

from tests.conftest import build_network
from tests.fixtures.investment_case.sam import sam_case as S

FIN_URL = "/api/simulation/finance"
CFG_URL = "/api/simulation/solver_config"
STUDY_URL = "/api/results/investment_case"
ABORT_URL = "/api/results/investment_case/abort"
REPORT_URL = "/api/results/investment_case/report"


def _s2_inputs() -> dict:
    return S.to_finance_case("s2").inputs.model_dump(mode="json")


def _solved(install_network, name: str | None = None):
    return install_network(build_network(solve=True), name=name)


def _put_fin(client, fin, if_match=None):
    headers = {"If-Match": if_match} if if_match is not None else {}
    return client.put(FIN_URL, json={"finance": fin}, headers=headers)


@contextmanager
def _live_thread(state: dict, key: str, record: bool = True):
    """A live fake solve (`key="thread"`) or study record under `key`."""
    release = threading.Event()
    t = threading.Thread(target=release.wait, daemon=True, name=f"fake-{key}")
    t.start()
    prev = state.get(key)
    if record:
        state[key] = {"status": "running", "error": None, "started_at": time.time(),
                      "thread": t, "stop_event": threading.Event()}
    else:
        state[key] = t
    try:
        yield
    finally:
        release.set()
        t.join(timeout=5)
        state[key] = prev if not record else None


def _fake_build(monkeypatch, *, hold: threading.Event | None = None, refuse: str | None = None,
                case_fn=None):
    """Replace the router's adapter closure factory (plan C1 seam)."""
    import routers.results as R
    from services.finance.case import FinanceRefused

    calls: list[dict] = []

    def factory(n, cfg, *, owner, lost_load):
        calls.append({"owner": owner, "finance": cfg.finance})

        def build():
            if hold is not None:
                assert hold.wait(timeout=30.0)
            if refuse:
                raise FinanceRefused(refuse, "stated by the test")
            return (case_fn or (lambda: S.to_finance_case("s2")))()
        return build

    monkeypatch.setattr(R, "_ic_build_case", factory)
    monkeypatch.setattr(R, "_ic_case_hash", lambda case: "case-hash-for-test")
    return calls


def _poll(client, until=lambda b: b.get("status") != "running", timeout: float = 30.0) -> dict:
    deadline = time.time() + timeout
    body: dict = {}
    while time.time() < deadline:
        r = client.get(STUDY_URL)
        if r.status_code == 200:
            body = r.json()
            if until(body):
                return body
        time.sleep(0.05)
    raise AssertionError(f"investment case did not settle: {body}")


# ── the finance inputs route ─────────────────────────────────────────────────


def test_get_before_any_put_is_not_set(client, install_network):
    install_network(build_network())
    r = client.get(FIN_URL)
    assert r.status_code == 200
    body = r.json()
    assert body["finance"] is None and body["status"] == "not_set" and body["digest"]


def test_put_validates_stores_and_the_digest_guards_the_edit(client, install_network,
                                                             session_state):
    install_network(build_network())
    d0 = client.get(FIN_URL).json()["digest"]
    fin = _s2_inputs()
    r = _put_fin(client, fin, if_match=d0)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["status"] == "ok" and body["finance"] == fin
    assert session_state(client)["solver_config"].finance == fin
    assert client.get(FIN_URL).json() == body
    # A stale If-Match (the digest before the edit) is refused, nothing stored.
    r = _put_fin(client, {**fin, "analysis_years": 10}, if_match=d0)
    assert r.status_code == 412, r.text
    assert r.json()["detail"]["code"] == "finance_changed"
    assert session_state(client)["solver_config"].finance["analysis_years"] == fin["analysis_years"]
    # The current digest (weak and quoted forms too) is accepted.
    r = _put_fin(client, {**fin, "analysis_years": 10}, if_match=f'W/"{body["digest"]}"')
    assert r.status_code == 200, r.text
    assert r.json()["finance"]["analysis_years"] == 10
    # null clears.
    r = _put_fin(client, None)
    assert r.status_code == 200 and r.json()["status"] == "not_set"
    assert session_state(client)["solver_config"].finance is None


def test_put_refuses_invalid_inputs_with_field_paths(client, install_network, session_state):
    install_network(build_network())
    bad = {**_s2_inputs(), "capex_phasing": [0.5, 0.2]}
    r = _put_fin(client, bad)
    assert r.status_code == 422, r.text
    detail = r.json()["detail"]
    assert detail["code"] == "finance_inputs_invalid"
    assert "capex_phasing" in detail["message"]
    r = _put_fin(client, {**_s2_inputs(), "analysis_years": 0})
    assert r.status_code == 422
    assert ["analysis_years"] in [e["loc"] for e in r.json()["detail"]["errors"]]
    assert session_state(client)["solver_config"].finance is None
    # An unknown (mistyped) key is refused, not silently dropped (review B3).
    r = _put_fin(client, {**_s2_inputs(), "terminalvalue": {"method": "fixed", "value": 5}})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "finance_inputs_invalid"
    assert session_state(client)["solver_config"].finance is None


def test_put_body_is_closed_and_required(client, install_network):
    install_network(build_network())
    assert client.put(FIN_URL, json={}).status_code == 422
    assert client.put(FIN_URL, json=_s2_inputs()).status_code == 422          # unwrapped
    assert client.put(FIN_URL, json={"finance": None, "x": 1}).status_code == 422


def test_put_is_refused_while_a_solve_is_in_flight(client, install_network, session_state):
    install_network(build_network())
    with _live_thread(session_state(client), "thread", record=False):
        r = _put_fin(client, _s2_inputs())
        assert r.status_code == 409, r.text
        assert r.json()["detail"]["code"] == "solver_in_flight"
    assert session_state(client)["solver_config"].finance is None


def test_the_solver_config_put_keeps_finance_and_refuses_a_change(client, install_network,
                                                                   session_state):
    install_network(build_network())
    fin = _s2_inputs()
    assert _put_fin(client, fin).status_code == 200
    # A partial PUT without the key keeps it.
    r = client.put(CFG_URL, json={"voll": 1234.0})
    assert r.status_code == 200, r.text
    assert r.json()["finance"] == fin
    # A full echo of the GET (the settings form) is accepted unchanged.
    full = client.get(CFG_URL).json()
    assert full["finance"] == fin
    r = client.put(CFG_URL, json=full)
    assert r.status_code == 200, r.text
    assert session_state(client)["solver_config"].finance == fin
    # A change — or a null — is refused, nothing written.
    for changed in ({**fin, "analysis_years": 3}, None):
        r = client.put(CFG_URL, json={"voll": 99.0, "finance": changed})
        assert r.status_code == 422, r.text
        assert r.json()["detail"]["code"] == "finance_via_dedicated_route"
    cfg = session_state(client)["solver_config"]
    assert cfg.finance == fin and cfg.voll == 1234.0


def test_finance_is_in_no_solve_fingerprint(client, install_network, session_ctx):
    """A finance edit never marks the dispatch stale: dispatch freshness is
    topology-based (`dispatch_status`), and the one config hash a solve
    records (the adequacy report's `assumptions_hash`) drops `finance`."""
    import dataclasses

    from services.adequacy.report import _config_hash
    from services.dispatch_status import dispatch_digest, dispatch_status

    _solved(install_network)
    ctx = session_ctx(client)
    n = ctx.network
    assert dispatch_status(n) == "fresh"
    before_cfg = ctx.solver_state["solver_config"]
    before_digest = dispatch_digest(n)
    assert _put_fin(client, _s2_inputs()).status_code == 200
    after_cfg = session_ctx(client).solver_state["solver_config"]
    assert after_cfg.finance is not None
    assert dispatch_status(n) == "fresh" and n.is_solved
    assert dispatch_digest(n) == before_digest
    assert _config_hash(after_cfg) == _config_hash(before_cfg)
    assert _config_hash(dataclasses.replace(before_cfg, voll=before_cfg.voll + 1)) \
        != _config_hash(before_cfg)                                         # the hash bites


# ── persistence ──────────────────────────────────────────────────────────────


def test_finance_persists_with_the_project_and_round_trips_the_bundle(
        client, install_network, session_ctx):
    install_network(build_network(), name="fin_persist")
    assert client.post("/api/projects/fin_persist",
                       params={"force": True, "rebind": True}).status_code == 200
    fin = _s2_inputs()
    assert _put_fin(client, fin).status_code == 200
    assert client.post("/api/projects/fin_persist",
                       params={"expect": "fin_persist"}).status_code == 200
    # Switch away (a fresh scratch network, default config), then load: the
    # finance inputs come back from disk.
    install_network(build_network())
    assert session_ctx(client).solver_state["solver_config"].finance is None
    assert client.get("/api/projects/fin_persist").status_code == 200
    assert session_ctx(client).solver_state["solver_config"].finance == fin
    assert client.get(FIN_URL).json()["finance"] == fin
    # Bundle export → import as another project.
    b = client.get("/api/projects/fin_persist/bundle")
    assert b.status_code == 200, b.text
    cfg = json.loads(zipfile.ZipFile(io.BytesIO(b.content)).read("solver_config.json"))
    assert cfg["finance"] == fin
    r = client.post("/api/projects/import_bundle?name=fin_imported",
                    files={"file": ("b.zip", b.content, "application/zip")})
    assert r.status_code == 200, r.text
    assert session_ctx(client).solver_state["solver_config"].finance == fin
    assert client.get(FIN_URL).json()["status"] == "ok"


# ── P3 hygiene ───────────────────────────────────────────────────────────────


def _edge(install_network, name=None):
    from tests.fixtures.investment_case.edge_15min import build_edge_15min

    return install_network(build_edge_15min(), name=name)


def _firm(fee_periods):
    return {"poc_link": "import", "timezone": "UTC",
            "connection": {"kind": "firm", "import_cap_mw": 60.0,
                           "available_from": "2030-01-01",
                           "capacity_fee": {"id": "cap_fee", "kind": "capacity",
                                            "unit": "per_kw_year", "periods": fee_periods}}}


def test_a_windowed_capacity_fee_is_refused_at_save_not_at_the_solve(client, install_network,
                                                                     session_state):
    _edge(install_network)
    r = client.put(CFG_URL, json={"commercial": _firm([{"name": "winter", "rate": 50.0,
                                                        "months": [1, 2]}])})
    assert r.status_code == 422, r.text
    detail = r.json()["detail"]
    assert detail["code"] == "connection_capacity_fee_unsupported"
    assert "time windows" in detail["message"]
    assert session_state(client)["solver_config"].commercial is None
    # The supported shape still saves.
    r = client.put(CFG_URL, json={"commercial": _firm([{"name": "all", "rate": 50.0}])})
    assert r.status_code == 200, r.text


def test_a_per_kwh_capacity_fee_is_refused_at_save(client, install_network):
    _edge(install_network)
    body = _firm([{"name": "all", "rate": 0.05}])
    body["connection"]["capacity_fee"]["unit"] = "per_kwh"
    r = client.put(CFG_URL, json={"commercial": body})
    assert r.status_code == 422, r.text
    assert r.json()["detail"]["code"] == "connection_capacity_fee_unsupported"
    assert "per_kwh" in r.json()["detail"]["message"]


def _tariff() -> dict:
    import pathlib

    return json.loads((pathlib.Path(__file__).parent / "fixtures" / "investment_case"
                       / "de_tariff_capacity_tou.json").read_text())


def test_tariff_ref_load_issue_unit():
    from models.commercial import CommercialConfig
    from services.commercial import hashing as H
    from services.commercial.binding import tariff_ref_load_issue

    tariff = _tariff()
    good = H.library_item_digest(CommercialConfig.model_validate(
        {"poc_link": "import", "import_tariff": tariff}).import_tariff)
    ref = {"kind": "tariff", "id": "t", "version": 1, "hash": good}
    base = {"poc_link": "import", "import_tariff": tariff, "import_tariff_ref": ref}
    assert tariff_ref_load_issue(base) is None
    assert tariff_ref_load_issue({"poc_link": "import", "import_tariff_ref": ref}) is None
    assert tariff_ref_load_issue(None) is None
    edited = json.loads(json.dumps(tariff))
    edited["items"][0]["periods"][0]["rate"] = 0.99
    issue = tariff_ref_load_issue({**base, "import_tariff": edited})
    assert issue["code"] == "import_tariff_ref_conflict_on_load"
    assert issue["reason"] == "inline_differs_from_ref" and issue["id"] == "t"


def test_the_load_path_rechecks_the_tariff_ref_hash(client, install_network, session_ctx):
    """A project whose saved inline tariff no longer hashes to its ref (edited
    on disk, or a bundle) is flagged on load; the frontend's `replacesInline`
    trusted the PUT, which a file on disk never went through."""
    _edge(install_network, name="tariff_ref_load")
    assert client.post("/api/projects/tariff_ref_load",
                       params={"force": True, "rebind": True}).status_code == 200
    edited = _tariff()
    edited["items"][0]["periods"][0]["rate"] = 0.99
    session_ctx(client).solver_state["solver_config"].commercial = {
        "poc_link": "import", "timezone": "UTC", "import_tariff": edited,
        "import_tariff_ref": {"kind": "tariff", "id": "t", "version": 1, "hash": "0" * 64}}
    assert client.post("/api/projects/tariff_ref_load",
                       params={"expect": "tariff_ref_load"}).status_code == 200
    r = client.get("/api/projects/tariff_ref_load")
    assert r.status_code == 200, r.text
    codes = [i["code"] for i in r.json()["library_issues"]]
    assert "import_tariff_ref_conflict_on_load" in codes, r.json()["library_issues"]


# ── the study routes ─────────────────────────────────────────────────────────


def test_204_before_a_run(client, install_network):
    install_network(build_network())
    assert client.get(STUDY_URL).status_code == 204
    assert client.get(REPORT_URL).status_code == 204
    assert client.get("/api/results/investment_case/export.xlsx").status_code == 204
    r = client.post(ABORT_URL)
    assert r.status_code == 404


def test_post_refusals(client, install_network, monkeypatch):
    _fake_build(monkeypatch)
    install_network(build_network())
    r = client.post(STUDY_URL, json={})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "not_solved"
    _solved(install_network)
    r = client.post(STUDY_URL, json={})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "finance_inputs_missing"
    assert _put_fin(client, {**_s2_inputs(), "tax_pack_id": "xx_nowhere"}).status_code == 200
    r = client.post(STUDY_URL, json={})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "tax_pack_not_found"
    assert client.post(STUDY_URL, json={"bogus": 1}).status_code == 422


def test_409_during_a_solve_or_another_study(client, install_network, session_state,
                                             monkeypatch):
    _fake_build(monkeypatch)
    _solved(install_network)
    assert _put_fin(client, _s2_inputs()).status_code == 200
    with _live_thread(session_state(client), "thread", record=False):
        r = client.post(STUDY_URL, json={})
        assert r.status_code == 409, r.text
        assert "solve" in r.json()["detail"]
    with _live_thread(session_state(client), "eh_study"):
        r = client.post(STUDY_URL, json={})
        assert r.status_code == 409, r.text
        assert "Energy Hub" in r.json()["detail"]
    assert session_state(client).get("investment_case") is None


def test_a_run_stores_the_report_and_it_is_current(client, install_network, session_state,
                                                   monkeypatch):
    from services.finance.engine import run_case

    calls = _fake_build(monkeypatch)
    _solved(install_network)
    fin = _s2_inputs()
    assert _put_fin(client, fin).status_code == 200
    r = client.post(STUDY_URL, json={"owner": "owner"})
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "running"
    body = _poll(client)
    assert body["status"] == "done", body
    assert body["stages_done"] == ["build_case", "load_pack", "run_engine", "assemble", "store"]
    assert "thread" not in body and "stop_event" not in body
    assert calls == [{"owner": "owner", "finance": fin}]
    assert body["report"]["present"] and body["report"]["stale"] is False

    rep = client.get(REPORT_URL)
    assert rep.status_code == 200, rep.text
    out = rep.json()
    expected = run_case(S.to_finance_case("s2")).metrics
    assert out["project_irr_pre_tax"] == pytest.approx(expected["project_pre_tax_irr"])
    assert out["min_dscr"] == pytest.approx(expected["min_dscr"])
    assert out["llcr"] == pytest.approx(expected["llcr"])
    # No tax pack → the tax section and every post-tax headline are not established.
    assert out["project_irr_post_tax"] is None and out["npv_at_wacc"] is None
    assert out["lcoe_finance_consistent_eur_per_mwh"] is None
    assert out["completeness"]["tax"] == "not_established"
    assert out["completeness"]["project"] == "ok" and out["completeness"]["debt"] == "ok"
    for name in ("design", "commercial", "dispatch_modes", "tax_equity", "uncertainty"):
        assert out["completeness"][name] == "skipped"
    stored = session_state(client)["investment_case_report"]
    assert isinstance(stored, dict) and stored["excludes_shed_cost"] is True
    prov = stored["sections"]["project"]["payload"]["provenance"]
    assert prov["finance_case_hash"] == "case-hash-for-test"
    assert set(prov["inputs"]) == {"finance", "value_flows", "commercial", "solver_config",
                                   "dispatch", "packs"}
    assert stored["cashflow_lines"], "cashflow lines missing"
    # detail=full: the sections with payloads and the cashflow lines (WP4.7b).
    full = client.get(REPORT_URL, params={"detail": "full"}).json()
    assert full["sections"]["debt"]["status"] == "ok" and full["cashflow_lines"]
    assert full["project_irr_pre_tax"] == out["project_irr_pre_tax"]
    # IC P4 WP4.6d: the stored report as a workbook.
    import io

    import openpyxl
    x = client.get("/api/results/investment_case/export.xlsx")
    assert x.status_code == 200 and "spreadsheetml" in x.headers["content-type"]
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    about = {r[0].value: r[1].value for r in wb["About"].iter_rows() if r[0].value}
    assert about["Assumptions hash"] == stored["assumptions_hash"]
    assert "CashflowLines" in wb.sheetnames


def test_staleness_follows_each_input(client, install_network, session_ctx, monkeypatch):
    import dataclasses

    _fake_build(monkeypatch)
    _solved(install_network)
    fin = _s2_inputs()
    assert _put_fin(client, fin).status_code == 200
    assert client.post(STUDY_URL, json={}).status_code == 200
    assert _poll(client)["report"]["stale"] is False

    # 1. a finance edit
    assert _put_fin(client, {**fin, "analysis_years": 20}).status_code == 200
    rep = client.get(STUDY_URL).json()["report"]
    assert rep["stale"] is True and rep["changed"] == ["finance"]
    assert _put_fin(client, fin).status_code == 200
    assert client.get(STUDY_URL).json()["report"]["stale"] is False

    # 2. the value flows (and 3. the rest of the commercial config)
    st = session_ctx(client).solver_state
    cfg0 = st["solver_config"]
    st["solver_config"] = dataclasses.replace(
        cfg0, commercial={"poc_link": "x", "value_flows": {"participants": []}})
    rep = client.get(STUDY_URL).json()["report"]
    assert rep["stale"] is True and set(rep["changed"]) == {"value_flows", "commercial"}
    st["solver_config"] = dataclasses.replace(
        cfg0, commercial={"poc_link": "x", "value_flows": None})
    assert set(client.get(STUDY_URL).json()["report"]["changed"]) == {"commercial"}
    st["solver_config"] = cfg0
    assert client.get(STUDY_URL).json()["report"]["stale"] is False

    # 3b. the solver config the case reads at build time (WP4.6b review B1)
    for field, val in (("discount_rate", 0.03), ("voll", 9000.0), ("dsr_share_of_load", 0.2)):
        st["solver_config"] = dataclasses.replace(cfg0, **{field: val})
        rep = client.get(STUDY_URL).json()["report"]
        assert rep["stale"] is True and rep["changed"] == ["solver_config"], field
    st["solver_config"] = cfg0
    assert client.get(STUDY_URL).json()["report"]["stale"] is False

    # 4. the dispatch (a re-solve or a restored result)
    n = session_ctx(client).network
    n.generators_t.p.iloc[0, 0] += 1.0
    rep = client.get(STUDY_URL).json()["report"]
    assert rep["stale"] is True and rep["changed"] == ["dispatch"]
    n.generators_t.p.iloc[0, 0] -= 1.0

    # 5. a pack: the stored hash names a pack version the current one differs from
    st["investment_case_report"]["assumptions_hash"] = "0" * 32
    rep = client.get(STUDY_URL).json()["report"]
    assert rep["stale"] is True


def test_a_stored_report_without_a_record_is_served_and_checked(client, install_network,
                                                                session_ctx, monkeypatch):
    """After a project reload the report is back (result state) and the study
    record is not (study records are never persisted): GET still answers."""
    _fake_build(monkeypatch)
    _solved(install_network)
    assert _put_fin(client, _s2_inputs()).status_code == 200
    assert client.post(STUDY_URL, json={}).status_code == 200
    _poll(client)
    st = session_ctx(client).solver_state
    st["investment_case"] = None
    body = client.get(STUDY_URL).json()
    assert body["status"] == "idle" and body["report"]["stale"] is False


def test_a_refused_case_is_a_status_and_a_report_never_a_500(client, install_network,
                                                             monkeypatch):
    _fake_build(monkeypatch, refuse="template_not_annual")
    _solved(install_network)
    assert _put_fin(client, _s2_inputs()).status_code == 200
    assert client.post(STUDY_URL, json={}).status_code == 200
    body = _poll(client)
    assert body["status"] == "refused" and body["error_code"] == "template_not_annual"
    out = client.get(REPORT_URL).json()
    assert out["project_irr_pre_tax"] is None and out["min_dscr"] is None
    assert out["completeness"]["project"] == "not_established"


def test_abort_mid_run_stores_nothing_and_holds_the_mesh_meanwhile(client, install_network,
                                                                   session_state, monkeypatch):
    hold = threading.Event()
    _fake_build(monkeypatch, hold=hold)
    _solved(install_network)
    assert _put_fin(client, _s2_inputs()).status_code == 200
    try:
        assert client.post(STUDY_URL, json={}).status_code == 200
        running = _poll(client, until=lambda b: b.get("stage") == "build_case")
        assert running["status"] == "running"
        # The run holds the study mesh: a second run, a solve and another study wait.
        r = client.post(STUDY_URL, json={})
        assert r.status_code == 409 and "already running" in r.json()["detail"]
        r = client.post("/api/simulation/run")
        assert r.status_code == 409, r.text
        assert "investment-case run" in json.dumps(r.json()["detail"])
        r = client.post(ABORT_URL)
        assert r.status_code == 200 and r.json() == {"status": "running", "aborting": True}
    finally:
        hold.set()
    body = _poll(client)
    assert body["status"] == "aborted", body
    assert "run_engine" not in body["stages_done"]
    assert session_state(client).get("investment_case_report") is None
    assert client.get(REPORT_URL).status_code == 204
    # Idempotent after the end.
    r = client.post(ABORT_URL)
    assert r.status_code == 200 and r.json()["aborting"] is False

