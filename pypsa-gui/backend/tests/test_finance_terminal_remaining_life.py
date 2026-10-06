"""
IC D10 (plan `docs/superpowers/plans/2026-10-06-ic-s0b-replacements-terminal.md` §4 WP-C; S6): the
`remaining_life_annuity` terminal value — GS's `annuity_pv` salvage, computed by the engine.

At the last operating year, TV = Σ over assets and parts of base × annuity(r_a, L) ×
annuity_pv_factor(r, remaining): the purchase of each part still alive at the horizon, valued on the annuity
the LP charged for it. `remaining = start + L − (cod_year + analysis_years)`, `start` the first service year
of the part's last purchase with the UNROUNDED k·L (GS's `last_buy + L − H`); `r_a` the asset's own rate,
else the LP's; `r` the LP's.

The identity (GS's `BY_CONSTRUCTION`): on a flat case with no escalation and no tax, WACC = the LP rate,
financial close one year before COD, `capex_phasing = [1.0]`, `analysis_years = H` and
`replacement_rule = "part_lifetimes"`, PV(capex + replacements − TV) is exactly the PV of the LP's
annuities, so the case NPV is the LP objective saving × the annuity factor.
"""
from __future__ import annotations

import io
import math
from datetime import date

import numpy as np
import pytest
from pydantic import ValidationError

from models.finance import FinanceInputs, TerminalValueRule
from services.finance.case import (
    AssetFinance, AssetPart, FinanceCase, LpBasis, StorageYear, Template, TemplateLine,
)
from services.finance.engine import run_case
from services.finance.tax import DepreciationClass, TaxLayer, sl_half_year

COD = date(2030, 1, 1)
NO_TAX = (TaxLayer(name="corp", rate=0.0,
                   depreciation=(DepreciationClass("all", 1.0, sl_half_year(10)),)),)
RLA = {"method": "remaining_life_annuity"}
OM_KEY = "asset:fom:StorageUnit:bess"


def crf(r, life):
    return 1.0 / life if r == 0 else r / (1.0 - (1.0 + r) ** -life)


def apf(r, y):
    if y <= 0:
        return 0.0
    return float(y) if r == 0 else (1.0 - (1.0 + r) ** -y) / r


def _battery(power=400_000.0, energy=1_200_000.0, lp=10.0, le=15.0):
    return AssetFinance("bess", "StorageUnit", power + energy, max(lp, le), "battery",
                        (AssetPart("power", power, lp, None), AssetPart("energy", energy, le, None)))


def _fin(**over) -> FinanceInputs:
    kw = dict(financial_close=date(2029, 1, 1), cod_by_asset={}, analysis_years=25,
              contingency_share=0.0, capex_phasing=[1.0],
              escalation={c: 0.0 for c in ("opex", "tariff", "export", "capex")},
              degradation_by_asset={"bess": 0.0}, tax_losses="offset_other_income",
              financing_fee_tax="not_deducted", wacc_nominal=0.07, cost_of_equity=0.07,
              inflation=0.0, replacement_rule="part_lifetimes", terminal_value=RLA)
    kw.update(over)
    return FinanceInputs(**kw)


def _case(fin=None, assets=None, *, rate=0.07, asset_rates=None, lp=True, storage=False,
          saving=300_000.0) -> FinanceCase:
    assets = assets if assets is not None else (_battery(),)
    st = {"bess": StorageYear(900.0, 1000.0, 40_000.0, 0.0, om_keys=(OM_KEY,))} if storage else {}
    # A flat site: the option's bill and the baseline's (the counterfactual), saving `saving` a year.
    actual = Template(2030, (TemplateLine("bill", "energy_import", -1_000_000.0 + saving, "tariff"),
                             TemplateLine(OM_KEY, "fom", -0.0, "opex", source="asset")),
                      storage=st)
    base = Template(2030, (TemplateLine("bill", "energy_import", -1_000_000.0, "tariff"),))
    return FinanceCase(inputs=fin or _fin(), owner="o", base_year=2030, cod=COD, templates=(actual,),
                       assets=tuple(assets), counterfactual=(base,),
                       lp_basis=LpBasis(discount_rate=rate,
                                        asset_discount_rates=asset_rates or {}) if lp else None)


def _tv(r) -> float:
    return float(r.op.terminal[-1])


# ── GS parity, rebuilt from hand values ──────────────────────────────────────


def test_gs_parity_on_the_seed_ledger_numbers():
    """GS's seed ledger (technology-data v0.14.0): battery inverter 213.9279 EUR/kW · 10 y, battery storage
    189.861 EUR/kWh · 25 y (= the horizon), rooftop PV 883.8138 EUR/kW; PV at 30 y here (in a 25-y
    horizon, as the plan states). 1 MW / 2 h battery, 2 MW PV, 7 %. The inverter re-bought at t = 10 and 20
    serves years 21–30: 5 years left at H = 25; the PV has 5 left; the storage block none."""
    inv, sto, pv_kw = 213_927.9, 2 * 189_861.0, 883_813.8
    bess = AssetFinance("bess", "StorageUnit", inv + sto, 25.0, "battery",
                        (AssetPart("power", inv, 10.0, None), AssetPart("energy", sto, 25.0, None)))
    pv = AssetFinance("pv", "Generator", 2 * pv_kw, 30.0, "solar",
                      (AssetPart("investment", 2 * pv_kw, 30.0, None),))
    r = run_case(_case(assets=(bess, pv)), layers=NO_TAX)
    want = {("bess", "power"): inv * crf(0.07, 10) * apf(0.07, 5),
            ("bess", "energy"): 0.0,
            ("pv", "investment"): 2 * pv_kw * crf(0.07, 30) * apf(0.07, 5)}
    got = {(t.asset, t.part): t.value for t in r.op.terminal_terms}
    assert got.keys() == want.keys()
    for k, v in want.items():
        assert got[k] == pytest.approx(v, rel=1e-9, abs=1e-9), k
    assert _tv(r) == pytest.approx(sum(want.values()), rel=1e-9)
    inverter = next(t for t in r.op.terminal_terms if t.part == "power")
    assert (inverter.start_year, inverter.remaining_years) == (2050, 5.0)    # serves 2050–2059
    # The re-purchases are GS's t = 10, 20 (financial close 2029 = t 0).
    assert [x.year for x, _v in r.op.replacement_items] == [2039, 2049]


def test_the_asset_rate_annuitises_and_the_lp_rate_discounts():
    r = run_case(_case(asset_rates={"bess": 0.05}), layers=NO_TAX)
    power = next(t for t in r.op.terminal_terms if t.part == "power")
    assert (power.asset_rate, power.rate) == (0.05, 0.07)
    assert power.value == pytest.approx(400_000.0 * crf(0.05, 10) * apf(0.07, 5), rel=1e-12)


def test_a_replacement_base_is_its_escalated_cost():
    esc = {c: 0.0 for c in ("opex", "tariff", "export")} | {"capex": 0.02}
    r = run_case(_case(_fin(escalation=esc)), layers=NO_TAX)
    terms = {t.part: t for t in r.op.terminal_terms}
    assert terms["power"].base == pytest.approx(400_000.0 * 1.02 ** (2049 - 2030), rel=1e-12)
    assert terms["energy"].base == pytest.approx(1_200_000.0 * 1.02 ** (2044 - 2030), rel=1e-12)
    # energy: k = 1, start 2045, 15 y → 2059; 5 years past the horizon 2055.
    assert terms["energy"].remaining_years == 5.0


# ── the identity (GS's BY_CONSTRUCTION) ──────────────────────────────────────


IDENTITY_CASES = [
    (25, ((400_000.0, 10.0), (1_200_000.0, 15.0))),       # the battery
    (25, ((1_000_000.0, 30.0),)),                          # PV 30 y in 25 y
    (20, ((1_000_000.0, 10.0),)),                          # a 10-y part in H = 20
]


@pytest.mark.parametrize("rate", [0.07, 0.03, 0.0])
@pytest.mark.parametrize("horizon,parts", IDENTITY_CASES)
def test_the_case_npv_is_the_lp_saving_times_the_annuity_factor(rate, horizon, parts):
    a = AssetFinance("a", "Generator", sum(c for c, _ in parts), max(life for _, life in parts),
                     None, tuple(AssetPart(f"p{i}", c, life, None) for i, (c, life) in enumerate(parts)))
    saving = 300_000.0
    case = _case(_fin(analysis_years=horizon, wacc_nominal=rate, cost_of_equity=rate,
                      degradation_by_asset={}), assets=(a,), rate=rate, saving=saving)
    r = run_case(case, layers=NO_TAX)
    annuity = sum(c * crf(rate, life) for c, life in parts)
    lp_pv = annuity * apf(rate, horizon)
    n = r.tl.n
    d = (1.0 + rate) ** np.arange(n)
    pv_costs = float(np.sum((r.op.capex + r.op.replacement - r.op.terminal) / d))
    assert pv_costs == pytest.approx(lp_pv, rel=1e-9)
    want = (saving - annuity) * apf(rate, horizon)
    for key in ("project_pre_tax_npv", "project_post_tax_npv"):
        assert r.metrics[key] == pytest.approx(want, rel=1e-9, abs=1e-6), key


def test_the_review_hand_check_on_the_battery_at_7_percent():
    r = run_case(_case(), layers=NO_TAX)
    d = 1.07 ** np.arange(r.tl.n)
    pv_costs = float(np.sum((r.op.capex + r.op.replacement - r.op.terminal) / d))
    assert pv_costs == pytest.approx(2_199_084.18, abs=0.005)


# ── not established, the model, the LCOS ─────────────────────────────────────


def test_an_infinite_lifetime_is_terminal_part_unknown():
    pv = AssetFinance("pv", "Generator", 1_000_000.0, None, "solar",
                      (AssetPart("investment", 1_000_000.0, math.inf, None),))
    r = run_case(_case(assets=(_battery(), pv)), layers=NO_TAX)
    assert r.op.terminal is None
    assert {"terminal_value_missing", "terminal_part_unknown:pv:investment"} <= \
        set(r.op.reasons["terminal"])
    assert r.cash["project_pre_tax"] is None


@pytest.mark.parametrize("make,reason", [
    (lambda: _case(lp=False), "terminal_needs_lp_rate"),
    (lambda: _case(rate=None), "terminal_needs_lp_rate"),
    (lambda: _case(_fin(analysis_years=15), assets=(AssetFinance(
        "bess", "StorageUnit", 1_600_000.0, 15.0, "battery",
        (AssetPart("power", 400_000.0, None, None), AssetPart("energy", 1_200_000.0, 15.0, None))),)),
     "terminal_part_unknown:bess:power"),
    (lambda: _case(_fin(replacement_rule="fixed", replacement_capex=[
        (2039, "bess", 400_000.0), (2044, "bess", 1_200_000.0), (2049, "bess", 400_000.0)])),
     "terminal_needs_part_lifetimes:bess"),
])
def test_not_established_with_a_reason(make, reason):
    r = run_case(make(), layers=NO_TAX)
    assert r.op.terminal is None
    assert "terminal_value_missing" in r.op.reasons["terminal"]
    assert reason in r.op.reasons["terminal"]


def test_fixed_rule_single_part_values_its_last_entry():
    """Under `fixed`, a single-part asset's last purchase is its last entry: 10 y bought in 2039
    (serving 2040–2049) on a 15-y axis to 2044 → 5 years left."""
    a = AssetFinance("gen", "Generator", 1_000_000.0, 10.0)
    case = _case(_fin(replacement_rule="fixed", analysis_years=15,
                      replacement_capex=[(2039, "gen", 900_000.0)]), assets=(a,))
    (t,) = run_case(case, layers=NO_TAX).op.terminal_terms
    assert (t.start_year, t.remaining_years, t.base) == (2040, 5.0, 900_000.0)
    assert t.value == pytest.approx(900_000.0 * crf(0.07, 10) * apf(0.07, 5), rel=1e-12)


def test_the_method_takes_no_value():
    assert TerminalValueRule(method="remaining_life_annuity").value is None
    with pytest.raises(ValidationError):
        TerminalValueRule(method="remaining_life_annuity", value=1.0)


def test_the_lcos_nets_the_part_terms_under_this_method_only():
    rla = run_case(_case(storage=True), layers=NO_TAX)
    none = run_case(_case(_fin(terminal_value={"method": "none"}), storage=True), layers=NO_TAX)
    fixed = run_case(_case(_fin(terminal_value={"method": "fixed", "value": 1e6}), storage=True),
                     layers=NO_TAX)
    a, b, c = (x.lcos["assets"]["bess"] for x in (rla, none, fixed))
    pv_tv = _tv(rla) / 1.07 ** (rla.tl.n - 1)
    assert a["pv_terminal"] == pytest.approx(pv_tv, rel=1e-12)
    assert b["pv_terminal"] is None and c["pv_terminal"] is None
    num = lambda x: x["lcos_nominal_per_mwh"] * x["pv_discharge_mwh"]           # noqa: E731
    assert num(a) == pytest.approx(num(b) - pv_tv, rel=1e-12)
    assert c["lcos_nominal_per_mwh"] == b["lcos_nominal_per_mwh"]
    assert "remaining_life_annuity" in rla.lcos["basis"]


# ── the report and the workbook ──────────────────────────────────────────────


def test_the_report_lists_the_terms_and_the_workbook_carries_the_method():
    from openpyxl import load_workbook

    from services.finance.export_xlsx import build_workbook
    from services.finance.report import assemble_finance_sections

    case = _case()
    r = run_case(case, layers=NO_TAX)
    rep = assemble_finance_sections(r, case)
    block = rep.sections["project"].payload["terminal_value"]
    assert block["method"] == "remaining_life_annuity"
    assert block["value"] == pytest.approx(_tv(r), rel=1e-12)
    assert {(t["asset"], t["part"]) for t in block["terms"]} == {("bess", "power"), ("bess", "energy")}
    power = next(t for t in block["terms"] if t["part"] == "power")
    assert power["value"] == pytest.approx(400_000.0 * crf(0.07, 10) * apf(0.07, 5), rel=1e-12)
    assert "LP" in block["rate_basis"] and "escalated" in block["base_basis"]
    ws = load_workbook(io.BytesIO(build_workbook(rep)))["About"]
    rows = {ws.cell(i, 1).value: ws.cell(i, 2).value for i in range(1, ws.max_row + 1)}
    assert rows["Terminal value method"] == "remaining_life_annuity"
    # Another method: the block names it, with no terms.
    case2 = _case(_fin(terminal_value={"method": "none"}))
    p2 = assemble_finance_sections(run_case(case2, layers=NO_TAX), case2).sections["project"].payload
    assert p2["terminal_value"]["method"] == "none" and p2["terminal_value"]["terms"] == []
