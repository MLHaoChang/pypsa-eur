"""
The in-repo Excel formula evaluator (`tests/xlsx_formula_eval.py`) checked
against Excel's documented NPV semantics (plan S5; review v1 S9).

Ported verbatim from `test_proforma_xlsx.py` (U2 gate C8 b): the workbook
targets (`test_engine_adapter_workbook.py`, `test_u2_wp9_workbooks.py`)
evaluate the guided workbooks with this evaluator, so its own tests outlive
the pro forma's file, which WP10 deletes.
"""
from __future__ import annotations

import openpyxl
import pytest

from tests.xlsx_formula_eval import FormulaEvaluator, excel_npv


def test_the_evaluator_matches_excels_documented_npv_examples():
    """
    Microsoft's NPV function page, both examples:
    * ``=NPV(10%, -10000, 3000, 4200, 6800)`` is 1,188.44 (the -10,000 is
      paid at the END of the first period, so it is discounted too);
    * ``=NPV(8%, B3:B7) + B2`` with an initial -40,000 and 8,000, 9,200,
      10,000, 12,000, 14,500 is 1,922.06, and with a -9,000 loss in year six
      ``=NPV(8%, B3:B7, -9000) + B2`` is -3,749.47.
    """
    assert excel_npv(0.10, [-10000, 3000, 4200, 6800]) == pytest.approx(1188.44, abs=0.005)
    flows = [8000, 9200, 10000, 12000, 14500]
    assert excel_npv(0.08, flows) - 40000 == pytest.approx(1922.06, abs=0.005)
    assert excel_npv(0.08, flows + [-9000]) - 40000 == pytest.approx(-3749.47, abs=0.005)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    for i, v in enumerate([-40000] + flows, start=2):
        ws[f"B{i}"] = v
    ws["A1"] = 0.08
    ws["C1"] = "=NPV(A1,B3:B7)+B2"
    ws["C2"] = "=NPV(A1,B3:B7,-9000)+B2"
    ws["C3"] = "=NPV(10%,-10000,3000,4200,6800)".replace("10%", "0.1")
    ev = FormulaEvaluator(wb)
    assert ev.value("S", "C1") == pytest.approx(1922.06, abs=0.005)
    assert ev.value("S", "C2") == pytest.approx(-3749.47, abs=0.005)
    assert ev.value("S", "C3") == pytest.approx(1188.44, abs=0.005)


def test_the_evaluator_refuses_what_it_does_not_implement():
    wb = openpyxl.Workbook()
    wb.active.title = "S"
    wb["S"]["A1"] = "=IRR(B1:B3)"
    with pytest.raises(ValueError):
        FormulaEvaluator(wb).value("S", "A1")
