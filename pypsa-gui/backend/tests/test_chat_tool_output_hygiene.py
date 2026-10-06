"""
What chat tools hand to people and to the provider — the register's
"lower-confidence notes", verified and closed (2026-10-06), plus one
interaction the #82 merge created.

* `export_to_csv` / `export_to_excel` wrote model-supplied cells unmodified.
  openpyxl stores ANY string starting with "=" as a live formula, so a bus named
  `=HYPERLINK("https://…","open")` became a working link in a file the user was
  invited to open; in CSV, a leading = + - @ tab or CR is a formula to Excel.
* `gridspine_export_handoff_bundle` returned `str(path)` — the server's storage
  root plus the org and project UUIDs — to the model.
* `apply_demand_from_excel` declared a `replace` flag that nothing read; the
  tool always replaces.
* A bare exception's text (a `FileNotFoundError` naming its path) reached the
  provider with the absolute storage path in it.
* The chat undo snapshot was taken BEFORE the live-study gate refused the tool —
  the reverse of the HTTP middleware's order — exporting a network the study
  was re-solving, for an undo step that changes nothing.
"""
from __future__ import annotations

import contextlib
import csv
import io
import types

import pypsa
import pytest
from openpyxl import load_workbook

from services import chat_service, chat_tools, undo_service
# AUTO_APPROVE_TIERS lives in harness.confirm; patch it there, not via the
# chat_service alias (tests/test_harness_layout.py's tripwire).
from harness import confirm as harness_confirm
from services.pypsa_service import PyPSAService

HYPERLINK = '=HYPERLINK("https://example.invalid","open")'


@pytest.fixture
def captured_export(monkeypatch):
    seen: dict = {}

    def _save(payload, filename, mime):
        seen.update(payload=payload, filename=filename, mime=mime)
        return {"file_id": "x", "filename": filename}

    monkeypatch.setattr(chat_tools, "_save_agent_export", _save)
    return seen


# ── Formula injection ────────────────────────────────────────────────────────


def test_a_csv_cell_that_reads_as_a_formula_is_written_as_text(captured_export):
    chat_tools.export_to_csv(
        rows=[[HYPERLINK, "+1+1", "@SUM(A1)", "-5", "+3.2", 7, "plain"]],
        columns=["=cmd|' /C calc'!A0", "b", "c", "d", "e", "f", "g"],
        filename="out.csv")

    cells = list(csv.reader(io.StringIO(captured_export["payload"].decode())))
    header, row = cells
    assert header[0] == "'=cmd|' /C calc'!A0"
    assert row[:3] == ["'" + HYPERLINK, "'+1+1", "'@SUM(A1)"]
    # A string that IS a number is data, not a formula: left as the number.
    assert row[3:] == ["-5", "+3.2", "7", "plain"]


def test_an_xlsx_cell_that_reads_as_a_formula_is_stored_as_a_string(captured_export):
    chat_tools.export_to_excel(
        {"S": [["name", "value"], [HYPERLINK, 3]]}, "out.xlsx")

    ws = load_workbook(io.BytesIO(captured_export["payload"]))["S"]
    cell = ws["A2"]
    assert cell.data_type == "s", (
        "openpyxl stored the string as a FORMULA (data_type 'f'); Excel would "
        "evaluate it on open"
    )
    assert cell.value == HYPERLINK, "the text itself must survive unchanged"
    assert ws["B2"].value == 3


# ── What the model is told ───────────────────────────────────────────────────


def test_the_handoff_bundle_tool_returns_a_route_not_a_server_path(tmp_path, monkeypatch):
    from services import gridspine_service

    bundle = tmp_path / "org-uuid" / "project-uuid" / "h12.zip"
    bundle.parent.mkdir(parents=True)
    bundle.write_bytes(b"PK")

    @contextlib.contextmanager
    def _acting():
        yield None, None

    monkeypatch.setattr(chat_tools, "_acting", _acting)
    monkeypatch.setattr(chat_tools, "_gridspine_project",
                        lambda db, user, pid: types.SimpleNamespace(name="My Hub"))
    monkeypatch.setattr(gridspine_service, "export_handoff_bundle",
                        lambda project, hour: bundle)

    result = chat_tools.gridspine_export_handoff_bundle("My Hub", 12)

    assert str(tmp_path) not in repr(result)
    assert result == {"download_url": "/api/gridspine/My%20Hub/bundles/12",
                      "filename": "h12.zip", "bytes": 2}


def test_apply_demand_from_excel_offers_no_flag_it_ignores():
    import inspect

    from services.chat_tools_schema import TOOLS

    schema = next(t for t in TOOLS if t["name"] == "apply_demand_from_excel")
    assert "replace" not in schema["input_schema"]["properties"]
    assert "replace" not in inspect.signature(chat_tools.apply_demand_from_excel).parameters
    assert "REPLACING" in schema["description"], "say what the tool always does"


def test_a_bare_exception_does_not_send_the_storage_root_to_the_provider():
    from settings import get_settings

    root = get_settings().projects_root
    exc = FileNotFoundError(2, "No such file or directory",
                            str(root / "org-uuid" / "project-uuid" / "asset_health.json"))

    content = chat_service._error_result_content(None, exc, "tool_error")

    assert str(root) not in content
    assert "asset_health.json" in content, "the file name is what the model needs"
    assert "<project storage>" in content


# ── Undo snapshot vs the live-study gate ─────────────────────────────────────


def test_a_tool_refused_by_a_live_study_costs_no_undo_step(
    tmp_projects_dir, install_network, monkeypatch,
):
    from tests.test_live_network_untouched import _LiveRecord

    n = pypsa.Network()
    n.add("Bus", "B1")
    n.add("Generator", "G1", bus="B1", p_nom=100.0)
    install_network(n)
    undo_service.clear()
    monkeypatch.setattr(harness_confirm, "AUTO_APPROVE_TIERS",
                        frozenset({"write", "destructive"}))
    exported: list[int] = []
    from services import network_undo
    real_push = network_undo.push_undo_snapshot
    monkeypatch.setattr(network_undo, "push_undo_snapshot",
                        lambda: (exported.append(1), real_push())[1])

    with _LiveRecord(PyPSAService.get_solver_state(), "fmea_sweep"):
        frames = list(chat_service._dispatch_real_tool_call(
            chat_service.ChatSession(),
            {"id": "t", "name": "update_component",
             "input": {"component_class": "Generator", "name": "G1",
                       "attrs": {"p_nom": 150.0}}}, []))

    kinds = [p.get("error_kind") for ev, p in frames if ev == "tool_error"]
    assert kinds == ["study_in_flight"], frames
    assert exported == [], (
        "the network was exported for an undo snapshot while the study was "
        "re-solving it, for an edit the gate then refused"
    )
    assert undo_service.depth() == 0
