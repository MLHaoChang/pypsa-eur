"""
QA: phase 1 of the study-report plan, end to end — a real Energy Hub study
over HTTP, then the evidence-only report created through the report routes,
listed, fetched, its figure served, exported to Word and re-opened, carried
by a save-as, and deleted.

Plan: docs/superpowers/plans/2026-09-28-llm-report-generation-increment-1.md
(phase 1 gate: WP1 store/routes, WP2 evidence, WP4 figures, WP5 writer +
``POST /{name}/reports`` with ``mode: "evidence_only"``). The unit and route
tests cover each piece with hand-built reports; this driver pins the JOURNEY
on the numbers the study engines actually produced: every ``REPORT_SECTIONS``
entry is a section of the document, the ranked failure mode is in the table
and the Pareto figure is embedded, a stage the budget skipped is stated as
not established (never omitted, never a zero), the ``.docx`` the blob route
serves re-opens with python-docx and carries the ``sec:<id>`` bookmarks the
round trip (increment 3) will anchor on, and ``reports/`` travels with the
project.

Runs in CI through ``tests/run_qa_drivers.py``. Exit 0 = every step passed.
"""
from __future__ import annotations

import io
import pathlib
import sys
import time
import zipfile

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

# THE load-bearing import, and it must come first: `qa_support` pins the
# sandbox before anything imports `main` or `settings`.
from tests import qa_support          # noqa: E402
from services.reports.evidence import fmea_top_modes  # noqa: E402

from tests.eh_stage_fixtures import VOLL, certifiable_weak_network  # noqa: E402

PROJECT = "qa_reports_phase1"
PROJECT_COPY = "qa_reports_phase1_copy"
BUDGET_SOLVES = 8
STAGES = ["apply_pack", "ens_solve", "fmea_top", "assemble"]

PASS = 0
FAIL = 0


def _step(label: str, ok: bool, msg: str = "") -> None:
    global PASS, FAIL
    if ok:
        PASS += 1
        print(f"  [PASS] {label}" + (f" — {msg}" if msg else ""))
    else:
        FAIL += 1
        print(f"  [FAIL] {label}" + (f" — {msg}" if msg else ""))


def _poll(path: str, *, timeout: float = 600.0) -> dict:
    c = qa_support.client()
    t0 = time.time()
    body: dict = {}
    while time.time() - t0 < timeout:
        r = c.get(path)
        if r.status_code == 204:
            return {"status": "no-content"}
        if r.status_code != 200:
            return {"status": f"http-{r.status_code}", "error": r.text[:300]}
        body = r.json()
        if body.get("status") != "running":
            return body
        time.sleep(0.25)
    return {"status": "timeout", **body}


def _docx_text(blob: bytes) -> tuple[str, int]:
    from docx import Document

    doc = Document(io.BytesIO(blob))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts), len(doc.inline_shapes)


def _document_xml(blob: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        return z.read("word/document.xml").decode("utf-8")


# ── the journey ───────────────────────────────────────────────────────────

def section_1_real_study() -> dict | None:
    print("\n[1] a real strong_grid study on the certifiable weak network")
    c = qa_support.client()
    # Install UNBOUND, then save: the first save claims the name AND binds the
    # session context's storage_dir (the lesson qa_reports_phase0 records).
    qa_support.install_network(certifiable_weak_network())
    qa_support.save_project(PROJECT)

    r = c.put("/api/simulation/solver_config",
              json={"solver_name": "highs", "voll": VOLL})
    _step("the solver config takes the VoLL", r.status_code == 200, r.text[:200])
    r = c.post("/api/results/eh_study", json={
        "archetype": "strong_grid", "budget_solves": BUDGET_SOLVES, "stages": STAGES})
    _step("the study starts", r.status_code == 200, r.text[:200])
    if r.status_code != 200:
        return None
    study = _poll("/api/results/eh_study")
    _step("the study finishes", study.get("status") == "done",
          f"status={study.get('status')} error={str(study.get('error'))[:200]}")
    if study.get("status") != "done":
        return None
    r = c.get("/api/results/eh_reference_design")
    _step("the report is persisted for GET", r.status_code == 200, r.text[:200])
    if r.status_code != 200:
        return None
    rep = r.json()
    comp = rep.get("completeness") or {}
    fmea = (rep.get("sections") or {}).get("fmea_top") or {}
    top = fmea_top_modes(fmea.get("payload"))
    _step("fmea_top is established with at least one ranked mode",
          comp.get("fmea_top") == "ok" and len(top) >= 1,
          f"status={comp.get('fmea_top')} rows={len(top)}")
    skipped = [k for k, v in comp.items() if v != "ok"]
    _step("at least one stage was not established (the budget skipped it)",
          len(skipped) >= 1, f"skipped={skipped}")
    return rep


def section_2_evidence_only_report(rep: dict) -> tuple[str, dict] | None:
    print("\n[2] POST /reports (evidence_only) → listed → fetched → figure")
    from models.energy_hub import REPORT_SECTIONS
    from services.reports.docx_writer import SECTION_TITLES

    c = qa_support.client()
    r = c.post(f"/api/projects/{PROJECT}/reports",
               json={"mode": "evidence_only", "title": "QA phase-1 client report"})
    _step("the route answers 200", r.status_code == 200, r.text[:300])
    if r.status_code != 200:
        return None
    body = r.json()
    rid = body.get("report_id") or ""
    _step("with a report id and the evidence_only mode",
          len(rid) == 16 and body.get("mode") == "evidence_only", str(body)[:160])
    doc = body.get("document") or {}
    by_id = {s.get("section_id"): s for s in doc.get("sections") or []}
    missing = [n for n in REPORT_SECTIONS if n not in by_id]
    _step("every REPORT_SECTIONS entry is a section of the document",
          not missing, f"missing={missing}")
    wrong_title = [n for n in REPORT_SECTIONS
                   if n in by_id and by_id[n].get("heading") != SECTION_TITLES[n]]
    _step("each under its report title", not wrong_title, str(wrong_title))
    fmea = by_id.get("fmea_top") or {}
    blocks = fmea.get("blocks") or []
    _step("fmea_top is ok with a TableRef and a FigureRef",
          fmea.get("status") == "ok"
          and any(b.get("type") == "table_ref" and b.get("table_id") == "fmea_top"
                  for b in blocks)
          and any(b.get("type") == "figure_ref" and b.get("figure_id") == "fmea_pareto"
                  for b in blocks),
          f"status={fmea.get('status')} blocks={[b.get('type') for b in blocks]}")
    comp = rep.get("completeness") or {}
    skipped = [n for n in REPORT_SECTIONS if comp.get(n) != "ok"]
    bad = []
    for n in skipped:
        s = by_id.get(n) or {}
        callouts = [b for b in s.get("blocks") or []
                    if b.get("type") == "callout" and b.get("kind") == "not_established"]
        if s.get("status") not in ("not_established", "skipped") or len(callouts) != 1:
            bad.append((n, s.get("status"), len(callouts)))
    _step("every skipped stage is a not_established/skipped section with one "
          "not_established callout", not bad, str(bad)[:200])
    _step("the document's mode is evidence_only and every section is code-sourced",
          doc.get("mode") == "evidence_only"
          and all(s.get("source") == "code" for s in doc.get("sections") or []))
    headline = (doc.get("tables") or {}).get("headline") or {}
    lole = next((row for row in headline.get("rows") or []
                 if row and str(row[0]).startswith("MC LOLE")), None)
    _step("the headline MC LOLE reads 'not established' (mc_certify was skipped)",
          rep.get("mc_lole_h") is None and lole is not None
          and lole[1] == "not established", f"mc_lole_h={rep.get('mc_lole_h')} row={lole}")
    _step("the document's hash is the meta's", doc.get("evidence_hash") == body.get("evidence_hash"))

    r = c.get(f"/api/projects/{PROJECT}/reports")
    ids = [m.get("report_id") for m in (r.json() if r.status_code == 200 else [])]
    _step("GET /reports lists it", r.status_code == 200 and ids == [rid], str(ids))
    r = c.get(f"/api/projects/{PROJECT}/reports/{rid}")
    _step("GET /reports/{id} returns version 1",
          r.status_code == 200 and r.json().get("version") == 1,
          f"HTTP {r.status_code}")
    r = c.get(f"/api/projects/{PROJECT}/reports/{rid}/figures/fmea_pareto")
    _step("GET …/figures/fmea_pareto is a PNG",
          r.status_code == 200 and r.content[:8] == b"\x89PNG\r\n\x1a\n"
          and r.headers.get("content-type", "").startswith("image/png"),
          f"HTTP {r.status_code} {r.headers.get('content-type')}")
    return rid, doc


def section_3_export(rep: dict, rid: str, doc: dict) -> None:
    print("\n[3] POST …/export → agent_export chip → blob → Word")
    c = qa_support.client()
    r = c.post(f"/api/projects/{PROJECT}/reports/{rid}/export",
               json={"filename": "qa-phase1-report"})
    _step("the export answers 200", r.status_code == 200, r.text[:300])
    if r.status_code != 200:
        return
    meta = r.json()
    _step("it is an agent_export chip named as asked",
          meta.get("kind") == "agent_export"
          and meta.get("filename") == "qa-phase1-report.docx", str(meta)[:200])
    r = c.get(f"/api/projects/{PROJECT}/uploads/{meta.get('file_id')}/blob")
    _step("the blob route serves it", r.status_code == 200, f"HTTP {r.status_code}")
    _step("with the Word MIME",
          r.headers.get("content-type", "").startswith(
              "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
          r.headers.get("content-type", ""))
    if r.status_code != 200:
        return
    blob = r.content
    text, pictures = _docx_text(blob)
    _step("the document re-opens with python-docx", bool(text))
    fmea = (rep.get("sections") or {}).get("fmea_top") or {}
    top = fmea_top_modes(fmea.get("payload"))
    _step("it carries the study's top failure mode by name",
          bool(top) and str(top[0].get("name")) in text,
          f"name={top[0].get('name') if top else None}")
    _step("the Link-primary disclosure is present", "Link-primary" in text)
    _step("the headline MC LOLE row reads 'not established'",
          "MC LOLE (h/yr)\nnot established" in text)
    _step("the skipped stages are stated, not omitted",
          text.count("Not established:") >= 3,
          f"statements={text.count('Not established:')}")
    figures = doc.get("figures") or {}
    _step("one picture per figure the report references (Pareto + capacity mix "
          "when sizing is established)",
          "fmea_pareto" in figures and pictures == len(figures),
          f"pictures={pictures} figures={sorted(figures)}")
    xml = _document_xml(blob)
    _step("the fmea_top heading carries its sec: bookmark",
          'w:name="sec:fmea_top"' in xml)
    _step("every EH section has a bookmark",
          all(f'w:name="sec:{n}"' in xml for n in (
              "target", "certification", "frontier", "sizing", "fmea_top", "tea")))
    r = c.get(f"/api/projects/{PROJECT}/uploads")
    kinds = {u.get("file_id"): u.get("kind") for u in (r.json() if r.status_code == 200 else [])}
    _step("the uploads list shows the chip as an agent export",
          kinds.get(meta.get("file_id")) == "agent_export", str(kinds)[:200])


def section_4_save_as_and_delete(rid: str) -> None:
    print("\n[4] save-as carries reports/; DELETE removes the original")
    c = qa_support.client()
    r = c.post(f"/api/projects/{PROJECT_COPY}", params={"rebind": "true", "force": True})
    _step("save-as to a second name succeeds", r.status_code == 200, r.text[:200])
    r = c.get(f"/api/projects/{PROJECT_COPY}/reports")
    ids = [m.get("report_id") for m in (r.json() if r.status_code == 200 else [])]
    _step("the new project lists the report", ids == [rid], str(ids))
    r = c.get(f"/api/projects/{PROJECT_COPY}/reports/{rid}/figures/fmea_pareto")
    _step("and serves its figure", r.status_code == 200, f"HTTP {r.status_code}")

    r = c.delete(f"/api/projects/{PROJECT}/reports/{rid}")
    _step("DELETE on the original answers 200",
          r.status_code == 200 and r.json().get("deleted") is True, r.text[:200])
    r = c.get(f"/api/projects/{PROJECT}/reports")
    _step("the original no longer lists it",
          r.status_code == 200 and r.json() == [], r.text[:200])
    r = c.get(f"/api/projects/{PROJECT}/reports/{rid}")
    _step("and GET on it is 404 report_not_found",
          r.status_code == 404
          and (r.json().get("detail") or {}).get("error_kind") == "report_not_found",
          f"HTTP {r.status_code}")
    r = c.get(f"/api/projects/{PROJECT_COPY}/reports/{rid}")
    _step("the copy keeps its own", r.status_code == 200, f"HTTP {r.status_code}")


def main() -> int:
    print("=" * 60)
    print("QA: reports phase 1 — EH study → evidence-only report → .docx")
    print("=" * 60)
    crashed = False
    try:
        qa_support.reset_backend()
        qa_support.delete_project(PROJECT, PROJECT_COPY)
        rep = section_1_real_study()
        if rep is not None:
            created = section_2_evidence_only_report(rep)
            if created is not None:
                rid, doc = created
                section_3_export(rep, rid, doc)
                section_4_save_as_and_delete(rid)
    except Exception as exc:                                     # noqa: BLE001
        crashed = True
        import traceback
        print(f"\n  [FAIL] the journey raised {type(exc).__name__}: {exc}")
        traceback.print_exc()
    finally:
        try:
            qa_support.delete_project(PROJECT, PROJECT_COPY)
            qa_support.reset_backend()
        except Exception as exc:                                 # noqa: BLE001
            print(f"  ! cleanup raised {type(exc).__name__}: {exc}")

    total = PASS + FAIL + (1 if crashed else 0)
    print("\n" + "=" * 60)
    print(f"Total: {total}")
    print(f"Pass:  {PASS}")
    print(f"Fail:  {FAIL + (1 if crashed else 0)}")
    print("=" * 60)
    return 1 if (FAIL or crashed) else 0


if __name__ == "__main__":
    sys.exit(main())
