"""
U2 WP9 — the guided workbooks on the Investment Case engine.

Plan: docs/superpowers/plans/2026-10-05-guided-study-u2-engine-rewire.md WP9
(§2 C8; owner decision 9). `engine_adapter.case_workbook` is the target file
`test_engine_adapter_workbook.py`; here are the routes and the report:

* the case route's workbook is the engine's (`build_workbook` first, the
  guided sheets appended), its live NPV the case's to 1e-6;
* the report workbook values each option on the engine's cash series, stored
  with the report record at assembly (the reads render the record and load
  no network); a record stored before WP9 still renders, on the earlier
  sheet, and says so;
* rule rows 21-34 name the guided study as their author in the HTML, DOCX and
  XLSX appendix, never "library" (gate U2-WP8a N9);
* the ledger CSV carries `engine_path` (WP5 status: "no engine_path column
  yet (WP9)").
"""
from __future__ import annotations

import csv
import io
import pathlib

import openpyxl
import pytest

from tests.xlsx_formula_eval import FormulaEvaluator

SHEET = "Cash flows (formulas)"


@pytest.fixture
def studies_on(monkeypatch):
    from tests.study_s4_support import enable_studies

    yield from enable_studies(monkeypatch)


@pytest.fixture
def fake(monkeypatch):
    from services import solver_service
    from tests.test_study_runner import FakeSolver

    solver = FakeSolver()
    monkeypatch.setattr(solver_service, "run_simulation", solver)
    return solver


def _find(ws, label):
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == label:
                return cell
    raise AssertionError(f"{label!r} not on {ws.title}")


def _npv_cell(ws):
    label = _find(ws, "NPV")
    return ws.cell(label.row, label.column + 1)


# ── the routes ───────────────────────────────────────────────────────────

def test_the_case_route_serves_the_engines_workbook_with_the_guided_sheets(
        client, api_project, studies_on, fake):
    from services.http_filenames import content_disposition
    # Money-year 01: the report tests' `_run` states the custom tariff in 2025 EUR.
    from tests.test_study_report_routes import _run

    name = "wp9-case"
    sid = _run(client, api_project, name)
    base = f"/api/projects/{name}/studies/{sid}"
    case = client.get(f"{base}/options/bess_2h/case").json()
    x = client.get(f"{base}/options/bess_2h/case.xlsx")
    assert x.status_code == 200, x.text
    study = client.get(base).json()
    assert x.headers["content-disposition"] == content_disposition(
        f"{study['name']} - bess_2h - investment case.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames[:2] == ["About", "Summary"]
    assert wb.sheetnames[-4:] == ["KPIs", SHEET, "Assumptions", "Provenance"]
    about = {r[0].value: r[1].value for r in wb["About"].iter_rows()}
    assert about["Project"] == study["name"]
    got = FormulaEvaluator(wb).value(SHEET, _npv_cell(wb[SHEET]).coordinate)
    assert got == pytest.approx(case["kpis"]["npv"], abs=1e-6)
    prov = {r[0].value: r[1].value for r in wb["Provenance"].iter_rows() if r[0].value}
    assert prov["model_hash"] == case["provenance"]["model_hash"]
    assert prov["project_ref"] == case["provenance"]["project_ref"]
    assert prov["export_series"].startswith("decision-study:")


def test_the_report_workbook_values_each_option_on_the_engines_cash_series(
        client, api_project, studies_on, fake, project_storage_dir):
    from routers.studies import REPORT_AUX
    from services.study import store
    from tests.test_study_report_routes import _full, _post

    name = "wp9-report"
    sid = _full(client, api_project, name)
    _post(client, name, sid)
    base = f"/api/projects/{name}/studies/{sid}"
    record = store.load_aux(project_storage_dir(name), sid, REPORT_AUX)
    assert set(record["engine_cash"]) == {k for k, c in record["cases"].items()
                                          if c["status"] == "ok"} != set()
    x = client.get(f"{base}/report.xlsx")
    assert x.status_code == 200, x.text
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    ev = FormulaEvaluator(wb)
    for oid in record["engine_cash"]:
        case = client.get(f"{base}/options/{oid}/case").json()
        ws = wb[f"Cash flows {oid}"]
        _find(ws, "Operating cash (engine)")
        assert ev.value(ws.title, _npv_cell(ws).coordinate) == pytest.approx(
            case["kpis"]["npv"], abs=1e-6)
        kpis = wb[f"KPIs {oid}"]
        row = _find(kpis, "npv").row
        assert kpis[f"B{row}"].value.startswith(f"='Cash flows {oid}'!")

    # A record stored before WP9 (no engine series) still renders: the
    # earlier sheet over the stored case, and it says so.
    del record["engine_cash"]
    store.save_aux(project_storage_dir(name), sid, REPORT_AUX, record)
    x = client.get(f"{base}/report.xlsx")
    assert x.status_code == 200, x.text
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    for oid in [k for k, c in record["cases"].items() if c["status"] == "ok"]:
        ws = wb[f"Cash flows {oid}"]
        with pytest.raises(AssertionError):
            _find(ws, "Operating cash (engine)")
        _find(ws, "Savings")
        assert any(isinstance(c.value, str) and "assembled before" in c.value
                   for row in ws.iter_rows() for c in row)


# ── the report workbook, in process ──────────────────────────────────────

@pytest.fixture(scope="module")
def golden():
    from tests import test_engine_adapter_case_golden as G

    n, cfg, compiled, ledger = G._solved_ic("bess_2h")
    return G._bundle("bess_2h"), ledger


def test_the_report_sheet_is_the_case_workbooks_view_of_the_same_series(golden):
    from services.study import engine_adapter as A
    from services.study.report_xlsx import write_report_xlsx
    from tests.test_study_decision_report import _report

    bundle, ledger = golden
    series = A.cash_series(bundle)
    report = _report(ledger=ledger)
    wb = openpyxl.load_workbook(io.BytesIO(write_report_xlsx(
        report, {"bess_2h": bundle.view}, ledger, engine_cash={"bess_2h": series})))
    case_wb = openpyxl.load_workbook(io.BytesIO(A.case_workbook(bundle, ledger)))
    ws, cs = wb["Cash flows bess_2h"], case_wb[SHEET]
    assert [[c.value for c in r] for r in ws.iter_rows()] == \
        [[c.value for c in r] for r in cs.iter_rows()]
    got = FormulaEvaluator(wb).value(ws.title, _npv_cell(ws).coordinate)
    assert got == pytest.approx(bundle.result.metrics["project_pre_tax_npv"], abs=1e-6)


def test_an_engine_series_is_never_paired_with_another_case(golden):
    """The stored series belongs to its case: another case id falls back."""
    from services.study import engine_adapter as A
    from services.study.report_xlsx import write_report_xlsx
    from tests.test_study_decision_report import _report

    bundle, ledger = golden
    series = {**A.cash_series(bundle), "case_id": "another.case"}
    wb = openpyxl.load_workbook(io.BytesIO(write_report_xlsx(
        _report(ledger=ledger), {"bess_2h": bundle.view}, ledger,
        engine_cash={"bess_2h": series})))
    with pytest.raises(AssertionError):
        _find(wb["Cash flows bess_2h"], "Operating cash (engine)")


# ── rows 21-34: the guided study is the rule's author (gate WP8a N9) ─────

_RULES = ("financial_close_year", "contingency_share", "escalation_tariff", "salvage_rule",
          "tax_pack", "value_flows_template")


def _edited_ledger():
    from services.study import ledger as LG
    from tests.test_study_decision_report import _ledger

    return LG.apply_user_row(_ledger(), "battery_storage_eur_per_kwh", 450.0,
                             unit="EUR/kWh", changed_by="test")


def _provenance_column(ws) -> dict:
    head = [c.value for c in ws[1]]
    j = head.index("provenance")
    return {r[0].value: r[j].value for r in ws.iter_rows(min_row=2) if r[0].value}


def test_the_workbooks_name_each_default_by_its_author():
    from services.study.report_xlsx import write_report_xlsx
    from services.study.workbook_view import assumptions
    from tests.test_study_decision_report import _report

    ledger = _edited_ledger()
    own = openpyxl.Workbook()
    assumptions(own.active, ledger)
    rep = openpyxl.load_workbook(io.BytesIO(write_report_xlsx(_report(ledger=ledger), {},
                                                              ledger)))
    for ws in (own.active, rep["Assumptions"]):
        prov = _provenance_column(ws)
        for key in _RULES:
            assert prov[key] == "guided study rule", (key, prov[key])
        for key in ("tariff", "demand_charge_price", "export_series"):
            assert prov[key] == "tariff", key
        for key in ("battery_inverter_eur_per_kw", "discount_rate",
                    "sizing_limit_connection_multiple"):
            assert prov[key] == "defaults pack", key
        assert prov["battery_storage_eur_per_kwh"] == "user"
        assert "library" not in set(prov.values())


def test_a_legacy_ledgers_defaults_are_the_earlier_study_library():
    """A ledger seeded before U2 (`technology-data v0.14.0`) names its library."""
    from services.study.workbook_view import provenance_label
    from tests.u2_targets import legacy_ledger

    # U2 WP10: the legacy library is deleted; its ledger is frozen (gate C8 a).
    legacy = legacy_ledger("de")
    assert not legacy.ledger_version.startswith("generic-defaults")
    row = next(r for r in legacy.rows if r.key == "battery_storage_eur_per_kwh")
    assert provenance_label(row, legacy.ledger_version) == "study library (earlier version)"


def test_the_html_and_docx_appendix_say_default_range_not_library_range():
    import docx

    from services.study.render_docx import render_docx
    from services.study.render_html import render_html
    from tests.test_study_decision_report import _report

    report = _report(ledger=_edited_ledger())
    html = render_html(report, charts={})
    assert "library range" not in html and "default range" in html
    doc = docx.Document(io.BytesIO(render_docx(report, charts={})))
    text = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "library range" not in text and "default range" in text


# ── the ledger CSV ───────────────────────────────────────────────────────

def test_the_ledger_csv_carries_engine_path_and_round_trips_it():
    from services.study import ledger as LG
    from tests.test_study_decision_report import _ledger

    led = _ledger()
    assert any(r.engine_path for r in led.rows)
    text = LG.ledger_to_csv(led)
    rows = {r["key"]: r for r in csv.DictReader(io.StringIO(text))}
    assert rows["salvage_rule"]["engine_path"] == "finance.terminal_value"
    assert rows["battery_storage_degradation_calendar_pct_per_year"]["engine_path"] == ""
    assert LG.ledger_rows_from_csv(text) == led.rows


# ── gate WP8b N2: the HTML and DOCX name engines in words ────────────────

def test_the_report_pages_name_each_figures_engine_in_words():
    from services.study.render_html import ENGINE_LABELS, provenance
    from tests.test_u2_wp8b_gate_conditions import _engine_case_without_lcos, _report_over

    report = _report_over(_engine_case_without_lcos())
    fig = report.sections["economics"].facts["case_npv"]
    assert fig.engine == "finance_engine"
    shown = provenance(fig)
    assert shown.startswith("finance engine;") and "finance_engine" not in shown, shown
    assert ENGINE_LABELS["cash_flow_expander"].startswith("the earlier ")


def test_the_engine_labels_mirror_the_frontends():
    import re

    from services.study.render_html import ENGINE_LABELS

    ts = (pathlib.Path(__file__).resolve().parents[2] / "frontend" / "src" / "utils"
          / "decisionVocabulary.ts").read_text(encoding="utf-8")
    block = ts[ts.index("export const ENGINE_LABELS"):]
    block = block[:block.index("\n}")]
    front = dict(re.findall(r"^\s+(\w+): '([^']*)',", block, flags=re.M))
    assert front and ENGINE_LABELS == front
