"""LC-06 S11, LC-D21: a generated report's facts, inspected in the exported document.

The seam is the copilot's report tools (``generate_report``, ``get_report``,
``get_report_table``, ``regenerate_report_section``, ``diff_report_versions``,
``export_report_docx``) over a stored study report and a scripted model, the way
``test_chat_report_tools.py`` runs them. The exported Word file is opened with
``python-docx`` and its tables and paragraphs read, not the tool's own replies.

Hand oracle, the stored report (archetype ``strong_grid``, ENS cap 10 per ten
thousand, Monte-Carlo LOLE 3.21 h/yr, nothing else established):
- the document's headline table has the rows Archetype = strong_grid, ENS cap
  10.00, Achieved ENS = not established, MC LOLE (h/yr) = 3.21 h/yr, and every
  section the study did not run says "Not established";
- the scripted summary "The plan meets 3.2 h/yr." is audited as verified;
- a regenerated summary that claims 9.9 h/yr is a new version, its 9.9 is audited
  as unverified and listed in the document's "Numbers to check" appendix, and neither
  the headline table nor any other section changes; version 1 still reads as it did.
"""
# ruff: noqa: F811 -- shared public harness fixtures
import io

from docx import Document

from services import chat_tools as T
from services.llm_fake import FakeProvider
from tests.test_chat_report_tools import _draft, _install_provider, _poll, _script, _store_eh_report, project  # noqa: F401
from tests.test_workflow_tools import helpers_state  # noqa: F401

SECTIONS = ["executive_summary", "fmea_top"]
HEADLINE = [
    ["Archetype", "strong_grid"], ["ENS cap (‱ of demand)", "10.00"],
    ["Achieved ENS (‱ of demand)", "not established"], ["Achieved shed hours (h/yr)", "not established"],
    ["MC LOLE (h/yr)", "3.21 h/yr"], ["Cost at target", "not established"], ["LCOE", "not established"],
    ["LCOH", "not established"], ["Pack hash", "p"], ["Assumptions hash", "a"]]


def _docx(client, report_id, version=None):
    exported = T.export_report_docx(report_id, version=version)
    response = client.get(T.get_file_delivery(exported["file_id"])["download_url"])
    assert response.status_code == 200, response.text
    return Document(io.BytesIO(response.content))


def _generated(monkeypatch):
    _store_eh_report()
    _install_provider(monkeypatch, FakeProvider(_script()))
    T.generate_report(sections=SECTIONS)
    assert _poll()["status"] == "done"
    return T.list_reports()[0]["report_id"]


def test_the_exported_document_holds_the_stored_figures_and_says_what_was_not_run(project, client, monkeypatch):
    report_id = _generated(monkeypatch)
    doc = _docx(client, report_id)
    headline = [[c.text for c in row.cells] for row in doc.tables[0].rows]
    assert headline == [["Item", "Value"]] + HEADLINE
    paragraphs = [p.text for p in doc.paragraphs]
    assert "The plan meets 3.2 h/yr." in paragraphs
    assert paragraphs.count("Not established: this section was not run in this study.") >= 6
    assert T.get_report(report_id)["sections"][0]["audit"] == {"unverified": [], "verified": ["3.2 h/yr"]}
    assert T.get_report_table(report_id, "headline")["items"] == HEADLINE


def test_a_regenerated_summary_cannot_change_the_admitted_numbers(project, client, monkeypatch):
    report_id = _generated(monkeypatch)
    prose = _draft("executive_summary", "The plan meets 9.9 h/yr.")
    _install_provider(monkeypatch, FakeProvider([prose, prose, prose]))
    assert T.regenerate_report_section(report_id, "executive_summary")["report_id"] == report_id
    assert _poll()["version"] == 2

    second = T.get_report(report_id)
    assert second["version"] == 2
    assert second["sections"][0]["blocks"][0] == {"type": "paragraph", "md": "The plan meets 9.9 h/yr."}
    assert second["sections"][0]["audit"] == {"unverified": ["9.9 h/yr"], "verified": []}
    assert T.get_report_table(report_id, "headline")["items"] == HEADLINE
    headline = [[c.text for c in row.cells] for row in _docx(client, report_id).tables[0].rows]
    assert headline == [["Item", "Value"]] + HEADLINE

    # The document flags it in the "Numbers to check" appendix and, in version 1, has none.
    assert "Executive summary: 9.9 h/yr" in [p.text for p in _docx(client, report_id).paragraphs]
    assert not any(p.text.startswith("Executive summary: ") for p in _docx(client, report_id, 1).paragraphs)

    first = T.get_report(report_id, version=1)
    assert first["sections"][0]["blocks"][0] == {"type": "paragraph", "md": "The plan meets 3.2 h/yr."}
    changes = {s["section_id"]: s["change"] for s in T.diff_report_versions(report_id, 1, 2)["sections"]}
    assert changes.pop("executive_summary") == "changed"
    assert set(changes.values()) == {"unchanged"}
