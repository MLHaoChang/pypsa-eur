"""
Offline extraction of a pinned, attributed DEA reference subset; never fetches.

Usage: python build_dea_reference.py POWER_XLSX FUELS_XLSX OUTPUT_JSON
Requires the repository's existing openpyxl dependency. XLSX files are untrusted
data: read-only cached values, no formula evaluation, no external links/macros.
"""
from __future__ import annotations

import hashlib
import json
import math
import sys
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

SOURCES = {
    "power": {
        "url": "https://ens.dk/media/8615/download",
        "sha256": "afb78dede13a17d17dca371ce74689bad98e55149ca58996fb3a535ac2704bb9",
        "landing_page": "https://ens.dk/en/analyses-and-statistics/technology-data-generation-electricity-and-district-heating",
    },
    "fuels": {
        "url": "https://ens.dk/media/6444/download",
        "sha256": "07ac3ef166883591a438e739ac69c7768b5419964abdd43e668130bd7a55ffee",
        "landing_page": "https://ens.dk/en/analyses-and-statistics/technology-data-renewable-fuels",
    },
}


def extract(power_path: Path, fuels_path: Path) -> dict:
    books = {}
    for key, path in [("power", power_path), ("fuels", fuels_path)]:
        raw = path.read_bytes()
        if hashlib.sha256(raw).hexdigest() != SOURCES[key]["sha256"]:
            raise ValueError(f"{key}: workbook differs from inspected release")
        book = load_workbook(BytesIO(raw), read_only=True, data_only=True, keep_links=False)
        licence = " ".join(str(book["Intro"].cell(r, c).value or "")
                           for r in range(1, 25) for c in range(1, 7))
        if "Attribution 4.0 International" not in licence:
            raise ValueError(f"{key}: expected workbook licence not found")
        books[key] = book

    specs = [
        ("onwind", "power", "20 Onshore turbines", 2020, 17, 25, 10, None,
         "MW_e output", "output_electric", True, "Total nominal investment includes source-specific grid connection, land/neighbour settlements and decommissioning of existing turbines.",
         ["Variable O&M must be priced separately", "Source forced/planned outage estimates are not a site-specific failure/repair law"]),
        ("solar-utility", "power", "22 Utility-scale PV", 2025, 17, 26, 10, None,
         "MW_e AC output", "output_electric_ac", True, "Fixed-tilt total investment includes modules, inverter, balance of plant, installation, soft costs and grid connection; FOM includes land rent.",
         ["PV degradation is a separate profile/finance assumption", "Do not substitute DC module capacity for AC capacity"]),
        ("solar-rooftop-commercial", "power", "22 Rooftop PV comm.&industrial", 2025,
         17, 26, 10, None, "MW_e AC output", "output_electric_ac", "unknown",
         "Commercial/industrial rooftop total system investment; separate from residential PV and from the historical 50/50 blend.",
         ["PV inverter replacement and degradation need separate project treatment", "Do not substitute DC module capacity for AC capacity"]),
        ("electrolysis-aec-100mw", "fuels", "80 AEC 100 MW", 2020, 27, 38, 23, 15,
         "MW_e total electricity input", "input_electric", "unknown", "100 MW AEC system investment and FOM, source system boundary; published planning estimates, not OEM quotations.",
         ["Stack replacement is excluded from FOM (source note P)", "Water/electricity input prices are excluded", "Operating-hour degradation, heat and water flows need separate modelling"]),
        ("electrolysis-pem-100mw", "fuels", "80 PEMEC 100 MW", 2020, 27, 38, 23, 15,
         "MW_e total electricity input", "input_electric", "unknown", "100 MW PEM system investment and FOM, source system boundary; published planning estimates, not OEM quotations.",
         ["Stack replacement is excluded from FOM (source note P)", "Water/electricity input prices are excluded", "Operating-hour degradation, heat and water flows need separate modelling"]),
    ]
    # capacity_basis_code / includes_grid_connection use the typed cost-unit
    # vocabulary shared with defaults-pack values; "unknown" where the source
    # text does not say whether grid connection is in the investment.
    entries = []
    for (technology, source, sheet, money_year, cap, fom, life, eff, capacity, capacity_code,
         grid_connection, included, gaps) in specs:
        ws = books[source][sheet]
        for col in range(3, 8):
            year = ws.cell(3, col).value
            if year not in (2025, 2030, 2035, 2040, 2050):
                continue
            # Exact inspected labels: the updated PV sheet leaves 2025/2035
            # blank; the two fuels sheets spell the 2030 label "crtl".
            # C:G are dated projection columns, H:K lower/upper uncertainty.
            expected_estimate = (None if source == "power" and technology.startswith("solar-")
                                 and year in (2025, 2035) else
                                 "crtl" if source == "fuels" and year == 2030 else "ctrl")
            if ws.cell(4, col).value != expected_estimate:
                raise ValueError(f"{sheet}: non-central projection column")
            cells = {}
            for name, row, expected in [
                ("overnight", cap, "investment"),
                ("fom_share", fom, "Fixed O&M"),
                ("lifetime", life, "lifetime"),
                *([("efficiency", eff, "LHV-based efficiency")] if eff else []),
            ]:
                label = str(ws.cell(row, 2).value)
                if expected.lower() not in label.lower():
                    raise ValueError(f"{sheet}:{row}: unexpected row label")
                cell = ws.cell(row, col)
                if isinstance(cell.value, bool) or not isinstance(cell.value, (float, int)) or not math.isfinite(cell.value):
                    raise ValueError(f"{sheet}:{cell.coordinate}: invalid numeric source")
                cells[name] = {"cell": cell.coordinate, "label_cell": ws.cell(row, 2).coordinate,
                               "label": label, "value": cell.value,
                               "number_format": cell.number_format}
            entries.append({
                "technology": technology, "projection_year": year, "source": source,
                "sheet": sheet, "year_cell": ws.cell(3, col).coordinate,
                "estimate_cell": ws.cell(4, col).coordinate,
                "estimate_label": ws.cell(4, col).value,
                "currency": "EUR", "currency_year": money_year,
                "price_basis": "real", "capacity_basis": capacity,
                "capacity_basis_code": capacity_code,
                "includes_grid_connection": grid_connection,
                "applicability": "DEA Danish representative technology benchmark; not country-calibrated EU vendor/site data",
                "included_costs": included, "unmapped_terms": gaps,
                "cells": cells,
            })
    for book in books.values():
        book.close()
    return {
        "schema_version": 1, "reference_id": "eu-dea-2026-08",
        "publisher": "Danish Energy Agency and Energinet",
        "retrieved_on": "2026-10-21", "source_release": "August 2026",
        "licence": "CC-BY-4.0", "licence_url": "https://creativecommons.org/licenses/by/4.0/",
        "licence_evidence": "Both inspected XLSX Intro sheets explicitly state Attribution 4.0 International",
        "attribution": "Danish Energy Agency / Energinet technology catalogues, August 2026; selected cached cells extracted and transformed by GridSpine. No endorsement is implied.",
        "modifications": "Selected central projection cells only; native reader converts units and derives absolute FOM/capex shares. No interpolation or inflation/FX conversion.",
        "status": "curated_benchmark_reference_not_project_defaults",
        "sources": SOURCES, "entries": entries,
    }


if __name__ == "__main__":
    if len(sys.argv) != 4:
        raise SystemExit(__doc__)
    output = extract(Path(sys.argv[1]), Path(sys.argv[2]))
    Path(sys.argv[3]).write_text(json.dumps(output, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Extracted {len(output['entries'])} explicit technology/year references")
